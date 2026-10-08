# -*- coding: utf-8 -*-
"""Build the Document Registration BRD from the Marriage BRD v8 template (styles, numbering, header / footer, TOC)
and the Document Registration process diagrams in ./ProcessDiagram.

Run:  python "_make_document_registration_brd.py"
"""
from __future__ import annotations

import copy
import html
import importlib
import re
import sys
import tempfile
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from numpy.lib.stride_tricks import sliding_window_view
from PIL import Image

import _brd_doc_content as C
import _brd_stamp_master as M

HERE = Path(__file__).resolve().parent
PLAN = HERE.parents[1]
sys.path.insert(0, str(PLAN / "Acts_Rules" / "Document"))
import _stamp_article_inputs_data as SD  # noqa: E402

TEMPLATE = PLAN / "Finalized BRD" / "Marriage" / "RFP" / "BRD_Marriage_BRD_Final_v8.docx"
OUT = HERE / f"BRD_Document_Registration_v{C.VERSION}.docx"
PD = HERE / "ProcessDiagram"

FONT = "Times New Roman"
TABLE_W, TABLE_IND = 8647, 1413
IMG_W_IN = 7.0
IMG_MAX_H_IN = 8.0
IMG_OUT_W = 3000
ABSTRACT = {3: 12, 4: 13, 5: 16}
UNNUMBERED_IND = {3: 1440, 4: 1800, 5: 2880}
H2_NUM = 7
BULLET_NUM = 26

sys.stdout.reconfigure(encoding="utf-8")


# ---------------------------------------------------------------------------------------------------------------------
# Diagram steps (read from the diagram generator scripts so the tables always match the diagrams)
# ---------------------------------------------------------------------------------------------------------------------
DIAGRAMS = [
    ("pre", "_make_citizen_preregistration_diagram", "citizen_pre_registration"),
    ("scr", "_make_scrutiny_diagram", "scrutiny"),
    ("pay", "_make_payment_esign_diagram", "payment_esign"),
    ("prs", "_make_presentation_registration_diagram", "presentation_registration"),
    ("drr", "_make_district_registrar_referral_diagram", "district_registrar_referral"),
    ("rfd", "_make_refund_process_diagram", "refund_process"),
]
_SEP, _START = "\x01", "\x00"


def _clean(s: str) -> str:
    s = re.sub(r"<br\s*/?>\s*•\s*", "; ", s)
    s = re.sub(r"<br\s*/?>", " ", s)
    s = re.sub(r"<div[^>]*>", " ", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s).replace("•", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s.replace(":;", ":").replace(": ;", ":").replace(" ;", ";")


def extract_steps() -> dict:
    sys.path.insert(0, str(PD))
    base = importlib.import_module("_make_citizen_preregistration_diagram")
    record: list = []
    orig_node, orig_vertex = base.Diagram.node, base.Diagram._vertex

    def node(self, cid, lane, row, kind, value, *a, **k):
        record.append(("node", lane, row, kind, value))
        return orig_node(self, cid, lane, row, kind, value, *a, **k)

    def vertex(self, cid, value, style, x, y, w, h):
        cid_s = str(cid)
        if cid_s.startswith("sec_"):
            record.append(("sec", None, float(cid_s[4:]), "sec", value))
        elif cid_s.startswith("lane_"):
            record.append(("lane", cid_s[5:], 0, "lane", value))
        return orig_vertex(self, cid, value, style, x, y, w, h)

    base.Diagram.node, base.Diagram._vertex = node, vertex
    out = {}
    try:
        for key, mod_name, fn in DIAGRAMS:
            mod = importlib.import_module(mod_name)
            saved = mod.label, mod.lane_title
            mod.label = lambda num, text, ref="": f"{_START}{num}{_SEP}{text}{_SEP}{ref}"
            mod.lane_title = lambda name, sub="": name
            record.clear()
            getattr(mod, fn)()
            mod.label, mod.lane_title = saved
            lanes = {r[1]: _clean(r[4]) for r in record if r[0] == "lane"}
            secs = sorted((r[2], _clean(r[4])) for r in record if r[0] == "sec")
            steps = []
            for kind_, lane, row, kind, value in record:
                if kind_ != "node" or not isinstance(value, str) or not value.startswith(_START):
                    continue
                num, text, ref = value[1:].split(_SEP)
                if not num:
                    continue
                sec = secs[0][1] if secs else ""
                for srow, stext in secs:
                    if srow <= row + 0.01:
                        sec = stext
                steps.append({"num": num, "text": _clean(text), "ref": _clean(ref), "lane": lanes.get(lane, lane),
                              "kind": kind, "section": sec})
            if key != "drr":
                steps.sort(key=lambda s: (int(re.match(r"\d+", s["num"]).group()), s["num"]))
            out[key] = {"stem": mod.STEM, "sections": [s for _, s in secs], "steps": steps}
    finally:
        base.Diagram.node, base.Diagram._vertex = orig_node, orig_vertex
    return out


# ---------------------------------------------------------------------------------------------------------------------
# Diagram slicing: page-sized parts with the lane header repeated on every part
# ---------------------------------------------------------------------------------------------------------------------
def slice_diagram(png: Path, out_dir: Path) -> list[Path]:
    Image.MAX_IMAGE_PIXELS = None
    im = Image.open(png).convert("RGB")
    a = np.asarray(im)
    H, W = a.shape[:2]
    ink = np.empty(H, dtype=np.int64)
    used_cols = np.zeros(W, dtype=bool)
    for y0 in range(0, H, 1024):
        blk = a[y0:y0 + 1024].astype(np.int16)
        ink[y0:y0 + 1024] = (np.abs(blk - 252).max(axis=2) > 6).sum(axis=1)
        used_cols |= (blk.min(axis=2) < 245).any(axis=0)
    right = min(W, int(np.nonzero(used_cols)[0].max()) + 12)
    bottom = min(H, int(np.nonzero(ink > 0)[0].max()) + 12)

    hdr_rows = np.nonzero(ink > 0.85 * ink.max())[0]
    h0 = h1 = int(hdr_rows[0])
    for r in hdr_rows[1:]:
        if r - h1 > 4:
            break
        h1 = int(r)
    hdr_top, hdr_bot = max(0, h0 - 1), h1 + 3
    hdr_img = im.crop((0, hdr_top, right, hdr_bot))
    hh = hdr_bot - hdr_top

    cap_first = int(right * IMG_MAX_H_IN / IMG_W_IN)
    cap_next = cap_first - hh
    k = 21
    bad = sliding_window_view(np.pad(ink, (k // 2, k // 2), mode="edge"), k).max(axis=1)

    def best_cut(lo, hi):
        seg = bad[lo:hi]
        cand = np.nonzero(seg <= seg.min() + 2)[0]
        return lo + int(cand[-1])

    n, y = 1, 0
    while bottom - y > (cap_first if n == 1 else cap_next):
        y = best_cut(y + int((cap_first if n == 1 else cap_next) * 0.55), y + (cap_first if n == 1 else cap_next))
        n += 1
    target = (bottom + (n - 1) * hh) / n

    cuts, y = [], 0
    for i in range(n - 1):
        cap = cap_first if i == 0 else cap_next
        ideal = int(y + target - (0 if i == 0 else hh))
        lo, hi = max(y + int(cap * 0.5), ideal - int(cap * 0.15)), min(y + cap, ideal + int(cap * 0.15))
        c = best_cut(lo, max(hi, lo + 1))
        cuts.append((y, c))
        y = c
    cuts.append((y, bottom))

    paths = []
    for i, (y0, y1) in enumerate(cuts):
        part = im.crop((0, y0, right, y1))
        if i > 0:
            canvas = Image.new("RGB", (right, hh + part.height), "white")
            canvas.paste(hdr_img, (0, 0))
            canvas.paste(part, (0, hh))
            part = canvas
        if part.width > IMG_OUT_W:
            part = part.resize((IMG_OUT_W, round(part.height * IMG_OUT_W / part.width)), Image.LANCZOS)
        part = part.quantize(colors=256, method=Image.Quantize.FASTOCTREE)
        p = out_dir / f"{png.stem}_part{i + 1}.png"
        part.save(p, optimize=True)
        paths.append(p)
    return paths


# ---------------------------------------------------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _rpr(bold=False, size=24):
    rpr = _el("w:rPr")
    fonts = _el("w:rFonts", ascii=FONT, hAnsi=FONT, cs=FONT)
    rpr.append(fonts)
    if bold:
        rpr.append(_el("w:b"))
    rpr.append(_el("w:sz", val=size))
    rpr.append(_el("w:szCs", val=size))
    return rpr


ID_RE = re.compile(r"(\b[A-Z]{2,4}(?:-[A-Z0-9]+)+\b)")


def _runs(parent, text, size=24, bold_all=False):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = bold_all or (part.startswith("**") and part.endswith("**"))
        seg = part[2:-2] if part.startswith("**") and part.endswith("**") else part
        r = _el("w:r")
        r.append(_rpr(bold, size))
        for tok in ID_RE.split(seg):
            if not tok:
                continue
            chunks = tok.split("-") if ID_RE.fullmatch(tok) else [tok]
            for i, chunk in enumerate(chunks):
                if i:
                    r.append(_el("w:noBreakHyphen"))
                t = _el("w:t")
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                t.text = chunk
                r.append(t)
        parent.append(r)


class Builder:
    def __init__(self, doc: Document):
        self.doc = doc
        body = doc.element.body
        self.sect = body[-1]
        self.numbering = doc.part.numbering_part.element
        self.cur = {3: None, 4: None, 5: None}
        self.table_style = doc.styles["Table Grid"].style_id
        self.figures = 0

    # -- numbering --------------------------------------------------------------------------------------------------
    def new_num(self, abstract_id: int) -> int:
        nums = self.numbering.findall(qn("w:num"))
        num_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        num = _el("w:num", numId=num_id)
        num.append(_el("w:abstractNumId", val=abstract_id))
        lo = _el("w:lvlOverride", ilvl=0)
        lo.append(_el("w:startOverride", val=1))
        num.append(lo)
        nums[-1].addnext(num)
        return num_id

    # -- paragraphs -------------------------------------------------------------------------------------------------
    def para(self, text="", style=None, ind_left=None, jc=None, num=None, spacing=None, bold_all=False):
        p = _el("w:p")
        ppr = _el("w:pPr")
        if style:
            ppr.append(_el("w:pStyle", val=style))
        if num is not None:
            npr = _el("w:numPr")
            npr.append(_el("w:ilvl", val=0))
            npr.append(_el("w:numId", val=num))
            ppr.append(npr)
        if spacing:
            ppr.append(_el("w:spacing", after=spacing[0], line=spacing[1], lineRule=spacing[2]))
        if ind_left is not None:
            ppr.append(_el("w:ind", left=ind_left))
        if jc:
            ppr.append(_el("w:jc", val=jc))
        ppr.append(_rpr())
        p.append(ppr)
        _runs(p, text, bold_all=bold_all)
        self.sect.addprevious(p)
        return p

    def body(self, text):
        return self.para(text, ind_left=1440, jc="both")

    def bullet(self, text):
        return self.para(text, style="ListBullet", num=BULLET_NUM)

    def heading(self, level: int, text: str, numbered: bool = True):
        num, ind = None, None
        if level == 2:
            num = H2_NUM if numbered else None
            self.cur = {3: None, 4: None, 5: None}
        elif numbered:
            if self.cur[level] is None:
                self.cur[level] = self.new_num(ABSTRACT[level])
            num = self.cur[level]
            for deeper in range(level + 1, 6):
                self.cur[deeper] = None
        else:
            ind = UNNUMBERED_IND[level]
        return self.para(text, style=f"Heading{level}", num=num, ind_left=ind)

    # -- tables -----------------------------------------------------------------------------------------------------
    def table(self, headers, rows, weights=None, size=None):
        n = len(headers)
        weights = weights or [1] * n
        tot = sum(weights)
        widths = [int(TABLE_W * w / tot) for w in weights]
        widths[-1] += TABLE_W - sum(widths)
        size = size or (24 if n <= 4 else 20)
        tbl = _el("w:tbl")
        tpr = _el("w:tblPr")
        tpr.append(_el("w:tblStyle", val=self.table_style))
        tpr.append(_el("w:tblW", w=TABLE_W, type="dxa"))
        tpr.append(_el("w:tblInd", w=TABLE_IND, type="dxa"))
        tpr.append(_el("w:tblLayout", type="fixed"))
        tpr.append(_el("w:tblLook", val="04A0", firstRow=1, lastRow=0, firstColumn=1, lastColumn=0, noHBand=0, noVBand=1))
        tbl.append(tpr)
        grid = _el("w:tblGrid")
        for w in widths:
            grid.append(_el("w:gridCol", w=w))
        tbl.append(grid)
        for ri, row in enumerate([headers] + [list(r) for r in rows]):
            tr = _el("w:tr")
            if ri == 0:
                trpr = _el("w:trPr")
                trpr.append(_el("w:tblHeader"))
                tr.append(trpr)
            for ci in range(n):
                tc = _el("w:tc")
                tcpr = _el("w:tcPr")
                tcpr.append(_el("w:tcW", w=widths[ci], type="dxa"))
                tc.append(tcpr)
                p = _el("w:p")
                ppr = _el("w:pPr")
                ppr.append(_el("w:spacing", after=0))
                ppr.append(_rpr(ri == 0, size))
                p.append(ppr)
                val = row[ci] if ci < len(row) else ""
                _runs(p, "" if val is None else str(val), size=size, bold_all=(ri == 0))
                tc.append(p)
                tr.append(tc)
            tbl.append(tr)
        self.sect.addprevious(tbl)
        spacer = self.para("")
        return tbl, spacer

    # -- images -----------------------------------------------------------------------------------------------------
    def image(self, path: Path, caption: str):
        p = self.doc.add_paragraph(style="Normal (Web)")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(IMG_W_IN))
        self.figures += 1
        self.para(f"Figure: {caption}", jc="center")


def set_cell_text(cell, text):
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    p = paras[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)


def set_para_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def prepare_template(doc: Document):
    body = doc.element.body
    kids = list(body)
    for el in kids[37:-1]:
        body.remove(el)

    set_para_text(doc.paragraphs[1], C.MODULE)
    cover = doc.tables[0]
    for (field, value), row in zip(C.COVER, cover.rows[1:]):
        set_cell_text(row.cells[0], field)
        set_cell_text(row.cells[1], value)
    ver = doc.tables[1]
    for row in list(ver.rows)[2:]:
        row._tr.getparent().remove(row._tr)
    for i, v in enumerate(C.VERSIONS):
        if i > 0:
            ver._tbl.append(copy.deepcopy(ver.rows[1]._tr))
        for cell, val in zip(ver.rows[i + 1].cells, v):
            set_cell_text(cell, val)

    sdt = next(el for el in body if el.tag == qn("w:sdt"))
    content = sdt.find(qn("w:sdtContent"))
    paras = content.findall(qn("w:p"))
    for p in paras[1:]:
        content.remove(p)
    for child in list(content):
        if child.tag != qn("w:p"):
            content.remove(child)
    fp = _el("w:p")
    for kind, extra in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '), ("separate", None),
                        ("text", "Right-click and choose Update Field to refresh the table of contents."), ("end", None)):
        r = _el("w:r")
        if kind == "instr":
            t = _el("w:instrText")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = extra
            r.append(t)
        elif kind == "text":
            t = _el("w:t")
            t.text = extra
            r.append(t)
        else:
            fc = _el("w:fldChar", fldCharType=kind)
            if kind == "begin":
                fc.set(qn("w:dirty"), "true")
            r.append(fc)
        fp.append(r)
    content.append(fp)

    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
            for t in rel.target_part.element.iter(qn("w:t")):
                if t.text and "Marriage Registration Module" in t.text:
                    t.text = t.text.replace("Marriage Registration Module", C.MODULE)

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(_el("w:updateFields", val="true"))
    cp = doc.core_properties
    cp.title = "Business Requirements Document — Document Registration Module"
    cp.subject = "Kaveri 3.0 — Document Registration"
    cp.author = "Nandha Kumar"


def drop_unused_images(doc: Document):
    xml = doc.element.xml
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image") and f'"{rid}"' not in xml:
            doc.part.drop_rel(rid)


def section_title(text: str) -> str:
    m = re.match(r"^([A-Z](?:[–-][A-Z])?\.)\s*(.*)$", text)
    prefix, rest = (m.group(1), m.group(2)) if m else ("", text)
    words = [w.lower() if (w.isupper() and len(w) > 1 and not re.search(r"\d", w)) else w for w in rest.split(" ")]
    rest = " ".join(words)
    rest = rest[:1].upper() + rest[1:]
    for a, b in (("district registrar", "District Registrar"), ("sub-registrar", "Sub-Registrar"), ("e-sign", "e-Sign")):
        rest = re.sub(a, b, rest, flags=re.I)
    return f"{prefix} {rest}".strip()


def step_rows(steps):
    rows = []
    for s in steps:
        note = s["ref"] or ("Decision" if s["kind"] == "decision" else "")
        rows.append((s["num"], s["text"], s["lane"], note))
    return rows


# ---------------------------------------------------------------------------------------------------------------------
# Stamp duty inputs (from Acts_Rules/Document/_stamp_article_inputs_data.py, the source of the input workbook)
# ---------------------------------------------------------------------------------------------------------------------
PURPOSE = {"I": "Identify", "C": "Calculate", "E": "Exemption"}
SHEET_REFS = (("See Lookup_Lists: ", "See 3.v list: "), ("see Jurisdiction_Slab", "see 3.v.d"),
              ("Family_Definitions", "family definitions table"), ("Guidance Value module", "Guidance Value service"))


def brd_text(s: str) -> str:
    for a, b in SHEET_REFS:
        s = s.replace(a, b)
    return s


def build_stamp_inputs(b: Builder):
    b.heading(4, "Other lookup lists")
    b.body(C.LOOKUP_INTRO)
    gen05 = next(i for i in SD.INPUTS if i[0] == "GEN-05")
    b.body("Party category (GEN-05 executant, GEN-06 claimant):")
    b.table(["#", "Party category"], [(k + 1, c.strip()) for k, c in enumerate(gen05[5].split(","))], [0.6, 6])
    b.body("Relationship of claimant to executant (GEN-07) and family definitions:")
    heads = ["Relationship"] + [f"{a} {d}" for a, d in SD.FAMILY_MATRIX_ARTICLES]
    b.table(heads, [[rel] + SD.FAMILY_MATRIX[rel] for rel in SD.RELATIONSHIPS], [2.6, 1.3, 1.3, 1.3, 1.3, 1.5])
    for t in C.FAMILY_NOTES:
        b.bullet(t)
    b.body("Local body limits (PRP-03) and family-rate fixed duty (₹):")
    names = dict(SD.LOCAL_LIMITS)
    b.table(["Code", "Local body limits"] + SD.JURISDICTION_SLAB_COLS,
            [[c, names[c]] + [f"{v:,}" for v in vals] for c, *vals in SD.JURISDICTION_SLAB], [0.8, 3.2, 1.2, 1.3, 1.2, 1.4])
    b.body(C.SLAB_NOTE)
    b.body("Allotting authority (AGR-04, lease-cum-sale under Article 5(d) / 5(da)):")
    b.table(["#", "Authority"], [(k + 1, a) for k, a in enumerate(SD.ALLOTTING_AUTHORITIES)], [0.6, 6])
    b.body("Form of security (LON-03):")
    forms = ["Deposit of title deeds", "Pawn / pledge of movable property"] + SD.SECURITY_FORMS + \
        ["Mortgage of crop", "Further charge"]
    b.table(["#", "Form of security"], [(k + 1, f) for k, f in enumerate(forms)], [0.6, 6])

    b.heading(4, "Input catalogue")
    b.body(C.INPUT_CATALOGUE_INTRO)
    groups: dict[str, list] = {}
    for i in SD.INPUTS:
        groups.setdefault(i[1], []).append(i)
    for g, items in groups.items():
        b.heading(5, f"{g} ({items[0][0].split('-')[0]}, {len(items)} fields)")
        rows = [(iid, brd_text(f"**{label}**: {q}"), brd_text(f"{dtype}; {allowed}"), brd_text(src), PURPOSE.get(p, p))
                for iid, _g, label, q, dtype, allowed, src, p in items]
        b.table(["ID", "Field and question to user", "Data type and allowed values", "Source", "Purpose"], rows,
                [1.2, 3.8, 2.6, 1.5, 1.5])


def article_rule_rows():
    rows = []
    for r in SD.RULES:
        inputs = []
        for tag, ids in (("I", r["id_in"]), ("C", r["calc_in"]), ("E", r["adj_in"])):
            if ids:
                inputs.append(f"{tag}: " + ", ".join(ids))
        duty = r["duty"]
        caps = [f"Min {r['min']}" if r["min"] else "", f"Max {r['max']}" if r["max"] else ""]
        caps = "; ".join(c for c in caps if c)
        if caps:
            duty += f" ({caps})"
        if r["s45a"] == "Yes":
            duty += "; Sec. 45-A market value check"
        notes = "; ".join(x for x in (r["proviso"], f"Exempt: {r['exempt']}" if r["exempt"] else "") if x)
        rows.append((r["key"], brd_text(f"**{r['instrument']}** — {r['cond']}"), "; ".join(inputs) or "—", brd_text(duty),
                     brd_text(notes)))
    return rows


# ---------------------------------------------------------------------------------------------------------------------
def build():
    steps = extract_steps()
    doc = Document(str(TEMPLATE))
    prepare_template(doc)
    b = Builder(doc)
    STEP_W = [0.75, 4.55, 2.0, 2.7]

    b.heading(2, "Executive summary")
    for t in C.EXEC_SUMMARY:
        b.para(t, style="ListParagraph", spacing=(0, 300, "atLeast"), jc="both")

    b.heading(2, "Scope")
    for t in C.SCOPE:
        b.bullet(t)

    # 3. Legal ---------------------------------------------------------------------------------------------------------
    b.heading(2, "Legal and regulatory reference")
    b.heading(3, "Applicable Acts")
    b.body(C.ACTS_INTRO)
    b.table(["Act", "Description"], C.ACTS, [2.2, 5])
    b.heading(3, "Relevant sections followed by the Department for Document Registration")
    b.heading(4, "Registration Act, 1908 (selected sections, with Karnataka amendments)")
    b.table(["Section", "Topic", "Relevance", "Refer section 7 / 8 for Implementation"], C.SECTIONS_REG_ACT, [1.5, 3, 2.6, 2])
    b.heading(4, "Karnataka Stamp Act, 1957 (selected sections)")
    b.table(["Section", "Topic", "Relevance", "Refer section 7 / 8 for Implementation"], C.SECTIONS_STAMP_ACT, [1.4, 3, 2.6, 2])
    b.heading(3, "Relevant rules followed by the Department for Document Registration")
    b.heading(4, "Karnataka Registration Rules, 1965")
    b.table(["Rule", "Requirement", "Refer section 7 for Implementation"], C.RULES_REG, [1.8, 4, 2.4])
    b.heading(4, "Karnataka Stamp Rules, e-Stamping Rules, Undervaluation Rules and Franking Rules")
    b.table(["Rule", "Requirement", "Refer section 7 for Implementation"], C.RULES_STAMP, [2.4, 4, 2])
    b.heading(3, "Relevant notifications and amendments")
    b.table(["Instrument", "Date / No.", "Effect", "Refer section 7 for Implementation"], C.NOTIFICATIONS, [2.2, 1.8, 3.6, 2])
    b.heading(3, "Stamp duty Article determination")
    b.body(C.ARTICLE_INTRO)
    b.heading(4, "Transaction intents")
    b.table(["Code", "Transaction Intent"], M.INTENTS, [1, 5])
    b.heading(4, "Nature of document and Article")
    b.table(["Code", "Nature of Document", "Intent", "Article(s)"], M.NATURES, [1, 4.4, 1.1, 1.2])
    b.heading(4, "Sub-classification masters")
    b.body("Agreement subject (Article 5):")
    b.table(["Code", "Agreement relates to", "Article"], M.AGREEMENT_SUBJECTS, [1, 5, 1.2])
    b.body("Power of Attorney purpose (Article 41):")
    b.table(["Code", "Purpose", "Article"], M.POA_PURPOSES, [1, 5, 1.4])
    b.body("Conveyance type (Article 20):")
    b.table(["Code", "Conveyance type", "Article"], M.CONVEYANCE_TYPES, [1, 5, 1.2])
    build_stamp_inputs(b)
    b.heading(4, "General stamp duty rules")
    b.table(["Rule", "Section / Article", "Topic", "Legal provision (summary)", "System implementation"], M.GENERAL_RULES,
            [1.0, 1.2, 1.5, 3.2, 2.6])
    b.heading(3, "Statutory registers, books and forms")
    b.table(["Register / form", "Statutory ref", "Kaveri 3.0 treatment"], C.REGISTERS, [2, 2, 4.2])
    b.heading(3, "Sakala — Karnataka Guarantee of Services")
    b.body(C.SAKALA_TEXT)

    # 4–6 -------------------------------------------------------------------------------------------------------------
    b.heading(2, "Stakeholders and actors")
    b.table(["Actor", "Description", "Primary goals", "Channel involvement"], C.STAKEHOLDERS, [1.8, 3, 2.6, 1.7])
    b.heading(2, "Definitions and glossary")
    b.table(["Term", "Definition"], C.GLOSSARY, [2, 5])
    b.heading(2, "Current state")
    b.heading(3, "As-Is pain points")
    b.body(C.PAIN_INTRO)
    b.table(["Sr.No", "Pain Point", "Description", "Source", "Addressed in (this BRD)"],
            [(i + 1, p, d, C.PAIN_SOURCE, a) for i, (p, d, a) in enumerate(C.PAIN_POINTS)], [0.6, 1.8, 3.2, 1.6, 2])

    # 7. Future state ---------------------------------------------------------------------------------------------------
    b.heading(2, "Future state (To-Be)")
    b.body(C.TOBE_INTRO)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for key, _mod, _fn in DIAGRAMS:
            P, S = C.PROCESSES[key], steps[key]
            b.heading(3, P["title"])
            b.heading(4, "Channel model")
            b.table(["Service Type", "Online Activities", "Office Activities", "Mode"], P["channel"], [2, 3.3, 3.3, 1.3])
            b.heading(4, "Process diagram")
            b.body(f"Source: ProcessDiagram/{S['stem']}.drawio (full-resolution PNG alongside).")
            parts = slice_diagram(PD / f"{S['stem']}.png", tmp)
            for i, path in enumerate(parts):
                suffix = f" (part {i + 1} of {len(parts)})" if len(parts) > 1 else ""
                b.image(path, f"{P['figure']}{suffix}")
            print(f"{key}: {len(parts)} diagram parts, {len(S['steps'])} steps")
            b.heading(4, "Process flow")
            b.body(P["flow_intro"])
            if S["sections"]:
                groups: dict[str, list] = {}
                for s in S["steps"]:
                    groups.setdefault(s["section"], []).append(s)
                for sec in S["sections"]:
                    rows = step_rows(groups.get(sec, []))
                    if sec == S["sections"][-1]:
                        rows += P.get("extra_rows", [])
                    if not rows:
                        continue
                    b.heading(5, section_title(sec), numbered=False)
                    b.table(["#", "Step", "Lane", "Notes"], rows, STEP_W)
            else:
                b.table(["#", "Step", "Lane", "Notes"], step_rows(S["steps"]), STEP_W)
            b.body(P["key"])
            b.heading(4, "Application status model")
            b.body(f"Application statuses used in the {P['title'].lower()} process:")
            b.table(["Status", "Description", "Actor", "Next states"],
                    [(s,) + C.STATUS_MODEL[s] for s in P["statuses"]], [2, 3.4, 1.6, 2.6])

        b.heading(3, "Application status model (end-to-end)")
        b.body("Consolidated status model across the six processes. Every change of status is audited (NFR-DOC-AUD-001) "
               "and notified to the applicant (FR-DOC-NTF-001).")
        b.table(["Status", "Description", "Actor", "Next states"], [(s,) + v for s, v in C.STATUS_MODEL.items()],
                [2, 3.4, 1.6, 2.6])
        b.heading(3, "What is new in Kaveri 3.0")
        b.table(["#", "Capability", "What is new in Kaveri 3.0"],
                [(i + 1, c, w) for i, (c, w) in enumerate(C.WHATS_NEW)], [0.5, 2.3, 5])
        b.heading(4, "Rectified As-Is pain points")
        b.table(["Sr.No", "Pain Point (As-Is)", "How rectified in Kaveri 3.0"],
                [(i + 1, p, h) for i, (p, h) in enumerate(C.RECTIFIED)], [0.6, 2.6, 4.6])

        # 8. Functional requirements ----------------------------------------------------------------------------------
        b.heading(2, "Functional requirements")
        fr_headers = ["Req ID", "Requirement", "Priority", "Acceptance criteria"]
        fr_w = [2.6, 3.3, 1.2, 2.2]
        for title, intro, groups in C.FR:
            b.heading(3, title)
            if intro:
                b.body(intro)
            if len(groups) == 1:
                b.table(fr_headers, groups[0][1], fr_w)
            else:
                for gtitle, rows in groups:
                    b.heading(4, gtitle)
                    b.table(fr_headers, rows, fr_w)
        b.heading(3, "Business rules")
        b.table(["Rule ID", "Description", "Statutory ref", "System enforcement"], C.BUSINESS_RULES, [1.6, 3.0, 1.7, 2.3])
        b.heading(3, "User interface (high-level)")
        b.table(["Screen / step", "Purpose", "Channel", "Statutory alignment", "Notes"], C.UI_SCREENS, [1.8, 2.5, 1.0, 1.6, 1.9])
        b.body("Wireframe links: [Figma / prototype URLs]")
        b.body("Bilingual: all labels in English and Kannada — content manager sign-off.")
        b.heading(3, "Integrations")
        b.table(["Integration", "Direction", "Purpose", "Channel", "Owner", "Status"], C.INTEGRATIONS, [2.2, 1.1, 2.6, 1.1, 1.6, 1.1])
        b.body("Interface requirements: API specifications to be finalised by the Architect for each integration above, "
               "including the Guidance Value, Stays & Liabilities and Fee Calculation microservice contracts.")
        b.heading(3, "Data requirements")
        b.heading(4, "Core entities (logical)")
        for t in C.DATA_ENTITIES:
            b.bullet(t)
        b.heading(4, "Retention")
        b.body(C.RETENTION)
        b.heading(4, "Migration (high level)")
        b.table(["Topic"], C.MIGRATION)
        b.heading(3, "Requirements traceability matrix (RTM) - template")
        b.table(["Req ID", "Act / Rule / Form", "Requirement summary", "Section", "UI screen", "Test case ID", "Status"],
                [r + ("TBD", "Draft") for r in C.RTM], [2.1, 1.5, 2.0, 0.9, 1.4, 1.0, 1.1])

    # 9–14 --------------------------------------------------------------------------------------------------------------
    b.heading(2, "Non-functional requirements")
    b.body(C.NFR_INTRO)
    for title, intro, rows in C.NFR:
        b.heading(3, title)
        b.body(intro)
        b.table(fr_headers, rows, fr_w)
    b.heading(2, "Risk and Mitigation Strategy")
    b.body(C.RISK_INTRO)
    b.table(["Risk ID", "Risk", "Mitigation", "Related requirements"], C.RISKS, [1.9, 2.2, 2.6, 2.6])
    b.heading(2, "System Fallbacks & Error Handling")
    b.body(C.FALLBACK_INTRO)
    b.table(fr_headers, C.FALLBACKS, fr_w)
    b.heading(2, "Training and Change Management")
    b.body(C.TRAINING["intro"])
    b.heading(3, "Target audience")
    b.body(C.TRAINING["audience"])
    b.heading(3, "Training delivery")
    for t in C.TRAINING["delivery"]:
        b.bullet(t)
    b.heading(3, "Citizen change management")
    b.body(C.TRAINING["citizen"])
    b.heading(3, "Post-Go-Live support")
    b.body(C.TRAINING["support"])
    b.heading(2, "Appendix A — References")
    for t in C.REFERENCES:
        b.bullet(t)
    b.heading(2, "Appendix B — Article rules and stamp duty inputs")
    b.body(C.APPENDIX_B_INTRO)
    b.table(["Ref", "Instrument and identification condition", "Inputs", "Proper stamp duty", "Provisos / exemptions"],
            article_rule_rows(), [1.3, 3.2, 1.9, 2.6, 2.4])
    b.heading(2, "Acceptance and sign-off of BRD")
    b.table(["Role", "Name", "Signature / Date"], C.SIGNOFF, [2, 3, 3])

    drop_unused_images(doc)
    doc.save(str(OUT))
    print("Wrote", OUT, f"({OUT.stat().st_size / 1e6:.1f} MB, {b.figures} figures)")


if __name__ == "__main__":
    build()
