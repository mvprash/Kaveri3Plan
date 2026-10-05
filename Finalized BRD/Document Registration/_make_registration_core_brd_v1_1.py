# -*- coding: utf-8 -*-
"""Create BRD_Registration_Core_v1.1.docx — Registration core: Registration,
Appointment, Status tracking (27-08-2026 discussion; Schedule Sr.12, sub-modules
#1, #2, #19).

Section structure mirrors Finalized BRD/Marriage/RFP/BRD_Marriage_BRD_v8.docx,
which is also used as the style template (fonts, heading styles, header).
Process diagrams (§7.2) follow the CVC Guidance Value Fixation BRD v1.9 style;
regenerate them with ProcessDiagram/_make_registration_core_process_diagrams.py
and ProcessDiagram/_export_registration_core_png.py before running this script.

Sources:
  Requirement Discussions/Daily Reports/Consolidated_Requirement_Discussions_25082026_to_11092026.docx (§2)
  Requirement Discussions/Daily Reports/Document_Registration_requirement_27082026_v1.1.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
BASE = ROOT / "Finalized BRD" / "Document Registration"
TEMPLATE = ROOT / "Finalized BRD" / "Marriage" / "RFP" / "BRD_Marriage_BRD_v8.docx"
DIAGRAM_DIR = BASE / "ProcessDiagram"
DST = BASE / "BRD_Registration_Core_v1.1.docx"

MODULE = "Registration Core — Registration, Appointment and Status Tracking"
VERSION = "1.1"
DATE = "05-10-2026"

# Column widths in inches; portrait usable width is 7.1".
COLUMN_WIDTHS: dict[tuple[str, ...], list[float]] = {
    ("#", "Step", "Lane", "Notes"): [0.4, 3.0, 1.4, 2.3],
    ("Sr.No", "Pain Point", "Description", "Source", "Addressed in (this BRD)"): [0.5, 1.4, 2.2, 1.6, 1.4],
    ("Status", "Description", "Actor", "Next states"): [1.5, 2.6, 1.2, 1.8],
    ("Rule ID", "Description", "Statutory ref", "System enforcement"): [1.0, 3.0, 1.5, 1.6],
    ("Risk ID", "Risk", "Mitigation", "Related requirements"): [1.0, 2.1, 2.5, 1.5],
    ("Req ID", "Act/Rule/Form", "Requirement summary", "BRD section", "UI screen", "Test case ID", "Status"): [1.0, 1.2, 1.7, 0.6, 1.1, 0.8, 0.7],
    ("ID", "User story", "Requirements"): [0.9, 4.7, 1.5],
    ("#", "Capability", "What is new in Kaveri 3.0"): [0.4, 1.9, 4.8],
    ("Sr.No", "Pain Point (As-Is)", "How rectified in Kaveri 3.0"): [0.5, 2.4, 4.2],
    ("Section", "Topic", "Relevance", "Refer section 7 / 8 for implementation"): [1.1, 2.3, 2.1, 1.6],
    ("Field", "Value"): [2.2, 4.9],
    ("Term", "Definition"): [2.0, 5.1],
}


# ---------------------------------------------------------------------------
# Template + writing helpers
# ---------------------------------------------------------------------------
def load_template() -> Document:
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    for rid, rel in list(doc.part.rels.items()):
        if "image" in rel.reltype:
            doc.part.drop_rel(rid)
    hdr = doc.sections[0].header.paragraphs[0]
    text = f"Kaveri 3.0 | BRD | {MODULE}"
    if hdr.runs:
        hdr.runs[0].text = text
        for r in hdr.runs[1:]:
            r.text = ""
    else:
        hdr.add_run(text)
    return doc


def shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    lines = text.split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(9)


class W:
    def __init__(self, doc: Document) -> None:
        self.doc = doc

    def h(self, text: str, level: int) -> None:
        self.doc.add_heading(text, level=level)

    def p(self, text: str, bold_prefix: str | None = None, italic: bool = False) -> None:
        para = self.doc.add_paragraph()
        if bold_prefix:
            para.add_run(bold_prefix).bold = True
        run = para.add_run(text)
        run.italic = italic

    def bullets(self, items: list[str]) -> None:
        for it in items:
            self.doc.add_paragraph(it, style="List Bullet")

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        t.style = "Table Grid"
        for i, hd in enumerate(headers):
            cell_text(t.rows[0].cells[i], hd, bold=True)
            shade(t.rows[0].cells[i], "D9E2F3")
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                cell_text(t.rows[ri].cells[ci], val)
        widths = widths or COLUMN_WIDTHS.get(tuple(headers))
        if widths:
            t.autofit = False
            for row in t.rows:
                for ci, wd in enumerate(widths):
                    row.cells[ci].width = Inches(wd)
        self.doc.add_paragraph()

    def fr(self, rows: list[list[str]]) -> None:
        self.table(["Req ID", "Requirement", "Priority", "Acceptance criteria"], rows, [1.25, 3.1, 0.7, 2.05])

    def caption(self, text: str) -> None:
        para = self.doc.add_paragraph(text, style="Caption")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------
def front_matter(w: W) -> None:
    w.doc.add_paragraph("Business Requirements Document (BRD)", style="Title")
    w.h(MODULE, 1)

    w.h("Document control", 2)
    w.table(
        ["Field", "Value"],
        [
            ["Document ID", "BRD-K3-DOC-REG-001"],
            ["Version", VERSION],
            ["Status", "Draft — for department review"],
            ["Module", "Document Registration — Registration core (Schedule Sr.12)"],
            ["Sub-modules in scope", "#1 Registration; #2 Appointment; #19 Status tracking"],
            ["Discussion date", "27-08-2026 (session held 26-08-2026 to 27-08-2026) — current issues and process walkthrough"],
            ["Attendees", "Kaveri IT Cell, AIGR (Computerisation), Project Manager, Assistant Manager (CMS)"],
            ["Legal basis (primary)", "The Registration Act, 1908; The Registration (Karnataka Amendment) Act, 2023; The Karnataka Stamp Act, 1957; The Transfer of Property Act, 1882 (Sec. 54)"],
            ["State rules (primary)", "The Karnataka Registration Rules, 1965; Karnataka Stamp Rules, 1958; Karnataka Stamp (Payment of Duty by Means of e-Stamping) Rules, 2009"],
            ["Author (BA)", "Nandha Kumar"],
            ["Product Owner", "M V Prashanth"],
            ["Domain expert / SRO reviewer", "Prabhakar Naik, Committee members"],
            ["Target audience", "Kaveri IT Cell, Department of Stamps and Registration, Government of Karnataka"],
            ["Last updated", DATE],
        ],
    )
    w.p("Version history:")
    w.table(
        ["Version", "Date", "Author", "Summary of change", "Approver"],
        [
            ["1.0", DATE, "Nandha Kumar", "Initial BRD for Registration core (Registration, Appointment, Status tracking) from the 27-08-2026 discussion", "M V Prashanth"],
            [VERSION, DATE, "Nandha Kumar", "§7.2 process diagrams redrawn in the CVC BRD style and plain language — Process A (application), B (SR check and payment), C (appointment and registration) with step tables", "M V Prashanth"],
        ],
    )
    w.p("Related documents:")
    w.table(
        ["ID", "Title", "Location"],
        [
            ["BRD-K3-DOC-REG-001", "This document", "Finalized BRD/Document Registration/"],
            ["BRD-K3-DOC-001", "BRD_Document_Registration_v1.1 — Legal and regulatory reference for Schedule Sr.12–15", "Finalized BRD/Document Registration/"],
            ["REQ-27082026", "Document_Registration_requirement_27082026_v1.1.docx", "Requirement Discussions/Daily Reports/"],
            ["REQ-CONS-01", "Consolidated_Requirement_Discussions_25082026_to_11092026.docx — §2 (27-08-2026)", "Requirement Discussions/Daily Reports/"],
            ["PD-REG-01", "Process_A_Registration_Application; Process_B_SR_Check_and_Payment; Process_C_Appointment_and_Registration (.drawio / .png)", "Finalized BRD/Document Registration/ProcessDiagram/"],
            ["PD-REG-02 / 03", "Registration_Appeal_Part_XII; Registration_22BCD_Forged_Document (exception tracks)", "Finalized BRD/Document Registration/ProcessDiagram/"],
            ["SD-ISSUES", "ServiceDeskIssuesList.xlsx (OverallList + Categorized)", "Requirement Discussions/ServiceDesk Issues/"],
        ],
    )


def executive_summary(w: W) -> None:
    w.h("1. Executive summary", 2)
    w.p(
        "This document assesses the existing Kaveri 2.0 document registration intake, appointment "
        "and status-tracking process, as walked through with the department on 27-08-2026, and the "
        "ServiceDesk issues logged against it. Citizens currently choose an article from opaque "
        "dropdowns, re-enter information the system could derive, see the base guidance rate, and "
        "cannot reliably generate a summary or book an appointment after payment; department users "
        "see applications stuck at intermediate steps."
    )
    w.p(
        "Based on that analysis, the document proposes a future-state Registration core for Kaveri 3.0: "
        "a requirement-based article selection, KSRSAC-driven village and Index-II details, a visual "
        "property map, transaction-aware party labels, a simplified valuation screen that hides the base "
        "rate, a mandatory pre-submission preview, an appointment step that opens immediately after "
        "payment, and a single, explicit application status model shared by citizens and offices."
    )
    w.p(
        "The approach is intended to reduce wrong-article and wrong-duty registrations, eliminate the "
        "recurring save / summary / schedule failures seen in ServiceDesk, and provide a scalable "
        "foundation for the related Document Registration BRDs (stamp duty and fee calculation, CVC / "
        "GIS valuation, refusal and appeal, Rule 17 filing)."
    )


def scope(w: W) -> None:
    w.h("2. Scope", 2)
    w.h("2.1 In scope", 3)
    w.bullets([
        "Document classification — requirement-based article / sub-article selection and instrument mapping.",
        "Village index selection — KSRSAC integration for Index-II village details; geofencing feasibility.",
        "Property details — structured boundaries, removal of East-West / North-South fields for agricultural land, visual property map.",
        "Party information — transaction-based party labels, service-specific Aadhaar OTP templates, gender auto-fill from salutation, conditional schedule allocation.",
        "Property valuation (intake) — selection of physical characteristics only; base rate hidden from citizens.",
        "Review and submission — dynamic consideration caption and mandatory pre-submission preview.",
        "Sub-Registrar verification and send-for-payment; payment visibility on summary.",
        "Appointment — slot booking after payment, rescheduling, check-in at SRO and presentation queue.",
        "Presentation, examination and registration completion on the happy path (Secs. 32A–35, 52, 58–61).",
        "Status tracking — one status model across citizen portal, SRO and District Registrar views; ageing and stuck-application controls.",
        "Notifications, MIS, audit trail and bilingual UI (English + Kannada) for the above.",
    ])
    w.h("2.2 Out of scope (covered by other BRDs)", 3)
    w.bullets([
        "Stamp duty and registration fee computation rules and guideline value calculation (Schedule Sr.13; 28-08-2026 discussion) — this BRD consumes their results.",
        "Valuation Module (CVC) and GIS valuation (Sr.14).",
        "Sec. 45-A undervaluation reference and Stamp Act impounding case-work — only the resulting statuses are defined here.",
        "Refusal and appeal workflows under Part XII (Secs. 71–77) and Karnataka Secs. 22-B / 22-C / 22-D — only the entry status is defined here (see PD-REG-02 / 03).",
        "Rule 17(2) / 17(3) filing, old pending release, scanning, FRUITS, memo transmission and other Document sub-modules (Sr.15 onwards).",
        "Private attendance / registration at residence (Sec. 31).",
    ])
    w.h("2.3 Assumptions", 3)
    w.bullets([
        "Sub-Registrar and District Registrar offices operate under the Registration Act, 1908 (as amended in Karnataka) and the Karnataka Registration Rules, 1965.",
        "Stamp duty and registration fee amounts are supplied by the Sr.13 calculation engine; guideline rates by the CVC master (Sr.14).",
        "KSRSAC exposes village boundary and Index-II village data through an approved API; Bhoomi exposes RTC / survey data for agricultural land.",
        "Citizen login, roles and office postings are provided by the User Management module.",
    ])
    w.h("2.4 Constraints", 3)
    w.bullets([
        "Documents falling under Sec. 22-B (forged / prohibited / attached property) must not be registered.",
        "Presentation time limits (Secs. 23, 25) and fine provisions (Rules 51–55) must be surfaced before an appointment is confirmed.",
        "Appointment calendars must respect office hours and holidays (Rules 3, 5).",
        "Aadhaar authentication must follow UIDAI regulations and department approval; SMS templates must be DLT-registered.",
        "Hiding the base rate from citizens must not change the statutory value used for duty; the Sub-Registrar must still see it.",
    ])


def legal(w: W) -> None:
    w.h("3. Legal and regulatory reference", 2)
    w.h("3.1 Applicable Acts", 3)
    w.p(
        "The Department of Stamps and Registration (IGR, District Registrars and Sub-Registrars) "
        "administers the following Acts for the Registration core:"
    )
    w.table(
        ["Act", "Act No.", "Relevance to Registration core"],
        [
            ["The Registration Act, 1908", "Central Act 16 of 1908", "Compulsory / optional registration; presentation time and place; examination; endorsements and Sec. 60 certificate; refusal / appeal; registration fees — core Registration and Status tracking"],
            ["The Registration Act, 1908 — office structure", "Central Act 16 of 1908", "Districts / sub-districts; Registrars and Sub-Registrars; place of registration — Appointment office selection and jurisdiction routing"],
            ["The Registration (Karnataka Amendment) Act, 2023", "Karnataka Act 47 of 2024", "Refusal / cancellation of forged or prohibited documents; appeal; penalties — registration intake validation"],
            ["The Karnataka Stamp Act, 1957", "Karnataka Act 34 of 1957", "Stamp duty chargeable; payment modes; impounding; undervaluation reference — consumed from the 28-08 discussion"],
            ["The Transfer of Property Act, 1882", "Central Act 4 of 1882", "Sec. 54 — sale of immovable property of value above ₹100 requires a registered instrument; drives article / transaction-type selection"],
        ],
    )

    w.h("3.2 Relevant sections followed by the Department for Document Registration", 3)
    w.h("3.2.1 The Registration Act, 1908", 4)
    w.table(
        ["Section", "Topic", "Relevance", "Refer section 7 / 8 for implementation"],
        [
            ["Secs. 5–7", "Districts and sub-districts; Registrars and Sub-Registrars; offices", "SRO / district master for jurisdiction and appointment calendar", "7.2.1 step 7; 7.2.3; FR-REG-012; FR-APT-002"],
            ["Secs. 10–12", "Absence / vacancy of Registrar or Sub-Registrar", "Charge arrangement; SRO-initiated reschedule", "FR-APT-006"],
            ["Secs. 17–18", "Documents of which registration is compulsory / optional", "Article selection; compulsory-registration guidance", "7.2.1 steps 2–4; FR-REG-001, FR-REG-004"],
            ["Secs. 19–22", "Language; interlineations; description of property and maps; reference to government maps / surveys", "Structured property description, boundaries and map", "7.2.1 step 8; FR-REG-020–025"],
            ["Secs. 23, 25", "Time for presentation (4 months); delay with fine", "Appointment window check and fine warning", "FR-APT-003; BR-REG-007"],
            ["Secs. 28–30", "Place for registering documents relating to land / other documents", "Jurisdictional SRO derived from property location", "FR-REG-012; BR-REG-004"],
            ["Sec. 31", "Registration at private residence", "Out of scope (2.2)", "—"],
            ["Secs. 32–33, 32A", "Persons to present; powers of attorney; photograph and fingerprints", "Presentant / representative capture; biometrics at SRO", "FR-REG-035; FR-REG-081"],
            ["Secs. 34–35", "Enquiry before registration; admission / denial of execution", "Examination checklist; admission status", "7.2.3 steps 7–9; FR-REG-082"],
            ["Secs. 51–52", "Register books; duties on presentation (endorse time, presentant, receipt)", "Presentation time recorded at call / check-in", "FR-REG-080"],
            ["Secs. 58–61", "Endorsements; certificate of registration (Sec. 60); copy and return of document", "Digital endorsements; Check and Register; Sec. 60 certificate; return", "FR-REG-083–086"],
            ["Secs. 71–77", "Refusal to register — reasons, appeal, application, suit", "Refused status entry point (workflow in PD-REG-02)", "7.3; FR-REG-087"],
            ["Sec. 78", "Fees for registration", "Registration fee shown on summary / payment (computed by Sr.13)", "FR-REG-070, FR-REG-073"],
        ],
    )
    w.h("3.2.2 The Registration (Karnataka Amendment) Act, 2023", 4)
    w.table(
        ["Section", "Topic", "Relevance", "Refer section 7 / 8 for implementation"],
        [
            ["Sec. 22-B", "Refusal to register forged / prohibited / attached-property documents", "Screening at SR verification; mandatory refusal", "FR-REG-062; BR-REG-005"],
            ["Secs. 22-C, 22-D", "Cancellation by District Registrar; appeal to IGR", "Post-registration status reversal (PD-REG-03)", "7.3 (exception statuses)"],
            ["Secs. 81-A, 81-B", "Penalties for registering in contravention", "Audit trail for screening decisions", "NFR-REG-AUD-001"],
        ],
    )
    w.h("3.2.3 The Karnataka Stamp Act, 1957 (touch points)", 4)
    w.table(
        ["Section", "Topic", "Relevance", "Refer section 7 / 8 for implementation"],
        [
            ["Sec. 3; Schedule", "Instruments chargeable; article-wise duty", "Article / sub-article master drives duty rule selection", "FR-REG-003"],
            ["Sec. 10", "Duties how to be paid", "Payment modes after send-for-payment (e-Stamp, challan, gateway)", "FR-REG-070"],
            ["Secs. 33–39", "Examination and impounding", "Impounded status (case-work out of scope)", "7.3; FR-REG-087"],
            ["Sec. 45-A", "Undervaluation — reference by registering officer", "Referred u/s 45-A status (case-work out of scope)", "7.3; FR-REG-087"],
        ],
    )
    w.h("3.2.4 The Transfer of Property Act, 1882", 4)
    w.table(
        ["Section", "Topic", "Relevance", "Refer section 7 / 8 for implementation"],
        [["Sec. 54", "Sale — registered instrument required for immovable property above ₹100", "Conveyance path in guided article selection", "FR-REG-001, FR-REG-004"]],
    )

    w.h("3.3 Relevant rules followed by the Department for Document Registration", 3)
    w.h("3.3.1 The Karnataka Registration Rules, 1965", 4)
    w.table(
        ["Rule", "Requirement", "Refer section 7 / 8 for implementation"],
        [
            ["Ch. II — Rules 3, 5", "Office hours and holidays", "Appointment / slot calendar — FR-APT-002"],
            ["Ch. VI — Rules 13–15", "Territorial divisions; survey / city survey description of property", "Village index and property details (KSRSAC / Bhoomi) — FR-REG-010–016, FR-REG-020–025"],
            ["Ch. IX — Rules 37, 40–46", "Office where a document may be registered; presentation; photograph; examination", "Jurisdiction, presentation and examination — FR-REG-012, FR-REG-080–082"],
            ["Ch. IX — Rules 51–55", "Suspension for fine / stamp; fines; condonation", "Suspended status; fine warning on appointment — FR-APT-003, FR-REG-087"],
            ["Ch. XII / XVI — Rules 71–73, 78–79, 94, 104", "Executant examination; thumb impression / photograph; endorsements; Sec. 60 certificate", "Registration completion — FR-REG-081, FR-REG-083–085"],
            ["Ch. XVII / XXV — Rules 110–118, 175–188", "Receipts; return of document; appeal against refusal", "Status tracking termini — FR-REG-086, 7.3"],
        ],
    )
    w.h("3.3.2 Stamp rules (payment visibility)", 4)
    w.table(
        ["Rule", "Requirement", "Refer section 7 / 8 for implementation"],
        [["Karnataka Stamp Rules, 1958 / e-Stamping Rules, 2009", "Impressed / franking / e-Stamp payment procedures at presentation", "Payment modes and e-Stamp capture — FR-REG-070–071"]],
    )

    w.h("3.4 Relevant notifications issued by the Department for Document Registration", 3)
    w.p("Gazette notifications and amendments cited for the Registration core:")
    w.table(
        ["Instrument", "Date / No.", "Effect", "Refer section 7 / 8 for implementation"],
        [
            ["RGN 2/2002-03", "1 Apr 2002; w.e.f. 4 Apr 2002", "Rule 19-A document sheets; Rules 22-A–22-C; Rule 40 photograph / digital photo", "Document format and photo at presentation — FR-REG-081"],
            ["RD 403 ESR 85 as amended by RD/46/MNMU/2025", "27 May 1986; amended 29 Aug 2025 (w.e.f. 31 Aug 2025)", "Table of Registration Fees under Sec. 78 — fee revision", "Registration fee on summary / payment — FR-REG-073"],
            ["RD 380 MUNOMU 2008", "8 Apr 2009", "Karnataka Stamp (Payment of Duty by Means of e-Stamping) Rules, 2009", "e-Stamp payment and verification at SRO — FR-REG-070"],
            ["Stamp Act Sec. 45-B (Act 8 of 2003) / Undervaluation Rules, 1977 (GSR 81)", "w.e.f. 1 Apr 2003 / 2 Mar 1977", "CVC market value guidelines; Sec. 45-A reference / Form I", "Hidden base rate and computed market value — FR-REG-040–043"],
        ],
    )


def stakeholders(w: W) -> None:
    w.h("4. Stakeholders and actors", 2)
    w.table(
        ["Actor", "Description", "Primary goals", "Channel involvement"],
        [
            ["Citizen / Presentant", "Person who prepares and presents the document (executant, claimant or authorised representative)", "Choose the right article, enter details once, pay, book a slot, track status", "Portal intake, payment, appointment; office presentation"],
            ["Executants / Claimants (parties)", "Parties to the instrument — labelled per transaction (Seller / Purchaser, Donor / Donee, etc.)", "Authenticate, appear, admit execution, receive registered document", "Aadhaar OTP on portal; biometrics and signatures at SRO"],
            ["Identifying witnesses", "Persons identifying the executants at the SRO", "Identify parties; sign endorsements", "At SRO"],
            ["Sub-Registrar (SR)", "Registering officer for the jurisdiction", "Verify application, send for payment, examine, admit, Check and Register, return document", "Department portal (SRO)"],
            ["SRO staff / operator", "Office staff supporting kiosk, scanning and queue", "Check-in, biometric capture support, scanning", "Department portal (SRO)"],
            ["District Registrar (DR)", "Supervising officer for SROs in the district", "Maintain SRO / holiday / slot masters; monitor ageing; handle escalations and exceptions", "Department portal (DR)"],
            ["IGR / Kaveri IT Cell", "Department administration and system owner", "Article master, templates, configuration, MIS", "Admin console"],
            ["KSRSAC", "Karnataka State Remote Sensing Applications Centre", "Village boundaries, Index-II village details, map tiles", "System integration"],
            ["Bhoomi (Revenue)", "Land records — RTC / survey data", "Survey number, hissa and extent for agricultural land", "System integration"],
            ["UIDAI / Aadhaar ASA", "Aadhaar authentication service", "OTP / e-KYC for parties", "System integration"],
            ["Treasury (Khajane-II) / payment gateway / e-Stamp CRA", "Collection of stamp duty and fees", "Payment, reconciliation, e-Stamp verification", "System integration"],
            ["ServiceDesk", "Support desk for Kaveri", "Fewer tickets for save / summary / schedule failures", "Support"],
        ],
    )


def glossary(w: W) -> None:
    w.h("5. Definitions and glossary", 2)
    w.table(
        ["Term", "Definition"],
        [
            ["Article / Sub-article", "Entry in the Karnataka Stamp Act Schedule identifying the instrument type (e.g., Conveyance, Gift, Partition 39-b, Release) and its duty rule"],
            ["Requirement-based selection", "Guided set of plain-language questions that resolves to an article / sub-article instead of a raw dropdown"],
            ["Executant / Claimant", "Party executing the document / party in whose favour it is executed; displayed with transaction-specific labels"],
            ["Presentant", "Person presenting the document for registration (Sec. 32)"],
            ["Village index / Index-II", "Village-wise index of registered documents; Index-II village details are retrieved from KSRSAC"],
            ["Property schedule", "Description of each property in the document (Schedule A, B, …) including location, extent and boundaries"],
            ["Schedule allocation", "Allocation of property schedules to individual parties (e.g., in partition or settlement deeds)"],
            ["Consideration amount", "Value stated in the document; caption varies by article (sale consideration, loan amount, rent, etc.)"],
            ["Guidance (guideline) value / Market value", "Rate notified by the Central Valuation Committee under Stamp Act Sec. 45-B; market value used for duty"],
            ["SD / RF", "Stamp duty / Registration fee"],
            ["Appointment / Slot", "Date and time booked at the jurisdictional SRO for presentation"],
            ["Check-in (Kiosk)", "Marking the parties' presence at the SRO on the appointed date"],
            ["Check and Register", "SR action that completes registration and assigns the registration number"],
            ["Sec. 60 certificate", "Certificate of registration endorsed on the document"],
            ["Minute Book", "Register of suspensions, refusals and other deviations"],
            ["11E sketch", "Revenue survey sketch for sub-division of agricultural land, required for certain transactions"],
            ["Geofencing", "Validating a captured map location against the selected village boundary polygon"],
            ["KSRSAC", "Karnataka State Remote Sensing Applications Centre"],
            ["RTC", "Record of Rights, Tenancy and Crops (Bhoomi)"],
            ["DLT", "Distributed Ledger Technology registration for SMS templates (TRAI)"],
        ],
    )


def current_state(w: W) -> None:
    w.h("6. Current state", 2)
    w.p(
        "In Kaveri 2.0 the citizen logs in, selects the document registration service and the SRO, "
        "and fills a sequence of screens: article / sub-article (standard dropdowns), village index, "
        "property details (including East-West / North-South measurement text fields for every property "
        "type), party information, property valuation (base guidance rate displayed), and a summary. "
        "The SR verifies the data and sends the application for payment; after payment the citizen "
        "schedules an appointment, visits the SRO, marks presence at the kiosk and presents the document. "
        "The department workflow then moves through numbered steps ending in Check and Register (Step 7)."
    )
    w.h("6.1 As-Is pain points", 3)
    w.p(
        "Pain points evidenced from the 27-08-2026 walkthrough and ServiceDesk tickets "
        "(ServiceDeskIssuesList.xlsx — OverallList + Categorized). The Addressed in column maps each "
        "item to the To-Be process and requirements."
    )
    w.table(
        ["Sr.No", "Pain Point", "Description", "Source", "Addressed in (this BRD)"],
        [
            ["1", "Article / sub-article selection and instrument mapping", "Standard dropdowns lead to wrong article; wrong or double stamp duty", "Discussion; tickets 17714, 31257, 93306", "7.2.1; FR-REG-001–005; BR-REG-001"],
            ["2", "Village / index / road master gaps", "Village index absent for split villages; wrong village shown; road names wrong; road option missing for agricultural land", "Tickets 13949, 29600, 95135, 31819", "FR-REG-010–015"],
            ["3", "Survey number not fetched from Bhoomi", "Sy. No. lookup fails and blocks property entry", "Tickets 20170, 30045", "FR-REG-016; FB-REG-002"],
            ["4", "Property schedule / boundary capture fails", "Unable to save schedule; Save and Continue blocked after 11E sketch; property missing in department summary", "Tickets 24450, 24247, 95367", "FR-REG-024–026"],
            ["5", "Irrelevant boundary fields for agricultural land", "East-West / North-South text fields shown for agricultural land; no visual of boundaries", "Discussion", "FR-REG-020–023"],
            ["6", "Party information save / display errors", "Unable to save party; claimant missing in summary; executant name 'None'; claimant photo in executant place", "Tickets 30276, 25784, 28927, 29588, 22789, 27709", "FR-REG-034"],
            ["7", "Generic party labels", "Parties labelled generically regardless of transaction type", "Discussion", "FR-REG-030"],
            ["8", "Generic Aadhaar OTP message; repeated data entry", "OTP SMS does not name the service; gender entered even when implied by salutation", "Discussion", "FR-REG-031–032"],
            ["9", "Schedule allocation shown for all transactions", "Schedule allocation feature displayed even where the transaction does not need it", "Discussion", "FR-REG-033"],
            ["10", "Property valuation / fee after valuation", "Market value blank; fee zero and Save disabled after Valuate; fee zero in summary", "Tickets 29357, 29532, 27923, 28596, 26292", "FR-REG-040–043; FB-REG-005"],
            ["11", "Base valuation rate exposed to citizens", "Citizens see base rate rather than just selecting property characteristics", "Discussion", "FR-REG-040–041"],
            ["12", "Pre-submission summary / review broken; static caption", "Summary cannot be generated / viewed; wrong consideration / market value; sub-article and challan missing; 'Consideration Amount' label static", "Discussion; tickets 27360, 27805, 30609, 12009, 6187, 10531, 92782, 31168, 29793", "FR-REG-050–055"],
            ["13", "Appointment / schedule after intake", "Unable to schedule; schedule not reflecting; schedule option missing after payment", "Tickets 27283, 30478, 30908, 28240, 28812", "7.2.3; FR-APT-001–004"],
            ["14", "Status / step stuck", "Check and Register (Step 7) disabled; applications stuck in Step 5; pending items show in Step 1 instead of Minute Book", "Tickets 94299, 94264, 92559, 23243, 25183, 22599, 31762", "7.3; FR-REG-084; FR-STS-003–004"],
        ],
    )


def add_figure(w: W, png: Path, caption: str) -> None:
    sec = w.doc.sections[-1]
    max_w = sec.page_width - sec.left_margin - sec.right_margin
    max_h = Inches(8.2)
    with Image.open(png) as im:
        ratio = im.width / im.height
    width = min(max_w, int(max_h * ratio))
    w.doc.add_picture(str(png), width=width)
    w.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    w.caption(caption)


def process_diagrams(w: W) -> None:
    w.h("7.2 Process Diagrams", 3)
    w.p(
        "The Registration core is shown as three simple process diagrams, drawn the same way as the CVC "
        "Guidance Value Fixation BRD: one vertical lane per actor, steps flowing top to bottom, and the "
        "legal reference in grey under each step. Numbered boxes match the numbered steps listed with "
        "each figure. Orange boxes are exceptions that continue in a separate process."
    )
    w.table(
        ["Process", "What it covers", "Figure"],
        [
            ["A — Filling the application", "Citizen fills the online application (27-08-2026 intake screens) and submits it", "Figure 1"],
            ["B — Sub-Registrar check and payment", "SR checks the application, sends it for payment; citizen pays", "Figure 2"],
            ["C — Appointment and registration", "Citizen books a slot, comes to the office, SR registers and returns the document", "Figure 3"],
        ],
        [2.1, 4.0, 1.0],
    )

    w.h("7.2.1 Process A — Filling the Document Registration application", 4)
    w.bullets([
        "Trigger: citizen wants to register a document (sale, gift, partition, mortgage, lease, release, …).",
        "Channel: citizen portal (online). Outside systems: KSRSAC (village details and map), Bhoomi (survey number and extent for farm land), Aadhaar (OTP).",
        "Result: application number generated and the application sent to the jurisdictional Sub-Registrar (Process B).",
    ])
    w.table(
        ["#", "Step (plain language)", "Who", "Reference / requirement"],
        [
            ["1", "Log in and choose “Register a document”", "Citizen", "User Management login"],
            ["2", "Answer a few simple questions about the deal", "Citizen", "FR-REG-001"],
            ["3", "Kaveri suggests the matching document type (article) in plain words", "Kaveri System", "Stamp Act Schedule; Sec. 17; FR-REG-001–003"],
            ["4", "Is the suggested type correct? If not, change the answers", "Citizen", "FR-REG-002"],
            ["5", "Choose where the property is: district, taluk, village", "Citizen", "Rules 13–15; FR-REG-010"],
            ["6", "KSRSAC sends village details and map; Bhoomi sends survey number and extent for farm land", "Outside systems", "FR-REG-011, FR-REG-014–016"],
            ["7", "Kaveri finds the right Sub-Registrar office for the property", "Kaveri System", "Sec. 28; FR-REG-012"],
            ["8", "Enter boundaries (East, West, North, South) and check the property on the map; no size boxes for farm land", "Citizen", "Secs. 21–22; FR-REG-020–026"],
            ["9", "Add the parties with labels for this deal (e.g. Seller and Buyer); gender filled from Mr / Mrs; allocate properties to parties only where the deal needs it", "Citizen", "Secs. 32–33; FR-REG-030, 032–036"],
            ["10", "Aadhaar sends an OTP that names this service; each party confirms", "Outside systems", "FR-REG-031"],
            ["11", "Pick the property type (residential, commercial, dry, wet …); the rate is not shown", "Citizen", "FR-REG-040"],
            ["12", "Kaveri works out market value, stamp duty and registration fee", "Kaveri System", "Stamp Act Secs. 45-A, 45-B; Sec. 78; FR-REG-041–042"],
            ["13", "Check the full summary (preview) and submit", "Citizen", "FR-REG-050–054"],
            ["14", "Kaveri gives an application number and sends it to the Sub-Registrar", "Kaveri System", "FR-REG-055; status Submitted for verification"],
        ],
    )
    add_figure(w, DIAGRAM_DIR / "Process_A_Registration_Application.png",
               "Figure 1 — Process A: Filling the Document Registration application. "
               "Source: ProcessDiagram/Process_A_Registration_Application.drawio")

    w.h("7.2.2 Process B — Sub-Registrar check and payment", 4)
    w.bullets([
        "Trigger: application submitted at the end of Process A.",
        "Channel: Sub-Registrar workbench (office); citizen portal for corrections and payment; Bank / Treasury / e-Stamp for payment.",
        "Result: payment received, receipt issued and appointment booking opened (Process C) — or the application is refused with reasons.",
    ])
    w.table(
        ["#", "Step (plain language)", "Who", "Reference / requirement"],
        [
            ["1", "Check the property and parties against the banned / attached property list", "Kaveri System", "Sec. 22-B; FR-REG-062"],
            ["2", "Open the application; check the details and the property value", "Sub-Registrar", "FR-REG-060, FR-REG-043"],
            ["3", "Is everything correct? If corrections are needed, the citizen corrects only the marked parts and resubmits; if forged / banned, refuse with written reasons", "Sub-Registrar / Citizen", "Secs. 22-B, 71; FR-REG-060–061"],
            ["4", "Send to the applicant for payment", "Sub-Registrar", "FR-REG-060; status Sent for payment"],
            ["5", "Pay stamp duty and fee online, by challan or e-Stamp", "Citizen", "Stamp Act Sec. 10; FR-REG-070"],
            ["6", "Payment received? If not, try again (no double payment)", "Bank / Treasury", "FR-REG-071–072; FB-REG-001"],
            ["7", "Give the receipt and open appointment booking right away", "Kaveri System", "FR-APT-001; status Payment completed"],
        ],
    )
    add_figure(w, DIAGRAM_DIR / "Process_B_SR_Check_and_Payment.png",
               "Figure 2 — Process B: Sub-Registrar check and payment. "
               "Source: ProcessDiagram/Process_B_SR_Check_and_Payment.drawio")

    w.h("7.2.3 Process C — Appointment and registration at the Sub-Registrar office", 4)
    w.bullets([
        "Trigger: payment completed at the end of Process B.",
        "Channel: citizen portal for booking; kiosk and Sub-Registrar workbench at the office. The District Registrar keeps the office calendar up to date.",
        "Result: document registered, registration number given and the registered document returned — or held / kept pending as an exception.",
    ])
    w.table(
        ["#", "Step (plain language)", "Who", "Reference / requirement"],
        [
            ["—", "Keep office working days, holidays and daily slot limits up to date", "District Registrar", "Rules 3, 5; FR-APT-002"],
            ["1", "Pick a date and time at the Sub-Registrar office", "Citizen", "FR-APT-001–002"],
            ["2", "Is the date within 4 months of signing the document? If not, show the late-fee warning; the citizen accepts the fine or picks an earlier date", "Kaveri System / Citizen", "Secs. 23, 25; Rules 51–55; FR-APT-003"],
            ["3", "Confirm the slot; send SMS and an appointment slip with QR code", "Kaveri System", "FR-APT-004, FR-APT-009; status Appointment booked"],
            ["4", "Can all parties come on that day? If not, change the date", "Citizen", "FR-APT-005–006"],
            ["5", "Come to the office on the day; scan the QR code at the kiosk", "Citizen", "FR-APT-007"],
            ["6", "Mark presence and add to the Sub-Registrar’s queue in slot order", "Kaveri System", "FR-APT-007; status Checked-in"],
            ["7", "Call the parties; note the time the document is presented", "Sub-Registrar", "Sec. 52; FR-REG-080"],
            ["8", "Full stamp duty paid and value not too low? If not, the document is held (impound / value check — separate process)", "Sub-Registrar", "Stamp Act Secs. 33–39, 45-A; FR-REG-082, FR-REG-087"],
            ["9", "All parties present and agree they signed? If not, keep it pending in the Minute Book; parties book a new date", "Sub-Registrar", "Secs. 34–35; Rules 51–55; FR-REG-087, FR-STS-004"],
            ["10", "Take photo and fingerprints of all parties", "Sub-Registrar", "Sec. 32A; Rule 40; FR-REG-081"],
            ["11", "Parties sign the endorsements", "Citizen", "Secs. 58–59; FR-REG-083"],
            ["12", "Click “Check and Register”; sign the certificate digitally", "Sub-Registrar", "Sec. 60; FR-REG-084–085"],
            ["13", "Give the registration number, update books and index, scan the document", "Kaveri System", "Sec. 51; FR-REG-085; status Registered"],
            ["14", "Return the registered document and digital copy", "Sub-Registrar", "Sec. 61; Rules 110–118; FR-REG-086; status Document returned"],
        ],
    )
    add_figure(w, DIAGRAM_DIR / "Process_C_Appointment_and_Registration.png",
               "Figure 3 — Process C: Appointment and registration at the Sub-Registrar office. "
               "Source: ProcessDiagram/Process_C_Appointment_and_Registration.drawio")


def future_state(w: W) -> None:
    w.h("7. Future state (To-Be)", 2)
    w.p(
        "The To-Be Registration core is a single service with online intake and in-person presentation "
        "at the jurisdictional SRO. It documents the channel model, three process diagrams with step tables "
        "(A — filling the application, B — Sub-Registrar check and payment, C — appointment and registration), "
        "the application status model, and what is new in Kaveri 3.0. Exception tracks (refusal, appeal, Sec. 22-B, Sec. 45-A, impounding) are "
        "entered from the statuses defined in 7.3 and specified in their own BRDs."
    )
    w.h("7.1 Channel models", 3)
    w.table(
        ["Service Type", "Online Activities", "Office Activities", "Mode"],
        [
            ["Document registration (standard)", "Guided article selection, village / property / party / valuation capture, Aadhaar OTP, preview and submit, payment, appointment booking, status tracking", "SR verification and send-for-payment, check-in, examination, biometrics, endorsements, Check and Register, Sec. 60 certificate, return of document", "Online intake + In-person presentation"],
            ["Exception handling (entry only)", "Status visibility and notifications", "Suspension, impounding, Sec. 45-A reference, refusal — per separate BRDs", "Office"],
        ],
    )

    process_diagrams(w)

    w.h("7.2.4 Intake screens — what changes", 4)
    w.p("Kaveri 3.0 behaviour for each intake screen discussed on 27-08-2026 (Process A):")
    w.table(
        ["Screen", "Kaveri 2.0 (As-Is)", "Kaveri 3.0 (To-Be)", "Requirements"],
        [
            ["A. Document classification", "Standard article / sub-article dropdowns", "Requirement-based guided selection resolving to article / sub-article; master-driven instrument mapping", "FR-REG-001–005"],
            ["B. Village index selection", "Manual cascading selection; gaps for split villages", "KSRSAC auto-populates Index-II village details; geofencing evaluated for location accuracy", "FR-REG-010–016"],
            ["C. Property details", "East-West / North-South text fields for all property types", "Fields removed for agricultural land; structured E / W / N / S boundaries; visual property map", "FR-REG-020–026"],
            ["D. Party information", "Generic party labels; generic OTP SMS; manual gender; schedule allocation always shown", "Labels per transaction; service-specific OTP template; gender from salutation; schedule allocation only where required", "FR-REG-030–036"],
            ["E. Property valuation", "Base valuation rate shown", "Citizen selects physical characteristics only (commercial, residential, dry, wet, …); rate hidden", "FR-REG-040–043"],
            ["F. Review and submission", "Static 'Consideration Amount' caption; unreliable summary", "Dynamic caption per document nature; mandatory comprehensive preview before submission for verification", "FR-REG-050–055"],
        ],
    )

    w.h("7.3 Application Status Model", 3)
    w.p("Statuses shared by citizen portal, SRO and District Registrar views:")
    w.table(
        ["Status", "Description", "Actor", "Next states"],
        [
            ["Draft", "Saved, not submitted", "Citizen", "Submitted for verification; Withdrawn"],
            ["Submitted for verification", "Preview confirmed; application number generated", "Citizen / System", "Under SR verification"],
            ["Under SR verification", "SR reviewing data, valuation and Sec. 22-B screening result", "Sub-Registrar", "Returned for correction; Sent for payment; Refused"],
            ["Returned for correction", "SR remarks recorded; citizen edits flagged sections only", "Citizen", "Submitted for verification; Withdrawn"],
            ["Sent for payment", "SD / RF payable shown to citizen", "System", "Payment pending confirmation; Payment completed"],
            ["Payment pending confirmation", "Gateway response awaited; pay action locked", "System", "Payment completed; Sent for payment (on failure)"],
            ["Payment completed", "Receipt generated; appointment booking enabled", "System", "Appointment booked"],
            ["Appointment booked", "Slot confirmed at jurisdictional SRO", "Citizen", "Appointment rescheduled; Checked-in; Appointment lapsed"],
            ["Appointment rescheduled", "Slot changed by citizen or SRO", "Citizen / Sub-Registrar", "Checked-in; Appointment lapsed"],
            ["Appointment lapsed", "No-show; payment retained", "System", "Appointment booked"],
            ["Checked-in", "Presence marked at kiosk; in SR queue", "Citizen / System", "Under examination"],
            ["Under examination", "Presented; presentation time recorded (Sec. 52)", "Sub-Registrar", "Admitted; Suspended; Impounded; Referred u/s 45-A; Refused"],
            ["Admitted", "Execution admitted; biometrics and endorsements in progress", "Sub-Registrar", "Registered"],
            ["Suspended (pending)", "Pending fine / party appearance; entered in Minute Book (Rules 51–55)", "Sub-Registrar", "Under examination; Refused"],
            ["Impounded / Referred u/s 45-A", "Stamp Act Secs. 33–39 / 45-A — case-work in separate BRD", "Sub-Registrar / District Registrar", "Under examination; Refused"],
            ["Refused", "Reasons recorded (Sec. 71 or Sec. 22-B); order copy to citizen", "Sub-Registrar", "Under appeal (separate BRD)"],
            ["Registered", "Check and Register done; registration number; Sec. 60 certificate signed", "Sub-Registrar / System", "Document returned"],
            ["Document returned", "Registered document and digital copy delivered (terminal)", "Sub-Registrar", "—"],
            ["Withdrawn", "Withdrawn by citizen before examination (terminal)", "Citizen", "—"],
        ],
    )

    w.h("7.4 What is new in Kaveri 3.0", 3)
    w.p("Material enhancements compared with the Kaveri 2.0 registration intake, appointment and status flow:")
    w.table(
        ["#", "Capability", "What is new in Kaveri 3.0"],
        [
            ["1", "Guided article selection", "Requirement-based questions replace raw dropdowns; master-driven instrument mapping"],
            ["2", "KSRSAC village data", "Index-II village details auto-populated; geofencing evaluated"],
            ["3", "Visual property map", "Property shown on map with E / W / N / S boundaries; agricultural E-W / N-S fields removed"],
            ["4", "Transaction-aware parties", "Seller / Purchaser, Donor / Donee, etc.; schedule allocation only where needed"],
            ["5", "Smart forms", "Gender from salutation; service-specific Aadhaar OTP SMS"],
            ["6", "Simplified valuation", "Citizens pick physical characteristics; base rate hidden; SR sees full computation"],
            ["7", "Mandatory preview", "Comprehensive summary with dynamic consideration caption before submission"],
            ["8", "Appointment after payment", "Slot booking opens immediately; reschedule, check-in and no-show handling"],
            ["9", "Explicit status model", "One status set across citizen, SRO and DR; ageing watchdog; no stuck steps"],
        ],
    )
    w.h("7.4.1 Rectified As-Is pain points", 4)
    w.table(
        ["Sr.No", "Pain Point (As-Is)", "How rectified in Kaveri 3.0"],
        [
            ["1", "Article / sub-article selection and instrument mapping", "Guided selection + effective-dated article master (FR-REG-001–005)"],
            ["2", "Village / index / road master gaps", "KSRSAC integration; split-village mapping; road master for all property types (FR-REG-010–015)"],
            ["3", "Survey number not fetched from Bhoomi", "Bhoomi integration with retry and clear error (FR-REG-016; FB-REG-002)"],
            ["4", "Property schedule / boundary capture fails", "Atomic, resumable save; 11E step non-blocking (FR-REG-024–026)"],
            ["5", "Irrelevant boundary fields for agricultural land", "Fields removed; map view (FR-REG-020–023)"],
            ["6", "Party information save / display errors", "Per-party persistence; single summary data model (FR-REG-034, FR-REG-054)"],
            ["7", "Generic party labels", "Labels from article master (FR-REG-030)"],
            ["8", "Generic OTP message; repeated data entry", "Service-specific template; gender auto-fill (FR-REG-031–032)"],
            ["9", "Schedule allocation shown for all transactions", "Shown only for flagged articles (FR-REG-033)"],
            ["10", "Valuation / fee after valuation", "No zero fee; explicit error and retry (FR-REG-042)"],
            ["11", "Base rate exposed to citizens", "Characteristics-only UI (FR-REG-040–041)"],
            ["12", "Summary broken; static caption", "Mandatory preview; dynamic caption (FR-REG-050–055)"],
            ["13", "Appointment after intake", "Booking enabled on payment; immediate reflection (FR-APT-001–004)"],
            ["14", "Status / step stuck", "Explicit next states; watchdog; Minute Book routing (FR-STS-003–004; FR-REG-084)"],
        ],
    )


def functional(w: W) -> None:
    w.h("8. Functional requirements", 2)
    w.p(
        "Functional requirements are organised by intake screen and lifecycle stage, aligned with 7 "
        "(To-Be). Section 7 is the process authority; requirements below cite the related step or status."
    )

    w.h("8.1 Document classification (article / sub-article)", 3)
    w.p("(Ref: 7.2.4 screen A; Registration Act Secs. 17–18; TPA Sec. 54; Stamp Act Sec. 3 and Schedule)")
    w.fr([
        ["FR-REG-001", "System shall provide a requirement-based selection: the citizen answers plain-language questions (nature of transaction, property type, relationship between parties, consideration involved) and the system proposes the matching article / sub-article with a plain-language description for confirmation", "Must", "Proposed article and description shown before confirmation; raw dropdown is not the primary entry"],
        ["FR-REG-002", "System shall offer keyword search over articles as a secondary path for experienced users and SRO staff", "Should", "Search returns article, sub-article and description"],
        ["FR-REG-003", "Article / sub-article master shall be configurable and effective-dated by IGR administrators and shall drive party labels, consideration caption, schedule-allocation flag and the duty / fee rule reference", "Must", "Master changes audited; effective date respected; tickets 17714, 31257, 93306 not reproducible"],
        ["FR-REG-004", "System shall indicate whether the selected instrument is compulsorily registrable (Sec. 17; TPA Sec. 54) or optional (Sec. 18)", "Should", "Indicator shown on selection and preview"],
        ["FR-REG-005", "Selected article and sub-article shall appear on both the citizen and department summaries", "Must", "Sub-article name present in department summary (ticket 92782)"],
    ])

    w.h("8.2 Village index and jurisdiction", 3)
    w.p("(Ref: 7.2.1 steps 5–7; Registration Act Secs. 5–7, 28–30; Rules 13–15, 37)")
    w.fr([
        ["FR-REG-010", "System shall capture property location through cascading district → taluk → hobli → village / ward selection from the location master", "Must", "Only valid combinations selectable"],
        ["FR-REG-011", "System shall integrate with KSRSAC to retrieve and auto-populate Index-II village details for the selected location", "Must", "Village code and Index-II details populated without manual entry; source stamped on record"],
        ["FR-REG-012", "System shall derive the jurisdictional SRO from the property location (Sec. 28); where properties lie in more than one sub-district, the citizen shall choose among those SROs", "Must", "Non-jurisdictional SRO not selectable"],
        ["FR-REG-013", "Geofencing (subject to feasibility study): system shall allow a map pin / GPS location and validate it against the KSRSAC village boundary, warning on mismatch", "Could", "Feasibility report accepted by department before build"],
        ["FR-REG-014", "Location master shall support split / renamed villages with effective-dated old-to-new mapping", "Must", "Tickets 13949, 29600 not reproducible"],
        ["FR-REG-015", "Road / locality master shall be available for agricultural and non-agricultural properties where the guidance value depends on it", "Should", "Tickets 95135, 31819 not reproducible"],
        ["FR-REG-016", "For agricultural land, system shall fetch survey number, hissa and extent from Bhoomi (RTC)", "Must", "Fetch success recorded; on failure, clear bilingual error and retry (FB-REG-002)"],
    ])

    w.h("8.3 Property details", 3)
    w.p("(Ref: 7.2.4 screen C; Registration Act Secs. 21–22; Rules 13–15)")
    w.fr([
        ["FR-REG-020", "For agricultural property, system shall not display the 'East to West' and 'North to South' measurement text fields", "Must", "Fields absent for agricultural property types"],
        ["FR-REG-021", "System shall capture East, West, North and South boundaries as structured fields for every property schedule", "Must", "All four boundaries mandatory"],
        ["FR-REG-022", "System shall display the property on a map (KSRSAC / survey layer) indicating East, West, North and South boundaries, on the property screen and in the preview", "Should", "Map renders with four boundary labels; text-only fallback when map service unavailable (FB-REG-004)"],
        ["FR-REG-023", "For non-agricultural property, system shall retain dimension fields (East-West, North-South) and built-up area", "Must", "Fields shown only for non-agricultural types"],
        ["FR-REG-024", "Property schedule save shall be atomic and resumable; the 11E sketch step shall not block Save and Continue once mandatory data is present", "Must", "Tickets 24450, 24247 not reproducible"],
        ["FR-REG-025", "System shall support multiple property schedules per document with unique schedule identifiers (Schedule A, B, …)", "Must", "Schedules listed in preview and department summary"],
        ["FR-REG-026", "Property details shall appear in the department summary", "Must", "Ticket 95367 not reproducible"],
    ])

    w.h("8.4 Party information", 3)
    w.p("(Ref: 7.2.4 screen D; Registration Act Secs. 32–35)")
    w.fr([
        ["FR-REG-030", "Party role labels shall change with the transaction type from the article master (e.g., Seller / Purchaser for sale, Donor / Donee for gift, Mortgagor / Mortgagee, Lessor / Lessee, Releasor / Releasee; default Executant / Claimant)", "Must", "Labels consistent on screens, preview, endorsements and summary"],
        ["FR-REG-031", "Aadhaar OTP authentication shall use service-specific SMS / message templates naming the service and purpose", "Must", "DLT-approved template per service; message names Document Registration and the action"],
        ["FR-REG-032", "Gender shall be auto-populated from the selected salutation (e.g., Mr → Male; Mrs / Ms / Smt / Kum → Female); for neutral salutations (Dr, Adv) gender remains selectable; Aadhaar e-KYC data takes precedence", "Must", "Gender pre-filled for gendered salutations"],
        ["FR-REG-033", "Schedule allocation (allocating property schedules to parties) shall appear only for transaction types flagged in the article master (e.g., partition, settlement) and be hidden for all others [interpretation to be confirmed by department]", "Must", "Feature absent for non-flagged articles"],
        ["FR-REG-034", "Each party shall be saved independently; summaries shall show all executants, claimants and witnesses with the correct name and photo in the correct position", "Must", "Tickets 30276, 25784, 28927, 29588, 22789, 27709 not reproducible"],
        ["FR-REG-035", "System shall capture representative capacity (power-of-attorney holder, guardian, authorised signatory) with PoA reference (Secs. 32–33)", "Must", "PoA details on preview and endorsement"],
        ["FR-REG-036", "System shall capture identifying witness details for presentation at the SRO [number and mandatory stage to be confirmed]", "Should", "Witness details on summary"],
    ])

    w.h("8.5 Property valuation (intake)", 3)
    w.p("(Ref: 7.2.4 screen E; Stamp Act Secs. 45-A, 45-B; computation rules in the Sr.13 / Sr.14 BRDs)")
    w.fr([
        ["FR-REG-040", "Citizen shall select only the property's physical characteristics (e.g., commercial, residential, dry, wet, construction type, floor) to proceed; the base valuation rate shall not be displayed", "Must", "No base rate visible on any citizen screen or preview"],
        ["FR-REG-041", "System shall compute market value, stamp duty and registration fee in the background from the CVC guideline master and show the resulting market value and amounts payable", "Must", "Values shown; computation traceable in audit"],
        ["FR-REG-042", "Valuation shall never produce a blank market value or zero fee silently; on failure, Save shall be disabled with an explicit error and a retry", "Must", "Tickets 29357, 29532, 27923, 28596, 26292 not reproducible"],
        ["FR-REG-043", "SR department view shall show base rate, selected characteristics and computed value for scrutiny", "Must", "Rate visible only to department roles"],
    ])

    w.h("8.6 Review and submission", 3)
    w.p("(Ref: 7.2.4 screen F; 7.2.1 step 13)")
    w.fr([
        ["FR-REG-050", "The 'Consideration Amount' label shall change with the nature of the document from the article master (e.g., Sale consideration, Loan amount, Annual rent / premium, Market value of gifted property)", "Must", "Caption matches article on entry screen, preview and summary"],
        ["FR-REG-051", "Before submission for verification, system shall show a mandatory comprehensive preview: article, property schedules and map, parties with roles and photos, valuation, consideration, SD / RF and payment details", "Must", "Submit disabled until preview viewed and declaration accepted"],
        ["FR-REG-052", "Preview shall be downloadable as a bilingual PDF", "Should", "PDF matches on-screen preview"],
        ["FR-REG-053", "Citizen shall be able to jump from a preview section to edit it and return to the preview", "Should", "Edits reflected immediately in preview"],
        ["FR-REG-054", "Citizen and department summaries shall be generated from the same data model and shall always be available", "Must", "Tickets 27360, 27805, 30609, 12009, 6187, 10531, 31168, 29793 not reproducible"],
        ["FR-REG-055", "On submission, system shall generate the application number and set status Submitted for verification", "Must", "Number on screen, SMS and dashboard"],
    ])

    w.h("8.7 SR verification and send for payment", 3)
    w.p("(Ref: 7.2.2 steps 1–4; Registration Act Sec. 22-B)")
    w.fr([
        ["FR-REG-060", "SR shall have a queue of submitted applications for the office, view complete data and valuation, and act: Send for payment, Return for correction (with remarks) or Refuse (with recorded reasons)", "Must", "Action, remarks and actor audited"],
        ["FR-REG-061", "On Return for correction, data shall be preserved and only flagged sections editable; resubmission returns to the same office queue", "Must", "Unflagged sections read-only"],
        ["FR-REG-062", "System shall screen the property and parties against Sec. 22-B prohibited / attached-property lists and flag hits to the SR", "Must", "Screening result shown in SR view and audited"],
    ])

    w.h("8.8 Fees and payments", 3)
    w.p("(Ref: 7.2.2 steps 5–7; Registration Act Sec. 78; Stamp Act Sec. 10; e-Stamping Rules, 2009)")
    w.fr([
        ["FR-REG-070", "After Send for payment, system shall show the SD / RF breakdown and the permitted payment modes (online gateway / Treasury challan / e-Stamp certificate)", "Must", "Breakdown matches Sr.13 engine output"],
        ["FR-REG-071", "System shall reconcile payment, generate a receipt and show challan / e-Stamp details on the summary", "Must", "Tickets 31168, 29793 not reproducible"],
        ["FR-REG-072", "On gateway timeout, system shall poll for final status and lock the pay action until a terminal status is received", "Must", "No duplicate debit (FB-REG-001)"],
        ["FR-REG-073", "Registration fee master shall reflect RD/46/MNMU/2025 (w.e.f. 31-08-2025) and later notifications with effective dates", "Must", "Fee per current notification"],
    ])

    w.h("8.9 Appointment", 3)
    w.p("(Ref: 7.2.3 steps 1–6; Registration Act Secs. 10–12, 23, 25; Rules 3, 5, 51–55)")
    w.fr([
        ["FR-APT-001", "Appointment booking shall be enabled immediately on successful payment and reachable from the receipt and dashboard", "Must", "Tickets 27283, 30478, 30908, 28812 not reproducible"],
        ["FR-APT-002", "Slot calendar per SRO shall be derived from office hours and the holiday master, with slot capacity configurable by DR / SR", "Must", "Holidays and non-working hours not bookable"],
        ["FR-APT-003", "System shall check the presentation window (Sec. 23 — four months from execution) and warn when the chosen date falls in the Sec. 25 delay period with fine", "Must", "Warning text cites Sec. 25 and fine rule"],
        ["FR-APT-004", "Booking shall reflect immediately on the citizen dashboard and the SRO day list", "Must", "Ticket 28240 not reproducible"],
        ["FR-APT-005", "Citizen shall be able to reschedule or cancel before a configurable cut-off [reschedule limit to be confirmed]", "Should", "Reschedule audited; parties notified"],
        ["FR-APT-006", "SRO shall be able to reschedule appointments in bulk on office closure or officer absence (Secs. 10–12), notifying citizens", "Must", "Bulk reschedule with reason"],
        ["FR-APT-007", "Kiosk check-in on arrival shall mark presence, set status Checked-in and place the application in the SR queue in slot order", "Must", "Queue ordered by slot and check-in time"],
        ["FR-APT-008", "Unattended appointments shall lapse at day end; citizen may rebook without repaying", "Must", "Status Appointment lapsed; payment retained"],
        ["FR-APT-009", "System shall issue an appointment slip with QR code (application number, SRO, date, time, parties to appear)", "Should", "QR scannable at kiosk"],
    ])

    w.h("8.10 Presentation, examination and registration", 3)
    w.p("(Ref: 7.2.3 steps 7–14; Registration Act Secs. 32A–35, 52, 58–61; Rules 40–46, 71–73, 78–79, 94, 104, 110–118)")
    w.fr([
        ["FR-REG-080", "System shall record presentation date, time and presentant when the SR calls the application (Sec. 52)", "Must", "Presentation timestamp on endorsement"],
        ["FR-REG-081", "System shall capture photograph and fingerprints of presentant and parties (Sec. 32A; Rule 40)", "Must", "Biometric references stored and masked in logs"],
        ["FR-REG-082", "SR shall complete an examination checklist: identity, admission of execution (Secs. 34–35), SD / RF sufficiency, presentation time limit", "Must", "Checklist complete before Admit"],
        ["FR-REG-083", "System shall generate endorsements digitally (Secs. 58–59) for party signature / thumb impression", "Must", "Endorsement text per Rules 71–73, 78–79, 94"],
        ["FR-REG-084", "Check and Register shall be enabled once all mandatory endorsements, payments and biometrics are complete; when disabled, the screen shall list the missing prerequisites", "Must", "Tickets 94299, 94264, 92559 not reproducible"],
        ["FR-REG-085", "On registration, system shall assign registration number and book / volume / page, generate the Sec. 60 certificate signed with the SR's DSC, scan the document and update Index-II", "Must", "Certificate signature verifiable"],
        ["FR-REG-086", "System shall record return of the registered document and provide a digital copy (Sec. 61; Rules 110–118)", "Must", "Status Document returned"],
        ["FR-REG-087", "Exceptions (suspension, impounding, Sec. 45-A reference, refusal) shall move the application to the explicit status in 7.3 and record it in the Minute Book; detailed workflows follow the related BRDs", "Must", "No exception leaves the application in an intermediate step"],
    ])

    w.h("8.11 Status tracking", 3)
    w.p("(Ref: 7.3)")
    w.fr([
        ["FR-STS-001", "A single status model (7.3) shall be used by the citizen portal, SRO and District Registrar views", "Must", "Same status label everywhere"],
        ["FR-STS-002", "Citizen shall see a timeline with each status, date / time, actor role, remarks and the next action required", "Must", "Timeline bilingual"],
        ["FR-STS-003", "Every status shall have defined next states; a watchdog shall flag applications idle beyond a configurable period to the SR and DR", "Must", "Tickets 23243, 25183 not reproducible; ageing alert raised"],
        ["FR-STS-004", "Suspended and other pending documents shall appear in the Minute Book / pending register and not in the intake queue", "Must", "Tickets 22599, 31762 not reproducible"],
        ["FR-STS-005", "Any party shall be able to track by application or registration number with OTP to the registered mobile", "Must", "OTP-gated tracking"],
        ["FR-STS-006", "Each status change shall emit an event for notifications and MIS", "Must", "Event per transition"],
        ["FR-STS-007", "Department users shall correct status only through controlled actions with reason; no back-end edits", "Must", "All corrections audited"],
    ])

    w.h("8.12 User stories", 3)
    w.p("Derived from the 27-08-2026 discussion sections and mapped Acts / Rules (Consolidated report §2.6):")
    w.table(
        ["ID", "User story", "Requirements"],
        [
            ["US-REG-01", "As a citizen, I want to select article / sub-article through a requirement-based interface instead of opaque dropdowns, so that I choose the correct instrument for my transaction.", "FR-REG-001–005"],
            ["US-REG-02", "As a citizen / Sub-Registrar, I want Index-II village details retrieved via KSRSAC (and geofencing evaluated for location accuracy), so that village and location masters are accurate at intake.", "FR-REG-010–016"],
            ["US-REG-03", "As a citizen, I want to enter agricultural property boundaries without East-West / North-South text fields and see boundaries on a visual map, so that property schedule capture matches field practice.", "FR-REG-020–026"],
            ["US-REG-04", "As a citizen, I want party labels that change with transaction type (Seller / Purchaser, Donor / Donee, etc.), so that party roles are clear for the deed type.", "FR-REG-030"],
            ["US-REG-05", "As a citizen, I want to authenticate with Aadhaar OTP using service-specific SMS templates and have gender auto-filled from salutation, so that party capture is faster and less error-prone.", "FR-REG-031–032"],
            ["US-REG-06", "As a citizen, I want schedule allocation shown only when the transaction type requires it, so that unnecessary steps are hidden.", "FR-REG-033"],
            ["US-REG-07", "As a citizen, I want to complete valuation by selecting physical characteristics without seeing the base rate, so that the valuation UI stays simple while duty still uses official rates.", "FR-REG-040–043"],
            ["US-REG-08", "As a citizen, I want to review a full document summary with dynamic consideration captions before submission, so that I can correct errors before verification.", "FR-REG-050–055"],
            ["US-REG-09", "As a citizen, I want to book my appointment as soon as I pay and see it confirmed, so that I can plan my visit to the SRO.", "FR-APT-001–009"],
            ["US-REG-10", "As a citizen / SR / DR, I want one clear status and timeline for every application, so that nothing gets stuck unnoticed.", "FR-STS-001–007"],
        ],
    )

    w.h("8.13 Notifications", 3)
    w.fr([
        ["FR-REG-090", "System shall send SMS / email on submission, return for correction, send for payment, payment success / failure, appointment booked / rescheduled / lapsed, registration completed and document ready", "Should", "Bilingual EN / KN DLT-approved templates"],
        ["FR-REG-091", "System shall send an appointment reminder to all parties one day before the slot", "Should", "Reminder delivered to each party's mobile"],
    ])

    w.h("8.14 Reports and MIS", 3)
    w.fr([
        ["FR-REG-095", "SRO daily appointment list with check-in and no-show status", "Must", "Exportable"],
        ["FR-REG-096", "Applications by status with ageing per SRO / district", "Must", "Drill-down to application"],
        ["FR-REG-097", "Slot utilisation and no-show rate per SRO", "Should", "Period filter"],
        ["FR-REG-098", "Return-for-correction reasons and registrations by article / sub-article", "Should", "Period and office filters"],
        ["FR-REG-099", "Integration health (KSRSAC, Bhoomi, Aadhaar, payment) failure counts", "Should", "Daily summary to IT Cell"],
    ])

    w.h("8.15 Business rules", 3)
    w.table(
        ["Rule ID", "Description", "Statutory ref", "System enforcement"],
        [
            ["BR-REG-001", "Article / sub-article must be selected from the active master; it determines labels, caption, schedule allocation and duty rule", "Stamp Act Sec. 3, Schedule", "Hard stop; FR-REG-001–003"],
            ["BR-REG-002", "Village and Index-II details must come from KSRSAC unless the fallback is active", "Rules 13–15", "FR-REG-011; FB-REG-002"],
            ["BR-REG-003", "Every property schedule must have four boundaries; agricultural schedules do not carry E-W / N-S dimensions", "Secs. 21–22; Rules 13–15", "FR-REG-020–021"],
            ["BR-REG-004", "Application may be presented only at the jurisdictional SRO", "Secs. 28–30; Rule 37", "FR-REG-012"],
            ["BR-REG-005", "Sec. 22-B hits must be decided by the SR before send for payment", "Karnataka Act 47 of 2024 Sec. 22-B", "FR-REG-062"],
            ["BR-REG-006", "No payment before SR sends for payment; no appointment before payment completed", "—", "FR-REG-070; FR-APT-001"],
            ["BR-REG-007", "Appointment beyond the four-month window shows the delay-fine warning", "Secs. 23, 25; Rules 51–55", "FR-APT-003"],
            ["BR-REG-008", "Base rate is never shown on citizen screens", "Stamp Act Sec. 45-B (CVC)", "FR-REG-040"],
            ["BR-REG-009", "Submission requires preview viewed and declaration accepted", "—", "FR-REG-051"],
            ["BR-REG-010", "Check and Register only after all endorsements, payments and biometrics are complete", "Secs. 58–60; Rules 71–73, 78–79", "FR-REG-084"],
            ["BR-REG-011", "Pending (suspended) documents are tracked in the Minute Book", "Rules 51–55", "FR-STS-004"],
        ],
    )

    w.h("8.16 User interface (high-level)", 3)
    w.table(
        ["Screen / step", "Purpose", "Channel", "Statutory alignment", "Notes"],
        [
            ["Guided article selection", "Resolve article / sub-article", "Online", "Secs. 17–18; Stamp Act Schedule", "7.2.4 A; Process A steps 2–4"],
            ["Village index", "Location and Index-II village", "Online", "Rules 13–15", "7.2.4 B; Process A steps 5–7"],
            ["Property details + map", "Schedules, boundaries, map", "Online", "Secs. 21–22", "7.2.4 C; Process A step 8"],
            ["Party information", "Parties, OTP, schedule allocation", "Online", "Secs. 32–33", "7.2.4 D; Process A steps 9–10"],
            ["Property valuation", "Characteristics; computed value", "Online", "Stamp Act Secs. 45-A / 45-B", "7.2.4 E; Process A steps 11–12"],
            ["Preview and submit", "Comprehensive review and declaration", "Online", "—", "7.2.4 F; Process A steps 13–14"],
            ["SR verification queue", "Review, send for payment, return, refuse", "Office", "Secs. 22-B, 71", "8.7; Process B steps 1–4"],
            ["Payment", "SD / RF payment and receipt", "Online", "Sec. 78; Stamp Act Sec. 10", "8.8; Process B steps 5–7"],
            ["Appointment booking", "Slots, reschedule, slip", "Online", "Secs. 23, 25; Rules 3, 5", "7.2.3; Process C steps 1–4"],
            ["Kiosk check-in", "Mark presence", "Office", "—", "FR-APT-007; Process C steps 5–6"],
            ["Examination and registration", "Checklist, biometrics, endorsements, Check and Register", "Office", "Secs. 32A–35, 52, 58–61", "8.10; Process C steps 7–14"],
            ["Status timeline", "Track application", "Both", "—", "8.11"],
        ],
    )
    w.p("Wireframe links: [Figma / prototype URLs]")
    w.p("Bilingual: all labels [EN / KN] — content manager sign-off.")

    w.h("8.17 Integrations", 3)
    w.table(
        ["Integration", "Direction", "Purpose", "Channel", "Owner", "Status"],
        [
            ["KSRSAC", "Inbound", "Village boundaries, Index-II village details, map layer, geofence polygon", "Online", "", "To be confirmed"],
            ["Bhoomi (RTC)", "Inbound", "Survey number, hissa, extent for agricultural land", "Online", "", ""],
            ["UIDAI / Aadhaar ASA", "Outbound", "OTP authentication / e-KYC of parties", "Online", "", ""],
            ["SMS gateway (DLT) / Email", "Outbound", "Service-specific OTP and status notifications", "Both", "", ""],
            ["Payment gateway / Treasury (Khajane-II)", "Outbound", "Stamp duty and registration fee collection; reconciliation", "Online", "", ""],
            ["e-Stamp CRA", "Inbound", "e-Stamp certificate verification", "Both", "", ""],
            ["CVC guideline master (internal)", "Internal", "Base rate lookup for valuation", "Both", "", ""],
            ["Stamp duty / fee engine (Sr.13, internal)", "Internal", "SD / RF computation", "Both", "", ""],
            ["DSC / eSign", "Outbound", "SR signature on Sec. 60 certificate", "Office", "", ""],
            ["Biometric devices / kiosk", "Inbound", "Photo, fingerprint, check-in", "Office", "", ""],
            ["Scanning Module (shared)", "Internal", "Scan registered document", "Office", "", ""],
        ],
    )
    w.p("Interface requirements: [API list TBD by Architect]")

    w.h("8.18 Data requirements", 3)
    w.h("8.18.1 Core entities (logical)", 4)
    w.bullets([
        "Application, ArticleSelection, PropertySchedule (location, boundaries, map reference), Party (role, representative, biometric reference), ScheduleAllocation, Valuation, Payment, Appointment, CheckIn, Examination, Endorsement, Registration (number, book / volume / page), Certificate, StatusEvent, AuditEvent.",
        "Masters: Article / sub-article (effective-dated), Location (district → village, split-village mapping), Road / locality, SRO and jurisdiction, Office hours and holidays, Slot capacity, Notification templates.",
    ])
    w.h("8.18.2 Retention", 4)
    w.p("Permanent preservation of registered documents, registers, indexes and certificates; application drafts and lapsed appointments per department retention policy [to be confirmed].")
    w.h("8.18.3 Migration (high level)", 4)
    w.table(
        ["Topic"],
        [
            ["Article / sub-article and location masters from Kaveri 2.0, cleansed against KSRSAC village codes"],
            ["SRO, holiday and slot masters"],
            ["In-flight Kaveri 2.0 applications (submitted, paid, appointment booked) mapped to the 7.3 status model"],
        ],
    )

    w.h("8.19 Requirements traceability matrix (RTM) - template", 3)
    w.table(
        ["Req ID", "Act/Rule/Form", "Requirement summary", "BRD section", "UI screen", "Test case ID", "Status"],
        [
            ["FR-REG-001", "Secs. 17–18; TPA Sec. 54", "Requirement-based article selection", "8.1", "Guided article selection", "TC-REG-___", "Draft"],
            ["FR-REG-003", "Stamp Act Sec. 3, Schedule", "Effective-dated article master", "8.1", "Admin", "TC-REG-___", "Draft"],
            ["FR-REG-011", "Rules 13–15", "KSRSAC Index-II village details", "8.2", "Village index", "TC-REG-___", "Draft"],
            ["FR-REG-012", "Secs. 28–30; Rule 37", "Jurisdictional SRO", "8.2", "Village index", "TC-REG-___", "Draft"],
            ["FR-REG-013", "—", "Geofencing (feasibility)", "8.2", "Village index", "TC-REG-___", "Draft"],
            ["FR-REG-020", "Secs. 21–22", "No E-W / N-S for agricultural", "8.3", "Property details", "TC-REG-___", "Draft"],
            ["FR-REG-022", "Secs. 21–22", "Visual property map", "8.3", "Property details", "TC-REG-___", "Draft"],
            ["FR-REG-030", "—", "Transaction-based party labels", "8.4", "Party information", "TC-REG-___", "Draft"],
            ["FR-REG-031", "—", "Service-specific Aadhaar OTP template", "8.4", "Party information", "TC-REG-___", "Draft"],
            ["FR-REG-032", "—", "Gender from salutation", "8.4", "Party information", "TC-REG-___", "Draft"],
            ["FR-REG-033", "—", "Conditional schedule allocation", "8.4", "Party information", "TC-REG-___", "Draft"],
            ["FR-REG-040", "Stamp Act Sec. 45-B", "Hide base rate", "8.5", "Property valuation", "TC-REG-___", "Draft"],
            ["FR-REG-050", "—", "Dynamic consideration caption", "8.6", "Preview and submit", "TC-REG-___", "Draft"],
            ["FR-REG-051", "—", "Mandatory pre-submission preview", "8.6", "Preview and submit", "TC-REG-___", "Draft"],
            ["FR-REG-062", "Sec. 22-B", "Prohibited-property screening", "8.7", "SR verification", "TC-REG-___", "Draft"],
            ["FR-REG-073", "Sec. 78; RD/46/MNMU/2025", "Fee master per notification", "8.8", "Payment", "TC-REG-___", "Draft"],
            ["FR-APT-001", "—", "Booking enabled on payment", "8.9", "Appointment booking", "TC-APT-___", "Draft"],
            ["FR-APT-002", "Rules 3, 5", "Office hours / holiday calendar", "8.9", "Appointment booking", "TC-APT-___", "Draft"],
            ["FR-APT-003", "Secs. 23, 25", "Presentation window warning", "8.9", "Appointment booking", "TC-APT-___", "Draft"],
            ["FR-REG-081", "Sec. 32A; Rule 40", "Photo and fingerprints", "8.10", "Examination", "TC-REG-___", "Draft"],
            ["FR-REG-084", "Secs. 58–60", "Check and Register prerequisites", "8.10", "Examination", "TC-REG-___", "Draft"],
            ["FR-REG-085", "Sec. 60", "Registration number and certificate", "8.10", "Examination", "TC-REG-___", "Draft"],
            ["FR-STS-003", "—", "Ageing watchdog / no stuck steps", "8.11", "Status timeline", "TC-STS-___", "Draft"],
            ["FR-STS-004", "Rules 51–55", "Minute Book routing", "8.11", "SR dashboard", "TC-STS-___", "Draft"],
        ],
    )


def nfr(w: W) -> None:
    w.h("9. Non-functional requirements", 2)
    w.p(
        "To ensure the Registration core operates reliably, securely and efficiently, the following "
        "non-functional parameters shall be incorporated into design, build and acceptance testing."
    )
    w.h("9.1 Performance and Scalability", 3)
    w.p("Targets below are load-test gates before production [values to be confirmed with the Architect].")
    w.fr([
        ["NFR-REG-PERF-001", "The system shall process standard user-interface interactions within 2 seconds and complete external API calls (KSRSAC, Bhoomi, Aadhaar, payment gateway) within 5 seconds", "Must", "Measured at p95 under NFR-REG-PERF-002 load"],
        ["NFR-REG-PERF-002", "The portal shall support the peak concurrent citizen and SRO sessions defined by the department without degradation", "Must", "Load-test evidence at the agreed concurrency"],
        ["NFR-REG-SCALE-001", "Infrastructure shall scale automatically for surges such as auspicious days, month-end and the eve of guidance value revisions", "Must", "Surge test at 3× baseline without manual scale-up"],
        ["NFR-REG-PERF-003", "Preview / summary generation shall complete within 3 seconds for documents with up to 20 parties and 20 property schedules", "Should", "Measured at p95"],
    ])
    w.h("9.2 Security and Data Privacy", 3)
    w.fr([
        ["NFR-REG-SEC-001", "All data in transit shall use TLS 1.3 and all data at rest (PII, documents, biometrics) shall be encrypted with AES-256", "Must", "Confirmed in hosting design"],
        ["NFR-REG-PRIV-001", "Aadhaar numbers and biometric references shall be masked in logs and unauthorised views", "Must", "No raw values for unauthorised roles"],
        ["NFR-REG-PRIV-002", "Base guidance rate shall be withheld from citizen-facing APIs as well as screens", "Must", "API responses to citizen role contain no base rate"],
        ["NFR-REG-AUD-001", "All status changes, SR actions, Sec. 22-B screening decisions and master changes shall generate an immutable, timestamped audit log tied to the actor", "Must", "Append-only audit with previous / next state and reason"],
    ])
    w.h("9.3 System Availability and Error Handling", 3)
    w.fr([
        ["NFR-REG-AVA-001", "If an external integration is unavailable, the system shall show a clear bilingual message and pause the workflow safely without losing data", "Must", "Application resumable; no data loss"],
        ["NFR-REG-PAY-001", "On payment-gateway timeout, the system shall poll for final status and prevent duplicate payment", "Must", "No double-debit; aligns with FR-REG-072"],
    ])
    w.h("9.4 Security Audit and Compliance (VAPT Policy)", 3)
    w.fr([
        ["NFR-REG-VAPT-001", "The application shall undergo comprehensive VAPT before production deployment", "Must", "VAPT report accepted before Go-Live"],
        ["NFR-REG-VAPT-002", "VAPT scope shall cover web application, APIs, payment integration and responsive interfaces, including OWASP Top 10, business-logic flaws and unauthorised data access", "Must", "Scope statement and evidence"],
        ["NFR-REG-VAPT-003", "Critical / high findings shall be remediated and re-verified before Go-Live; automated scanning alone is insufficient", "Must", "Zero open Critical / High at Go-Live"],
        ["NFR-REG-VAPT-004", "VAPT shall be repeated at least annually and on significant architectural or feature change", "Must", "Annual and change-triggered VAPT scheduled"],
    ])


def risks(w: W) -> None:
    w.h("10. Risk and Mitigation Strategy", 2)
    w.p("Operational, technical and adoption risks for the Registration core and their mitigations:")
    w.table(
        ["Risk ID", "Risk", "Mitigation", "Related requirements"],
        [
            ["RS-REG-001", "Guided questions resolve to the wrong article because the question set or mapping is incomplete", "Department-validated question tree; SR can correct the article on return for correction; MIS on corrected articles to refine the tree", "FR-REG-001–003; FR-REG-098"],
            ["RS-REG-002", "KSRSAC or Bhoomi downtime blocks intake", "Resumable drafts; controlled fallback with SR verification flag", "FR-REG-011, FR-REG-016; FB-REG-002"],
            ["RS-REG-003", "Geofencing inaccuracy (GPS drift, boundary data quality) produces false warnings", "Feasibility study first; warnings only, not hard stops", "FR-REG-013"],
            ["RS-REG-004", "Slot capacity misconfigured, causing crowding or idle SROs", "DR-managed capacity; utilisation MIS; bulk reschedule", "FR-APT-002, FR-APT-006; FR-REG-097"],
            ["RS-REG-005", "Hiding the base rate causes citizen distrust of the computed value", "Show computed market value and amounts; public guidance value lookup (Sr.14); SR sees full computation", "FR-REG-040–043"],
            ["RS-REG-006", "Article master errors cascade to labels, captions and duty", "Effective-dated master with maker-checker approval and regression tests per article", "FR-REG-003; NFR-REG-AUD-001"],
        ],
    )


def fallbacks(w: W) -> None:
    w.h("11. System Fallbacks & Error Handling", 2)
    w.p(
        "The system shall handle exceptions and integration failures without losing citizen data. "
        "9.3 states the availability principles; the controls below specify behaviour."
    )
    w.fr([
        ["FB-REG-001", "Payment gateway timeout: poll the gateway for final status; lock the pay action until Success or Failed", "Must", "Same control as NFR-REG-PAY-001"],
        ["FB-REG-002", "KSRSAC / Bhoomi unavailable: save draft, show bilingual message and retry; after a configurable period allow manual entry flagged for SR verification", "Must", "Flag visible to SR; source recorded as manual"],
        ["FB-REG-003", "Aadhaar OTP failure: allow retry within limits; otherwise defer authentication to biometric verification at the SRO", "Must", "Deferred status visible to SR"],
        ["FB-REG-004", "Map service unavailable: show text boundaries and continue; map rendered later in preview when available", "Should", "Workflow not blocked"],
        ["FB-REG-005", "Valuation / fee engine failure: block Save with explicit error and retry; never persist zero fee", "Must", "Aligns with FR-REG-042"],
        ["FB-REG-006", "SMS / email gateway delay: queue and retry notifications; workflow steps not halted", "Must", "Retry queue drains"],
        ["FB-REG-007", "DSC failure at Check and Register: keep status Admitted with 'Pending SR signature'; retry signing without re-entering data", "Must", "Retry resumes signing only"],
    ])


def training(w: W) -> None:
    w.h("12. Training and Change Management", 2)
    w.p(
        "Moving from Kaveri 2.0 to Kaveri 3.0 changes how citizens select articles, enter property and "
        "party details, and book appointments, and how SROs verify and track applications."
    )
    w.h("12.1 Target audience", 3)
    w.p("Sub-Registrars, SRO staff (kiosk, scanning), District Registrars, Kaveri IT Cell / ServiceDesk, and document writers / advocates who prepare documents for citizens.")
    w.h("12.2 Training delivery", 3)
    w.bullets([
        "Role-based workshops for SRO staff on the verification queue, return for correction, appointment day list, check-in and the new status model.",
        "SOPs, quick-reference cards and video tutorials for guided article selection, property map, schedule allocation and preview.",
        "Briefing for document writers on article questions, transaction-based party labels and hidden base rate.",
    ])
    w.h("12.3 Citizen change management", 3)
    w.p("Tooltips, dynamic help text and the guided article questions explain each step; the preview and status timeline show what happens next.")
    w.h("12.4 Post-Go-Live support", 3)
    w.p("Dedicated hyper-care support for the first 90 days, monitoring integration failures, stuck-status alerts and appointment issues.")


def appendix(w: W) -> None:
    w.h("Appendix A — References", 2)
    w.bullets([
        "The Registration Act, 1908 (Central Act 16 of 1908)",
        "The Registration (Karnataka Amendment) Act, 2023 (Karnataka Act 47 of 2024)",
        "The Karnataka Registration Rules, 1965",
        "The Karnataka Stamp Act, 1957 and Schedule; Karnataka Stamp Rules, 1958; e-Stamping Rules, 2009",
        "The Transfer of Property Act, 1882 — Sec. 54",
        "Notifications RGN 2/2002-03; RD 403 ESR 85 as amended by RD/46/MNMU/2025; RD 380 MUNOMU 2008",
        "Document_Registration_requirement_27082026_v1.1.docx; Consolidated_Requirement_Discussions_25082026_to_11092026.docx",
        "ServiceDeskIssuesList.xlsx (OverallList + Categorized)",
        "OWASP Top 10 — https://owasp.org/ (VAPT scope, NFR-REG-VAPT-002)",
    ])
    w.h("Acceptance and sign-off of BRD", 2)
    w.table(
        ["Role", "Name", "Signature / Date"],
        [["Product Owner", "", ""], ["Domain Expert", "", ""], ["Business Analyst", "", ""]],
    )


def build() -> None:
    doc = load_template()
    w = W(doc)
    front_matter(w)
    executive_summary(w)
    scope(w)
    legal(w)
    stakeholders(w)
    glossary(w)
    current_state(w)
    future_state(w)
    functional(w)
    nfr(w)
    risks(w)
    fallbacks(w)
    training(w)
    appendix(w)
    doc.core_properties.title = f"BRD — {MODULE}"
    doc.core_properties.author = "Nandha Kumar"
    doc.save(str(DST))
    print(f"Wrote {DST}")


if __name__ == "__main__":
    build()
