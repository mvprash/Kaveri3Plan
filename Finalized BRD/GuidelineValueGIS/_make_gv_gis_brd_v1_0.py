"""Generate BRD — Guideline Value Calculation & GIS Valuation (v1.0).

Source: Requirement Discussions/Daily Reports/Consolidated_Requirement_Discussions_25082026_to_11092026.docx
        Section 4 (01-09-2026 — Guideline Value, CVC Valuation Module & GIS Valuation).
Template / section structure: Finalized BRD/Marriage/RFP/BRD_Marriage_BRD_v8.docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
TEMPLATE = ROOT / "Finalized BRD" / "Marriage" / "RFP" / "BRD_Marriage_BRD_v8.docx"
DIAGRAMS = HERE / "Process Diagram"
OUT = HERE / "BRD_Guideline_Value_GIS_Valuation_v1.0.docx"

HEADER_TEXT = "Kaveri 3.0 | BRD  | Guideline Value Calculation & GIS Valuation"
HEADER_FILL = "D9E2F3"


# --------------------------------------------------------------------------- helpers
def clear_body(doc) -> None:
    body = doc.element.body
    for el in list(body.iterchildren()):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/image") or rel.reltype.endswith("/comments"):
            doc.part.drop_rel(rid)


def set_header(doc) -> None:
    for sec in doc.sections:
        paras = sec.header.paragraphs
        if paras:
            p = paras[0]
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            if p.runs:
                p.runs[0].text = HEADER_TEXT
            else:
                p.add_run(HEADER_TEXT)


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def write_cell(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        write_cell(c, h, bold=True)
        shade(c, HEADER_FILL)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            write_cell(cells[i], v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def h(doc, text, level):
    return doc.add_heading(text, level=level)


def para(doc, text, style="Normal"):
    return doc.add_paragraph(text, style=style)


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def figure(doc, png: Path, caption: str, max_w=6.3, max_h=8.4):
    with Image.open(png) as im:
        w, hgt = im.size
    width = max_w
    if width * hgt / w > max_h:
        width = max_h * w / hgt
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(png), width=Inches(width))
    para(doc, caption, "Caption")


REQ4 = ["Req ID", "Requirement", "Priority", "Acceptance criteria"]
REQ4_W = [1.0, 3.6, 0.6, 1.9]


# --------------------------------------------------------------------------- content
def build(doc):
    para(doc, "Business Requirements Document (BRD)", "Title")
    h(doc, "Guideline Value Calculation and GIS Valuation", 1)

    table(doc, ["Field", "Value"], [
        ["Document ID", "BRD-K3-GVC-GIS-001"],
        ["Version", "1"],
        ["Status", "Draft — based on the 01-09-2026 requirement discussion (Guideline value calculation, Valuation Module (CVC) and GIS valuation)"],
        ["Module", "Guideline Value Calculation (#5) and GIS Valuation (#22)"],
        ["Legal basis (primary)", "The Karnataka Stamp Act, 1957 — Secs. 2(ac), 2(mm), 3 + Schedule, 27, 45-A, 45-B"],
        ["State rules (primary)", "Karnataka Stamp (Prevention of Undervaluation of Instruments) Rules, 1977; Karnataka Registration Rules — Rules 13–15 (supporting)"],
        ["Related BRD", "BRD-K3-CVC-GVF-001 — Valuation Module (CVC) — Guidance Value Fixation (upstream owner of guideline rates)"],
        ["Source discussion", "Consolidated_Requirement_Discussions_25082026_to_11092026 — Section 4 (01-09-2026); Document_Registration_requirement_01092026_v2.docx"],
        ["Author (BA)", "Nandha Kumar"],
        ["Product Owner", "M V Prashanth"],
        ["Domain expert / reviewer", "DIGR — Valuation; Kaveri IT Cell; Committee members"],
        ["Target audience", "Kaveri IT Cell, Department of Stamps and Registration, Government of Karnataka; KSRSAC (GIS integration)"],
        ["Last updated", "05-10-2026"],
    ], [1.8, 4.7])

    h(doc, "Document control", 2)
    table(doc, ["Version", "Date", "Author", "Summary of change", "Approver"], [
        ["1", "05-10-2026", "Nandha Kumar",
         "Initial BRD for Guideline Value Calculation and GIS Valuation from the 01-09-2026 discussion: legal mapping (Stamp Act Secs. 2(mm), 3, 27, 45-A, 45-B; Undervaluation Rules 1977; Registration Rules 13–15); As-Is pain points with ServiceDesk evidence; To-Be Process A (guideline value calculation with KSRSAC GIS valuation) and Process B (Sec. 45-A undervaluation reference); functional and non-functional requirements; user-story traceability",
         "Prashanth"],
    ], [0.6, 0.9, 1.0, 3.2, 0.8])

    # 1 ---------------------------------------------------------------------
    h(doc, "1. Executive summary", 2)
    para(doc, "This document presents an assessment of the existing guideline (market) value calculation used at the time of document registration in Kaveri 2.0, and defines the Kaveri 3.0 requirements for guideline value calculation and for GIS-based valuation integrated with the Karnataka State Remote Sensing Applications Centre (KSRSAC).", "List Paragraph")
    para(doc, "Guideline value calculation applies the market value guidelines published by the Central Valuation Committee (CVC) under Sec. 45-B of the Karnataka Stamp Act, 1957 to the property described in the instrument. The higher of the guideline value and the consideration becomes the market value (Sec. 2(mm)), on which ad valorem stamp duty is charged under Sec. 3 and the Schedule. Where the Sub-Registrar believes the market value is not truly set forth, the undervaluation procedure under Sec. 45-A and the Karnataka Stamp (Prevention of Undervaluation of Instruments) Rules, 1977 applies.", "List Paragraph")
    para(doc, "GIS valuation is an implementation layer on top of the CVC guideline rates. The property is located on the KSRSAC cadastral layer and matched spatially to the correct CVC valuation area, road and rural / urban zone, so that the right rate is applied and the valuation is evidenced with a map snapshot.", "List Paragraph")
    para(doc, "The document identifies the current pain points (blank or wrong market value, wrong rural / urban building rate, unit-conversion failures, duty calculator failures, adjudicated value not respected, CVC rates missing in logins, broken revaluation and Sec. 45-A workflows), proposes the To-Be workflows, and defines functional, non-functional and data requirements with traceability to the Acts, Rules and user stories discussed.", "List Paragraph")

    # 2 ---------------------------------------------------------------------
    h(doc, "2. Scope", 2)
    para(doc, "In scope:")
    bullets(doc, [
        "Guideline value calculation for all instruments that require market value — citizen self-valuation on the portal, DEO-assisted valuation at the SRO and Sub-Registrar verification.",
        "Property identification by territorial division and survey description (Registration Rules 13–15): district, taluk, hobli / town, village / ward, survey number + Pot Hissa, city survey / CTS, khata / UPIN / ULPIN, or pick on map.",
        "GIS valuation: KSRSAC parcel lookup, boundary validation, spatial match to CVC valuation area / road / rural-urban zone, map display and stored GIS map snapshot.",
        "Consumption of the Active, versioned CVC guideline rate master effective on the date of execution (rates are owned and published by the CVC Valuation Module — BRD-K3-CVC-GVF-001).",
        "Valuation computation: land, building / structure and apartment components; unit conversion (acre / gunta / cent / hectare / sq.ft / sq.m); itemised breakup.",
        "Market value determination (higher of guideline value and consideration, Sec. 2(mm)) and hand-off to the stamp duty and registration fee engine with the Schedule article.",
        "Sub-Registrar verification: accept, send back for revaluation, or initiate a Sec. 45-A reference.",
        "Sec. 45-A undervaluation reference: SR estimate and communication, Form I and reference to the Deputy Commissioner, hearing and order, suo motu examination, appeal to the Regional Commissioner with 50% deposit, duty difference with interest, refund and release of the pending document.",
        "GIS / valuation master alignment between Kaveri, the CVC rate master and KSRSAC layers (village, road, survey crosswalk).",
        "Notifications, MIS, audit trail, bilingual UI (English + Kannada) and role-based access.",
    ])
    para(doc, "Out of scope (covered elsewhere):")
    bullets(doc, [
        "Estimation, publication and revision of guideline rates by the CVC and sub-committees, and fixation for new individual projects — BRD-K3-CVC-GVF-001.",
        "Stamp duty / registration fee article master, exemptions and concessions — Document Registration fee engine (this BRD supplies the market value and article).",
        "Revision of Deputy Commissioner orders by the Chief Controlling Revenue Authority under Sec. 53-A — to be addressed in the Stamp Act adjudication BRD.",
    ])

    # 3 ---------------------------------------------------------------------
    h(doc, "3. Legal and regulatory reference", 2)
    h(doc, "3.1 Applicable Acts", 3)
    para(doc, "The Department of Stamps and Registration, through the Sub-Registrars (registering officers), the Deputy Commissioners / District Registrars and the Central Valuation Committee, administers the following Act and Rules for guideline value calculation:")
    table(doc, ["Act / Rules", "Description"], [
        ["The Karnataka Stamp Act, 1957",
         "Sec. 2(ac) defines the Central Valuation Committee; Sec. 2(mm) defines market value (open-market price on the date of execution or the consideration, whichever is higher); Sec. 3 + Schedule levy ad valorem duty on market value for listed instruments; Sec. 27 requires the consideration and all facts affecting duty to be fully and truly set forth; Sec. 45-A sets out the undervaluation procedure; Sec. 45-B constitutes the CVC for estimation, publication and revision of market value guidelines. Notifications / amendments: Act 24 of 1999, Act 8 of 2003, Act 7 of 2006, Act 7 of 2007, Act 17 of 2007 — see 3.4"],
        ["Karnataka Stamp (Prevention of Undervaluation of Instruments) Rules, 1977",
         "Operational procedure when Sec. 45-A is invoked — Form I (property particulars and market value), reference to the Deputy Commissioner, inquiry, order and appeal. Notification GSR 81 / RD 73 EST 74; amended by RD 264 MUNOMU 99 (18-08-1999) — see 3.4"],
        ["The Registration Act, 1908 and Karnataka Registration Rules (Rules 13–15) — supporting",
         "Description of property by territorial division and survey / city survey particulars, needed to identify the property and pick the correct guideline rate"],
    ], [2.2, 4.3])

    h(doc, "3.2 Relevant sections — Karnataka Stamp Act, 1957", 3)
    table(doc, ["Section", "Topic", "Relevance", "Refer section 7 for Implementation"], [
        ["Sec. 2(ac)", "Definition — Central Valuation Committee", "CVC (constituted under Sec. 45-B) is the statutory owner of the guideline rates consumed by the calculation", "7.1 step 6 (Active CVC rate version)"],
        ["Sec. 2(mm)", "Definition — Market value", "Price the property would fetch in the open market on the date of execution, or the consideration stated in the instrument, whichever is higher. Proviso: for instruments executed by / on behalf of / in favour of the State or Central Government, a local authority, a statutory authority or a wholly Government-owned body, market value is the consideration set forth (substituted by Act 8 of 2003 — 3.4)", "7.1 steps 6 and 10 (execution-date rate; higher-of rule; Government proviso)"],
        ["Sec. 3 + Schedule", "Instruments chargeable with duty", "Ad valorem duty on market value for listed instruments (conveyance, gift, exchange, settlement, partition, lease, trust and others)", "7.1 step 11 (Schedule article and duty)"],
        ["Sec. 27", "Facts affecting duty to be set forth", "Consideration and all facts affecting chargeability / amount of duty must be fully and truly set forth in the instrument — captured as declared consideration", "7.1 step 9 (consideration capture)"],
        ["Sec. 45-A(1)", "Instruments undervalued — reference", "For instruments in clauses (a)–(o) (see 3.5), if the registering officer has reason to believe, having regard to the CVC guidelines or otherwise, that market value is not truly set forth, he arrives at the estimated market value, communicates it to the parties and, unless duty is paid on that value, keeps registration pending and refers the matter with a copy of the instrument to the Deputy Commissioner (substituted by Act 24 of 1999 — 3.4)", "7.1 steps 13–14; 7.2 steps 1–4"],
        ["Sec. 45-A(2)", "Determination by Deputy Commissioner", "After reasonable opportunity of hearing and inquiry as prescribed, DC determines market value and duty, as far as may be within 90 days of receipt of reference; difference payable with 12% p.a. interest if not paid within 90 days of the order (interest not applicable to instruments executed before 31-03-2006 — Act 7 of 2006)", "7.2 steps 5–7, 11–12"],
        ["Sec. 45-A(3)", "Suo motu examination", "DC may suo motu, within two years from the date of registration, call for and examine any instrument in sub-section (1) not already referred; same procedure and interest as (2); not applicable to instruments registered before the Karnataka Stamp (Amendment) Act, 1975", "7.2 DC start (suo motu) → steps 5–7"],
        ["Sec. 45-A(4)", "Communication of order", "DC order communicated to the person liable to pay duty; copy to the registering officer", "7.2 step 7"],
        ["Sec. 45-A(5)", "Appeal", "Appeal to the Regional Commissioner (Act 17 of 2007); not admitted unless 50% of the duty difference determined by the DC is deposited; deposit refunded if, after appeal or remand, duty borne is sufficient; difference payable with 12% p.a. interest if not paid within 90 days of DC order or 60 days of appellate order (not applicable to instruments executed before 18-08-1999)", "7.2 steps 8–12"],
        ["Sec. 45-B", "Constitution of Central Valuation Committee", "CVC under the IGR & Commissioner of Stamps estimates, publishes and revises market value guidelines for the purpose of Sec. 45-A; final authority on policy, methodology and administration; may constitute district / sub-district market valuation sub-committees", "7.1 step 6 (rates consumed); upstream process — BRD-K3-CVC-GVF-001"],
    ], [0.8, 1.3, 2.9, 1.5])

    h(doc, "3.3 Relevant rules", 3)
    para(doc, "Karnataka Stamp (Prevention of Undervaluation of Instruments) Rules, 1977 — made under the rule-making power for the manner of holding inquiry under Sec. 45-A(2) and (3). Rule numbers below are to be confirmed against the notified Rules text before sign-off.")
    table(doc, ["Rule / provision", "Requirement", "Refer section 7 for Implementation"], [
        ["Reference procedure (Sec. 45-A(1))", "Registering officer refers the undervalued instrument to the Deputy Commissioner with a copy of the instrument and the property particulars", "7.2 step 4"],
        ["Form I", "Statement of property particulars and market value (survey / location, extent, nature, consideration, estimated market value and reasons) to accompany the reference", "7.2 step 4; 3.6 forms mapping"],
        ["Inquiry and order (Sec. 45-A(2)/(3))", "Notice to parties, hearing, inquiry and written order determining market value and duty payable", "7.2 steps 5–7"],
        ["Appeal (Sec. 45-A(5))", "Appeal to the Regional Commissioner within the prescribed time with proof of 50% deposit; hearing and disposal; remand to DC where ordered", "7.2 steps 8–10"],
        ["Amendment RD 264 MUNOMU 99 (18-08-1999)", "Workflow follows the Rules as amended; obsolete provisional / Rule 6 path is not implemented", "7.2 (status model has no provisional path)"],
    ], [1.8, 3.2, 1.5])
    para(doc, "Karnataka Registration Rules (supporting):")
    table(doc, ["Rule", "Requirement", "Refer section 7 for Implementation"], [
        ["Rules 13–15", "Property described by territorial divisions (district, taluk, village / town / ward) and survey particulars (survey number, Pot Hissa / sub-division, city survey number), with extent and boundaries — needed to locate the property and select the correct guideline rate", "7.1 steps 1–2; 8.1"],
    ], [1.2, 3.8, 1.5])

    h(doc, "3.4 Relevant notifications and amendments", 3)
    para(doc, "Gazette notifications and amending Acts cited in the Karnataka Stamp Act, 1957 and the Undervaluation Rules. Each instrument is also cross-referenced on the related section rows in 3.2–3.3.")
    table(doc, ["Instrument", "Date / No.", "Effect", "Refer section 7 for Implementation"], [
        ["Karnataka Stamp (Prevention of Undervaluation of Instruments) Rules, 1977", "GSR 81 / RD 73 EST 74", "Original notification of the Undervaluation Rules (Form I, reference, inquiry, appeal)", "7.2"],
        ["RD 264 MUNOMU 99", "18-08-1999", "Amendments to the Undervaluation Rules; removes the provisional / Rule 6 path", "7.2 (no provisional path)"],
        ["Act No. 24 of 1999", "w.e.f. 18-08-1999", "Substituted Sec. 45-A(1); 90-day DC determination; 50% pre-deposit and refund provisos for appeal", "7.2 steps 4, 6, 9, 11"],
        ["Act No. 8 of 2003", "w.e.f. 01-04-2003", "Substituted Sec. 2(mm) (market value — higher of open-market value and consideration; Government proviso); strengthened Sec. 45-B CVC framework; substituted the third proviso to Sec. 45-A(5)", "7.1 steps 6, 10; 7.2 step 11"],
        ["Act No. 7 of 2006", "w.e.f. 01-04-2006", "12% p.a. interest on duty difference if unpaid within 90 days of DC order; not applicable to instruments executed before 31-03-2006", "7.2 step 11"],
        ["Act No. 7 of 2007", "w.e.f. 01-04-2007", "Added Agreement (Art. 5(f)), Award (Art. 11(a)) and Trust (Art. 54(A)(iii)) to Sec. 45-A(1); Release reference to Art. 45(a)", "3.5; 7.2 step 1"],
        ["Act No. 17 of 2007", "deemed w.e.f. 05-01-2007", "Appellate authority under Sec. 45-A(5) is the Regional Commissioner", "7.2 steps 9–10"],
        ["Act No. 9 of 2009 / Act No. 8 of 2010", "w.e.f. 01-04-2009 / 01-04-2010", "Lease (Art. 30(vi)), Power of Attorney (Art. 41(e)/(ea)/(eb)) and Transferable Development Rights (Art. 20(7)) brought under Sec. 45-A(1)", "3.5; 7.2 step 1"],
        ["CVC market value guideline notifications", "Per revision cycle (gazette effective date)", "Published guideline rates applied from the effective date; consumed as Active rate versions", "7.1 step 6; BRD-K3-CVC-GVF-001"],
    ], [1.9, 1.1, 2.3, 1.2])

    h(doc, "3.5 Instruments covered by Sec. 45-A(1)", 3)
    para(doc, "The undervaluation procedure applies only to the following instruments. The list shall be held as a configurable master linked to the Schedule articles.")
    table(doc, ["Clause", "Instrument", "Article / reference"], [
        ["(a)", "Conveyance", "Sec. 2(1)(d)"],
        ["(b)", "Gift", "Art. 28(a)"],
        ["(c)", "Exchange of property", "Art. 26"],
        ["(d)", "Settlement", "Art. 48-A(I)"],
        ["(e)", "Reconstitution of partnership", "Art. 40-B(a)"],
        ["(f)", "Dissolution of partnership", "Art. 40-C(a)"],
        ["(g)", "Agreement to sell", "Art. 5(e)(i)"],
        ["(h)", "Lease", "Art. 30(vi)"],
        ["(i)", "Power of Attorney", "Art. 41(e), (ea), (eb)"],
        ["(j)", "Release", "Art. 45(a)"],
        ["(k)", "Conveyance under decree or final order of a Civil Court", "—"],
        ["(l)", "Agreement", "Art. 5(f)"],
        ["(m)", "Award", "Art. 11(a)"],
        ["(n)", "Trust", "Art. 54(A)(iii)"],
        ["(o)", "Transferable Development Rights", "Art. 20(7)"],
    ], [0.8, 3.5, 2.2])

    h(doc, "3.6 Statutory forms and outputs mapping", 3)
    table(doc, ["Form / output", "Ref", "Purpose", "Generated by"], [
        ["Valuation summary with GIS map snapshot", "Sec. 2(mm), 3, 27", "Itemised guideline value, consideration, market value, duty and fee with property map (Kaveri 3.0 artefact)", "System on valuation; attached to the application"],
        ["Communication of estimated market value", "Sec. 45-A(1)", "SR's estimated market value and extra duty communicated to the parties", "System on SR estimate (portal + SMS / e-mail)"],
        ["Form I", "Undervaluation Rules, 1977", "Property particulars and market value statement accompanying the reference", "System from valuation record; SR confirms"],
        ["Reference to Deputy Commissioner", "Sec. 45-A(1)", "Electronic reference with Form I and copy of the instrument", "System on SR referral"],
        ["Notice of hearing", "Sec. 45-A(2)/(3); Rules 1977", "Notice to parties for hearing / inquiry", "System on DC action"],
        ["Order determining market value and duty", "Sec. 45-A(2)/(3)/(4)", "Order to the person liable; copy to the registering officer", "System on DC digital signature"],
        ["Appeal memorandum and 50% deposit receipt", "Sec. 45-A(5)", "Appeal to the Regional Commissioner with pre-deposit proof", "Citizen portal + payment gateway"],
        ["Appellate order", "Sec. 45-A(5)", "Value decided or remand to DC", "System on appellate authority digital signature"],
    ], [1.8, 1.3, 2.2, 1.2])

    # 4 ---------------------------------------------------------------------
    h(doc, "4. Stakeholders and actors", 2)
    table(doc, ["Actor", "Description", "Primary goals", "Channel involvement"], [
        ["Citizen / Presentant / Party", "Executant, claimant or person liable to pay duty", "Value the property correctly, see duty and fee before registration, pay any difference", "Online self-valuation; office visit for registration"],
        ["Data Entry Operator (DEO) at SRO", "Office operator assisting walk-in parties", "Enter property and valuation details on behalf of the party", "Assisted (office)"],
        ["Sub-Registrar (SR)", "Registering officer under the Registration Act, 1908", "Verify valuation, send back for revaluation, estimate value and refer under Sec. 45-A", "SR workbench"],
        ["Deputy Commissioner / District Registrar", "Authority under Sec. 45-A(2)/(3) (District Registrar where notified / delegated)", "Hear parties, determine market value and duty, suo motu examination", "DC workbench"],
        ["Regional Commissioner", "Appellate authority under Sec. 45-A(5)", "Hear and decide appeals or remand", "Appellate workbench"],
        ["Central Valuation Committee / DIGR — Valuation", "Owner of guideline rates under Sec. 45-B", "Publish Active rate versions; act on rate-gap and correction requests", "CVC Valuation Module (upstream)"],
        ["KSRSAC", "Karnataka State Remote Sensing Applications Centre — GIS provider", "Provide cadastral parcel, boundary and road layers and spatial lookup services", "System integration"],
        ["Valuation / GIS administrator", "Department master-data owner", "Keep village, road and survey masters aligned between Kaveri, CVC rates and KSRSAC", "Admin console"],
        ["Payment gateway / Treasury", "Fee and duty collection", "Collect duty, extra duty, 50% deposit, duty difference with interest; refunds", "Online"],
        ["Kaveri IT Cell", "Product and operations team", "Operate, monitor and support the module", "Back office"],
    ], [1.5, 1.7, 2.0, 1.3])

    # 5 ---------------------------------------------------------------------
    h(doc, "5. Definitions and glossary", 2)
    table(doc, ["Term", "Definition"], [
        ["Guideline value (guidance value)", "Value computed by applying the CVC market value guideline rates to the property particulars"],
        ["Market value", "Higher of the open-market (guideline-based) value on the date of execution and the consideration (Sec. 2(mm)); for Government-party instruments, the consideration"],
        ["Consideration", "Amount set forth in the instrument as the price / value (Sec. 27)"],
        ["CVC", "Central Valuation Committee constituted under Sec. 45-B, chaired by the IGR & Commissioner of Stamps"],
        ["Active rate version", "The published, gazette-linked CVC guideline rate set effective on a given date"],
        ["Valuation area / zone", "Geographic unit (village, ward, layout, road stretch) to which a CVC rate applies"],
        ["Rural / urban zone", "Classification of the property location (Gram Panchayat vs Town / City Municipal Council, City Corporation, Development Authority) that drives building rates"],
        ["KSRSAC", "Karnataka State Remote Sensing Applications Centre — GIS / cadastral data provider"],
        ["Parcel", "Spatial polygon of a survey / property unit returned by KSRSAC"],
        ["GIS map snapshot", "Stored image with parcel, valuation zone, coordinates, layer version and timestamp used as valuation evidence"],
        ["Pot Hissa", "Sub-division of a survey number"],
        ["City survey / CTS", "Urban survey number system"],
        ["UPIN / ULPIN", "Unique (Land) Parcel Identification Number"],
        ["Gunta / cent", "Land area units: 1 acre = 40 guntas = 100 cents = 4,046.86 sq.m; 1 gunta ≈ 101.17 sq.m; 1 cent ≈ 40.47 sq.m"],
        ["Manual area selection", "Valuation area chosen by the user (with reason) when the GIS parcel is not found or lies outside the selected village; flagged for SR"],
        ["Revaluation (send back)", "SR returns the application to the party to correct property features / valuation"],
        ["Undervaluation reference", "Reference of an instrument to the Deputy Commissioner under Sec. 45-A(1)"],
        ["Adjudicated value", "Market value determined by the DC (or appellate authority) under Sec. 45-A; supersedes the guideline-derived value for that instrument"],
        ["Duty difference", "Additional duty payable on the adjudicated value over duty already paid"],
        ["Suo motu examination", "DC's own-motion examination of a registered instrument within two years (Sec. 45-A(3))"],
    ], [1.8, 4.7])

    # 6 ---------------------------------------------------------------------
    h(doc, "6. Current state", 2)
    para(doc, "Guideline value is calculated in Kaveri 2.0 from CVC rates selected through village / area masters, without spatial validation. Rate data, building-rate zone mapping, unit conversion and the duty calculator are loosely coupled, and the Sec. 45-A undervaluation process is partly manual. The pain points below are evidenced from ServiceDesk tickets and categorized issues discussed on 01-09-2026.")
    h(doc, "6.1 As-Is pain points", 3)
    para(doc, "The Addressed in column maps each item to the To-Be process (7) and to functional / non-functional requirements, risks and fallbacks (8–11).")
    table(doc, ["Sr.No", "Pain Point", "Description", "Source", "Addressed in (this BRD)"], [
        ["1", "Wrong / blank market (guideline) value shown", "Market value blank on sale deed; valuation not showing; agriculture land valuation not showing; wrong market value; market value wrong in EC; market value not in SR login", "Service Desk 29357, 29532, 27923, 31438, 18690, 29931", "7.1 steps 6–10; 8.3 FR-GVC-010–015; 8.4 FR-GVC-024; 8.5 FR-GVC-037; FB-GVC-003"],
        ["2", "Wrong building rate applied for rural properties", "Rural properties calculated on TMC / CMC rates instead of rural rates", "Service Desk 29393; Categorized 93332", "7.1 step 5; 8.2 FR-GIS-003; 8.4 FR-GVC-021; BR-GVC-004"],
        ["3", "Unit-level valuation fails (gunta / area)", "Valuation for 1 gunta not showing; gunta / cent conversion issues", "Service Desk 31788, 24718, 26498, 6097", "8.1 FR-GVC-004; 8.4 FR-GVC-020; BR-GVC-005"],
        ["4", "Duty / fee calculator fails after valuation", "Unable to calculate stamp duty and registration fee; citizen valuates but fee is zero and Save is disabled; fee zero in summary", "Service Desk 23369, 28596, 26292", "7.1 step 11; 8.5 FR-GVC-033; BR-GVC-007; FB-GVC-004"],
        ["5", "Adjudicated Sec. 45-A value not respected", "Summary shows market value higher than the DR-adjudicated value under Sec. 45-A; citizen sees SR value higher than the undervaluation amount", "Service Desk 10531, 31773", "7.2 step 7; 8.5 FR-GVC-035; 8.8 FR-UV-014; BR-UV-005"],
        ["6", "Instrument-specific gaps", "Request for market-value estimation for trust deed; partition / gift fee and duty wrong on articles", "Service Desk 31534; Categorized fee issues 31257, 93306", "3.5; 8.5 FR-GVC-034; 8.7 FR-UV-001"],
        ["7", "CVC rates missing in SR and Citizen login", "\"In CVC module some rates are not displaying in SR Login and Citizen login\"", "Service Desk 31481 (P1), 31521", "8.3 FR-GVC-012, FR-GVC-015; 8.12 FR-GVC-062"],
        ["8", "Cannot correct CVC rates in Citizen login", "\"Correction Rate in CVC Citizen login\"", "Service Desk 21795", "8.3 FR-GVC-013 (correction routed to CVC); BR-GVC-002"],
        ["9", "Generic CVC module failures", "\"CVC issue\"; payment issue after CVC valuation", "Service Desk 15278, 31574 (P1)", "8.5 FR-GVC-033; 9.3 NFR-GVC-AVA-001; FB-GVC-004"],
        ["10", "Revaluation / send-back workflow broken", "Application moved for revaluation then stuck / needs withdrawal; after evaluation, SR cannot send the application further", "Service Desk 23267, 30960, 31484", "7.1 steps 13–15; 7.1.4 status model; 8.6 FR-GVC-042, FR-GVC-043"],
        ["11", "Undervaluation (Sec. 45-A) data / payment stuck", "Undervaluation data fetch failure; unable to pay for undervaluation; pending Sec. 45-A cases in Kaveri-1", "Service Desk 25098, 31071, 27140", "7.2; 8.7–8.9 FR-UV-005, FR-UV-024, FR-UV-026; FB-GVC-005"],
        ["12", "Master misalignment — village, road, survey", "Wrong road names, missing village splits and failed survey fetch lead to blank or wrong market value; no spatial check that the selected rate area matches the property", "Discussion 01-09-2026 (US-GIS-01 to US-GIS-03)", "7.1 steps 3–5; 8.2; 8.10 FR-GIS-020–024"],
    ], [0.4, 1.3, 2.0, 1.2, 1.6])

    # 7 ---------------------------------------------------------------------
    h(doc, "7. Future state (To-Be)", 2)
    para(doc, "The To-Be model has two processes. Process A — Guideline Value Calculation with GIS Valuation — is used for every instrument that requires market value. Process B — Undervaluation Reference under Sec. 45-A — is invoked from Process A when the Sub-Registrar has reason to believe the market value is not truly set forth, or independently when the Deputy Commissioner acts suo motu. Guideline rates are always read from the Active CVC rate version published by the CVC Valuation Module.")

    h(doc, "7.1 Process A — Guideline value calculation with GIS valuation", 3)
    h(doc, "7.1.1 Channel models", 4)
    para(doc, "Valuation can be performed by the citizen on the portal or by a DEO at the SRO on the party's behalf. In both channels the same valuation engine, GIS lookup and CVC rate version are used, and the Sub-Registrar verifies the result.")
    table(doc, ["Service type", "Online activities", "Office activities", "Mode"], [
        ["Citizen self-valuation", "Property identification, GIS parcel lookup, property features, consideration, valuation summary with map, duty and fee", "SR verification; Sec. 45-A reference where required", "Online"],
        ["DEO-assisted valuation", "—", "DEO enters property details on behalf of the party; same GIS lookup and valuation; SR verification", "Office (assisted)"],
        ["Public guideline value search", "View CVC rates and map for an area without starting an application", "—", "Online (no login)"],
    ], [1.5, 2.3, 2.0, 0.7])

    h(doc, "7.1.2 Process diagram", 4)
    figure(doc, DIAGRAMS / "Process_A_Guideline_Value_Calculation.png", "Figure 1: Process A — Guideline Value Calculation with GIS Valuation")
    table(doc, ["#", "Step", "Lane", "Notes"], [
        ["1", "Select district, taluk, hobli / town and ward / village for the property", "Citizen / DEO", "Registration Rule 13 territorial divisions; cascading active masters"],
        ["2", "Enter survey number + Pot Hissa / city survey, or khata / UPIN / ULPIN, or pick the property on the map", "Citizen / DEO", "Registration Rule 15 property description"],
        ["3", "Find the parcel; return boundary (polygon), centre point and nearest road", "KSRSAC GIS", "Spatial lookup service"],
        ["4", "Decision: parcel found and inside the selected village?", "Kaveri System", "If No → correct the details, or choose the valuation area manually with a reason (flagged for SR)"],
        ["5", "Match parcel to CVC valuation area, road and rural / urban zone (spatial match)", "Kaveri System", "GIS valuation layer; prevents neighbouring village / road / urban slab being applied"],
        ["6", "Fetch the Active CVC guideline rate version for the date of execution", "Kaveri System", "Sec. 45-B, Sec. 2(mm); read-only from CVC Valuation Module"],
        ["7", "Select property type and features (agricultural dry / wet / garden; non-agricultural residential / commercial / industrial; building)", "Citizen / DEO", "Drives the rate head; revaluation from SR returns here"],
        ["8", "Calculate guideline value: extent (base unit) × rate + building / structure value", "Kaveri System", "Unit conversion; zone-specific building rate"],
        ["9", "Enter the consideration / value set forth in the instrument", "Citizen / DEO", "Sec. 27"],
        ["10", "Market value = higher of guideline value and consideration", "Kaveri System", "Sec. 2(mm), Sec. 3; Government proviso; adjudicated value overrides where present"],
        ["11", "Calculate stamp duty and registration fee; show summary with GIS map snapshot", "Kaveri System", "Schedule article from the fee engine"],
        ["12", "Review the valuation summary and submit the application", "Citizen / DEO", "Summary is the evidence pack for SR"],
        ["13", "Decision: valuation details correct?", "Sub-Registrar", "If No → send back for revaluation with remarks (returns to step 7)"],
        ["14", "Decision: reason to believe market value not truly set forth?", "Sub-Registrar", "If Yes → go to Process B (Sec. 45-A undervaluation)"],
        ["15", "Accept valuation; continue to payment and registration", "Sub-Registrar", "End of Process A"],
    ], [0.4, 2.6, 1.1, 2.4])
    para(doc, "Key characteristics: territorial and survey-based identification; KSRSAC parcel lookup with boundary check; manual valuation-area fallback only with reason and SR flag; spatial match to CVC area, road and rural / urban zone; Active CVC rate version by execution date; itemised valuation; higher-of rule with Government proviso; duty and fee shown with GIS map snapshot before submission; SR accept / send back / refer.")

    h(doc, "7.1.3 Valuation calculation logic", 4)
    table(doc, ["Component", "Rule", "Reference"], [
        ["Extent normalisation", "All extents converted to the CVC notified base unit (sq.m / acre) through a conversion master — acre, gunta, ana, cent, hectare, sq.ft, sq.m; small extents (e.g. 1 gunta) supported without rounding to zero", "FR-GVC-004, FR-GVC-020"],
        ["Land value", "Normalised extent × CVC land rate for the matched valuation area and property type", "Sec. 45-B rates; FR-GVC-020"],
        ["Building / structure value", "Built-up area × CVC building rate for the zone category (rural / TMC / CMC / Corporation) and construction type, as per CVC methodology", "FR-GVC-021"],
        ["Apartment / flat", "Undivided share of land and built-up area valued per the CVC apartment methodology", "FR-GVC-022"],
        ["Additions / deductions", "Only those notified in the Active rate version (e.g. corner site, main-road frontage)", "FR-GVC-023"],
        ["Guideline value", "Sum of the components above, rounded per department rule [rounding rule TBD]", "FR-GVC-024, FR-GVC-025"],
        ["Market value", "max(guideline value, consideration); Government-party instruments: consideration; adjudicated Sec. 45-A value: overrides", "Sec. 2(mm); FR-GVC-031, 032, 035"],
        ["Duty and fee", "Market value passed with the Schedule article to the fee engine", "Sec. 3 + Schedule; FR-GVC-033"],
    ], [1.5, 3.7, 1.3])

    h(doc, "7.1.4 Valuation status model", 4)
    table(doc, ["Status", "Description", "Actor", "Next states"], [
        ["Draft valuation", "Valuation started, not submitted", "Citizen / DEO", "Property identified"],
        ["Property identified", "Territorial division and survey / identifier captured", "Citizen / DEO", "Parcel matched / Manual area selected"],
        ["Parcel matched", "KSRSAC parcel found inside the selected village and matched to a CVC area / zone", "System", "Rate fetched"],
        ["Manual area selected", "Parcel not found or outside the village; user selected area with reason; flagged for SR", "Citizen / DEO", "Rate fetched"],
        ["Rate fetched", "Active CVC rate version for the execution date retrieved", "System", "Valuation computed / Rate not available"],
        ["Rate not available", "No rate for area / property type; rate-gap logged to CVC", "System", "Rate fetched (after CVC action) / SR-assessed"],
        ["Valuation computed", "Guideline value computed with breakup", "System", "Market value determined"],
        ["Market value determined", "Consideration captured; market value and duty / fee computed", "System", "Submitted"],
        ["Submitted", "Valuation summary submitted with the application", "Citizen / DEO", "Pending SR verification"],
        ["Pending SR verification", "Awaiting SR scrutiny", "SR", "Accepted / Sent back for revaluation / Referred under Sec. 45-A"],
        ["Sent back for revaluation", "Returned with remarks; data retained; SR may withdraw before citizen acts", "SR", "Valuation computed (after correction) / Pending SR verification (withdrawn)"],
        ["Referred under Sec. 45-A", "Process B started; registration pending", "SR", "Accepted (after duty paid) / Adjudicated"],
        ["Adjudicated", "DC / appellate value locked for this instrument", "DC / Regional Commissioner", "Accepted"],
        ["Accepted", "Valuation accepted; proceeds to payment and registration", "SR", "Locked"],
        ["Locked", "Valuation frozen on registration; persisted to document, index and EC", "System", "—"],
    ], [1.5, 2.8, 1.0, 1.2])

    h(doc, "7.2 Process B — Undervaluation reference under Sec. 45-A", 3)
    h(doc, "7.2.1 Trigger model", 4)
    table(doc, ["Trigger", "Initiated by", "Statutory basis", "Entry step"], [
        ["Reference during registration", "Sub-Registrar from Process A step 14", "Sec. 45-A(1)", "Step 1"],
        ["Suo motu examination", "Deputy Commissioner, within two years from the date of registration, for an instrument not already referred", "Sec. 45-A(3)", "DC start → step 5"],
    ], [1.6, 2.4, 1.2, 1.3])
    h(doc, "7.2.2 Process diagram", 4)
    figure(doc, DIAGRAMS / "Process_B_Sec45A_Undervaluation.png", "Figure 2: Process B — Undervaluation Reference under Section 45-A")
    table(doc, ["#", "Step", "Lane", "Notes"], [
        ["1", "Arrive at the estimated market value (CVC guidelines or otherwise) and record reasons", "Sub-Registrar", "Sec. 45-A(1); instrument must be in 3.5"],
        ["2", "Communicate estimated value and extra duty to the parties (SMS / e-mail / portal)", "Kaveri System", "Sec. 45-A(1)"],
        ["3", "Decision: pay duty on the SR's estimated value?", "Party", "Yes → collect extra duty; continue registration (end)"],
        ["4", "Keep registration pending; prepare Form I and refer to the DC with a copy of the instrument", "Sub-Registrar", "Sec. 45-A(1); 1977 Rules"],
        ["5", "Issue notice, give the parties a hearing and hold the inquiry", "Deputy Commissioner", "Sec. 45-A(2); 1977 Rules. Suo motu cases (Sec. 45-A(3)) enter here"],
        ["6", "Order market value and duty payable (as far as may be within 90 days)", "Deputy Commissioner", "Sec. 45-A(2)"],
        ["7", "Send order to the party and a copy to the SR; lock adjudicated value as the market value", "Kaveri System", "Sec. 45-A(4)"],
        ["8", "Decision: appeal against the order?", "Party", "No → step 11"],
        ["9", "Deposit 50% of the duty difference and file the appeal", "Party", "Sec. 45-A(5) proviso"],
        ["10", "Hear and decide the appeal (value decided, or remand: DC determines value again)", "Regional Commissioner", "Sec. 45-A(5); remand returns to step 6"],
        ["11", "Calculate duty difference + 12% p.a. interest if paid late; refund deposit if duty sufficient", "Kaveri System", "Sec. 45-A(2), (5); interest cut-off dates"],
        ["12", "Pay the duty difference (and interest, if any)", "Party", "Payment gateway"],
        ["13", "Release the pending document and complete registration (or close the suo motu case)", "Sub-Registrar", "End of Process B"],
    ], [0.4, 2.6, 1.2, 2.3])
    para(doc, "Key characteristics: only Sec. 45-A(1) instruments; SR estimate with reasons and electronic communication; pay-now option; Form I auto-filled from the valuation record; electronic reference to the DC with 90-day ageing; suo motu window of two years; DC order with digital signature and automatic copy to SR; adjudicated value locked and reflected in summary and payment; appeal admitted only on 50% deposit; interest and refund computed by rule; pending document released only after payment.")

    h(doc, "7.2.3 Undervaluation case status model", 4)
    table(doc, ["Status", "Description", "Actor", "Next states"], [
        ["SR estimate recorded", "Estimated market value and reasons recorded", "SR", "Communicated to party"],
        ["Communicated to party", "Estimated value and extra duty sent", "System", "Extra duty paid / Referred to DC"],
        ["Extra duty paid", "Party paid duty on SR estimate", "Party", "Closed — registration continues"],
        ["Referred to DC", "Form I and instrument copy sent; registration pending", "SR", "Hearing scheduled"],
        ["Suo motu initiated", "DC called for a registered instrument (≤ 2 years)", "DC", "Hearing scheduled"],
        ["Hearing scheduled", "Notice issued to parties", "DC", "Under inquiry"],
        ["Under inquiry", "Hearing / inquiry in progress; 90-day ageing tracked", "DC", "Order issued"],
        ["Order issued", "Market value and duty determined; order communicated; adjudicated value locked", "DC", "Under appeal / Duty difference payable"],
        ["Under appeal", "Appeal filed with 50% deposit", "Party", "Appeal decided / Remanded"],
        ["Remanded", "Returned to DC for fresh determination", "Regional Commissioner", "Under inquiry"],
        ["Appeal decided", "Value decided by appellate authority", "Regional Commissioner", "Duty difference payable / Refund due"],
        ["Duty difference payable", "Difference + interest (if any) computed", "System", "Paid"],
        ["Refund due", "Duty borne found sufficient; deposit to be refunded", "System", "Closed"],
        ["Paid", "Difference and interest paid", "Party", "Closed"],
        ["Closed", "Document released / suo motu case closed", "SR / System", "—"],
    ], [1.5, 2.8, 1.1, 1.1])

    h(doc, "7.3 What is new in Kaveri 3.0", 3)
    para(doc, "Material enhancements compared with the Kaveri 2.0 valuation and undervaluation features (6). 7.3.1 maps each As-Is pain point from 6.1 to the Kaveri 3.0 closure.")
    table(doc, ["#", "Capability", "What is new in Kaveri 3.0"], [
        ["1", "GIS valuation (KSRSAC)", "Parcel lookup, boundary check and spatial match to CVC area / road / rural-urban zone; map display and stored map snapshot"],
        ["2", "Single source of rates", "Read-only consumption of Active, versioned CVC rates by execution date — identical in Citizen, DEO and SR logins"],
        ["3", "Transparent valuation", "Itemised breakup with rate version, unit conversion and higher-of explanation before submission"],
        ["4", "Reliable duty hand-off", "Market value passed with Schedule article; Save gated on a valid computed duty and fee"],
        ["5", "Adjudicated value respected", "Sec. 45-A adjudicated value locked and used in summary, payment, document and EC"],
        ["6", "Digitised Sec. 45-A", "End-to-end reference, Form I, DC hearing and order, suo motu, appeal with 50% deposit, interest and refund"],
        ["7", "Revaluation workflow", "Send back with remarks, withdraw, resubmit — no stuck applications"],
        ["8", "Master alignment", "Village / road / survey crosswalk between Kaveri, CVC and KSRSAC with effective-dated splits and reconciliation reports"],
        ["9", "Public guideline value search", "Citizens can see rates and map for an area before applying"],
    ], [0.4, 1.8, 4.3])
    h(doc, "7.3.1 Rectified As-Is pain points", 4)
    table(doc, ["Sr.No", "Pain Point (As-Is)", "How rectified in Kaveri 3.0"], [
        ["1", "Wrong / blank market value", "Rates read from the Active CVC version; rate-gap detection instead of blank value; market value persisted to document, index and EC"],
        ["2", "Wrong building rate for rural properties", "GIS-derived rural / urban zone drives the building rate; rural parcels cannot receive TMC / CMC rates"],
        ["3", "Unit-level valuation fails", "Conversion master and base-unit calculation; 1 gunta and other small extents supported"],
        ["4", "Duty / fee calculator fails", "Market value and article passed to the fee engine; Save enabled only on valid computed duty / fee; retry on fee-engine failure"],
        ["5", "Adjudicated value not respected", "Adjudicated value overrides guideline value in summary and payment"],
        ["6", "Instrument-specific gaps", "Valuation basis per Schedule article including trust, partition, gift; Sec. 45-A instrument master"],
        ["7", "CVC rates missing in logins", "Single rate service for all logins; missing-rate alerts and reconciliation report"],
        ["8", "Cannot correct CVC rates in Citizen login", "Rates are read-only; correction request routed to CVC with evidence"],
        ["9", "Generic CVC module / payment failures", "Resilient integration with clear errors, resumable state and payment polling"],
        ["10", "Revaluation / send-back broken", "Explicit status model with send back, withdraw and atomic transitions"],
        ["11", "Sec. 45-A data / payment stuck", "Digitised case workflow, online payment of difference, migration of pending Kaveri-1 cases"],
        ["12", "Master misalignment", "Crosswalk masters, effective-dated village splits, road-name correction workflow, KSRSAC sync"],
    ], [0.5, 2.0, 4.0])

    # 8 ---------------------------------------------------------------------
    h(doc, "8. Functional requirements", 2)
    para(doc, "Functional requirements are aligned with section 7 (To-Be). Section 7 is the process authority; section 8 states testable requirements (FR-GVC-* for guideline value calculation, FR-GIS-* for GIS valuation, FR-UV-* for Sec. 45-A undervaluation) and cites the related step where applicable.")

    h(doc, "8.1 Property identification and jurisdiction", 3)
    table(doc, REQ4, [
        ["FR-GVC-001", "System shall capture property location through cascading active masters: district → taluk → hobli / town → village / ward (Registration Rule 13) and derive the jurisdictional SRO", "Must", "Only active masters selectable; SRO derived automatically; 7.1 step 1"],
        ["FR-GVC-002", "System shall capture the property identifier: survey number with Pot Hissa / sub-division (rural), city survey / CTS / municipal number (urban), khata / UPIN / ULPIN, or pick on map (Registration Rule 15)", "Must", "At least one identifier mandatory; 7.1 step 2"],
        ["FR-GVC-003", "System shall capture property type — agricultural (dry / wet / garden) or non-agricultural (residential / commercial / industrial; site / flat / building) — mapped to CVC rate heads", "Must", "Property type determines rate head; 7.1 step 7"],
        ["FR-GVC-004", "System shall capture extent with unit (acre-gunta-ana, cent, hectare, sq.ft, sq.m) and convert to the base unit through a configurable conversion master", "Must", "1 acre = 40 guntas = 4,046.86 sq.m; 1 gunta valuation returns a non-zero value (31788)"],
        ["FR-GVC-005", "System shall capture building particulars — built-up area per floor, construction type, year of construction and CVC-notified amenities", "Must", "Fields shown only for property types with structures"],
        ["FR-GVC-006", "System should prefill survey extent and owner particulars from integrated land / property records (RTC, urban property records) where available [integration list TBD]", "Should", "Prefilled values editable with audit; failed fetch does not block manual entry"],
    ], REQ4_W)

    h(doc, "8.2 GIS parcel lookup and spatial rate matching", 3)
    table(doc, REQ4, [
        ["FR-GIS-001", "System shall send the property identifiers to the KSRSAC GIS service and receive the parcel polygon, centre point, nearest road and administrative boundary", "Must", "Response within NFR-GVC-PERF-002; 7.1 step 3"],
        ["FR-GIS-002", "System shall validate that the parcel lies inside the selected village / ward; if not found or outside, the user shall correct the details or choose the valuation area manually with a mandatory reason, and the application shall be flagged for SR", "Must", "Manual selection requires reason; flag visible on SR workbench; 7.1 step 4"],
        ["FR-GIS-003", "System shall spatially match the parcel to the CVC valuation area, road and rural / urban zone (Gram Panchayat, TMC, CMC, City Corporation, Development Authority)", "Must", "Rural parcel never receives TMC / CMC building rate (29393, 93332); 7.1 step 5"],
        ["FR-GIS-004", "Where a parcel spans more than one valuation area, system shall apply the CVC-approved rule [rule TBD by CVC] and display the areas involved", "Should", "Rule configurable; all intersecting areas shown"],
        ["FR-GIS-005", "System shall display an interactive map with the parcel, valuation area boundary and rate label to citizen, DEO and SR", "Must", "Map visible in valuation and SR verification screens"],
        ["FR-GIS-006", "System shall store a GIS map snapshot (image, coordinates, KSRSAC layer version, timestamp) with the valuation and print it on the valuation summary", "Must", "Snapshot retrievable for the life of the record; 7.1 step 11"],
        ["FR-GIS-007", "System should allow the user to pick the parcel on the map and derive the village and survey number from the selection", "Should", "Map pick fills identifiers; still validated by FR-GIS-002"],
        ["FR-GIS-008", "System shall record the KSRSAC request / response identifier and layer version for every lookup", "Must", "Audit query by application number"],
    ], REQ4_W)

    h(doc, "8.3 CVC guideline rate consumption", 3)
    table(doc, REQ4, [
        ["FR-GVC-010", "System shall read guideline rates only from the published, Active, versioned rate master of the CVC Valuation Module (BRD-K3-CVC-GVF-001); no local rate tables or overrides", "Must", "Single rate service; Sec. 45-B"],
        ["FR-GVC-011", "System shall select the rate version effective on the date of execution of the instrument, including historic versions for back-dated execution", "Must", "Execution date before the latest effective date picks the earlier version; Sec. 2(mm)"],
        ["FR-GVC-012", "The same Active rates shall be returned to Citizen, DEO and SR logins; any rate not resolvable shall raise an exception alert to the valuation administrator", "Must", "No login-specific divergence (31481, 31521); alert logged"],
        ["FR-GVC-013", "Rates shall be read-only in Citizen, DEO and SR logins; a rate correction request with evidence shall be routed to the CVC / DR workflow", "Must", "No edit controls (21795); request tracked to closure in CVC module"],
        ["FR-GVC-014", "System should provide a public guideline value search by area with map, without login", "Should", "Shows Active version and effective date"],
        ["FR-GVC-015", "If no rate exists for the matched area / property type, system shall block automatic valuation, display \"Rate not available\", log a rate-gap request to CVC / DR, and allow SR-assessed value with reason [policy to be confirmed by DIGR — Valuation]", "Must", "No blank or zero market value shown; rate-gap listed in FR-GVC-062"],
    ], REQ4_W)

    h(doc, "8.4 Valuation computation", 3)
    table(doc, REQ4, [
        ["FR-GVC-020", "System shall compute land value as normalised extent × applicable CVC land rate in the notified unit", "Must", "Unit tests for acre, gunta, cent, sq.ft, sq.m"],
        ["FR-GVC-021", "System shall compute building / structure value as built-up area × CVC building rate for the zone category and construction type, per CVC methodology", "Must", "Zone category from FR-GIS-003"],
        ["FR-GVC-022", "System should value apartments / flats using undivided share of land and built-up area per CVC apartment methodology", "Should", "Configurable per CVC notification"],
        ["FR-GVC-023", "System should apply additions / deductions only where notified in the Active rate version", "Should", "No unnotified adjustment applied"],
        ["FR-GVC-024", "System shall display an itemised valuation breakup — each component, rate, unit, extent and rate version", "Must", "Breakup shown on summary and SR screen"],
        ["FR-GVC-025", "System should round values per the department rounding rule [TBD]", "Should", "Configurable"],
        ["FR-GVC-026", "System shall recompute the valuation when any property input changes before submission and invalidate stale results", "Must", "Summary never shows outdated value"],
    ], REQ4_W)

    h(doc, "8.5 Market value determination and duty hand-off", 3)
    table(doc, REQ4, [
        ["FR-GVC-030", "System shall capture the consideration set forth in the instrument (Sec. 27)", "Must", "Mandatory for instruments with consideration; 7.1 step 9"],
        ["FR-GVC-031", "System shall set market value as the higher of guideline value and consideration and display both with the basis (Sec. 2(mm))", "Must", "Basis label shown; 7.1 step 10"],
        ["FR-GVC-032", "For instruments executed by / for the State or Central Government, local authority, statutory authority or wholly Government-owned body, system shall set market value to the consideration on party-category flag with SR confirmation (Sec. 2(mm) proviso)", "Must", "Flag audited; SR confirmation recorded"],
        ["FR-GVC-033", "System shall pass market value and the selected Schedule article to the stamp duty and registration fee engine and display duty and fee; Save / Proceed shall be enabled only when a valid duty and fee (or valid exemption) is returned", "Must", "No zero fee with Save disabled (23369, 28596, 26292); 7.1 step 11"],
        ["FR-GVC-034", "System shall apply the valuation basis configured for each Schedule article (conveyance, gift, partition, settlement, exchange, trust, lease, power of attorney, etc.)", "Must", "Trust deed valuation supported (31534); partition / gift duty correct (31257, 93306)"],
        ["FR-GVC-035", "Where an adjudicated Sec. 45-A value exists for the instrument, it shall supersede the guideline-derived value in summary and payment", "Must", "Summary equals adjudicated value (10531, 31773)"],
        ["FR-GVC-036", "System shall generate a valuation summary (PDF) with property, GIS snapshot, rate version, breakup, consideration, market value, duty and fee", "Must", "Downloadable by citizen and SR; 7.1 step 12"],
        ["FR-GVC-037", "Market value shall be persisted to the registered document record, index and Encumbrance Certificate data", "Should", "EC shows the registered market value (18690)"],
    ], REQ4_W)

    h(doc, "8.6 SR verification and revaluation", 3)
    table(doc, REQ4, [
        ["FR-GVC-040", "SR workbench shall display the valuation summary, GIS map, manual-area flags, rate version and breakup", "Must", "All items visible in one view; 7.1 step 13"],
        ["FR-GVC-041", "SR shall Accept, Send back for revaluation with remarks, or Initiate Sec. 45-A reference", "Must", "Decision and remarks audited"],
        ["FR-GVC-042", "Send back shall return the application to the property-features step with all data retained; SR may withdraw a send-back before the party acts", "Must", "No stuck application (23267)"],
        ["FR-GVC-043", "After SR acceptance, the application shall move to payment and the next registration stage in a single atomic transition", "Must", "No case where SR cannot forward after evaluation (30960, 31484)"],
        ["FR-GVC-044", "System shall retain every valuation version with actor, timestamp and reason", "Must", "History viewable by SR and audit roles"],
    ], REQ4_W)

    h(doc, "8.7 Sec. 45-A reference — Sub-Registrar", 3)
    table(doc, REQ4, [
        ["FR-UV-001", "Sec. 45-A reference shall be available only for instruments in the configurable Sec. 45-A(1) master (clauses (a)–(o), 3.5)", "Must", "Option hidden for other instruments"],
        ["FR-UV-002", "SR shall record the estimated market value and reasons; system computes extra duty", "Must", "Reasons mandatory; 7.2 step 1"],
        ["FR-UV-003", "System shall communicate the estimated value and extra duty to the parties via portal, SMS and e-mail", "Must", "Delivery logged; 7.2 step 2"],
        ["FR-UV-004", "Party may pay the extra duty online, after which registration continues", "Must", "Payment closes the case; 7.2 step 3"],
        ["FR-UV-005", "If the party does not pay, SR shall refer the case electronically to the DC with Form I auto-filled from the valuation record and a copy of the instrument; the document is marked Pending — Sec. 45-A", "Must", "Form I generated; reference visible in DC queue (25098); 7.2 step 4"],
        ["FR-UV-006", "A document pending under Sec. 45-A shall not be completed or released until the DC order is complied with", "Must", "Release action blocked"],
    ], REQ4_W)

    h(doc, "8.8 Deputy Commissioner determination and suo motu", 3)
    table(doc, REQ4, [
        ["FR-UV-010", "DC workbench shall list references with Form I, instrument copy, valuation summary and ageing against 90 days", "Must", "Overdue cases highlighted"],
        ["FR-UV-011", "DC shall issue notice of hearing, schedule hearings and record proceedings and evidence", "Must", "Notice sent via FR-GVC-051; 7.2 step 5"],
        ["FR-UV-012", "DC shall pass an order determining market value and duty, digitally signed; the order is communicated to the party and a copy sent to the SR (Sec. 45-A(4))", "Must", "Order PDF in party and SR views; 7.2 steps 6–7"],
        ["FR-UV-013", "DC may initiate suo motu examination of a registered instrument within two years from the date of registration, if not already referred; system enforces the window and excludes instruments registered before the 1975 amendment", "Must", "Outside window blocked (Sec. 45-A(3))"],
        ["FR-UV-014", "The adjudicated value shall be locked on the valuation record and shown to SR and citizen", "Must", "Feeds FR-GVC-035"],
    ], REQ4_W)

    h(doc, "8.9 Appeal, duty difference, interest and refund", 3)
    table(doc, REQ4, [
        ["FR-UV-020", "Party may file an appeal to the Regional Commissioner online; the appeal is admitted only after deposit of 50% of the duty difference", "Must", "Appeal blocked without deposit; 7.2 step 9"],
        ["FR-UV-021", "Appellate authority shall decide the value or remand to the DC", "Must", "Remand returns case to Under inquiry; 7.2 step 10"],
        ["FR-UV-022", "System shall compute the duty difference and 12% p.a. interest when unpaid beyond 90 days of the DC order or 60 days of the appellate order, applying the interest cut-off dates (DC order: instruments executed before 31-03-2006 exempt; appeal: before 18-08-1999 exempt)", "Must", "Interest calculation unit-tested for each cut-off; 7.2 step 11"],
        ["FR-UV-023", "System shall refund the deposit where the duty borne is found sufficient after appeal or remand", "Must", "Refund instruction generated"],
        ["FR-UV-024", "Party shall pay the difference and interest online; on success the SR releases the pending document or closes the suo motu case", "Must", "Undervaluation payment succeeds (31071); 7.2 steps 12–13"],
        ["FR-UV-025", "The workflow shall follow the Undervaluation Rules as amended by RD 264 MUNOMU 99; no provisional / Rule 6 path", "Must", "No provisional status present"],
        ["FR-UV-026", "Pending Sec. 45-A cases from Kaveri 1.0 / 2.0 shall be migrated with their current status, orders and payments", "Must", "Migrated cases actionable (27140)"],
    ], REQ4_W)

    h(doc, "8.10 GIS and valuation master alignment", 3)
    table(doc, REQ4, [
        ["FR-GIS-020", "System shall maintain crosswalk masters linking Kaveri village / ward / road / survey codes, CVC valuation areas and KSRSAC layer identifiers", "Must", "Every active CVC area mapped or listed as exception"],
        ["FR-GIS-021", "Village splits, merges and renames shall be effective-dated with history retained", "Must", "Valuation by execution date uses the correct master"],
        ["FR-GIS-022", "Valuation administrator should correct road names through a maker-checker workflow", "Should", "Change audited"],
        ["FR-GIS-023", "System shall produce reconciliation reports: villages without GIS polygon, CVC areas without rates, mismatched roads, failed survey fetches", "Must", "Scheduled and on demand"],
        ["FR-GIS-024", "System shall synchronise KSRSAC layers on a schedule and record the layer version", "Must", "Version visible in FR-GIS-008"],
    ], REQ4_W)

    h(doc, "8.11 Notifications", 3)
    table(doc, REQ4, [
        ["FR-GVC-050", "SMS / e-mail on valuation sent back for revaluation and on acceptance", "Should", "Bilingual EN / KN templates"],
        ["FR-GVC-051", "SMS / e-mail / portal on Sec. 45-A events: estimated value communication, reference to DC, hearing notice, order, appeal decision, remand, payment due and refund", "Must", "Delivery logged"],
        ["FR-GVC-052", "Reminder before the 90-day / 60-day interest threshold", "Should", "Sent at configurable days before threshold"],
    ], REQ4_W)

    h(doc, "8.12 Reports and MIS", 3)
    table(doc, REQ4, [
        ["FR-GVC-060", "Valuation exceptions report — manual area selections, SR-assessed values, overrides — by office and period", "Must", "Drill-down to application"],
        ["FR-GVC-061", "Send-back / revaluation statistics by SR office and reason", "Should", ""],
        ["FR-GVC-062", "Rate-gap and missing-rate report for CVC / DR", "Must", "Feeds CVC action list"],
        ["FR-GVC-063", "Sec. 45-A pendency report with ageing against 90 days, by DC / district", "Must", "Overdue highlighted"],
        ["FR-GVC-064", "Revenue impact of Sec. 45-A — extra duty collected, difference, interest, refunds", "Must", "Reconciles with payment gateway"],
        ["FR-GVC-065", "GIS match-rate dashboard — parcel found, outside village, not found — by district", "Should", ""],
    ], REQ4_W)

    h(doc, "8.13 Business rules", 3)
    table(doc, ["Rule ID", "Description", "Statutory ref", "System enforcement"], [
        ["BR-GVC-001", "Guideline rates come only from the Active CVC rate version", "Sec. 45-B", "Read-only service; FR-GVC-010"],
        ["BR-GVC-002", "Rates cannot be edited in Citizen, DEO or SR logins", "Sec. 45-B(2)", "No edit; FR-GVC-013"],
        ["BR-GVC-003", "Rate version is selected by the date of execution", "Sec. 2(mm)", "FR-GVC-011"],
        ["BR-GVC-004", "Building rate follows the GIS-derived rural / urban zone", "CVC methodology", "FR-GIS-003, FR-GVC-021"],
        ["BR-GVC-005", "Extent is converted to the base unit before rate application", "—", "FR-GVC-004, FR-GVC-020"],
        ["BR-GVC-006", "Market value = higher of guideline value and consideration, except Government-party instruments", "Sec. 2(mm)", "FR-GVC-031, FR-GVC-032"],
        ["BR-GVC-007", "Application cannot be saved / submitted without a valid computed duty and fee", "Sec. 3 + Schedule", "FR-GVC-033"],
        ["BR-GVC-008", "Manual valuation-area selection requires a reason and is flagged for SR", "Registration Rules 13–15", "FR-GIS-002"],
        ["BR-GVC-009", "Valuation is locked on registration", "—", "Status Locked; 7.1.4"],
        ["BR-UV-001", "Sec. 45-A reference only for instruments in Sec. 45-A(1)", "Sec. 45-A(1)", "FR-UV-001"],
        ["BR-UV-002", "SR estimate must record reasons", "Sec. 45-A(1)", "FR-UV-002"],
        ["BR-UV-003", "Registration stays pending until duty on the estimate is paid or the DC order is complied with", "Sec. 45-A(1)", "FR-UV-006"],
        ["BR-UV-004", "Suo motu only within two years of registration and for instruments not already referred", "Sec. 45-A(3)", "FR-UV-013"],
        ["BR-UV-005", "Adjudicated value supersedes guideline value for that instrument", "Sec. 45-A(2)", "FR-GVC-035, FR-UV-014"],
        ["BR-UV-006", "Appeal admitted only after 50% deposit of the duty difference", "Sec. 45-A(5)", "FR-UV-020"],
        ["BR-UV-007", "Interest at 12% p.a. beyond 90 days (DC order) / 60 days (appellate order), subject to cut-off dates", "Sec. 45-A(2), (3), (5)", "FR-UV-022"],
        ["BR-UV-008", "No provisional / Rule 6 path", "RD 264 MUNOMU 99", "FR-UV-025"],
    ], [1.0, 2.9, 1.2, 1.4])

    h(doc, "8.14 User interface (high-level)", 3)
    table(doc, ["Screen / step", "Purpose", "Actor", "Statutory alignment", "Notes"], [
        ["Property location", "Cascading territorial selection", "Citizen / DEO", "Registration Rule 13", "FR-GVC-001"],
        ["Property identifier / map pick", "Survey / city survey / khata / UPIN or map pick", "Citizen / DEO", "Registration Rule 15", "FR-GVC-002, FR-GIS-007"],
        ["GIS result and map", "Parcel, valuation area, rate label; manual-area fallback with reason", "Citizen / DEO", "—", "FR-GIS-001–006"],
        ["Property type and features", "Rate head, extent, building particulars", "Citizen / DEO", "CVC methodology", "FR-GVC-003–005"],
        ["Consideration", "Declared consideration", "Citizen / DEO", "Sec. 27", "FR-GVC-030"],
        ["Valuation summary", "Breakup, market value basis, duty, fee, map snapshot", "Citizen / DEO", "Sec. 2(mm), 3", "FR-GVC-024, 031, 033, 036"],
        ["SR valuation verification", "Accept / send back / refer under Sec. 45-A", "SR", "Sec. 45-A(1)", "FR-GVC-040–043"],
        ["Sec. 45-A estimate and reference", "Estimate with reasons, Form I, refer to DC", "SR", "Sec. 45-A(1); Rules 1977", "FR-UV-002–005"],
        ["DC case workbench", "Hearing, inquiry, order, suo motu", "DC", "Sec. 45-A(2)–(4)", "FR-UV-010–014"],
        ["Appeal", "File appeal with 50% deposit; appellate decision", "Party / Regional Commissioner", "Sec. 45-A(5)", "FR-UV-020–021"],
        ["Duty difference payment", "Pay difference + interest; refund status", "Party", "Sec. 45-A(2), (5)", "FR-UV-022–024"],
        ["Public guideline value search", "Rates and map by area", "Public", "Sec. 45-B", "FR-GVC-014"],
        ["Master alignment console", "Crosswalk, splits, road names, reconciliation", "Valuation / GIS admin", "—", "FR-GIS-020–024"],
    ], [1.4, 1.9, 1.0, 1.1, 1.1])
    para(doc, "Wireframe links: [Figma / prototype URLs]")
    para(doc, "Bilingual: all labels [EN / KN] — content manager sign-off.")

    h(doc, "8.15 Integrations", 3)
    table(doc, ["Integration", "Direction", "Purpose", "Owner", "Status"], [
        ["KSRSAC GIS services", "Outbound / Inbound", "Parcel lookup, boundary, nearest road, admin boundary, map tiles, layer sync", "KSRSAC / Kaveri IT Cell", "TBD"],
        ["CVC Valuation Module", "Inbound", "Active versioned guideline rates; rate-gap and correction requests outbound", "CVC / Kaveri IT Cell", "Internal"],
        ["Stamp duty and registration fee engine", "Internal", "Duty and fee from market value and Schedule article", "Document Registration", "Internal"],
        ["Land / property records (RTC, urban property records)", "Inbound", "Prefill survey extent and particulars [list TBD]", "Revenue / Urban Development", "TBD"],
        ["Payment gateway / Treasury", "Outbound", "Duty, extra duty, 50% deposit, difference + interest, refunds", "Finance", "TBD"],
        ["SMS / e-mail gateway", "Outbound", "Notifications (8.11)", "Platform", "Existing"],
        ["DSC / signing service", "Outbound", "DC and appellate orders", "Platform", "Existing"],
        ["Kaveri master data", "Inbound", "District, taluk, hobli, village, ward, SRO, Schedule articles", "Platform", "Existing"],
    ], [1.6, 0.9, 2.3, 1.1, 0.6])
    para(doc, "Interface requirements: [API list TBD by Architect — KSRSAC lookup / layer sync, CVC rate service, fee engine, payment and refund APIs]")

    h(doc, "8.16 Data requirements", 3)
    h(doc, "8.16.1 Core entities (logical)", 4)
    bullets(doc, [
        "PropertyLocation (district, taluk, hobli / town, village / ward, SRO), PropertyIdentifier (survey, Pot Hissa, city survey, khata, UPIN / ULPIN), PropertyFeatures (type, extent, unit, building particulars).",
        "GISLookup (KSRSAC request / response ID, parcel polygon, centroid, nearest road, layer version, match result), GISSnapshot (image, coordinates, timestamp), ManualAreaSelection (reason, actor).",
        "RateVersionRef (CVC version ID, effective date, valuation area, rate head, rate), ValuationRecord (version, guideline value, components), ValuationComponent, MarketValueDecision (consideration, basis, Government flag, adjudicated override), DutyFeeResult.",
        "SRValuationDecision (accept / send back / refer, remarks), UndervaluationCase (trigger, SR estimate, reasons), FormI, DCReference, Hearing, DCOrder, Appeal, Deposit, DutyDifference (interest), Refund.",
        "MasterCrosswalk (Kaveri code, CVC area, KSRSAC ID, effective from / to), RateGapRequest.",
    ])
    h(doc, "8.16.2 Retention", 4)
    para(doc, "Valuation records, GIS snapshots and Sec. 45-A case records shall be preserved for the life of the registered document record.")
    h(doc, "8.16.3 Migration (high level)", 4)
    table(doc, ["Topic"], [
        ["Pending Sec. 45-A cases from Kaveri 1.0 / 2.0 with orders, deposits and payments (27140)"],
        ["Historic CVC rate versions needed for back-dated execution dates (from CVC module)"],
        ["Village / road / survey crosswalk baseline between Kaveri, CVC and KSRSAC"],
    ], [6.5])

    h(doc, "8.17 User stories and traceability (01-09-2026 discussion)", 3)
    para(doc, "User stories recorded in the 01-09-2026 discussion, mapped to the implementing requirements. Stories owned by the CVC rate-fixation process are traced to BRD-K3-CVC-GVF-001 and consumed here through the rate interface.")
    table(doc, ["US ID", "User story", "Legal / source ref", "Implemented by"], [
        ["US-GV-01", "As a Sub-Registrar, I want to compare the consideration with the guidelines published under Sec. 45-B, so that if undervalued I can estimate the value, collect the extra duty, or refer to the Deputy Commissioner", "Sec. 45-A", "7.1 steps 13–14; 7.2; FR-GVC-041, FR-UV-001–006"],
        ["US-GV-02", "As a citizen / SR, I want to use the official guidelines published and revised under Sec. 45-B, so that duty is computed on current notified rates", "Sec. 45-B", "FR-GVC-010–013"],
        ["US-GV-03", "As a citizen / SR, I want to calculate ad valorem duty on market value for Schedule instruments, so that the correct article and rate are applied", "Sec. 3 + Schedule", "FR-GVC-033, FR-GVC-034"],
        ["US-GV-04", "As a SR, I want to capture property particulars and market value in Form I when starting a Sec. 45-A reference", "Undervaluation Rules 1977", "FR-UV-005"],
        ["US-GV-05", "As a citizen / SR, I want to enter survey number, territorial division and property description, so that the correct guideline rate is picked", "Registration Rules 13–15", "FR-GVC-001, FR-GVC-002"],
        ["US-CVC-01", "As IGR / CVC, I want to estimate, publish and revise guidelines and constitute sub-committees", "Sec. 45-B", "Upstream — BRD-K3-CVC-GVF-001"],
        ["US-CVC-02", "As a CVC data-entry officer, I want to publish rates for SR and citizen logins, so that rates do not go missing", "Sec. 45-B", "Upstream — BRD-K3-CVC-GVF-001; consumption FR-GVC-010, FR-GVC-012"],
        ["US-CVC-03", "As an administrator, I want CVC maintained as the statutory body in Sec. 2(ac)", "Sec. 2(ac)", "Upstream — BRD-K3-CVC-GVF-001"],
        ["US-CVC-04", "As a SR, I want to apply CVC guidelines at registration, so that duty is charged on the higher of consideration and guideline value", "Sec. 45-A; Sec. 2(mm)", "FR-GVC-031"],
        ["US-CVC-05", "As a citizen / SR, I want to see and use the DR's adjudicated Sec. 45-A value instead of a higher guideline figure", "Sec. 45-A", "FR-GVC-035, FR-UV-014"],
        ["US-CVC-06", "As a SR / DC, I want Form I capture, DC determination and appeal under the 1977 Rules without data or payment stuck states", "Undervaluation Rules 1977", "FR-UV-005, FR-UV-010–024"],
        ["US-CVC-07", "As a CVC administrator, I want the Valuation Module to follow Sec. 45-B as strengthened by Act 8 of 2003", "Act 8 of 2003", "Upstream — BRD-K3-CVC-GVF-001"],
        ["US-CVC-08", "As a DC / DR, I want the undervaluation workflow as amended by RD 264 MUNOMU 99 (no provisional / Rule 6 path)", "RD 264 MUNOMU 99", "FR-UV-025"],
        ["US-GIS-01", "As a citizen / SR, I want to locate the property on GIS (KSRSAC) and apply the matching CVC rate, so that the correct spatial rate is used", "GIS + CVC rates", "FR-GIS-001–006"],
        ["US-GIS-02", "As a citizen / SR, I want survey / Pot Hissa / city survey mapped to the GIS layer, so that even 1 gunta is valued with the correct rural / urban rate", "Registration Rules 13–15", "FR-GIS-003, FR-GVC-004, FR-GVC-021"],
        ["US-GIS-03", "As a valuation / GIS administrator, I want village, road and survey masters aligned between Kaveri, CVC and KSRSAC", "GIS master alignment", "FR-GIS-020–024"],
    ], [0.9, 3.0, 1.1, 1.5])

    h(doc, "8.18 Requirements traceability matrix (RTM) — template", 3)
    table(doc, ["Req ID", "Act / Rule", "Requirement summary", "BRD section", "UI screen", "Test case ID", "Status"], [
        ["FR-GVC-001", "Registration Rule 13", "Cascading territorial selection", "8.1", "Property location", "TC-GVC-___", "Draft"],
        ["FR-GIS-002", "Registration Rules 13–15", "Parcel inside village; manual area with reason", "8.2", "GIS result and map", "TC-GIS-___", "Draft"],
        ["FR-GIS-003", "CVC methodology", "Spatial match to area / road / zone", "8.2", "GIS result and map", "TC-GIS-___", "Draft"],
        ["FR-GVC-011", "Sec. 2(mm)", "Rate version by execution date", "8.3", "Valuation summary", "TC-GVC-___", "Draft"],
        ["FR-GVC-031", "Sec. 2(mm)", "Higher of guideline value and consideration", "8.5", "Valuation summary", "TC-GVC-___", "Draft"],
        ["FR-GVC-033", "Sec. 3 + Schedule", "Duty / fee hand-off; Save gate", "8.5", "Valuation summary", "TC-GVC-___", "Draft"],
        ["FR-GVC-035", "Sec. 45-A", "Adjudicated value overrides", "8.5", "Valuation summary", "TC-GVC-___", "Draft"],
        ["FR-UV-005", "Sec. 45-A(1); Rules 1977", "Form I and reference to DC", "8.7", "Sec. 45-A estimate and reference", "TC-UV-___", "Draft"],
        ["FR-UV-013", "Sec. 45-A(3)", "Suo motu within two years", "8.8", "DC case workbench", "TC-UV-___", "Draft"],
        ["FR-UV-020", "Sec. 45-A(5)", "Appeal on 50% deposit", "8.9", "Appeal", "TC-UV-___", "Draft"],
        ["FR-UV-022", "Sec. 45-A(2), (5)", "Duty difference and interest", "8.9", "Duty difference payment", "TC-UV-___", "Draft"],
    ], [0.9, 1.1, 1.6, 0.6, 1.1, 0.7, 0.5])

    # 9 ---------------------------------------------------------------------
    h(doc, "9. Non-functional requirements", 2)
    para(doc, "The following non-functional parameters apply to guideline value calculation, GIS valuation and the Sec. 45-A workflow on all channels. Owners to validate: Solution Architect, DevOps / SDC, Security, DBA, Ops and Product Owner.")
    h(doc, "9.1 Performance and scalability", 3)
    table(doc, REQ4, [
        ["NFR-GVC-PERF-001", "Valuation computation (rate fetch + calculation + duty / fee) shall complete within 2 seconds", "Must", "p95 under NFR-GVC-PERF-003 load"],
        ["NFR-GVC-PERF-002", "KSRSAC parcel lookup shall return within 5 seconds; map tiles render within 3 seconds", "Must", "p95; timeout handled by FB-GVC-001"],
        ["NFR-GVC-PERF-003", "System shall support at least 5,000 concurrent citizen valuation sessions and 2,500 concurrent office sessions", "Must", "Load-test evidence"],
        ["NFR-GVC-SCALE-001", "Infrastructure shall scale for a 300% surge around guideline-rate revision effective dates and financial year-end", "Must", "Surge test at 3× baseline"],
    ], [1.2, 3.4, 0.6, 1.9])
    h(doc, "9.2 Security, data integrity and privacy", 3)
    table(doc, REQ4, [
        ["NFR-GVC-SEC-001", "Data in transit shall use TLS 1.3 and data at rest AES-256", "Must", "Confirmed in hosting design"],
        ["NFR-GVC-INT-001", "Published rate versions consumed by valuation shall be immutable; every valuation shall store the rate version ID used", "Must", "Recalculation reproduces the same value"],
        ["NFR-GVC-AUD-001", "All valuation decisions, manual-area selections, SR decisions, Sec. 45-A orders and payments shall generate an immutable, timestamped audit log with actor ID", "Must", "Append-only audit store"],
        ["NFR-GVC-PRIV-001", "Party PII shall be masked in logs and unauthorised views", "Must", "Log review"],
    ], [1.2, 3.4, 0.6, 1.9])
    h(doc, "9.3 System availability and error handling", 3)
    table(doc, REQ4, [
        ["NFR-GVC-AVA-001", "If KSRSAC, the fee engine or the payment gateway is unavailable, the system shall show a clear bilingual message and keep the valuation resumable without data loss", "Must", "Outage test; state resumable"],
        ["NFR-GVC-PAY-001", "On payment timeout, the system shall poll for the final status and lock the pay action to prevent duplicate payment", "Must", "No double debit"],
    ], [1.2, 3.4, 0.6, 1.9])
    h(doc, "9.4 Security audit and compliance (VAPT policy)", 3)
    table(doc, REQ4, [
        ["NFR-GVC-VAPT-001", "The module shall undergo comprehensive VAPT (web, API, GIS integration, payment) covering OWASP Top 10 and business-logic flaws before production", "Must", "VAPT report accepted by Security"],
        ["NFR-GVC-VAPT-002", "Critical / high findings shall be remediated and re-tested before Go-Live; annual and change-triggered VAPT thereafter", "Must", "Zero open Critical / High at Go-Live"],
    ], [1.2, 3.4, 0.6, 1.9])

    # 10 --------------------------------------------------------------------
    h(doc, "10. Risk and mitigation strategy", 2)
    para(doc, "Operational, technical and adoption risks for guideline value calculation and GIS valuation, with mandatory mitigations.")
    table(doc, ["Risk ID", "Risk", "Mitigation", "Related requirements"], [
        ["RS-GVC-001", "KSRSAC cadastral data is incomplete or misaligned with Kaveri village / CVC area masters, producing wrong or no spatial match", "Crosswalk masters, reconciliation reports and manual-area fallback with reason and SR flag; phased district roll-out after match-rate threshold is met", "FR-GIS-002, FR-GIS-020–024, FR-GVC-065"],
        ["RS-GVC-002", "CVC rate gaps for some areas / property types at go-live", "Rate-gap detection and routing to CVC; pre-go-live coverage report from CVC module", "FR-GVC-015, FR-GVC-062"],
        ["RS-GVC-003", "Divergence between valuation and fee engine causes zero / wrong duty", "Single market value hand-off contract; Save gate; regression suite per Schedule article", "FR-GVC-033, FR-GVC-034"],
        ["RS-GVC-004", "Migration of legacy pending Sec. 45-A cases is incomplete", "Case-by-case reconciliation with DC offices before cut-over", "FR-UV-026"],
        ["RS-GVC-005", "Users find GIS map selection difficult", "Identifier-first search with map as confirmation; guided help and tooltips", "FR-GVC-002, FR-GIS-007"],
    ], [0.9, 1.9, 2.4, 1.3])

    # 11 --------------------------------------------------------------------
    h(doc, "11. System fallbacks and error handling", 2)
    para(doc, "The system shall handle exceptions and integration failures without losing user data. 9.3 states the availability NFR; this section specifies the operational fallbacks.")
    table(doc, REQ4, [
        ["FB-GVC-001", "KSRSAC unavailable or timed out: allow manual valuation-area selection with reason, flag for SR, and retry GIS lookup in background; record that the snapshot is pending", "Must", "Valuation not blocked; flag visible to SR"],
        ["FB-GVC-002", "Parcel not found / outside village: show bilingual guidance to correct identifiers; allow manual area with reason", "Must", "FR-GIS-002"],
        ["FB-GVC-003", "Rate not available: never show blank or zero market value; show \"Rate not available\", log rate-gap and route per FR-GVC-015", "Must", "No blank value"],
        ["FB-GVC-004", "Fee engine failure: retain valuation, show retry, keep Save disabled until valid duty / fee returned", "Must", "No zero-fee submission"],
        ["FB-GVC-005", "Sec. 45-A payment timeout: poll gateway, lock duplicate payment; release document only on confirmed payment", "Must", "No stuck or double payment"],
        ["FB-GVC-006", "Notification gateway delay: queue and retry; statutory timers (90 / 60 days) run from order date regardless of delivery", "Must", "Workflow continues"],
    ], REQ4_W)

    # 12 --------------------------------------------------------------------
    h(doc, "12. Training and change management", 2)
    para(doc, "Moving from Kaveri 2.0 valuation to GIS-based valuation and a digitised Sec. 45-A workflow needs structured change management for department staff and citizens.")
    h(doc, "12.1 Target audience", 3)
    para(doc, "Sub-Registrars, DEOs, District Registrars / Deputy Commissioners, Regional Commissioner offices, valuation / GIS administrators and IT Helpdesk.")
    h(doc, "12.2 Training delivery", 3)
    bullets(doc, [
        "Role-based workshops on GIS lookup, manual-area fallback, reading the valuation breakup and SR send-back / referral decisions.",
        "DC and appellate training on the Sec. 45-A workbench, hearings, orders, interest and refunds.",
        "SOPs, quick-reference cards and video tutorials in English and Kannada.",
    ])
    h(doc, "12.3 Citizen change management", 3)
    para(doc, "In-screen help explaining guideline value, the higher-of rule, the map and what to do when the parcel is not found; public guideline value search to set expectations before registration.")
    h(doc, "12.4 Post-Go-Live support", 3)
    para(doc, "Dedicated hyper-care for 90 days after launch, focused on GIS match issues, rate gaps and fee-engine hand-off.")

    # Appendix --------------------------------------------------------------
    h(doc, "Appendix A — References", 2)
    bullets(doc, [
        "The Karnataka Stamp Act, 1957 (Secs. 2(ac), 2(mm), 3, 27, 45-A, 45-B and Schedule)",
        "Karnataka Stamp (Prevention of Undervaluation of Instruments) Rules, 1977 (GSR 81 / RD 73 EST 74) and amendment RD 264 MUNOMU 99",
        "Karnataka Registration Rules — Rules 13–15",
        "BRD-K3-CVC-GVF-001 — Valuation Module (CVC) — Guidance Value Fixation",
        "Consolidated_Requirement_Discussions_25082026_to_11092026 — Section 4 (01-09-2026); Document_Registration_requirement_01092026_v2.docx",
        "Process diagrams: Process_A_Guideline_Value_Calculation.drawio, Process_B_Sec45A_Undervaluation.drawio",
        "OWASP Top 10 — https://owasp.org/ (VAPT scope)",
    ])

    h(doc, "Acceptance and sign-off of BRD", 2)
    table(doc, ["Role", "Name", "Signature / Date"], [
        ["Product Owner", "", ""],
        ["Domain Expert (DIGR — Valuation)", "", ""],
        ["Business Analyst", "", ""],
    ], [2.2, 2.2, 2.1])


def main():
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    set_header(doc)
    doc.core_properties.title = "BRD — Guideline Value Calculation and GIS Valuation"
    doc.core_properties.subject = "Kaveri 3.0"
    build(doc)
    doc.save(str(OUT))
    print("Saved", OUT)


if __name__ == "__main__":
    main()
