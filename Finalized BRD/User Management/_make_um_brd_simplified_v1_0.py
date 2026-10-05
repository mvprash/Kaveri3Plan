# -*- coding: utf-8 -*-
"""Create BRD_User_Management_Simplified_v1.0.docx — main User Management BRD in plain business language, organised by activity, with unique requirement IDs."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DST = Path(r"Finalized BRD/User Management/BRD_User_Management_Simplified_v1.0.docx")

DOC_VERSION = "1.0"
DOC_DATE = "05-10-2026"
MODULE = "UM"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
HEADER_FILL = "1F3A5F"


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

PURPOSE = (
    "This document sets out what the Department of Stamps and Registration needs from the User "
    "Management part of KAVERI 3.0, written in simple business language. It covers how every type "
    "of user is registered and signs in, how officers are posted, transferred, relieved and placed "
    "on leave, and how roles, sanctioned posts and reporting hierarchies are maintained."
)
BACKGROUND = (
    "KAVERI 3.0 replaces the separate user administration of Kaveri 2.0 with one common way of "
    "managing all users, roles, posts and offices. Nobody uses a password — citizens and other "
    "department users sign in with an OTP on their mobile, and DSR officers sign in with face or "
    "biometric check. An officer's access depends on the post they hold in an office."
)
TWO_PARTS = [
    "Kaveri application — where the Department's work is done: screens for administrators, offices, posts, roles, transfers, leave and reports.",
    "IAM (Identity and Access Management) — the sign-in system: keeps each user's login details and checks OTP, face or biometrics before letting the user in.",
]
HOW_TO_READ = (
    "Every requirement has a unique ID made of the module code (UM), an activity code and a running "
    "number — for example UM-TIN-02 is the second Transfer In requirement. IDs are never reused; a "
    "requirement that is withdrawn keeps its ID and is marked as withdrawn. The activity codes are:"
)
SCOPE_IN = [
    "Registration of citizens, DSR officers and other department users",
    "Sign-in for every type of user, post selection, additional charge, lost mobile and change of mobile or email",
    "Transfer In, Relieving / Transfer Out, temporary absence and temporary charge",
    "Roles and what each role can do; sanctioned posts for each office; office and officer reporting hierarchies",
    "Managing users (suspend, deactivate, search), notifications, reports and records",
    "Moving e-KYC-verified citizens from Kaveri 2.0 to KAVERI 3.0",
]
SCOPE_OUT = [
    "The work done inside other KAVERI modules (document registration, marriage registration, encumbrance search, etc.) — only who may use them is covered here",
    "Moving DSR officers and other department users from Kaveri 2.0 — they are created afresh",
]

OBJECTIVES = [
    "Let citizens register and sign in easily and safely, without passwords.",
    "Make sure every officer can work only on the post they actually hold, in the office they are posted to.",
    "Keep sanctioned posts, vacancies and postings correct at all times, including after transfers and relieving.",
    "Let superiors manage transfers, relieving, leave and temporary charge for the officers under them.",
    "Give the Department full control over roles and what each role can do.",
    "Keep a complete record of who did what and when, for audit and reports.",
    "Move existing e-KYC-verified citizens from Kaveri 2.0 so they need not register again.",
]

USERS = [
    ["Citizen (public user)", "Registers on their own", "Username chosen by the citizen", "Username + Captcha + OTP to mobile"],
    ["DSR Officer", "Created by an authorised administrator", "KGID", "KGID + Captcha + face or biometric (no OTP)"],
    ["Other Department user", "Created by an authorised administrator", "Department Code + Employee ID or KGID (e.g. REV-12345)", "Username + Captcha + OTP to mobile"],
]
USERS_NOTE = (
    "Application Admin is one of the roles under the DSR Officer user type. A DSR Officer with this "
    "role maintains roles, posts, sanctioned posts, offices and hierarchies."
)

# (activity code, section title, intro, [requirements])
SECTIONS = [
    ("CIT", "Citizen registration",
     "A citizen registers on their own, without visiting an office.",
     ["A citizen can register at any time; no approval is needed.",
      "The citizen chooses a Username. The system says at once if it is already taken and suggests other names.",
      "The citizen enters an email address and a mobile number. An OTP is sent to each; when both are entered correctly, the account is created.",
      "Aadhaar e-KYC must be completed successfully after the first sign-in. Until it is completed, the citizen cannot use any service — only the e-KYC step is shown, at every sign-in, until it succeeds.",
      "No password and no security questions are asked.",
      "The Username must be unique. The same mobile or email may be used on more than one account.",
      "The Username cannot be changed after registration."]),
    ("DSR", "DSR Officer registration",
     "An authorised administrator creates the officer. Giving a post at creation is optional; posts can be given later in a separate activity.",
     ["Only authorised administrators can create DSR Officers.",
      "The officer's KGID becomes the Username and must be unique. The mobile number is entered. No email and no security questions are needed.",
      "Giving a post at creation is not mandatory. The administrator may give one or more posts while creating the officer, or only create the officer and give the post later through the separate posting activity (Transfer In). When posts are given at creation, only posts with a vacancy, at offices under the administrator, are shown.",
      "An officer may hold more than one post at the same time, provided each post has a vacancy.",
      "There is no primary / secondary post and no end date on a post.",
      "The officer's face or biometric details are used for sign-in.",
      "The officer is informed by SMS when the account is created.",
      "An officer created without a post cannot sign in until a post is given."]),
    ("ODU", "Other Department user registration",
     "An authorised administrator creates users from other government departments.",
     ["Only authorised administrators can create Other Department users.",
      "The administrator selects the Department and enters Employee ID or KGID, official email of that department, and mobile. The Username is Department Code + Employee ID / KGID (for example REV-12345).",
      "Where a list of allowed official email domains is set up, personal email addresses are refused.",
      "Exactly one role is given, chosen only from Other Department roles. Posts do not apply to these users.",
      "An optional End Date may be entered. On that date the account stops automatically."]),
    ("LOG", "User logins",
     "How each type of user signs in, and the rules that keep sign-in safe.",
     ["No user signs in with a password.",
      "Citizens and Other Department users sign in with Username + Captcha + OTP sent to their registered mobile. The OTP is never sent to email.",
      "DSR Officers sign in with KGID + Captcha + face or biometric (for example fingerprint). No OTP is sent.",
      "OTP rules: 6 digits; valid 5 minutes for sign-in and 10 minutes for registration or reset; 3 wrong entries cancel the OTP; a new OTP can be asked for after 60 seconds, at most 3 times in 15 minutes.",
      "After 5 failed attempts the Username is locked for 15 minutes. An administrator may unlock earlier, with a reason.",
      "The user is signed out after 10 minutes without activity, and after 4 hours in any case.",
      "A user can be signed in at only one place at a time. A new sign-in closes the earlier one.",
      "The user can sign out at any time.",
      "A DSR Officer holding more than one post selects one post at sign-in, shown as \"Role — Post — Office (Office Code)\". If the officer holds only one post, it is selected automatically.",
      "After sign-in the officer can do only the work allowed for the selected post — not the work of their other posts.",
      "The screen header shows the selected post, role and office.",
      "An officer on approved leave or absence today cannot sign in. The message shows the type of absence and the date until which it applies."]),
    ("MOB", "Lost mobile and change of mobile or email",
     "Each type of user has its own rule.",
     ["A citizen who has lost their mobile can reset it: Username + Captcha, then Aadhaar e-KYC, then a PIN sent to the registered email, then an OTP to the new mobile. The citizen is informed by email.",
      "Repeated failures in the reset lock it for 30 minutes.",
      "A citizen can change mobile or email after signing in; an OTP is sent to the new mobile or email.",
      "A DSR Officer can change their own mobile after signing in; an OTP is sent to the new number.",
      "An Other Department user's mobile or email can be changed only by an administrator, with a reason.",
      "DSR Officers and Other Department users have no lost-mobile reset before sign-in.",
      "Users can upload a profile photo."]),
    ("ADC", "Additional charge (after sign-in)",
     "An officer may look after a vacant post directly under them in the same office, during the same session.",
     ["After sign-in, an officer can take additional charge of a post directly under their post, in the same office, only if nobody holds that post.",
      "If that post is vacant, the next post below it may also be offered. Example: a Sub-Registrar at SRO Yeshwanthapura can take FDA if no FDA is posted there, and SDA if no SDA is posted under that vacant FDA. If even one FDA is posted, FDA is not offered.",
      "No order is needed. Only one additional charge at a time.",
      "While on additional charge the officer can do only the work of that post (for example a Sub-Registrar acting as FDA cannot sign as Sub-Registrar). The officer can switch back at any time without signing out.",
      "The header shows which post is currently active."]),
    ("TIN", "Transfer In",
     "A superior places a DSR Officer in a post that has a vacancy. This is also how a post is given to an officer who was created without one.",
     ["A superior can place an officer only in offices under them, and only in posts that report directly to the superior's post.",
      "Transfer In is allowed only when the post has a vacancy (occupied is less than sanctioned strength). If the post is full, Transfer In is blocked.",
      "The Transfer Order / Reporting Order number is required (with upload where available). No Joining Date is entered.",
      "Transfer In takes effect immediately and the vacancy count is updated. The officer sees the new post at the next sign-in.",
      "A post whose holder is only on leave is not a vacancy."]),
    ("REL", "Relieving / Transfer Out",
     "A superior relieves an officer from one or more posts.",
     ["A superior can relieve only officers in offices under them, holding posts that report directly to the superior's post. Example: the District Registrar of DRO Bengaluru can relieve the Sub-Registrar of SRO Yeshwanthapura, but not of SRO Mysuru East; the IGR cannot relieve a Sub-Registrar even though the SRO is visible.",
      "Relieving Date, Relieving Order and Relieving Reason are required. The reason must be one of: Deputation, Transfer, Suspension, Superannuation, Death.",
      "The officer keeps the post until the end of the Relieving Date.",
      "Shortly after midnight the system removes the officer from the post and updates the vacancy counts. If this fails, Kaveri IT Cell is alerted.",
      "Any open session the officer has on that post is ended.",
      "If the officer has no post left, they cannot sign in until a new post is given."]),
    ("ABS", "Temporary absence and temporary charge",
     "A superior records leave or other absence, and may give the post's charge to another officer for that period.",
     ["A superior records Leave, OOD or Other absence against an officer's post, with reason, from-date and to-date, and an optional order. The officer cannot record their own absence.",
      "The same office and reporting rules as for relieving apply.",
      "During the absence the officer cannot sign in at all, on any post.",
      "An absence does not create a vacancy: Transfer In is not allowed because of it, and the post is not offered for additional charge.",
      "The superior may give temporary charge of the absent officer's post to another officer under them, even at another office under the superior (for example a District Registrar gives the charge of SRO A's Sub-Registrar to the Sub-Registrar of SRO B). Only one temporary charge per absent post at a time.",
      "The covering officer sees \"Temporary charge — Role — Post — Office — covering <name>\" at sign-in. Selecting it gives only that post's work.",
      "The absence and temporary charge end automatically after the to-date (11:59 PM). The superior can cancel earlier, with immediate effect."]),
    ("ROL", "Creating / modifying roles and their access",
     "Roles, and what each role can do, are maintained by the Application Admin.",
     ["There is one list of roles for all users, grouped by role category: Citizen, DSR, Other Department.",
      "Application Admin is one of the DSR roles.",
      "The Application Admin can add, change and deactivate roles. The starting list of roles is confirmed by the Domain Expert before go-live.",
      "A list of divisions is maintained (for example Admin, Vigilance, Enforcement, Computers, Intelligence & Audit, CVC).",
      "Each post is linked to its role(s). A post with no role cannot be sanctioned or given to an officer. Division posts get only their own division's role (for example DIGR Admin does not get Enforcement work).",
      "The Application Admin maintains the modules (for example Registration of Documents, Marriage Registration, Encumbrance Search, Certified Copy), their functions (View, Add, Edit, Approve, Sign, Print, Download) and which role may use which function.",
      "Only a role of the matching category can be given to a user (for example only Other Department roles to Other Department users).",
      "Users can see and do only what their role allows. Anything else is refused and recorded.",
      "Every change is recorded with who made it and when."]),
    ("SAN", "Creating / modifying sanctioned posts for each office",
     "Posts and sanctioned strength are maintained by the Application Admin.",
     ["The Application Admin maintains the list of posts (for example Sub-Registrar, FDA (Enforcement), DEO, District Registrar), each linked to a division.",
      "The Application Admin sets which posts can exist in which type of office (for example Sub-Registrar only at a Sub-Registrar Office, District Registrar only at a District Registrar Office).",
      "The Application Admin sets the sanctioned strength of each post in each office. Every change is recorded.",
      "For every post in every office the system shows sanctioned, occupied and vacant counts, and whether the post is fully vacant.",
      "Officers can be placed only in sanctioned posts with a vacancy; placing above sanctioned strength is blocked."]),
    ("HIE", "Creating / managing reporting hierarchies",
     "The office hierarchy and the officer reporting hierarchy are maintained by the Application Admin.",
     ["Office hierarchy: MS Building (Secretariat) → IGR Office (Head Office) → District Registrar Offices → Sub-Registrar Offices. Each office has a code, name, type and parent office.",
      "Officer reporting hierarchy, by post: Additional Chief Secretary / Principal Secretary / Secretary → IGR → divisions → posts under them. Each post reports to one parent post.",
      "The Application Admin can add, change, reorder, enable and disable offices and hierarchy entries. Every change is recorded.",
      "The hierarchies decide which offices a superior sees and for whom they can create users, transfer, relieve, record absence, give temporary charge, and take additional charge."]),
    ("USR", "Managing users",
     "Day-to-day administration of user accounts.",
     ["Administrators can create, change, suspend and deactivate users. Deactivation needs a reason.",
      "Administrators can search users by Username / KGID, user type, role, office, division and status.",
      "A wrong Username (including KGID or Employee ID) can be corrected only by an administrator, with a reason.",
      "Users are informed by SMS / email when their account is created, suspended or deactivated.",
      "Every administrator action is recorded.",
      "The Application Admin role is given to a DSR Officer in the same way as other DSR roles. No second approval (maker-checker) is needed for user creation or master changes."]),
    ("MIG", "Moving citizens from Kaveri 2.0",
     "Existing e-KYC-verified citizens are moved so that they need not register again.",
     ["Only Kaveri 2.0 citizens whose e-KYC is completed (e-KYC = true) are moved. Others must register afresh on KAVERI 3.0.",
      "Only four items are moved: the email ID (Kaveri 2.0 username), mobile number, email, and e-KYC status. Passwords are not moved.",
      "The Kaveri 2.0 email ID becomes the citizen's Username and cannot be changed.",
      "Aadhaar e-KYC is not asked again.",
      "At first sign-in the citizen uses email ID + Captcha + OTP to mobile, verifies the email by OTP and accepts the privacy notice. The account then becomes active.",
      "A moved citizen who has lost their mobile uses the normal lost-mobile reset.",
      "Records with a wrong or duplicate email ID, or an invalid mobile number, are set aside for correction. If the email is blank, the email ID is used as the email.",
      "The Kaveri 2.0 email ID is kept as a link to the citizen's old applications, documents and payments.",
      "Trial runs are done first. The counts must match (total = not moved + set aside + moved). Kaveri 2.0 citizen registration is frozen at cut-over, and go-live needs sign-off by Kaveri IT Cell and the Product Owner.",
      "Moved citizens receive an SMS / email in Kannada and English asking them to activate their account. It contains no OTP or link."]),
]

GENERAL = [
    ["Language and access", "Citizen screens are in Kannada and English and follow Government accessibility guidelines (GIGW). Registration and sign-in work on desktop and mobile browsers."],
    ["Speed", "Sign-in OTP reaches the mobile within 5 seconds; registration OTPs and the reset PIN within 30 seconds. Sign-in completes within 2 seconds after the OTP / face / biometric check."],
    ["Capacity", "At least 500 users signed in at the same time at launch; higher targets confirmed during testing with Kaveri IT Cell."],
    ["Security", "No passwords are stored. All connections are secure. Aadhaar e-KYC and biometrics follow the Aadhaar Act and UIDAI rules."],
    ["Data", "All personal data stays in India, on Government-approved (MeitY-empanelled) hosting, and follows the DPDP Act 2023."],
    ["Records", "Every create, change, sign-in, transfer, relieving, absence and charge is recorded with who did it and when. Absence and charge records are kept for at least seven years."],
    ["Night-time update", "The nightly update of postings and vacancies finishes before the first sign-in of the day. If it fails, Kaveri IT Cell is alerted."],
]

REPORTS = [
    "Active, inactive and suspended users.",
    "Sign-in attempts (successful and failed) for a chosen period.",
    "Roles given to users, and what each role can do.",
    "Sanctioned posts in each office: sanctioned, occupied, vacant, and fully vacant.",
    "Changes of mobile and email, and citizen lost-mobile resets.",
    "Additional charge taken and cleared, by officer.",
    "Temporary absence and temporary charge, by officer and period.",
    "Nightly posting update: officers removed from posts and vacancy changes.",
    "Transfer In and Relieving history, with orders and reasons.",
    "Posting history of each officer.",
    "Moving citizens from Kaveri 2.0: counts moved, not moved and set aside; activation progress after go-live.",
]

ACCEPTANCE = [
    "A citizen registers with Username and email and mobile OTPs, signs in with Username, Captcha and OTP, is asked for Aadhaar e-KYC at the first sign-in, and cannot use any service until e-KYC succeeds.",
    "A DSR Officer with two posts must select one post at sign-in; with one post it is selected automatically; the officer can do only that post's work.",
    "After sign-in at SRO Yeshwanthapura, additional charge offers only fully vacant posts under the officer at that same office — not vacant posts at another SRO.",
    "Transfer In to a full post is blocked; after relieving takes effect at midnight, Transfer In succeeds immediately with only a Transfer / Reporting Order.",
    "An officer relieved with Relieving Date D keeps the post until 11:59 PM on D and is removed from it shortly after midnight.",
    "A District Registrar records Leave for the Sub-Registrar of SRO A, who then cannot sign in; temporary charge given to the Sub-Registrar of SRO B appears at their sign-in and ends after the to-date.",
    "OTP and session rules work as stated: 5-minute OTP, 3 wrong entries, 15-minute lock after 5 failures, 10-minute idle sign-out, 4-hour limit, one session per user.",
    "Every active role has its module access set up and confirmed by the Domain Expert before go-live.",
    "A moved Kaveri 2.0 citizen signs in with their email ID, is not asked for e-KYC again, and becomes active after email OTP and accepting the privacy notice; the migration counts match.",
    "All general needs in Section 5 are met, security testing has no open critical or high findings, and user and administrator guides are delivered.",
]

GLOSSARY = [
    ["Kaveri application", "The part of KAVERI 3.0 where the Department's work is done (screens, offices, posts, roles, transfers, reports)."],
    ["IAM", "Identity and Access Management — the sign-in system that keeps login details and checks OTP, face or biometrics."],
    ["DSR", "Department of Stamps and Registration."],
    ["KGID", "Karnataka Government ID of an officer; used as the DSR Officer's Username."],
    ["IGR / DRO / SRO", "Inspector General of Registration office (Head Office) / District Registrar Office / Sub-Registrar Office."],
    ["Application Admin", "A DSR role; the officer holding it maintains roles, posts, sanctioned posts, offices and hierarchies."],
    ["Superior", "The officer whose post is directly above the officer's post in the reporting hierarchy, at an office under them."],
    ["Sanctioned strength", "The approved number of persons for a post in an office."],
    ["Occupied / Vacancy", "Number of officers currently holding the post / sanctioned strength minus occupied."],
    ["Fully vacant", "A post in an office that nobody holds (occupied = 0). Only such posts can be taken as additional charge."],
    ["Additional charge", "An officer looking after a fully vacant post directly under them in the same office, during their session."],
    ["Temporary charge", "An officer looking after the post of an absent officer, given by the superior for the period of absence."],
    ["Transfer In", "Placing an officer in a post with a vacancy, with a Transfer / Reporting Order."],
    ["Relieving / Transfer Out", "Removing an officer from a post from the end of the Relieving Date, with a Relieving Order and reason."],
    ["OOD", "Out of Duty / On Other Duty — a type of temporary absence."],
    ["OTP / PIN", "One-time code sent by SMS or email to confirm the user."],
    ["Captcha", "A simple check on the sign-in screen that the user is a person."],
    ["Aadhaar e-KYC", "Confirmation of a citizen's identity through UIDAI using Aadhaar."],
]


def rid(code: str, n: int) -> str:
    return f"{MODULE}-{code}-{n:02d}"


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------

def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def cell_text(cell, text: str, bold=False, size=10, white=False) -> None:
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if white:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_table(doc, header, rows, widths, bold_col=None):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell_text(table.rows[0].cells[i], h, bold=True, white=True)
        shade(table.rows[0].cells[i], HEADER_FILL)
    repeat_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for i, v in enumerate(values):
            cell_text(cells[i], v, bold=(i == bold_col))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()


def heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY


def para(doc, text):
    doc.add_paragraph(text).paragraph_format.space_after = Pt(6)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def header_footer(section) -> None:
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("Kaveri 3.0  |  BRD  |  User Management")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run()
    for kind, val in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    run.font.size = Pt(9)


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def main() -> None:
    codes = [s[0] for s in SECTIONS] + ["GEN", "RPT", "ACC"]
    assert len(codes) == len(set(codes)), "activity codes must be unique"

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
    header_footer(sec)

    heading(doc, "Business Requirements Document (BRD)", 0)
    heading(doc, "User Management Module", 1)

    heading(doc, "Document Control", 2)
    add_table(doc, ["Field", "Value"], [
        ["Document ID", "BRD-K3-UM-001-S"],
        ["Version", DOC_VERSION],
        ["Status", "Draft for review"],
        ["Module", "User Management"],
        ["Author (BA)", "Nandha Kumar"],
        ["Product Owner", "M V Prashanth"],
        ["Domain expert / reviewer", "Prabhakar Naik"],
        ["Target audience", "Department of Stamps and Registration, Government of Karnataka"],
        ["Last updated", DOC_DATE],
    ], [5.0, 12.0])

    heading(doc, "Version History", 2)
    add_table(doc, ["Version", "Date", "Author", "Description"], [
        [DOC_VERSION, DOC_DATE, "Nandha Kumar", "User Management BRD in plain business language, organised by activity, with unique requirement IDs."],
    ], [2.0, 2.5, 3.0, 9.5])

    heading(doc, "1. Introduction", 1)
    heading(doc, "1.1 Purpose", 2)
    para(doc, PURPOSE)
    heading(doc, "1.2 Background", 2)
    para(doc, BACKGROUND)
    heading(doc, "1.3 Two parts of KAVERI 3.0", 2)
    bullets(doc, TWO_PARTS)
    heading(doc, "1.4 Requirement IDs", 2)
    para(doc, HOW_TO_READ)
    code_rows = [[c, f"4.{i} {t}"] for i, (c, t, _, _) in enumerate(SECTIONS, start=1)]
    code_rows += [["GEN", "5 General needs"], ["RPT", "6 Reports"], ["ACC", "7 Acceptance"]]
    add_table(doc, ["Code", "Activity (section)"], code_rows, [2.5, 14.5], bold_col=0)
    heading(doc, "1.5 Scope", 2)
    para(doc, "In scope:")
    bullets(doc, SCOPE_IN)
    para(doc, "Out of scope:")
    bullets(doc, SCOPE_OUT)

    heading(doc, "2. Business Objectives", 1)
    bullets(doc, OBJECTIVES)

    heading(doc, "3. Who Uses the System", 1)
    add_table(doc, ["User type", "How the account is created", "Username", "How they sign in"],
              USERS, [3.6, 4.0, 4.2, 5.2])
    para(doc, USERS_NOTE)

    heading(doc, "4. Requirements by Activity", 1)
    for s, (code, title, intro, reqs) in enumerate(SECTIONS, start=1):
        heading(doc, f"4.{s} {title}", 2)
        para(doc, intro)
        rows = [[rid(code, n), text] for n, text in enumerate(reqs, start=1)]
        add_table(doc, ["Req ID", "Requirement"], rows, [2.6, 14.4], bold_col=0)

    heading(doc, "5. General Needs", 1)
    add_table(doc, ["Req ID", "Area", "What is needed"],
              [[rid("GEN", n), a, t] for n, (a, t) in enumerate(GENERAL, start=1)], [2.6, 3.2, 11.2], bold_col=0)

    heading(doc, "6. Reports", 1)
    para(doc, "The system shall provide the following reports, for a chosen period where relevant:")
    add_table(doc, ["Req ID", "Report"], [[rid("RPT", n), t] for n, t in enumerate(REPORTS, start=1)], [2.6, 14.4], bold_col=0)

    heading(doc, "7. Acceptance", 1)
    para(doc, "User Management will be accepted when the following have been shown to work and signed off by the Product Owner:")
    add_table(doc, ["ID", "Acceptance criterion"], [[rid("ACC", n), t] for n, t in enumerate(ACCEPTANCE, start=1)], [2.6, 14.4], bold_col=0)

    heading(doc, "8. Glossary", 1)
    add_table(doc, ["Term", "Meaning"], GLOSSARY, [4.0, 13.0])

    heading(doc, "9. Approval", 1)
    para(doc, "By signing below, the stakeholders confirm they have reviewed and approved this document.")
    add_table(doc, ["Name", "Role", "Signature", "Date"], [
        ["M V Prashanth", "Product Owner", "", ""],
        ["Prabhakar Naik", "Domain expert / reviewer", "", ""],
        ["", "Department of Stamps and Registration", "", ""],
    ], [4.5, 5.5, 4.0, 3.0])

    doc.save(str(DST))
    print(f"Saved {DST}")


if __name__ == "__main__":
    main()
