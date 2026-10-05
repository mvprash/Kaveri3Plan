import copy
import docx
from docx.table import Table

PATH = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.0.docx"
d = docx.Document(PATH)

prf = next(t for t in d.tables if len(t.rows) > 1 and t.rows[1].cells[0].text == "UM-PRF-01")
last = prf.rows[-1]
new_tr = copy.deepcopy(last._tr)
last._tr.addnext(new_tr)
row = prf.rows[-1]
for cell, text in zip(row.cells, ["UM-PRF-02",
                                  "Users can add or update other details in their profile, such as address and additional contact numbers."]):
    p = cell.paragraphs[0]
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)

d.save(PATH)
print([(r.cells[0].text, r.cells[1].text) for r in prf.rows])
