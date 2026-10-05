import copy
import docx
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.2.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.3.docx"
DATE = "05-10-2026"
d = docx.Document(SRC)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find_row(first_cell):
    hits = [r for t in d.tables for r in t.rows if r.cells[0].text == first_cell]
    assert len(hits) == 1, (first_cell, len(hits))
    return hits[0]


def find_para(text):
    hits = [p for p in d.paragraphs if p.text == text]
    assert len(hits) == 1, (text, len(hits))
    return hits[0]


def add_row_after(row, values):
    tr = copy.deepcopy(row._tr)
    row._tr.addnext(tr)
    new = _Row(tr, row._parent)
    for cell, text in zip(new.cells, values):
        set_text(cell.paragraphs[0], text)
    return new


def fill_table(table, rows):
    template = table.rows[1]._tr
    for r in list(table.rows[1:]):
        r._tr.getparent().remove(r._tr)
    for values in rows:
        tr = copy.deepcopy(template)
        table._tbl.append(tr)
        for cell, text in zip(table.rows[-1].cells, values):
            set_text(cell.paragraphs[0], text)


for old, new in [("4.14 Managing users", "4.15 Managing users"),
                 ("4.15 Moving citizens from Kaveri 2.0", "4.16 Moving citizens from Kaveri 2.0")]:
    set_text(find_para(old), new)
    set_text(next(r for r in d.tables[2].rows if r.cells[1].text == old).cells[1].paragraphs[0], new)

add_row_after(find_row("HIE"), ["DSG", "4.14 Creating designations and mapping them to DSR Officers"])

scope = find_para("Roles and what each role can do; sanctioned posts for each office; office and officer "
                  "reporting hierarchies")
scope_new = copy.deepcopy(scope._p)
scope._p.addnext(scope_new)
set_text(Paragraph(scope_new, scope._parent), "Designations, and mapping each DSR Officer to one designation")

heading = find_para("4.13 Creating / managing reporting hierarchies")._p
block = [heading]
for _ in range(3):
    block.append(block[-1].getnext())
anchor = block[-1]
for el in block:
    clone = copy.deepcopy(el)
    anchor.addnext(clone)
    anchor = clone
new_heading = Paragraph(anchor.getprevious().getprevious().getprevious(), heading.getparent())
new_intro = Paragraph(new_heading._p.getnext(), heading.getparent())
new_table = Table(new_intro._p.getnext(), heading.getparent())
set_text(new_heading, "4.14 Creating designations and mapping them to DSR Officers")
set_text(new_intro, "Designations are created by an authorised administrator, and each designation is "
                    "mapped to a DSR Officer.")
fill_table(new_table, [
    ("UM-DSG-01", "An authorised administrator maintains the list of designations. Each designation has a "
                  "code and a name, and can be added, changed, enabled and disabled."),
    ("UM-DSG-02", "An authorised administrator maps a designation to a DSR Officer."),
    ("UM-DSG-03", "The mapping is always one to one: a DSR Officer has only one designation, and a "
                  "designation is mapped to only one DSR Officer."),
    ("UM-DSG-04", "Designations apply only to DSR Officers. Other Department users and citizens have no "
                  "designation."),
    ("UM-DSG-05", "Every change to designations and to their mapping is recorded with who made it and when."),
])

add_row_after(find_row("UM-RPT-11"), ["UM-RPT-12", "Designations and the DSR Officer mapped to each."])
add_row_after(find_row("Authorised administrator"), [
    "Designation", "A title created by an authorised administrator and mapped to exactly one DSR Officer; "
                   "each DSR Officer has only one designation."])

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.3")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)
add_row_after(d.tables[1].rows[-1], [
    "1.3", DATE, "Nandha Kumar",
    "New activity 4.14 Creating designations and mapping them to DSR Officers (UM-DSG-01 to 05; "
    "mapping is always one to one). Managing users and Moving citizens from Kaveri 2.0 renumbered to "
    "4.15 and 4.16. Report UM-RPT-12 and glossary entry for Designation added."])

d.save(DST)
print("Saved", DST)
