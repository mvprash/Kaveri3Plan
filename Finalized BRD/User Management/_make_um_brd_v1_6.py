import copy
import re

import docx
from docx.table import _Row
from docx.text.paragraph import Paragraph

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.5.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.6.docx"
DATE = "05-10-2026"
d = docx.Document(SRC)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def add_row_after(row, values):
    tr = copy.deepcopy(row._tr)
    row._tr.addnext(tr)
    new = _Row(tr, row._parent)
    for cell, text in zip(new.cells, values):
        set_text(cell.paragraphs[0], text)
    return new


removed = []
for cap in list(d.paragraphs):
    m = re.match(r"Figure 4\.(\d+): Process diagram", cap.text)
    if not m or int(m.group(1)) < 10:
        continue
    pic = cap._p.getprevious()
    assert pic.xpath(".//pic:pic") and not Paragraph(pic, cap._parent).text, cap.text
    pic.getparent().remove(pic)
    cap._p.getparent().remove(cap._p)
    removed.append(int(m.group(1)))
assert removed == list(range(10, 16)), removed

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.6")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)
add_row_after(d.tables[1].rows[-1], [
    "1.6", DATE, "Nandha Kumar",
    "Process diagrams removed from activities 4.10 to 4.15; diagrams remain for 4.1 to 4.9."])

d.save(DST)
print("Saved", DST)
