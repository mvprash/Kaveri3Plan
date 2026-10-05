import copy
import docx

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.1.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.2.docx"
DATE = "05-10-2026"
d = docx.Document(SRC)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find_row(first_cell):
    hits = [r for t in d.tables for r in t.rows if r.cells[0].text == first_cell]
    assert len(hits) == 1, (first_cell, len(hits))
    return hits[0]


def set_para(old, new):
    hits = [p for p in d.paragraphs if p.text == old]
    assert len(hits) == 1, (old, len(hits))
    set_text(hits[0], new)


set_para(
    "Application Admin is one of the roles under the DSR Officer user type. A DSR Officer with this role "
    "maintains roles, posts, sanctioned posts, offices and hierarchies.",
    "Roles, posts, sanctioned posts, offices and hierarchies are maintained by authorised administrators.")
set_para("Roles, and what each role can do, are maintained by the Application Admin.",
         "Roles, and what each role can do, are maintained by an authorised administrator.")
set_para("Posts and sanctioned strength are maintained by the Application Admin.",
         "Posts and sanctioned strength are maintained by an authorised administrator.")
set_para("The office hierarchy and the officer reporting hierarchy are maintained by the Application Admin.",
         "The office hierarchy and the officer reporting hierarchy are maintained by an authorised administrator.")

set_text(find_row("UM-ROL-02").cells[1].paragraphs[0],
         "Roles, posts and hierarchies are maintained only by authorised administrators.")
for req in ["UM-ROL-03", "UM-ROL-06", "UM-SAN-01", "UM-SAN-02", "UM-SAN-03", "UM-HIE-03"]:
    p = find_row(req).cells[1].paragraphs[0]
    assert p.text.startswith("The Application Admin "), req
    set_text(p, "An authorised administrator " + p.text[len("The Application Admin "):])
set_text(find_row("UM-USR-06").cells[1].paragraphs[0],
         "No second approval (maker-checker) is needed for user creation or master changes.")

term = find_row("Application Admin")
set_text(term.cells[0].paragraphs[0], "Authorised administrator")
set_text(term.cells[1].paragraphs[0],
         "A user given the authority to create users and to maintain roles, posts, sanctioned posts, "
         "offices and hierarchies.")

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.2")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)

history = d.tables[1]
last = history.rows[-1]
last._tr.addnext(copy.deepcopy(last._tr))
for cell, text in zip(history.rows[-1].cells, [
        "1.2", DATE, "Nandha Kumar",
        "Roles, posts and hierarchies are maintained by authorised administrators; the separate "
        "Application Admin role is no longer named (UM-ROL-02 reworded)."]):
    set_text(cell.paragraphs[0], text)

d.save(DST)
print("Saved", DST)
