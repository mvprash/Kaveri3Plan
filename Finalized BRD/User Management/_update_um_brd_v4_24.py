# -*- coding: utf-8 -*-
"""Create BRD_User_Management_v4.24.docx — add Citizen user data migration from Kaveri 2.0 to Kaveri 3.0 (Section 4.8, FR-UM-088–FR-UM-096)."""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

SRC = Path(r"Finalized BRD/User Management/BRD_User_Management_v4.23.docx")
DST = Path(r"Finalized BRD/User Management/BRD_User_Management_v4.24.docx")

DOC_VERSION = "1.1"
DOC_DATE = "03-10-2026"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell: _Cell, text: str) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(text)
        return
    set_paragraph_text(paragraphs[0], text)
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)


def strip_bookmarks(el) -> None:
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in el.iter(qn(tag)):
            bm.getparent().remove(bm)


def clone_paragraph(template: Paragraph, text: str):
    new_p = copy.deepcopy(template._p)
    strip_bookmarks(new_p)
    para = Paragraph(new_p, template._parent)
    set_paragraph_text(para, text)
    return new_p


def clone_table(template: Table, header: list[str] | None, rows: list[list[str]]):
    new_tbl = copy.deepcopy(template._tbl)
    strip_bookmarks(new_tbl)
    table = Table(new_tbl, template._parent)
    data_tr = table.rows[1]._tr
    for row in list(table.rows)[1:]:
        new_tbl.remove(row._tr)
    if header:
        for i, val in enumerate(header):
            set_cell_text(table.rows[0].cells[i], val)
    for values in rows:
        tr = copy.deepcopy(data_tr)
        new_tbl.append(tr)
        row = table.rows[-1]
        for i, val in enumerate(values):
            set_cell_text(row.cells[i], val)
    return new_tbl


def add_table_row_clone(table: Table, values: list[str]) -> None:
    new_tr = copy.deepcopy(table.rows[-1]._tr)
    table._tbl.append(new_tr)
    row = table.rows[-1]
    for i, val in enumerate(values):
        set_cell_text(row.cells[i], val)


def find_req_row(doc, req_id: str):
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text.strip() == req_id:
                return row
    raise KeyError(req_id)


def find_paragraph(doc, startswith: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def insert_after(anchor, elements) -> None:
    for el in elements:
        anchor.addnext(el)
        anchor = el


# ---------------------------------------------------------------------------
# Section 4.8 content
# ---------------------------------------------------------------------------

S48_INTRO = (
    "Citizen (Public user) accounts registered on Kaveri 2.0 shall be migrated into the single "
    "Kaveri 3.0 User Master so that existing citizens do not have to register again and their "
    "ownership of migrated Kaveri 2.0 applications, documents and payments is preserved. "
    "Migration is a one-time bulk load with mock runs and a final delta load at cutover. Kaveri "
    "2.0 passwords and security questions are not migrated because Kaveri 3.0 is passwordless "
    "(FR-UM-005) and does not use security questions (FR-UM-055 retired). Migrated citizens are "
    "loaded in a Pending Activation state and become Active only after a one-time first-login "
    "activation that brings them to the same identity standard as newly registered citizens — "
    "Aadhaar e-KYC (FR-UM-085) and verified mobile and email (FR-UM-063). Migration of DSR "
    "Officers and Other Department users is out of scope of this section; those users are "
    "created through Section 4.6."
)

S48_FRS = [
    [
        "FR-UM-088",
        "The system shall support migration of all Public user (Citizen) accounts from Kaveri 2.0 "
        "into the Kaveri 3.0 User Master with User Category = Public (Citizen). Migration shall be "
        "performed as a one-time bulk load followed by a final delta load at cutover, using the "
        "data mapping in 4.8.1. The following shall be migrated where present in Kaveri 2.0: "
        "Kaveri 2.0 User ID, login ID / username, citizen name, mobile number, email address, "
        "address, registration date, last login date, and account status. The following shall NOT "
        "be migrated: passwords or password hashes, security questions and answers, OTP and "
        "session logs, and any Aadhaar number stored in clear text. Migrated accounts shall be "
        "loaded with Account Status = Pending Activation (or Inactive per FR-UM-090) and with a "
        "Migration Source = Kaveri 2.0 flag and the migration batch ID.",
        "High",
    ],
    [
        "FR-UM-089",
        "The Kaveri 2.0 login ID shall be carried over as the Kaveri 3.0 Username when it satisfies "
        "the Kaveri 3.0 Username format rules and is unique across the whole User Master "
        "(FR-UM-004, FR-UM-062), including DSR Officer KGIDs and Other Department Usernames. Where "
        "the Kaveri 2.0 login ID is invalid in Kaveri 3.0 or collides with an existing Username, "
        "the account shall still be migrated, flagged Username Change Required, and the citizen "
        "shall choose a new preferred Username during first-login activation (FR-UM-091). Until "
        "then the citizen may log in using the Kaveri 2.0 login ID, which shall be held as the "
        "Legacy Kaveri 2.0 Username and shall remain searchable by administrators (FR-UM-021). "
        "Where two Kaveri 2.0 accounts collide with each other, the account with the most recent "
        "last login date shall retain the login ID and the other shall be flagged Username Change "
        "Required.",
        "High",
    ],
    [
        "FR-UM-090",
        "The system shall map Kaveri 2.0 account status to Kaveri 3.0 as follows: Active in "
        "Kaveri 2.0 → Pending Activation; Inactive, Blocked, Suspended or Deactivated in Kaveri "
        "2.0 → Inactive, with the Kaveri 2.0 status and reason retained for reference. Inactive "
        "migrated accounts shall not be able to log in and may be reactivated only by an "
        "authorised administrator with reason and audit trail (FR-UM-020); reactivation sets the "
        "account to Pending Activation. Accounts deleted in Kaveri 2.0 shall not be migrated.",
        "High",
    ],
    [
        "FR-UM-091",
        "A migrated citizen in Pending Activation shall log in with Username (or Legacy Kaveri 2.0 "
        "Username) + Captcha + OTP to the migrated mobile number (FR-UM-005, FR-UM-010). Before "
        "the home page is shown, the system shall require a one-time activation in which the "
        "citizen: (a) completes Aadhaar e-KYC (FR-UM-085); (b) verifies the migrated email address "
        "by OTP, or enters and verifies an email address by OTP where none was migrated or the "
        "migrated email was invalid (FR-UM-063); (c) chooses a new preferred Username where the "
        "account is flagged Username Change Required (FR-UM-062, FR-UM-089); (d) reviews and "
        "confirms the migrated profile details; and (e) accepts the Kaveri 3.0 privacy notice and "
        "consent for processing of personal data. The account shall become Active only after all "
        "applicable steps succeed; if activation is abandoned the account remains Pending "
        "Activation and activation restarts at the next login. Any Kaveri 2.0 name that differs "
        "from the Aadhaar e-KYC name shall be replaced by the e-KYC name, and the earlier name "
        "shall be retained in the audit trail.",
        "High",
    ],
    [
        "FR-UM-092",
        "A migrated citizen in Pending Activation who no longer has access to the migrated mobile "
        "number may use the Citizen lost-mobile reset (FR-UM-056) — Aadhaar e-KYC, then a PIN sent "
        "to the migrated email address, then OTP verification of the new mobile — and shall then "
        "continue first-login activation (FR-UM-091). Because the migrated email has not yet been "
        "verified in Kaveri 3.0, a successful PIN entry in this flow shall also mark the email as "
        "verified. Where the citizen has access to neither the migrated mobile nor the migrated "
        "email, the system shall not offer any other self-service path; assisted recovery is an "
        "open decision for the Department (Section 8).",
        "High",
    ],
    [
        "FR-UM-093",
        "Before loading, the migration process shall validate every Kaveri 2.0 citizen record. "
        "Records shall be rejected to the migration exception register (not loaded) when: the "
        "Kaveri 2.0 User ID is missing or duplicated; the mobile number is missing or not a valid "
        "10-digit Indian mobile number; or the citizen name is missing. Records shall be loaded "
        "but flagged when: the email address is missing or has an invalid format (email left "
        "blank and captured at activation — FR-UM-091); or the login ID needs to change "
        "(FR-UM-089). The same mobile number or email address appearing on more than one Kaveri "
        "2.0 account shall not be a rejection reason, because mobile and email are not unique in "
        "Kaveri 3.0 (FR-UM-004). Data cleansing rules (trimming, case normalisation of email, "
        "removal of country code from mobile) shall be applied consistently and documented.",
        "High",
    ],
    [
        "FR-UM-094",
        "The system shall store the Kaveri 2.0 User ID against each migrated citizen as an "
        "immutable Legacy Kaveri 2.0 User ID and shall maintain a cross-reference of Kaveri 2.0 "
        "User ID to Kaveri 3.0 User ID. This cross-reference shall be made available to the "
        "migration of other Kaveri modules (registration applications, documents, encumbrance "
        "searches, payments) so that every migrated transaction is linked to the correct Kaveri "
        "3.0 citizen account. The cross-reference shall not be editable by any user.",
        "High",
    ],
    [
        "FR-UM-095",
        "Migration runs shall be repeatable and idempotent, keyed on the Kaveri 2.0 User ID, so "
        "that re-running a batch updates rather than duplicates accounts. Each run shall be "
        "executed first as a mock run in a non-production environment, and shall produce a "
        "reconciliation report (Section 6) showing source count, loaded count, flagged count and "
        "rejected count, by Kaveri 2.0 status. At cutover, citizen registration and profile "
        "changes in Kaveri 2.0 shall be frozen, a final delta load shall be run, and go-live shall "
        "proceed only after the reconciliation has been signed off by Kaveri IT Cell and the "
        "Product Owner. An Application Admin shall be able to view the exception register, "
        "correct a rejected record and reprocess it, or close it with a reason; all such actions "
        "shall be audit-logged (FR-UM-022).",
        "High",
    ],
    [
        "FR-UM-096",
        "After go-live, the system shall notify each migrated citizen in Pending Activation by SMS "
        "to the migrated mobile, and by email where a valid email was migrated, in Kannada and "
        "English, that the account is available on Kaveri 3.0 and must be activated at first "
        "login. The notification shall show the Username to use (current or Legacy Kaveri 2.0 "
        "Username) and shall not contain any OTP, PIN, password or activation link. Notifications "
        "shall be sent in throttled batches so that SMS and email gateways and the login service "
        "are not overloaded.",
        "Medium",
    ],
]

S481_HEADER = ["Kaveri 2.0 data", "Kaveri 3.0 User Master field", "Migrated?", "Rule"]
S481_ROWS = [
    ["User ID", "Legacy Kaveri 2.0 User ID", "Yes", "Immutable; key for idempotent re-runs and cross-reference to other modules (FR-UM-094, FR-UM-095)"],
    ["Login ID / username", "Username (or Legacy Kaveri 2.0 Username)", "Yes", "Carried over if valid and unique; otherwise flagged Username Change Required (FR-UM-089)"],
    ["Citizen name", "Name", "Yes", "Mandatory; replaced by Aadhaar e-KYC name at activation if different (FR-UM-091)"],
    ["Mobile number", "Registered mobile", "Yes", "Mandatory, valid 10-digit Indian mobile; not unique (FR-UM-004, FR-UM-093)"],
    ["Email address", "Registered email", "Yes", "Loaded unverified; verified by OTP at activation; invalid email left blank (FR-UM-091, FR-UM-093)"],
    ["Address", "Address", "Yes", "Migrated as-is where present; citizen confirms at activation"],
    ["Registration date / last login date", "Legacy registration date / Legacy last login date", "Yes", "Reference only; last login date used to resolve login ID collisions (FR-UM-089)"],
    ["Account status", "Account Status", "Yes", "Mapped to Pending Activation or Inactive (FR-UM-090)"],
    ["Password / password hash", "—", "No", "Kaveri 3.0 is passwordless (FR-UM-005)"],
    ["Security questions and answers", "—", "No", "Not used in Kaveri 3.0 (FR-UM-055 retired)"],
    ["Aadhaar number (if stored)", "—", "No", "Aadhaar is captured only through e-KYC at activation (FR-UM-085)"],
    ["OTP / session / login logs", "—", "No", "Retained in Kaveri 2.0 archive per Department retention policy"],
]

S482_ROWS = [
    ["1", "Extract all Kaveri 2.0 citizen accounts into the migration staging area", "Kaveri IT Cell (migration tool)", "Encrypted extract; data stays in India (Section 5)"],
    ["2", "Apply cleansing and validation rules", "System", "Reject to exception register or load with flag (FR-UM-093)"],
    ["3", "Map fields, Username and status", "System", "Per 4.8.1, FR-UM-089 and FR-UM-090"],
    ["4", "Load into User Master as Pending Activation / Inactive; build Kaveri 2.0 → Kaveri 3.0 User ID cross-reference", "System", "Idempotent on Kaveri 2.0 User ID; batch ID recorded (FR-UM-094, FR-UM-095)"],
    ["5", "Generate reconciliation report and review exceptions", "Kaveri IT Cell / Application Admin", "Correct and reprocess, or close with reason (FR-UM-095)"],
    ["6", "Repeat steps 1–5 as mock runs until reconciliation is accepted", "Kaveri IT Cell", "Non-production environment"],
    ["7", "Cutover: freeze Kaveri 2.0 citizen registration and profile changes; run final delta load", "Kaveri IT Cell", "Delta updates existing accounts; no duplicates (FR-UM-095)"],
    ["8", "Sign off final reconciliation; go-live", "Kaveri IT Cell, Product Owner", "Go-live blocked until sign-off"],
    ["9", "Send activation notifications to migrated citizens", "System", "Throttled SMS / email, bilingual, no OTP or link (FR-UM-096)"],
]

S483_ROWS = [
    ["1", "Enter Username or Legacy Kaveri 2.0 Username + Captcha", "Citizen", "FR-UM-005, FR-UM-011, FR-UM-089"],
    ["2", "Verify OTP sent to migrated mobile", "System", "Lost mobile → FR-UM-056 via FR-UM-092"],
    ["3", "Detect Pending Activation; start activation", "System", "Home page not shown until activation completes (FR-UM-091)"],
    ["4", "Complete Aadhaar e-KYC", "Citizen / UIDAI", "FR-UM-085; e-KYC name replaces Kaveri 2.0 name if different"],
    ["5", "Verify migrated email by OTP, or enter and verify a new email", "Citizen", "FR-UM-063, FR-UM-091(b)"],
    ["6", "Choose new preferred Username (only if flagged Username Change Required)", "Citizen", "Availability checked across User Master (FR-UM-062, FR-UM-089)"],
    ["7", "Review and confirm profile; accept privacy notice and consent", "Citizen", "DPDP Act 2023 (Section 5)"],
    ["8", "Set Account Status = Active; audit-log; notify citizen", "System", "Subsequent logins follow FR-UM-005 normally"],
]

NFR_ROWS = [
    ["Security", "Kaveri 2.0 citizen data shall be extracted, transferred and staged in encrypted form, staging data shall be accessible only to named migration staff, and staging copies shall be purged after go-live sign-off per the Department retention policy (FR-UM-088–FR-UM-095)."],
    ["Performance", "Full citizen migration (bulk load) shall be completed within the cutover window agreed with Kaveri IT Cell; the final delta load at cutover shall not extend the planned Kaveri 2.0 downtime."],
    ["Data Integrity", "Migration reconciliation shall balance exactly: source count = loaded count + rejected count, with zero unexplained differences, before go-live (FR-UM-095)."],
]

REPORT_TEXT = (
    "The system shall provide a citizen migration report (FR-UM-088–FR-UM-096) showing, per "
    "migration batch: source, loaded, flagged and rejected counts by Kaveri 2.0 status; the "
    "exception register with rejection reason and resolution; accounts flagged Username Change "
    "Required; and, after go-live, activation progress — Pending Activation, Active and Inactive "
    "counts and activations per day."
)

UAT_TEXT = (
    "UAT — FR-UM-088–FR-UM-096: An active Kaveri 2.0 citizen is migrated as Pending Activation "
    "with no password or security questions; logs in with the migrated login ID + Captcha + OTP "
    "to the migrated mobile; is not shown the home page until Aadhaar e-KYC and email OTP "
    "verification are completed; and is then Active. A citizen whose login ID collides with a "
    "DSR KGID is prompted to choose a new Username at activation. Re-running a migration batch "
    "creates no duplicate accounts, and the reconciliation report balances (source = loaded + "
    "rejected)."
)

RISK_ROWS = [
    ["Kaveri 2.0 citizen data quality (invalid mobile, missing or wrong email, duplicate accounts)", "High", "Validation and cleansing rules (FR-UM-093); exception register with correction and reprocessing (FR-UM-095); repeated mock runs before cutover"],
    ["Migrated citizen has lost access to both migrated mobile and migrated email", "Medium", "FR-UM-092 covers mobile loss via email PIN; Department to decide an assisted-recovery policy (e.g. verified request at SRO) before go-live"],
    ["Low activation uptake or login surge after go-live", "Medium", "Throttled bilingual notifications (FR-UM-096); activation progress report (Section 6); capacity validated during performance testing"],
    ["Migrated transactions linked to the wrong citizen", "High", "Immutable Legacy Kaveri 2.0 User ID and read-only cross-reference used by all module migrations (FR-UM-094); reconciliation sign-off (FR-UM-095)"],
]

GLOSSARY_ROWS = [
    ["Pending Activation", "Account Status of a citizen migrated from Kaveri 2.0 who has not yet completed first-login activation (Aadhaar e-KYC, email verification, Username change if required); the home page is not available until activation is complete (FR-UM-091)"],
    ["Legacy Kaveri 2.0 User ID", "Immutable identifier of the citizen in Kaveri 2.0, stored on the migrated account and used to link migrated transactions from other modules (FR-UM-094)"],
    ["Legacy Kaveri 2.0 Username", "Kaveri 2.0 login ID retained on a migrated account; usable for login until a new Username is chosen when the account is flagged Username Change Required (FR-UM-089)"],
    ["Migration exception register", "List of Kaveri 2.0 citizen records rejected or flagged during migration validation, with reason and resolution (FR-UM-093, FR-UM-095)"],
]


def main() -> None:
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))
    t = doc.tables

    # Document control and version history
    for row in t[0].rows:
        if row.cells[0].text.strip() == "Version":
            set_cell_text(row.cells[1], DOC_VERSION)
        elif row.cells[0].text.strip() == "Last updated":
            set_cell_text(row.cells[1], DOC_DATE)
    add_table_row_clone(
        t[1],
        [
            DOC_VERSION,
            DOC_DATE,
            "Nandha Kumar",
            "Added Section 4.8 Citizen User Data Migration from Kaveri 2.0 (FR-UM-088–FR-UM-096): "
            "data mapping, Username and status mapping, validation and exception handling, "
            "Kaveri 2.0 → Kaveri 3.0 User ID cross-reference, mock runs, cutover and "
            "reconciliation, first-login activation with Aadhaar e-KYC, and activation "
            "notifications. Related updates to scope, NFRs, reports, acceptance criteria, risks "
            "and glossary.",
        ],
    )

    # Scope bullet
    scope_tpl = find_paragraph(doc, "Administrative user management, audit logging")
    scope_tpl._p.addnext(
        clone_paragraph(scope_tpl, "Migration of Citizen (Public user) accounts from Kaveri 2.0 to Kaveri 3.0, including first-login activation of migrated citizens (Section 4.8)")
    )

    # Cross-references in existing requirements
    row085 = find_req_row(doc, "FR-UM-085")
    set_cell_text(
        row085.cells[1],
        row085.cells[1].text
        + " Citizens migrated from Kaveri 2.0 shall complete Aadhaar e-KYC during first-login activation (FR-UM-091).",
    )
    for row in t[13].rows:
        if row.cells[0].text.startswith("Public users"):
            set_cell_text(
                row.cells[2],
                row.cells[2].text + "; migrated Kaveri 2.0 citizens complete first-login activation (FR-UM-091)",
            )

    # Section 4.8 after the 4.7 requirements table
    h2_tpl = find_paragraph(doc, "4.7 Administrative User Management")
    h3_tpl = find_paragraph(doc, "4.6.1 DSR Officer User Creation")
    body_tpl = find_paragraph(doc, "The system shall provide dedicated step-by-step workflows")
    fr_tpl = t[51]
    step_tpl = t[42]

    new_elements = [
        clone_paragraph(h2_tpl, "4.8 Citizen User Data Migration from Kaveri 2.0"),
        clone_paragraph(body_tpl, S48_INTRO),
        clone_table(fr_tpl, None, S48_FRS),
        clone_paragraph(h3_tpl, "4.8.1 Citizen Data Mapping — Kaveri 2.0 to Kaveri 3.0 (FR-UM-088–FR-UM-090, FR-UM-093)"),
        clone_table(step_tpl, S481_HEADER, S481_ROWS),
        clone_paragraph(h3_tpl, "4.8.2 Migration Run and Cutover (FR-UM-093–FR-UM-096)"),
        clone_table(step_tpl, None, S482_ROWS),
        clone_paragraph(h3_tpl, "4.8.3 First-Login Activation of Migrated Citizens (FR-UM-091, FR-UM-092)"),
        clone_table(step_tpl, None, S483_ROWS),
    ]
    insert_after(fr_tpl._tbl, new_elements)

    # NFRs
    for values in NFR_ROWS:
        add_table_row_clone(t[52], values)

    # Reports
    rpt_tpl = find_paragraph(doc, "The system shall provide an officer posting and service history report")
    rpt_tpl._p.addnext(clone_paragraph(rpt_tpl, REPORT_TEXT))

    # Acceptance criteria
    uat_tpl = find_paragraph(doc, "UAT — Role–Module–Function coverage")
    uat_tpl._p.addnext(clone_paragraph(uat_tpl, UAT_TEXT))

    # Risks and glossary
    for values in RISK_ROWS:
        add_table_row_clone(t[53], values)
    for values in GLOSSARY_ROWS:
        add_table_row_clone(t[54], values)

    doc.save(str(DST))
    print(f"Saved {DST}")


if __name__ == "__main__":
    main()
