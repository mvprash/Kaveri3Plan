import copy
import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches, Pt
from docx.table import _Row
from docx.text.paragraph import Paragraph
from PIL import Image

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.4.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.5.docx"
DIAGRAMS = Path(r"Finalized BRD\User Management\NewProcessDiagram")
DATE = "05-10-2026"
d = docx.Document(SRC)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def add_row_after(row, values):
    tr = copy.deepcopy(row._tr)
    row._tr.addnext(tr)
    new = _Row(tr, row._parent)
    for cell, text in zip(new.cells, values):
        set_text(cell.paragraphs[0], text)
    return new


def new_para_after(p):
    el = copy.deepcopy(p._p)
    for child in list(el):
        if not child.tag.endswith("pPr"):
            el.remove(child)
    p._p.addnext(el)
    return Paragraph(el, p._parent)


sec = d.sections[0]
max_w = sec.page_width - sec.left_margin - sec.right_margin
max_h = sec.page_height - sec.top_margin - sec.bottom_margin - Inches(1.2)

pngs = {int(m.group(1)): f for f in DIAGRAMS.glob("UM_4_*.png") if (m := re.match(r"UM_4_(\d+)_", f.name))}
done = []
for p in list(d.paragraphs):
    m = re.match(r"4\.(\d+) (.*)", p.text)
    if p.style.name != "Heading 2" or not m:
        continue
    num, title = int(m.group(1)), m.group(2)
    png = pngs[num]
    intro = Paragraph(p._p.getnext(), p._parent)
    assert intro.style.name == "Normal" and intro.text, p.text

    pic = new_para_after(intro)
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    w, h = Image.open(png).size
    scale = min(max_w / w, max_h / h)
    pic.add_run().add_picture(str(png), width=Emu(int(w * scale)), height=Emu(int(h * scale)))

    cap = new_para_after(pic)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(f"Figure 4.{num}: Process diagram — {title}")
    run.italic = True
    run.font.size = Pt(9)
    done.append(num)

assert sorted(done) == sorted(pngs) == list(range(1, 16)), (done, sorted(pngs))

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.5")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)
add_row_after(d.tables[1].rows[-1], [
    "1.5", DATE, "Nandha Kumar",
    "Process diagram added for each activity (4.1 to 4.15), shown under the activity's introduction. "
    "Editable draw.io files and PNGs are kept in the NewProcessDiagram folder."])

d.save(DST)
print("Saved", DST, "with", len(done), "diagrams")
