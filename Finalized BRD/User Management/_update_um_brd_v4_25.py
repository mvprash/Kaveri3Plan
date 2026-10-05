# -*- coding: utf-8 -*-
"""Create BRD_User_Management_v4.25.docx from v4.23 — Citizen user data migration from Kaveri 2.0 (Section 4.8, FR-UM-088–FR-UM-096).

Kaveri 2.0 uses the email ID as the username. Migrated fields: email ID (as Username), mobile
number, email, e-KYC status. Only accounts with e-KYC status = true are migrated.
"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

SRC = Path(r"Finalized BRD/User Management/BRD_User_Management_v4.23.docx")
DST = Path(r"Finalized BRD/User Management/BRD_User_Management_v4.25.docx")

DOC_VERSION = "1.2"
DOC_DATE = "03-10-2026"


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell: _Cell, text: str) -> None:
    paragraphs = cell.paragraphs
    if not paragraphs:
        cell.add_paragraph(text)
        return
    set_paragraph_text(paragraphs[0], text)
    for para in paragraphs[1:]:
        para._p.getparent().remove(para._p)


def strip_bookmarks(el) -> None:
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in list(el.iter(qn(tag))):
            bm.getparent().remove(bm)


def clone_paragraph(template: Paragraph, text: str):
    new_p = copy.deepcopy(template._p)
    strip_bookmarks(new_p)
    set_paragraph_text(Paragraph(new_p, template._parent), text)
    return new_p


def clone_table(template: Table, header: list[str] | None, rows: list[list[str]]):
    new_tbl = copy.deepcopy(template._tbl)
    strip_bookmarks(new_tbl)
    table = Table(new_tbl, template._parent)
    data_tr = table.rows[1]._tr
    for row in list(table.rows)[1:]:
        new_tbl.remove(row._tr)
    if header:
        for i, val in enumerate(header):
            set_cell_text(table.rows[0].cells[i], val)
    for values in rows:
        new_tbl.append(copy.deepcopy(data_tr))
        row = table.rows[-1]
        for i, val in enumerate(values):
            set_cell_text(row.cells[i], val)
    return new_tbl


def add_table_row_clone(table: Table, values: list[str]) -> None:
    table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
    row = table.rows[-1]
    for i, val in enumerate(values):
        set_cell_text(row.cells[i], val)


def find_req_row(doc, req_id: str):
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text.strip() == req_id:
                return row
    raise KeyError(req_id)


def find_paragraph(doc, startswith: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            return p
    raise KeyError(startswith)


def insert_after(anchor, elements) -> None:
    for el in elements:
        anchor.addnext(el)
        anchor = el


# ---------------------------------------------------------------------------
# Section 4.8 content
# ---------------------------------------------------------------------------

S48_INTRO = (
    "Citizen (Public user) accounts registered on Kaveri 2.0 shall be migrated into the single "
    "Kaveri 3.0 User Master so that existing e-KYC-verified citizens do not have to register "
    "again. In Kaveri 2.0 the citizen's email ID is the username. Only four fields are migrated "
    "— the email ID (as Username), mobile number, email, and e-KYC status — and only accounts "
    "whose Kaveri 2.0 e-KYC status is true are migrated. Citizens whose e-KYC status is not true "
    "are not migrated and shall register afresh on Kaveri 3.0 (FR-UM-001). Kaveri 2.0 passwords "
    "and security questions are not migrated because Kaveri 3.0 is passwordless (FR-UM-005) and "
    "does not use security questions (FR-UM-055 retired). Migrated accounts are loaded as "
    "Pending Activation and become Active after a one-time first-login activation. Migration is "
    "a one-time bulk load with mock runs and a final delta load at cutover. Migration of DSR "
    "Officers and Other Department users is out of scope of this section; those users are "
    "created through Section 4.6."
)

S48_FRS = [
    [
        "FR-UM-088",
        "The system shall migrate Public user (Citizen) accounts from Kaveri 2.0 into the Kaveri "
        "3.0 User Master with User Category = Public (Citizen). Only Kaveri 2.0 accounts whose "
        "e-KYC status is true shall be migrated; accounts whose e-KYC status is false, blank or "
        "any other value shall not be migrated and shall only be counted in the reconciliation "
        "(FR-UM-095). Exactly four fields shall be migrated, per the data mapping in 4.8.1: "
        "(1) Kaveri 2.0 username (email ID) as the Kaveri 3.0 Username; (2) mobile number; "
        "(3) email; and (4) e-KYC status. No other Kaveri 2.0 data shall be migrated — in "
        "particular passwords or password hashes, security questions and answers, Aadhaar number, "
        "and OTP or session logs. Migrated accounts shall be loaded with Account Status = Pending "
        "Activation, a Migration Source = Kaveri 2.0 flag, and the migration batch ID.",
        "High",
    ],
    [
        "FR-UM-089",
        "The Kaveri 2.0 username (email ID) shall become the Kaveri 3.0 Username of the migrated "
        "citizen, stored in lower case after trimming spaces. For migrated citizens this replaces "
        "the preferred-Username choice of FR-UM-062; the Username shall not be changeable by the "
        "citizen after migration (FR-UM-062), even if the citizen later changes the registered "
        "email address (FR-UM-013). The Username shall be unique across the whole User Master "
        "(FR-UM-004). A Kaveri 2.0 username that is not a valid email ID format, that duplicates "
        "another Kaveri 2.0 username when compared case-insensitively, or that is already used as "
        "a Kaveri 3.0 Username shall not be loaded and shall be placed in the migration exception "
        "register (FR-UM-093).",
        "High",
    ],
    [
        "FR-UM-090",
        "For each migrated account the system shall record the e-KYC status as Completed with "
        "e-KYC Source = Kaveri 2.0 and the migration date. This migrated status shall satisfy the "
        "Aadhaar e-KYC requirement of FR-UM-085 for that account, and Aadhaar e-KYC shall not be "
        "repeated during first-login activation (FR-UM-091). A fresh Aadhaar e-KYC shall still be "
        "performed whenever the citizen uses the lost-mobile reset (FR-UM-056).",
        "High",
    ],
    [
        "FR-UM-091",
        "A migrated citizen in Pending Activation shall log in with Username (Kaveri 2.0 email ID) "
        "+ Captcha + OTP to the migrated mobile number (FR-UM-005, FR-UM-010). Before the home "
        "page is shown, the system shall require a one-time activation in which the citizen: "
        "(a) verifies the registered email address by OTP (FR-UM-063); (b) reviews the migrated "
        "mobile number and email; and (c) accepts the Kaveri 3.0 privacy notice and consent for "
        "processing of personal data. The account shall become Active only after all steps "
        "succeed; if activation is abandoned the account remains Pending Activation and "
        "activation restarts at the next login. Activation shall be audit-logged.",
        "High",
    ],
    [
        "FR-UM-092",
        "A migrated citizen in Pending Activation who no longer has access to the migrated mobile "
        "number may use the Citizen lost-mobile reset (FR-UM-056) — fresh Aadhaar e-KYC, then a "
        "PIN sent to the migrated email, then OTP verification of the new mobile — and shall then "
        "continue first-login activation (FR-UM-091). A successful PIN entry in this flow shall "
        "also mark the email as verified.",
        "High",
    ],
    [
        "FR-UM-093",
        "Before loading, the migration process shall validate every Kaveri 2.0 citizen record "
        "whose e-KYC status is true. A record shall be rejected to the migration exception "
        "register (not loaded) when: the username is missing, is not a valid email ID format, or "
        "is a duplicate (FR-UM-089); or the mobile number is missing or not a valid 10-digit "
        "Indian mobile number. Where the email field is blank or invalid, the username email ID "
        "shall be used as the registered email. The same mobile number or email appearing on more "
        "than one account shall not be a rejection reason, because mobile and email are not "
        "unique in Kaveri 3.0 (FR-UM-004). Cleansing rules (trimming, lower-casing of email IDs, "
        "removal of country code from mobile) shall be applied consistently and documented.",
        "High",
    ],
    [
        "FR-UM-094",
        "The system shall store the Kaveri 2.0 username (email ID) against each migrated citizen "
        "as an immutable Legacy Kaveri 2.0 Username and shall maintain a cross-reference of Kaveri "
        "2.0 username to Kaveri 3.0 User ID. This cross-reference shall be made available to the "
        "migration of other Kaveri modules (registration applications, documents, payments) so "
        "that migrated transactions are linked to the correct Kaveri 3.0 citizen account. The "
        "cross-reference shall not be editable by any user.",
        "High",
    ],
    [
        "FR-UM-095",
        "Migration runs shall be repeatable and idempotent, keyed on the Kaveri 2.0 username, so "
        "that re-running a batch updates rather than duplicates accounts. Each run shall first be "
        "executed as a mock run in a non-production environment and shall produce a "
        "reconciliation report (Section 6) showing: total Kaveri 2.0 citizen accounts; accounts "
        "excluded because e-KYC status is not true; accounts rejected to the exception register; "
        "and accounts migrated. At cutover, citizen registration and profile changes in Kaveri "
        "2.0 shall be frozen, a final delta load shall be run, and go-live shall proceed only "
        "after the reconciliation has been signed off by Kaveri IT Cell and the Product Owner. An "
        "Application Admin shall be able to view the exception register, correct a rejected "
        "record and reprocess it, or close it with a reason; all such actions shall be "
        "audit-logged (FR-UM-022).",
        "High",
    ],
    [
        "FR-UM-096",
        "After go-live, the system shall notify each migrated citizen in Pending Activation by SMS "
        "to the migrated mobile and by email to the registered email, in Kannada and English, "
        "that the account is available on Kaveri 3.0, that the Kaveri 2.0 email ID is the "
        "Username, and that the account must be activated at first login. The notification shall "
        "not contain any OTP, PIN, password or activation link. Notifications shall be sent in "
        "throttled batches so that SMS and email gateways and the login service are not "
        "overloaded.",
        "Medium",
    ],
]

S481_HEADER = ["Kaveri 2.0 data", "Kaveri 3.0 User Master field", "Migrated?", "Rule"]
S481_ROWS = [
    ["Username (email ID)", "Username; Legacy Kaveri 2.0 Username", "Yes", "Trimmed and lower-cased; valid email format; unique; not changeable after migration (FR-UM-089, FR-UM-094)"],
    ["Mobile number", "Registered mobile", "Yes", "Mandatory, valid 10-digit Indian mobile; not unique (FR-UM-004, FR-UM-093)"],
    ["Email", "Registered email", "Yes", "Verified by OTP at activation; if blank or invalid, the username email ID is used (FR-UM-091, FR-UM-093)"],
    ["e-KYC status", "e-KYC Status = Completed (Source: Kaveri 2.0)", "Yes — only if true", "Accounts with e-KYC status not true are not migrated (FR-UM-088, FR-UM-090)"],
    ["Password / password hash", "—", "No", "Kaveri 3.0 is passwordless (FR-UM-005)"],
    ["Security questions and answers", "—", "No", "Not used in Kaveri 3.0 (FR-UM-055 retired)"],
    ["Aadhaar number and any other profile data", "—", "No", "Only the four fields above are migrated (FR-UM-088)"],
]

S482_ROWS = [
    ["1", "Extract all Kaveri 2.0 citizen accounts (username, mobile, email, e-KYC status) into the migration staging area", "Kaveri IT Cell (migration tool)", "Encrypted extract; data stays in India (Section 5)"],
    ["2", "Exclude accounts whose e-KYC status is not true", "System", "Counted as excluded in reconciliation (FR-UM-088, FR-UM-095)"],
    ["3", "Apply cleansing and validation rules", "System", "Reject to exception register (FR-UM-093)"],
    ["4", "Load into User Master as Pending Activation with e-KYC Completed; build Kaveri 2.0 username → Kaveri 3.0 User ID cross-reference", "System", "Idempotent on Kaveri 2.0 username; batch ID recorded (FR-UM-090, FR-UM-094, FR-UM-095)"],
    ["5", "Generate reconciliation report and review exceptions", "Kaveri IT Cell / Application Admin", "Correct and reprocess, or close with reason (FR-UM-095)"],
    ["6", "Repeat steps 1–5 as mock runs until reconciliation is accepted", "Kaveri IT Cell", "Non-production environment"],
    ["7", "Cutover: freeze Kaveri 2.0 citizen registration and profile changes; run final delta load", "Kaveri IT Cell", "Delta updates existing accounts; no duplicates (FR-UM-095)"],
    ["8", "Sign off final reconciliation; go-live", "Kaveri IT Cell, Product Owner", "Go-live blocked until sign-off"],
    ["9", "Send activation notifications to migrated citizens", "System", "Throttled SMS / email, bilingual, no OTP or link (FR-UM-096)"],
]

S483_ROWS = [
    ["1", "Enter Username (Kaveri 2.0 email ID) + Captcha", "Citizen", "FR-UM-005, FR-UM-011, FR-UM-089"],
    ["2", "Verify OTP sent to migrated mobile", "System", "Lost mobile → FR-UM-056 via FR-UM-092"],
    ["3", "Detect Pending Activation; start activation", "System", "Home page not shown until activation completes (FR-UM-091); no repeat e-KYC (FR-UM-090)"],
    ["4", "Verify registered email by OTP", "Citizen", "FR-UM-063, FR-UM-091(a)"],
    ["5", "Review migrated mobile and email; accept privacy notice and consent", "Citizen", "DPDP Act 2023 (Section 5)"],
    ["6", "Set Account Status = Active; audit-log; notify citizen", "System", "Subsequent logins follow FR-UM-005 normally"],
]

NFR_ROWS = [
    ["Security", "Kaveri 2.0 citizen data shall be extracted, transferred and staged in encrypted form, staging data shall be accessible only to named migration staff, and staging copies shall be purged after go-live sign-off per the Department retention policy (FR-UM-088–FR-UM-095)."],
    ["Performance", "Full citizen migration (bulk load) shall be completed within the cutover window agreed with Kaveri IT Cell; the final delta load at cutover shall not extend the planned Kaveri 2.0 downtime."],
    ["Data Integrity", "Migration reconciliation shall balance exactly — total Kaveri 2.0 citizen accounts = excluded (e-KYC not true) + rejected + migrated — with zero unexplained differences before go-live (FR-UM-095)."],
]

REPORT_TEXT = (
    "The system shall provide a citizen migration report (FR-UM-088–FR-UM-096) showing, per "
    "migration batch: total Kaveri 2.0 citizen accounts, accounts excluded because e-KYC status "
    "is not true, accounts rejected with reason and resolution, and accounts migrated; and, "
    "after go-live, activation progress — Pending Activation and Active counts and activations "
    "per day."
)

UAT_TEXT = (
    "UAT — FR-UM-088–FR-UM-096: A Kaveri 2.0 citizen with e-KYC status true is migrated as "
    "Pending Activation with the Kaveri 2.0 email ID as Username, the migrated mobile and email, "
    "and e-KYC Completed — with no password or security questions; logs in with the email ID + "
    "Captcha + OTP to the migrated mobile; is not asked to repeat Aadhaar e-KYC; is not shown "
    "the home page until email OTP verification and consent are completed; and is then Active. A "
    "Kaveri 2.0 citizen with e-KYC status false is not migrated and can self-register. Re-running "
    "a migration batch creates no duplicate accounts, and the reconciliation report balances "
    "(total = excluded + rejected + migrated)."
)

RISK_ROWS = [
    ["Kaveri 2.0 citizen data quality (invalid email ID usernames, case-variant duplicates, invalid mobile numbers)", "High", "Validation and cleansing rules (FR-UM-089, FR-UM-093); exception register with correction and reprocessing (FR-UM-095); repeated mock runs before cutover"],
    ["Citizens with e-KYC status not true lose their Kaveri 2.0 account and must register afresh", "Medium", "Communicate before cutover; excluded count reported (FR-UM-095); self-registration with Aadhaar e-KYC available (FR-UM-001, FR-UM-085)"],
    ["Migrated transactions linked to the wrong citizen", "High", "Immutable Legacy Kaveri 2.0 Username and read-only cross-reference used by all module migrations (FR-UM-094); reconciliation sign-off (FR-UM-095)"],
]

GLOSSARY_ROWS = [
    ["Pending Activation", "Account Status of a citizen migrated from Kaveri 2.0 who has not yet completed first-login activation (email OTP verification and consent); the home page is not available until activation is complete (FR-UM-091)"],
    ["Legacy Kaveri 2.0 Username", "The citizen's Kaveri 2.0 username (email ID), which becomes the Kaveri 3.0 Username and is retained immutably to link migrated transactions from other modules (FR-UM-089, FR-UM-094)"],
    ["Migration exception register", "List of e-KYC-verified Kaveri 2.0 citizen records rejected during migration validation, with reason and resolution (FR-UM-093, FR-UM-095)"],
]

HISTORY_ROWS = [
    [
        "1.1",
        DOC_DATE,
        "Nandha Kumar",
        "Added Section 4.8 Citizen User Data Migration from Kaveri 2.0 (FR-UM-088–FR-UM-096).",
    ],
    [
        DOC_VERSION,
        DOC_DATE,
        "Nandha Kumar",
        "Revised Section 4.8 to confirmed Kaveri 2.0 data: Kaveri 2.0 email ID is the username "
        "and becomes the Kaveri 3.0 Username; only username, mobile number, email and e-KYC "
        "status are migrated; only accounts with e-KYC status true are migrated and e-KYC is not "
        "repeated at first login. Related updates to scope, NFRs, reports, acceptance criteria, "
        "risks and glossary.",
    ],
]


def main() -> None:
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))
    t = doc.tables

    for row in t[0].rows:
        if row.cells[0].text.strip() == "Version":
            set_cell_text(row.cells[1], DOC_VERSION)
        elif row.cells[0].text.strip() == "Last updated":
            set_cell_text(row.cells[1], DOC_DATE)
    for values in HISTORY_ROWS:
        add_table_row_clone(t[1], values)

    scope_tpl = find_paragraph(doc, "Administrative user management, audit logging")
    scope_tpl._p.addnext(
        clone_paragraph(scope_tpl, "Migration of e-KYC-verified Citizen (Public user) accounts from Kaveri 2.0 to Kaveri 3.0, including first-login activation of migrated citizens (Section 4.8)")
    )

    row085 = find_req_row(doc, "FR-UM-085")
    set_cell_text(
        row085.cells[1],
        row085.cells[1].text
        + " Citizens migrated from Kaveri 2.0 with e-KYC status true are treated as e-KYC completed and are not asked to repeat e-KYC at first login (FR-UM-090).",
    )
    for row in t[13].rows:
        if row.cells[0].text.startswith("Public users"):
            set_cell_text(row.cells[1], row.cells[1].text + "; migrated Kaveri 2.0 citizens: Kaveri 2.0 email ID (FR-UM-089)")
            set_cell_text(row.cells[2], row.cells[2].text + "; migrated Kaveri 2.0 citizens complete first-login activation (FR-UM-091)")

    h2_tpl = find_paragraph(doc, "4.7 Administrative User Management")
    h3_tpl = find_paragraph(doc, "4.6.1 DSR Officer User Creation")
    body_tpl = find_paragraph(doc, "The system shall provide dedicated step-by-step workflows")
    fr_tpl = t[51]
    step_tpl = t[42]

    insert_after(
        fr_tpl._tbl,
        [
            clone_paragraph(h2_tpl, "4.8 Citizen User Data Migration from Kaveri 2.0"),
            clone_paragraph(body_tpl, S48_INTRO),
            clone_table(fr_tpl, None, S48_FRS),
            clone_paragraph(h3_tpl, "4.8.1 Citizen Data Mapping — Kaveri 2.0 to Kaveri 3.0 (FR-UM-088–FR-UM-090, FR-UM-093)"),
            clone_table(step_tpl, S481_HEADER, S481_ROWS),
            clone_paragraph(h3_tpl, "4.8.2 Migration Run and Cutover (FR-UM-093–FR-UM-096)"),
            clone_table(step_tpl, None, S482_ROWS),
            clone_paragraph(h3_tpl, "4.8.3 First-Login Activation of Migrated Citizens (FR-UM-091, FR-UM-092)"),
            clone_table(step_tpl, None, S483_ROWS),
        ],
    )

    for values in NFR_ROWS:
        add_table_row_clone(t[52], values)

    rpt_tpl = find_paragraph(doc, "The system shall provide an officer posting and service history report")
    rpt_tpl._p.addnext(clone_paragraph(rpt_tpl, REPORT_TEXT))

    uat_tpl = find_paragraph(doc, "UAT — Role–Module–Function coverage")
    uat_tpl._p.addnext(clone_paragraph(uat_tpl, UAT_TEXT))

    for values in RISK_ROWS:
        add_table_row_clone(t[53], values)
    for values in GLOSSARY_ROWS:
        add_table_row_clone(t[54], values)

    doc.save(str(DST))
    print(f"Saved {DST}")


if __name__ == "__main__":
    main()
