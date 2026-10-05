import copy

import docx
from docx.oxml.ns import qn
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.6.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.7.docx"
DATE = "05-10-2026"
TOTAL_W = 9637
d = docx.Document(SRC)
body = d.element.body


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def add_row_after(row, values):
    tr = copy.deepcopy(row._tr)
    row._tr.addnext(tr)
    new = _Row(tr, row._parent)
    for cell, text in zip(new.cells, values):
        set_text(cell.paragraphs[0], text)
    return new


anchor = next(p for p in d.paragraphs if p.style.name == "Heading 1" and p.text.startswith("2. "))
template_tbl = d.tables[2]._tbl
template_blank = next(p for p in d.paragraphs if p.style.name == "Normal" and not p.text)._p


def para(text, style):
    p = anchor.insert_paragraph_before(text, style=style)
    return p


def blank():
    anchor._p.addprevious(copy.deepcopy(template_blank))


def table(headers, rows, widths):
    widths = [round(TOTAL_W * w / sum(widths)) for w in widths]
    tbl = copy.deepcopy(template_tbl)
    head_tc = copy.deepcopy(tbl.tr_lst[0].tc_lst[0])
    body_tc = copy.deepcopy(tbl.tr_lst[1].tc_lst[1])
    head_tr = copy.deepcopy(tbl.tr_lst[0])
    body_tr = copy.deepcopy(tbl.tr_lst[1])
    for tr in list(tbl.tr_lst):
        tbl.remove(tr)
    grid = tbl.find(qn("w:tblGrid"))
    for gc in list(grid):
        grid.remove(gc)
    for w in widths:
        gc = grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(w)})
        grid.append(gc)

    def build(tr_template, tc_template, values):
        tr = copy.deepcopy(tr_template)
        for tc in list(tr.tc_lst):
            tr.remove(tc)
        for w, v in zip(widths, values):
            tc = copy.deepcopy(tc_template)
            tc.tcPr.find(qn("w:tcW")).set(qn("w:w"), str(w))
            tr.append(tc)
        tbl.append(tr)
        row = _Row(tr, Table(tbl, anchor._parent))
        for cell, v in zip(row.cells, values):
            set_text(cell.paragraphs[0], v)

    build(head_tr, head_tc, headers)
    for r in rows:
        build(body_tr, body_tc, r)
    anchor._p.addprevious(tbl)
    blank()


para("1.6 Legal and regulatory reference", "Heading 2")
para("User Management records the service actions taken on DSR Officers — appointment, posting, transfer, "
     "relieving, leave and charge — and keeps the records used for audit and investigation. These actions are "
     "governed by the Acts, rules and orders listed below.", "Normal")

para("i. Applicable Acts", "Heading 3")
table(["Act", "Why it applies to User Management"], [
    ("Constitution of India — Art. 309, 149, 151",
     "Art. 309: rules for recruitment and conditions of service of State employees. Art. 149 and 151: audit "
     "of State accounts by the Comptroller and Auditor-General (CAG)."),
    ("Karnataka State Civil Services Act, 1978",
     "Rules for categories of posts, pay scales, recruitment and conditions of service."),
    ("Registration Act, 1908",
     "Appointment of the Inspector-General of Registration, Registrars and Sub-Registrars, and of in-charge "
     "officers when a registering officer is absent or the office is vacant."),
    ("CAG's (Duties, Powers and Conditions of Service) Act, 1971",
     "Audit of the Department's receipts and expenditure by the Accountant General."),
    ("Karnataka Lokayukta Act, 1984", "Investigation of complaints against public servants."),
    ("Prevention of Corruption Act, 1988", "Criminal investigation of public servants by Lokayukta Police."),
    ("Income-tax Act, 2025 (Income-tax Act, 1961 for earlier years)",
     "Search, survey and information powers of the Income Tax Department."),
    ("Aadhaar Act, 2016", "Aadhaar e-KYC of citizens and biometric sign-in (UM-GEN-04)."),
    ("Digital Personal Data Protection Act, 2023", "Protection of users' personal data (UM-GEN-05)."),
], [4, 6])

para("ii. Relevant sections followed by the Department", "Heading 3")
table(["Act and section", "What it provides"], [
    ("Constitution — Art. 309 (proviso)",
     "The Governor may make rules on recruitment and conditions of service until a law is made."),
    ("KSCS Act, 1978 — Sec. 3(1) and 8",
     "The State Government may, by notification, make rules on categories of posts, pay scales, recruitment "
     "and conditions of service."),
    ("Registration Act, 1908 — Sec. 3(1)",
     "The State Government appoints the Inspector-General of Registration."),
    ("Registration Act, 1908 — Sec. 6",
     "The State Government appoints Registrars of districts and Sub-Registrars of sub-districts — the basis "
     "for posting orders of registering officers."),
    ("Registration Act, 1908 — Sec. 10(1)",
     "When a Registrar is absent (not on duty) or the office is temporarily vacant, a person appointed by the "
     "IGR acts as Registrar."),
    ("Registration Act, 1908 — Sec. 11",
     "A Registrar absent on duty in the district may appoint a Sub-Registrar or other person to do the "
     "Registrar's duties, except those under Sec. 68 and 72."),
    ("Registration Act, 1908 — Sec. 12",
     "When a Sub-Registrar is absent or the office is temporarily vacant, a person appointed by the District "
     "Registrar acts as Sub-Registrar."),
    ("CAG's Act, 1971 — Sec. 13, 16, 18",
     "Audit of expenditure (Sec. 13) and receipts (Sec. 16); power to inspect offices and call for records "
     "(Sec. 18)."),
    ("KL Act, 1984 — Sec. 7, 9, 10–12",
     "Matters that can be investigated, complaint procedure, power to call for documents and records, and "
     "reports to the competent authority."),
    ("PC Act, 1988 — Sec. 7, 13, 17, 17A, 19",
     "Offences by public servants, investigation by police, prior approval before inquiry, and sanction for "
     "prosecution."),
    ("Income-tax Act, 2025 — Sec. 247; 1961 Act — Sec. 131, 133, 133A",
     "Search and seizure; power to summon, call for information and inspect records."),
], [4, 6])

para("iii. Relevant rules followed by the Department", "Heading 3")
table(["Rule", "What it provides"], [
    ("KCSR Rule 8(49)",
     "Definition of transfer: movement of a Government servant from one headquarters station to another."),
    ("KCSR Rule 20-A", "A Government servant may be transferred from one post to another."),
    ("KCSR Rule 12",
     "Charge must be made over at the office headquarters, with both the relieving and the relieved officer "
     "present, unless the transferring authority orders otherwise."),
    ("KCSR Rule 23",
     "Pay begins when the officer takes charge; charge handed over in the forenoon or afternoon (12 noon is "
     "treated as forenoon)."),
    ("KCSR Rule 8(24) and Chapter VII", "Joining time on transfer and pay during joining time."),
    ("KCSR Chapter X",
     "Kinds of leave: Earned Leave (R.112), Half Pay Leave (R.114), Commuted Leave (R.116), Leave Not Due "
     "(R.117), Extraordinary Leave (R.118). Casual leave is governed by Government Orders."),
    ("KCSR Rule 68", "Full additional charge of another post, and the charge allowance."),
    ("KCSR Rule 32", "An officer may be put in charge of the current duties of a vacant post."),
    ("KCS (General Recruitment) Rules, 1977; Cadre and Recruitment Rules of the Department",
     "Posts, cadres and how they are filled."),
    ("KCS (CCA) Rules, 1957; KCS (Conduct) Rules, 2021", "Discipline, suspension and conduct of employees."),
    ("Regulations on Audit and Accounts, 2007",
     "Procedure for audit, Inspection Reports and the Department's replies."),
], [4, 6])

para("iv. Relevant notifications issued by the Department", "Heading 3")
table(["Order / circular", "What it says"], [
    ("IGR Circular No. ಸಿಬ್ಬಂದಿ-2-160/2024-25/48458 dated 21.12.2024",
     "On transfer, or on a court stay or other order, an officer must obtain a Movement Order from Head Office "
     "before reporting for duty. Head Office must know who took charge on joining and to whom charge was "
     "handed over on relieving. Non-compliance invites disciplinary action."),
    ("G.O. No. DPAR 22 SeVaNe 2013 dated 07.06.2013, and the orders for each year's general transfers",
     "Transfer period, limits and tenure for Group A to D employees."),
    ("DPAR Circular dated 10.07.2026",
     "A transfer order must also give the next posting. Waiting for a posting should not exceed one month "
     "(up to three months only with the Chief Minister's prior approval)."),
    ("Government Orders on casual leave", "Casual leave, which is not a kind of leave under KCSR."),
], [4, 6])

para("v. Service actions mapped to User Management activities", "Heading 3")
table(["Service action", "Governing Acts, rules and orders", "Where it is handled in this BRD"], [
    ("Employment / appointment", "Art. 309; KSCS Act Sec. 3, 8; Registration Act Sec. 3, 6; recruitment rules",
     "4.2 DSR Officer registration; 4.11 Sanctioned posts; 4.13 Designations"),
    ("Transfer", "KCSR R.8(49), R.20-A; DPAR transfer guidelines; IGR Circular 21.12.2024",
     "4.8 Assigning the DSR Officer to Posts — Transfer In (movement order, UM-ASG-02)"),
    ("Relieving", "KCSR R.12, R.23; IGR Circular 21.12.2024",
     "4.8 Assigning the DSR Officer to Posts — Transfer Out / Relieving"),
    ("Joining", "KCSR Chapter VII, R.8(24), R.23; IGR Circular 21.12.2024",
     "4.8 Assigning the DSR Officer to Posts — Transfer In"),
    ("Leave", "KCSR Chapter X; Government Orders on casual leave",
     "4.9 Temporary absence and temporary charge"),
    ("Additional charge", "KCSR R.68, R.32; Registration Act Sec. 10, 11, 12",
     "4.7 Additional charge; 4.9 Temporary charge"),
    ("Audit", "Art. 149, 151; CAG's Act Sec. 13, 16, 18; Regulations on Audit and Accounts, 2007",
     "Records (UM-GEN-06) and reports (Section 6)"),
    ("Investigation — Lokayukta", "KL Act Sec. 7, 9, 10–12; PC Act Sec. 7, 13, 17, 17A, 19",
     "Records (UM-GEN-06) and reports (Section 6)"),
    ("Investigation — Income Tax", "Income-tax Act, 2025 Sec. 247; 1961 Act Sec. 131, 133, 133A",
     "Records (UM-GEN-06) and reports (Section 6)"),
], [2.2, 4, 3.8])

para("Note: The wording above is a summary. KCSR Chapter VII and Chapter X, KCSR Rule 68, the Lokayukta and "
     "Prevention of Corruption Act sections, and the Income-tax Act numbering are to be checked against the "
     "latest official texts before they are cited in any order.", "Normal")
blank()

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.7")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)
add_row_after(d.tables[1].rows[-1], [
    "1.7", DATE, "Nandha Kumar",
    "New section 1.6 Legal and regulatory reference: applicable Acts, relevant sections, rules and Department "
    "notifications, and service actions mapped to User Management activities (based on KCSR Service Rules v3)."])

d.save(DST)
print("Saved", DST)
