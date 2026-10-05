import copy
import docx

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.0.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.1.docx"
DATE = "05-10-2026"
d = docx.Document(SRC)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find_row(first_cell):
    hits = [r for t in d.tables for r in t.rows if r.cells[0].text == first_cell]
    assert len(hits) == 1, (first_cell, len(hits))
    return hits[0]


rol4 = find_row("UM-ROL-04")
rol4._tr.getparent().remove(rol4._tr)

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.1")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)

history = d.tables[1]
last = history.rows[-1]
tr = copy.deepcopy(last._tr)
last._tr.addnext(tr)
new_row = history.rows[-1]
for cell, text in zip(new_row.cells, [
        "1.1", DATE, "Nandha Kumar",
        "DSR Officer registration revised (post optional at creation). Profile updates added as a separate "
        "activity. Transfer In and Relieving / Transfer Out triggered from the IGR office by uploading the "
        "movement order. UM-ROL-04 (list of divisions) removed."]):
    set_text(cell.paragraphs[0], text)

d.save(DST)
print("Saved", DST)
