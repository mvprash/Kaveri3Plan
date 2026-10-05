import copy
import re
import docx
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph

SRC = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.3.docx"
DST = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.4.docx"
DATE = "05-10-2026"
d = docx.Document(SRC)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def find_row(first_cell):
    hits = [r for t in d.tables for r in t.rows if r.cells[0].text == first_cell]
    assert len(hits) == 1, (first_cell, len(hits))
    return hits[0]


def find_para(text):
    hits = [p for p in d.paragraphs if p.text == text]
    assert len(hits) == 1, (text, len(hits))
    return hits[0]


def add_row_after(row, values):
    tr = copy.deepcopy(row._tr)
    row._tr.addnext(tr)
    new = _Row(tr, row._parent)
    for cell, text in zip(new.cells, values):
        set_text(cell.paragraphs[0], text)
    return new


def fill_table(table, rows):
    template = table.rows[1]._tr
    for r in list(table.rows[1:]):
        r._tr.getparent().remove(r._tr)
    for values in rows:
        table._tbl.append(copy.deepcopy(template))
        for cell, text in zip(table.rows[-1].cells, values):
            set_text(cell.paragraphs[0], text)


TITLE = "4.8 Assigning the DSR Officer to Posts"

rel_heading = find_para("4.9 Relieving / Transfer Out")._p
rel_block = [rel_heading]
for _ in range(3):
    rel_block.append(rel_block[-1].getnext())
assert rel_block[2].tag.endswith("tbl") and not Paragraph(rel_block[3], None).text
for el in rel_block:
    el.getparent().remove(el)

heading = find_para("4.8 Transfer In")
intro = Paragraph(heading._p.getnext(), heading._parent)
table = Table(intro._p.getnext(), heading._parent)
assert table.rows[1].cells[0].text == "UM-TIN-01"
set_text(heading, TITLE)
set_text(intro, "The superior assigns a DSR Officer to posts, or removes the officer from posts, after the IGR "
                "office has uploaded the movement order. The superior first selects the process: Transfer In or "
                "Transfer Out / Relieving. This is also how a post is given to an officer who was created "
                "without one.")
fill_table(table, [
    ("UM-ASG-01", "The superior selects the process for assigning: Transfer In (placing the officer in a post) "
                  "or Transfer Out / Relieving (removing the officer from one or more posts). The details asked "
                  "for depend on the process selected."),
    ("UM-ASG-02", "Both processes are triggered from the IGR office. An authorised user at the IGR office "
                  "uploads the movement order for the officer; only then can the superior carry out the selected "
                  "process. The movement order is linked to the Transfer In or Relieving."),
    ("UM-ASG-03", "A superior can act only on offices under them, and only on posts that report directly to the "
                  "superior's post. Example: the District Registrar of DRO Bengaluru can relieve the "
                  "Sub-Registrar of SRO Yeshwanthapura, but not of SRO Mysuru East; the IGR cannot relieve a "
                  "Sub-Registrar even though the SRO is visible."),
    ("UM-ASG-04", "Transfer In: allowed only when the post has a vacancy (occupied is less than sanctioned "
                  "strength). If the post is full, Transfer In is blocked."),
    ("UM-ASG-05", "Transfer In: the Transfer Order / Reporting Order number is required (with upload where "
                  "available). No Joining Date is entered."),
    ("UM-ASG-06", "Transfer In: takes effect immediately and the vacancy count is updated. The officer sees the "
                  "new post at the next sign-in."),
    ("UM-ASG-07", "Transfer In: a post whose holder is only on leave is not a vacancy."),
    ("UM-ASG-08", "Transfer Out / Relieving: Relieving Date, Relieving Order and Relieving Reason are required. "
                  "The reason must be one of: Deputation, Transfer, Suspension, Superannuation, Death."),
    ("UM-ASG-09", "Transfer Out / Relieving: the officer keeps the post until the end of the Relieving Date."),
    ("UM-ASG-10", "Transfer Out / Relieving: shortly after midnight the system removes the officer from the post "
                  "and updates the vacancy counts. If this fails, Kaveri IT Cell is alerted."),
    ("UM-ASG-11", "Transfer Out / Relieving: any open session the officer has on that post is ended."),
    ("UM-ASG-12", "Transfer Out / Relieving: if the officer has no post left, they cannot sign in until a new "
                  "post is given."),
])

codes = d.tables[2]
tin = next(r for r in codes.rows if r.cells[0].text == "TIN")
set_text(tin.cells[0].paragraphs[0], "ASG")
set_text(tin.cells[1].paragraphs[0], TITLE)
rel = next(r for r in codes.rows if r.cells[0].text == "REL")
rel._tr.getparent().remove(rel._tr)


def renumber(p):
    m = re.match(r"4\.(\d+) (.*)", p.text)
    if m and int(m.group(1)) >= 10:
        set_text(p, "4.%d %s" % (int(m.group(1)) - 1, m.group(2)))


for p in d.paragraphs:
    if p.style.name == "Heading 2":
        renumber(p)
for r in codes.rows:
    renumber(r.cells[1].paragraphs[0])

ids = find_para("Every requirement has a unique ID made of the module code (UM), an activity code and a running "
                "number — for example UM-TIN-02 is the second Transfer In requirement. IDs are never reused; a "
                "requirement that is withdrawn keeps its ID and is marked as withdrawn. The activity codes are:")
set_text(ids, ids.text.replace("UM-TIN-02 is the second Transfer In requirement",
                               "UM-ASG-02 is the second requirement for Assigning the DSR Officer to Posts"))
set_text(find_para("Transfer In, Relieving / Transfer Out, temporary absence and temporary charge"),
         "Assigning DSR Officers to posts (Transfer In and Transfer Out / Relieving), temporary absence and "
         "temporary charge")

dsr3 = find_row("UM-DSR-03").cells[1].paragraphs[0]
assert "through the separate posting activity (Transfer In)" in dsr3.text
set_text(dsr3, dsr3.text.replace("through the separate posting activity (Transfer In)",
                                 "through the separate activity Assigning the DSR Officer to Posts (Transfer In)"))

control = {r.cells[0].text: r for r in d.tables[0].rows}
set_text(control["Version"].cells[1].paragraphs[0], "1.4")
set_text(control["Last updated"].cells[1].paragraphs[0], DATE)
add_row_after(d.tables[1].rows[-1], [
    "1.4", DATE, "Nandha Kumar",
    "Transfer In and Relieving / Transfer Out merged into one activity, 4.8 Assigning the DSR Officer to Posts "
    "(code ASG); the superior selects the process. UM-TIN-01..06 and UM-REL-01..07 replaced by UM-ASG-01..12. "
    "Later sections renumbered 4.9 to 4.15."])

d.save(DST)
print("Saved", DST)
