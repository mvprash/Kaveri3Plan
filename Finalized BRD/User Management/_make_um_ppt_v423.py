# -*- coding: utf-8 -*-
"""Build BRD_User_Management_v4.23.pptx — a review deck for the KAVERI 3.0
User Management BRD.

Every process/structure diagram in this deck is drawn as native PowerPoint
shapes (numbered step chips, hierarchy trees, comparison cards) rather than
embedded pictures of the .drawio exports — editable, legible at a distance,
and free of the stale "Section 6.x" headers baked into the BRD's own raster
figures. Content is sourced from BRD_User_Management_v4.23.docx.
"""
from __future__ import annotations

import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

DST = "BRD_User_Management_v4.23.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

MARGIN = Inches(0.62)
CONTENT_W = SLIDE_W - 2 * MARGIN
BODY_TOP = Inches(1.72)
BODY_BOTTOM = Inches(6.85)

NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x1B, 0x6C, 0xA8)
GOLD = RGBColor(0xC8, 0x9B, 0x3C)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)
RULE = RGBColor(0xD4, 0xDD, 0xE6)
BAND = RGBColor(0xF2, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Segoe UI"


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def add_rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, line_w=0.75):
    box = slide.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = Pt(line_w)
    box.shadow.inherit = False
    return box


def add_line(slide, x1, y1, x2, y2, color=RULE, weight=1.25):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    conn.shadow.inherit = False
    return conn


def add_text(
    slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.0, wrap=True,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = wrap
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    para = frame.paragraphs[0]
    para.alignment = align
    para.line_spacing = spacing
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    return box


def add_bullets(slide, x, y, w, h, items, size=15, color=INK, spacing=1.18, space_after=9, marker="—"):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(items):
        if isinstance(item, tuple):
            lead, rest = item
        else:
            lead, rest = None, item
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.line_spacing = spacing
        para.space_after = Pt(space_after)
        if marker:
            tick = para.add_run()
            tick.text = f"{marker}  "
            tick.font.size = Pt(size)
            tick.font.bold = True
            tick.font.color.rgb = GOLD
            tick.font.name = FONT
        if lead:
            head = para.add_run()
            head.text = f"{lead}  "
            head.font.size = Pt(size)
            head.font.bold = True
            head.font.color.rgb = NAVY
            head.font.name = FONT
        run = para.add_run()
        run.text = rest
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = FONT
    return box


def set_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def node_box(slide, cx, cy, w, h, text, size=10.5, fill=BAND, line=RULE, bold=False, color=INK):
    x, y = cx - w / 2, cy - h / 2
    box = add_rect(slide, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.02
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    return box


# --------------------------------------------------------------------------
# slide scaffolding
# --------------------------------------------------------------------------
class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.number = 0

    def _bare(self):
        return self.prs.slides.add_slide(self.blank)

    def chrome(self, slide, title: str, kicker: str | None = None) -> None:
        self.number += 1
        if kicker:
            add_text(slide, MARGIN, Inches(0.46), CONTENT_W, Inches(0.26), kicker.upper(),
                     size=10.5, bold=True, color=BLUE)
            title_y = Inches(0.74)
        else:
            title_y = Inches(0.55)
        add_text(slide, MARGIN, title_y, CONTENT_W, Inches(0.55), title, size=27, bold=True, color=NAVY)
        add_rect(slide, MARGIN, Inches(1.46), Inches(1.15), Pt(3.2), fill=GOLD)
        add_rect(slide, MARGIN, Inches(7.02), CONTENT_W, Pt(0.9), fill=RULE)
        add_text(slide, MARGIN, Inches(7.12), Inches(9.0), Inches(0.26),
                 "KAVERI 3.0  ·  User Management BRD v4.23  ·  Department of Stamps and Registration, Government of Karnataka",
                 size=9, color=MUTED)
        add_text(slide, SLIDE_W - MARGIN - Inches(1.0), Inches(7.12), Inches(1.0), Inches(0.26),
                 str(self.number), size=9, bold=True, color=BLUE, align=PP_ALIGN.RIGHT)

    # ---- slide types ----
    def title_slide(self) -> None:
        slide = self._bare()
        add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
        add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, fill=GOLD)
        add_text(slide, Inches(1.15), Inches(1.62), Inches(11.0), Inches(0.34),
                 "KAVERI 3.0  ·  DEPARTMENT OF STAMPS AND REGISTRATION, GOVERNMENT OF KARNATAKA",
                 size=12, bold=True, color=GOLD)
        add_text(slide, Inches(1.15), Inches(2.30), Inches(11.9), Inches(1.0),
                 "Business Requirements Document", size=40, bold=True, color=WHITE)
        add_text(slide, Inches(1.15), Inches(3.35), Inches(11.9), Inches(0.8),
                 "User Management Module", size=32, color=RGBColor(0x9E, 0xC4, 0xE4))
        add_rect(slide, Inches(1.15), Inches(4.20), Inches(2.1), Pt(3.2), fill=GOLD)
        add_text(slide, Inches(1.15), Inches(4.62), Inches(11.0), Inches(0.9),
                 "Identity  ·  Passwordless authentication  ·  Aadhaar e-KYC  ·  Sanctioned post occupancy  ·  RBAC  ·  Officer lifecycle",
                 size=15, color=RGBColor(0xC8, 0xD7, 0xE6))
        add_text(slide, Inches(1.15), Inches(5.85), Inches(11.0), Inches(0.9),
                 "Document BRD-K3-UM-001   |   Version 4.23   |   08 September 2026\n"
                 "Author: Nandha Kumar (Business Analyst)   |   Status: In review — pending Domain Expert sign-off",
                 size=12.5, color=RGBColor(0xA9, 0xBC, 0xCE), spacing=1.35)
        set_notes(slide, "Walkthrough of the User Management BRD v4.23 for the KAVERI 3.0 platform. v4.23 is a "
                  "consistency-review pass over v4.18: it renumbers Functional Requirements as Section 4 (closing "
                  "the historical Section 2/5 gap) and folds in the authentication and lifecycle changes made in "
                  "v4.19-v4.22. Every diagram in this deck is redrawn as native shapes from the BRD's own process "
                  "and structure tables.")
        self.number += 1

    def divider(self, number: str, title: str, blurb: str) -> None:
        slide = self._bare()
        self.number += 1
        add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY)
        add_rect(slide, 0, 0, Inches(0.18), SLIDE_H, fill=GOLD)
        add_text(slide, Inches(1.15), Inches(2.55), Inches(2.0), Inches(1.3), number,
                 size=76, bold=True, color=RGBColor(0x1D, 0x4E, 0x80))
        add_text(slide, Inches(2.80), Inches(2.86), Inches(9.4), Inches(0.9), title, size=34, bold=True, color=WHITE)
        add_rect(slide, Inches(2.80), Inches(3.86), Inches(1.6), Pt(3.2), fill=GOLD)
        add_text(slide, Inches(2.80), Inches(4.20), Inches(9.15), Inches(1.0), blurb, size=14.5,
                 color=RGBColor(0xB6, 0xC9, 0xDA), spacing=1.28)
        add_text(slide, SLIDE_W - MARGIN - Inches(1.0), Inches(7.02), Inches(1.0), Inches(0.26),
                 str(self.number), size=9, bold=True, color=RGBColor(0x5E, 0x82, 0xA6), align=PP_ALIGN.RIGHT)

    def bullets_slide(self, title, kicker, items, lead_in=None, size=15) -> None:
        slide = self._bare()
        self.chrome(slide, title, kicker)
        top = BODY_TOP
        if lead_in:
            add_text(slide, MARGIN, top, CONTENT_W, Inches(0.6), lead_in, size=14, color=MUTED, spacing=1.25)
            top = top + Inches(0.72)
        add_bullets(slide, MARGIN, top, CONTENT_W, BODY_BOTTOM - top, items, size=size)
        return slide

    def two_column_slide(self, title, kicker, left_head, left_items, right_head, right_items) -> None:
        slide = self._bare()
        self.chrome(slide, title, kicker)
        gap = Inches(0.5)
        col_w = (CONTENT_W - gap) / 2
        for index, (head, items) in enumerate(((left_head, left_items), (right_head, right_items))):
            x = MARGIN + index * (col_w + gap)
            add_rect(slide, x, BODY_TOP, col_w, Inches(0.42), fill=NAVY)
            add_text(slide, x + Inches(0.18), BODY_TOP + Inches(0.09), col_w - Inches(0.3), Inches(0.28),
                     head, size=12.5, bold=True, color=WHITE)
            add_bullets(slide, x, BODY_TOP + Inches(0.66), col_w, BODY_BOTTOM - BODY_TOP - Inches(0.66),
                        items, size=13.5, spacing=1.16, space_after=8)
        return slide

    def metrics_slide(self, title, kicker, metrics, footnote=None) -> None:
        slide = self._bare()
        self.chrome(slide, title, kicker)
        gap = Inches(0.34)
        per_row = 4
        card_w = (CONTENT_W - gap * (per_row - 1)) / per_row
        card_h = Inches(1.92)
        for index, (value, label) in enumerate(metrics):
            row, col = divmod(index, per_row)
            x = MARGIN + col * (card_w + gap)
            y = BODY_TOP + Inches(0.18) + row * (card_h + gap)
            add_rect(slide, x, y, card_w, card_h, fill=BAND, line=RULE)
            add_rect(slide, x, y, card_w, Pt(3.4), fill=GOLD)
            add_text(slide, x + Inches(0.22), y + Inches(0.34), card_w - Inches(0.44), Inches(0.7),
                     value, size=36, bold=True, color=NAVY)
            add_text(slide, x + Inches(0.22), y + Inches(1.12), card_w - Inches(0.44), Inches(0.66),
                     label, size=11.5, color=MUTED, spacing=1.15)
        if footnote:
            add_text(slide, MARGIN, Inches(6.52), CONTENT_W, Inches(0.4), footnote, size=11.5, color=MUTED, italic=True)
        return slide

    def table_slide(self, title, kicker, headers, rows, col_widths, lead_in=None,
                    header_size=11.5, body_size=11, row_h=Inches(0.38)) -> None:
        slide = self._bare()
        self.chrome(slide, title, kicker)
        top = BODY_TOP
        if lead_in:
            add_text(slide, MARGIN, top, CONTENT_W, Inches(0.5), lead_in, size=13.5, color=MUTED, spacing=1.22)
            top = top + Inches(0.60)

        total = sum(col_widths)
        widths = [Inches(13.333 - 1.24) * (w / total) for w in col_widths]
        shape = slide.shapes.add_table(len(rows) + 1, len(headers), MARGIN, top, CONTENT_W, row_h * (len(rows) + 1))
        table = shape.table
        table.first_row = True
        table.horz_banding = False
        for index, width in enumerate(widths):
            table.columns[index].width = Emu(int(width))

        for index, head in enumerate(headers):
            cell = table.cell(0, index)
            cell.text = head
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(header_size)
            para.font.bold = True
            para.font.color.rgb = WHITE
            para.font.name = FONT

        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(value)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r % 2 else BAND
                cell.margin_left = Inches(0.1)
                cell.margin_right = Inches(0.1)
                cell.margin_top = Inches(0.04)
                cell.margin_bottom = Inches(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(body_size)
                para.font.color.rgb = INK
                para.font.name = FONT
                para.font.bold = c == 0 and len(headers) > 2
        return slide

    # ---- native diagram slide types ----
    def flow_slide(self, title, kicker, steps, key_text, lead_in=None, notes=None) -> None:
        """A numbered step-chip flow: gold circles above rounded boxes, connected
        by small gold arrows, wrapping to a second row when there are more than
        five steps, with a closing "Key characteristics" band."""
        slide = self._bare()
        self.chrome(slide, title, kicker)
        top = BODY_TOP
        if lead_in:
            add_text(slide, MARGIN, top, CONTENT_W, Inches(0.32), lead_in, size=12, italic=True, color=MUTED)
            top = top + Inches(0.40)

        numbered = list(enumerate(steps, start=1))
        if len(numbered) <= 5:
            rows = [numbered]
        else:
            split = -(-len(numbered) // 2)
            rows = [numbered[:split], numbered[split:]]

        cols = 5
        gap_x = Inches(0.20)
        box_w = (CONTENT_W - gap_x * (cols - 1)) / cols
        box_h = Inches(1.30) if len(rows) == 1 else Inches(1.05)
        circle_d = Inches(0.36)
        row_gap = Inches(0.26)

        row_y = top + circle_d / 2 + Inches(0.02)
        for row in rows:
            box_top = row_y + circle_d / 2
            for i, (num, label) in enumerate(row):
                x = MARGIN + i * (box_w + gap_x)
                cx = x + box_w / 2
                circ = add_rect(slide, cx - circle_d / 2, row_y, circle_d, circle_d, fill=GOLD, shape=MSO_SHAPE.OVAL)
                tf = circ.text_frame
                tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = str(num)
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = NAVY
                r.font.name = FONT
                node_box(slide, cx, box_top + box_h / 2, box_w, box_h, label, size=10.5)
                if i < len(row) - 1:
                    ax = x + box_w + Inches(0.01)
                    ay = box_top + box_h / 2 - Inches(0.08)
                    arrow = add_rect(slide, ax, ay, gap_x - Inches(0.02), Inches(0.16), fill=GOLD,
                                     shape=MSO_SHAPE.RIGHT_ARROW)
            row_y = box_top + box_h + row_gap

        kc_top = row_y + Inches(0.06)
        kc_h = BODY_BOTTOM - kc_top
        if kc_h < Inches(0.75):
            kc_h = Inches(0.75)
        add_rect(slide, MARGIN, kc_top, CONTENT_W, kc_h, fill=BAND, line=RULE)
        add_text(slide, MARGIN + Inches(0.22), kc_top + Inches(0.13), CONTENT_W - Inches(0.44), Inches(0.24),
                 "KEY CHARACTERISTICS", size=10, bold=True, color=NAVY)
        add_text(slide, MARGIN + Inches(0.22), kc_top + Inches(0.42), CONTENT_W - Inches(0.44), kc_h - Inches(0.54),
                 key_text, size=11, color=MUTED, spacing=1.2)
        if notes:
            set_notes(slide, notes)

    def chip_flow_slide(self, title, kicker, rows, footnote=None, notes=None) -> None:
        """Un-numbered chip rows for resolution/derivation logic (e.g. how a
        user category resolves to session claims), one labelled row per case."""
        slide = self._bare()
        self.chrome(slide, title, kicker)
        label_w = Inches(1.95)
        chips_x0 = MARGIN + label_w + Inches(0.14)
        avail_w = CONTENT_W - label_w - Inches(0.14)
        max_chips = max(len(r[1]) for r in rows)
        gap_x = Inches(0.16)
        chip_w = (avail_w - gap_x * (max_chips - 1)) / max_chips
        chip_h = Inches(0.62)
        row_h = Inches(0.92)
        y = BODY_TOP + Inches(0.15)
        for label, chips in rows:
            add_rect(slide, MARGIN, y, label_w, chip_h, fill=NAVY)
            add_text(slide, MARGIN + Inches(0.12), y, label_w - Inches(0.24), chip_h, label, size=11.5,
                     bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, spacing=1.05)
            for i, chip in enumerate(chips):
                cx = chips_x0 + i * (chip_w + gap_x)
                node_box(slide, cx + chip_w / 2, y + chip_h / 2, chip_w, chip_h, chip, size=10)
                if i < len(chips) - 1:
                    ax = cx + chip_w + Inches(0.01)
                    ay = y + chip_h / 2 - Inches(0.08)
                    add_rect(slide, ax, ay, gap_x - Inches(0.02), Inches(0.16), fill=GOLD, shape=MSO_SHAPE.RIGHT_ARROW)
            y = y + row_h
        if footnote:
            add_rect(slide, MARGIN, y + Inches(0.05), CONTENT_W, Inches(0.62), fill=BAND, line=RULE)
            add_text(slide, MARGIN + Inches(0.2), y + Inches(0.17), CONTENT_W - Inches(0.4), Inches(0.4),
                     footnote, size=11, color=MUTED, spacing=1.15)
        if notes:
            set_notes(slide, notes)

    def comparison_slide(self, title, kicker, card_a, card_b, footnote=None, lead_in=None, notes=None) -> None:
        """Two bordered cards side by side with a gold VS badge between them —
        used for two-rule comparisons (vacancy tests, span vs authority)."""
        slide = self._bare()
        self.chrome(slide, title, kicker)
        top = BODY_TOP
        if lead_in:
            add_text(slide, MARGIN, top, CONTENT_W, Inches(0.5), lead_in, size=13, color=MUTED, spacing=1.2)
            top = top + Inches(0.58)
        card_h = BODY_BOTTOM - top - (Inches(0.75) if footnote else Inches(0))
        gap = Inches(0.55)
        card_w = (CONTENT_W - gap) / 2

        for index, card in enumerate((card_a, card_b)):
            x = MARGIN + index * (card_w + gap)
            add_rect(slide, x, top, card_w, card_h, fill=WHITE, line=RULE, line_w=1.1)
            add_rect(slide, x, top, card_w, Inches(0.5), fill=NAVY)
            add_text(slide, x + Inches(0.2), top, card_w - Inches(0.4), Inches(0.5), card["title"], size=13,
                     bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
            ry = top + Inches(0.68)
            for row_label, row_text in card["rows"]:
                add_text(slide, x + Inches(0.24), ry, card_w - Inches(0.48), Inches(0.26), row_label.upper(),
                         size=9.5, bold=True, color=BLUE)
                add_text(slide, x + Inches(0.24), ry + Inches(0.28), card_w - Inches(0.48), Inches(0.62), row_text,
                         size=11.5, color=INK, spacing=1.16)
                ry = ry + Inches(0.86)

        badge_d = Inches(0.62)
        bx = MARGIN + card_w + gap / 2 - badge_d / 2
        by = top + card_h / 2 - badge_d / 2
        badge = add_rect(slide, bx, by, badge_d, badge_d, fill=GOLD, shape=MSO_SHAPE.OVAL)
        tf = badge.text_frame
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "VS"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY
        r.font.name = FONT

        if footnote:
            fy = top + card_h + Inches(0.16)
            add_rect(slide, MARGIN, fy, CONTENT_W, Inches(0.56), fill=BAND, line=RULE)
            add_text(slide, MARGIN + Inches(0.2), fy + Inches(0.1), CONTENT_W - Inches(0.4), Inches(0.36),
                     footnote, size=11, color=MUTED, spacing=1.1)
        if notes:
            set_notes(slide, notes)

    def tree_slide(self, title, kicker, levels, edges, footnote=None, notes=None) -> None:
        """A top-down hierarchy tree. `levels` is a list of rows, each row a
        list of (key, label, width_in) node tuples centred on that row's y.
        `edges` is a list of (parent_key, child_key) pairs drawn as straight
        connector lines between box edges."""
        slide = self._bare()
        self.chrome(slide, title, kicker)
        n_levels = len(levels)
        top = BODY_TOP + Inches(0.15)
        bottom = BODY_BOTTOM - (Inches(0.7) if footnote else Inches(0.1))
        level_h = Inches(0.56)
        gap = (bottom - top - level_h * n_levels) / max(1, n_levels - 1)
        centers: dict[str, tuple] = {}
        y = top + level_h / 2
        for row in levels:
            n = len(row)
            row_gap = Inches(0.35)
            total_w = sum(w for _, _, w in row) + row_gap * (n - 1)
            x = MARGIN + (CONTENT_W - total_w) / 2
            for key, label, w in row:
                cx = x + w / 2
                node_box(slide, cx, y, w, level_h, label, size=10.5, bold=(n_levels > 0 and key in [e[0] for e in edges] and False))
                centers[key] = (cx, y, w, level_h)
                x = x + w + row_gap
            y = y + level_h + gap

        for parent, child in edges:
            if parent not in centers or child not in centers:
                continue
            px, py, pw, ph = centers[parent]
            cx, cy, cw, ch = centers[child]
            add_line(slide, px, py + ph / 2, cx, cy - ch / 2, color=RULE, weight=1.5)

        if footnote:
            fy = bottom + Inches(0.12)
            add_rect(slide, MARGIN, fy, CONTENT_W, Inches(0.56), fill=BAND, line=RULE)
            add_text(slide, MARGIN + Inches(0.2), fy + Inches(0.1), CONTENT_W - Inches(0.4), Inches(0.36),
                     footnote, size=11, color=MUTED, spacing=1.1)
        if notes:
            set_notes(slide, notes)


# --------------------------------------------------------------------------
# deck content
# --------------------------------------------------------------------------
def build() -> Presentation:
    deck = Deck()

    deck.title_slide()

    deck.bullets_slide(
        "Agenda", "Contents",
        [
            ("1.", "Introduction, scope and business objectives"),
            ("2.", "Identity and passwordless authentication — Citizens, DSR Officers, Other Department users"),
            ("3.", "Posts, roles and role-based access control"),
            ("4.", "DSR Officer lifecycle — creation, transfer, occupancy and temporary absence"),
            ("5.", "Non-functional requirements, reporting, acceptance and sign-off"),
        ],
        lead_in="A walkthrough of BRD-K3-UM-001 v4.23. Every diagram in this deck is redrawn from the BRD's own "
                "process step tables and structure rules — not a picture of the .drawio export.",
        size=16,
    )

    deck.table_slide(
        "Document control", "BRD-K3-UM-001",
        ["Field", "Value"],
        [
            ["Document ID", "BRD-K3-UM-001"],
            ["Version", "4.23  (08 September 2026)"],
            ["Status", "In review — pending Domain Expert sign-off"],
            ["Module", "User Management"],
            ["Author (BA)", "Nandha Kumar"],
            ["Product Owner", "Prashanth"],
            ["Domain expert / reviewer", "Prabhakar Naik"],
            ["Target audience", "Kaveri IT Cell, Department of Stamps and Registration, Government of Karnataka"],
            ["Legal basis (primary)", "IT Act 2000; Indian Registration Act 1908 (Sub-Registrar appointment); Aadhaar Act 2016"],
            ["State rules (primary)", "Karnataka e-Governance hosting/security norms; MeitY / CERT-In / STQC / GIGW; GOs for office/post creation"],
        ],
        col_widths=[3.0, 9.1], row_h=Inches(0.44),
    )

    deck.bullets_slide(
        "What's changed since v4.18", "Revision history — v4.19 to v4.23",
        [
            ("Authentication —", "DSR Officers now authenticate with Username (KGID) + Captcha + Face authentication or "
             "Biometrics only — no OTP (FR-UM-006). Other Department users drop biometrics and stay OTP-only (FR-UM-007)."),
            ("Identity proofing —", "Aadhaar e-KYC is mandatory for Citizen self-registration and lost-mobile reset "
             "(FR-UM-085); the five security questions are retired (FR-UM-055)."),
            ("Contact recovery —", "DSR Officers can now change their own registered mobile after login, OTP-verified, "
             "with no administrator involved (FR-UM-086). Other Department mobile changes remain administrator-only."),
            ("Usernames —", "Other Department Username is now Department Code + Employee ID/KGID, not KGID alone "
             "(FR-UM-062, FR-UM-064)."),
            ("Session policy tightened —", "OTP resend cooldown 60s (was 30s); idle timeout 10 minutes (was 15); "
             "absolute session 4 hours (was 8) — FR-UM-072, FR-UM-074, FR-UM-075."),
            ("Transfer In simplified —", "Joining Date and the reserved/future-dated path are retired (FR-UM-061, "
             "FR-UM-067); Transfer In now takes effect immediately once capacity is available (FR-UM-060)."),
            ("Section numbering —", "v4.23 closes the historical Section 2/5 gap: Functional Requirements move from "
             "Section 6 to Section 4; NFR / Reporting / Acceptance / Risks / Glossary / Approval renumber 5-10."),
        ],
        size=13.5,
    )

    # ---------------- Section 1 ----------------
    deck.divider("01", "Introduction and Scope", "Why the module exists, what it replaces, and where its boundaries lie.")

    deck.two_column_slide(
        "Purpose and background", "Sections 1.1 – 1.2",
        "PURPOSE",
        [
            "Defines the business requirements for the KAVERI 3.0 User Management module.",
            "Covers identity, passwordless authentication, sanctioned post occupancy, RBAC and officer lifecycle.",
            "Applies to Citizens, DSR Officers and Other Department users.",
            "Is the agreed basis for design, development, testing and sign-off.",
        ],
        "BACKGROUND",
        [
            "KAVERI 3.0 needs one platform service for users, roles, posts and module access.",
            "Replaces fragmented Kaveri 2.0 user administration.",
            "One User Master and one Role Master — no separate stores per category.",
            "No passwords anywhere; OTP for Citizens/Other Department, Face/Biometrics for DSR Officers; Application "
            "Admin maintains privilege mapping.",
        ],
    )

    deck.bullets_slide(
        "Scope", "Section 1.3",
        [
            "User registration and profile management for Citizens, DSR Officers and Other Department users (single User Master).",
            "Passwordless authentication, session policy, Citizen lost-mobile reset, DSR post selection and additional charge (§4.2–4.5).",
            "Unified Role Master; Posts and Sanctioned Posts masters; office and officer hierarchies; RBAC via Module, Function and Resource mapping (§4.5).",
            "DSR lifecycle: post assignment, Transfer Out / In, occupancy refresh, Temporary Absence and Temporary Charge (§4.6).",
            "Administrative user management, audit logging and reporting (§4.7–6).",
        ],
        lead_in="In scope for the User Management module:", size=15,
    )

    deck.bullets_slide(
        "Business objectives", "Section 2",
        [
            "Provide a secure and reliable mechanism for users to register and access the system.",
            "Enable administrators to manage user accounts and access rights efficiently.",
            "Enforce role-based access control to protect sensitive data and functionality.",
            "Maintain the DSR Officer Hierarchy Master aligned to the departmental organisation chart.",
            "Support Module Function and Resource masters with Role–Module–Function mapping and runtime enforcement.",
            "Reduce support overhead on account access issues through reliable OTP delivery.",
            "Ensure compliance with organisational security and data privacy policies.",
            "Provide auditability of user actions for security and compliance reporting.",
        ],
        size=14.5,
    )

    deck.metrics_slide(
        "The module at a glance", "Scale of the specification",
        [
            ("82", "Active functional requirements (5 formally retired)"),
            ("3", "User categories in a single User Master"),
            ("19", "Approved process and structure diagrams"),
            ("8", "Business modules under the Module Master"),
            ("40", "Posts in the Posts Master across 8 divisions"),
            ("56", "Seed roles — 6 Citizen, 10 Other Department, 40 DSR"),
            ("0", "Passwords stored — OTP / Face-Biometric only"),
            ("7 yrs", "Minimum audit-log retention"),
        ],
        footnote="Counts taken from BRD v4.23. FR-UM-001–FR-UM-087 issued; FR-UM-015, 031, 055, 061 and 067 are "
                 "formally retired. Detailed requirement text remains authoritative in Section 4 of the document.",
    )

    deck.table_slide(
        "Stakeholders", "Section 3",
        ["Name / Role", "Department", "Responsibility"],
        [
            ["Kaveri IT Cell", "Engineering", "Reviews technical feasibility"],
            ["Citizens (Public users)", "External", "Self-register and access citizen portal services"],
            ["DSR Officers & Other Department users", "Government",
             "Access departmental modules via OTP (Other Department) or Captcha + Face/Biometric authentication (DSR Officers — no OTP)"],
        ],
        lead_in="Product Owner (Prashanth) and Domain Expert (Prabhakar Naik) sign off the document — see Document control.",
        col_widths=[3.6, 2.4, 6.1], row_h=Inches(0.6),
    )

    # ---------------- Section 2 ----------------
    deck.divider("02", "Identity and Authentication", "One User Master, three categories, and no passwords anywhere in the platform.")

    deck.table_slide(
        "Three user categories, one User Master", "Section 4.5.2",
        ["User category", "Username", "Authentication", "Lost-mobile / contact reset"],
        [
            ["Public (Citizen)", "Preferred Username chosen at registration (FR-UM-062)",
             "Username + Captcha + OTP to mobile",
             "Aadhaar e-KYC + PIN to registered email + OTP to new mobile — no security questions (FR-UM-056, FR-UM-085)"],
            ["DSR Officer", "KGID (FR-UM-062, FR-UM-064)",
             "Username (KGID) + Captcha + Face authentication or Biometrics — no OTP (FR-UM-006), then post selection (FR-UM-052)",
             "Self-service after login — OTP-verified mobile change, no administrator needed (FR-UM-086)"],
            ["Other Department", "Department Code + Employee ID/KGID (FR-UM-062, FR-UM-064)",
             "Username + Captcha + OTP to mobile — no biometrics (FR-UM-007)",
             "Not available — administrator changes the mobile with reason and audit (FR-UM-065)"],
        ],
        lead_in="Differentiation is by User Category on users and Role Category on roles — there are no separate masters per category.",
        col_widths=[1.9, 2.7, 3.7, 3.9], body_size=10.5, row_h=Inches(1.05),
    )

    deck.chip_flow_slide(
        "How identity resolves to access", "Section 4.5.1  ·  S-01 Access Model",
        [
            ("PUBLIC (CITIZEN)", ["User Category = Citizen", "Citizen role(s) from Role Master", "Session claims"]),
            ("DSR OFFICER", ["User Category = DSR", "Post occupancy (Post + Office)", "Post–Role mapping — FR-UM-052 selects one post", "Session claims"]),
            ("OTHER DEPARTMENT", ["User Category = Other Dept", "Exactly one Role — FR-UM-029, FR-UM-034", "Session claims"]),
        ],
        footnote="The Username is the single unique login identifier across the whole User Master (FR-UM-004, FR-UM-062) — "
                 "email and mobile carry no uniqueness constraint. Password-based login is out of scope (FR-UM-009).",
        notes="One User Master and one Role Master. Category drives the creation path, the authentication factors and "
              "the eligible roles.",
    )

    deck.flow_slide(
        "Citizen self-registration", "Section 4.1.1  ·  P-01 Citizen Self-Registration  ·  FR-UM-001, FR-UM-062, FR-UM-063, FR-UM-085",
        [
            "Open registration; enter name — instant self-registration",
            "Choose preferred Username; system checks availability across the User Master",
            "Enter email address and mobile number",
            "System dispatches a separate OTP to the email and to the mobile",
            "Enter both OTPs — email and mobile",
            "Complete mandatory Aadhaar e-KYC",
            "Save — account created, Citizen role assigned; no security questions anywhere",
        ],
        key_text="No approval workflow. Email and mobile are each OTP-verified and Aadhaar e-KYC must also succeed "
                  "before the account is created — the five security questions are retired (FR-UM-055).",
    )

    deck.flow_slide(
        "Login across all categories", "Section 4.5.2.1  ·  P-02 Login (all categories)  ·  FR-UM-004–FR-UM-007, FR-UM-010, FR-UM-011, FR-UM-062",
        [
            "Enter Username + Captcha",
            "Validate Captcha; look up the account by Username",
            "Dispatch login OTP to the registered mobile — Citizens and Other Department only",
            "Citizen/Other Dept enter OTP; DSR completes Face authentication or Biometrics — no OTP",
            "Citizen who cannot receive the OTP leaves login for lost-mobile reset",
            "On success, enter home under the assigned post (DSR selects post first — FR-UM-052)",
            "After login, Citizen or DSR may update their mobile from profile",
        ],
        key_text="The account resolves from the Username alone — no category selector is shown. DSR Officers "
                  "authenticate with Captcha + Face/Biometrics only; the login OTP is never sent to email (FR-UM-010).",
    )

    deck.flow_slide(
        "Citizen lost-mobile reset — Aadhaar e-KYC plus PIN", "Section 4.5.2.2  ·  P-03 Citizen Lost-Mobile Reset  ·  FR-UM-056, FR-UM-085",
        [
            "Choose “Lost / changed mobile number” — Citizens only",
            "Enter Username + Captcha",
            "System initiates Aadhaar e-KYC for identity verification",
            "Complete Aadhaar e-KYC successfully",
            "System sends a single-use, time-limited PIN to the registered email",
            "Enter the PIN",
            "Enter the new mobile number",
            "System sends an OTP to the new mobile; enter it",
            "Mobile updated; registered email notified; audit logged",
            "Continue login with OTP to the new mobile",
        ],
        key_text="Two independent proofs — a successful Aadhaar e-KYC and a single-use email PIN — are required before "
                  "the mobile may change; the new number is itself OTP-verified. Security questions are retired.",
    )

    deck.flow_slide(
        "Mobile number change", "Section 4.5.2.3  ·  P-04 Departmental Mobile Change  ·  FR-UM-065, FR-UM-086",
        [
            "User reports a lost/changed mobile — DSR Officers use self-service instead",
            "Administrator locates the Other Department user by Username",
            "Administrator enters the new mobile number and a mandatory reason",
            "System verifies the new number by OTP before saving",
            "Save — mobile updated; audit-logged; user notified",
            "User logs in with the updated mobile (Other Dept: OTP · DSR: Face/Biometrics)",
        ],
        key_text="DSR Officers can now change their own mobile after login — OTP-verified, no administrator required "
                  "(FR-UM-086). Other Department users still have no self-service path (FR-UM-065).",
    )

    deck.table_slide(
        "OTP and session policy", "FR-UM-069 – FR-UM-076",
        ["Control", "Rule", "Requirement"],
        [
            ["Login OTP validity", "5 minutes from dispatch (IST)", "FR-UM-069"],
            ["Registration OTP / reset PIN validity", "10 minutes from dispatch", "FR-UM-069"],
            ["Code length", "6 numeric digits", "FR-UM-070"],
            ["Incorrect entries per code", "Maximum 3, then the code is invalidated", "FR-UM-071"],
            ["Resend cooldown", "60 seconds; maximum 3 resends per channel per 15 minutes", "FR-UM-072"],
            ["Failed-login lockout", "5 failures lock the Username for 15 minutes", "FR-UM-073"],
            ["Idle timeout", "10 minutes of inactivity ends the session", "FR-UM-074"],
            ["Absolute session limit", "4 hours from login, even if the user is active", "FR-UM-075"],
            ["Concurrent sessions", "One per Username — last login wins", "FR-UM-076"],
        ],
        lead_in="Applies to every user category. The 10-minute idle timeout and 4-hour session cap matter most on shared SRO counter machines, where the session post is fixed for the whole session.",
        col_widths=[3.5, 6.6, 2.0], row_h=Inches(0.4),
    )

    # ---------------- Section 3 ----------------
    deck.divider("03", "Posts, Roles and Access Control", "Access follows the post an officer occupies — not a privilege attached to the person.")

    deck.flow_slide(
        "DSR Officer post selection at login", "Section 4.5.2.4  ·  P-05 DSR Login Post Selection  ·  FR-UM-052",
        [
            "Authenticate — Username (KGID) + Captcha + Face/Biometrics, no OTP",
            "Load active post occupancies (relieved posts excluded)",
            "One active post auto-selects; two or more show a mandatory list",
            "Each option labelled Role — Post Name — Office Name (Office Code)",
            "Officer selects exactly one assigned post and confirms",
            "Module Function claims resolve from the assigned post; enter home",
            "Home/header displays the assigned Post with mapped Role(s)",
        ],
        key_text="Mandatory only when more than one occupancy is active. Claims derive from the selected post alone "
                  "via Post–Role mapping — never from the union of other posts the officer holds.",
    )

    deck.flow_slide(
        "Additional charge of an unoccupied subordinate post", "Section 4.5.2.5  ·  P-06 Additional Charge After Login  ·  FR-UM-053, FR-UM-054",
        [
            "Officer is logged in under the assigned Post + Office",
            "Opens “Additional charge” from the header — post-login only, no logout",
            "System reads Hierarchy Master children of the assigned post",
            "Lists wholly unoccupied subordinate posts (cascade rule)",
            "Shows Role + Post options, plus “clear additional charge”",
            "Officer selects a post, switches back, or clears — one context at a time",
            "Module Function claims recompute for the active context only",
            "Header updates to show assigned post + active context",
            "Selection or reversion is audit-logged",
        ],
        key_text="A subordinate post qualifies only when wholly unoccupied (Occupied = 0); any occupant blocks both "
                  "the post and the cascade beneath it. A Sub-Registrar acting as FDA loses SR signing until switch-back.",
    )

    deck.comparison_slide(
        "Two distinct vacancy tests", "Section 4.5.3  ·  S-02 Two Vacancy Tests  ·  FR-UM-066",
        {
            "title": "(a)  Available capacity",
            "rows": [
                ("Rule", "Occupied < Sanctioned Strength"),
                ("Used for", "Post assignment at user creation, and Transfer In"),
                ("Example", "Strength 2, Occupied 1 → has available capacity"),
            ],
        },
        {
            "title": "(b)  Wholly unoccupied",
            "rows": [
                ("Rule", "Occupied = 0, irrespective of sanctioned strength"),
                ("Used for", "Post-login additional charge only (FR-UM-053)"),
                ("Example", "Strength 2, Occupied 1 → NOT wholly unoccupied"),
            ],
        },
        footnote="Occupied includes active occupancies only. The reserved/future-dated Transfer In path (FR-UM-061, "
                 "FR-UM-067) is retired — Transfer In now requires available capacity at the time of recording.",
    )

    deck.table_slide(
        "Role Master — DSR roles by division", "Sections 4.5.3 – 4.5.4",
        ["Division (FR-UM-077)", "Roles under Role Category = DSR"],
        [
            ["Secretariat", "ACS / Principal Secretary / Secretary"],
            ["Top Management", "IGR"],
            ["Admin, Law & Computers", "DIGR (Admin, Law & Computers), AIGR (Admin), HQA (Admin), Sub Registrar (Admin), Accountant Superintendent (Admin), FDA (Admin), SDA (Admin), Typist (Admin), HQA (RTI), FDA (RTI), SDA (RTI), Statistical Inspector"],
            ["Vigilance", "DIGR (Vigilance), Law Officer"],
            ["Computers", "AIGR (Computers), System Integrator, PMU, Application Developer, HQA / Project Manager (Comp), Sub Registrar (Computers), FDA (Computers), SDA (Computers)"],
            ["Enforcement", "DIGR (Enforcement), DRO, HQA (Enforcement), Sub-Registrar (SR), FDA (Enforcement), SDA (Enforcement), DEO"],
            ["Intelligence & Audit", "DIGR (Intelligence), AIGR (Audit), HQA (Audit), Superintendent (Audit), FDA (Audit), SDA (Audit), Typist (Audit)"],
            ["DIGR CVC", "DIGR CVC, JD Town Planning"],
        ],
        lead_in="Divisions come from the Division Master, not free text. DIV-ADMIN (“Admin, Law & Computers”) and "
                "DIV-COMPUTERS (“Computers”) are distinct divisions — v4.23 clarifies this is not a duplication.",
        col_widths=[2.6, 9.5], body_size=10, row_h=Inches(0.40),
    )

    deck.table_slide(
        "The privilege chain", "Section 4.5.6",
        ["Level", "Master", "Maintained by", "Purpose"],
        [
            ["1", "Module Master", "Application Admin", "Business modules — Registration of Documents, Marriage Registration, Encumbrance Search, Certified Copy, Stamp Duty / Payments, Firm / Society Registration, User Management, MIS / Dashboards"],
            ["2", "Module Function Master", "Application Admin", "Functions under each module — VIEW, ADD, EDIT, APPROVE, SIGN, PRINT, APPLY, ISSUE, ADMIN"],
            ["3", "Resource Master", "Application Admin", "APIs and URLs linked to each Module Function, with an Is Public flag for unauthenticated endpoints"],
            ["4", "Role–Module–Function mapping", "Application Admin", "Which roles may perform which Module Functions; role names must match the Role Master exactly (FR-UM-050)"],
        ],
        lead_in="Access is modelled as User → Role(s) → Module Function(s) → Resource(s). DSR organisational roles are never named after application services.",
        col_widths=[0.8, 2.6, 2.2, 6.5], body_size=10.5, row_h=Inches(0.86),
    )

    deck.flow_slide(
        "Runtime access enforcement", "Section 4.5.6  ·  S-03 Runtime Enforcement  ·  FR-UM-038, FR-UM-041, FR-UM-050, FR-UM-051",
        [
            "User authenticates — Citizen/Other Dept: OTP; DSR: Face/Biometrics, no OTP",
            "Session role(s) load for the active context (assigned post or additional charge)",
            "Role–Module–Function mappings resolve into session claims",
            "User calls an API or opens a URL",
            "Resource Master is checked for the matching Type + Method + Path",
            "Required Module Function is obtained from the matched resource",
            "Allow if session claims include that function — else HTTP 403, audited",
            "UI hides menus/buttons for functions not in session claims",
        ],
        key_text="Deny by default. Only resources flagged Is Public bypass authentication; Application Admin is a "
                  "system-level actor authorised for FN-UM-ADMIN outside Role–Module–Function mapping (FR-UM-051).",
    )

    deck.tree_slide(
        "DSR Officer Hierarchy Master", "Section 4.5.7  ·  S-04 Officer Hierarchy Tree  ·  FR-UM-043",
        levels=[
            [("acs", "ACS / Principal Secretary / Secretary", Inches(4.6))],
            [("igr", "IGR", Inches(2.4))],
            [("digr", "DIGR (Enforcement)", Inches(3.2))],
            [("dro", "District Registrar (DRO)", Inches(3.2))],
            [("sr", "Sub-Registrar (SR)", Inches(3.0))],
            [("fda", "FDA (Enforcement)", Inches(2.9)), ("deo", "Data Entry Operator (DEO)", Inches(2.9))],
            [("sda", "SDA (Enforcement)", Inches(2.9))],
        ],
        edges=[("acs", "igr"), ("igr", "digr"), ("digr", "dro"), ("dro", "sr"), ("sr", "fda"), ("sr", "deo"), ("fda", "sda")],
        footnote="Authority is granted by immediate parentage only — IGR may act on DIGR, DIGR (Enforcement) on DRO, "
                 "DRO on Sub-Registrar, SR on FDA and DEO. IGR cannot relieve a Sub-Registrar directly.",
    )

    deck.tree_slide(
        "Office Hierarchy Master", "Section 4.5.8  ·  S-05 Office Hierarchy Tree  ·  FR-UM-059",
        levels=[
            [("ms", "MS Building (Secretariat)", Inches(3.6))],
            [("igro", "IGR Office (Head Office)", Inches(3.4))],
            [("drob", "DRO Bengaluru", Inches(2.8)), ("drom", "DRO Mysuru", Inches(2.8))],
            [("sroy", "SRO Yeshwanthapura", Inches(2.5)), ("sroj", "SRO Jayanagar", Inches(2.5)), ("srom", "SRO Mysuru East", Inches(2.5))],
        ],
        edges=[("ms", "igro"), ("igro", "drob"), ("igro", "drom"), ("drob", "sroy"), ("drob", "sroj"), ("drom", "srom")],
        footnote="Maintained separately from the Officer Hierarchy; it scopes which offices a superior can see for "
                 "Transfer Out and Transfer In. MS Building is the Secretariat root — office of POST-ACS-SEC.",
    )

    deck.comparison_slide(
        "Office span is visibility, not authority", "Section 4.5.8  ·  S-06 Span vs Action  ·  FR-UM-057, FR-UM-059, FR-UM-043",
        {
            "title": "Office span — visibility",
            "rows": [
                ("Rule", "The actor's session Office plus all descendant offices in the Office Hierarchy (FR-UM-059)"),
                ("Grants", "Visibility of occupancies in those offices — nothing more"),
                ("Example", "IGR at Head Office sees every DRO and SRO statewide"),
            ],
        },
        {
            "title": "Immediate parent — authority",
            "rows": [
                ("Rule", "The actor's session Post must be the immediate parent of the target Post (FR-UM-043)"),
                ("Grants", "The actual right to relieve, Transfer In, or record absence"),
                ("Example", "DRO may act on Sub-Registrar; IGR may not — SR does not report to IGR"),
            ],
        },
        footnote="Both tests must pass. Seeing a descendant office in the tree does not grant authority over the posts in it.",
    )

    # ---------------- Section 4 ----------------
    deck.divider("04", "Officer Lifecycle", "Creation, transfer, the midnight occupancy refresh, and temporary absence — Transfer In now takes effect immediately.")

    deck.flow_slide(
        "DSR Officer creation with post assignment", "Section 4.6.1  ·  P-07 DSR Officer User Creation  ·  FR-UM-002, FR-UM-017, FR-UM-030, FR-UM-045–048, FR-UM-062, FR-UM-066(a)",
        [
            "Open “Add DSR Department User” — authorised admin only",
            "Enter KGID (becomes the Username) + particulars — no email, no security questions",
            "Assign one or more sanctioned posts with available capacity — at least one required",
            "System shows roles available via Post–Role mapping for each post",
            "Upload approval letter — recommended, not mandatory",
            "Capture face authentication template or biometrics — mandatory",
            "Review post occupancies and mapped roles; confirm",
            "Save — account active; occupied count updated per assigned post",
        ],
        key_text="No optional End Date or Deputation Reason is captured at creation (removed v4.19) — later movement "
                  "uses Transfer Out / Transfer In instead (FR-UM-057, FR-UM-060).",
    )

    deck.flow_slide(
        "Other Department user creation", "Section 4.6.2  ·  P-08 Other Department User Creation  ·  FR-UM-003, FR-UM-029, FR-UM-033, FR-UM-034, FR-UM-062, FR-UM-064",
        [
            "Open “Add Other Department User” — authorised admin only",
            "Enter Employee ID/KGID + Department Code — Username = Dept Code + ID",
            "Enter parent department, designation, official email, mobile",
            "Assign exactly one role (Role Category = Other Department)",
            "Optionally enter Account End Date",
            "Upload authorisation letter / NOC — recommended, not mandatory",
            "Review role and End Date; confirm",
            "Save — account active with module access for the assigned role",
            "If End Date is reached — deactivate user; block login",
        ],
        key_text="Biometrics are not required for this category — Other Department authenticates with Username + "
                  "Captcha + OTP only (FR-UM-007). No security questions are captured (no self-service reset).",
    )

    deck.flow_slide(
        "Transfer Out and relieving", "Section 4.6.3  ·  P-09 Transfer Out / Relieving  ·  FR-UM-057, FR-UM-058, FR-UM-059, FR-UM-043, FR-UM-068, FR-UM-087",
        [
            "Open Transfer Out / Relieving",
            "System lists only offices in the actor's office span",
            "Within those, lists occupancies where the actor's post is immediate parent",
            "Superior selects the officer / post occupancy to relieve",
            "Enter Relieving Date, mandatory Relieving Reason, and Relieving Order",
            "Confirm relieving — audit-logged",
            "Occupancy remains active through the Relieving Date",
            "Midnight job after the Relieving Date de-allocates the mapping",
            "If no post occupancies remain — login is blocked/limited until reassigned",
        ],
        key_text="Relieving Reason is now mandatory — Deputation, Transfer, Suspension, Superannuation or Death "
                  "(FR-UM-087) — alongside Relieving Date and Order.",
    )

    deck.flow_slide(
        "Transfer In — immediate, capacity-gated", "Section 4.6.4  ·  P-10 Transfer In  ·  FR-UM-060, FR-UM-066(a)",
        [
            "Open Transfer In",
            "System lists offices in the actor's office span",
            "Select target Post + Office where the actor's post is immediate parent",
            "Proceed only if available capacity; block if at full strength",
            "Select / identify the officer being transferred in",
            "Enter Transfer Order / Reporting Order — mandatory",
            "Confirm — occupancy active immediately; occupied count +1",
            "Officer may select the post at the very next login",
            "Occupancy refresh job handles relieving de-allocation only",
        ],
        key_text="Transfer In no longer captures a Joining Date — the reserved/future-dated path is retired "
                  "(FR-UM-061, FR-UM-067). A post at full strength blocks Transfer In until relieving frees capacity.",
    )

    deck.flow_slide(
        "Worked example — relieving then Transfer In", "Section 4.6.4  ·  P-12 Handover Timeline  ·  DRO Bengaluru → SRO Yeshwanthapura, Relieving Date 31-Aug-2026",
        [
            "Relieving Date entered & confirmed (Day 0)",
            "Relieving Date ends — post still usable that day",
            "Midnight refresh job de-allocates the mapping",
            "Capacity freed — Transfer In now allowed",
            "Transfer In recorded — occupancy active immediately",
            "New officer selects the post at next login",
        ],
        key_text="There is no reservation against a future relieving — Transfer In of the incoming officer cannot be "
                  "recorded while the post remains at full strength. IGR at Head Office cannot perform this Transfer "
                  "In: Sub-Registrar does not report immediately to IGR.",
    )

    deck.flow_slide(
        "The midnight occupancy refresh job", "Section 4.6.5  ·  P-11 Occupancy Refresh Job  ·  FR-UM-068 — idempotent, runs shortly after 12:00 AM IST",
        [
            "Job starts shortly after 12:00 AM IST — idempotent",
            "De-allocates occupancies whose Relieving Date ended the previous day",
            "Recalculates occupied count, capacity and the wholly-unoccupied flag",
            "Refreshes each officer's effective post assignments for login selection",
            "Writes an occupancy-refresh audit record",
            "On failure — raises an operational alert",
        ],
        key_text="No reserved Transfer In activation step and no optional End Date processing remain in this job — "
                  "both were retired in v4.19/v4.20. Transfer In itself now takes effect immediately when recorded.",
    )

    deck.flow_slide(
        "Temporary absence — Leave, OOD and cover", "Section 4.6.6  ·  P-13 Temporary Absence and Temporary Charge  ·  FR-UM-079 – FR-UM-084",
        [
            "Superior opens Temporary Absence / Leave / OOD — not officer self-service",
            "System lists office span, then occupancies with immediate-parent post",
            "Select occupancy; enter type, reason code, from/to dates",
            "Confirm absence — Occupied remains unchanged",
            "From from_date — login is denied for the absent Username",
            "Optionally assign Temporary Charge to another subordinate (any office)",
            "Cover officer authenticates; sees own posts + Temporary charge row",
            "Cover officer selects Temporary charge — claims from covered post only",
            "After to_date — absence and charge clear; absentee may log in again",
        ],
        key_text="Absence does not free the slot — Transfer In must not treat the post as vacant. While an effective "
                  "absence exists the officer's login is blocked entirely (distinct from FR-UM-053 additional charge).",
    )

    # ---------------- Section 5 ----------------
    deck.divider("05", "Quality, Reporting and Sign-off", "Non-functional requirements, reports, acceptance criteria and open risks.")

    deck.two_column_slide(
        "Non-functional requirements", "Section 5 — 24 requirements across 6 categories",
        "SECURITY AND PERFORMANCE",
        [
            "No password storage anywhere; DSR Officers authenticate with Captcha + Face/Biometrics only — no OTP; Citizens and Other Department use OTP to mobile only.",
            "TLS 1.2 or higher for all data in transit; DSR biometric/face data and Citizen Aadhaar e-KYC both comply with the Aadhaar Act 2016 and UIDAI / MeitY norms.",
            "OTP dispatch within 5 seconds; login completes within 2 seconds after OTP or face/biometric verification.",
            "Minimum 500 concurrent authenticated sessions at launch, validated in performance testing with Kaveri IT Cell.",
        ],
        "AVAILABILITY, AUDIT AND COMPLIANCE",
        [
            "The occupancy refresh job completes shortly after midnight so relieving is reflected before the first login of the day; failure raises an operational alert.",
            "Disaster recovery targets RPO ≤ 24 hours and RTO ≤ 4 hours.",
            "All create/update/delete, login, mobile/email change, Aadhaar e-KYC, Transfer Out/In and occupancy-refresh actions are audit-logged with timestamp and actor.",
            "Audit logs retained for at least 7 years; bilingual Kannada/English citizen UI, GIGW accessibility, India data residency, DPDP Act 2023 compliance.",
        ],
    )

    deck.bullets_slide(
        "Reporting requirements", "Section 6",
        [
            "Active, inactive and suspended users; role and permission assignments across all users.",
            "Audit log of login attempts, successful and failed, over a selected date range.",
            "Sanctioned post occupancy — sanctioned strength, occupied count, remaining capacity and wholly-unoccupied flag per office (FR-UM-066).",
            "Role-to-Module mapping showing which modules and functions each role holds.",
            "A single contact-change and recovery report covering Citizen Aadhaar e-KYC / lost-mobile recovery, DSR self-service and administrator-initiated mobile changes, and email changes.",
            "Additional charge report (FR-UM-053) and Temporary Absence / Temporary Charge report (FR-UM-079–084).",
            "Occupancy-refresh report per midnight run, with before and after occupied counts (FR-UM-068).",
            "Transfer Out / Transfer In history and officer posting / service history.",
        ],
        size=14,
    )

    deck.bullets_slide(
        "Acceptance criteria", "Section 7",
        [
            ("FR-UM-052", "an officer with two occupancies must choose a post before home; a single occupancy auto-selects."),
            ("FR-UM-053", "additional charge lists only wholly unoccupied posts at the same office; privileges switch until switch-back."),
            ("FR-UM-058 / 068", "relieving holds through 23:59 IST of the Relieving Date; the midnight job then de-allocates and updates counts."),
            ("FR-UM-060 / 066(a)", "Transfer In to a full post is blocked; once capacity frees it succeeds immediately with a Transfer Order only — no Joining Date."),
            ("FR-UM-069–076", "OTP validity, code length, lockout, 10-minute idle timeout, 4-hour session cap and single active session all enforced."),
            ("FR-UM-079–084", "absence blocks login, keeps Occupied unchanged, and temporary charge appears for the cover officer until the to_date."),
            ("Role mapping", "Role–Module–Function mapping completed and Domain Expert confirmed for every active role and module before go-live."),
        ],
        lead_in="Accepted when Section 4 is implemented and passes QA including these UAT scenarios, Section 5 is verified, "
                "the Product Owner signs off UAT, and security review closes with no critical or high findings.",
        size=13.5,
    )

    deck.table_slide(
        "Key risks and mitigations", "Section 8 — 24 risks logged",
        ["Risk", "Impact", "Mitigation"],
        [
            ["Role–Module–Function mapping incomplete at go-live", "High", "Only SR, FDA (Enforcement), DEO, Citizen roles and one Other Department role are mapped; ~37 DSR roles and 3 of 8 modules have none. Application Admin, with Domain Expert sign-off, must complete the matrix before go-live."],
            ["Aadhaar e-KYC failure or UIDAI downtime", "Medium", "Registration and lost-mobile reset both depend on Aadhaar e-KYC (FR-UM-085); provide clear user messaging, retry, and operational monitoring of UIDAI connectivity — there is no security-question fallback."],
            ["Unauthorised relieving or Transfer In", "High", "Scope to offices under the actor (FR-UM-059), then immediate-parent posts only (FR-UM-043, FR-UM-057)."],
            ["Vacancy tests confused in build", "High", "FR-UM-066 defines available capacity and wholly unoccupied separately; additional charge uses Occupied = 0."],
            ["Transfer In without available capacity", "High", "Block Transfer In when Occupied ≥ Sanctioned Strength (FR-UM-060, FR-UM-066(a)); the Joining Date / reservation path is retired."],
            ["Leave / OOD treated as a vacancy", "High", "FR-UM-081 keeps Occupied unchanged; temporary charge is superior-assigned and distinct from FR-UM-053."],
            ["Account takeover via lost-mobile reset", "High", "Two independent proofs — Aadhaar e-KYC and a single-use PIN to email — plus OTP verification of the new number, rate limiting and audit."],
            ["Re-authentication may slow high-volume SRO counters", "Medium", "10-minute idle timeout plus Captcha + face/biometrics on every fresh DSR login adds overhead; measure end-to-end re-authentication time in performance testing with Kaveri IT Cell."],
        ],
        col_widths=[3.6, 1.0, 7.5], body_size=9.5, row_h=Inches(0.62),
    )

    slide = deck.bullets_slide(
        "Approval and next steps", "Section 10",
        [
            "Domain Expert review of the officer hierarchy, office hierarchy and Post–Role mapping seed data.",
            "Application Admin to complete Role–Module–Function mapping for every active role and module before go-live.",
            "Re-export the header text on all nineteen process/structure diagrams — their .drawio sources still show the old §6.x-style section numbers, though captions now read correctly.",
            "Align the P-13 diagram's “InCharge Officer” label with the requirement text's “Cover officer” in the .drawio source.",
            "Department to confirm the PII retention and purge policy for dormant Citizen accounts.",
            "Performance testing with Kaveri IT Cell to validate peak-load and re-authentication timings at SRO counters.",
            "Product Owner sign-off on UAT, then baseline the BRD for design and development.",
        ],
        lead_in="Signatories: Prashanth (Product Owner) · Prabhakar Naik (Domain Expert) · Kaveri IT Cell Lead (IT Security / Engineering)",
        size=13,
    )
    add_rect(slide, MARGIN, Inches(6.14), CONTENT_W, Inches(0.66), fill=BAND, line=RULE)
    add_text(slide, MARGIN + Inches(0.24), Inches(6.30), CONTENT_W - Inches(0.48), Inches(0.4),
             "Source of record: BRD_User_Management_v4.23.docx  ·  Editable diagram sources: ProcessDiagrams/User_Management/*.drawio",
             size=11.5, color=MUTED)

    return deck.prs


def main() -> None:
    prs = build()
    prs.save(DST)
    print(f"{DST} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
