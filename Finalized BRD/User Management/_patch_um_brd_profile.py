import copy
import re
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

PATH = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.0.docx"
d = docx.Document(PATH)
body = d.element.body


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def strip_bookmarks(el):
    for b in el.xpath(".//w:bookmarkStart|.//w:bookmarkEnd"):
        b.getparent().remove(b)


# Renumber section 4.6+ headings and the code table before inserting the new section.
for p in d.paragraphs:
    if p.style.name == "Heading 2":
        m = re.match(r"4\.(\d+) (.*)", p.text)
        if m and int(m.group(1)) >= 6:
            set_text(p, f"4.{int(m.group(1)) + 1} {m.group(2)}")

codes = d.tables[2]
for row in codes.rows[1:]:
    cell = row.cells[1].paragraphs[0]
    m = re.match(r"4\.(\d+) (.*)", cell.text)
    if m and int(m.group(1)) >= 6:
        set_text(cell, f"4.{int(m.group(1)) + 1} {m.group(2)}")

mob_row = next(r for r in codes.rows if r.cells[0].text == "MOB")
new_code_row = copy.deepcopy(mob_row._tr)
mob_row._tr.addnext(new_code_row)
prf_code = [r for r in codes.rows if r._tr is new_code_row][0]
set_text(prf_code.cells[0].paragraphs[0], "PRF")
set_text(prf_code.cells[1].paragraphs[0], "4.6 Profile updates")

mob_tbl = d.tables[8]
assert mob_tbl.rows[1].cells[0].text == "UM-MOB-01"
photo_row = mob_tbl.rows[-1]
assert photo_row.cells[0].text == "UM-MOB-07"
photo_text = photo_row.cells[1].text

# Locate the MOB heading and intro paragraph preceding the table.
elems = list(body.iterchildren())
ti = elems.index(mob_tbl._tbl)
heading_el = next(e for e in reversed(elems[:ti]) if e.tag.endswith("}p")
                  and Paragraph(e, d).text.startswith("4.5 "))
between = [e for e in elems[elems.index(heading_el) + 1:ti] if e.tag.endswith("}p")]
intro_el = next(e for e in between if Paragraph(e, d).text.strip())

new_heading = copy.deepcopy(heading_el)
new_intro = copy.deepcopy(intro_el)
new_tbl = copy.deepcopy(mob_tbl._tbl)
for el in (new_heading, new_intro, new_tbl):
    strip_bookmarks(el)

mob_tbl._tbl.remove(photo_row._tr)

anchor = mob_tbl._tbl
for el in (new_heading, new_intro, new_tbl):
    anchor.addnext(el)
    anchor = el

set_text(Paragraph(new_heading, d), "4.6 Profile updates")
set_text(Paragraph(new_intro, d), "Users can update their own profile details after signing in.")
t = Table(new_tbl, d)
for r in list(t.rows)[2:]:
    new_tbl.remove(r._tr)
set_text(t.rows[1].cells[0].paragraphs[0], "UM-PRF-01")
set_text(t.rows[1].cells[1].paragraphs[0], photo_text)

d.save(PATH)
print("Saved", PATH)
