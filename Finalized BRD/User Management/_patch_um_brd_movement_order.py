import copy
import docx

PATH = r"Finalized BRD\User Management\BRD_User_Management_Simplified_v1.0.docx"
d = docx.Document(PATH)


def set_text(p, text):
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def replace_para(old_start, new):
    hits = [p for p in d.paragraphs if p.text.startswith(old_start)]
    assert len(hits) == 1, (old_start, len(hits))
    set_text(hits[0], new)


def find_row(first_cell):
    hits = [r for t in d.tables for r in t.rows if r.cells[0].text == first_cell]
    assert len(hits) == 1, (first_cell, len(hits))
    return hits[0]


def set_cell(first_cell, col, new):
    set_text(find_row(first_cell).cells[col].paragraphs[0], new)


def add_row_after(first_cell, values):
    row = find_row(first_cell)
    tr = copy.deepcopy(row._tr)
    row._tr.addnext(tr)
    new_row = next(r for t in d.tables for r in t.rows if r._tr is tr)
    for cell, text in zip(new_row.cells, values):
        set_text(cell.paragraphs[0], text)


replace_para("Let superiors manage transfers",
             "Let the IGR office trigger transfers and relieving through movement orders, and let superiors "
             "carry them out and manage leave and temporary charge for the officers under them.")

replace_para("A superior places a DSR Officer in a post that has a vacancy.",
             "Transfer In is triggered from the IGR office by uploading the movement order. The superior then "
             "places the DSR Officer in a post that has a vacancy. This is also how a post is given to an "
             "officer who was created without one.")
set_cell("UM-TIN-03", 1,
         "The movement order uploaded by the IGR office is linked to the Transfer In. The Transfer Order / "
         "Reporting Order number is also required (with upload where available). No Joining Date is entered.")
add_row_after("UM-TIN-05", ["UM-TIN-06",
                            "Transfer In is triggered from the IGR office. An authorised user at the IGR office "
                            "uploads the movement order for the officer; only then can the superior carry out "
                            "the Transfer In."])

replace_para("A superior relieves an officer from one or more posts.",
             "Relieving is triggered from the IGR office by uploading the movement order. The superior then "
             "relieves the officer from one or more posts.")
set_cell("UM-REL-02", 1,
         "The movement order uploaded by the IGR office is linked to the Relieving. Relieving Date, Relieving "
         "Order and Relieving Reason are also required. The reason must be one of: Deputation, Transfer, "
         "Suspension, Superannuation, Death.")
add_row_after("UM-REL-06", ["UM-REL-07",
                            "Relieving / Transfer Out is triggered from the IGR office. An authorised user at the "
                            "IGR office uploads the movement order for the officer; only then can the superior "
                            "relieve the officer."])

set_cell("UM-RPT-09", 1, "Transfer In and Relieving history, with movement orders, other orders and reasons.")
set_cell("UM-ACC-04", 1,
         "Transfer In and Relieving cannot be done until the IGR office has uploaded the movement order. "
         "Transfer In to a full post is blocked; after relieving takes effect at midnight, Transfer In succeeds "
         "immediately with the movement order and a Transfer / Reporting Order.")

set_cell("Transfer In", 1,
         "Placing an officer in a post with a vacancy, triggered by a movement order from the IGR office, with a "
         "Transfer / Reporting Order.")
set_cell("Relieving / Transfer Out", 1,
         "Removing an officer from a post from the end of the Relieving Date, triggered by a movement order from "
         "the IGR office, with a Relieving Order and reason.")
add_row_after("Relieving / Transfer Out", ["Movement order",
                                           "The order uploaded by the IGR office that triggers Transfer In and "
                                           "Relieving / Transfer Out of an officer."])

d.save(PATH)
print("Saved", PATH)
