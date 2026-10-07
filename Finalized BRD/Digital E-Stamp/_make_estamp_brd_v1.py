# -*- coding: utf-8 -*-
"""Generate BRD_Digital_EStamp_v1.docx.

Section structure, styles and table layouts follow
Finalized BRD/Marriage/RFP/BRD_Marriage_BRD_Final_v8.docx.
Template catalogue and document natures are read from the template / stamp-article data modules
so the BRD stays aligned with Legal Formats/Template.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
HERE = ROOT / "Finalized BRD" / "Digital E-Stamp"
REF = ROOT / "Finalized BRD" / "Marriage" / "RFP" / "BRD_Marriage_BRD_Final_v8.docx"
OUT = HERE / "BRD_Digital_EStamp_v1.docx"

sys.path.insert(0, str(ROOT / "Legal Formats" / "Template" / "Document Registration"))
sys.path.insert(0, str(ROOT / "Legal Formats" / "Template" / "Digital E-Stamp"))
sys.path.insert(0, str(ROOT / "Acts_Rules" / "Document"))
sys.path.insert(0, str(HERE))
import _template_data as TD  # noqa: E402
import _make_estamp_templates as ET  # noqa: E402
import _stamp_article_inputs_data as SA  # noqa: E402
import _estamp_diagrams as DG  # noqa: E402

# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------
doc = None
RTM: list[tuple] = []


def new_doc():
    global doc
    tmp = HERE / "_base_tmp.docx"
    shutil.copy(REF, tmp)
    doc = Document(str(tmp))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for sec in doc.sections:
        hdr = sec.header
        hdr.is_linked_to_previous = False
        p = hdr.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.add_run("Kaveri 3.0 | BRD | Digital E-Stamp Module")
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    tmp.unlink(missing_ok=True)


def H(text, level):
    return doc.add_heading(text, level=level)


def P(text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
    p.add_run(text)
    return p


def LP(text):
    return P(text, "List Paragraph")


def B(items):
    for it in items:
        if isinstance(it, tuple):
            P(it[1], "List Bullet", bold_prefix=it[0] + ": ")
        else:
            P(it, "List Bullet")


def _shade(cell, color="D9E2F3"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def T(headers, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(size)
        _shade(c)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            lines = str(v).split("\n")
            para = cells[i].paragraphs[0]
            for j, ln in enumerate(lines):
                if j:
                    para = cells[i].add_paragraph()
                rr = para.add_run(ln)
                rr.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def IMG(path, width=6.8):
    doc.add_picture(str(path), width=Inches(width))


def TOC():
    p = doc.add_paragraph()
    for kind, txt in (("begin", None), ("instr", 'TOC \\o "1-4" \\h \\z \\u'), ("separate", None),
                      ("text", "Right-click and choose Update Field to refresh the table of contents."),
                      ("end", None)):
        r = p.add_run()
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            r._r.append(it)
        elif kind == "text":
            r.text = txt
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def FR(prefix, start, section, screen, source, items, with_ac=True):
    """items: (requirement, priority, acceptance[, source override]). Returns next number."""
    rows = []
    n = start
    for it in items:
        req, pri, ac = it[0], it[1], it[2]
        src = it[3] if len(it) > 3 else source
        rid = f"{prefix}-{n:03d}"
        rows.append((rid, req, pri, ac) if with_ac else (rid, req, pri))
        RTM.append((rid, src, req, section, screen))
        n += 1
    if with_ac:
        T(["Req ID", "Requirement", "Priority", "Acceptance criteria"], rows, [1.0, 3.4, 0.6, 2.0])
    else:
        T(["Req ID", "Requirement", "Priority"], rows, [1.0, 5.0, 0.8])
    return n


# ---------------------------------------------------------------------------
# Derived data from template / stamp-article modules
# ---------------------------------------------------------------------------
def generated_in(t) -> str:
    reg = t["registration"]
    if ET.is_estamp(t):
        return "Digital E-Stamp (Part I-IV) and Document Registration (Part I-V)"
    if str(t["dn"]).startswith("DN-"):
        return "Document Registration"
    return "Template service only (not a Digital E-Stamp instrument)"


def _article_base(text: str) -> list[str]:
    return [m.group(0) for m in re.finditer(r"\b\d+(?:-[A-Z])?", re.sub(r"\([^)]*\)", "", str(text)))]


DN_BY_ARTICLE: dict[str, list[str]] = {}
for _dn, _name, _it, _art in SA.DOCUMENT_NATURES:
    for base in _article_base(_art)[:1]:
        DN_BY_ARTICLE.setdefault(base, []).append(_dn)

TEMPLATE_ROWS = []
DN_TEMPLATES: dict[str, list[str]] = {}
for t in TD.TEMPLATES:
    TEMPLATE_ROWS.append((t["id"], t["name"], t["dn"] or "-", t["article"] or "-", t["layout"],
                          generated_in(t)))
    covered = [t["dn"]] if str(t["dn"]).startswith("DN-") else []
    if covered:
        for base in _article_base(t["article"]):
            covered += DN_BY_ARTICLE.get(base, [])
    for dn in dict.fromkeys(covered):
        DN_TEMPLATES.setdefault(dn, []).append(t["id"])
DN_GAPS = [dn for dn, *_ in SA.DOCUMENT_NATURES if dn not in DN_TEMPLATES]
N_TEMPLATES = len(TD.TEMPLATES)
N_ESTAMP_TEMPLATES = len(ET.ESTAMP)
INTENTS = dict(SA.TRANSACTION_INTENTS)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------
def build():
    new_doc()
    doc.add_paragraph("Business Requirements Document (BRD)", style="Title")
    H("Digital E-Stamp Module", 1)

    # ------------------------------------------------------------------ control
    H("Document control", 2)
    T(["Field", "Value"], [
        ("Document ID", "BRD-K3-EST-001"),
        ("Version", "1"),
        ("Status", "Draft for Review Committee - initial BRD for Digital E-Stamp with shared Template Generation "
                   "Microservice"),
        ("Module", "Digital E-Stamp (deed generation and digital e-Stamp certificate)"),
        ("Legal basis (primary)", "The Karnataka Stamp Act, 1957 (Secs. 3, 4-6, 10, 10-A, 13-17, 27, 33-34, 47-52-A) "
                                  "and its Schedule; The Registration Act, 1908 (Secs. 17, 18); "
                                  "The Information Technology Act, 2000 (Secs. 3A, 4, 5)"),
        ("State rules (primary)", "The Karnataka Stamp (Payment of Duty by means of e-Stamping) Rules, 2009 "
                                  "(No. RD 380 MUNOMU 2008, dated 08-04-2009); The Karnataka Stamp Rules, 1958 "
                                  "(Rules 16, 17)"),
        ("Shared component", "Template Generation Microservice - consumed by Digital E-Stamp and Document "
                             "Registration to render PDFs from pre-designed templates"),
        ("Author (BA)", "Nandha Kumar"),
        ("Product Owner", "M V Prashanth"),
        ("Domain expert / SRO reviewer", "Committee members, Department of Stamps and Registration"),
        ("Target audience", "Kaveri IT Cell, Department of Stamps and Registration, Government of Karnataka"),
        ("Last updated", "07-10-2026"),
    ], [2.0, 4.8])
    T(["Version", "Date", "Author", "Summary of change", "Approver"], [
        ("1", "07-10-2026", "Nandha Kumar",
         "Initial BRD: Digital E-Stamp process (online and assisted), template-driven deed generation through the "
         "shared Template Generation Microservice, e-Stamp certificate, verify / lock / cancel / refund, "
         f"NFRs and template catalogue ({N_TEMPLATES} templates).", "Prashanth"),
    ], [0.7, 0.9, 1.1, 3.2, 0.9])
    H("Table of contents", 3)
    TOC()

    # ------------------------------------------------------------------ exec summary
    H("Executive summary", 2)
    LP("This document defines the business requirements for the Digital E-Stamp module of Kaveri 3.0. The module "
       "lets a citizen, or a department operator on the citizen's behalf, prepare an instrument (deed) from a "
       "pre-designed, legally vetted template, compute the stamp duty payable under the Karnataka Stamp Act, 1957, "
       "pay the duty through Khajane-II, have every party e-Sign, and receive a digitally sealed e-Stamp "
       "certificate with the deed attached.")
    LP("In Kaveri 2.0 the Digital e-Stamp captured only a free-text description of the instrument, and parties "
       "drafted the actual deed outside the system. Kaveri 3.0 replaces this with structured, template-driven deed "
       f"generation: {N_TEMPLATES} templates (of which {N_ESTAMP_TEMPLATES} are used by Digital E-Stamp) have been "
       "prepared from the Department's legal formats, each with a defined field dictionary.")
    LP("Deed rendering is delivered by a separate Template Generation Microservice. Both Digital E-Stamp and "
       "Document Registration call this service to turn structured data into a PDF, so a deed looks the same "
       "whether it starts in E-Stamp or in Document Registration, and a legal change to a template is made once.")
    LP("The proposed approach reduces drafting errors and disputes about wording, ensures that statutory "
       "particulars (Sec. 27 facts affecting duty, e-Stamping Rule 11 certificate particulars and the Rule 28 "
       "unique number on every page) are always present, prevents reuse of an e-Stamp through lock-on-use, and "
       "gives the Department one auditable source for stamp revenue collected through the digital channel.")

    # ------------------------------------------------------------------ scope
    H("Scope", 2)
    H("In scope", 3)
    B([
        "Digital E-Stamp for optionally registrable and non-registrable stamp instruments (Annexure 1 of the "
        "Kaveri 2.0 Digital e-Stamp FRS) - Online (citizen self-service) and Assisted (operator at SRO / Seva "
        "Kendra) channels.",
        "Two-level article mapping: nature of document, then stamp sub-article; one sub-article per instrument.",
        "Schedule (property) capture with fetch from Bhoomi, E-Swathu, E-Aasthi and ULMS; GL / GR and alienation "
        "restrictions; relaxed validations for listed sub-articles.",
        "Party, representative (POA, guardian of minor, organisation signatory) and witness capture with mobile OTP, "
        "Aadhaar e-KYC or DSC (non-e-KYC path), and C-DAC name matching.",
        "Stamp duty computation (market value, consideration, fixed, highest of all; minimum / maximum where "
        "applicable; family concessions per Schedule).",
        "Template-driven deed drafting, DRAFT preview and FINAL rendering through the Template Generation "
        "Microservice.",
        "Payment through Khajane-II including challan verification for indeterminate transactions.",
        "eSign by all parties, department seal on the final PDF, QR-code verification, download and DigiLocker push.",
        "Additional stamp duty on an issued e-Stamp (e-Stamping Rules 25-26) linked to the original UIN.",
        "Verify Digital E-Stamp (public, masked) and verify-and-lock by SRO / DR / DC of Stamps on use (Rules 29-30).",
        "Cancellation and refund of spoiled / unused digital e-Stamps (Rules 31-32; Secs. 47-52-A).",
        "Upload-own-deed path: e-Stamp certificate for a deed drafted outside Kaveri (PDF upload) where no template "
        "exists or the party prefers its own draft.",
        "Notifications, MIS, Sakala integration and audit.",
        "Template Generation Microservice: template registry and versioning, render API, PDF output standards, "
        "template administration (maker-checker) - as a shared platform component.",
    ])
    H("Out of scope", 3)
    B([
        "Issue of physical e-stamp certificates by the Central Record Keeping Agency (CRA) and its Authorised "
        "Collection Centres; Kaveri only verifies and locks such certificates through the CRA interface used by "
        "Document Registration.",
        "Compulsorily registrable instruments (Sec. 17, Registration Act) - drafted, stamped and registered in the "
        "Document Registration module, which consumes the same templates.",
        "Franking and impressed stamp paper sale.",
        "Adjudication of stamp duty (Sec. 31) and impounding (Sec. 33) workflows - covered in the Stamps / "
        "Adjudication BRD.",
        "Guidance value maintenance (covered in the Guidance Value BRD); E-Stamp only reads published values.",
    ])
    H("Assumptions", 3)
    B([
        "Government has notified, or will notify under Sec. 10(3) of the Stamp Act, that the Department's own "
        "Kaveri digital e-Stamp is an authorised electronic stamping system (see Appendix B, open point OP-01).",
        "Khajane-II, UIDAI e-KYC / eSign, DigiLocker and land-record interfaces used in Kaveri 2.0 remain "
        "available with the same contracts or better.",
        f"The {N_TEMPLATES} templates and the field workbook in Legal Formats/Template are the baseline and will be "
        "legally vetted by the Department before go-live.",
    ])
    H("Constraints", 3)
    B([
        "Aadhaar e-KYC and eSign are subject to UIDAI / ESP approvals and consent rules; Aadhaar numbers are stored "
        "only in the Aadhaar Data Vault.",
        "A stamp instrument must be stamped before or at execution (Sec. 17); the e-Stamp must therefore be issued "
        "before parties sign the deed.",
        "The e-Stamp certificate must appear on the face of the instrument with the instrument written below it, "
        "and the unique number must appear at the top centre of every page (Rules 27, 28; Secs. 13, 14).",
    ])
    H("Dependencies", 3)
    B([
        "Template Generation Microservice available in all environments before E-Stamp SIT.",
        "Stamp article master, sub-article duty rules and family matrix (Acts_Rules/Document) frozen and approved.",
        "Document Registration module exposes the 'e-Stamp UIN used' event so that lock-on-use is automatic.",
    ])

    # ------------------------------------------------------------------ legal
    H("Legal and regulatory reference", 2)
    H("Applicable Acts", 3)
    T(["Act", "Description"], [
        ("The Karnataka Stamp Act, 1957", "Levy of stamp duty on instruments (Sec. 3) per the Schedule; modes of "
         "payment including electronic stamping (Sec. 10(3)); manner of writing on stamped paper (Secs. 13-15); "
         "time of stamping (Sec. 17); facts affecting duty (Sec. 27); allowance and refund for spoiled stamps "
         "(Secs. 47-52-A)."),
        ("The Registration Act, 1908", "Distinguishes compulsorily registrable (Sec. 17) and optionally registrable "
         "(Sec. 18) documents - decides whether the deed is produced in Digital E-Stamp or in Document Registration."),
        ("The Information Technology Act, 2000", "Legal recognition of electronic records (Sec. 4) and electronic "
         "signatures (Secs. 3A, 5) - basis for eSign by parties and the department digital seal."),
        ("The Karnataka Guarantee of Services to Citizens Act, 2011 (Sakala)", "Time-bound delivery for notified "
         "services - e-Stamp issue, cancellation and refund."),
        ("The Aadhaar (Targeted Delivery of Financial and Other Subsidies, Benefits and Services) Act, 2016",
         "Consent-based authentication / e-KYC and Aadhaar Data Vault obligations."),
    ], [2.2, 4.6])

    H("Relevant sections followed by the Department for Digital E-Stamp", 3)
    H("Karnataka Stamp Act, 1957 (selected sections)", 4)
    T(["Section", "Topic", "Relevance", "Refer section for implementation"], [
        ("Sec. 3", "Instruments chargeable with duty", "Duty is determined by the Schedule article of the "
         "instrument - drives two-level article selection", "FR - Instrument and article selection"),
        ("Secs. 4-6", "Several instruments / distinct matters / several descriptions", "One sub-article per "
         "e-Stamp; where an instrument falls under several descriptions, highest duty applies", "Business rules "
         "BR-EST-004, BR-EST-005"),
        ("Sec. 10 / 10(3)", "Duties how to be paid; electronic stamping", "Legal basis for a computerised / "
         "electronic stamping system authorised by Government", "Assumptions; Appendix B OP-01"),
        ("Sec. 10-A", "Payment of duty in cash by challan", "Khajane-II challan as payment mode; endorsement of "
         "duty paid", "FR - Payment"),
        ("Secs. 13, 14, 15", "Writing on stamped paper; one instrument per stamp; contravention = unstamped",
         "Certificate on the face of the deed, deed written below; one deed per UIN; lock on use",
         "FR - Final rendering; Verify and lock"),
        ("Sec. 17", "Instruments executed in Karnataka stamped before or at execution", "e-Stamp issued before eSign "
         "of the deed; issue timestamp shown", "Status model; BR-EST-012"),
        ("Sec. 27", "Facts affecting duty to be set forth in instrument", "Consideration, market value and other "
         "chargeable facts must be printed in the deed from captured data", "FR - Deed content; FR-TGS schema"),
        ("Secs. 33-34", "Examination / impounding; inadmissibility if not duly stamped", "SRO verification of "
         "the e-Stamp before acting on the instrument", "FR - Verify and lock"),
        ("Secs. 47-52-A", "Allowance for spoiled / unused stamps; refund", "Cancellation and refund workflow "
         "through DC of Stamps", "FR - Cancellation and refund"),
    ], [0.9, 1.7, 2.6, 1.6])
    H("Registration Act, 1908 (selected sections)", 4)
    T(["Section", "Topic", "Relevance", "Refer section for implementation"], [
        ("Sec. 17", "Documents of which registration is compulsory", "Such sub-articles are not offered in Digital "
         "E-Stamp; user is redirected to Document Registration", "BR-EST-001"),
        ("Sec. 17(1A)", "Agreements for sale with possession (Sec. 53A TP Act)", "Agreement of sale where "
         "possession is delivered is routed to Document Registration", "BR-EST-002"),
        ("Sec. 18", "Documents of which registration is optional", "Primary population of Digital E-Stamp "
         "instruments", "Scope"),
    ], [0.9, 1.7, 2.6, 1.6])
    H("Information Technology Act, 2000 (selected sections)", 4)
    T(["Section", "Topic", "Relevance", "Refer section for implementation"], [
        ("Sec. 3A", "Electronic signature", "Aadhaar eSign by parties", "FR - eSign and execution"),
        ("Sec. 4", "Legal recognition of electronic records", "Final PDF is the original instrument", "FR - Output"),
        ("Sec. 5", "Legal recognition of electronic signatures", "DSC path for non-e-KYC parties and department "
         "seal", "FR - eSign and execution; certificate issue"),
    ], [0.9, 1.7, 2.6, 1.6])

    H("Relevant rules followed by the Department for Digital E-Stamp", 3)
    H("Karnataka Stamp (Payment of Duty by means of e-Stamping) Rules, 2009", 4)
    T(["Rule", "Requirement", "Refer section for implementation"], [
        ("Rule 2", "Definitions - e-stamp certificate, unique identification number, Central Record Keeping "
         "Agency", "Glossary"),
        ("Rule 11(a)-(k)", "Minimum particulars on the certificate: UIN, date and time, duty in words and "
         "figures, purchaser, parties, instrument and property description, issuing user / location, security "
         "mark, signature and seal", "e-Stamp certificate field mapping; FR - Certificate issue"),
        ("Rule 11(l)-(p)", "Lock to prevent reuse; cancel spoiled certificates; departmental search, view and "
         "MIS; certificate details on server", "FR - Verify and lock; Cancellation; MIS"),
        ("Rules 20-22", "Application (Form-3), payment modes and issue after applicant verifies entered details",
         "FR - Preview and confirm; Payment"),
        ("Rule 23(2)", "A4 paper (210 x 297 mm), left margin 3.5 cm, right margin 1.5 cm", "FR-TGS - Output "
         "standards"),
        ("Rule 24", "Details of issued certificate available on server to authorised officers", "FR - Verify"),
        ("Rules 25-26", "Additional stamp duty on the same document, separate certificate", "FR - Additional duty"),
        ("Rule 27", "Certificate on the face of the instrument; instrument written below; no second instrument",
         "FR - Final rendering; BR-EST-013"),
        ("Rule 28", "UIN typed at the top centre of each page of the instrument", "FR-TGS - Stamping"),
        ("Rules 29-30", "Registering officer verifies details by UIN and locks the certificate", "FR - Verify and "
         "lock"),
        ("Rules 31-32", "Refund of spoiled / unused certificate on Form-4; DC of Stamps cancels, endorses, locks "
         "and refunds by treasury", "FR - Cancellation and refund"),
        ("Rule 44", "Reports to the Department", "FR - Reports and MIS"),
    ], [1.2, 3.8, 1.8])
    H("Karnataka Stamp Rules, 1958", 4)
    T(["Rule", "Requirement", "Refer section for implementation"], [
        ("Rules 16, 17", "Procedure and evidence for allowance of spoiled stamps (applied to e-Stamp by Rule 31)",
         "FR - Cancellation and refund"),
    ], [1.2, 3.8, 1.8])

    H("Relevant notifications issued by the Department for Digital E-Stamp", 3)
    T(["Instrument", "Date / No.", "Effect", "Refer section for implementation"], [
        ("Karnataka e-Stamping Rules", "No. RD 380 MUNOMU 2008, dated 08-04-2009", "Framework for e-stamp "
         "certificates, verification, lock, cancellation and refund", "Whole module"),
        ("Karnataka Stamp Act Schedule (as amended to 2022)", "Schedule to the Act, as amended", "Article and "
         "sub-article duty rates, minimum / maximum duty, family concessions", "FR - Duty computation"),
        ("Kaveri 2.0 Digital e-Stamp FRS v1.3", "Department FRS", "Baseline functional behaviour carried into "
         "Kaveri 3.0", "Current state; FR"),
        ("Notification authorising Kaveri digital e-Stamp under Sec. 10(3)", "[To be provided by Department]",
         "Legal recognition of Department-issued digital e-Stamp", "Appendix B OP-01"),
        ("Sakala notification for e-Stamp services", "[To be provided by Department]", "Service timelines for "
         "issue, cancellation and refund", "FR - Sakala"),
    ], [1.8, 1.6, 2.2, 1.2])

    H("E-Stamp statutory forms mapping", 3)
    T(["Form", "Rule ref", "Purpose", "Generated by"], [
        ("Form-3 - Application for e-stamp", "Rules 20, 25", "Applicant particulars, instrument, parties, "
         "property, duty", "Captured digitally; application summary page rendered by Template Service"),
        ("Annexure A1 - e-Stamp certificate", "Rules 11, 22", "Certificate with statutory particulars, QR and UIN",
         "Template Service (certificate template CERT-EST-01)"),
        ("Form-4 - Application for refund", "Rule 31", "Cancellation / refund request", "Kaveri online form; "
         "acknowledgement PDF by Template Service"),
        ("Form-5 - Daily register of applications", "Rule 21(3)", "Register of certificates issued", "System "
         "register (MIS report)"),
        ("Form-6 - Daily register of duty collected", "Rule 44", "Duty collected and remitted", "MIS report "
         "reconciled with Khajane-II"),
        ("Deed (instrument) templates", "Sec. 27; Rules 27-28", "Body of the instrument under the certificate",
         f"Template Service ({N_TEMPLATES} templates - Appendix C)"),
    ], [1.8, 0.9, 2.1, 2.0])
    H("e-Stamp certificate (Annexure A1) - field mapping", 5)
    T(["Field (statutory)", "Mandatory", "Validation / notes", "Kaveri 3.0 field name"], [
        ("Unique identification number (Rule 11(a))", "Yes", "System generated; never reused; format "
         "KA-DES-<YYYYMMDD>-<10-digit sequence>-<check digit> [format to be confirmed]", "estamp_uin"),
        ("Date and time of issue (11(b))", "Yes", "Timestamp at UIN allotment after payment confirmation (IST)",
         "issue_datetime"),
        ("Stamp duty in figures and words (11(c))", "Yes", "Equals duty computed; words in English and Kannada",
         "duty_amount, duty_amount_words"),
        ("Name and address of purchaser (11(d))", "Yes", "Logged-in applicant (or first party in assisted mode)",
         "purchaser_name, purchaser_address"),
        ("Names of parties (11(e))", "Yes", "All executants and claimants; representatives shown as 'represented "
         "by'", "parties[].name, parties[].role"),
        ("Description of instrument (11(f))", "Yes", "Nature of document, sub-article, article number",
         "document_nature, stamp_article, sub_article"),
        ("Description of property (11(g))", "Conditional", "Auto-generated schedule summary; 'Not applicable' for "
         "miscellaneous articles", "schedule_summary"),
        ("User-id of issuing official (11(h))", "Yes", "'KAVERI-ONLINE' for self-service; operator user-id in "
         "assisted mode", "issued_by_user"),
        ("Code and location of issuing branch (11(i))", "Yes", "'Kaveri Digital E-Stamp' or SRO code for assisted",
         "issuing_office_code"),
        ("Security mark (11(j))", "Yes", "QR code with signed payload (UIN + hash); visible security pattern",
         "qr_payload, document_hash"),
        ("Signature and seal (11(k))", "Yes", "Department digital seal (DSC / HSM) on the certificate page",
         "department_seal"),
        ("Challan / transaction reference", "Yes", "Khajane-II challan number and date (FRS item 35)",
         "challan_no, challan_date"),
    ], [1.9, 0.8, 2.5, 1.6])

    H("Nature of document and stamp article mapping", 3)
    P("The table lists the 59 natures of document used for two-level article selection, with the templates "
      "that render each nature. The sub-article determination rules (questions, thresholds and the resulting "
      "sub-article) are maintained in Acts_Rules/Document/Stamp_Duty_Article_Determination_Inputs.xlsx.")
    T(["DN", "Nature of document", "Schedule article", "Transaction intent", "Templates"],
      [(dn, name, art, INTENTS.get(it, it),
        ", ".join(DN_TEMPLATES.get(dn, [])) or "No dedicated template yet - upload own deed")
       for dn, name, it, art in SA.DOCUMENT_NATURES], [0.6, 2.0, 0.8, 1.8, 1.6], size=8)
    P(f"{len(DN_GAPS)} natures of document have no dedicated template yet and use the upload-own-deed path "
      "until templates are designed (Appendix B, OP-09).")

    H("Sakala - Karnataka Guarantee of Services", 3)
    P("Digital E-Stamp issue is an instant online service. Cancellation and refund of a digital e-Stamp involve "
      "officer action and shall be tracked under Sakala with a Guaranteed Service Certificate (GSC) number "
      "issued on submission. Service codes and timelines are to be confirmed by the Department "
      "(Appendix B, OP-07).")

    # ------------------------------------------------------------------ stakeholders
    H("Stakeholders and actors", 2)
    T(["Actor", "Description", "Primary goals", "Channel involvement"], [
        ("Applicant / purchaser (citizen)", "Person obtaining the e-Stamp; usually a party", "Correct deed and duty, "
         "quick issue", "Online: self-service; Assisted: visits SRO / Seva Kendra"),
        ("Parties - executant(s) and claimant(s)", "Persons executing / benefiting from the instrument, incl. "
         "representatives (POA, guardian, authorised signatory)", "Verify own data, eSign", "OTP, e-KYC, eSign or DSC"),
        ("Witnesses", "Attesting witnesses where the article requires (minimum two)", "Identity capture and "
         "attestation", "OTP, e-KYC; eSign [to be confirmed - OP-05]"),
        ("Assisted-mode operator (DEO / Seva Kendra)", "Department operator entering data for a walk-in "
         "applicant", "Accurate entry; applicant confirmation", "Office"),
        ("Sub-Registrar / District Registrar", "Officer who verifies and locks an e-Stamp when the instrument is "
         "presented", "Prevent reuse; confirm duty", "Office (Document Registration)"),
        ("Deputy Commissioner of Stamps", "Authority for cancellation and refund (Rules 31-32)", "Decide "
         "refund within Sakala timeline", "Office"),
        ("Template Admin (IT Cell) and Legal reviewer (DSR)", "Maintain templates; approve legal wording",
         "Correct, versioned templates", "Template Admin portal"),
        ("Template Generation Microservice", "Shared platform service rendering PDFs", "Consistent, auditable "
         "documents", "System"),
        ("Document Registration module", "Consumer of templates; consumer of e-Stamp UIN at presentation",
         "Same deed output; lock on use", "System"),
        ("External services", "UIDAI (e-KYC / eSign), C-DAC (name match), Khajane-II, Bhoomi, E-Swathu, E-Aasthi, "
         "ULMS, DigiLocker, SMS / email gateways, Sakala", "Reliable integration", "System"),
    ], [1.6, 2.0, 1.5, 1.7])

    # ------------------------------------------------------------------ glossary
    H("Definitions and glossary", 2)
    T(["Term", "Definition"], [
        ("Digital E-Stamp", "Electronic stamp certificate issued by Kaveri on payment of stamp duty, digitally "
         "sealed, attached to the instrument as its first page"),
        ("e-Stamp UIN", "Unique identification number of an e-Stamp certificate (Rule 11(a)); printed on every page"),
        ("Nature of document (DN)", "First level of article selection, e.g. DN-34 Leave and Licence"),
        ("Stamp sub-article", "Second level: the exact Schedule entry that fixes the duty, e.g. 30(1)(i)"),
        ("Optionally registrable", "Instrument whose registration is not compulsory (Sec. 18, Registration Act)"),
        ("Template", "Pre-designed, legally vetted deed layout with a field schema, identified as T-xxx-nn"),
        ("Template version", "Approved, immutable revision of a template; applications are pinned to one version"),
        ("Template Generation Microservice (TGS)", "Shared service that renders DRAFT or FINAL PDFs from a "
         "template and a data payload"),
        ("Payload", "Structured JSON sent by a consumer module to the TGS; fields defined in the template field "
         "workbook"),
        ("DRAFT render", "Watermarked preview PDF; not valid as a stamped instrument"),
        ("FINAL render", "PDF with certificate page, UIN on every page, QR and signature boxes"),
        ("Lock (on use)", "Status change that prevents the same e-Stamp being used for another instrument (Rule 30)"),
        ("GL / GR", "Government Land / Government Restricted flags on agricultural land records"),
        ("Name match", "C-DAC similarity score between Aadhaar name and property-register name; threshold 80%"),
        ("Khajane-II", "Government of Karnataka integrated treasury and payment system"),
        ("eSign", "Aadhaar-based online electronic signature through an ESP"),
        ("DSC", "Digital Signature Certificate (Class 3) for non-e-KYC parties and department seal"),
        ("Assisted mode", "Department operator enters data on behalf of a walk-in applicant"),
        ("Upload-own-deed", "Path where the applicant uploads a self-drafted deed PDF to be placed under the "
         "certificate"),
        ("CRA", "Central Record Keeping Agency appointed under the 2009 Rules (e.g. SHCIL) for physical e-stamps"),
    ], [1.8, 5.0])

    # ------------------------------------------------------------------ current state
    H("Current state", 2)
    P("Kaveri 2.0 offers a Digital e-Stamp for optionally registrable articles (FRS v1.3): the user selects a "
      "nature of document and sub-article, adds schedule and parties, the system calculates duty, the user types "
      "a free-text description of the instrument, pays through Khajane-II, all parties eSign, and a certificate "
      "PDF with QR is issued and pushed to DigiLocker. Physical e-stamp certificates from the CRA continue in "
      "parallel and are verified at the SRO.")
    H("As-Is pain points", 3)
    PAIN = [
        ("No deed drafting", "Only a free-text description is captured; the real deed is drafted outside and may "
         "not match the e-Stamp", "FRS item 32", "FR-EST deed content; FR-TGS"),
        ("Inconsistent deeds", "Deeds drafted by different writers vary widely; mandatory Sec. 27 facts are often "
         "missing", "Workshop", "Template schema with mandatory fields"),
        ("Duplicate document logic", "Document Registration and E-Stamp produce documents with separate code and "
         "layouts", "Architecture review", "Shared Template Generation Microservice"),
        ("Costly wording changes", "Legal wording changes require code releases in multiple modules", "IT Cell",
         "Template versioning and maker-checker publishing"),
        ("Edit-after-payment ambiguity", "The FRS is ambiguous on whether data can be edited after payment",
         "FRS item 39", "BR-EST-011 - data frozen at preview"),
        ("Payment interruption", "A payment interruption cancels the whole application and the user re-enters "
         "everything", "FRS item 40", "Resumable application; challan verification; FB-EST-001"),
        ("No online refund", "No online cancellation / refund for unused digital e-Stamps", "Rules 31-32",
         "FR - Cancellation and refund"),
        ("Manual lock", "Lock on use depends on manual SRO action", "Rule 30", "Automatic lock on registration event"),
        ("No additional duty", "No additional-duty linkage to the original e-Stamp", "Rules 25-26",
         "FR - Additional duty"),
        ("Verification privacy", "Verify page exposes full party names", "Privacy review", "Masked verification view"),
        ("No Kannada deeds", "Kannada versions of deeds are not available", "Workshop", "Bilingual templates (OP-04)"),
        ("Limited MIS", "MIS is limited to counts and revenue", "FRS section 7", "Extended MIS incl. lock, refund, "
         "template usage"),
    ]
    T(["Sr.No", "Pain Point", "Description", "Source", "Addressed in (this BRD)"],
      [(i + 1, *p) for i, p in enumerate(PAIN)], [0.5, 1.3, 2.6, 1.0, 1.4])

    # ------------------------------------------------------------------ to-be
    H("Future state (To-Be)", 2)
    H("Digital E-Stamp - deed generation and certificate issue", 3)
    H("Channel models", 4)
    T(["Service Type", "Online Activities", "Office Activities", "Mode"], [
        ("Digital E-Stamp - Online", "Login, e-KYC, article selection, schedule, parties, witnesses, deed terms, "
         "preview, payment, eSign, download / DigiLocker", "None", "Online"),
        ("Digital E-Stamp - Assisted", "Parties complete OTP / e-KYC / eSign on their own phones", "Operator enters "
         "data, applicant confirms preview, payment at counter (Khajane-II)", "Assisted (In Person)"),
        ("Digital E-Stamp - Upload own deed", "Same as Online, but deed PDF uploaded instead of template drafting",
         "Optional assisted entry", "Online / Assisted"),
    ], [1.5, 2.6, 1.8, 0.9])
    H("Process Diagram", 4)
    IMG(HERE / "EStamp_ToBe_Process.png", 6.4)
    H("Common intake steps", 5)
    T(["#", "Step", "Lane", "Notes"], [
        (1, "Login with mobile OTP; choose Digital E-Stamp", "Citizen / Operator", "Existing Kaveri account"),
        (2, "Applicant e-KYC (Aadhaar OTP) or continue with DSC", "Citizen / UIDAI", "Non-e-KYC path: passport / "
         "enrolment ID + DSC"),
        (3, "Select nature of document, answer determination questions, system resolves sub-article",
         "E-Stamp", "Only optionally registrable sub-articles; one per instrument"),
        (4, "Choose drafting option: template (system lists templates for the DN / sub-article) or upload own deed",
         "Citizen", "Template list from TGS registry"),
        (5, "Capture schedule: property register fetch or manual (movable / miscellaneous)", "Citizen / Land "
         "records", "GL / GR / alienation checks; relaxed fields per annexure"),
        (6, "Add parties and representatives; OTP and e-KYC per party; schedule allocation", "Citizen / Parties",
         "Minimum parties per article; name match"),
        (7, "Add witnesses where required (minimum two)", "Citizen / Witnesses", "Article-driven"),
        (8, "Compute stamp duty and show calculation breakdown", "E-Stamp", "MV / consideration / fixed / highest"),
        (9, "Enter deed-specific terms (template fields, clause options)", "Citizen", "Field help and samples"),
        (10, "Render DRAFT PDF and preview; applicant confirms; data frozen", "TGS / Citizen", "Rule 22 "
         "verification by applicant"),
    ], [0.4, 3.0, 1.4, 2.0])
    H("Online", 5)
    T(["#", "Step", "Lane", "Notes"], [
        (11, "Pay through Khajane-II (UPI / net banking / card / challan)", "Citizen / Khajane-II",
         "Verify-challan option for pending status"),
        (12, "Allot UIN, record issue date-time and duty paid", "E-Stamp", "Sec. 17 - stamped before execution"),
        (13, "Render FINAL PDF: certificate page, deed below / following, UIN on each page, QR, signature boxes",
         "TGS", "Rules 27-28"),
        (14, "Notify parties (SMS / email) to eSign; each party eSigns", "Parties / ESP", "Name on eSign must match "
         "e-KYC name"),
        (15, "After last signature: apply department seal, publish certificate to Verify, push to DigiLocker",
         "E-Stamp", "SMS to all parties"),
        (16, "Download signed PDF from dashboard", "Citizen", "Re-download any time"),
    ], [0.4, 3.0, 1.4, 2.0])
    H("Offline (Assisted - In Person)", 5)
    T(["#", "Step", "Lane", "Notes"], [
        ("A1", "Operator logs in with department credentials and opens Assisted E-Stamp", "Operator", "Role-based"),
        ("A2", "Steps 2-9 entered by operator; each party completes OTP / e-KYC on own device", "Operator / Parties",
         "Operator cannot bypass OTP"),
        ("A3", "Printed or on-screen DRAFT shown to applicant; applicant confirms with OTP", "Applicant",
         "Rule 22 confirmation captured digitally"),
        ("A4", "Payment at counter through Khajane-II", "Operator / Khajane-II", "Receipt issued"),
        ("A5", "Steps 12-16 as Online; parties eSign remotely or at the counter", "Parties", "Operator user-id "
         "printed on certificate (Rule 11(h))"),
    ], [0.4, 3.0, 1.4, 2.0])
    H("Application Status Model", 5)
    T(["Status", "Description", "Actor", "Next states"], [
        ("Draft", "Application started", "Applicant / Operator", "Details in progress; Abandoned"),
        ("Details in progress", "Article, schedule, parties, witnesses being captured", "Applicant", "Duty computed"),
        ("Duty computed", "Duty calculated and shown", "System", "Draft deed generated; Details in progress"),
        ("Draft deed generated", "DRAFT PDF rendered", "TGS", "Confirmed; Details in progress"),
        ("Confirmed", "Applicant confirmed preview; data frozen", "Applicant", "Payment pending"),
        ("Payment pending", "Redirected to Khajane-II", "System", "Paid; Payment failed; Payment unconfirmed"),
        ("Payment unconfirmed", "No final status from Khajane-II", "System", "Paid; Payment failed"),
        ("Payment failed", "Payment declined", "System", "Payment pending (retry)"),
        ("Paid - UIN allotted", "UIN and issue time recorded", "System", "Final rendered"),
        ("Final rendered", "FINAL PDF ready for signatures", "TGS", "eSign pending"),
        ("eSign pending", "One or more parties yet to sign", "Parties", "Issued"),
        ("Issued", "All signed; seal applied; certificate published", "System", "Locked; Cancellation requested; "
         "Additional duty linked"),
        ("Additional duty linked", "Supplementary e-Stamp issued for same instrument", "System", "Locked"),
        ("Locked", "Used in registration / verified by officer", "SRO / DR", "Terminal"),
        ("Cancellation requested", "Form-4 submitted", "Applicant", "Cancelled; Cancellation rejected"),
        ("Cancellation rejected", "DC rejected with reason", "DC of Stamps", "Issued"),
        ("Cancelled", "Certificate cancelled and locked; endorsement on PDF", "DC of Stamps", "Refund paid"),
        ("Refund paid", "Refund credited by treasury", "Treasury", "Terminal"),
        ("Abandoned", "Unpaid draft past retention period", "System", "Terminal (purged)"),
    ], [1.5, 2.6, 1.2, 1.5])

    H("Template Generation Microservice (shared)", 3)
    H("Context and responsibilities", 4)
    B([
        "The service owns templates, their versions and the rendering of PDFs. It does not own business data, "
        "does not compute duty and does not sign on behalf of a party.",
        "Consumers (Digital E-Stamp, Document Registration) own the workflow, send a validated data payload and "
        "store the returned document reference against their application.",
        "Templates are DOCX layouts with merge fields, repeating groups (parties, schedules, witnesses, payments) "
        "and conditional clauses. Digital E-Stamp uses Part I to Part IV variants (Legal Formats/Template/Digital "
        "E-Stamp); Document Registration uses Part I to Part V variants (Legal Formats/Template/Document "
        "Registration).",
        "Each template is published with a JSON schema generated from the field dictionary; FINAL renders are "
        "allowed only for Approved versions.",
    ])
    H("Architecture", 4)
    IMG(HERE / "Template_Service_Architecture.png", 6.8)
    H("Render sequence", 4)
    T(["#", "Step", "Lane", "Notes"], [
        (1, "Consumer requests template list for a DN / sub-article and locale", "Consumer -> TGS", "Returns active "
         "approved versions"),
        (2, "Consumer pins template_id + version on the application", "Consumer", "Pinned for the life of the "
         "application"),
        (3, "Consumer posts payload in DRAFT mode", "Consumer -> TGS", "Schema validation; field-level errors"),
        (4, "TGS binds data, renders PDF with DRAFT watermark, stores it, returns document_id + hash", "TGS",
         "Sync for typical documents"),
        (5, "After payment the consumer posts FINAL mode with UIN, QR payload and signer list", "Consumer -> TGS",
         "Same payload hash as confirmed draft, plus certificate data"),
        (6, "TGS renders certificate page + instrument, UIN on each page, signature boxes; returns document_id + "
         "hash", "TGS", "PDF/A-2b"),
        (7, "Consumer runs eSign / seal on the returned PDF and stores the signed version", "Consumer", "Signed PDF "
         "stored as new object; TGS output unchanged"),
    ], [0.4, 3.0, 1.4, 2.0])

    H("Additional duty, verification, lock, cancellation and refund", 3)
    H("Process Diagram", 4)
    IMG(HERE / "EStamp_Verify_Lock_Cancel.png", 6.4)
    H("Process steps", 4)
    T(["#", "Step", "Lane", "Notes"], [
        (1, "Any person verifies by UIN or QR; masked details and status shown", "Public", "Rule 24"),
        (2, "Presentant quotes UIN in Document Registration pre-registration", "Document Registration",
         "UIN must be Issued and parties must match"),
        (3, "SR verifies details (Rule 29); on registration, E-Stamp locks the UIN (Rule 30)", "SRO / E-Stamp",
         "Lock event recorded with registration number"),
        (4, "Applicant needing more duty raises Additional Duty against the original UIN", "Applicant",
         "New UIN linked to original; same instrument"),
        (5, "Applicant applies for cancellation / refund (Form-4) with reason and declarations", "Applicant",
         "Only Issued and unlocked UINs"),
        (6, "DC of Stamps verifies, approves or rejects with reason", "DC of Stamps", "Secs. 47-52-A time limits"),
        (7, "On approval: certificate cancelled and locked, 'CANCELLED' endorsement rendered on PDF, refund "
         "initiated", "E-Stamp / Treasury", "Refund in name of purchaser (Rule 32(3))"),
    ], [0.4, 3.0, 1.4, 2.0])

    H("What is new in Kaveri 3.0", 3)
    T(["#", "Capability", "What is new in Kaveri 3.0"], [
        (1, "Deed drafting", f"Template-driven deed generation ({N_ESTAMP_TEMPLATES} E-Stamp templates) instead of "
         "free-text description"),
        (2, "Shared rendering", "One Template Generation Microservice for E-Stamp and Document Registration"),
        (3, "Template governance", "Versioned templates, maker-checker legal approval, pinned versions"),
        (4, "Bilingual output", "English and Kannada template variants (subject to OP-04)"),
        (5, "Assisted channel", "Operator-assisted E-Stamp with party OTP / eSign on own device"),
        (6, "Upload own deed", "e-Stamp certificate for a self-drafted deed"),
        (7, "Additional duty", "Supplementary e-Stamp linked to the original UIN"),
        (8, "Automatic lock on use", "Lock triggered by Document Registration event"),
        (9, "Online cancellation and refund", "Form-4 workflow with DC of Stamps approval and Sakala tracking"),
        (10, "Masked public verification", "QR / UIN verification without exposing full PII"),
        (11, "Resumable payment", "Application survives payment interruption; challan verification"),
    ], [0.4, 1.8, 4.6])
    H("Rectified As-Is pain points", 4)
    T(["Sr.No", "Pain Point (As-Is)", "How rectified in Kaveri 3.0"],
      [(i + 1, f"{p[0]}: {p[1]}", p[3]) for i, p in enumerate(PAIN)], [0.5, 3.4, 2.9])

    # ------------------------------------------------------------------ FRs
    H("Functional requirements", 2)
    H("Digital E-Stamp", 3)
    n = 1
    H("Module entry, login and applicant e-KYC", 4)
    n = FR("FR-EST", n, "FR - Module entry", "Login / Applicant e-KYC", "FRS 4.1; Aadhaar Act", [
        ("System shall allow a logged-in citizen (mobile OTP) or an authorised operator to start a Digital E-Stamp "
         "application.", "Must", "Unauthenticated users cannot start; operator role required for Assisted mode"),
        ("System shall perform Aadhaar OTP e-KYC for the applicant and auto-populate name, age and gender as "
         "non-editable.", "Must", "Fields locked after e-KYC; Aadhaar stored only in vault reference"),
        ("System shall allow a non-e-KYC path using passport / enrolment ID with DSC before proceeding.", "Must",
         "DSC validated (Class 3, not expired / revoked) before next step", "FRS items 25-27, 43"),
        ("System shall save the application automatically after each step and allow resume from the dashboard.",
         "Must", "Resume returns to last completed step with data intact"),
    ])
    H("Instrument (nature of document) and article selection", 4)
    n = FR("FR-EST", n, "FR - Article selection", "Article selection", "Stamp Act Sec. 3; FRS items 5-7", [
        ("System shall present two-level selection: nature of document, then stamp sub-article.", "Must",
         "Three-level article tree not shown"),
        ("System shall offer only optionally registrable / non-registrable sub-articles; compulsorily registrable "
         "sub-articles shall be redirected to Document Registration.", "Must", "Sec. 17 sub-articles not "
         "selectable; redirect message bilingual", "Registration Act Secs. 17-18"),
        ("System shall allow exactly one sub-article per e-Stamp application.", "Must", "Second selection replaces "
         "first after confirmation"),
        ("System shall resolve the sub-article through determination questions (amount thresholds, relationship, "
         "possession, purpose) as defined in the article determination workbook.", "Should", "Answers stored; "
         "resolved sub-article shown with Schedule reference"),
        ("System shall list the templates mapped to the selected DN / sub-article and locale, and the "
         "upload-own-deed option.", "Must", "List fetched from TGS registry; only Approved versions"),
    ])
    H("Schedule / property details", 4)
    n = FR("FR-EST", n, "FR - Schedule", "Schedule details", "FRS items 8-15", [
        ("System shall map each sub-article to Agricultural, Non-agricultural or Miscellaneous property type.",
         "Must", "Property entry screen matches mapping"),
        ("System shall fetch agricultural property from Bhoomi and non-agricultural property from E-Swathu, "
         "E-Aasthi or ULMS.", "Must", "Owner, extent, boundaries (where available) imported"),
        ("System shall block GL and GR properties and allow only full extent for alienated agricultural land.",
         "Must", "Blocked with bilingual reason"),
        ("System shall respect source validations (tax unpaid, under mutation) and apply relaxed validations "
         "(name / extent / boundary editable) only for listed sub-articles.", "Must", "Relaxation table "
         "configurable"),
        ("System shall ask Kaveri village / road only where market valuation is required (e.g. Joint Development "
         "Agreement).", "Must", "Other articles do not show valuation inputs"),
        ("System shall auto-generate the schedule description for immovable property; boundaries for Bhoomi "
         "property are user-entered.", "Must", "Description printed in deed schedule"),
        ("System shall allocate schedules to claimants automatically (single claimant or schedule) or manually.",
         "Must", "Every schedule allocated before duty computation"),
    ])
    H("Party, representative and witness details", 4)
    n = FR("FR-EST", n, "FR - Parties", "Parties and witnesses", "FRS items 16-29", [
        ("System shall import executants from the property register for immovable-property articles.", "Must",
         "Imported parties non-editable unless relaxed"),
        ("System shall verify each party's mobile number by OTP and support e-KYC per party.", "Must",
         "Unverified party blocks progress"),
        ("System shall support representatives: POA holder, guardian of minor, organisation authorised signatory, "
         "in the nine party-type scenarios.", "Must", "Representative capacity printed in deed"),
        ("System shall perform C-DAC name matching for fetched executants with a configurable 80% threshold and "
         "warn below threshold.", "Must", "Score stored; not applied to organisations or editable names"),
        ("System shall enforce the minimum number of parties per article.", "Must", "Cannot proceed below minimum"),
        ("System shall require witnesses where the article mandates them, and a minimum of two witnesses whenever "
         "witnesses are added.", "Must", "Witness OTP verification"),
    ])
    H("Stamp duty computation", 4)
    n = FR("FR-EST", n, "FR - Duty", "Duty calculation", "Stamp Act Schedule; FRS items 30-31", [
        ("System shall compute duty using the sub-article rule: market value, consideration, fixed, or highest of "
         "applicable values.", "Must", "Breakdown shows basis, rate and amount"),
        ("System shall apply minimum and maximum duty for the 102 sub-articles where applicable and not for the "
         "69 where not applicable.", "Must", "Configurable per sub-article"),
        ("System shall apply family-relationship concessions and local-limit slabs per the Schedule (e.g. gift, "
         "settlement, release, lease to family).", "Must", "Relationship matrix validated"),
        ("System shall round duty to the next rupee and show it in figures and words (English / Kannada).", "Must",
         "Matches certificate"),
        ("System shall recompute duty when any duty-affecting input changes before confirmation.", "Must",
         "Stale duty never shown at preview"),
    ])
    H("Deed content and drafting", 4)
    n = FR("FR-EST", n, "FR - Deed content", "Deed details", "Stamp Act Sec. 27", [
        ("System shall display template-specific fields (from the template schema) with help text, samples and "
         "validations.", "Must", "Field set equals schema of pinned template version"),
        ("System shall pre-fill party, schedule, consideration and duty fields from captured data; these cannot "
         "be retyped in the deed.", "Must", "Sec. 27 facts always come from captured data"),
        ("System shall allow selection of optional clauses and entry of additional covenants in a bounded free-text "
         "field.", "Should", "Free text length limited; profanity / script validation"),
        ("System shall support the upload-own-deed path: PDF upload (max 20 MB, text-searchable preferred) with "
         "declaration that it matches the captured data.", "Should", "Upload virus-scanned; page count recorded"),
    ])
    H("Draft preview and confirmation", 4)
    n = FR("FR-EST", n, "FR - Preview", "Preview", "e-Stamping Rule 22", [
        ("System shall request a DRAFT render from the Template Generation Microservice and show it in the "
         "browser.", "Must", "Watermark 'DRAFT - NOT A STAMPED INSTRUMENT' on every page"),
        ("System shall require the applicant (in Assisted mode, by OTP) to confirm the preview; data is frozen "
         "after confirmation.", "Must", "No edits after confirmation; changes need a new application"),
        ("System shall store the confirmed payload hash and template version.", "Must", "FINAL render uses the "
         "same payload hash"),
    ])
    H("Payment", 4)
    n = FR("FR-EST", n, "FR - Payment", "Payment", "Stamp Act Sec. 10-A; Rule 21; FRS 37-41", [
        ("System shall redirect to Khajane-II for payment of the computed duty.", "Must", "Amount cannot be "
         "altered in transit"),
        ("System shall provide 'Verify challan' for indeterminate payments and continue once paid.", "Must",
         "No second payment while status unresolved"),
        ("System shall keep the application resumable after a payment interruption instead of cancelling it.",
         "Must", "Same application, new payment attempt"),
        ("System shall allot the UIN and issue date-time only after Khajane-II confirms payment.", "Must",
         "No UIN for unpaid applications"),
    ])
    H("Final rendering, eSign and certificate issue", 4)
    n = FR("FR-EST", n, "FR - Certificate", "eSign / Certificate", "Rules 11, 27, 28; IT Act Secs. 3A, 5", [
        ("System shall request a FINAL render with certificate page first and the instrument following, UIN on "
         "each page and QR code.", "Must", "Rules 27-28 satisfied on every page"),
        ("System shall invite every party (and witnesses if required) to eSign; names must match e-KYC names.",
         "Must", "Mismatch blocks eSign with reason"),
        ("System shall allow DSC signing for non-e-KYC parties.", "Must", "Signature validated"),
        ("System shall apply the department digital seal after the last signature and mark the certificate "
         "Issued.", "Must", "Seal verifiable in PDF reader"),
        ("System shall send SMS / email to all parties on issue and push the document to each party's "
         "DigiLocker.", "Must", "DigiLocker failure retried; does not block issue"),
        ("System shall allow download of the signed PDF from the dashboard at any time.", "Must", "Same hash as "
         "stored"),
    ])
    H("Additional stamp duty", 4)
    n = FR("FR-EST", n, "FR - Additional duty", "Additional duty", "Rules 25-26", [
        ("System shall allow an applicant to raise additional duty against an Issued, unlocked UIN.", "Should",
         "Original UIN referenced"),
        ("System shall issue a separate certificate for the additional duty, linked to the original.", "Should",
         "Verify page shows both"),
    ])
    H("Verify, lock on use and Document Registration interface", 4)
    n = FR("FR-EST", n, "FR - Verify and lock", "Verify Digital E-Stamp", "Rules 24, 29, 30; Stamp Act Secs. 13-15, 33", [
        ("System shall provide public verification by UIN or QR showing masked particulars and status.", "Must",
         "No full Aadhaar / mobile / address shown"),
        ("System shall expose to Document Registration an API to validate a UIN (status, parties, amount, "
         "instrument) and fetch the signed PDF.", "Must", "Response within NFR limits"),
        ("System shall lock the UIN automatically when Document Registration records registration using it, "
         "and allow manual lock by SR / DR / DC.", "Must", "Locked UIN cannot be used again; lock event audited"),
        ("System shall reject use of a Locked or Cancelled UIN.", "Must", "Clear bilingual reason"),
    ])
    H("Cancellation and refund", 4)
    n = FR("FR-EST", n, "FR - Cancellation", "Cancellation / refund", "Rules 31-32; Secs. 47-52-A", [
        ("System shall allow the purchaser to apply for cancellation / refund of an Issued, unlocked UIN with "
         "reason and declarations (Form-4).", "Must", "Locked UIN cannot be cancelled"),
        ("System shall route the request to the jurisdictional DC of Stamps with Sakala GSC.", "Must", "GSC "
         "number on acknowledgement"),
        ("System shall let the DC approve or reject with reason and refund amount after deductions per Act.",
         "Must", "Decision audited"),
        ("On approval, system shall cancel and lock the UIN and re-render the certificate with a 'CANCELLED' "
         "endorsement.", "Must", "Verify page shows Cancelled"),
        ("System shall initiate refund to the purchaser through the treasury and record status.", "Must",
         "Refund status visible to applicant"),
    ])
    H("Notifications", 4)
    n = FR("FR-EST", n, "FR - Notifications", "Notifications", "Department policy", [
        ("System shall send bilingual SMS / email for: OTP, party invitation, payment success / failure, eSign "
         "pending reminder, issue, DigiLocker push, lock, cancellation decision and refund.", "Must", "Templates "
         "approved by DLT; reasons stated in OTP message"),
        ("System shall send eSign reminders to pending parties at configurable intervals.", "Should",
         "Default every 48 hours, maximum 5 reminders"),
    ], with_ac=True)
    H("Reports and MIS", 4)
    n = FR("FR-EST", n, "FR - MIS", "MIS reports", "Rule 44; FRS section 7", [
        ("System shall provide counts of applications received, issued and in progress by date range, article and "
         "office.", "Must", "Filters: from / to date, article, channel, office"),
        ("System shall provide revenue MIS by nature of document, day-wise and total, reconciled with Khajane-II "
         "(Form-6).", "Must", "Totals equal Khajane-II settlement"),
        ("System shall report e-KYC vs non-e-KYC applications, lock statistics, cancellations, refunds, additional "
         "duty, template usage and render failures.", "Should", "Export to Excel / PDF"),
        ("System shall provide the Form-5 daily register of certificates issued.", "Must", "Printable"),
    ])
    H("Sakala integration", 4)
    n = FR("FR-EST", n, "FR - Sakala", "GSC acknowledgement", "Karnataka Guarantee of Services Act, 2011", [
        ("System shall generate a GSC number for cancellation / refund applications and sync status to Sakala.",
         "Must", "Status sync on each transition"),
    ])

    H("Template Generation Microservice", 3)
    m = 1
    H("Template registry and versioning", 4)
    m = FR("FR-TGS", m, "FR - TGS registry", "Template Admin", "Department governance", [
        ("Service shall store each template with ID (e.g. T-SAL-01), name, layout (K / S / A), DN and "
         "sub-article mapping, locale, version and status (Draft, Under review, Approved, Retired).", "Must",
         "Registry query returns all attributes"),
        ("Approved versions shall be immutable; any change creates a new version.", "Must", "Edit of Approved "
         "version rejected"),
        ("Publishing shall require maker (Template Admin) and checker (DSR Legal reviewer) approval with effective "
         "date.", "Must", "Maker cannot approve own version"),
        ("Service shall keep in-flight applications on their pinned version; a retired version can still render "
         "FINAL for pinned applications unless flagged 'legally withdrawn'.", "Must", "Pinned render succeeds; "
         "withdrawn version returns explicit error"),
        ("Service shall generate a JSON schema per version from the field dictionary.", "Must", "Schema "
         "published with version"),
    ])
    H("Render API", 4)
    m = FR("FR-TGS", m, "FR - TGS render", "System", "Architecture", [
        ("Service shall render DRAFT and FINAL PDFs from template_id, version, locale and payload.", "Must",
         "Both modes available to authorised consumers"),
        ("Service shall validate the payload against the schema and return field-level errors.", "Must",
         "HTTP 422 with path, code and message per field"),
        ("Service shall support synchronous rendering and asynchronous jobs with callback for large documents.",
         "Must", "Async used above configured page / size limit"),
        ("Service shall return document_id, SHA-256 hash, page count and template version for each render.",
         "Must", "Values match stored object"),
        ("Service shall render repeating groups (parties, schedules, witnesses, payments) and conditional clauses.",
         "Must", "Golden-file tests per template"),
        ("Rendering shall be deterministic: same version and payload give identical content.", "Must", "Content "
         "hash equal excluding render timestamp"),
    ])
    H("Output standards", 4)
    m = FR("FR-TGS", m, "FR - TGS output", "System", "e-Stamping Rules 23, 27, 28", [
        ("Output shall be PDF/A-2b, A4, left margin 3.5 cm and right margin 1.5 cm, with embedded English and "
         "Kannada Unicode fonts.", "Must", "PDF/A validator passes; Kannada text selectable"),
        ("DRAFT output shall carry a diagonal watermark on every page.", "Must", "Watermark present on all pages"),
        ("FINAL output shall print the e-Stamp UIN at the top centre of every page, page 'x of y', QR on the "
         "certificate page and named signature boxes for each signer.", "Must", "Rules 27-28 checks pass"),
        ("Service shall be able to render endorsement overlays (e.g. CANCELLED, LOCKED - used in registration "
         "no. ...) on an existing document as a new version.", "Should", "Original preserved"),
    ])
    H("Security, storage and audit", 4)
    m = FR("FR-TGS", m, "FR - TGS security", "System", "IT security policy", [
        ("Service shall authenticate consumers with OAuth2 client credentials over mTLS and restrict each "
         "consumer to its allowed templates and modes.", "Must", "Unauthorised template returns 403"),
        ("Service shall store rendered PDFs in the object store and return time-limited signed URLs.", "Must",
         "URL expiry configurable (default 15 minutes)"),
        ("Service shall not keep payload PII beyond the render job; logs shall mask PII.", "Must", "Log review "
         "shows masked values"),
        ("Service shall audit every render: consumer, application reference, template version, mode, hash, "
         "time.", "Must", "Audit searchable by application reference"),
    ])

    H("Business rules", 3)
    T(["Rule ID", "Rule", "Source"], [
        ("BR-EST-001", "Compulsorily registrable sub-articles are not available in Digital E-Stamp.", "Reg. Act Sec. "
         "17; FRS item 6"),
        ("BR-EST-002", "Agreement of sale with possession delivered is routed to Document Registration.",
         "Reg. Act Sec. 17(1A)"),
        ("BR-EST-003", "One sub-article per e-Stamp application.", "FRS item 7"),
        ("BR-EST-004", "An instrument relating to several distinct matters needs separate duty for each - "
         "separate e-Stamp applications.", "Stamp Act Sec. 5"),
        ("BR-EST-005", "Where an instrument falls under several descriptions, the highest duty applies.",
         "Stamp Act Sec. 6"),
        ("BR-EST-006", "No transaction on GL / GR land; full extent only for alienated agricultural land.",
         "FRS item 10"),
        ("BR-EST-007", "Name-match threshold 80% (configurable); warning only, not a block, at capture.",
         "FRS item 24"),
        ("BR-EST-008", "eSign name must match e-KYC name; otherwise eSign is blocked.", "FRS item 42"),
        ("BR-EST-009", "Minimum two witnesses when witnesses are required or added.", "FRS item 28"),
        ("BR-EST-010", "Min / max duty applies only to configured sub-articles (102 yes / 69 no).", "FRS item 31"),
        ("BR-EST-011", "No change to captured data after preview confirmation.", "Rule 22; FRS item 39 clarified"),
        ("BR-EST-012", "UIN and issue time are allotted only after payment and before any party eSigns.",
         "Stamp Act Sec. 17"),
        ("BR-EST-013", "One instrument per UIN; a Locked or Cancelled UIN cannot be used.", "Stamp Act Secs. "
         "13-15; Rules 27, 30"),
        ("BR-EST-014", "FINAL render only from an Approved template version pinned on the application.",
         "Template governance"),
        ("BR-EST-015", "Cancellation only for Issued, unlocked UINs; refund in the name of the purchaser.",
         "Rules 31-32"),
        ("BR-EST-016", "Unpaid drafts are purged after a configurable period (default 30 days).", "Data "
         "minimisation [to be confirmed]"),
    ], [1.0, 4.2, 1.6])

    H("User interface (high-level)", 3)
    P("Wireframe links: [Figma / prototype URLs]. Mock screens per document nature: "
      "Acts_Rules/Document/Mock_Screens/index.html.")
    P("Bilingual: All labels [EN / KN] - content manager sign-off.")
    T(["Screen / step", "Purpose", "Channel", "Statutory alignment", "Notes"], [
        ("Login / start application", "Authenticated entry", "Both", "", "Common step 1"),
        ("Applicant e-KYC / DSC", "Identity of purchaser", "Both", "Rule 11(d)", "Common step 2"),
        ("Article selection", "DN and sub-article", "Both", "Stamp Act Sec. 3", "Determination questions"),
        ("Drafting option", "Template list / upload own deed", "Both", "", "From TGS registry"),
        ("Schedule details", "Property fetch / entry", "Both", "Rule 11(g)", "GL / GR checks"),
        ("Parties", "Executants, claimants, representatives", "Both", "Rule 11(e)", "OTP, e-KYC"),
        ("Witnesses", "Attestation", "Both", "", "Article-driven"),
        ("Duty calculation", "Breakdown", "Both", "Schedule", "Read-only"),
        ("Deed details", "Template fields and clauses", "Both", "Sec. 27", "Schema-driven form"),
        ("Preview and confirm", "DRAFT PDF", "Both", "Rule 22", "OTP confirmation in Assisted"),
        ("Payment", "Khajane-II", "Both", "Sec. 10-A; Rule 21", "Verify challan"),
        ("eSign tracker", "Signature status per party", "Both", "IT Act Sec. 3A", "Reminders"),
        ("Dashboard / download", "Signed PDF, DigiLocker status", "Both", "", ""),
        ("Verify Digital E-Stamp", "Public verification", "Public", "Rule 24", "Masked"),
        ("Cancellation / refund", "Form-4", "Online", "Rule 31", "GSC"),
        ("DC of Stamps worklist", "Decide refund", "Office", "Rule 32", ""),
        ("Template Admin", "Manage templates", "Office", "", "Maker-checker"),
    ], [1.4, 1.6, 0.7, 1.3, 1.8])

    H("Integrations", 3)
    P("Interface requirements: [API list to be finalised by the Architect - must include Template Generation "
      "Microservice, Khajane-II, UIDAI e-KYC / eSign, C-DAC name match, land records, DigiLocker, Document "
      "Registration, Sakala and SMS / email gateways].")
    T(["Integration", "Direction", "Purpose", "Channel", "Owner", "Status"], [
        ("Template Generation Microservice", "Outbound", "Template list, DRAFT / FINAL render, endorsement overlay",
         "Both", "Kaveri IT Cell", "New"),
        ("Khajane-II", "Outbound / callback", "Duty payment, challan verification, refund", "Both", "DSR / "
         "Treasury", "Existing"),
        ("UIDAI e-KYC (via AUA / KUA)", "Outbound", "Applicant, party and witness e-KYC", "Both", "DSR", "Existing"),
        ("eSign service provider", "Outbound", "Party eSign on FINAL PDF", "Both", "DSR", "Existing"),
        ("DSC / HSM signing", "Internal", "Department seal; non-e-KYC DSC", "Both", "Kaveri IT Cell", "Existing"),
        ("C-DAC name matching", "Outbound", "Name similarity score", "Both", "DSR", "Existing"),
        ("Bhoomi", "Outbound", "Agricultural property and GL / GR flags", "Both", "Revenue Dept", "Existing"),
        ("E-Swathu / E-Aasthi / ULMS", "Outbound", "Non-agricultural property", "Both", "RDPR / ULB", "Existing"),
        ("Document Registration module", "Inbound / outbound", "Validate UIN, fetch PDF, lock on use", "Both",
         "Kaveri IT Cell", "New"),
        ("DigiLocker", "Outbound", "Push issued document", "Both", "DSR", "Existing"),
        ("Sakala", "Outbound", "GSC and status for cancellation / refund", "Online", "DSR", "New"),
        ("SMS / email gateways", "Outbound", "Notifications", "Both", "DSR", "Existing"),
        ("CRA (physical e-stamp)", "Outbound", "Verify / lock physical e-stamps (Document Registration)", "Office",
         "DSR", "Existing - out of scope here"),
    ], [1.6, 0.9, 2.0, 0.6, 0.9, 0.8])
    H("Template Generation Microservice - API contract (business view)", 4)
    T(["Endpoint", "Method", "Purpose", "Key request fields", "Key response"], [
        ("/v1/templates", "GET", "List approved templates", "dn, sub_article, locale, consumer",
         "template_id, version, name, layout, schema_url"),
        ("/v1/templates/{id}/versions/{v}/schema", "GET", "Payload schema", "-", "JSON schema"),
        ("/v1/render", "POST", "Render DRAFT / FINAL", "template_id, version, locale, mode, consumer, "
         "application_ref, payload, certificate{uin, qr_payload}, signers[]", "document_id, sha256, pages, "
         "status or job_id"),
        ("/v1/jobs/{job_id}", "GET", "Async job status", "-", "status, document_id, errors[]"),
        ("/v1/documents/{document_id}", "GET", "Signed URL for PDF", "-", "url, expires_at, sha256"),
        ("/v1/documents/{document_id}/endorse", "POST", "Overlay endorsement", "endorsement_type, text, "
         "reference", "new document_id, sha256"),
        ("/v1/admin/templates", "POST / PUT", "Upload / submit / approve versions", "file, metadata, action",
         "version, status"),
    ], [1.7, 0.6, 1.2, 1.9, 1.4])

    H("Data requirements", 3)
    H("Core entities (logical)", 4)
    B([
        "EStampApplication, NatureOfDocument, StampSubArticle, DeterminationAnswer, Schedule (Property), "
        "ScheduleAllocation, Party, Representative, Witness, NameMatchResult, DutyComputation, DeedContent "
        "(template fields and clauses), TemplateRef (id, version, locale), RenderedDocument (document_id, mode, "
        "hash), Payment / Challan, EStampCertificate (UIN, issue time, amount, status), SignatureEvent, "
        "VerificationEvent, LockEvent, AdditionalDutyLink, CancellationRequest, Refund, Notification, AuditLog.",
        "Template Service: Template, TemplateVersion, TemplateSchema, ApprovalRecord, RenderJob, RenderAudit.",
    ])
    H("Payload field dictionary", 4)
    P("The Digital E-Stamp template payload is composed of the common Part I to Part IV blocks below plus "
      "template-specific fields. The full dictionary (field key, label, type, format, mandatory, source, sample) is "
      "in Legal Formats/Template/Digital E-Stamp/Digital_EStamp_Template_Fields.xlsx; the Document Registration "
      "variants (Part I to Part V) are described in Legal Formats/Template/Document Registration.")
    est_blocks = [b for b in ET.BLOCKS if any(b == x for t in ET.ESTAMP for x, _ in ET.estamp_blocks(t))]
    T(["Block", "Name", "Part", "Fields"],
      [(k, ET.BLOCKS[k]["name"], "Part III" if k == "A1" else ET.BLOCKS[k]["part"], len(ET.BLOCKS[k]["fields"]))
       for k in est_blocks], [0.8, 3.0, 1.6, 0.8])
    H("Retention", 4)
    B([
        "Issued certificates, signed deeds, lock and cancellation records: permanent.",
        "Unpaid drafts and DRAFT renders: purged after the configured period (default 30 days) [to be confirmed].",
        "Render audit and payment logs: as per Department retention policy [to be confirmed].",
    ])
    H("Migration (high level)", 4)
    T(["Topic"], [
        ("Kaveri 2.0 digital e-Stamp register (UIN, parties, amount, status, lock) for verification and lock "
         "continuity",),
        ("Kaveri 2.0 article / sub-article master and duty rules re-validated against the 2022 Schedule",),
    ], [6.8])

    H("Requirements traceability matrix (RTM) - template", 3)
    T(["Req ID", "Act/Rule/Form", "Requirement summary", "BRD section", "UI screen", "Test case ID", "Status"],
      [(r[0], r[1], (r[2][:90] + "...") if len(r[2]) > 90 else r[2], r[3], r[4],
        "TC-" + r[0].split("-")[1] + "-___", "Draft") for r in RTM],
      [0.8, 1.1, 2.2, 1.0, 0.9, 0.7, 0.5], size=7)

    # ------------------------------------------------------------------ NFR
    H("Non-functional requirements", 2)
    P("To ensure the Digital E-Stamp module and the Template Generation Microservice operate reliably, securely "
      "and efficiently, the following non-functional parameters shall be built in and verified before go-live.")
    H("Performance and Scalability", 3)
    P("Response time, concurrency and render throughput targets are mandatory load-test gates. Peaks are expected "
      "at month-end and financial-year-end.")
    T(["Req ID", "Requirement", "Priority", "Acceptance criteria"], [
        ("NFR-EST-PERF-001", "Standard UI interactions within 2 seconds; external API calls within 5 seconds.",
         "Must", "p95 at the concurrency in PERF-002"),
        ("NFR-EST-PERF-002", "Support 5,000 concurrent citizen sessions and 1,000 concurrent internal users.",
         "Must", "Load-test evidence without SLA breach"),
        ("NFR-TGS-PERF-001", "DRAFT render of up to 20 pages within 3 seconds; FINAL within 5 seconds (p95).",
         "Must", "Measured at 50 renders per second sustained"),
        ("NFR-TGS-PERF-002", "Service scales horizontally to serve both consumers at combined peak.", "Must",
         "Auto-scaling verified; no cross-consumer starvation (per-consumer quotas)"),
    ], [1.3, 3.2, 0.6, 1.7])
    H("Security and Data Privacy", 3)
    P("Encryption, PII protection, document integrity and auditability are baseline controls. Aadhaar / e-KYC use "
      "remains subject to Department and UIDAI approval.")
    T(["Req ID", "Requirement", "Priority", "Acceptance criteria"], [
        ("NFR-EST-SEC-001", "TLS 1.3 in transit; AES-256 at rest for data and documents.", "Must", "Confirmed in "
         "SDC design"),
        ("NFR-EST-SEC-002", "Final PDFs are tamper-evident: SHA-256 stored, department seal, signed QR payload.",
         "Must", "Altered PDF fails verification"),
        ("NFR-EST-SEC-003", "Service-to-service calls use mTLS and OAuth2 client credentials.", "Must", "No "
         "anonymous internal endpoints"),
        ("NFR-EST-PRIV-001", "Aadhaar numbers stored only in the Aadhaar Data Vault; masked everywhere else.",
         "Must", "Logs and DB views show masked values"),
        ("NFR-EST-PRIV-002", "Public verification shows only masked names and no contact details.", "Must",
         "Privacy review sign-off"),
    ], [1.3, 3.2, 0.6, 1.7])
    H("System Availability and Error Handling", 3)
    P("The module shall fail safely when a dependency is unavailable and shall not allow duplicate duty "
      "collection.")
    T(["Req ID", "Requirement", "Priority", "Acceptance criteria"], [
        ("NFR-EST-AVA-001", "99.5% monthly availability for E-Stamp and the Template Service (excluding planned "
         "maintenance).", "Must", "Monitoring report"),
        ("NFR-EST-AVA-002", "If an external integration is down, the user sees a bilingual reason and the "
         "application stays resumable with no data loss.", "Must", "Fault-injection test"),
        ("NFR-EST-PAY-001", "On payment timeout, poll Khajane-II until terminal status; pay action disabled "
         "meanwhile.", "Must", "No double debit in test"),
        ("NFR-TGS-AVA-001", "Template Service deployed active-active across two nodes / zones; RPO and RTO as "
         "per SDC standard [to be confirmed].", "Must", "Failover test"),
    ], [1.3, 3.2, 0.6, 1.7])
    H("Security Audit and Compliance (VAPT Policy)", 3)
    P("Vulnerability Assessment and Penetration Testing is mandatory. Automated scanning alone does not satisfy "
      "the gate.")
    T(["Req ID", "Requirement", "Priority", "Acceptance criteria"], [
        ("NFR-EST-VAPT-001", "Comprehensive VAPT by a CERT-In empanelled auditor before go-live and after major "
         "releases.", "Must", "Report accepted by Security"),
        ("NFR-EST-VAPT-002", "Scope covers web, APIs (including Template Service), payment flows and responsive "
         "UI; OWASP Top 10 and business-logic tests (UIN reuse, duty tampering).", "Must", "Scope statement and "
         "evidence"),
        ("NFR-EST-VAPT-003", "Critical and high findings closed before production.", "Must", "Retest evidence"),
    ], [1.3, 3.2, 0.6, 1.7])

    # ------------------------------------------------------------------ risks
    H("Risk and Mitigation Strategy", 2)
    P("Operational, technical, legal and adoption risks with their mitigations.")
    T(["Risk ID", "Risk", "Mitigation", "Related requirements"], [
        ("RS-EST-001", "User selects the wrong article and pays wrong duty", "Determination questions; plain-"
         "language help; duty breakdown before payment", "FR-EST article selection; BR-EST-004/005"),
        ("RS-EST-002", "Template wording legally inaccurate", "Legal reviewer approval (checker); versioning; "
         "rapid withdrawal flag", "FR-TGS registry"),
        ("RS-EST-003", "Template changes mid-application", "Version pinning; FINAL only from pinned version",
         "BR-EST-014"),
        ("RS-EST-004", "Reuse or forgery of e-Stamp", "Lock on use; signed QR; hash; public verification",
         "BR-EST-013; NFR-EST-SEC-002"),
        ("RS-EST-005", "Payment debited but UIN not allotted", "Reconciliation job with Khajane-II; verify challan; "
         "alerts", "NFR-EST-PAY-001; FB-EST-001"),
        ("RS-EST-006", "A party never eSigns, instrument stuck", "Reminders; cancellation and refund path",
         "FR-EST notifications; cancellation"),
        ("RS-EST-007", "Kannada text renders incorrectly", "Embedded Unicode fonts; golden-file tests", "FR-TGS "
         "output"),
        ("RS-EST-008", "Template Service outage stops both E-Stamp and Document Registration", "Active-active "
         "deployment; per-consumer quotas; queued async renders", "NFR-TGS-AVA-001; FB-EST-003"),
        ("RS-EST-009", "Legal standing of Department-issued digital e-Stamp challenged", "Obtain Sec. 10(3) "
         "notification before go-live", "Appendix B OP-01"),
    ], [0.9, 1.8, 2.5, 1.6])

    H("System Fallbacks & Error Handling", 2)
    P("The system shall handle exceptions and integration failures without losing citizen data.")
    T(["Req ID", "Requirement", "Priority", "Acceptance criteria"], [
        ("FB-EST-001", "Payment timeout: poll Khajane-II; show 'Verify challan'; never collect twice.", "Must",
         "Same as NFR-EST-PAY-001"),
        ("FB-EST-002", "eSign / DSC failure: keep status eSign pending; retry resumes signing only.", "Must",
         "No re-render or re-payment"),
        ("FB-EST-003", "Template Service unavailable: DRAFT shows 'Preview temporarily unavailable', application "
         "saved; after payment the FINAL render is queued and retried automatically with operations alert.",
         "Must", "Paid applications always reach Final rendered"),
        ("FB-EST-004", "Land-record service unavailable: retry; manual entry only for relaxed sub-articles.",
         "Must", "No bypass of GL / GR checks"),
        ("FB-EST-005", "UIDAI unavailable: offer DSC path or resume later.", "Must", "Application resumable"),
        ("FB-EST-006", "DigiLocker / SMS failure: retry queue; issue not blocked; email / in-app fallback.",
         "Should", "Retry log"),
    ], [1.0, 3.5, 0.6, 1.7])

    # ------------------------------------------------------------------ training
    H("Training and Change Management", 2)
    P("Moving from the Kaveri 2.0 description-only e-Stamp to template-driven deeds needs structured change "
      "management for citizens, document writers and department staff.")
    H("Target audience", 3)
    P("Citizens and document writers / advocates; assisted-mode operators; Sub-Registrars and District Registrars "
      "(verify and lock); Deputy Commissioners of Stamps (cancellation and refund); Template Admins and Legal "
      "reviewers; IT helpdesk.")
    H("Training delivery", 3)
    B([
        "Role-based workshops for operators, SRO staff and DC offices covering the status model, lock and refund.",
        "Template Admin training on versioning, schema generation, maker-checker and withdrawal.",
        "SOPs, quick-reference cards and short videos in English and Kannada.",
    ])
    H("Citizen change management", 3)
    P("Guided article selection, field-level help, sample filled deeds and a clear duty breakdown; public "
      "awareness of QR verification and that a used e-Stamp is locked.")
    H("Post-Go-Live support", 3)
    P("Dedicated hyper-care for 90 days covering payment reconciliation, eSign failures and template defects, "
      "with a fast-track template correction process.")

    # ------------------------------------------------------------------ appendices
    H("Appendix A - References", 2)
    B([
        "The Karnataka Stamp Act, 1957 and Schedule (as amended to 2022)",
        "The Karnataka Stamp (Payment of Duty by means of e-Stamping) Rules, 2009 - No. RD 380 MUNOMU 2008",
        "The Karnataka Stamp Rules, 1958",
        "The Registration Act, 1908",
        "The Information Technology Act, 2000",
        "Kaveri 2.0 Digital e-Stamp FRS v1.3",
        "Legal Formats/Template/Digital E-Stamp - Part I to IV templates and Digital_EStamp_Template_Fields.xlsx",
        "Legal Formats/Template/Document Registration - Part I to V templates and field workbook",
        "Acts_Rules/Document - Stamp_Duty_Article_Determination_Inputs.xlsx and Mock_Screens",
        "OWASP Top 10 - https://owasp.org/",
        "Sakala - https://sakala.kar.nic.in/; Karnataka Guarantee of Services to Citizens Act, 2011",
    ])
    H("Appendix B - Open points for Department decision", 2)
    T(["ID", "Open point", "Proposed default", "Owner"], [
        ("OP-01", "Legal instrument authorising the Department's Kaveri digital e-Stamp under Sec. 10(3) / 2009 "
         "Rules (Rules contemplate a CRA).", "Obtain notification or amendment before go-live", "DSR / Law"),
        ("OP-02", "Should Digital E-Stamp also sell e-Stamps for compulsorily registrable deeds (duty pre-payment)?",
         "No - handled in Document Registration", "DSR"),
        ("OP-03", "Upload-own-deed path - permitted for all articles?", "Yes, with declaration", "DSR"),
        ("OP-04", "Kannada templates - all templates or a priority list?", "Priority list for phase 1", "DSR / "
         "Legal"),
        ("OP-05", "Do witnesses eSign the digital deed?", "Yes, where the article requires witnesses", "DSR"),
        ("OP-06", "UIN format and check digit", "As proposed in certificate mapping", "IT Cell"),
        ("OP-07", "Sakala service codes and timelines for cancellation / refund", "As notified", "DSR / Sakala"),
        ("OP-08", "Retention period for unpaid drafts", "30 days", "DSR"),
        ("OP-09", f"Dedicated templates for the {len(DN_GAPS)} natures of document without one "
         f"({', '.join(DN_GAPS)})", "Design in priority order of transaction volume; upload own deed meanwhile",
         "DSR / Legal"),
    ], [0.6, 3.2, 2.0, 1.0])
    H("Appendix C - Template catalogue", 2)
    P(f"{N_TEMPLATES} templates are maintained in Legal Formats/Template. 'Generated in' shows which consumer "
      "module uses the template.")
    T(["Template ID", "Template name", "DN", "Article", "Layout", "Generated in"], TEMPLATE_ROWS,
      [0.8, 2.3, 0.6, 0.8, 0.5, 1.8], size=7)

    H("Acceptance and sign-off of BRD", 2)
    T(["Role", "Name", "Signature / Date"], [
        ("Product Owner", "", ""), ("Domain Expert", "", ""), ("Kaveri IT Cell", "", ""),
    ], [2.2, 2.4, 2.2])

    doc.save(str(OUT))
    print("saved", OUT)
    print("FR-EST", n - 1, "FR-TGS", m - 1, "RTM", len(RTM), "templates", N_TEMPLATES, "estamp", N_ESTAMP_TEMPLATES)


if __name__ == "__main__":
    DG.estamp_process(HERE / "EStamp_ToBe_Process.png")
    DG.verify_lock_process(HERE / "EStamp_Verify_Lock_Cancel.png")
    DG.template_service_architecture(HERE / "Template_Service_Architecture.png")
    build()
