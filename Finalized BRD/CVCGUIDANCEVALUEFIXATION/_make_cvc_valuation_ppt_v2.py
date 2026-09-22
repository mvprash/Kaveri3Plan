# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v2.pptx from v1.

Enriches slides already present in CVC_Valuation_Module_BRD_Presentation.pptx
with complete BRD detail for the same listed rows only:
  - Slide 3: Key Sections + Relevant Rules (as listed)
  - Slide 4: Key Notifications & Amendments (as listed)
Does NOT add extra rows or slides.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v2.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)

MARGIN = Inches(0.5)
CONTENT_W = Inches(12.333)

# Only the rows already listed on slide 3 of the source deck — fuller detail from BRD.
SECTIONS = [
    [
        "Sec. 2(ac)",
        "Definition — Central Valuation Committee",
        "CVC is the statutory owner of guideline rate masters and approvals",
        "7.i (actors / entity master); 7.ii–7.iii (approval authority)",
    ],
    [
        "Sec. 45-B",
        "Constitution of CVC; estimation, publication and revision of market value "
        "guidelines; district / sub-district sub-committees; CVC final authority for "
        "policy and methodology",
        "Primary legal basis for General Revision and for fixation of guidance value "
        "for new individual projects (DR acceptance of citizen request under Sec. 45-B "
        "framework)",
        "7.ii General Revision; 7.iii Individual Project Fixation",
    ],
]

RULES = [
    [
        "Stamp Act Sec. 45-B (practice)",
        "CVC under IGR; publish and revise market value guidelines; constitute market "
        "valuation sub-committees",
        "7.ii, 7.iii — revision / fixation workflows",
    ],
]

# Only the rows already listed on slide 4 of the source deck — fuller detail from BRD.
NOTIFICATIONS = [
    [
        "Act 8 of 2003",
        "w.e.f. 1-4-2003",
        "Substituted / strengthened Sec. 45-B CVC framework",
        "Legal foundation for Valuation Module revision cycles",
    ],
    [
        "CVC / IGR order & DIGR notification (operational)",
        "Per revision cycle",
        "Triggers General Revision; DIGR issues notification for fixing guidance value; "
        "CVC gazette fixes effective date",
        "7.ii.a–7.ii.l — General Revision process",
    ],
]


def _set_run(para, text: str, *, size: float, bold: bool, color: RGBColor, font: str = "Calibri"):
    para.clear()
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, x, y, w, h, text, *, size, bold=False, color=INK, font="Calibri",
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, cambria=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    _set_run(para, text, size=size, bold=bold, color=color, font="Cambria" if cambria else font)
    return box


def style_cell(cell, text: str, *, fill: RGBColor, font_color: RGBColor, size: float,
               bold: bool = False):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.text_frame.word_wrap = True
    cell.text_frame.margin_left = Emu(73152)
    cell.text_frame.margin_right = Emu(73152)
    cell.text_frame.margin_top = Emu(18288)
    cell.text_frame.margin_bottom = Emu(18288)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = cell.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    for p in list(cell.text_frame.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    _set_run(para, text, size=size, bold=bold, color=font_color, font="Calibri")


def add_table(slide, left, top, width, height, headers, rows, col_widths, font_size=10.5):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    for ci, h in enumerate(headers):
        style_cell(table.cell(0, ci), h, fill=NAVY, font_color=WHITE, size=11, bold=True)

    for ri, row in enumerate(rows):
        fill = WHITE if ri % 2 == 0 else ALT
        for ci, val in enumerate(row):
            style_cell(
                table.cell(ri + 1, ci),
                val,
                fill=fill,
                font_color=INK,
                size=font_size,
                bold=(ci == 0),
            )

    row_h = int(height / (len(rows) + 1))
    for row in table.rows:
        row.height = row_h
    return table_shape


def remove_shapes(slide, predicate):
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        if predicate(shape):
            sp_tree.remove(shape._element)


def rebuild_slide3(slide):
    def drop(shape):
        if shape.has_table:
            return True
        if hasattr(shape, "text") and shape.text.strip() == "Relevant Rules":
            return True
        return False

    remove_shapes(slide, drop)

    # Full detail for the two sections already listed on the source slide.
    add_table(
        slide,
        MARGIN,
        Inches(1.76),
        CONTENT_W,
        Inches(2.55),
        ["Section", "Topic", "BRD relevance", "Refer section 7"],
        SECTIONS,
        [Inches(1.35), Inches(3.55), Inches(4.05), Inches(3.05)],
        font_size=10,
    )

    add_textbox(
        slide,
        MARGIN,
        Inches(4.45),
        CONTENT_W,
        Inches(0.45),
        "Relevant Rules",
        size=30,
        bold=True,
        color=NAVY,
        cambria=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # Full detail for the single rule already listed on the source slide.
    add_table(
        slide,
        MARGIN,
        Inches(4.95),
        CONTENT_W,
        Inches(1.55),
        ["Rule / provision", "Requirement", "Refer section 7"],
        RULES,
        [Inches(3.4), Inches(5.2), Inches(3.4)],
        font_size=10.5,
    )


def rebuild_slide4(slide):
    """Key Notifications & Amendments — same two instruments, full BRD columns."""

    def drop(shape):
        return shape.has_table

    remove_shapes(slide, drop)

    add_table(
        slide,
        MARGIN,
        Inches(1.76),
        CONTENT_W,
        Inches(4.8),
        ["Instrument", "Date / No.", "Effect", "BRD relevance"],
        NOTIFICATIONS,
        [Inches(3.2), Inches(1.7), Inches(3.9), Inches(3.2)],
        font_size=11,
    )


def bump_version_strings(prs: Presentation):
    replacements = {
        "Version 1": "Version 2",
        "18-09-2026": "19-09-2026",
        "Document ID: BRD-K3-CVC-GVF-001     ·     Version 1     ·     Status: Draft     ·     Last updated 18-09-2026":
            "Document ID: BRD-K3-CVC-GVF-001     ·     Version 2     ·     Status: Draft     ·     Last updated 19-09-2026",
        "Document ID: BRD-K3-CVC-GVF-001  ·  Version 1  ·  18-09-2026":
            "Document ID: BRD-K3-CVC-GVF-001  ·  Version 2  ·  19-09-2026",
    }
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full = para.text
                new = full
                for old, repl in replacements.items():
                    if old in new:
                        new = new.replace(old, repl)
                if new != full and para.runs:
                    para.runs[0].text = new
                    for run in para.runs[1:]:
                        run.text = ""


def main():
    if not SRC.exists():
        raise SystemExit(f"Source presentation not found: {SRC}")

    # Prefer writing over DST; if locked (e.g. open in PowerPoint), write a sibling file.
    out = DST
    try:
        shutil.copy2(SRC, out)
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v2_updated.pptx"
        shutil.copy2(SRC, out)
        print(f"NOTE: {DST.name} is locked; writing {out.name} instead.")

    prs = Presentation(str(out))
    rebuild_slide3(prs.slides[2])
    rebuild_slide4(prs.slides[3])
    bump_version_strings(prs)
    try:
        prs.save(str(out))
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v2_updated.pptx"
        prs.save(str(out))
        print(f"NOTE: save locked; wrote {out.name} instead.")

    print(f"Wrote {out}")
    print(f"Slides: {len(prs.slides)} (unchanged count)")
    for label, s in (("Slide 3", prs.slides[2]), ("Slide 4", prs.slides[3])):
        print(f"--- {label} ---")
        for sh in s.shapes:
            if sh.has_table:
                tbl = sh.table
                print(f"TABLE {len(tbl.rows)}x{len(tbl.columns)}")
                for ri, row in enumerate(tbl.rows):
                    print(" ", [c.text[:55].replace("\n", " ") for c in row.cells])


if __name__ == "__main__":
    main()
