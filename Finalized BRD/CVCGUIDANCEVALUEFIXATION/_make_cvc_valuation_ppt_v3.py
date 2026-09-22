# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v3.pptx.

Starts from v1 and applies:
  - Slide 3: full detail for listed sections + listed rule, then notifications /
    amendments listed underneath in italic + accent color (with Effect).
  - Slide 4: full detail for the same listed notifications (Effect + BRD relevance).
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v3.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)
# Distinct accent for notification / amendment callouts under sections & rules
ACCENT = RGBColor(0x1F, 0x6B, 0x7A)

MARGIN = Inches(0.5)
CONTENT_W = Inches(12.333)

SECTIONS = [
    [
        "Sec. 2(ac)",
        "Definition — Central Valuation Committee",
        "CVC is the statutory owner of guideline rate masters and approvals",
        "7.i (actors / entity master); 7.ii–7.iii (approval authority)",
    ],
    [
        "Sec. 45-B",
        "Constitution of CVC; estimation, publication and revision of market value "
        "guidelines; district / sub-district sub-committees; CVC final authority for "
        "policy and methodology",
        "Primary legal basis for General Revision and for fixation of guidance value "
        "for new individual projects (DR acceptance of citizen request under Sec. 45-B "
        "framework)",
        "7.ii General Revision; 7.iii Individual Project Fixation",
    ],
]

RULES = [
    [
        "Stamp Act Sec. 45-B (practice)",
        "CVC under IGR; publish and revise market value guidelines; constitute market "
        "valuation sub-committees",
        "7.ii, 7.iii — revision / fixation workflows",
    ],
]

# Same instruments listed on slide 4 — shown under sections/rules on slide 3.
NOTIFICATIONS = [
    [
        "Act 8 of 2003",
        "w.e.f. 1-4-2003",
        "Substituted / strengthened Sec. 45-B CVC framework",
        "Legal foundation for Valuation Module revision cycles",
    ],
    [
        "CVC / IGR order & DIGR notification (operational)",
        "Per revision cycle",
        "Triggers General Revision; DIGR issues notification for fixing guidance value; "
        "CVC gazette fixes effective date",
        "7.ii.a–7.ii.l — General Revision process",
    ],
]


def _set_run(
    para,
    text: str,
    *,
    size: float,
    bold: bool,
    color: RGBColor,
    font: str = "Calibri",
    italic: bool = False,
):
    # Caller may clear; this only appends a run when used from add_runs helpers.
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def _replace_para_text(
    para,
    text: str,
    *,
    size: float,
    bold: bool,
    color: RGBColor,
    font: str = "Calibri",
    italic: bool = False,
):
    para.clear()
    _set_run(para, text, size=size, bold=bold, color=color, font=font, italic=italic)


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size,
    bold=False,
    color=INK,
    font="Calibri",
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    cambria=False,
    italic=False,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    _replace_para_text(
        para,
        text,
        size=size,
        bold=bold,
        color=color,
        font="Cambria" if cambria else font,
        italic=italic,
    )
    return box


def style_cell(
    cell,
    text: str,
    *,
    fill: RGBColor,
    font_color: RGBColor,
    size: float,
    bold: bool = False,
):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.text_frame.word_wrap = True
    cell.text_frame.margin_left = Emu(73152)
    cell.text_frame.margin_right = Emu(73152)
    cell.text_frame.margin_top = Emu(18288)
    cell.text_frame.margin_bottom = Emu(18288)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = cell.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    for p in list(cell.text_frame.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    _replace_para_text(para, text, size=size, bold=bold, color=font_color, font="Calibri")


def add_table(slide, left, top, width, height, headers, rows, col_widths, font_size=10.5):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    for ci, h in enumerate(headers):
        style_cell(table.cell(0, ci), h, fill=NAVY, font_color=WHITE, size=11, bold=True)

    for ri, row in enumerate(rows):
        fill = WHITE if ri % 2 == 0 else ALT
        for ci, val in enumerate(row):
            style_cell(
                table.cell(ri + 1, ci),
                val,
                fill=fill,
                font_color=INK,
                size=font_size,
                bold=(ci == 0),
            )

    row_h = int(height / (len(rows) + 1))
    for row in table.rows:
        row.height = row_h
    return table_shape


def remove_shapes(slide, predicate):
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        if predicate(shape):
            sp_tree.remove(shape._element)


def add_notification_callouts(slide, x, y, w, h):
    """Italic accent-colored lines: Instrument (Date) — Effect."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(36000)
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    # Label
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    para.space_after = Pt(4)
    _replace_para_text(
        para,
        "Related notifications / amendments",
        size=11,
        bold=True,
        color=ACCENT,
        font="Calibri",
        italic=True,
    )

    for instrument, date_no, effect, _brd in NOTIFICATIONS:
        line = f"•  {instrument}  ({date_no})  —  {effect}"
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(3)
        p.line_spacing = 1.15
        _set_run(p, line, size=10.5, bold=False, color=ACCENT, font="Calibri", italic=True)

    return box


def rebuild_slide3(slide):
    def drop(shape):
        if shape.has_table:
            return True
        if hasattr(shape, "text") and shape.text.strip() == "Relevant Rules":
            return True
        return False

    remove_shapes(slide, drop)

    # Compact sections + rules so italic notifications fit underneath.
    add_table(
        slide,
        MARGIN,
        Inches(1.55),
        CONTENT_W,
        Inches(2.15),
        ["Section", "Topic", "BRD relevance", "Refer section 7"],
        SECTIONS,
        [Inches(1.35), Inches(3.55), Inches(4.05), Inches(3.05)],
        font_size=9.5,
    )

    add_textbox(
        slide,
        MARGIN,
        Inches(3.78),
        CONTENT_W,
        Inches(0.38),
        "Relevant Rules",
        size=26,
        bold=True,
        color=NAVY,
        cambria=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    add_table(
        slide,
        MARGIN,
        Inches(4.18),
        CONTENT_W,
        Inches(1.15),
        ["Rule / provision", "Requirement", "Refer section 7"],
        RULES,
        [Inches(3.4), Inches(5.2), Inches(3.4)],
        font_size=10,
    )

    add_notification_callouts(slide, MARGIN, Inches(5.45), CONTENT_W, Inches(1.45))


def rebuild_slide4(slide):
    def drop(shape):
        return shape.has_table

    remove_shapes(slide, drop)

    add_table(
        slide,
        MARGIN,
        Inches(1.76),
        CONTENT_W,
        Inches(4.8),
        ["Instrument", "Date / No.", "Effect", "BRD relevance"],
        NOTIFICATIONS,
        [Inches(3.2), Inches(1.7), Inches(3.9), Inches(3.2)],
        font_size=11,
    )


def bump_version_strings(prs: Presentation, version: str = "3", date: str = "19-09-2026"):
    # Replace any prior Version N / date strings to the target version.
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full = para.text
                if not full:
                    continue
                new = full
                # Normalize known version lines
                if "BRD-K3-CVC-GVF-001" in new and "Version" in new:
                    if "Last updated" in new:
                        new = (
                            f"Document ID: BRD-K3-CVC-GVF-001     ·     Version {version}     ·     "
                            f"Status: Draft     ·     Last updated {date}"
                        )
                    else:
                        new = f"Document ID: BRD-K3-CVC-GVF-001  ·  Version {version}  ·  {date}"
                elif "Version 1" in new:
                    new = new.replace("Version 1", f"Version {version}")
                elif "Version 2" in new:
                    new = new.replace("Version 2", f"Version {version}")
                if "18-09-2026" in new:
                    new = new.replace("18-09-2026", date)
                if new != full and para.runs:
                    para.runs[0].text = new
                    for run in para.runs[1:]:
                        run.text = ""


def safe_write(prs: Presentation, path: Path) -> Path:
    try:
        prs.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_updated" + path.suffix)
        prs.save(str(alt))
        print(f"NOTE: {path.name} locked; wrote {alt.name}")
        return alt


def main():
    if not SRC.exists():
        raise SystemExit(f"Source presentation not found: {SRC}")

    out = DST
    try:
        shutil.copy2(SRC, out)
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v3_updated.pptx"
        shutil.copy2(SRC, out)
        print(f"NOTE: {DST.name} locked on copy; using {out.name}")

    prs = Presentation(str(out))
    rebuild_slide3(prs.slides[2])
    rebuild_slide4(prs.slides[3])
    bump_version_strings(prs, version="3", date="19-09-2026")
    out = safe_write(prs, out)

    print(f"Wrote {out}")
    print(f"Slides: {len(prs.slides)}")
    s = prs.slides[2]
    print("--- Slide 3 ---")
    for sh in s.shapes:
        if sh.has_table:
            tbl = sh.table
            print(f"TABLE {len(tbl.rows)}x{len(tbl.columns)}")
        elif hasattr(sh, "text") and sh.text.strip():
            t = sh.text.strip().replace("\n", " | ")
            if "notification" in t.lower() or "Act 8" in t or "DIGR" in t or "Relevant Rules" in t:
                print("TEXT:", t[:200])


if __name__ == "__main__":
    main()
