# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v5.pptx from v4.

Adds Market Valuation Sub-Committee composition from
Karnataka Stamp (Constitution of C.V.C. etc.) Rules, 2003 — Rule 4.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation_v4.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v5.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)
ACCENT = RGBColor(0x1F, 0x6B, 0x7A)

MARGIN = Inches(0.5)
CONTENT_W = Inches(12.333)
FOOTER_Y = Inches(7.14)

MEMBER_DEPTS = [
    "Department of Revenue",
    "Department of Survey and Settlement",
    "Public Works Department",
    "Municipal Councils or Town Panchayaths",
]


def _run(para, text, *, size, bold=False, italic=False, color=INK, font="Calibri"):
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def add_textbox(
    slide, x, y, w, h, text, *, size, bold=False, italic=False, color=INK,
    font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, cambria=False,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    para.clear()
    _run(
        para, text, size=size, bold=bold, italic=italic, color=color,
        font="Cambria" if cambria else font,
    )
    return box


def add_rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.adjustments[0] = 0.08
    return shape


def chrome(slide, kicker: str, title: str, page: str):
    add_textbox(slide, MARGIN, Inches(0.42), Inches(10.0), Inches(0.28),
                kicker, size=11.5, bold=True, color=MUTED)
    add_textbox(slide, MARGIN, Inches(0.70), Inches(11.5), Inches(0.55),
                title, size=24, bold=True, color=NAVY, cambria=True, anchor=MSO_ANCHOR.MIDDLE)
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.3), Inches(0.40), Inches(0.55), Inches(0.55))
    badge.fill.solid()
    badge.fill.fore_color.rgb = NAVY
    badge.line.fill.background()
    add_textbox(slide, Inches(12.3), Inches(0.40), Inches(0.55), Inches(0.55),
                page, size=15, bold=True, color=WHITE, cambria=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, MARGIN, FOOTER_Y, Inches(8.5), Inches(0.28),
                "VALUATION MODULE (CVC)  |  BRD WALKTHROUGH  ·  KAVERI 3.0", size=8.5, color=MUTED)
    add_textbox(slide, Inches(12.25), FOOTER_Y, Inches(0.6), Inches(0.28),
                page, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)


def insert_slide_at(prs: Presentation, index: int):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    sld_id_lst = prs.slides._sldIdLst
    sld_ids = list(sld_id_lst)
    last = sld_ids[-1]
    sld_id_lst.remove(last)
    sld_id_lst.insert(index, last)
    return prs.slides[index]


def add_role_card(slide, x, y, w, h, label, value, value_size=13):
    add_rect(slide, x, y, w, h, ALT)
    add_textbox(slide, x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.26),
                label, size=10, bold=True, color=ACCENT, italic=True)
    add_textbox(slide, x + Inches(0.18), y + Inches(0.38), w - Inches(0.36), Inches(0.55),
                value, size=value_size, bold=True, color=NAVY, cambria=True)


def build_subcommittee_slide(slide, page: str):
    chrome(
        slide,
        "LEGAL & REGULATORY REFERENCE",
        "Market Valuation Sub-Committees — Composition",
        page,
    )

    add_textbox(
        slide, MARGIN, Inches(1.28), CONTENT_W, Inches(0.32),
        "Karnataka Stamp (Constitution of C.V.C. etc.) Rules, 2003 — Rule 4  ·  "
        "Constituted by CVC for each sub-district and district for estimation & revision "
        "of market value guidelines",
        size=10, italic=True, color=MUTED,
    )

    card_w = Inches(3.95)
    card_h = Inches(1.05)
    gap = Inches(0.18)
    y = Inches(1.68)
    add_role_card(slide, MARGIN, y, card_w, card_h,
                  "HEAD (SUB-DISTRICT)", "Tahsildar of the concerned taluk")
    add_role_card(slide, MARGIN + card_w + gap, y, card_w, card_h,
                  "MEMBER SECRETARY", "Sub-Registrar of the sub-district")
    add_role_card(slide, MARGIN + 2 * (card_w + gap), y, card_w, card_h,
                  "OFFICE", "Office of the Sub-Registrar", value_size=13)

    # Members panel
    panel_y = Inches(2.90)
    add_rect(slide, MARGIN, panel_y, Inches(6.0), Inches(2.55), ALT)
    add_textbox(
        slide, MARGIN + Inches(0.22), panel_y + Inches(0.14), Inches(5.5), Inches(0.32),
        "Members drawn from", size=13, bold=True, color=NAVY, cambria=True,
    )
    box = slide.shapes.add_textbox(
        MARGIN + Inches(0.22), panel_y + Inches(0.55), Inches(5.5), Inches(1.8)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, dept in enumerate(MEMBER_DEPTS):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(8)
        para.clear()
        _run(para, f"•  {dept}", size=13, color=INK)

    # Hierarchy / control panel
    right_x = MARGIN + Inches(6.25)
    add_rect(slide, right_x, panel_y, Inches(6.08), Inches(2.55), ALT)
    add_textbox(
        slide, right_x + Inches(0.22), panel_y + Inches(0.14), Inches(5.6), Inches(0.32),
        "Administrative & supervisory control", size=13, bold=True, color=NAVY, cambria=True,
    )
    hierarchy = [
        ("Administrative control", "Registrar of the district"),
        ("Supervisory control", "Central Valuation Committee"),
        ("District Sub-Committees", "Function under the Registrar of the district"),
    ]
    hy = panel_y + Inches(0.55)
    for label, value in hierarchy:
        add_textbox(slide, right_x + Inches(0.22), hy, Inches(5.6), Inches(0.22),
                    label, size=10, bold=True, italic=True, color=ACCENT)
        add_textbox(slide, right_x + Inches(0.22), hy + Inches(0.22), Inches(5.6), Inches(0.32),
                    value, size=13, bold=True, color=NAVY, cambria=True)
        hy += Inches(0.62)

    add_textbox(
        slide, MARGIN, Inches(5.60), CONTENT_W, Inches(0.85),
        "Sub-Registrar (Member Secretary) looks after administration of the Sub-Committees, "
        "correspondence, and compilation of market-value data as per Sub-Committee resolutions.  "
        "CVC office (Rule 3): Office of the Inspector General of Registration and Commissioner "
        "of Stamps, or such other place as the Committee decides.",
        size=10.5, italic=True, color=ACCENT,
    )


def set_shape_number(shape, number: str, *, size: float | None = None, color=None, font=None):
    for para in shape.text_frame.paragraphs:
        if para.runs:
            para.runs[0].text = number
            if size is not None:
                para.runs[0].font.size = Pt(size)
            if color is not None:
                para.runs[0].font.color.rgb = color
            if font is not None:
                para.runs[0].font.name = font
            for r in para.runs[1:]:
                r.text = ""
        else:
            para.clear()
            _run(para, number, size=size or 15, bold=True, color=color or WHITE, font=font or "Cambria")


def renumber_content_slides(prs: Presentation):
    BADGE_LEFT_MIN = 10_500_000
    FOOTER_TOP_MIN = 6_200_000
    page = 2
    for idx in range(1, len(prs.slides) - 1):
        slide = prs.slides[idx]
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = shape.text.strip()
            if not re.fullmatch(r"\d{1,2}", t):
                continue
            if shape.left >= BADGE_LEFT_MIN and shape.top < 1_200_000:
                set_shape_number(shape, str(page), size=15, color=WHITE, font="Cambria")
            elif shape.left >= BADGE_LEFT_MIN and shape.top >= FOOTER_TOP_MIN:
                set_shape_number(shape, str(page), size=8.5, color=MUTED, font="Calibri")
        page += 1


def bump_version(prs: Presentation, version: str = "5", date: str = "19-09-2026"):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full = para.text
                if not full or "BRD-K3-CVC-GVF-001" not in full:
                    continue
                if "Last updated" in full:
                    new = (
                        f"Document ID: BRD-K3-CVC-GVF-001     ·     Version {version}     ·     "
                        f"Status: Draft     ·     Last updated {date}"
                    )
                else:
                    new = f"Document ID: BRD-K3-CVC-GVF-001  ·  Version {version}  ·  {date}"
                if para.runs:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""


def safe_save(prs: Presentation, path: Path) -> Path:
    try:
        prs.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_updated" + path.suffix)
        prs.save(str(alt))
        print(f"NOTE: {path.name} locked; wrote {alt.name}")
        return alt


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    out = DST
    try:
        shutil.copy2(SRC, out)
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v5_updated.pptx"
        shutil.copy2(SRC, out)

    prs = Presentation(str(out))

    # After CVC Composition (index 4 in v4) → insert Sub-Committee as next slide
    # v4: 0 Title, … 4 CVC Composition, 5 Stakeholders
    insert_at = 5
    slide = insert_slide_at(prs, insert_at)
    build_subcommittee_slide(slide, page="6")
    renumber_content_slides(prs)
    bump_version(prs, version="5", date="19-09-2026")

    out = safe_save(prs, out)
    print(f"Wrote {out}")
    print(f"Slides: {len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        texts = [
            sh.text.strip().split("\n")[0][:70]
            for sh in s.shapes
            if hasattr(sh, "text") and sh.text.strip()
        ]
        print(f"  {i}. {texts[1] if len(texts) > 1 else (texts[0] if texts else '')}")


if __name__ == "__main__":
    main()
