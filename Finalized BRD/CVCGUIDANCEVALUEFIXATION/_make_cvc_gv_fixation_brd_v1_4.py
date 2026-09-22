# -*- coding: utf-8 -*-
"""Create BRD_CVC_Guidance_Value_Fixation_v1.4.docx from v1.3.

Process / FR updates:
  (1) After SR refers evidence, SR does NOT propose first — takes matter to
      Sub-committee; then SR formally proposes; then preliminary notification
      is published (also via DIPR / Karnataka Rajya Patra) after Sub-committee.
  (2) Objections / public-opinion period shall be configurable.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\CVCGUIDANCEVALUEFIXATION")
SRC = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.3.docx"
DST = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.4.docx"
OUT_VERSION = "1.4"
OUT_DATE = "22-09-2026"


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def replace_paragraph_text(paragraph, new_text: str) -> None:
    # Preserve style; clear runs and write one run
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def update_document_control_version(doc: Document) -> None:
    for table in doc.tables:
        if len(table.columns) != 2:
            continue
        keys = [r.cells[0].text.strip() for r in table.rows]
        if "Document ID" not in keys:
            continue
        for row in table.rows:
            key = row.cells[0].text.strip()
            if key == "Version":
                set_cell_text(row.cells[1], OUT_VERSION)
            if key == "Last updated":
                set_cell_text(row.cells[1], OUT_DATE)


def append_version_history_row(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [c.text.strip() for c in table.rows[0].cells]
        if headers[:3] != ["Version", "Date", "Author"]:
            continue
        for r in table.rows[1:]:
            if r.cells[0].text.strip() == OUT_VERSION:
                return
        row = table.add_row()
        values = [
            OUT_VERSION,
            OUT_DATE,
            "Nandha Kumar",
            "General Revision flow: evidence → Sub-committee → SR proposal → "
            "preliminary notification (DIPR / Karnataka Rajya Patra); objections "
            "period configurable (FR-CVC-GR-004/006/007/009/017)",
            "Prashanth",
        ]
        for i, val in enumerate(values):
            if i < len(row.cells):
                set_cell_text(row.cells[i], val)
        return


def set_fr_text(doc: Document, req_id: str, new_text: str) -> bool:
    for table in doc.tables:
        if not table.rows or table.rows[0].cells[0].text.strip() != "Req ID":
            continue
        for row in table.rows[1:]:
            if row.cells[0].text.strip() == req_id:
                set_cell_text(row.cells[1], new_text)
                return True
    return False


def insert_fr_after(doc: Document, after_id: str, req_id: str, text: str, priority: str = "Must") -> None:
    for table in doc.tables:
        if not table.rows or table.rows[0].cells[0].text.strip() != "Req ID":
            continue
        ids = [r.cells[0].text.strip() for r in table.rows[1:]]
        if after_id not in ids:
            continue
        for r in table.rows[1:]:
            if r.cells[0].text.strip() == req_id:
                set_cell_text(r.cells[1], text)
                set_cell_text(r.cells[2], priority)
                return
        after_idx = None
        for i, r in enumerate(table.rows):
            if r.cells[0].text.strip() == after_id:
                after_idx = i
                break
        new_row = table.add_row()
        set_cell_text(new_row.cells[0], req_id, bold=True)
        set_cell_text(new_row.cells[1], text)
        set_cell_text(new_row.cells[2], priority)
        tr = new_row._tr
        table._tbl.remove(tr)
        table.rows[after_idx]._tr.addnext(tr)
        return
    raise SystemExit(f"Could not insert {req_id} after {after_id}")


def rewrite_process_steps(doc: Document) -> None:
    """Replace General Revision numbered steps that describe proposal / notice."""
    replacements = {
        "SR refers historical registration data": (
            "SR refers historical registration data (transaction data), KSRSAC GIS "
            "data / maps, Developer publications, RTC data, khata data, and town "
            "planning inputs (system evidence pack / attachments). SR does not lodge "
            "a formal rate proposal at this stage."
        ),
        "SR proposes the guidance value for each territorial unit": (
            "SR takes the evidence pack and discussion draft to the Sub-committee "
            "(members: Tahsildar, SR, Secretary, ADLR, PWD AEE, etc.). The "
            "Sub-committee discusses the evidence and methodology; minutes are recorded."
        ),
        "SR presents the proposal to the Sub-committee": (
            "After Sub-committee deliberation, SR formally proposes the guidance value "
            "for each territorial unit / rate slab in scope (proposal with justification, "
            "reflecting Sub-committee directions)."
        ),
        "SR publishes the proposed revised guideline value for public opinion": (
            "After the Sub-committee and SR proposal, a preliminary notification of the "
            "proposed revised guideline values is published for public objections / "
            "opinions. The preliminary notification shall also be published through "
            "integration with the Department of Information and Public Relations (DIPR), "
            "Government of Karnataka, and Karnataka Rajya Patra (e-gazette / official "
            "publication channel as applicable). The objections / public-opinion period "
            "shall be configurable in the system (default may be 15 days, but not hard-coded)."
        ),
        "After 15 days, the Sub-committee again discusses": (
            "After the configurable objections period closes, the Sub-committee again "
            "discusses and revises the proposed guidance value only for the objections "
            "received; revised proposal and objection disposal notes are saved."
        ),
    }

    # Also handle cascade / summary strings
    cascade_old_parts = [
        (
            "Cascade: DIGR notification → DR forwards to SR → SR proposes → "
            "Sub-committee → public opinion (15 days) → Sub-committee (objections) → "
            "DR → CVC approval → effective date → Kaveri adoption.",
            "Cascade: DIGR notification → DR forwards to SR → SR refers evidence → "
            "Sub-committee → SR proposes → preliminary notification (DIPR / Karnataka "
            "Rajya Patra) with configurable objections period → Sub-committee "
            "(objections) → DR → CVC approval → e-gazette → effective date → Kaveri adoption.",
        ),
        (
            "General Revision of market value guidelines for the entire State "
            "(IGR/CVC trigger → DIGR notification → DR → SR proposal → Sub-committee → "
            "15-day public opinion → CVC approval → effective date → Kaveri adoption).",
            "General Revision of market value guidelines for the entire State "
            "(IGR/CVC trigger → DIGR notification → DR → SR evidence → Sub-committee → "
            "SR proposal → preliminary notification via DIPR / Karnataka Rajya Patra "
            "(configurable objections period) → CVC approval → e-gazette → effective "
            "date → Kaveri adoption).",
        ),
        (
            "Publication of proposed revised guideline values for public opinion "
            "(15-day notice) and gazette / effective-date management.",
            "Publication of preliminary notification of proposed revised guideline "
            "values (via DIPR / Karnataka Rajya Patra) with a configurable objections "
            "period, and final e-gazette / effective-date management after CVC approval.",
        ),
        (
            "Digitised General Revision cycle with order → notification → SR proposal → "
            "Sub-committee → 15-day public opinion → CVC → effective date → "
            "automatic adoption.",
            "Digitised General Revision cycle with order → notification → SR evidence → "
            "Sub-committee → SR proposal → preliminary notification (DIPR / Karnataka "
            "Rajya Patra; configurable objections period) → CVC → e-gazette → effective "
            "date → automatic adoption.",
        ),
        (
            "Portal messaging for public opinion windows and for individual project "
            "fixation eligibility; FAQs linking Sec. 45-B and 15-day notice.",
            "Portal messaging for public opinion / objections windows (configurable "
            "period) and for individual project fixation eligibility; FAQs linking "
            "Sec. 45-B and preliminary notification.",
        ),
    ]

    for p in doc.paragraphs:
        text = p.text
        if not text.strip():
            continue
        for start, new in replacements.items():
            if text.strip().startswith(start) or start in text:
                # Prefer exact list-item match by startswith after strip
                if text.strip().startswith(start):
                    replace_paragraph_text(p, new)
                    break
        for old, new in cascade_old_parts:
            if text.strip() == old or old in text:
                replace_paragraph_text(p, text.replace(old, new) if old in text else new)
                break


def update_status_model(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        hdr = [c.text.strip() for c in table.rows[0].cells]
        if hdr[:3] != ["Status", "Meaning", "Owner"]:
            continue
        blob = " ".join(c.text for r in table.rows for c in r.cells)
        if "Public Opinion Open" not in blob:
            continue
        for row in table.rows[1:]:
            status = row.cells[0].text.strip()
            if status == "Proposal Draft":
                set_cell_text(row.cells[0], "Evidence Assembled", bold=True)
                set_cell_text(
                    row.cells[1],
                    "SR capturing evidence pack; no formal rate proposal yet",
                )
                set_cell_text(row.cells[2], "SR")
            elif status == "Sub-committee Review":
                set_cell_text(
                    row.cells[1],
                    "First Sub-committee sitting on evidence / discussion draft "
                    "(before formal SR proposal)",
                )
            elif status == "Public Opinion Open":
                set_cell_text(row.cells[0], "Preliminary Notification / Objections Open", bold=True)
                set_cell_text(
                    row.cells[1],
                    "SR proposal lodged; preliminary notification published "
                    "(incl. DIPR / Karnataka Rajya Patra); configurable objections "
                    "period running",
                )
                set_cell_text(row.cells[2], "SR / Public / DIPR")
            elif status == "Objection Disposal":
                set_cell_text(
                    row.cells[1],
                    "Post-objections-period Sub-committee revision",
                )
        # Insert Formal Proposal status after Sub-committee Review if missing
        statuses = [r.cells[0].text.strip() for r in table.rows]
        if "SR Proposal Lodged" not in statuses:
            # find Sub-committee Review row index
            idx = None
            for i, r in enumerate(table.rows):
                if r.cells[0].text.strip() == "Sub-committee Review":
                    idx = i
                    break
            if idx is not None:
                new_row = table.add_row()
                set_cell_text(new_row.cells[0], "SR Proposal Lodged", bold=True)
                set_cell_text(
                    new_row.cells[1],
                    "Formal SR proposal after Sub-committee; ready for preliminary notification",
                )
                set_cell_text(new_row.cells[2], "SR")
                tr = new_row._tr
                table._tbl.remove(tr)
                table.rows[idx]._tr.addnext(tr)
        return


def update_frs(doc: Document) -> None:
    set_fr_text(
        doc,
        "FR-CVC-GR-004",
        "SR shall assemble the Evidence Pack (FR-CVC-003) after referring "
        "registration, KSRSAC GIS, developer publications, RTC, khata and town-planning "
        "inputs. At this stage SR shall not lodge a formal guidance-value proposal; "
        "the pack is prepared for Sub-committee deliberation.",
    )
    set_fr_text(
        doc,
        "FR-CVC-GR-005",
        "System shall prevent SR from taking the matter to Sub-committee unless "
        "mandatory evidence checklist items are attached or explicitly marked N/A "
        "with reason.",
    )
    set_fr_text(
        doc,
        "FR-CVC-GR-006",
        "SR shall take the evidence pack / discussion draft to the Sub-committee "
        "(before formal proposal); system shall capture attendance, discussion points, "
        "directions on rates / methodology, and minutes.",
    )
    set_fr_text(
        doc,
        "FR-CVC-GR-007",
        "After Sub-committee deliberation, SR shall formally propose guidance values "
        "for each territorial unit / rate slab (with justification reflecting "
        "Sub-committee directions). Thereafter, SR shall issue a preliminary "
        "notification of the proposed revised guideline values for public objections / "
        "opinions. The objections / public-opinion period shall be configurable by "
        "authorised CVC / DIGR Admin per Revision Cycle (default may be 15 days but "
        "shall not be hard-coded). The preliminary notification shall also be "
        "published through integration with the Department of Information and Public "
        "Relations (DIPR), Government of Karnataka, and Karnataka Rajya Patra.",
    )
    set_fr_text(
        doc,
        "FR-CVC-GR-009",
        "After the configured objections period closes, Sub-committee shall reconvene; "
        "system shall allow revision only for rate lines that received objections "
        "(or as directed by Sub-committee minutes), with disposal remarks per objection.",
    )
    insert_fr_after(
        doc,
        "FR-CVC-GR-007",
        "FR-CVC-GR-017",
        "System shall integrate with the Department of Information and Public Relations "
        "(DIPR), Government of Karnataka, and Karnataka Rajya Patra to publish the "
        "preliminary notification of proposed guidance rates (post Sub-committee / "
        "SR proposal), and shall capture / link the publication particulars "
        "(number, date, URL/PDF) to the Revision Cycle before the objections window opens.",
        "Must",
    )


def update_rtm_range(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        for row in table.rows[1:]:
            cell0 = row.cells[0].text.strip()
            if "FR-CVC-GR-001" in cell0 and "016" in cell0 and "017" not in cell0:
                set_cell_text(row.cells[0], cell0.replace("016", "017"))
                return


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    update_document_control_version(doc)
    append_version_history_row(doc)
    rewrite_process_steps(doc)
    update_status_model(doc)
    update_frs(doc)
    update_rtm_range(doc)

    try:
        doc.save(str(DST))
    except PermissionError:
        alt = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.4_updated.docx"
        doc.save(str(alt))
        print(f"NOTE: locked; wrote {alt}")
        return

    print(f"Wrote {DST}")
    # Verify key FRs and steps
    for req in ("FR-CVC-GR-004", "FR-CVC-GR-006", "FR-CVC-GR-007", "FR-CVC-GR-009", "FR-CVC-GR-017"):
        for table in doc.tables:
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == req:
                    print(req, "→", row.cells[1].text.strip()[:110])
    print("--- process steps ---")
    capture = False
    for p in doc.paragraphs:
        if "Process steps — General Revision" in p.text:
            capture = True
            continue
        if capture:
            if p.style.name.startswith("Heading"):
                break
            if p.text.strip():
                print("-", p.text.strip()[:130])


if __name__ == "__main__":
    main()
