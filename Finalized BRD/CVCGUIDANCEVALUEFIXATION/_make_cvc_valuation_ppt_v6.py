# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v6.pptx from v5.

Adds Rule 6 — Guidelines for estimation of market value by the Sub-Committee
(also evidence pack factors when proposing new guidance value).
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation_v5.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v6.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)
ACCENT = RGBColor(0x1F, 0x6B, 0x7A)

MARGIN = Inches(0.45)
CONTENT_W = Inches(12.43)
FOOTER_Y = Inches(7.14)


def _run(para, text, *, size, bold=False, italic=False, color=INK, font="Calibri"):
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def add_textbox(
    slide, x, y, w, h, text, *, size, bold=False, italic=False, color=INK,
    font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, cambria=False,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    para.clear()
    _run(
        para, text, size=size, bold=bold, italic=italic, color=color,
        font="Cambria" if cambria else font,
    )
    return box


def add_rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.adjustments[0] = 0.06
    return shape


def chrome(slide, kicker: str, title: str, page: str):
    add_textbox(slide, MARGIN, Inches(0.38), Inches(10.2), Inches(0.26),
                kicker, size=11, bold=True, color=MUTED)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(11.5), Inches(0.48),
                title, size=22, bold=True, color=NAVY, cambria=True, anchor=MSO_ANCHOR.MIDDLE)
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.35), Inches(0.36), Inches(0.52), Inches(0.52))
    badge.fill.solid()
    badge.fill.fore_color.rgb = NAVY
    badge.line.fill.background()
    add_textbox(slide, Inches(12.35), Inches(0.36), Inches(0.52), Inches(0.52),
                page, size=14, bold=True, color=WHITE, cambria=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, MARGIN, FOOTER_Y, Inches(8.5), Inches(0.28),
                "VALUATION MODULE (CVC)  |  BRD WALKTHROUGH  ·  KAVERI 3.0", size=8.5, color=MUTED)
    add_textbox(slide, Inches(12.3), FOOTER_Y, Inches(0.6), Inches(0.28),
                page, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)


def insert_slide_at(prs: Presentation, index: int):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    sld_id_lst = prs.slides._sldIdLst
    sld_ids = list(sld_id_lst)
    last = sld_ids[-1]
    sld_id_lst.remove(last)
    sld_id_lst.insert(index, last)
    return prs.slides[index]


def add_card(slide, x, y, w, h, title: str, bullets: list[str], title_size=12, body_size=9.5):
    add_rect(slide, x, y, w, h, ALT)
    add_textbox(
        slide, x + Inches(0.14), y + Inches(0.10), w - Inches(0.28), Inches(0.28),
        title, size=title_size, bold=True, color=NAVY, cambria=True,
    )
    box = slide.shapes.add_textbox(
        x + Inches(0.14), y + Inches(0.40), w - Inches(0.28), h - Inches(0.50)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(2)
        para.line_spacing = 1.05
        para.clear()
        _run(para, f"•  {line}", size=body_size, color=INK)


def build_rule6_slide(slide, page: str):
    chrome(
        slide,
        "LEGAL & REGULATORY REFERENCE",
        "Guidelines for Estimating Guidance Value",
        page,
    )

    add_textbox(
        slide, MARGIN, Inches(1.12), CONTENT_W, Inches(0.42),
        "Karnataka Stamp (Constitution of C.V.C. etc.) Rules, 2003 — Rule 6.  "
        "Each Market Valuation Sub-Committee prepares average-rate statements for "
        "agricultural / non-agricultural lands and residential, commercial & industrial sites.  "
        "These factors are also the evidence basis when proposing a new guidance value.",
        size=10, italic=True, color=MUTED,
    )

    # Four category cards — Rule 6(1)(a)–(c) + 6(2)
    gap = Inches(0.14)
    card_w = (CONTENT_W - 3 * gap) / 4
    y = Inches(1.62)
    h = Inches(3.55)

    add_card(
        slide, MARGIN, y, card_w, h,
        "Lands — R.6(1)(a)",
        [
            "Dry / garden / wet classification",
            "Soil class in survey records",
            "Other valuation influencers",
            "Value of adjacent / vicinity lands",
            "Crop nature & 5-year avg. yield",
            "Nearness to road & market",
            "Distance from village; location; level",
            "Transport & irrigation (tank / well / pumpset)",
        ],
    )

    add_card(
        slide, MARGIN + card_w + gap, y, card_w, h,
        "House sites — R.6(1)(b)",
        [
            "General locality site values",
            "Proximity to road / rail / bus",
            "Proximity to market & shops",
            "Offices, hospitals, schools",
            "Development / industrial activity",
            "Land tax & local body valuation",
            "Other material features",
            "Special: bore-well, lawn, garden, pool",
        ],
    )

    add_card(
        slide, MARGIN + 2 * (card_w + gap), y, card_w, h,
        "Other properties — R.6(1)(c)",
        [
            "Nature & condition of property",
            "Purpose for which used",
            "Any other special valuation features",
        ],
        body_size=10.5,
    )

    add_card(
        slide, MARGIN + 3 * (card_w + gap), y, card_w, h,
        "Rate suggestions — R.6(2)",
        [
            "Non-agri / industrial: multiple of agri rate (village) or per sq.ft (town/city)",
            "Agricultural: dry / wet / garden; consider proximity to village",
            "Plantations (coconut / areca): treat as garden lands",
            "Buildings: PWD construction norms for the area",
            "Electricity, water, drainage = part of construction cost (not ‘special’)",
        ],
        body_size=9.5,
    )

    # Evidence / process callout
    add_rect(slide, MARGIN, Inches(5.30), CONTENT_W, Inches(1.45), ALT)
    add_textbox(
        slide, MARGIN + Inches(0.20), Inches(5.40), CONTENT_W - Inches(0.40), Inches(0.28),
        "Evidence for proposing new guidance value  ·  BRD linkage",
        size=12, bold=True, color=NAVY, cambria=True,
    )
    add_textbox(
        slide, MARGIN + Inches(0.20), Inches(5.72), CONTENT_W - Inches(0.40), Inches(0.90),
        "Rule 6 factors form the statutory checklist for Sub-Committee rate statements and for "
        "evidence packs attached to General Revision / Individual Project fixation proposals "
        "(BRD 7.ii.d / 7.iii.c).  Registrar verifies Sub-Committee statements; discrepancies "
        "are remitted for rectification (resubmit within 15 days); examined booklets / soft copies "
        "reach the Secretary, CVC in the first week of January of the next calendar year.",
        size=10.5, italic=True, color=ACCENT,
    )


def set_shape_number(shape, number: str, *, size: float | None = None, color=None, font=None):
    for para in shape.text_frame.paragraphs:
        if para.runs:
            para.runs[0].text = number
            if size is not None:
                para.runs[0].font.size = Pt(size)
            if color is not None:
                para.runs[0].font.color.rgb = color
            if font is not None:
                para.runs[0].font.name = font
            for r in para.runs[1:]:
                r.text = ""
        else:
            para.clear()
            _run(para, number, size=size or 14, bold=True, color=color or WHITE, font=font or "Cambria")


def renumber_content_slides(prs: Presentation):
    BADGE_LEFT_MIN = 10_500_000
    FOOTER_TOP_MIN = 6_200_000
    page = 2
    for idx in range(1, len(prs.slides) - 1):
        slide = prs.slides[idx]
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = shape.text.strip()
            if not re.fullmatch(r"\d{1,2}", t):
                continue
            if shape.left >= BADGE_LEFT_MIN and shape.top < 1_200_000:
                set_shape_number(shape, str(page), size=14, color=WHITE, font="Cambria")
            elif shape.left >= BADGE_LEFT_MIN and shape.top >= FOOTER_TOP_MIN:
                set_shape_number(shape, str(page), size=8.5, color=MUTED, font="Calibri")
        page += 1


def bump_version(prs: Presentation, version: str = "6", date: str = "19-09-2026"):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full = para.text
                if not full or "BRD-K3-CVC-GVF-001" not in full:
                    continue
                if "Last updated" in full:
                    new = (
                        f"Document ID: BRD-K3-CVC-GVF-001     ·     Version {version}     ·     "
                        f"Status: Draft     ·     Last updated {date}"
                    )
                else:
                    new = f"Document ID: BRD-K3-CVC-GVF-001  ·  Version {version}  ·  {date}"
                if para.runs:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""


def safe_save(prs: Presentation, path: Path) -> Path:
    try:
        prs.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_updated" + path.suffix)
        prs.save(str(alt))
        print(f"NOTE: {path.name} locked; wrote {alt.name}")
        return alt


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    out = DST
    try:
        shutil.copy2(SRC, out)
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v6_updated.pptx"
        shutil.copy2(SRC, out)

    prs = Presentation(str(out))

    # After Sub-Committees (index 5 in v5) → insert Rule 6 guidelines
    # v5: … 4 CVC, 5 Sub-Committee, 6 Stakeholders
    insert_at = 6
    slide = insert_slide_at(prs, insert_at)
    build_rule6_slide(slide, page="7")
    renumber_content_slides(prs)
    bump_version(prs, version="6", date="19-09-2026")

    out = safe_save(prs, out)
    print(f"Wrote {out}")
    print(f"Slides: {len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        texts = [
            sh.text.strip().split("\n")[0][:70]
            for sh in s.shapes
            if hasattr(sh, "text") and sh.text.strip()
        ]
        print(f"  {i}. {texts[1] if len(texts) > 1 else (texts[0] if texts else '')}")


if __name__ == "__main__":
    main()
