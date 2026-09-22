# -*- coding: utf-8 -*-
"""Create BRD_CVC_Guidance_Value_Fixation_v1.3.docx from v1.2.

Adds requirement: after approval of proposed rates, system shall integrate with
Department of Information and Public Relations (GoK) and Karnataka Rajya Patra
to publish the e-gazette.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\CVCGUIDANCEVALUEFIXATION")
SRC = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.2.docx"
DST = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.3.docx"
OUT_VERSION = "1.3"
OUT_DATE = "22-09-2026"

NEW_FR_ID = "FR-CVC-GR-016"
NEW_FR_TEXT = (
    "After approval of the proposed / revised guidance rates, the system shall "
    "integrate with the Department of Information and Public Relations (DIPR), "
    "Government of Karnataka, and Karnataka Rajya Patra to publish the e-gazette "
    "(Official Gazette notification) for the approved rates, and shall capture / "
    "link the returned e-gazette particulars (number, date, URL/PDF) to the "
    "approved rate set before Kaveri adoption from the notified effective date."
)
NEW_FR_PRIORITY = "Must"

INTEGRATION_SYSTEM = (
    "Department of Information and Public Relations (DIPR), Government of Karnataka "
    "/ Karnataka Rajya Patra (e-gazette)"
)
INTEGRATION_PURPOSE = (
    "Publish approved guidance-value notifications in the e-gazette after CVC / IGR "
    "approval; return gazette particulars for linkage to rate versions"
)


def shade_cell(cell, hex_fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def update_document_control_version(doc: Document) -> None:
    for table in doc.tables:
        if len(table.columns) != 2:
            continue
        keys = [r.cells[0].text.strip() for r in table.rows]
        if "Document ID" not in keys:
            continue
        for row in table.rows:
            key = row.cells[0].text.strip()
            if key == "Version":
                set_cell_text(row.cells[1], OUT_VERSION)
            if key == "Last updated":
                set_cell_text(row.cells[1], OUT_DATE)


def append_version_history_row(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [c.text.strip() for c in table.rows[0].cells]
        if headers[:3] != ["Version", "Date", "Author"]:
            continue
        # Avoid duplicate 1.3 if re-run
        for r in table.rows[1:]:
            if r.cells[0].text.strip() == OUT_VERSION:
                return
        row = table.add_row()
        values = [
            OUT_VERSION,
            OUT_DATE,
            "Nandha Kumar",
            "Added FR-CVC-GR-016: post-approval e-gazette publish integration with "
            "DIPR, Government of Karnataka and Karnataka Rajya Patra; Integrations "
            "table updated",
            "Prashanth",
        ]
        for i, val in enumerate(values):
            if i < len(row.cells):
                set_cell_text(row.cells[i], val)
        return
    raise SystemExit("Version history table not found")


def find_fr_table(doc: Document, containing_id: str):
    for table in doc.tables:
        if not table.rows:
            continue
        if table.rows[0].cells[0].text.strip() != "Req ID":
            continue
        ids = [r.cells[0].text.strip() for r in table.rows[1:]]
        if containing_id in ids:
            return table, ids
    return None, []


def insert_fr_after(table, after_id: str, req_id: str, text: str, priority: str) -> None:
    # If already present, skip
    for r in table.rows[1:]:
        if r.cells[0].text.strip() == req_id:
            set_cell_text(r.cells[1], text)
            set_cell_text(r.cells[2], priority)
            return

    # Find index of after_id
    after_idx = None
    for i, r in enumerate(table.rows):
        if r.cells[0].text.strip() == after_id:
            after_idx = i
            break
    if after_idx is None:
        raise SystemExit(f"FR {after_id} not found for insert")

    # Add row at end then move XML after after_idx
    new_row = table.add_row()
    set_cell_text(new_row.cells[0], req_id, bold=True)
    set_cell_text(new_row.cells[1], text)
    set_cell_text(new_row.cells[2], priority)

    tbl = table._tbl
    tr = new_row._tr
    tbl.remove(tr)
    # insert after the target row element
    target_tr = table.rows[after_idx]._tr
    target_tr.addnext(tr)


def add_integration_row(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        hdr = [c.text.strip() for c in table.rows[0].cells]
        if hdr[:2] != ["System", "Purpose"]:
            continue
        for r in table.rows[1:]:
            if "Rajya Patra" in r.cells[0].text or "DIPR" in r.cells[0].text:
                set_cell_text(r.cells[0], INTEGRATION_SYSTEM)
                set_cell_text(r.cells[1], INTEGRATION_PURPOSE)
                return
        row = table.add_row()
        set_cell_text(row.cells[0], INTEGRATION_SYSTEM)
        set_cell_text(row.cells[1], INTEGRATION_PURPOSE)
        return
    raise SystemExit("Integrations table not found")


def enrich_process_step(doc: Document) -> None:
    """Clarify gazette publish step in General Revision numbered list if present."""
    needle = (
        "New guidance rates are adopted in KAVERI from the effective date as in the gazette"
    )
    for p in doc.paragraphs:
        if needle in p.text and "Rajya Patra" not in p.text:
            # Keep original; add a preceding sibling sentence via rewriting
            old = p.text
            # If this is the adoption step, insert e-gazette publish ahead in same para
            p.clear()
            run = p.add_run(
                "On CVC approval of proposed rates, the system shall submit / publish "
                "the e-gazette notification via integration with the Department of "
                "Information and Public Relations (DIPR), Government of Karnataka, and "
                "Karnataka Rajya Patra; gazette particulars are linked to the approved "
                "rate set. "
                + old
            )
            return


def update_rtm_if_present(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        hdr = [c.text.strip() for c in table.rows[0].cells]
        if not hdr or "Req ID" not in hdr[0]:
            continue
        # RTM often has Legal ref columns
        if len(hdr) < 4:
            continue
        for r in table.rows[1:]:
            if "FR-CVC-GR-001" in r.cells[0].text or "FR-CVC-GR-001…015" in r.cells[0].text:
                # Expand range to include 016
                cell0 = r.cells[0].text.strip()
                if "016" not in cell0:
                    set_cell_text(
                        r.cells[0],
                        cell0.replace("015", "016").replace("…015", "…016")
                        if "015" in cell0
                        else "FR-CVC-GR-001…016",
                    )
                return


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    update_document_control_version(doc)
    append_version_history_row(doc)

    table, ids = find_fr_table(doc, "FR-CVC-GR-012")
    if table is None:
        raise SystemExit("CVC approval FR table not found")
    insert_fr_after(table, "FR-CVC-GR-012", NEW_FR_ID, NEW_FR_TEXT, NEW_FR_PRIORITY)

    add_integration_row(doc)
    enrich_process_step(doc)
    update_rtm_if_present(doc)

    try:
        doc.save(str(DST))
    except PermissionError:
        alt = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.3_updated.docx"
        doc.save(str(alt))
        print(f"NOTE: locked; wrote {alt}")
        return

    print(f"Wrote {DST}")
    table, ids = find_fr_table(doc, NEW_FR_ID)
    print("FR IDs in approval table:", ids)
    for r in table.rows:
        if r.cells[0].text.strip() in (NEW_FR_ID, "FR-CVC-GR-012", "FR-CVC-GR-013"):
            print(r.cells[0].text.strip(), "→", r.cells[1].text.strip()[:100])


if __name__ == "__main__":
    main()
