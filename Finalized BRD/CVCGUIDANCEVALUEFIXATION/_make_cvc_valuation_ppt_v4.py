# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v4.pptx from v3.

Adds one slide: Central Valuation Committee composition from
Karnataka Stamp (Constitution of C.V.C. etc.) Rules, 2003 — Rule 3
(source: user-provided gazette / KLJ extract photographs).
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation_v3.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v4.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)
ACCENT = RGBColor(0x1F, 0x6B, 0x7A)
GOLD = RGBColor(0xC8, 0x9B, 0x3C)

MARGIN = Inches(0.5)
CONTENT_W = Inches(12.333)
FOOTER_Y = Inches(7.14)

# Rule 3 — representatives (one from each), plus Chairman & Member Secretary.
MEMBERS_LEFT = [
    "(i) Directorate of Town Planning",
    "(ii) Directorate of Survey and Settlement",
    "(iii) Bangalore City Corporation",
    "(iv) Bangalore Development Authority",
    "(v) Income-tax Department",
    "(vi) Karnataka Public Works Department",
]
MEMBERS_RIGHT = [
    "(vii) Karnataka Irrigation Department",
    "(viii) Registration and Stamps Department",
    "(ix) Institute of Chartered Valuers",
    "(x) Federation of Karnataka Chamber of Commerce and Industries",
    "(xi) Any other person having expertise in the subject",
]


def _run(para, text, *, size, bold=False, italic=False, color=INK, font="Calibri"):
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size,
    bold=False,
    italic=False,
    color=INK,
    font="Calibri",
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    cambria=False,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    para.clear()
    _run(
        para,
        text,
        size=size,
        bold=bold,
        italic=italic,
        color=color,
        font="Cambria" if cambria else font,
    )
    return box


def add_rect(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.adjustments[0] = 0.08
    return shape


def chrome(slide, kicker: str, title: str, page: str):
    add_textbox(
        slide, MARGIN, Inches(0.42), Inches(10.0), Inches(0.28),
        kicker, size=11.5, bold=True, color=MUTED,
    )
    add_textbox(
        slide, MARGIN, Inches(0.70), Inches(11.5), Inches(0.55),
        title, size=26, bold=True, color=NAVY, cambria=True, anchor=MSO_ANCHOR.MIDDLE,
    )
    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(12.3), Inches(0.40), Inches(0.55), Inches(0.55)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = NAVY
    badge.line.fill.background()
    add_textbox(
        slide, Inches(12.3), Inches(0.40), Inches(0.55), Inches(0.55),
        page, size=15, bold=True, color=WHITE, cambria=True,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_textbox(
        slide, MARGIN, FOOTER_Y, Inches(8.5), Inches(0.28),
        "VALUATION MODULE (CVC)  |  BRD WALKTHROUGH  ·  KAVERI 3.0",
        size=8.5, color=MUTED,
    )
    add_textbox(
        slide, Inches(12.25), FOOTER_Y, Inches(0.6), Inches(0.28),
        page, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT,
    )


def insert_slide_at(prs: Presentation, index: int):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    sld_id_lst = prs.slides._sldIdLst
    sld_ids = list(sld_id_lst)
    last = sld_ids[-1]
    sld_id_lst.remove(last)
    sld_id_lst.insert(index, last)
    return prs.slides[index]


def add_role_card(slide, x, y, w, h, label, value):
    add_rect(slide, x, y, w, h, ALT)
    add_textbox(
        slide, x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.28),
        label, size=10, bold=True, color=ACCENT, italic=True,
    )
    add_textbox(
        slide, x + Inches(0.18), y + Inches(0.40), w - Inches(0.36), Inches(0.55),
        value, size=14, bold=True, color=NAVY, cambria=True,
    )


def add_member_column(slide, x, y, w, h, items):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(36000)
    tf.margin_right = Emu(18000)
    tf.margin_top = Emu(18000)
    tf.margin_bottom = Emu(18000)
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        para.space_after = Pt(6)
        para.line_spacing = 1.12
        para.clear()
        _run(para, item, size=12, bold=False, color=INK)


def build_cvc_members_slide(slide, page: str):
    chrome(
        slide,
        "LEGAL & REGULATORY REFERENCE",
        "Central Valuation Committee — Composition",
        page,
    )

    add_textbox(
        slide,
        MARGIN,
        Inches(1.32),
        CONTENT_W,
        Inches(0.35),
        "Karnataka Stamp (Constitution of Central Valuation Committee for Estimation, "
        "Publication and Revision of Market Value Guidelines of Properties) Rules, 2003 — Rule 3",
        size=10,
        italic=True,
        color=MUTED,
    )

    # Chairman / Member Secretary / Cap
    card_w = Inches(3.95)
    card_h = Inches(1.05)
    gap = Inches(0.18)
    y_cards = Inches(1.72)
    add_role_card(
        slide, MARGIN, y_cards, card_w, card_h,
        "CHAIRMAN",
        "Commissioner of Stamps",
    )
    add_role_card(
        slide, MARGIN + card_w + gap, y_cards, card_w, card_h,
        "MEMBER SECRETARY",
        "DIGR (Intelligence)",
    )
    add_role_card(
        slide, MARGIN + 2 * (card_w + gap), y_cards, card_w, card_h,
        "STRENGTH",
        "Total members ≤ 20",
    )

    # Members panel
    panel_y = Inches(2.95)
    panel_h = Inches(3.45)
    add_rect(slide, MARGIN, panel_y, CONTENT_W, panel_h, ALT)

    add_textbox(
        slide,
        MARGIN + Inches(0.22),
        panel_y + Inches(0.14),
        CONTENT_W - Inches(0.44),
        Inches(0.35),
        "Members — one representative from each of the following",
        size=13,
        bold=True,
        color=NAVY,
        cambria=True,
    )

    col_w = Inches(5.85)
    col_y = panel_y + Inches(0.55)
    col_h = Inches(2.35)
    add_member_column(slide, MARGIN + Inches(0.15), col_y, col_w, col_h, MEMBERS_LEFT)
    add_member_column(
        slide, MARGIN + col_w + Inches(0.25), col_y, col_w, col_h, MEMBERS_RIGHT
    )

    add_textbox(
        slide,
        MARGIN + Inches(0.22),
        panel_y + Inches(2.95),
        CONTENT_W - Inches(0.44),
        Inches(0.40),
        "Non-official members: term of two years, subject to the pleasure of the Government.  "
        "Member Secretary (DIGR–Intelligence) handles day-to-day administration, correspondence, "
        "and compilation / publication of market value guidelines.",
        size=10,
        italic=True,
        color=ACCENT,
    )


def set_shape_number(shape, number: str, *, size: float | None = None, color=None, font=None):
    for para in shape.text_frame.paragraphs:
        if para.runs:
            para.runs[0].text = number
            if size is not None:
                para.runs[0].font.size = Pt(size)
            if color is not None:
                para.runs[0].font.color.rgb = color
            if font is not None:
                para.runs[0].font.name = font
            for r in para.runs[1:]:
                r.text = ""
        else:
            para.clear()
            _run(
                para,
                number,
                size=size or 15,
                bold=True,
                color=color or WHITE,
                font=font or "Cambria",
            )


def renumber_content_slides(prs: Presentation):
    """Update only chrome page badge (top-right) and footer page (bottom-right).

    Do NOT touch other numeric shapes (process step chips, ranked lists, etc.).
    """
    # Empirically from this deck's chrome:
    BADGE_LEFT_MIN = 10_500_000  # ~11.5"+
    FOOTER_TOP_MIN = 6_200_000   # ~6.8"+

    page = 2
    for idx in range(1, len(prs.slides) - 1):
        slide = prs.slides[idx]
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            t = shape.text.strip()
            if not re.fullmatch(r"\d{1,2}", t):
                continue
            # Top-right circular badge
            if shape.left >= BADGE_LEFT_MIN and shape.top < 1_200_000:
                set_shape_number(shape, str(page), size=15, color=WHITE, font="Cambria")
            # Bottom-right footer page
            elif shape.left >= BADGE_LEFT_MIN and shape.top >= FOOTER_TOP_MIN:
                set_shape_number(shape, str(page), size=8.5, color=MUTED, font="Calibri")
        page += 1


def bump_version(prs: Presentation, version: str = "4", date: str = "19-09-2026"):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full = para.text
                if not full or "BRD-K3-CVC-GVF-001" not in full:
                    continue
                if "Last updated" in full:
                    new = (
                        f"Document ID: BRD-K3-CVC-GVF-001     ·     Version {version}     ·     "
                        f"Status: Draft     ·     Last updated {date}"
                    )
                else:
                    new = f"Document ID: BRD-K3-CVC-GVF-001  ·  Version {version}  ·  {date}"
                if para.runs:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""


def safe_save(prs: Presentation, path: Path) -> Path:
    try:
        prs.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_updated" + path.suffix)
        prs.save(str(alt))
        print(f"NOTE: {path.name} locked; wrote {alt.name}")
        return alt


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    out = DST
    try:
        shutil.copy2(SRC, out)
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v4_updated.pptx"
        shutil.copy2(SRC, out)

    prs = Presentation(str(out))

    # Insert after Definitions (index 4) — before Stakeholders — as legal Rule 3 reference.
    # Current: 0 Title, 1 Acts, 2 Sections, 3 Definitions, 4 Stakeholders...
    insert_at = 4
    slide = insert_slide_at(prs, insert_at)
    # Temporary page; renumber fixes all
    build_cvc_members_slide(slide, page="5")
    renumber_content_slides(prs)
    bump_version(prs, version="4", date="19-09-2026")

    out = safe_save(prs, out)
    print(f"Wrote {out}")
    print(f"Slides: {len(prs.slides)}")
    for i, s in enumerate(prs.slides, 1):
        texts = [
            sh.text.strip().split("\n")[0][:70]
            for sh in s.shapes
            if hasattr(sh, "text") and sh.text.strip()
        ]
        print(f"  {i}. {texts[1] if len(texts) > 1 else (texts[0] if texts else '')}")


if __name__ == "__main__":
    main()
