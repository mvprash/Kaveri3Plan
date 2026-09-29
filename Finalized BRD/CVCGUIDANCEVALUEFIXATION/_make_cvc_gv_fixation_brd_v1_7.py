# -*- coding: utf-8 -*-
"""Create BRD_CVC_Guidance_Value_Fixation_v1.7.docx from v1.6.

General Revision (Process A) restructured to the consolidated 15-stage workflow
of Rules 5 and 7 (Stages/Guideline_Value_Revision_Process_15_Steps.docx):
  CVC instructions -> Sub-Committee notice of intention -> 15-day objections ->
  Secretary processes objections -> Sub-Committee decides rates -> signed
  statement (booklet + soft copy) -> DR verification / 15-day rectification ->
  DR views, per-sub-district booklets -> CVC Secretary verification -> CVC
  decision -> attestation -> DR -> Sub-Committees -> publication and sale.
Effective date, DIPR / Karnataka Rajya Patra e-gazette and Kaveri adoption are
retained as system extension steps after stage 15.
Individual Project Fixation (Process B) is unchanged.
"""
from __future__ import annotations

import copy
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\CVCGUIDANCEVALUEFIXATION")
SRC = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.6.docx"
DST = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.7.docx"
OUT_VERSION = "1.7"
OUT_DATE = "29-09-2026"

TBC_END = "[period to be confirmed against Rule 7(2) text]"
TBC_FORMAT = "[soft-copy format to be confirmed against Rule 5(2) text]"

VERSION_SUMMARY = (
    "General Revision (Process A) restructured to the 15-stage statutory workflow of "
    "Rules 5 and 7 (Stages/Guideline_Value_Revision_Process_15_Steps): CVC instructions "
    "to Sub-Committees; notice of intention and 15-day objections before rates are "
    "decided; signed statement of average rates (booklet + soft copy); DR verification "
    "with 15-day rectification; CVC Secretary verification; CVC decision on "
    "Sub-Committee / DR suggestions; attestation, distribution, publication and sale. "
    "Effective date, DIPR / Karnataka Rajya Patra e-gazette and Kaveri adoption retained "
    "as system extension steps. FR-CVC-GR-001…022 re-baselined; status model, actors, "
    "glossary, business rules, data entities, notifications, MIS, UI updated. "
    "Process B unchanged."
)

# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

GR_STEPS = [
    "[Stage 1 — Rule 5(1)] CVC issues instructions and general policy guidelines to all "
    "Market Valuation Sub-Committees for estimation of market value guidelines for the next "
    "calendar year (system: CVC Secretary creates the Revision Cycle and attaches the "
    "instructions; all Sub-Committees, DRs and DIGR (Valuation) are notified). CVC may also "
    "instruct any Sub-Committee to revise rates at any time during the calendar year "
    "(mid-year Revision Cycle scoped to the selected Sub-Committees).",
    "[Stage 2 — Rule 5(2)] On receipt of instructions, the Sub-Committee publishes its "
    "intention to estimate or revise market values in local newspapers and on the notice "
    "boards of important offices (system: bilingual notice generated; publication proof "
    "recorded; notice also shown on the citizen portal).",
    "[Stage 3 — Rule 5(2)] The Sub-Committee allows 15 days for receipt of objections and "
    "suggestions from the public (online via citizen portal or offline at the SRO, "
    "digitised the same day).",
    "[Stage 4 — Rule 5(2)] The Secretary of the Sub-Committee (SR) processes the objections "
    "and suggestions and places them before the Sub-Committee for discussion.",
    "[Stage 5 — Rule 5(2)] The Sub-Committee meets as often as required, considers the "
    "public inputs (with the evidence pack and Rule 6 factors) and decides the estimated "
    "market value rates for the guidelines.",
    "[Stage 6 — Rule 5(2)] The Sub-Committee prepares a statement showing average rates of "
    "agricultural land and residential, commercial and industrial sites, with data arranged "
    "village-wise and local-body-wise.",
    "[Stage 7 — Rule 5(2)] The statement is prepared in the form specified by CVC, signed by "
    "the Sub-Committee Secretary and Chairman, and may record the Sub-Committee’s views on "
    "objections and suggestions.",
    "[Stage 8 — Rule 5(2)] The hard copy (booklet form) and the soft copy "
    f"{TBC_FORMAT} are sent to the concerned District Registrar.",
    "[Stage 9 — Rule 5(3)] The District Registrar verifies the statement. If any "
    "discrepancy or omission is found, it is immediately returned to the Sub-Committee for "
    "rectification or completion.",
    "[Stage 10 — Rule 5(3)] The Sub-Committee rectifies the discrepancy or supplies the "
    "omitted information and re-submits the statement within 15 days from the date of "
    "reference.",
    "[Stage 11 — Rule 5(3)] The District Registrar finally examines the data, records views "
    "regarding any improvement or change, and sends separate booklets and soft copies for "
    "each sub-district to the CVC Secretary.",
    "[Stage 12 — Rule 7(1)] The CVC Secretary verifies the statements received from the "
    "District Registrars and places them before the Central Valuation Committee.",
    "[Stage 13 — Rule 7(2)] The Committee discusses the district-wise estimations and "
    "considers suggestions from the Sub-Committees and District Registrars. As far as "
    f"possible before the end {TBC_END}, it takes the final decision, accepts or rejects "
    "suggestions, and records all decisions in the meeting proceedings.",
    "[Stage 14 — Rule 7(3)] Following approval for each sub-district or district, the CVC "
    "Secretary attests the statements and sends them to the concerned District Registrar. "
    "The District Registrar forwards them to the concerned Sub-Committees.",
    "[Stage 15 — Rule 7(3)] The approved statements are published in prominent offices in "
    "each sub-district, including the Sub-Registry Office; made available at the concerned "
    "District Registrar’s office; and printed in sufficient numbers for sale to the public "
    "at a price fixed by CVC.",
    "[System extension — not in Rule 5 / 7 text] CVC fixes the effective date of the "
    "approved rates. The system submits / publishes the e-gazette notification via "
    "integration with the Department of Information and Public Relations (DIPR), "
    "Government of Karnataka, and Karnataka Rajya Patra; gazette particulars are linked to "
    "the approved rate set.",
    "[System extension — not in Rule 5 / 7 text] New guidance rates are adopted in KAVERI "
    "from the effective date as in the gazette (automated or controlled publish of "
    "versioned masters to SR and Citizen logins); prior versions remain queryable.",
]

GR_STATUS_ROWS = [
    ("Cycle Draft", "Revision Cycle created; CVC instructions being prepared", "CVC Secretary"),
    ("Instructions Issued", "CVC instructions and general policy guidelines sent to Sub-Committees (Rule 5(1)); general or mid-year cycle", "CVC"),
    ("Notice of Intention Published", "Sub-Committee intention to estimate / revise published in newspapers and notice boards (Rule 5(2))", "Sub-Committee"),
    ("Objections Open", "15-day window for public objections / suggestions running", "Public / Sub-Committee Secretary"),
    ("Objections Under Processing", "Secretary processing objections for Sub-Committee agenda", "Sub-Committee Secretary (SR)"),
    ("Sub-Committee Deliberation", "Sub-Committee sittings; rates being decided", "Sub-Committee"),
    ("Statement Signed", "Statement of average rates in CVC format signed by Secretary and Chairman", "Sub-Committee"),
    ("Submitted to DR", "Booklet and soft copy submitted to District Registrar", "Sub-Committee / DR"),
    ("Returned for Rectification", "DR returned discrepancy / omission; 15-day rectification clock running (Rule 5(3))", "Sub-Committee"),
    ("DR Final Examination", "DR examining data and recording views on improvement / change", "DR"),
    ("With CVC Secretary", "Per-sub-district booklets / soft copies received; under verification (Rule 7(1))", "CVC Secretary"),
    ("Placed before CVC", "District-wise estimations on CVC agenda (Rule 7(2))", "CVC"),
    ("Approved", "CVC final decision recorded in proceedings; suggestions accepted / rejected", "CVC"),
    ("Attested and Transmitted", "CVC Secretary attested statements; sent to DR and forwarded to Sub-Committees (Rule 7(3))", "CVC Secretary / DR"),
    ("Published", "Approved statements displayed in sub-district offices / SRO, available at DR office, printed for sale (Rule 7(3))", "Sub-Committee / DR"),
    ("Gazette Linked", "Effective date fixed; e-gazette particulars captured (system extension)", "CVC / Admin"),
    ("Active in Kaveri", "Rates live from effective date", "System"),
    ("Returned / Clarification", "Returned to prior stage with remarks (e.g. CVC Secretary seeks clarification from DR)", "Prior owner"),
]

GR_FR_A = [
    ("FR-CVC-GR-001", "CVC (Chairman IGR, through the CVC Secretary) shall create a General Revision Cycle for the next calendar year and issue instructions and general policy guidelines to all Market Valuation Sub-Committees (Rule 5(1)). The instruction document shall be attached to the cycle and all Sub-Committees, DRs and DIGR (Valuation) notified.", "Must"),
    ("FR-CVC-GR-002", "CVC shall be able to instruct one or more specific Sub-Committees to revise rates at any time during the calendar year (Rule 5(1)); system shall create a mid-year Revision Cycle scoped to the selected Sub-Committees / jurisdictions.", "Must"),
    ("FR-CVC-GR-003", "On receipt of instructions, the Sub-Committee (through its Secretary, the SR) shall publish its intention to estimate or revise market values. System shall generate a bilingual notice of intention; record publication in local newspapers (paper, edition, date, clipping) and on notice boards of important offices (office, date, photo); and display the notice on the citizen portal (Rule 5(2)).", "Must"),
    ("FR-CVC-GR-004", "System shall compute the objections window as 15 days from the date of publication of the notice of intention (Rule 5(2)). The period shall be stored per cycle as a parameter (default 15 days per rule) and may be changed only by authorised CVC Admin with a recorded reason (e.g. rule amendment).", "Must"),
    ("FR-CVC-GR-005", "Public shall file objections / suggestions online (citizen portal) or offline at the SRO (digitised the same day) within the window, linked to the Sub-Committee jurisdiction and optionally to village / local body / property class. System shall issue an acknowledgement and flag late submissions.", "Must"),
    ("FR-CVC-GR-006", "The Secretary of the Sub-Committee shall process objections and suggestions (classify, consolidate, add remarks) and place them before the Sub-Committee as a discussion agenda (Rule 5(2)).", "Must"),
]

GR_FR_B = [
    ("FR-CVC-GR-007", "System shall support Sub-Committee sittings as often as required, capturing attendance, agenda, objections / suggestions considered, decisions and minutes. The Evidence Pack (FR-CVC-003) and Rule 6 factors shall be available on the Sub-Committee workbench.", "Must"),
    ("FR-CVC-GR-008", "The Sub-Committee shall decide the estimated market value rates. System shall prepare the statement of average rates of agricultural land and residential, commercial and industrial sites, arranged village-wise and local-body-wise, using the rate category masters (FR-CVC-010…012).", "Must"),
    ("FR-CVC-GR-009", "The statement shall be generated in the form specified by CVC (configurable template maintained by CVC Admin), shall record the Sub-Committee’s views on each objection / suggestion, and shall be signed by the Sub-Committee Secretary and Chairman (eSign / DSC).", "Must"),
    ("FR-CVC-GR-010", f"System shall generate the hard copy in booklet form (print-ready PDF) and the soft copy {TBC_FORMAT}, and submit both to the concerned District Registrar; dispatch of the physical booklet shall be recorded.", "Must"),
    ("FR-CVC-GR-011", "The District Registrar shall verify the statement and, if any discrepancy or omission is found, immediately return it to the Sub-Committee with a discrepancy list for rectification or completion (Rule 5(3)).", "Must"),
    ("FR-CVC-GR-012", "The Sub-Committee shall rectify the discrepancy / supply the omitted information and re-submit within 15 days from the date of reference (Rule 5(3)). System shall track the due date, send reminders and escalate breaches to the DR and DIGR (Valuation).", "Must"),
    ("FR-CVC-GR-013", "The District Registrar shall finally examine the data, record views regarding any improvement or change (per statement / rate line), and send separate booklets and soft copies for each sub-district to the CVC Secretary (Rule 5(3)).", "Must"),
]

GR_FR_C = [
    ("FR-CVC-GR-014", "The CVC Secretary shall verify the statements received from the District Registrars, may seek clarification from the DR, and shall place them before the Central Valuation Committee as a district-wise agenda (Rule 7(1)).", "Must"),
    ("FR-CVC-GR-015", f"CVC shall discuss the district-wise estimations on a committee workbench, consider suggestions from Sub-Committees and District Registrars, record acceptance / rejection of each suggestion, and record all decisions in the meeting proceedings (Rule 7(2)). A target decision date shall be configurable per cycle (as far as possible before the end {TBC_END}) with ageing alerts.", "Must"),
    ("FR-CVC-GR-016", "Following approval for each sub-district or district, the CVC Secretary shall attest the approved statements (digital signature) and send them to the concerned District Registrar, who shall forward them to the concerned Sub-Committees (Rule 7(3)). System shall record transmission and acknowledgement at each hop.", "Must"),
    ("FR-CVC-GR-017", "System shall record publication of approved statements in prominent offices in each sub-district, including the Sub-Registry Office (office, date, proof), their availability at the DR office, and make them viewable / downloadable on the citizen portal (Rule 7(3)).", "Must"),
    ("FR-CVC-GR-018", "System shall capture the sale price fixed by CVC and the number of copies printed per sub-district, and record sale of printed statements to the public (counter or payment gateway) (Rule 7(3)).", "Should"),
    ("FR-CVC-GR-019", "[System extension] On approval, CVC shall fix the effective date for the revised rates.", "Must"),
    ("FR-CVC-GR-020", "[System extension] After approval, the system shall integrate with the Department of Information and Public Relations (DIPR), Government of Karnataka, and Karnataka Rajya Patra to publish the e-gazette notification for the approved rates, and shall capture / link the returned gazette particulars (number, date, URL/PDF) to the approved rate set before Kaveri adoption.", "Must"),
    ("FR-CVC-GR-021", "From the effective date as in the gazette, new guidance rates shall be adopted automatically in Kaveri rate masters for SR and Citizen logins; prior versions remain queryable for audit / historical instruments.", "Must"),
    ("FR-CVC-GR-022", "System shall notify DIGR, DR, SR, Sub-Committees and configured stakeholders when rates become Active in Kaveri.", "Should"),
]

FR_HEADINGS = {
    "General Revision — trigger and cascade":
        "General Revision — CVC instructions, notice of intention and objections (Rule 5(1)–(2))",
    "General Revision — Sub-committee and public opinion":
        "General Revision — Sub-Committee statement and District Registrar verification (Rule 5(2)–(3))",
    "General Revision — CVC approval and Kaveri adoption":
        "General Revision — CVC decision, attestation, publication and Kaveri adoption (Rule 7)",
}

ACTORS = {
    "IGR / Chairman CVC": "Chairs CVC; CVC issues instructions and general policy guidelines to Sub-Committees to trigger General Revision or mid-year revision (Rule 5(1)); approves final guidance value for individual projects",
    "DIGR (CVC)": "Member Secretary of CVC (Rule 3) acting as CVC Secretary in General Revision: circulates CVC instructions; verifies DR statements and places them before CVC (Rule 7(1)); attests approved statements and sends them to DRs (Rule 7(3)); domain reviewer for this BRD",
    "District Registrar (DR) (Superintendent)": "General Revision: verifies Sub-Committee statements; returns discrepancies / omissions for rectification within 15 days; finally examines data, records views and sends separate booklets / soft copies per sub-district to CVC Secretary (Rule 5(3)); forwards attested statements to Sub-Committees and keeps approved statements available at DR office (Rule 7(3)). Individual Project: accepts citizen request under Sec. 45-B; conducts spot inspection and proposes value",
    "Sub-Registrar (SR) (Member Secretary)": "Secretary of the Sub-Committee (Rule 4): publishes notice of intention on behalf of the Sub-Committee; processes objections / suggestions and places them before the Sub-Committee; prepares and co-signs the statement of average rates; submits booklet and soft copy to DR; carries out rectification within 15 days",
    "Market Valuation Sub-committee": "Headed by Tahsildar (Chairman) with SR as Member Secretary; members from Revenue, Survey and Settlement, PWD, local bodies (Rule 4). Publishes intention to estimate / revise; considers public objections; decides estimated market value rates; statement signed by Chairman and Secretary; publishes approved statements in sub-district offices",
    "Secretary CVC": "CVC Secretary (DIGR Valuation). General Revision: see DIGR (CVC). Individual Project: reviews proposals using evidence pack; finalises guidance value for IGR approval",
    "Central Valuation Committee (CVC)": "Issues instructions and policy guidelines to Sub-Committees (Rule 5(1)); discusses district-wise estimations, considers Sub-Committee and DR suggestions, accepts / rejects them and records proceedings (Rule 7(2)); fixes price of printed statements (Rule 7(3)); fixes effective date (system extension)",
    "Citizen / Party / Developer": "Applies for fixation of guidance value for newly developed property; files objections / suggestions within 15 days of the Sub-Committee’s notice of intention (General Revision); may purchase printed approved statements",
    "Public": "May file objections / suggestions within 15 days of the Sub-Committee’s notice of intention; may view approved statements at sub-district offices, DR office and citizen portal",
}

GLOSSARY_UPDATE = {
    "General Revision": "Cycle to estimate / revise market value guidelines for the next calendar year, triggered by CVC instructions and general policy guidelines to Sub-Committees (Rule 5(1)); includes mid-year revision ordered by CVC for specific Sub-Committees.",
    "Sub-committee": "District / sub-district Market Valuation Sub-Committee constituted under Sec. 45-B / Rule 4, headed by the Tahsildar with the SR as Member Secretary.",
    "Public opinion period": "Fifteen (15) days allowed for public objections / suggestions from publication of the Sub-Committee’s notice of intention to estimate / revise market values (Rule 5(2)).",
}
GLOSSARY_ADD = [
    ("Notice of intention", "Sub-Committee notice, published in local newspapers and notice boards of important offices, announcing its intention to estimate or revise market values (Rule 5(2))."),
    ("Statement of average rates", "Sub-Committee statement of average rates of agricultural land and residential, commercial and industrial sites, arranged village-wise and local-body-wise, in CVC-specified form, signed by Secretary and Chairman; sent as booklet and soft copy (Rule 5(2))."),
    ("Rectification period", "Fifteen (15) days from the date of reference by the District Registrar within which the Sub-Committee must rectify discrepancies / omissions and re-submit (Rule 5(3))."),
    ("Attested statement", "CVC-approved statement attested by the CVC Secretary and sent to the DR and Sub-Committees for publication (Rule 7(3))."),
]

PARA_REPLACEMENTS = [
    # (startswith, new text)
    ("Based on the requirement discussion held between",
     "Based on the requirement discussion held between 31-08-2026 and 01-09-2026 "
     "(Document_Registration_requirement_01092026_v2), the workshop process walkthrough and "
     "the consolidated 15-stage workflow of Rules 5 and 7 (CVC Rules, 2003), two fixation "
     "paths are in scope: (1) General Revision — estimation / revision of market value "
     "guidelines triggered by CVC instructions to the Market Valuation Sub-Committees; and "
     "(2) Fixation of guidance value for a new individual project — citizen / party "
     "initiated request accepted by the District Registrar under Sec. 45-B."),
    ("The proposed To-Be workflows digitise committee routing",
     "The proposed To-Be workflows digitise the statutory routing (CVC → Sub-Committee "
     "notice of intention → 15-day public objections → Sub-Committee statement → District "
     "Registrar verification → CVC Secretary → CVC decision → attestation and publication), "
     "evidence packs (registration transaction data, KSRSAC GIS / maps, developer "
     "publications, RTC, khata, town planning), gazette effective-date adoption in Kaveri, "
     "and auditability of every rate version used for duty calculation."),
    ("General Revision of market value guidelines for the entire State",
     "General Revision of market value guidelines (CVC instructions → Sub-Committee notice of "
     "intention → 15-day objections → Sub-Committee rate decision and signed statement → DR "
     "verification and 15-day rectification → CVC Secretary → CVC decision → attestation, "
     "publication and sale of approved statements → effective date / e-gazette → Kaveri "
     "adoption), including mid-year revision for specific Sub-Committees."),
    ("Sub-committee and CVC committee workbenches",
     "Sub-Committee and CVC committee workbenches, minutes / proceedings, objection "
     "processing, statement generation (booklet + soft copy), DR verification and "
     "rectification loop, CVC Secretary verification and attestation."),
    ("Publication of preliminary notification of proposed revised guideline values",
     "Publication of the Sub-Committee’s notice of intention (newspapers / notice boards / "
     "portal) with a 15-day objections period; publication of approved statements in "
     "sub-district offices and DR office and printed copies for sale; final e-gazette / "
     "effective-date management after CVC approval."),
    ("Trigger: IGR as CVC Chairman issues an order",
     "Trigger: CVC (Chairman IGR) issues instructions and general policy guidelines to all "
     "Market Valuation Sub-Committees for estimation of market value guidelines for the next "
     "calendar year (Rule 5(1)); CVC may also instruct any Sub-Committee to revise rates at "
     "any time during the calendar year."),
    ("Cascade: DIGR notification",
     "Cascade: CVC instructions → Sub-Committee notice of intention (newspapers / notice "
     "boards) → 15-day objections → Secretary processes objections → Sub-Committee decides "
     "rates → signed statement (booklet + soft copy) → DR verification (15-day "
     "rectification) → DR views, per-sub-district booklets → CVC Secretary verification → "
     "CVC decision → attestation → DR → Sub-Committees → publication and sale → effective "
     "date / e-gazette → Kaveri adoption (system extension)."),
    ("Channels: Officer workbenches (IGR, DIGR",
     "Channels: Officer workbenches (CVC / CVC Secretary, DR, Sub-Committee Secretary (SR), "
     "Sub-Committee members); Citizen portal for viewing the notice of intention, filing "
     "objections / suggestions during the 15-day period, and viewing approved statements."),
    ("Digitised General Revision cycle with order",
     "Digitised General Revision cycle aligned to Rules 5 and 7: CVC instructions → notice "
     "of intention → 15-day objections → Sub-Committee statement → DR verification → CVC "
     "Secretary → CVC decision → attestation and publication → e-gazette → effective date → "
     "automatic adoption."),
    ("Objection capture and Sub-committee disposal",
     "Objection / suggestion capture during the 15-day window after the notice of intention, "
     "with Sub-Committee views recorded in the signed statement."),
    ("IGR / CVC dashboard:",
     "IGR / CVC dashboard: open cycles by stage, CVC agenda, decisions and proceedings, "
     "gazette linking."),
    ("DIGR notification console.",
     "CVC Secretary console: issue CVC instructions, verify DR statements, prepare CVC "
     "agenda, attest and transmit approved statements."),
    ("DR inbox: forward to SR",
     "DR inbox: verify Sub-Committee statements, return for rectification (15-day clock), "
     "record views, send per-sub-district booklets to CVC Secretary, forward attested "
     "statements; individual project accept / inspect / propose."),
    ("SR workbench: evidence pack",
     "Sub-Committee Secretary (SR) workbench: publish notice of intention, process "
     "objections, evidence pack, prepare statement, generate booklet / soft copy, "
     "rectification, record publication of approved statements."),
    ("Sub-committee / CVC meeting screens",
     "Sub-Committee / CVC meeting screens: attendance, rate decisions, suggestions "
     "accepted / rejected, minutes / proceedings, e-signing."),
    ("Citizen portal: apply for individual project fixation",
     "Citizen portal: apply for individual project fixation; view notice of intention; file "
     "objections / suggestions within 15 days; view / purchase approved statements; track "
     "status."),
    ("Process walkthrough of General Revision",
     "Process walkthrough of General Revision (15 statutory stages + 2 system extension "
     "steps) and Individual Project (6 steps) with sandbox revision cycle."),
    ("Job aids: evidence pack checklist",
     "Job aids: evidence pack checklist; statement template (CVC form); DR verification "
     "checklist; Sub-Committee minutes template; gazette linking checklist."),
    ("Portal messaging for public opinion",
     "Portal messaging for notice of intention and 15-day objections window, availability of "
     "approved statements, and individual project fixation eligibility; FAQs linking "
     "Sec. 45-B and Rules 5 and 7."),
    ("Workshop process (user-confirmed)",
     "Workshop process (user-confirmed): Fixation of guidance value for new individual "
     "project (6 steps). General Revision superseded by the 15-stage workflow below."),
]

APPENDIX_ADD = (
    "Finalized BRD/CVCGUIDANCEVALUEFIXATION/Stages/Guideline_Value_Revision_Process_15_Steps.docx "
    "— consolidated 15-stage General Revision workflow based on Rules 5 and 7 (basis for v1.7)."
)

BUSINESS_RULES = {
    "BR-CVC-02": "Public objection / suggestion period shall be 15 days from publication of the Sub-Committee’s notice of intention (Rule 5(2)); stored per cycle and changeable only by authorised CVC Admin with recorded reason.",
    "BR-CVC-03": "Discrepancies / omissions referred by the District Registrar shall be rectified and the statement re-submitted within 15 days from the date of reference (Rule 5(3)).",
}
BUSINESS_RULES_ADD = [
    ("BR-CVC-07", "A Sub-Committee statement may be submitted to the DR only when signed by both the Sub-Committee Secretary and Chairman, in the form specified by CVC."),
    ("BR-CVC-08", "Only statements approved by CVC and attested by the CVC Secretary may be published, sold, gazetted or activated in Kaveri."),
]

LEGAL_RULE_ROWS = [
    ("CVC Rules 2003 — Rule 5", "CVC instructions to Sub-Committees (general / mid-year); notice of intention; 15-day objections; statement of average rates signed by Secretary and Chairman; booklet + soft copy to DR; DR verification and 15-day rectification; per-sub-district booklets to CVC Secretary", "7.ii — General Revision stages 1–11"),
    ("CVC Rules 2003 — Rule 7", "CVC Secretary verification; CVC discussion of district-wise estimations and Sub-Committee / DR suggestions; proceedings; attestation; publication in sub-district offices / DR office; printed copies for sale at CVC price", "7.ii — General Revision stages 12–15"),
]

ENTITY_UPDATES = {
    "RevisionCycle": "cycle_id, cycle_type (General / Mid-year), target_calendar_year, scope (Sub-Committees / jurisdictions), instruction_ref, objection_period_days, target_decision_date, status, dates",
    "DigirNotification": None,  # replaced by CvcInstruction
    "SrProposal": None,  # replaced by RateStatement
    "ProposedRateLine": "line_id, statement_id, valuation_area_id, village / local_body keys, rate_category (Agri/Non-Agri), sub_type (Dry/Wet/Bhagayat or Residential/Commercial/Industrial), proposed_rate, unit, extent_sqm_ref, dr_views, cvc_outcome, remarks",
    "SubCommitteeSitting": "sitting_id, cycle_id, sub_committee_id, members / attendance, objections_considered, decisions, minutes",
    "PublicNotice": "notice_id, cycle_id, sub_committee_id, notice_type (Intention / Approved statement), newspaper and notice-board publications (proof), start, end (15 days), portal_url",
    "Objection": "objection_id, notice_id, objector, channel (online / SRO), linked village / local body / rate lines, text, secretary_remarks, sub_committee_views",
    "CvcDecision": "decision_id, cycle_id, district / sub-district, suggestions (accepted / rejected), proceedings_ref, outcome, effective_date, gazette_ref",
}
ENTITY_ADD = [
    ("CvcInstruction", "instruction_id, cycle_id, policy_guidelines_doc, target Sub-Committees, issued_on"),
    ("RateStatement", "statement_id, cycle_id, sub_committee_id, form_template_id, status, signed_by_secretary / chairman (eSign), booklet_pdf, soft_copy_ref, version"),
    ("DrVerification", "verification_id, statement_id, discrepancies, referred_on, rectification_due (referred_on + 15 days), resubmitted_on, dr_views, sent_to_cvc_secretary_on"),
    ("AttestedStatement", "attestation_id, statement_id, attested_by (CVC Secretary), sent_to_dr_on, forwarded_to_sub_committee_on"),
    ("StatementPublication", "publication_id, attestation_id, office (sub-district / SRO / DR), displayed_on, proof, sale_price, copies_printed, copies_sold"),
]

INTEGRATION_ADD = (
    "eSign / Digital Signature (DSC)",
    "Signing of Sub-Committee statements by Secretary and Chairman and attestation of approved statements by CVC Secretary",
)

N_001 = ("System shall send SMS / e-Mail on key events: Revision Cycle opened / CVC instructions issued; notice of intention published; objections window open / close; objection acknowledgement; statement submitted to DR; returned for rectification and 15-day due-date reminders; sent to CVC Secretary; CVC decision; attested statements transmitted; approved statements published; rates Active in Kaveri; individual application status changes.")
N_002 = ("The notice of intention and the approved statements shall appear on the citizen portal; approved statements shall be downloadable, with information on purchase of printed copies.")
R_001 = ("MIS: Revision Cycles by stage and district; notices of intention published; objection counts per Sub-Committee; Sub-Committee sittings; DR verification and rectification ageing against the 15-day limit; statements pending with CVC Secretary / CVC; publication and sale status.")

RISK_ADD = (
    "Sub-Committee misses 15-day rectification or CVC decision deadline",
    "Statements delayed; old rates persist",
    "Due-date tracking, reminders and escalation to DR / DIGR / IGR; ageing MIS",
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def header(table) -> list[str]:
    return [c.text.strip() for c in table.rows[0].cells] if table.rows else []


def find_table(doc, hdr_prefix: list[str], contains: str | None = None):
    for t in doc.tables:
        if header(t)[: len(hdr_prefix)] != hdr_prefix:
            continue
        if contains is None:
            return t
        blob = " ".join(c.text for r in t.rows for c in r.cells)
        if contains in blob:
            return t
    raise SystemExit(f"Table not found: {hdr_prefix} / {contains}")


def clone_row(table, template_tr, values: list[str], bold_first: bool = True):
    tr = copy.deepcopy(template_tr)
    table._tbl.append(tr)
    row = table.rows[-1]
    for i, v in enumerate(values):
        set_cell_text(row.cells[i], v, bold=(bold_first and i == 0))
    return row


def rebuild_table(table, rows: list[tuple], bold_first: bool = True) -> None:
    template = copy.deepcopy(table.rows[1]._tr)
    for r in list(table.rows[1:]):
        table._tbl.remove(r._tr)
    for values in rows:
        clone_row(table, template, list(values), bold_first)


def append_rows(table, rows: list[tuple], bold_first: bool = True) -> None:
    existing = {r.cells[0].text.strip() for r in table.rows[1:]}
    template = table.rows[-1]._tr
    for values in rows:
        if values[0] in existing:
            continue
        clone_row(table, template, list(values), bold_first)


# ---------------------------------------------------------------------------
# Updates
# ---------------------------------------------------------------------------


def update_document_control(doc) -> None:
    t = find_table(doc, ["Field", "Value"], "Document ID")
    for row in t.rows:
        key = row.cells[0].text.strip()
        if key == "Version":
            set_cell_text(row.cells[1], OUT_VERSION)
        elif key == "Last updated":
            set_cell_text(row.cells[1], OUT_DATE)
        elif key == "Status":
            set_cell_text(
                row.cells[1],
                "Draft — based on Sr.14 discussion (31-08-2026 to 01-09-2026); General "
                "Revision aligned to the 15-stage workflow of Rules 5 and 7; Individual "
                "Project fixation per workshop process",
            )
        elif key == "State rules (primary)":
            text = row.cells[1].text.strip()
            if "Rules 5" not in text:
                set_cell_text(
                    row.cells[1],
                    "Karnataka Stamp (Constitution of Central Valuation Committee …) Rules, "
                    "2003 — Rules 3–7 (Rule 5: Sub-Committee estimation procedure; Rule 7: "
                    "CVC approval and publication); " + text,
                )

    vh = find_table(doc, ["Version", "Date", "Author"])
    if OUT_VERSION not in [r.cells[0].text.strip() for r in vh.rows[1:]]:
        row = vh.add_row()
        for i, v in enumerate([OUT_VERSION, OUT_DATE, "Nandha Kumar", VERSION_SUMMARY, "Prashanth"]):
            set_cell_text(row.cells[i], v)


def update_paragraphs(doc) -> None:
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if t in FR_HEADINGS:
            replace_paragraph_text(p, FR_HEADINGS[t])
            continue
        for start, new in PARA_REPLACEMENTS:
            if t.startswith(start):
                replace_paragraph_text(p, new)
                break


def rewrite_gr_steps(doc) -> None:
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.strip().startswith("Process steps — General Revision"))
    steps = []
    for p in paras[start + 1:]:
        if p.style.name.startswith("Heading"):
            break
        if p.style.name == "List Number":
            steps.append(p)
    if not steps:
        raise SystemExit("General Revision steps not found")
    template = steps[0]
    prev = steps[-1]._p
    for text in GR_STEPS:
        new_p = copy.deepcopy(template._p)
        prev.addnext(new_p)
        prev = new_p
        replace_paragraph_text(Paragraph(new_p, template._parent), text)
    for p in steps:
        p._p.getparent().remove(p._p)


def update_status_model(doc) -> None:
    t = find_table(doc, ["Status", "Meaning", "Owner"], "Order Issued")
    rebuild_table(t, GR_STATUS_ROWS)


def update_frs(doc) -> None:
    targets = [
        (find_table(doc, ["Req ID", "Requirement", "Priority"], marker), rows)
        for rows, marker in ((GR_FR_A, "FR-CVC-GR-001"), (GR_FR_B, "FR-CVC-GR-006"), (GR_FR_C, "FR-CVC-GR-011"))
    ]
    for t, rows in targets:
        rebuild_table(t, rows)

    t = find_table(doc, ["Req ID", "Requirement", "Priority"], "FR-CVC-N-001")
    for r in t.rows[1:]:
        rid = r.cells[0].text.strip()
        if rid == "FR-CVC-N-001":
            set_cell_text(r.cells[1], N_001)
        elif rid == "FR-CVC-N-002":
            set_cell_text(r.cells[1], N_002)

    t = find_table(doc, ["Req ID", "Requirement", "Priority"], "FR-CVC-R-001")
    for r in t.rows[1:]:
        if r.cells[0].text.strip() == "FR-CVC-R-001":
            set_cell_text(r.cells[1], R_001)


def update_actors_glossary(doc) -> None:
    t = find_table(doc, ["Actor", "Role in Guidance Value Fixation"])
    for r in t.rows[1:]:
        key = r.cells[0].text.strip()
        if key in ACTORS:
            set_cell_text(r.cells[1], ACTORS[key])

    g = find_table(doc, ["Term", "Definition"])
    for r in g.rows[1:]:
        key = r.cells[0].text.strip()
        if key in GLOSSARY_UPDATE:
            set_cell_text(r.cells[1], GLOSSARY_UPDATE[key])
    append_rows(g, GLOSSARY_ADD)


def update_legal(doc) -> None:
    t = find_table(doc, ["Rule / provision", "Requirement"])
    append_rows(t, LEGAL_RULE_ROWS)

    n = find_table(doc, ["Instrument", "Date / No."])
    for r in n.rows[1:]:
        if r.cells[0].text.strip().startswith("CVC / IGR order"):
            set_cell_text(r.cells[0], "CVC instructions and general policy guidelines (Rule 5(1))", bold=True)
            set_cell_text(r.cells[1], "Per revision cycle (annual for next calendar year; or mid-year)")
            set_cell_text(r.cells[2], "Triggers General Revision in Sub-Committees; CVC approves statements (Rule 7); gazette fixes effective date")
            set_cell_text(r.cells[3], "7.ii — General Revision process (stages 1–15 + system extension)")


def update_business_rules(doc) -> None:
    t = find_table(doc, ["Rule ID", "Rule"])
    for r in t.rows[1:]:
        rid = r.cells[0].text.strip()
        if rid in BUSINESS_RULES:
            set_cell_text(r.cells[1], BUSINESS_RULES[rid])
    append_rows(t, BUSINESS_RULES_ADD)


def update_entities(doc) -> None:
    t = find_table(doc, ["Entity", "Key attributes"])
    for r in list(t.rows[1:]):
        key = r.cells[0].text.strip()
        if key in ENTITY_UPDATES:
            if ENTITY_UPDATES[key] is None:
                t._tbl.remove(r._tr)
            else:
                set_cell_text(r.cells[1], ENTITY_UPDATES[key])
    # keep new General Revision entities grouped directly after RevisionCycle
    existing = {r.cells[0].text.strip() for r in t.rows[1:]}
    anchor = t.rows[1]._tr
    for name, attrs in ENTITY_ADD:
        if name in existing:
            continue
        tr = copy.deepcopy(anchor)
        anchor.addnext(tr)
        anchor = tr
        row = next(r for r in t.rows if r._tr is tr)
        set_cell_text(row.cells[0], name, bold=True)
        set_cell_text(row.cells[1], attrs)


def update_integrations_rtm_risk(doc) -> None:
    t = find_table(doc, ["System", "Purpose"])
    if not any("eSign" in r.cells[0].text for r in t.rows):
        clone_row(t, t.rows[-1]._tr, list(INTEGRATION_ADD), bold_first=False)
    for r in t.rows[1:]:
        if r.cells[0].text.strip().startswith("Department of Information and Public Relations"):
            set_cell_text(
                r.cells[1],
                "Publish e-gazette notification for CVC-approved guidance values (system "
                "extension after Rule 7(3) publication); optionally place newspaper notice of "
                "intention; return gazette particulars for linkage to rate versions",
            )

    rtm = find_table(doc, ["Req ID", "Legal anchor", "Process"])
    for r in rtm.rows[1:]:
        c0 = r.cells[0].text.strip()
        if c0.startswith("FR-CVC-GR-001"):
            set_cell_text(r.cells[0], "FR-CVC-GR-001…022", bold=True)
            set_cell_text(r.cells[1], "Sec. 45-B; CVC Rules 2003 Rules 5 and 7")
        elif c0.startswith("FR-CVC-004/005/014"):
            set_cell_text(r.cells[0], "FR-CVC-004/005; FR-CVC-GR-021", bold=True)

    risk = find_table(doc, ["Risk", "Impact", "Mitigation"])
    if not any(RISK_ADD[0] == r.cells[0].text.strip() for r in risk.rows):
        clone_row(risk, risk.rows[-1]._tr, list(RISK_ADD), bold_first=False)

    us = find_table(doc, ["US ID", "Summary", "BRD coverage"])
    for r in us.rows[1:]:
        if r.cells[0].text.strip() == "US-CVC-02":
            set_cell_text(r.cells[2], "FR-CVC-004, FR-CVC-GR-021, FR-CVC-IP-006")


def add_appendix_reference(doc) -> None:
    paras = doc.paragraphs
    for p in paras:
        if "Guideline_Value_Revision_Process_15_Steps" in p.text:
            return
    target = next(p for p in paras if p.text.strip().startswith("Workshop process (user-confirmed)"))
    new_p = copy.deepcopy(target._p)
    target._p.addnext(new_p)
    replace_paragraph_text(Paragraph(new_p, target._parent), APPENDIX_ADD)


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    update_document_control(doc)
    update_paragraphs(doc)
    rewrite_gr_steps(doc)
    update_status_model(doc)
    update_frs(doc)
    update_actors_glossary(doc)
    update_legal(doc)
    update_business_rules(doc)
    update_entities(doc)
    update_integrations_rtm_risk(doc)
    add_appendix_reference(doc)

    try:
        doc.save(str(DST))
    except PermissionError:
        alt = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.7_updated.docx"
        doc.save(str(alt))
        print(f"NOTE: locked; wrote {alt}")
        return
    print(f"Wrote {DST}")


if __name__ == "__main__":
    main()
