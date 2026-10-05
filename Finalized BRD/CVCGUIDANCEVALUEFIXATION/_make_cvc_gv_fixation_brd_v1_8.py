# -*- coding: utf-8 -*-
"""Create BRD_CVC_Guidance_Value_Fixation_v1.8.docx from v1.7.

Adds the Process A (General Revision) and Process B (Individual Project
Fixation) swimlane process diagrams (PNG exports of the draw.io files in
Process Diagram/) after the respective process steps.
"""
from __future__ import annotations

import copy
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\CVCGUIDANCEVALUEFIXATION")
SRC = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.7.docx"
DST = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.8.docx"
DIAGRAMS = BASE / "Process Diagram"
OUT_VERSION = "1.8"
OUT_DATE = "30-09-2026"

MAX_W_IN = 6.5
MAX_H_IN = 8.2

FIGURES = [
    {
        "after_heading": "Process steps — General Revision",
        "png": DIAGRAMS / "Process_A_General_Revision.png",
        "intro": "Figure 1 shows the General Revision steps above as a swimlane process diagram, "
                 "one lane per actor. Dashed grey boxes are Kaveri system steps added beyond the "
                 "Rule 5 / 7 text.",
        "caption": "Figure 1 — Process A: General Revision of guidance value (Rules 5 and 7). "
                   "Source: Process Diagram/Process_A_General_Revision.drawio",
    },
    {
        "after_heading": "Process steps — Individual Project Fixation",
        "png": DIAGRAMS / "Process_B_Individual_Project_Fixation.png",
        "intro": "Figure 2 shows the Individual Project Fixation steps above as a swimlane process "
                 "diagram, one lane per actor.",
        "caption": "Figure 2 — Process B: Fixing guidance value for a new individual project "
                   "(Sec. 45-B). Source: Process Diagram/Process_B_Individual_Project_Fixation.drawio",
    },
]

APPENDIX_ADD = (
    "Finalized BRD/CVCGUIDANCEVALUEFIXATION/Process Diagram/ — Process_A_General_Revision.drawio / .png "
    "and Process_B_Individual_Project_Fixation.drawio / .png (swimlane process diagrams, Figures 1 and 2)."
)

VERSION_SUMMARY = (
    "Added swimlane process diagrams: Figure 1 (Process A — General Revision, after General "
    "Revision process steps) and Figure 2 (Process B — Individual Project Fixation, after its "
    "process steps); draw.io sources and PNGs stored in Process Diagram/; Appendix A updated."
)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def update_document_control(doc) -> None:
    for table in doc.tables:
        keys = [r.cells[0].text.strip() for r in table.rows]
        if "Document ID" in keys:
            for row in table.rows:
                key = row.cells[0].text.strip()
                if key == "Version":
                    set_cell_text(row.cells[1], OUT_VERSION)
                elif key == "Last updated":
                    set_cell_text(row.cells[1], OUT_DATE)
            break
    for table in doc.tables:
        if [c.text.strip() for c in table.rows[0].cells][:3] != ["Version", "Date", "Author"]:
            continue
        if OUT_VERSION in [r.cells[0].text.strip() for r in table.rows[1:]]:
            return
        row = table.add_row()
        for i, v in enumerate([OUT_VERSION, OUT_DATE, "Nandha Kumar", VERSION_SUMMARY, "Prashanth"]):
            set_cell_text(row.cells[i], v)
        return


def picture_size(png: Path):
    with Image.open(png) as im:
        w, h = im.size
    width = MAX_W_IN
    if width * h / w > MAX_H_IN:
        width = MAX_H_IN * w / h
    return Inches(width)


def new_paragraph_after(doc, anchor_el, style: str) -> Paragraph:
    p = doc.add_paragraph(style=style)
    anchor_el.addnext(p._p)
    return p


def insert_figure(doc, fig: dict) -> None:
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.strip() == fig["after_heading"])
    last = paras[start]
    for p in paras[start + 1:]:
        if p.style.name.startswith("Heading"):
            break
        if p.text.strip() == fig["caption"]:
            return
        last = p

    intro = new_paragraph_after(doc, last._p, "Normal")
    intro.add_run(fig["intro"])

    pic = new_paragraph_after(doc, intro._p, "Normal")
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(fig["png"]), width=picture_size(fig["png"]))

    cap = new_paragraph_after(doc, pic._p, "Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(fig["caption"])


def add_appendix_reference(doc) -> None:
    paras = doc.paragraphs
    if any("Process Diagram/" in p.text and p.style.name == "List Number" for p in paras):
        return
    target = next(p for p in paras if "Guideline_Value_Revision_Process_15_Steps" in p.text)
    new_p = copy.deepcopy(target._p)
    target._p.addnext(new_p)
    replace_paragraph_text(Paragraph(new_p, target._parent), APPENDIX_ADD)


def main():
    for fig in FIGURES:
        if not fig["png"].exists():
            raise SystemExit(f"Missing diagram PNG: {fig['png']} (run Process Diagram/_export_png.py)")
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    update_document_control(doc)
    for fig in FIGURES:
        insert_figure(doc, fig)
    add_appendix_reference(doc)

    try:
        doc.save(str(DST))
    except PermissionError:
        alt = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.8_updated.docx"
        doc.save(str(alt))
        print(f"NOTE: locked; wrote {alt}")
        return
    print(f"Wrote {DST}")


if __name__ == "__main__":
    main()
