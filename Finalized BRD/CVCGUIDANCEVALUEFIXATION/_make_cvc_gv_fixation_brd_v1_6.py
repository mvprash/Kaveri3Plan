# -*- coding: utf-8 -*-
"""Create BRD_CVC_Guidance_Value_Fixation_v1.6.docx from v1.5.

Individual Property / Project fixation: integrate with ULMS to extract
geo-fenced polygon details using ULPIN as input.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\CVCGUIDANCEVALUEFIXATION")
SRC = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.5.docx"
DST = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.6.docx"
OUT_VERSION = "1.6"
OUT_DATE = "22-09-2026"

FR_IP_008 = (
    "For Individual Property / Project fixation, the system shall integrate with "
    "ULMS (Unified Land Management System / notified ULMS interface) to extract "
    "geo-fenced polygon details of the property. The input to ULMS shall be the "
    "ULPIN (Unique Land Parcel Identification Number). On successful response, "
    "the system shall store the returned geo-fence (polygon coordinates) against "
    "the application / valuation area and make it available for inspection, "
    "evidence and rate fixation (aligned to FR-CVC-009)."
)

INTEGRATION_SYSTEM = "ULMS (ULPIN-based land parcel / geo-fence service)"
INTEGRATION_PURPOSE = (
    "Individual Property fixation: accept ULPIN as input and retrieve geo-fenced "
    "polygon details for the parcel; persist coordinates on the application / "
    "valuation area"
)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(size)


def update_document_control_version(doc: Document) -> None:
    for table in doc.tables:
        if len(table.columns) != 2:
            continue
        keys = [r.cells[0].text.strip() for r in table.rows]
        if "Document ID" not in keys:
            continue
        for row in table.rows:
            key = row.cells[0].text.strip()
            if key == "Version":
                set_cell_text(row.cells[1], OUT_VERSION)
            if key == "Last updated":
                set_cell_text(row.cells[1], OUT_DATE)


def append_version_history_row(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [c.text.strip() for c in table.rows[0].cells]
        if headers[:3] != ["Version", "Date", "Author"]:
            continue
        for r in table.rows[1:]:
            if r.cells[0].text.strip() == OUT_VERSION:
                return
        row = table.add_row()
        values = [
            OUT_VERSION,
            OUT_DATE,
            "Nandha Kumar",
            "Added FR-CVC-IP-008: Individual Property fixation integration with ULMS "
            "using ULPIN to extract geo-fenced polygon details; Integrations table updated",
            "Prashanth",
        ]
        for i, val in enumerate(values):
            if i < len(row.cells):
                set_cell_text(row.cells[i], val)
        return


def find_ip_fr_table(doc: Document):
    for table in doc.tables:
        if not table.rows or table.rows[0].cells[0].text.strip() != "Req ID":
            continue
        ids = [r.cells[0].text.strip() for r in table.rows[1:]]
        if "FR-CVC-IP-001" in ids:
            return table
    return None


def add_ip_fr(doc: Document) -> None:
    table = find_ip_fr_table(doc)
    if table is None:
        raise SystemExit("Individual Project FR table not found")

    for r in table.rows[1:]:
        if r.cells[0].text.strip() == "FR-CVC-IP-008":
            set_cell_text(r.cells[1], FR_IP_008)
            set_cell_text(r.cells[2], "Must")
            break
    else:
        # Insert after IP-001 (application capture) so ULPIN/ULMS is early in flow
        after_idx = None
        for i, r in enumerate(table.rows):
            if r.cells[0].text.strip() == "FR-CVC-IP-001":
                after_idx = i
                break
        new_row = table.add_row()
        set_cell_text(new_row.cells[0], "FR-CVC-IP-008", bold=True)
        set_cell_text(new_row.cells[1], FR_IP_008)
        set_cell_text(new_row.cells[2], "Must")
        if after_idx is not None:
            tr = new_row._tr
            table._tbl.remove(tr)
            table.rows[after_idx]._tr.addnext(tr)

    # Enrich IP-001 to mention ULPIN capture
    for r in table.rows[1:]:
        if r.cells[0].text.strip() != "FR-CVC-IP-001":
            continue
        text = r.cells[1].text.strip()
        if "ULPIN" not in text:
            set_cell_text(
                r.cells[1],
                text.rstrip(".")
                + ". Application shall capture ULPIN where available; system shall "
                "use ULPIN to fetch the geo-fenced polygon from ULMS (FR-CVC-IP-008).",
            )
        break


def add_integration(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        hdr = [c.text.strip() for c in table.rows[0].cells]
        if hdr[:2] != ["System", "Purpose"]:
            continue
        for r in table.rows[1:]:
            if "ULMS" in r.cells[0].text or "ULPIN" in r.cells[0].text:
                set_cell_text(r.cells[0], INTEGRATION_SYSTEM)
                set_cell_text(r.cells[1], INTEGRATION_PURPOSE)
                return
        row = table.add_row()
        set_cell_text(row.cells[0], INTEGRATION_SYSTEM)
        set_cell_text(row.cells[1], INTEGRATION_PURPOSE)
        return
    raise SystemExit("Integrations table not found")


def update_data_entity(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        hdr = [c.text.strip() for c in table.rows[0].cells]
        if hdr[:2] != ["Entity", "Key attributes"]:
            continue
        for row in table.rows[1:]:
            if row.cells[0].text.strip() == "IndividualFixationApp":
                attrs = row.cells[1].text.strip()
                if "ULPIN" not in attrs:
                    set_cell_text(
                        row.cells[1],
                        attrs.rstrip(".")
                        + "; ulpin; ulms_geo_fence_polygon; ulms_fetch_status / timestamp",
                    )
                return


def update_rtm(doc: Document) -> None:
    for table in doc.tables:
        if not table.rows:
            continue
        for row in table.rows[1:]:
            cell0 = row.cells[0].text.strip()
            if cell0.startswith("FR-CVC-IP-001") and "008" not in cell0:
                set_cell_text(row.cells[0], cell0.replace("007", "008"))
                return


def enrich_process_step(doc: Document) -> None:
    """Mention ULMS/ULPIN in Individual Project process steps if present."""
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if t.startswith("Party / citizen applies") and "ULPIN" not in t:
            p.runs[0].text = (
                t.rstrip(".")
                + ", including ULPIN where available for ULMS geo-fence retrieval "
                "(FR-CVC-IP-008)."
            )
            for r in p.runs[1:]:
                r.text = ""
            return
        if "applies for fixation of guidance value" in t.lower() and "ULPIN" not in t:
            if p.runs:
                p.runs[0].text = (
                    t.rstrip(".")
                    + "; capture ULPIN and fetch geo-fenced polygon from ULMS "
                    "(FR-CVC-IP-008)."
                )
                for r in p.runs[1:]:
                    r.text = ""
            return


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    update_document_control_version(doc)
    append_version_history_row(doc)
    add_ip_fr(doc)
    add_integration(doc)
    update_data_entity(doc)
    update_rtm(doc)
    enrich_process_step(doc)

    try:
        doc.save(str(DST))
    except PermissionError:
        alt = BASE / "BRD_CVC_Guidance_Value_Fixation_v1.6_updated.docx"
        doc.save(str(alt))
        print(f"NOTE: locked; wrote {alt}")
        return

    print(f"Wrote {DST}")
    table = find_ip_fr_table(doc)
    for r in table.rows[1:]:
        rid = r.cells[0].text.strip()
        if rid in ("FR-CVC-IP-001", "FR-CVC-IP-008"):
            print(rid, "→", r.cells[1].text.strip()[:130])


if __name__ == "__main__":
    main()
