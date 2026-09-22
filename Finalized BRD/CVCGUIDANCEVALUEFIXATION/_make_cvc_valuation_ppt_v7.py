# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v7.pptx from v6.

Rebuilds the Rule 6 guidelines slide to include electronic data-source tags
(department / system) for each estimation factor group.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation_v6.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v7.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x2A, 0x37)
MUTED = RGBColor(0x5C, 0x6B, 0x7B)
ACCENT = RGBColor(0x1F, 0x6B, 0x7A)
TAG_BG = RGBColor(0xD7, 0xEB, 0xF0)

MARGIN = Inches(0.45)
CONTENT_W = Inches(12.43)
FOOTER_Y = Inches(7.14)


def _run(para, text, *, size, bold=False, italic=False, color=INK, font="Calibri"):
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return run


def add_textbox(
    slide, x, y, w, h, text, *, size, bold=False, italic=False, color=INK,
    font="Calibri", align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, cambria=False,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    para.clear()
    _run(
        para, text, size=size, bold=bold, italic=italic, color=color,
        font="Cambria" if cambria else font,
    )
    return box


def add_rect(slide, x, y, w, h, fill, rounded=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.06
    return shape


def chrome(slide, kicker: str, title: str, page: str):
    add_textbox(slide, MARGIN, Inches(0.38), Inches(10.2), Inches(0.26),
                kicker, size=11, bold=True, color=MUTED)
    add_textbox(slide, MARGIN, Inches(0.62), Inches(11.5), Inches(0.48),
                title, size=22, bold=True, color=NAVY, cambria=True, anchor=MSO_ANCHOR.MIDDLE)
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.35), Inches(0.36), Inches(0.52), Inches(0.52))
    badge.fill.solid()
    badge.fill.fore_color.rgb = NAVY
    badge.line.fill.background()
    add_textbox(slide, Inches(12.35), Inches(0.36), Inches(0.52), Inches(0.52),
                page, size=14, bold=True, color=WHITE, cambria=True,
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, MARGIN, FOOTER_Y, Inches(8.5), Inches(0.28),
                "VALUATION MODULE (CVC)  |  BRD WALKTHROUGH  ·  KAVERI 3.0", size=8.5, color=MUTED)
    add_textbox(slide, Inches(12.3), FOOTER_Y, Inches(0.6), Inches(0.28),
                page, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title: str, bullets: list[str], tags: list[str], body_size=9.0):
    add_rect(slide, x, y, w, h, ALT)

    add_textbox(
        slide, x + Inches(0.12), y + Inches(0.08), w - Inches(0.24), Inches(0.26),
        title, size=11.5, bold=True, color=NAVY, cambria=True,
    )

    # Factor bullets (leave room for e-source tag strip at bottom)
    bullet_h = h - Inches(1.15)
    box = slide.shapes.add_textbox(
        x + Inches(0.12), y + Inches(0.36), w - Inches(0.24), bullet_h
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.space_after = Pt(1.5)
        para.line_spacing = 1.02
        para.clear()
        _run(para, f"•  {line}", size=body_size, color=INK)

    # Electronic source tag strip
    tag_y = y + h - Inches(0.72)
    add_rect(slide, x + Inches(0.08), tag_y, w - Inches(0.16), Inches(0.62), TAG_BG)
    add_textbox(
        slide, x + Inches(0.14), tag_y + Inches(0.04), w - Inches(0.28), Inches(0.18),
        "e-Source / Department", size=8, bold=True, italic=True, color=ACCENT,
    )
    tag_line = "  ".join(f"#{t}" for t in tags)
    add_textbox(
        slide, x + Inches(0.14), tag_y + Inches(0.22), w - Inches(0.28), Inches(0.36),
        tag_line, size=8.5, bold=True, italic=True, color=NAVY,
    )


def clear_slide(slide):
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)


def find_rule6_slide(prs: Presentation):
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if hasattr(sh, "text") and "Guidelines for Estimating Guidance Value" in sh.text:
                return i, slide
    raise SystemExit("Rule 6 slide not found in source deck")


def build_rule6_slide(slide, page: str):
    chrome(
        slide,
        "LEGAL & REGULATORY REFERENCE",
        "Guidelines for Estimating Guidance Value",
        page,
    )

    add_textbox(
        slide, MARGIN, Inches(1.08), CONTENT_W, Inches(0.38),
        "Karnataka Stamp (Constitution of C.V.C. etc.) Rules, 2003 — Rule 6.  "
        "Average-rate factors for Sub-Committee statements — also the evidence checklist "
        "when proposing new guidance value.  Tags (#) show where data can be extracted electronically.",
        size=9.5, italic=True, color=MUTED,
    )

    gap = Inches(0.12)
    card_w = (CONTENT_W - 3 * gap) / 4
    y = Inches(1.50)
    h = Inches(3.70)

    add_card(
        slide, MARGIN, y, card_w, h,
        "Lands — R.6(1)(a)",
        [
            "Dry / garden / wet classification",
            "Soil class in survey records",
            "Other valuation influencers",
            "Value of adjacent / vicinity lands",
            "Crop nature & 5-year avg. yield",
            "Nearness to road & market",
            "Distance from village; location; level",
            "Transport & irrigation facilities",
        ],
        tags=["Bhoomi/RTC", "SSLR", "KSRSAC-GIS", "Agriculture", "Irrigation", "Kaveri-Regn"],
        body_size=8.5,
    )

    add_card(
        slide, MARGIN + card_w + gap, y, card_w, h,
        "House sites — R.6(1)(b)",
        [
            "General locality site values",
            "Proximity to road / rail / bus",
            "Proximity to market & shops",
            "Offices, hospitals, schools",
            "Development / industrial activity",
            "Land tax & local body valuation",
            "Other material / special features",
        ],
        tags=["Kaveri-Regn", "KSRSAC-GIS", "ULB/e-Aasthi", "DTCP/BDA/UDA", "Transport"],
        body_size=8.5,
    )

    add_card(
        slide, MARGIN + 2 * (card_w + gap), y, card_w, h,
        "Other properties — R.6(1)(c)",
        [
            "Nature & condition of property",
            "Purpose for which used",
            "Any other special valuation features",
        ],
        tags=["Kaveri-Regn", "ULB-BldgPlan", "Khata", "KSRSAC-GIS"],
        body_size=9.5,
    )

    add_card(
        slide, MARGIN + 3 * (card_w + gap), y, card_w, h,
        "Rate suggestions — R.6(2)",
        [
            "Non-agri / industrial: agri multiple or per sq.ft",
            "Agricultural: dry / wet / garden + village proximity",
            "Plantations: treat as garden lands",
            "Buildings: PWD construction norms",
            "Civic amenities in construction cost",
        ],
        tags=["Bhoomi", "PWD-SoR", "Horticulture", "BESCOM/ESCOM", "BWSSB/ULB"],
        body_size=8.5,
    )

    # Bottom evidence + master source legend
    add_rect(slide, MARGIN, Inches(5.32), CONTENT_W, Inches(1.48), ALT)
    add_textbox(
        slide, MARGIN + Inches(0.18), Inches(5.40), CONTENT_W - Inches(0.36), Inches(0.24),
        "Evidence for proposing new guidance value  ·  Electronic extraction map",
        size=11.5, bold=True, color=NAVY, cambria=True,
    )
    add_textbox(
        slide, MARGIN + Inches(0.18), Inches(5.68), CONTENT_W - Inches(0.36), Inches(1.00),
        "Rule 6 factors = statutory checklist for Sub-Committee statements and for evidence packs "
        "on General Revision / Individual Project proposals (BRD 7.ii.d / 7.iii.c).  "
        "Primary e-feeds (as available): Kaveri registration transactions · Bhoomi/RTC & SSLR survey · "
        "KSRSAC GIS layers · ULB khata / property tax (e-Aasthi) · DTCP / BDA / UDA planning · "
        "PWD Schedule of Rates · Agriculture / Horticulture / Irrigation dept. systems · "
        "utility masters (BESCOM/ESCOM, BWSSB/KUWSDB).  "
        "Registrar verifies statements (15-day rectification); booklets / soft copies to Secretary, CVC "
        "in first week of January.",
        size=9.5, italic=True, color=ACCENT,
    )


def bump_version(prs: Presentation, version: str = "7", date: str = "19-09-2026"):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                full = para.text
                if not full or "BRD-K3-CVC-GVF-001" not in full:
                    continue
                if "Last updated" in full:
                    new = (
                        f"Document ID: BRD-K3-CVC-GVF-001     ·     Version {version}     ·     "
                        f"Status: Draft     ·     Last updated {date}"
                    )
                else:
                    new = f"Document ID: BRD-K3-CVC-GVF-001  ·  Version {version}  ·  {date}"
                if para.runs:
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""


def get_page_number(slide) -> str:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        t = shape.text.strip()
        if re.fullmatch(r"\d{1,2}", t) and shape.left >= 10_500_000 and shape.top < 1_200_000:
            return t
    return "7"


def safe_save(prs: Presentation, path: Path) -> Path:
    try:
        prs.save(str(path))
        return path
    except PermissionError:
        alt = path.with_name(path.stem + "_updated" + path.suffix)
        prs.save(str(alt))
        print(f"NOTE: {path.name} locked; wrote {alt.name}")
        return alt


def main():
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    out = DST
    try:
        shutil.copy2(SRC, out)
    except PermissionError:
        out = BASE / "CVC_Valuation_Module_BRD_Presentation_v7_updated.pptx"
        shutil.copy2(SRC, out)

    prs = Presentation(str(out))
    _idx, slide = find_rule6_slide(prs)
    page = get_page_number(slide)
    clear_slide(slide)
    build_rule6_slide(slide, page=page)
    bump_version(prs, version="7", date="19-09-2026")

    out = safe_save(prs, out)
    print(f"Wrote {out}  (Rule 6 slide index {_idx + 1}, page {page})")
    s = prs.slides[_idx]
    for sh in s.shapes:
        if hasattr(sh, "text") and "#" in sh.text:
            print("TAGS:", sh.text.replace("\n", " | ")[:160])


if __name__ == "__main__":
    main()
