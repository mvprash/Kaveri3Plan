# -*- coding: utf-8 -*-
"""Build CVC_Valuation_Module_BRD_Presentation_v8.pptx from v7, aligned to BRD v1.8.

- General Revision (Process A) re-aligned to the 15-stage workflow of Rules 5 and 7
  (+ Kaveri system steps): trigger / channel, 17-step process flow, 18-status model,
  actors, glossary, legal references, highlights.
- New slides with the swimlane process diagrams for Process A and Process B
  (Process Diagram/*.png, BRD v1.8 Figures 1 and 2).
- Process B slides mention ULPIN / ULMS boundary fetch.
- New slide listing the details captured for each valuation area (BRD FR-CVC-008 to 012).
- Version 8 / 30-09-2026; badge and footer page numbers renumbered.
"""
from __future__ import annotations

import copy
import re
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(__file__).resolve().parent
SRC = BASE / "CVC_Valuation_Module_BRD_Presentation_v7.pptx"
DST = BASE / "CVC_Valuation_Module_BRD_Presentation_v8.pptx"
DIAGRAMS = BASE / "Process Diagram"
VERSION, DATE = "8", "30-09-2026"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
ALT = RGBColor(0xEE, 0xF2, 0xF8)
SYS_BG = RGBColor(0xF5, 0xF5, 0xF5)
MUTED = RGBColor(0x5B, 0x6B, 0x84)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xB8, 0x54, 0x50)

# Process A diagram geometry (logical px, see Process Diagram/_make_process_diagrams.py);
# PNG is exported at 2x with the graph origin at y=0.
A_LANE_TOP, A_LANE_HEADER_END, A_FIRST_ROW, A_ROW_H, A_SPLIT_ROW, PNG_SCALE = 90, 154, 154, 160, 11, 2
A_PNG_BORDER = 10

# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

AREA_DETAILS = [
    ("Area / extent",
     "Size of the valuation area in square metres (m²). m² is the standard stored unit; the area can be shown in "
     "other units without changing the stored value."),
    ("Location (geo-fence)",
     "Geo-coordinates of the area boundary, mapped as a polygon. Used to identify the area on the map, overlay "
     "evidence (e.g. KSRSAC) and apply the correct rate."),
    ("Rate categories",
     "Guidance rates under two categories: Non-Agricultural and Agricultural."),
    ("Non-Agricultural rates",
     "Separate rates for Residential, Commercial and Industrial use."),
    ("Agricultural rates",
     "Separate rates for Dry, Wet and Bhagayat (garden) land."),
]

GR_FLOW = [
    ("CVC", "Sends instructions and guidelines to all Sub-Committees (yearly, or to one Sub-Committee mid-year)", "Rule 5(1)"),
    ("Sub-Committee", "Announces the plan to revise rates in newspapers and on office notice boards", "Rule 5(2)"),
    ("Public", "Sends objections or suggestions within 15 days", "Rule 5(2)"),
    ("Sub-Committee Secretary (SR)", "Sorts objections and places them before the Sub-Committee", "Rule 5(2)"),
    ("Sub-Committee", "Meets, considers public views and decides the new rates", "Rule 5(2)"),
    ("Sub-Committee", "Prepares statement of average rates, village-wise and local-body-wise", "Rule 5(2)"),
    ("Chairman + Secretary", "Sign the statement in CVC format, with views on objections", "Rule 5(2)"),
    ("Sub-Committee Secretary (SR)", "Sends printed booklet and soft copy to the District Registrar", "Rule 5(2)"),
    ("District Registrar", "Checks the statement; returns mistakes or missing data", "Rule 5(3)"),
    ("Sub-Committee", "Corrects and resubmits within 15 days", "Rule 5(3)"),
    ("District Registrar", "Final review with remarks; booklet per sub-district to CVC Secretary", "Rule 5(3)"),
    ("CVC Secretary", "Checks the statements and places them before the CVC", "Rule 7(1)"),
    ("CVC", "Discusses district-wise; accepts / rejects suggestions; records decisions", "Rule 7(2)"),
    ("CVC Secretary → DR", "Attests approved statements; DR forwards them to Sub-Committees", "Rule 7(3)"),
    ("Sub-Committee / DR", "Display approved rates in offices; printed copies sold at CVC price", "Rule 7(3)"),
    ("CVC + Kaveri", "Fix effective date; publish e-Gazette via DIPR / Karnataka Rajya Patra", "Kaveri system step"),
    ("Kaveri System", "New rates live from the effective date; old rates kept for history", "Kaveri system step"),
]

GR_STATUS = [
    ("Cycle Draft", "Revision Cycle created; instructions being prepared", "CVC Secretary"),
    ("Instructions Issued", "CVC instructions sent to Sub-Committees (general / mid-year)", "CVC"),
    ("Notice of Intention Published", "Plan to revise published in newspapers / notice boards", "Sub-Committee"),
    ("Objections Open", "15-day objection / suggestion window running", "Public / SC Secretary"),
    ("Objections Under Processing", "Secretary sorting objections for the agenda", "SC Secretary (SR)"),
    ("Sub-Committee Deliberation", "Sittings in progress; rates being decided", "Sub-Committee"),
    ("Statement Signed", "Statement signed by Secretary and Chairman", "Sub-Committee"),
    ("Submitted to DR", "Booklet and soft copy with District Registrar", "Sub-Committee / DR"),
    ("Returned for Rectification", "Mistakes returned; 15-day clock running", "Sub-Committee"),
    ("DR Final Examination", "DR recording views on improvement / change", "DR"),
    ("With CVC Secretary", "Per-sub-district booklets under verification", "CVC Secretary"),
    ("Placed before CVC", "District-wise estimations on CVC agenda", "CVC"),
    ("Approved", "Decision recorded; suggestions accepted / rejected", "CVC"),
    ("Attested and Transmitted", "Attested statements sent to DR → Sub-Committees", "CVC Secretary / DR"),
    ("Published", "Displayed in offices; printed copies on sale", "Sub-Committee / DR"),
    ("Gazette Linked", "Effective date fixed; e-Gazette particulars captured", "CVC / Admin"),
    ("Active in Kaveri", "Rates live from effective date", "System"),
    ("Returned / Clarification", "Returned to prior stage with remarks", "Prior owner"),
]

GLOSSARY_UPDATE = {
    "General Revision": "Cycle to estimate / revise market value guidelines for the next calendar year, triggered by CVC instructions to Sub-Committees (Rule 5(1)); includes mid-year revision for specific Sub-Committees.",
    "Sub-committee": "Market Valuation Sub-Committee (Rule 4) headed by the Tahsildar, with the Sub-Registrar as Member Secretary.",
    "Public opinion period": "Fifteen (15) days for public objections / suggestions after the Sub-Committee’s notice of intention (Rule 5(2)).",
}
GLOSSARY_ADD = [
    ("Notice of intention", "Sub-Committee notice in local newspapers and office notice boards announcing its plan to estimate / revise market values (Rule 5(2))."),
    ("Statement of average rates", "Signed Sub-Committee statement of average rates (agricultural land; residential, commercial, industrial sites), village-wise and local-body-wise; sent as booklet + soft copy."),
    ("Rectification period", "Fifteen (15) days from the District Registrar’s reference to correct and resubmit the statement (Rule 5(3))."),
    ("Attested statement", "CVC-approved statement attested by the CVC Secretary and sent to DR and Sub-Committees for publication (Rule 7(3))."),
]

RULE_ROWS = [
    ("CVC Rules 2003 — Rule 5", "CVC instructions; notice of intention; 15-day objections; signed statement of average rates (booklet + soft copy) to DR; DR verification and 15-day rectification"),
    ("CVC Rules 2003 — Rule 7", "CVC Secretary verification; CVC decision on Sub-Committee / DR suggestions; attestation; publication in offices; printed copies for sale"),
]
NOTIFICATION_OLD = "CVC / IGR order & DIGR notification"
NOTIFICATION_NEW = ("•  CVC instructions to Sub-Committees — Rule 5(1)  (per revision cycle: yearly or mid-year)  —  "
                    "Triggers General Revision; CVC approves statements (Rule 7); gazette fixes effective date")

ACTORS = {
    "IGR / Chairman CVC": "Chairs CVC; CVC instructions to Sub-Committees start General Revision (Rule 5(1)); approves final value for individual projects.",
    "DIGR (Valuation) / Secretary CVC": "CVC Secretary: checks DR statements and places them before CVC (Rule 7(1)); attests approved statements (Rule 7(3)); reviews individual-project proposals.",
    "District Registrar": "Checks Sub-Committee statements; returns mistakes for 15-day rectification; sends booklets to CVC Secretary; forwards attested statements; accepts Sec. 45-B requests.",
    "Sub-Registrar": "Sub-Committee Secretary: publishes notice of intention; sorts objections; prepares and co-signs the statement; sends booklet + soft copy to DR.",
    "Market Valuation Sub-committee": "Tahsildar (Chairman), SR (Secretary), Revenue, Survey, PWD, local bodies: considers objections, decides rates, signs statement, publishes approved rates.",
    "Central Valuation Committee": "Issues instructions (Rule 5(1)); discusses district-wise rates, accepts / rejects suggestions, records decisions (Rule 7(2)); fixes sale price and effective date.",
    "Citizen / Party / Developer": "Applies for individual project fixation (with ULPIN); files objections within 15 days of the notice of intention; may buy approved statements.",
}

TEXT_REPLACEMENTS = [
    # (startswith, new text)
    ("State-wide (or notified jurisdiction) periodic revision",
     "Estimation / revision of guidance values for the next calendar year (Rules 5 and 7), triggered by CVC "
     "instructions to Sub-Committees: notice of intention → 15-day objections → signed Sub-Committee statement → "
     "DR verification → CVC Secretary → CVC decision → attestation and publication → effective date / e-Gazette → "
     "Kaveri adoption."),
    ("Citizen / party-initiated request for a newly developed property",
     "Citizen / party-initiated request for a newly developed property (ULPIN given; boundary fetched from ULMS), "
     "accepted by the District Registrar under Sec. 45-B: DR inspection → Secretary CVC review → Chairman CVC (IGR) "
     "approval → Kaveri adoption."),
    ("Trigger: IGR as CVC Chairman issues an order",
     "Trigger: CVC (Chairman IGR) issues instructions and general policy guidelines to all Market Valuation "
     "Sub-Committees for the next calendar year (Rule 5(1)); CVC may also order any Sub-Committee to revise rates mid-year."),
    ("IGR · DIGR · DR · SR · SUB-COMMITTEE · CVC", "CVC · CVC SECRETARY · DR · SUB-COMMITTEE · SR"),
    ("IGR order → DIGR notification",
     "CVC instructions → notice of intention → 15-day objections → Secretary sorts objections → Sub-Committee "
     "decides rates → signed statement (booklet + soft copy) → DR check and 15-day rectification → CVC Secretary → "
     "CVC decision → attestation → publication and sale → effective date / e-Gazette → Kaveri adoption."),
    ("Historical registration (transaction) data, KSRSAC GIS",
     "Rule 6 factors with historical registration (transaction) data, KSRSAC GIS / maps, developer publications, "
     "RTC data, khata data, town planning inputs — available on the Sub-Committee workbench."),
    ("VIEW · OBJECT · TRACK", "VIEW · OBJECT · BUY · TRACK"),
    ("View proposed revised guideline values",
     "View the Sub-Committee’s notice of intention; file objections / suggestions within 15 days (online or at the "
     "SRO); view / download approved statements and buy printed copies; track cycle status."),
    ("Public objections are routed to the Sub-committee",
     "Objections go to the Sub-Committee Secretary and Sub-Committee; their views are recorded in the signed "
     "statement. The portal does not propose or approve rates."),
    ("Apply for fixation of guidance value for the newly developed property",
     "Apply for fixation of guidance value for the newly developed property with ULPIN (boundary fetched from ULMS), "
     "property particulars and project documents; track status online."),
    ("Party / citizen applies for fixation of guidance value (application",
     "Party / citizen applies (application, ULPIN, property particulars, project documents); boundary fetched from ULMS"),
    ("Order → notification → SR proposal",
     "CVC instructions → notice of intention → 15-day objections → signed statement → DR verification → CVC "
     "Secretary → CVC decision → attestation and publication → effective date → automatic adoption."),
    ("Objection capture & disposal", "Objection capture (15-day window)"),
    ("Sub-committee disposal limited to objected items",
     "Objections / suggestions captured online or at the SRO after the notice of intention; Sub-Committee views "
     "recorded in the signed statement."),
]

HIGHLIGHT_6 = ("Signed statements & DR verification",
               "eSign by Chairman and Secretary; booklet + soft copy; DR 15-day rectification tracking; CVC Secretary "
               "attestation before publication.")

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def set_text(shape_or_tf, text: str) -> None:
    tf = shape_or_tf.text_frame if hasattr(shape_or_tf, "text_frame") else shape_or_tf
    paras = tf.paragraphs
    first = paras[0]
    if first.runs:
        first.runs[0].text = text
        for r in first.runs[1:]:
            r.text = ""
    else:
        first.add_run().text = text
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def set_cell(cell, text: str) -> None:
    set_text(cell.text_frame, text)


def find_slide(prs, title: str, kicker: str | None = None):
    for slide in prs.slides:
        texts = [sh.text_frame.text.strip() for sh in slide.shapes if sh.has_text_frame]
        if title in texts and (kicker is None or any(t.startswith(kicker) for t in texts)):
            return slide
    raise SystemExit(f"Slide not found: {title}")


def add_textbox(slide, x, y, w, h, text, *, size, bold=False, italic=False, color=NAVY, font="Calibri",
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
    r.font.color.rgb = color
    return box


def add_shape(slide, kind, x, y, w, h, fill, line=None, dashed=False):
    shp = slide.shapes.add_shape(kind, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
        if dashed:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    return shp


def clear_body(slide, keep: int = 6) -> None:
    """Remove everything except the first `keep` chrome shapes (kicker, title, badge, footer)."""
    tree = slide.shapes._spTree
    for shp in list(slide.shapes)[keep:]:
        tree.remove(shp._element)


def move_slide(prs, slide, new_index: int) -> None:
    lst = prs.slides._sldIdLst
    el = next(i for i in list(lst) if prs.part.related_part(i.rId) is slide.part)
    lst.remove(el)
    lst.insert(new_index, el)


def new_slide_like(prs, template_slide, kicker: str, title: str, after_slide):
    slide = prs.slides.add_slide(template_slide.slide_layout)
    for shp in list(template_slide.shapes)[:6]:
        slide.shapes._spTree.append(copy.deepcopy(shp._element))
    shapes = list(slide.shapes)
    set_text(shapes[0], kicker)
    set_text(shapes[1], title)
    idx = list(prs.slides).index(after_slide)
    move_slide(prs, slide, idx + 1)
    return slide


def rebuild_table_rows(shape, rows: list[tuple], total_h=None) -> None:
    tbl = shape.table
    tbl_el = tbl._tbl
    templates = [copy.deepcopy(tbl.rows[i]._tr) for i in range(1, min(3, len(tbl.rows)))]
    for r in list(tbl.rows)[1:]:
        tbl_el.remove(r._tr)
    for i, values in enumerate(rows):
        tr = copy.deepcopy(templates[i % len(templates)])
        tbl_el.append(tr)
        row = tbl.rows[len(tbl.rows) - 1]
        for c, v in enumerate(values):
            set_cell(row.cells[c], v)
    if total_h is not None:
        h = int(total_h / len(tbl.rows))
        for r in tbl.rows:
            r.height = h


def renumber(prs) -> None:
    for i, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if not shp.has_text_frame or not re.fullmatch(r"\d{1,2}", shp.text_frame.text.strip()):
                continue
            if shp.left >= Inches(11.5) and (shp.top < Inches(1.2) or shp.top > Inches(6.8)):
                set_text(shp, str(i))


def bump_version(prs) -> None:
    for slide in prs.slides:
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            for p in shp.text_frame.paragraphs:
                if "BRD-K3-CVC-GVF-001" not in p.text:
                    continue
                if "Last updated" in p.text:
                    new = (f"Document ID: BRD-K3-CVC-GVF-001     ·     Version {VERSION}     ·     "
                           f"Status: Draft     ·     Last updated {DATE}")
                else:
                    new = f"Document ID: BRD-K3-CVC-GVF-001  ·  Version {VERSION}  ·  {DATE}"
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""


# ---------------------------------------------------------------------------
# Slide updates
# ---------------------------------------------------------------------------


def apply_text_replacements(prs) -> None:
    for slide in prs.slides:
        for shp in slide.shapes:
            if not shp.has_text_frame:
                continue
            t = shp.text_frame.text.strip()
            for start, new in TEXT_REPLACEMENTS:
                if t.startswith(start):
                    set_text(shp, new)
                    break


def update_legal(prs) -> None:
    slide = find_slide(prs, "Karnataka Stamp Act, 1957 — Key Sections")
    tables = [s for s in slide.shapes if s.has_table]
    rules = next(t for t in tables if t.table.cell(0, 0).text.strip().startswith("Rule"))
    existing = [(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in list(rules.table.rows)[1:]]
    rebuild_table_rows(rules, existing + RULE_ROWS)
    row_h = Inches(0.42)
    for r in rules.table.rows:
        r.height = row_h
    bottom = rules.top + row_h * len(rules.table.rows)
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip().startswith("Related notifications"):
            shp.top = bottom + Inches(0.12)
            for p in shp.text_frame.paragraphs:
                if NOTIFICATION_OLD in p.text:
                    p.runs[0].text = NOTIFICATION_NEW
                    for r in p.runs[1:]:
                        r.text = ""


def update_glossary(prs) -> None:
    slide = find_slide(prs, "Definitions and Glossary")
    shp = next(s for s in slide.shapes if s.has_table)
    rows = [(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in list(shp.table.rows)[1:]]
    rows = [(k, GLOSSARY_UPDATE.get(k, v)) for k, v in rows]
    rows += [r for r in GLOSSARY_ADD if r[0] not in {k for k, _ in rows}]
    rebuild_table_rows(shp, rows, total_h=Inches(5.1))


def update_actors(prs) -> None:
    slide = find_slide(prs, "Stakeholders and Actors")
    shapes = [s for s in slide.shapes if s.has_text_frame]
    for i, shp in enumerate(shapes):
        name = shp.text_frame.text.strip()
        if name in ACTORS and i + 1 < len(shapes):
            set_text(shapes[i + 1], ACTORS[name])


def build_gr_flow(prs) -> None:
    slide = find_slide(prs, "Process Flow — General Revision")
    set_text(list(slide.shapes)[1], "Process Flow — General Revision (15 Stages + Kaveri Steps)")
    clear_body(slide)
    per_row, box_w, box_h, pitch = 6, Inches(1.86), Inches(1.22), Inches(2.09)
    x0, y0, row_gap = Inches(0.5), Inches(1.72), Inches(1.62)
    for i, (actor, text, ref) in enumerate(GR_FLOW):
        row, col = divmod(i, per_row)
        x, y = x0 + pitch * col, y0 + row_gap * row
        system = ref.startswith("Kaveri")
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h,
                  SYS_BG if system else ALT, line=MUTED if system else None, dashed=system).adjustments[0] = 0.08
        badge = add_shape(slide, MSO_SHAPE.OVAL, x + (box_w - Inches(0.32)) // 2, y - Inches(0.16),
                          Inches(0.32), Inches(0.32), MUTED if system else GOLD)
        add_textbox(slide, badge.left, badge.top, badge.width, badge.height, str(i + 1), size=10.5, bold=True,
                    color=WHITE, font="Cambria", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        box = add_textbox(slide, x + Inches(0.07), y + Inches(0.2), box_w - Inches(0.14), box_h - Inches(0.26),
                          actor, size=8.5, bold=True, color=GOLD if not system else MUTED, align=PP_ALIGN.CENTER)
        tf = box.text_frame
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size, r.font.name, r.font.color.rgb = Pt(8.2), "Calibri", NAVY
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ref
        r.font.size, r.font.italic, r.font.name, r.font.color.rgb = Pt(7.5), True, "Calibri", MUTED
        if col < per_row - 1 and i < len(GR_FLOW) - 1:
            add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + box_w + Inches(0.0), y + box_h // 2 - Inches(0.08),
                      Inches(0.23), Inches(0.16), GOLD)
    add_textbox(slide, Inches(0.5), Inches(6.62), Inches(12.3), Inches(0.3),
                "Rows read left to right. Step 9 returns mistakes to step 10 (15-day rectification); step 13 can return "
                "to step 12 for clarification. Grey dashed boxes = Kaveri system steps (not in the Rule 5 / 7 text).",
                size=9, italic=True, color=MUTED)


def split_process_a_png(tmp: Path) -> tuple[Path, Path]:
    im = Image.open(DIAGRAMS / "Process_A_General_Revision.png")
    s = PNG_SCALE
    split = (A_PNG_BORDER + A_FIRST_ROW + A_SPLIT_ROW * A_ROW_H) * s
    header = im.crop((0, (A_PNG_BORDER + A_LANE_TOP) * s, im.width, (A_PNG_BORDER + A_LANE_HEADER_END) * s))
    part1 = im.crop((0, 0, im.width, split))
    lower = im.crop((0, split, im.width, im.height))
    part2 = Image.new("RGB", (im.width, header.height + lower.height), "white")
    part2.paste(header, (0, 0))
    part2.paste(lower, (0, header.height))
    p1, p2 = tmp / "process_a_part1.png", tmp / "process_a_part2.png"
    part1.save(p1)
    part2.save(p2)
    return p1, p2


def add_picture_fit(slide, path: Path, x, y, max_w, max_h):
    with Image.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pic = slide.shapes.add_picture(str(path), x, y, int(w * scale), int(h * scale))
    pic.line.color.rgb = RGBColor(0xC8, 0xD0, 0xDC)
    pic.line.width = Pt(0.75)
    return pic


def notes_card(slide, x, y, w, h, title: str, lines: list[str]) -> None:
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, ALT).adjustments[0] = 0.04
    box = add_textbox(slide, x + Inches(0.15), y + Inches(0.12), w - Inches(0.3), h - Inches(0.24),
                      title, size=12, bold=True, font="Cambria")
    for line in lines:
        p = box.text_frame.add_paragraph()
        p.space_before = Pt(5)
        r = p.add_run()
        r.text = line
        r.font.size, r.font.name, r.font.color.rgb = Pt(9.5), "Calibri", NAVY


def add_diagram_slides(prs, tmp: Path) -> None:
    flow_a = find_slide(prs, "Process Flow — General Revision (15 Stages + Kaveri Steps)")
    slide = new_slide_like(prs, flow_a, "7.I GENERAL REVISION", "Swimlane Process Diagram — General Revision", flow_a)
    add_textbox(slide, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.3),
                "One lane per actor; flow runs top to bottom. Full-size: BRD v1.8 Figure 1  ·  "
                "Process Diagram/Process_A_General_Revision.drawio",
                size=10, italic=True, color=MUTED)
    p1, p2 = split_process_a_png(tmp)
    top, max_h = Inches(1.95), Inches(4.95)
    add_textbox(slide, Inches(0.5), Inches(1.72), Inches(4.8), Inches(0.22), "Part 1 — Stages 1 to 11",
                size=9.5, bold=True, color=GOLD)
    pic1 = add_picture_fit(slide, p1, Inches(0.5), top, Inches(4.9), max_h)
    x2 = pic1.left + pic1.width + Inches(0.2)
    add_textbox(slide, x2, Inches(1.72), Inches(4.8), Inches(0.22), "Part 2 — Stages 12 to 15 and Kaveri steps",
                size=9.5, bold=True, color=GOLD)
    pic2 = add_picture_fit(slide, p2, x2, top, Inches(4.9), max_h)
    nx = pic2.left + pic2.width + Inches(0.2)
    notes_card(slide, nx, Inches(1.72), Inches(12.83) - nx, Inches(5.18), "Lanes", [
        "CVC (Chairman: IGR)",
        "CVC Secretary — DIGR (Valuation)",
        "District Registrar",
        "Market Valuation Sub-Committee (Chairman: Tahsildar)",
        "Sub-Committee Secretary — SR",
        "Public / Citizens",
        "Kaveri System",
        "Red dashed arrows = return paths (rectification; CVC clarification).",
        "Grey dashed boxes = Kaveri system steps.",
    ])

    flow_b = find_slide(prs, "Process Flow — Individual Project Fixation (6 Steps)")
    slide = new_slide_like(prs, flow_b, "7.II INDIVIDUAL PROJECT FIXATION",
                           "Swimlane Process Diagram — Individual Project Fixation", flow_b)
    add_textbox(slide, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.3),
                "One lane per actor; flow runs top to bottom. Full-size: BRD v1.8 Figure 2  ·  "
                "Process Diagram/Process_B_Individual_Project_Fixation.drawio",
                size=10, italic=True, color=MUTED)
    pic = add_picture_fit(slide, DIAGRAMS / "Process_B_Individual_Project_Fixation.png",
                          Inches(0.5), Inches(1.8), Inches(7.0), Inches(5.15))
    nx = pic.left + pic.width + Inches(0.3)
    notes_card(slide, nx, Inches(1.8), Inches(12.83) - nx, Inches(5.1), "Reading the diagram", [
        "Lanes: Citizen / Developer · Kaveri System (with ULMS) · District Registrar · CVC Secretary · IGR (Chairman, CVC).",
        "Kaveri fetches the property boundary from ULMS using the ULPIN.",
        "DR check has three outcomes: accepted → site inspection; needs changes → applicant fixes and resubmits; "
        "not eligible → rejected with written reasons.",
        "CVC Secretary can return the proposal to the DR with remarks; IGR can return it to the CVC Secretary.",
        "On IGR approval the start date is fixed and the rate goes live in Kaveri; the applicant is informed.",
    ])


def style_cell(cell, text: str, *, size=9.5, bold=False, color=NAVY, fill=None, align=PP_ALIGN.LEFT) -> None:
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.04)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, "Calibri", color
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def tree_box(slide, x, y, w, h, text, fill, color=NAVY, bold=False, size=9):
    box = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    box.adjustments[0] = 0.15
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, "Calibri", color
    return box


def add_area_master_slide(prs) -> None:
    paths = find_slide(prs, "Two Guidance Value Fixation Paths in Kaveri 3.0")
    slide = new_slide_like(prs, paths, "FUTURE STATE (TO-BE)",
                           "Details Captured for Each Valuation Area", paths)
    add_textbox(slide, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.3),
                "For every valuation area, Kaveri 3.0 records the details below. "
                "They are used by both General Revision and Individual Project Fixation.",
                size=10, italic=True, color=MUTED)

    tx, ty, tw = Inches(0.5), Inches(1.85), Inches(8.55)
    shp = slide.shapes.add_table(len(AREA_DETAILS) + 1, 2, tx, ty, tw, Inches(4.9))
    tbl = shp.table
    tbl.first_row = tbl.horz_banding = False
    for c, w in enumerate((Inches(2.3), Inches(6.25))):
        tbl.columns[c].width = w
    tbl.rows[0].height = Inches(0.45)
    for c, head in enumerate(("Detail", "What is captured")):
        style_cell(tbl.cell(0, c), head, size=11, bold=True, color=WHITE, fill=NAVY)
    for i, (detail, text) in enumerate(AREA_DETAILS, 1):
        fill = ALT if i % 2 == 0 else WHITE
        tbl.rows[i].height = Inches(0.89)
        style_cell(tbl.cell(i, 0), detail, size=11, bold=True, fill=fill)
        style_cell(tbl.cell(i, 1), text, size=11, fill=fill)

    cx, cy, cw, ch = Inches(9.3), Inches(1.85), Inches(3.53), Inches(4.9)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch, ALT).adjustments[0] = 0.04
    add_textbox(slide, cx + Inches(0.15), cy + Inches(0.12), cw - Inches(0.3), Inches(0.3),
                "At a glance", size=12, bold=True, font="Cambria")

    root_w = cw - Inches(0.3)
    tree_box(slide, cx + Inches(0.15), cy + Inches(0.5), root_w, Inches(0.42),
             "Valuation area", NAVY, WHITE, bold=True, size=10)
    tree_box(slide, cx + Inches(0.15), cy + Inches(1.02), root_w, Inches(0.36),
             "Area / extent in m²", WHITE)
    tree_box(slide, cx + Inches(0.15), cy + Inches(1.46), root_w, Inches(0.36),
             "Geo-fence (boundary on map)", WHITE)

    add_textbox(slide, cx + Inches(0.15), cy + Inches(1.95), root_w, Inches(0.22),
                "Guidance rates by category", size=9.5, bold=True, color=GOLD)
    col_w = (root_w - Inches(0.15)) // 2
    branches = (("Non-Agricultural", ("Residential", "Commercial", "Industrial")),
                ("Agricultural", ("Dry", "Wet", "Bhagayat (garden)")))
    for b, (head, leaves) in enumerate(branches):
        bx = cx + Inches(0.15) + b * (col_w + Inches(0.15))
        tree_box(slide, bx, cy + Inches(2.22), col_w, Inches(0.5), head, GOLD, WHITE, bold=True)
        for j, leaf in enumerate(leaves):
            tree_box(slide, bx + Inches(0.12), cy + Inches(2.85) + j * Inches(0.44),
                     col_w - Inches(0.12), Inches(0.36), leaf, WHITE)
    add_textbox(slide, cx + Inches(0.15), cy + Inches(4.2), root_w, Inches(0.6),
                "Rate categories and types are maintained by CVC / DIGR Admin as masters, "
                "with English and Kannada labels, so new types can be added later.",
                size=8.5, italic=True, color=MUTED)


def update_status_model(prs) -> None:
    slide = find_slide(prs, "Application / Cycle Status Model", "7.I")
    tables = sorted([s for s in slide.shapes if s.has_table], key=lambda s: s.left)
    half = (len(GR_STATUS) + 1) // 2
    for shp, rows in zip(tables, (GR_STATUS[:half], GR_STATUS[half:])):
        rebuild_table_rows(shp, rows, total_h=Inches(4.9))


def update_highlights(prs) -> None:
    slide = find_slide(prs, "Highlights in Kaveri 3.0")
    shapes = list(slide.shapes)
    rects = sorted([s for s in shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE],
                   key=lambda s: (s.top, s.left))
    if len(rects) >= 6:
        return
    card1, card2, card5 = rects[0], rects[1], rects[-1]
    dx = card2.left - card1.left
    members = [s for s in shapes if card5.left <= s.left < card5.left + card5.width
               and card5.top <= s.top < card5.top + card5.height]
    new = []
    for s in members:
        el = copy.deepcopy(s._element)
        slide.shapes._spTree.append(el)
        new.append(slide.shapes[-1])
    for s in new:
        s.left = s.left + dx
    texts = sorted([s for s in new if s.has_text_frame and s.text_frame.text.strip()], key=lambda s: (s.top, s.left))
    for s in texts:
        t = s.text_frame.text.strip()
        if t == "5":
            set_text(s, "6")
        elif t.startswith("Objection capture"):
            set_text(s, HIGHLIGHT_6[0])
        else:
            set_text(s, HIGHLIGHT_6[1])


def main():
    prs = Presentation(str(SRC))
    bump_version(prs)
    apply_text_replacements(prs)
    update_legal(prs)
    update_glossary(prs)
    update_actors(prs)
    build_gr_flow(prs)
    update_status_model(prs)
    update_highlights(prs)
    with tempfile.TemporaryDirectory() as tmp:
        add_diagram_slides(prs, Path(tmp))
        add_area_master_slide(prs)
        renumber(prs)
        try:
            prs.save(str(DST))
            out = DST
        except PermissionError:
            out = DST.with_name(DST.stem + "_updated.pptx")
            prs.save(str(out))
    print(f"Wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
