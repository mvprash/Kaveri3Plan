# -*- coding: utf-8 -*-
"""Merge Marriage BRD v3 + FRS/NFRS v2.0 into one combined Word document."""
from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\Marriage\RFP")
BRD = BASE / "BRD_Marriage_BRD_v3.docx"
FRS = BASE / "FRS-NFR" / "FRS_and_NFRS_Marriage_v2.0.docx"
DST = BASE / "BRD_and_FRS_NFRS_Marriage_Combined.docx"

COMBINED_TITLE = (
    "Business Requirements Document (BRD), Functional Requirements Specification "
    "(FRS) and Non-Functional Requirements (NFRs)"
)
COMBINED_HEADER = (
    "Kaveri 3.0 | BRD + FRS & NFR | Marriage Registration Module"
)


def set_para_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""


def set_cell_text(cell: _Cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        cell.add_paragraph(text)
        return
    set_para_text(paras[0], text)
    for p in paras[1:]:
        set_para_text(p, "")


def add_version_row(table, values: list[str]) -> None:
    table._tbl.append(deepcopy(table.rows[-1]._tr))
    row = table.rows[-1]
    for ci, val in enumerate(values):
        if ci < len(row.cells):
            set_cell_text(row.cells[ci], val)


def set_all_headers(doc: Document, text: str) -> None:
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            for p in header.paragraphs:
                if p.text.strip():
                    set_para_text(p, text)
                    break
            else:
                if header.paragraphs:
                    set_para_text(header.paragraphs[0], text)


def update_front_matter(doc: Document) -> None:
    # Title (first Title / BRD title para)
    for p in doc.paragraphs[:6]:
        t = p.text.strip()
        if "Business Requirements Document" in t or (p.style and p.style.name == "Title"):
            set_para_text(p, COMBINED_TITLE)
            break

    # Document control — BRD table remains primary; mark as combined package
    ctrl = doc.tables[0]
    set_cell_text(ctrl.rows[1].cells[1], "BRD-FRS-NFRS-K3-MRG-001")
    set_cell_text(ctrl.rows[2].cells[1], "3 / FRS-NFR 2.0")
    set_cell_text(
        ctrl.rows[3].cells[1],
        "Draft / In review — Combined BRD (v3) + FRS and NFRs (v2.0)",
    )

    hist = doc.tables[1]
    last = hist.rows[-1].cells[3].text.strip()
    if "Combined BRD" not in last:
        add_version_row(
            hist,
            [
                "3+2.0",
                "09-09-2026",
                "Nandha Kumar",
                "Combined BRD_Marriage_BRD_v3.docx and FRS_and_NFRS_Marriage_v2.0.docx into one RFP package document",
                "Prashanth",
            ],
        )

    # Related documents — this file is the combined package
    if len(doc.tables) >= 3:
        rel = doc.tables[2]
        if len(rel.rows) >= 2:
            set_cell_text(rel.rows[1].cells[0], "BRD-FRS-NFRS-K3-MRG-001")
            set_cell_text(rel.rows[1].cells[1], "This document (combined BRD + FRS and NFRs)")
            set_cell_text(rel.rows[1].cells[2], DST.name)
        # Keep / refresh FRS companion row as source reference
        for row in rel.rows[1:]:
            joined = " ".join(c.text for c in row.cells).upper()
            if "FRS" in joined and "THIS DOCUMENT" not in row.cells[1].text.upper():
                set_cell_text(row.cells[0], "FRS-NFRS-K3-MRG-001")
                set_cell_text(row.cells[1], "Source FRS and NFRs (also appended below)")
                set_cell_text(row.cells[2], "FRS-NFR/FRS_and_NFRS_Marriage_v2.0.docx")
                break

    set_all_headers(doc, COMBINED_HEADER)


def main() -> None:
    if not BRD.exists():
        raise FileNotFoundError(BRD)
    if not FRS.exists():
        raise FileNotFoundError(FRS)

    master = Document(str(BRD))
    # Ensure FRS starts on a new page
    master.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    composer = Composer(master)
    composer.append(Document(str(FRS)))

    # Post-merge front-matter updates on the combined document
    update_front_matter(composer.doc)
    composer.save(str(DST))

    # Verify
    out = Document(str(DST))
    headings = [
        p.text.strip()
        for p in out.paragraphs
        if p.style and (p.style.name.startswith("Heading") or p.style.name == "Title") and p.text.strip()
    ]
    blips = list(out.element.body.iter(qn("a:blip")))
    print(f"Wrote: {DST}")
    print(f"  Size: {DST.stat().st_size:,} bytes")
    print(f"  Paragraphs: {len(out.paragraphs)}  Tables: {len(out.tables)}  Images: {len(blips)}")
    print(f"  Title: {headings[0][:90] if headings else '(none)'}")
    print(f"  First headings: {headings[:4]}")
    print(f"  Has FRS heading: {any('Functional Requirements Specification' in h for h in headings)}")
    print(f"  Has FR §1: {any(h.startswith('1. Functional') for h in headings)}")
    print(f"  Has NFR §8: {any(h.startswith('8. Non-functional') for h in headings)}")
    print(f"  Has Appendix: {any(h.startswith('12. Appendix') or h.startswith('Appendix') for h in headings)}")
    print(f"  Doc ID: {out.tables[0].rows[1].cells[1].text.strip()}")
    print(f"  Version: {out.tables[0].rows[2].cells[1].text.strip()}")
    print(f"  Status: {out.tables[0].rows[3].cells[1].text.strip()}")


if __name__ == "__main__":
    main()
