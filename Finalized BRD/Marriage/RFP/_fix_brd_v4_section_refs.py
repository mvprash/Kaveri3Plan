# -*- coding: utf-8 -*-
"""Retarget BRD_Marriage_BRD_v4 section cites to the hierarchical outline → v6.

Focus: §3 Legal mapping tables + §6 As-Is pain points (and consistent cites elsewhere).

1. Fix Word heading list (hierarchical %1.%2.%3. + ilvl from Heading style).
2. Strip leftover manual FRS numbers from heading text.
3. Replace old roman (7.i…) / companion-FRS cites with current outline numbers.
4. Rewrite narrative that still treats FRS as a separate document.
"""
from __future__ import annotations

import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Finalized BRD\Marriage\RFP")
SRC = BASE / "BRD_Marriage_BRD_v4.docx"
DST = BASE / "BRD_Marriage_BRD_v6.docx"
OUT_VERSION = "6"

# ---------------------------------------------------------------------------
# Cite replacement map (longest / most-specific first)
# ---------------------------------------------------------------------------
# Old roman To-Be → hierarchical §7.*
ROMAN_REPLACEMENTS: list[tuple[str, str]] = [
    ("7.vi.a", "7.6.1"),
    ("7.vi", "7.6"),
    ("7.iv.b.D", "7.4.2.4"),
    ("7.iv.b.C", "7.4.2.3"),
    ("7.iv.b.B", "7.4.2.2"),
    ("7.iv.b.A", "7.4.2.1"),
    ("7.iv.b", "7.4.2"),
    ("7.iv.a", "7.4.1"),
    ("7.iv", "7.4"),
    ("7.iii.b.C", "7.3.2.3"),
    ("7.iii.b.B", "7.3.2.2"),
    ("7.iii.b.A", "7.3.2.1"),
    ("7.iii.b.x", "7.3.2"),
    ("7.iii.b", "7.3.2"),
    ("7.iii.a", "7.3.1"),
    ("7.iii", "7.3"),
    ("7.ii.b.D", "7.2.2.4"),
    ("7.ii.b.C", "7.2.2.3"),
    ("7.ii.b.B", "7.2.2.2"),
    ("7.ii.b.A", "7.2.2.1"),
    ("7.ii.b.x", "7.2.2"),
    ("7.ii.b", "7.2.2"),
    ("7.ii.a", "7.2.1"),
    ("7.ii", "7.2"),
    ("7.i.b.D", "7.1.2.4"),
    ("7.i.b.C", "7.1.2.3"),
    ("7.i.b.B", "7.1.2.2"),
    ("7.i.b.A", "7.1.2.1"),
    ("7.i.b", "7.1.2"),
    ("7.i.a", "7.1.1"),
    ("7.i–7.v", "7.1–7.5"),
    ("7.i-7.v", "7.1–7.5"),
    ("7.i", "7.1"),
    ("7.v", "7.5"),
    ("7.b.x", "7.1.2"),
]

# Null-and-void used decimal 7.4 in FR notes; under new outline that is Parsi (7.4).
# After roman→decimal, remaining Null/Void "7.4" tokens must become 7.5.
# Match bare 7.4 only (not 7.4.2 / 7.45). Parsi owns 7.4.*; Null/Void → 7.5.
NULL_VOID_SEVEN_FOUR = re.compile(r"(?<!\d)7\.4(?!\.\d)(?!\d)")

# Exact full-string / phrase replacements (applied before generic maps)
PHRASE_REPLACEMENTS: list[tuple[str, str]] = [
    (
        "see 8.7 and 11. (FRS and NFR document)",
        "see 8.8 (Sakala integration)",
    ),
    (
        "Pain points evidenced from Kaveri 2.0 workshops, ServiceDesk tickets and department discussions. "
        "The Addressed in column maps each item to the To-Be process in this BRD (7) and to functional / "
        "non-functional requirements, fallbacks and risks in the companion FRS and NFRs document "
        "(FRS_and_NFRS_Marriage_v1.22.docx — 8–18).",
        "Pain points evidenced from Kaveri 2.0 workshops, ServiceDesk tickets and department discussions. "
        "The Addressed in column maps each item to the To-Be process (§7) and to functional / "
        "non-functional requirements, fallbacks and risks in this document (§8–§11).",
    ),
    (
        "This section summarises material enhancements in Kaveri 3.0 compared with the legacy Kaveri 2.0 "
        "Marriage Registration module (6). Capability highlights are listed below; 7.vi.a maps each As-Is "
        "pain point from 6.1 to the Kaveri 3.0 closure. Cross-references point to To-Be process in this BRD "
        "(7.i–7.v) and to functional / non-functional requirements in FRS_and_NFRS_Marriage_v1.22.docx.",
        "This section summarises material enhancements in Kaveri 3.0 compared with the legacy Kaveri 2.0 "
        "Marriage Registration module (§6). Capability highlights are listed below; 7.6.1 maps each As-Is "
        "pain point from 6.1 to the Kaveri 3.0 closure. Cross-references point to To-Be process in this BRD "
        "(§7.1–§7.5) and to functional / non-functional requirements in §§8–9 of this document.",
    ),
    (
        "Functional requirements are organized by service, aligned with 7 (To-Be). Section 7 is the process "
        "authority — channel models, diagrams, step tables and status models. Section 8 states testable "
        "functional requirements (FR-HMA-* / FR-SMA-*) and does not restate those process steps. Where a "
        "requirement implements a To-Be step, the FR cites the 7 reference. Cross-cutting post-registration, "
        "notification and MIS requirements follow the service-specific blocks.",
        "Functional requirements are organized by service, aligned with §7 (To-Be). Section 7 is the process "
        "authority — channel models, diagrams, step tables and status models. Section 8 states testable "
        "functional requirements (FR-HMA-* / FR-SMA-* / FR-PMA-*) and does not restate those process steps. "
        "Where a requirement implements a To-Be step, the FR cites the §7 reference. Cross-cutting "
        "post-registration, notification and MIS requirements follow the service-specific blocks.",
    ),
    ("(Ref: 7.4)", "(Ref: 7.5)"),  # Null and Void FR heading — before generic maps
    ("see 1.7", "see 8.7"),
    ("in 1.2.7", "in 8.2.7"),
    ("of 1.3.2", "of 8.3.2"),
    ("Addressed in (BRD / FRS·NFR)", "Addressed in (this BRD)"),
]

# Pain-point / RTM style FRS·NFR section retargets (after roman map).
# Apply longest / most-specific first. Old combined-FRS numbers → current §§8–11.
FRS_SECTION_REPLACEMENTS: list[tuple[str, str]] = [
    ("FRS/NFRS: 8.1.3–8.1.5", "8.1.3–8.1.5"),
    ("FRS/NFRS: 8.1.4–8.1.5", "8.1.4–8.1.5"),
    ("FRS/NFRS: 8.2.5–8.2.7", "8.2.5–8.2.7"),
    ("FRS/NFRS: 8.1.16", "8.1.16"),
    ("FRS/NFRS: 8.1.15", "8.1.15"),
    ("FRS/NFRS: 8.1.14", "8.1.14"),
    ("FRS/NFRS: 8.1.13", "8.1.13"),
    ("FRS/NFRS: 8.1.11", "8.1.11"),
    ("FRS/NFRS: 8.1.10", "8.1.10"),
    ("FRS/NFRS: 8.1.9", "8.1.9"),
    ("FRS/NFRS: 8.1.6", "8.1.6"),
    ("FRS/NFRS: 8.1.3", "8.1.3"),
    ("FRS/NFRS: 8.1.2", "8.1.2"),
    ("FRS/NFRS: 8.2.8", "8.2.8"),
    ("FRS/NFRS: 10 UI", "8.11 UI"),
    ("FRS/NFRS: 17 FB-MRG-003", "11 FB-MRG-003"),
    ("FRS/NFRS: 8.5", "8.6"),  # notifications (shifted: old 8.5 → new 8.6)
    ("FRS/NFRS: 8.6", "8.7"),  # reports / MIS (shifted)
    ("FRS/NFRS: 8.4", "8.5"),  # post-registration (shifted)
    ("15.4 NFR-MRG-VAPT-002", "9.4 NFR-MRG-VAPT-002"),
    ("11 Integrations", "8.12 Integrations"),
    ("12.1 Core entities", "8.13.1 Core entities"),
    ("16 RS-MRG-003", "10 RS-MRG-003"),
    ("17 FB-MRG-001", "11 FB-MRG-001"),
    ("8.3.iii", "8.3.3"),
    # Bare "8.6 FR-HMA-042 (cycle-time MIS)" in same column — old reports §
    ("8.6 FR-HMA-042", "8.7 FR-HMA-042"),
]

HEADING_STRIP = [
    (re.compile(r"^1\.2\.1\s+"), ""),
    (re.compile(r"^1\.4\.1\s+"), ""),
    (re.compile(r"^8\.2\s+"), ""),
    (re.compile(r"^8\.3\s+"), ""),
    (re.compile(r"^8\.4\s+"), ""),
]


def set_para_text(paragraph: Paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for r in paragraph.runs[1:]:
        r.text = ""


def set_cell_text(cell: _Cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        cell.add_paragraph(text)
        return
    set_para_text(paras[0], text)
    for p in paras[1:]:
        set_para_text(p, "")


def add_version_row(table, values: list[str]) -> None:
    table._tbl.append(deepcopy(table.rows[-1]._tr))
    row = table.rows[-1]
    for ci, val in enumerate(values):
        if ci < len(row.cells):
            set_cell_text(row.cells[ci], val)


# Bare old-FRS section numbers (1.2.7 → 8.2.7). Do not run on TLS / Kaveri 1.x version text.
FRS_ONE_PREFIX = re.compile(r"(?<![0-9.])1\.(\d+(?:\.\d+)*)")


def retarget_old_frs_one_prefix(text: str) -> str:
    if not text:
        return text
    if "TLS" in text or re.search(r"Kaveri\s+1\.", text):
        return text
    # Only rewrite when the cell/paragraph looks like an FRS section cite
    if not re.search(r"(?<![0-9.])1\.\d+", text):
        return text
    # Skip if it's only embedded inside already-correct 7.x/8.x (no bare 1. at token start)
    if not re.search(r"(?:^|[\s;/,(])1\.\d+", text):
        return text
    return FRS_ONE_PREFIX.sub(r"8.\1", text)


def rewrite_text(text: str, *, null_void_context: bool = False) -> str:
    if not text:
        return text
    out = text
    for old, new in PHRASE_REPLACEMENTS:
        out = out.replace(old, new)
    for old, new in ROMAN_REPLACEMENTS:
        out = out.replace(old, new)
    for old, new in FRS_SECTION_REPLACEMENTS:
        out = out.replace(old, new)
    # Drop leftover "FRS/NFRS:" prefix if any remain
    out = out.replace("FRS/NFRS: ", "")
    out = out.replace("FRS_and_NFRS_Marriage_v1.22.docx", "this document (§§8–11)")
    # Null/Void decimal 7.4 → 7.5 (Parsi now owns 7.4)
    if null_void_context:
        out = NULL_VOID_SEVEN_FOUR.sub("7.5", out)
    # Old separate-FRS §1.x → combined §8.x
    out = retarget_old_frs_one_prefix(out)
    return out


def is_null_void_context(text: str) -> bool:
    t = text.lower()
    return any(
        k in t
        for k in (
            "null",
            "void",
            "nullity",
            "court order",
            "fr-hma-094",
            "sec. 11",
            "sec. 12",
            "7.4 step",
            "7.4 steps",
            "; 7.4",
            "7.4;",
            "(ref: 7.4)",
        )
    )


def fix_heading_list_definition(doc: Document) -> None:
    """Make abstractNum for heading list hierarchical (%1.%2.%3.)."""
    root = doc.part.numbering_part._element
    abs_id = None
    for num in root.findall(qn("w:num")):
        if num.get(qn("w:numId")) == "7":
            abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            break
    if abs_id is None:
        raise RuntimeError("numId 7 not found")
    for absn in root.findall(qn("w:abstractNum")):
        if absn.get(qn("w:abstractNumId")) != abs_id:
            continue
        for lvl in absn.findall(qn("w:lvl")):
            ilvl = int(lvl.get(qn("w:ilvl")))
            lt = lvl.find(qn("w:lvlText"))
            if lt is None:
                continue
            # %1. / %1.%2. / %1.%2.%3. ...
            parts = "".join(f"%{i}." for i in range(1, ilvl + 2))
            # trim trailing style: Word often wants "1.2.3" without final weirdness;
            # keep trailing period after last number to match prior "%1." style → "1.2.3."
            lt.set(qn("w:val"), parts)
        break


def fix_heading_ilvls(doc: Document) -> int:
    """Set numPr ilvl from Heading N style (H2→0 … H5→3)."""
    fixed = 0
    style_to_ilvl = {
        "Heading 2": "0",
        "Heading 3": "1",
        "Heading 4": "2",
        "Heading 5": "3",
        "Heading 6": "4",
    }
    for p in doc.paragraphs:
        if not p.style or p.style.name not in style_to_ilvl:
            continue
        if p.text.strip() == "Document control":
            continue  # keep unnumbered
        pPr = p._p.get_or_add_pPr()
        numPr = pPr.find(qn("w:numPr"))
        if numPr is None:
            numPr = OxmlElement("w:numPr")
            ilvl_el = OxmlElement("w:ilvl")
            ilvl_el.set(qn("w:val"), style_to_ilvl[p.style.name])
            numId_el = OxmlElement("w:numId")
            numId_el.set(qn("w:val"), "7")
            numPr.append(ilvl_el)
            numPr.append(numId_el)
            pPr.append(numPr)
            fixed += 1
            continue
        ilvl_el = numPr.find(qn("w:ilvl"))
        want = style_to_ilvl[p.style.name]
        if ilvl_el is None:
            ilvl_el = OxmlElement("w:ilvl")
            ilvl_el.set(qn("w:val"), want)
            numPr.insert(0, ilvl_el)
            fixed += 1
        elif ilvl_el.get(qn("w:val")) != want:
            ilvl_el.set(qn("w:val"), want)
            fixed += 1
        numId_el = numPr.find(qn("w:numId"))
        if numId_el is not None and numId_el.get(qn("w:val")) != "7":
            numId_el.set(qn("w:val"), "7")
    return fixed


def strip_heading_manual_numbers(doc: Document) -> int:
    n = 0
    for p in doc.paragraphs:
        if not (p.style and p.style.name.startswith("Heading") and p.text.strip()):
            continue
        text = p.text
        new = text
        for rx, repl in HEADING_STRIP:
            new = rx.sub(repl, new)
        if new != text:
            set_para_text(p, new)
            n += 1
    return n


def replace_in_paragraph(p: Paragraph) -> bool:
    full = p.text
    if not full.strip():
        return False
    # Skip rewriting heading titles except leftover strips (handled separately)
    if p.style and p.style.name.startswith("Heading"):
        return False
    new = rewrite_text(full, null_void_context=is_null_void_context(full))
    if new == full:
        return False
    # Prefer run-preserving replace when single-run or uniform
    if len(p.runs) == 1:
        p.runs[0].text = new
    else:
        set_para_text(p, new)
    return True


def replace_in_cell(cell: _Cell) -> int:
    """Replace across whole cell text; write back to first paragraph if changed."""
    full = cell.text
    if not full.strip():
        return 0
    new = rewrite_text(full, null_void_context=is_null_void_context(full))
    if new == full:
        return 0
    # Cell may have multiple paragraphs; rebuild simply into first para
    # Preserve multi-para only if unchanged structure needed — most cite cells are single-para
    if len(cell.paragraphs) == 1:
        set_para_text(cell.paragraphs[0], new)
    else:
        # Join was \n between paras; split new on \n if same count, else collapse
        parts = new.split("\n")
        if len(parts) == len(cell.paragraphs):
            for para, part in zip(cell.paragraphs, parts):
                set_para_text(para, part)
        else:
            set_cell_text(cell, new)
    return 1


def update_document_control(doc: Document) -> None:
    ctrl = doc.tables[0]
    set_cell_text(ctrl.rows[2].cells[1], OUT_VERSION)
    set_cell_text(
        ctrl.rows[3].cells[1],
        "Final — §3 Legal & §6 As-Is pain-point cross-references retargeted to hierarchical outline",
    )
    for row in ctrl.rows:
        if row.cells[0].text.strip().lower().startswith("last updated"):
            set_cell_text(row.cells[1], "09-09-2026")
            break
    hist = doc.tables[1]
    last_ver = hist.rows[-1].cells[0].text.strip()
    if last_ver != OUT_VERSION:
        add_version_row(
            hist,
            [
                OUT_VERSION,
                "09-09-2026",
                "Nandha Kumar",
                "Updated §3 (Legal Acts/Rules/Notifications implementation maps) and §6 (As-Is pain points): "
                "retargeted 7.i–7.vi / companion FRS_and_NFRS_Marriage_v1.22 cites to hierarchical §7–§11; "
                "fixed heading list levels",
                "Prashanth",
            ],
        )


def verify_sections_3_and_6(doc: Document) -> None:
    """Assert the screenshot targets (§3 legal maps + §6 pain points) are clean."""
    # §3 Hindu Act map (table with Refer section 7…)
    legal = None
    for t in doc.tables:
        hdr = " ".join(c.text for c in t.rows[0].cells)
        if "Refer section 7" in hdr and "Section" in t.rows[0].cells[0].text:
            # first such table is Hindu Act (T3 in v4)
            if legal is None:
                legal = t
                break
    assert legal is not None, "§3 legal implementation table not found"
    impl = legal.rows[1].cells[3].text.strip()
    assert impl.startswith("7.1.2.1"), f"§3 still has old cite: {impl}"
    assert "7.i" not in impl, impl

    # §6 pain intro
    intro = next(p.text for p in doc.paragraphs if "Pain points evidenced" in p.text)
    assert "FRS_and_NFRS_Marriage_v1.22" not in intro, intro
    assert "companion FRS" not in intro, intro
    assert "§7" in intro or "§8" in intro, intro

    pain = next(
        t
        for t in doc.tables
        if t.rows and "Addressed in" in t.rows[0].cells[-1].text
    )
    hdr = pain.rows[0].cells[-1].text.strip()
    assert "FRS" not in hdr.upper() or "this BRD" in hdr, hdr
    row1 = pain.rows[1].cells[-1].text.strip()
    assert "7.vi" not in row1 and "7.i.a" not in row1 and "FRS/NFRS:" not in row1, row1
    assert "7.6" in row1 and "8.11" in row1, row1
    print("§3 / §6 screenshot targets OK")
    print(f"  §3 sample: {impl}")
    print(f"  §6 intro: {intro[:120]}…")
    print(f"  §6 row1: {row1[:120]}…")


def simulate_outline(doc: Document) -> list[str]:
    counters = [0] * 9
    lines = []
    style_ilvl = {"Heading 2": 0, "Heading 3": 1, "Heading 4": 2, "Heading 5": 3, "Heading 6": 4}
    for p in doc.paragraphs:
        if not (p.style and p.style.name in style_ilvl and p.text.strip()):
            continue
        if p.text.strip() == "Document control":
            lines.append(f"— {p.text.strip()}")
            continue
        ilvl = style_ilvl[p.style.name]
        counters[ilvl] += 1
        for j in range(ilvl + 1, 9):
            counters[j] = 0
        num = ".".join(str(counters[k]) for k in range(ilvl + 1))
        lines.append(f"{num} {p.text.strip()}")
    return lines


def verify(doc: Document) -> None:
    # Exclude version-history table (historical change notes may keep old labels)
    parts: list[str] = [p.text for p in doc.paragraphs]
    for ti, table in enumerate(doc.tables):
        if ti == 1:
            continue
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    blob = "\n".join(parts)

    stale = []
    for pat in [
        r"7\.vi\.a",
        r"7\.vi\b",
        r"7\.iv\b",
        r"7\.iii\b",
        r"7\.ii\b",
        r"7\.i\.b",
        r"7\.i\.a",
        r"7\.i\b",
        r"7\.v\b",
        r"7\.b\.x",
        r"FRS_and_NFRS_Marriage_v1\.22",
        r"FRS and NFR document",
        r"companion FRS",
        r"see 1\.7",
        r"in 1\.2\.7",
        r"of 1\.3\.2",
        r"FRS/NFRS:",
        r"^1\.2\.1 ",
        r"^1\.4\.1 ",
        r"^8\.2 Security",
        r"^8\.3 System",
        r"^8\.4 Security Audit",
        r"7\.4 step",
        r"(?:^|[\n;/])1\.(?:1|2|3|4|8|9)(?:\.\d+)*\b",
    ]:
        if re.search(pat, blob, re.M):
            stale.append(pat)

    outline = simulate_outline(doc)
    assert any(l.startswith("7.1 ") and "Hindu Marriage" in l for l in outline), outline[30:40]
    assert any(l.startswith("7.4 ") and "Parsi" in l for l in outline), [l for l in outline if l.startswith("7.")]
    assert any(l.startswith("7.5 ") and "Null" in l for l in outline)
    assert any("Functional requirements" in l and l.split()[0].rstrip(".") == "8" for l in outline)

    print("Outline sample (top of §§6–9):")
    for l in outline:
        if re.match(r"^[6-9](?:\.\d+)* ", l) and l.count(".") <= 1:
            print(" ", l)

    if stale:
        print("WARNING — leftover patterns:", stale)
        for pat in stale[:6]:
            m = re.search(pat, blob, re.M)
            if m:
                start = max(0, m.start() - 30)
                print(f"  e.g. {pat}: {blob[start:m.end()+50]!r}")
    else:
        print("Verification OK — no stale roman / companion-FRS patterns found")


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    fix_heading_list_definition(doc)
    n_ilvl = fix_heading_ilvls(doc)
    n_strip = strip_heading_manual_numbers(doc)

    n_para = 0
    for p in doc.paragraphs:
        if replace_in_paragraph(p):
            n_para += 1

    n_cell = 0
    for table in doc.tables:
        for row in table.rows:
            seen_in_row: set[int] = set()
            for cell in row.cells:
                # Dedupe horizontally merged grid slots only (ids stable while row.cells live)
                tc_id = id(cell._tc)
                if tc_id in seen_in_row:
                    continue
                seen_in_row.add(tc_id)
                n_cell += replace_in_cell(cell)

    update_document_control(doc)
    doc.save(str(DST))

    # Reload for verify
    out = Document(str(DST))
    print(f"Wrote {DST}")
    print(f"  Heading ilvl fixes: {n_ilvl}")
    print(f"  Heading title strips: {n_strip}")
    print(f"  Paragraphs updated: {n_para}")
    print(f"  Cells updated: {n_cell}")
    print(f"  Version: {out.tables[0].rows[2].cells[1].text.strip()}")
    verify_sections_3_and_6(out)
    verify(out)


if __name__ == "__main__":
    main()
