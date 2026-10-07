"""Builds the Digital E-Stamp variants of the Kaveri legal document templates (Part I to Part IV only)
and Digital_EStamp_Template_Fields.xlsx.

Field definitions, clauses and layout helpers come from ../Document Registration
(_template_data.py and _make_templates.py), so both template sets stay aligned.

Digital E-Stamp variant:
    Part I   - Digital e-Stamp certificate (e-Stamping Rules 2009, Rule 11 particulars, QR, department seal)
    Part II  - Schedules
    Part III - Title flow / consideration / terms (affidavit body for layout A)
    Part IV  - Executants, witnesses and e-Sign execution
No Part V - an e-Stamp instrument is not endorsed by the Sub-Registrar. The e-Stamp number is printed at the
top of every page (Rule 28) and the instrument is written below the certificate (Rule 27).

Run:  python _make_estamp_templates.py
"""

import copy
import os
import re
import shutil
import sys
import tempfile
from collections import OrderedDict, defaultdict

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

HERE = os.path.dirname(os.path.abspath(__file__))
DOCREG = os.path.join(os.path.dirname(HERE), "Document Registration")
sys.path.insert(0, DOCREG)
import _make_templates as M  # noqa: E402
import _template_data as D  # noqa: E402

WORKBOOK = os.path.join(HERE, "Digital_EStamp_Template_Fields.xlsx")
REG_WORKBOOK = os.path.join(os.path.dirname(HERE), "Legal_Document_Template_Fields.xlsx")
SAMPLE_DIR = os.path.join(HERE, "_Sample_Rendered")
F = D.F

SOURCES = dict(D.SOURCES, EST="Kaveri Digital E-Stamp module (certificate / UIN)",
               DSR="Department digital seal (DSC / HSM)")

BLOCKS = copy.deepcopy(D.BLOCKS)
for f in BLOCKS["K4"]["fields"]:
    if f["key"] == "photo":
        f.update(fmt="From Aadhaar e-KYC", src="KYC", req="C",
                 desc="Aadhaar e-KYC photo; blank for DSC (non e-KYC) parties")
BLOCKS["K1"]["name"] = "Part I - Digital e-Stamp certificate"
BLOCKS["E1"] = dict(name="Part I - Digital e-Stamp certificate (Rule 11 particulars)", part="Part I", fields=[
    F("estamp_issue_time", "e-Stamp issue time", "Time", "HH:MM:SS (IST)", src="EST",
      desc="Rule 11(b) - time at which the UIN was allotted after payment", sample="12:38:11"),
    F("stamp_duty_in_words", "Stamp duty in words", src="SYS", desc="Rule 11(c)",
      sample="Rupees Six Hundred Only"),
    F("purchaser_name", "Purchaser of e-Stamp", src="KYC", desc="Rule 11(d) - logged-in applicant",
      sample="DHANUSH C V"),
    F("purchaser_address", "Purchaser address", src="KYC", desc="Rule 11(d)",
      sample="345 Marathahalli, Bangalore, Karnataka"),
    F("property_brief", "Brief description of property", req="C", src="SYS",
      desc="Rule 11(g) - schedule summary; 'Not applicable' when no schedule",
      sample="Khata No. 111, Nelamangala, Bangalore Rural"),
    F("duty_basis", "Basis of duty", "Dropdown", "Market value / Consideration / Fixed / Highest of all", src="SYS",
      sample="Fixed"),
    F("challan_date", "Challan / payment date", "Date", "DD-Mon-YYYY", src="PAY", sample="05-Sep-2026"),
    F("issued_by_user", "Issued by (user id)", src="EST",
      desc="Rule 11(h) - KAVERI-ONLINE for self-service; operator user id in assisted mode", sample="KAVERI-ONLINE"),
    F("issuing_office", "Issuing office code / location", src="EST", desc="Rule 11(i)",
      sample="Kaveri Digital E-Stamp"),
    F("qr_code", "QR code (verification)", "Image", "Signed payload: UIN + document hash", src="EST",
      desc="Rule 11(j) security mark", sample="[QR]"),
    F("department_seal", "Department digital seal field", src="DSR", desc="Rule 11(k) - applied after last e-Sign",
      sample="DeptSealField"),
])
BLOCKS["E4"] = dict(name="Part IV - Execution by e-Sign", part="Part IV", fields=[
    F("esign_completed_on", "e-Sign completed on", "DateTime", "DD-MM-YYYY HH:MM", src="SYS",
      desc="Time of the last party e-Sign", sample="05-09-2026 13:05"),
])

PART_ORDER = {"Header": 0, "Header / Footer": 0, "Part I": 1, "Part II": 2, "Part III": 3, "Part IV": 4}


# ---------------------------------------------------------------------------
# Selection
# ---------------------------------------------------------------------------
def _registration_index():
    ws = load_workbook(REG_WORKBOOK, read_only=True)["Template_Index"]
    rows = list(ws.iter_rows(values_only=True))
    tid, reg = rows[0].index("Template ID"), rows[0].index("Registration")
    return {r[tid]: r[reg] or "" for r in rows[1:] if r[tid]}


REGISTRATION = _registration_index()


def is_estamp(t):
    """Optionally registrable (Sec 18) stamp instrument, per the Registration column of the template workbook."""
    return (t["layout"] in ("K", "S", "A") and str(t["dn"]).startswith("DN-")
            and REGISTRATION.get(t["id"], "").startswith("Optional"))


def estamp_condition(t):
    return REGISTRATION.get(t["id"], "")


ESTAMP = [t for t in D.TEMPLATES if is_estamp(t)]


# ---------------------------------------------------------------------------
# Field catalogue
# ---------------------------------------------------------------------------
def estamp_blocks(t):
    out = [("HDR", "Required"), ("PSM", "Required"), ("K1", "Required"), ("E1", "Required")]
    if t["layout"] == "A":
        out.append(("A1", "Required"))
    else:
        out.append({"immovable": ("K2", "Required"), "optional": ("K2", "Optional"),
                    "movable": ("K2M", "Required")}.get(t["schedule"], (None, None)))
        if t["consideration"]:
            out.append(("K3", "Required"))
        out.append(("K3T", "Required"))
    out += [("K4", "Required"), ("K5", "Optional" if t["layout"] == "A" else "Required"), ("E4", "Required")]
    return [b for b in out if b[0]]


def field_rows(t):
    rows = []
    for bid, appl in estamp_blocks(t):
        blk = BLOCKS[bid]
        part = "Part III" if bid == "A1" else blk["part"]
        for i, f in enumerate(blk["fields"], 1):
            r = dict(f)
            r.update(fid=f"{bid}-{i:02d}", block=bid, block_name=blk["name"], part=part, appl=appl)
            if appl == "Optional" and r["req"] == "M":
                r["req"] = "C"
            rows.append(r)
    for i, f in enumerate(t["fields"], 1):
        r = dict(f)
        r.update(fid=f"{t['id']}-F{i:02d}", block="SPEC", block_name="Template specific", part="Part III",
                 appl="Required")
        rows.append(r)
    return rows


# ---------------------------------------------------------------------------
# Document parts
# ---------------------------------------------------------------------------
def new_document(t):
    doc = M.new_document(dict(t, layout="K"))
    hp = doc.sections[0].header.paragraphs[0]
    p = hp.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Digital e-Stamp No.: {{ estamp_cert_no }}")
    r.bold = True
    r.font.size = Pt(9)
    return doc


def part_i(doc):
    M.part_heading(doc, "Part I", "(Digital e-Stamp Certificate)")
    M.para(doc, "GOVERNMENT OF KARNATAKA - DEPARTMENT OF STAMPS AND REGISTRATION", bold=True, align="c",
           space_after=0)
    M.para(doc, "DIGITAL e-STAMP CERTIFICATE", bold=True, size=12, align="c", space_after=6)
    tb = M.table(doc, 4, (4.2, 4.3, 4.2, 4.3))
    rows = [
        ("Digital e-Stamp Certificate No.", "{{ estamp_cert_no }}", "Issue Date & Time",
         "{{ estamp_issue_date }} {{ estamp_issue_time }}"),
        ("Purchaser of e-Stamp", ["{{ purchaser_name }}", "{{ purchaser_address }}"], "Issued by / Office",
         ["{{ issued_by_user }}", "{{ issuing_office }}"]),
        ("Description of Document", "{{ document_description }}", "As Per Stamp Act",
         "{{ stamp_article }} {{ stamp_article_description }}"),
        ("Executor/Executant Details", "{{ executant_names }}", "Executee/Claimant Details", "{{ claimant_names }}"),
        ("Description of Property", "{{ property_brief }}", "Basis of Duty", "{{ duty_basis }}"),
        ("Stamp Duty Amount", ["₹ {{ stamp_duty_amount }}", "({{ stamp_duty_in_words }})"],
         "Challan Number / Date", ["{{ challan_number }}", "{{ challan_date }}"]),
    ]
    for r in rows:
        cells = M.add_row(tb, r)
        for i in (0, 2):
            M.set_cell_bg(cells[i], "F2F2F2")
            cells[i].paragraphs[0].runs[0].bold = True
    M.para(doc, space_after=2)
    tb = M.table(doc, 2, (4.0, 13.0))
    M.add_row(tb, ["{{ qr_code }}",
                   ["This Digital e-Stamp certificate is issued under the Karnataka Stamp (Payment of Duty by means of "
                    "e-Stamping) Rules, 2009 for the instrument written below it. No other instrument shall be written "
                    "on or using this certificate (Rule 27). Verify by scanning the QR code or at "
                    "https://kaveri.karnataka.gov.in using the certificate number.",
                    "Department seal: {{ department_seal }}"]])
    M.para(doc, space_after=2)
    M.end_of_part(doc, "Part I")


def part_ii(doc, t):
    if t["layout"] == "A":
        M.part_heading(doc, "Part II", "(Schedules)")
        M.para(doc, "Not applicable - an affidavit / declaration does not convey or affect a scheduled property.",
               italic=True)
        M.end_of_part(doc, "Part II")
    else:
        M.part_ii(doc, t)


def part_iii_affidavit(doc, t):
    M.part_heading(doc, "Part III", "(Statement on Oath / Declaration and Verification)")
    M.tag(doc, "{%p if affidavit_before %}")
    M.para(doc, "BEFORE {{ affidavit_before|upper }}", bold=True, align="c")
    M.tag(doc, "{%p endif %}")
    M.para(doc, "{{ case_reference }}", align="c")
    M.specific_section(doc, t, "Particulars")
    M.para(doc, "I, {{ deponent_name }}, {{ deponent_relation_type }} {{ deponent_relation_name }}, aged about "
                "{{ deponent_age }} years, {{ deponent_occupation }}, residing at {{ deponent_address }} "
                "(ID: {{ deponent_id_no }}), do hereby solemnly affirm and state on oath as follows:", align="j")
    M.tag(doc, "{%p for st in statements %}")
    M.para(doc, "{{ loop.index }}. {{ st }}", align="j")
    M.tag(doc, "{%p endfor %}")
    M.para(doc, "This affidavit is made {{ purpose }}.", align="j")
    M.heading(doc, "VERIFICATION")
    M.para(doc, "I, the deponent above named, do hereby verify that the contents of paragraphs 1 to "
                "{{ statements|length }} of this affidavit are true and correct to the best of my knowledge and belief "
                "and nothing material has been concealed therefrom. Verified at {{ verification_place }} on "
                "{{ verification_date }}.", align="j")
    M.para(doc, "Sworn / affirmed before me: {{ attesting_officer }} {{ attesting_officer_name }} "
                "(Reg. No. {{ attesting_reg_no }})", size=9)
    M.end_of_part(doc, "Part III")


def part_iv(doc, t):
    if t["layout"] == "A":
        M.part_heading(doc, "Part IV", "(Deponent and Witnesses)")
        M.heading(doc, "Deponent Details")
    else:
        M.part_heading(doc, "Part IV", "(Executants and Witnesses)")
        M.heading(doc, "Executant Details")
    M.party_rows(doc, t)
    if t["layout"] == "A":
        M.tag(doc, "{%p if witnesses %}")
        M.witness_rows(doc)
        M.tag(doc, "{%p endif %}")
    else:
        M.witness_rows(doc)
    M.para(doc, "IN WITNESS WHEREOF the parties have executed this instrument by Aadhaar e-Sign / Digital Signature "
                "on the dates recorded against their signatures, the last signature being affixed on "
                "{{ esign_completed_on }}, at {{ execution_place }}. Stamp duty of ₹ {{ stamp_duty_amount }} has been "
                "paid by Digital e-Stamp No. {{ estamp_cert_no }} before execution.", align="j")
    M.end_of_part(doc, "Part IV")


def build(t):
    doc = new_document(t)
    part_i(doc)
    part_ii(doc, t)
    if t["layout"] == "A":
        part_iii_affidavit(doc, t)
    else:
        M.part_iii(doc, t)
    part_iv(doc, t)
    return doc


def template_path(t):
    fam = f"{t['family']} {D.FAMILIES[t['family']]}"
    return os.path.join(HERE, fam, f"{t['id']} {M.safe_name(t['name'])} - Digital E-Stamp.docx")


# ---------------------------------------------------------------------------
# Sample context
# ---------------------------------------------------------------------------
def build_context(t):
    ctx = {}
    groups = defaultdict(list)
    for r in field_rows(t):
        if r["group"]:
            groups[r["group"]].append(r)
        else:
            ctx[r["key"]] = M.sample_value(r)
    for g, fs in groups.items():
        if fs[0]["key"] is None:
            ctx[g] = [M.sample_value(fs[0]), M.sample_value(fs[0]) + " (2)"]
        else:
            ctx[g] = [{f["key"]: M.sample_value(f) for f in fs} for _ in range(2)]
    ctx["document_title"] = t["title"]
    ctx.update(stamp_article=t["article"], stamp_article_description=t["name"], document_description=t["name"],
               stamp_duty_amount="500.00", stamp_duty_in_words="Rupees Five Hundred Only")
    e = dict(ctx["parties"][0], side="Executant", party_type=(t["exec_role"] or "Executant").split(" /")[0])
    c = dict(ctx["parties"][1], side="Claimant", party_type=(t["claim_role"] or "Claimant").split(" /")[0],
             name="Priya Singh", esign_field="SigClaim_370521_PartIV")
    ctx["parties"] = [e, c] if t["claim_role"] else [e]
    return ctx


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
def build_workbook(generated):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Kaveri 3.0 - Digital E-Stamp Templates : Field Workbook", True),
        ("", False),
        ("Purpose", True),
        (f"Fields captured for the {len(ESTAMP)} Digital E-Stamp template variants. Each variant has Part I to "
         "Part IV only: Part I Digital e-Stamp certificate, Part II Schedules, Part III Title Flow / Consideration / "
         "Terms (affidavit body for affidavits), Part IV Executants, Witnesses and e-Sign execution. Part V "
         "(Sub-Registrar endorsement) is not part of an e-Stamp instrument.", False),
        ("Statutory basis: Karnataka Stamp (Payment of Duty by means of e-Stamping) Rules, 2009 - Rule 11 "
         "certificate particulars, Rule 27 instrument written below the certificate, Rule 28 certificate number on "
         "every page (printed in the page header).", False),
        ("Selection: only optionally registrable instruments (Registration = 'Optional (Sec 18, Registration Act "
         "1908)...' in ../Legal_Document_Template_Fields.xlsx, sheet Template_Index). Where that entry carries a "
         "caveat (e.g. compulsory where possession is delivered), the caveat is shown in Template_Index.", False),
        ("", False),
        ("Sheets", True),
        ("Template_Index - one row per Digital E-Stamp template with article, availability condition and file.", False),
        ("Common_Fields - shared Part I to Part IV blocks (incl. E1 certificate particulars and E4 e-Sign execution).",
         False),
        ("Template_Fields - fields specific to each template (Part III Particulars).", False),
        ("Full_Field_List - every field of every template. Filter on Template ID.", False),
        ("Clauses - standard opening, recitals and clauses.", False),
        ("", False),
        ("Requirement codes", True),
        ("M = Mandatory, C = Conditional, O = Optional.", False),
        ("", False),
        ("Source codes", True),
    ] + [(f"{k} = {v}", False) for k, v in SOURCES.items()] + [
        ("", False),
        ("Field definitions are shared with the Document Registration templates "
         "(../Document Registration/_template_data.py). Regenerate with: python _make_estamp_templates.py", False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True, size=12 if i == 1 else 11, color="1F4E78")
    ws.column_dimensions["A"].width = 130

    ws = wb.create_sheet("Template_Index")
    rows = []
    for t in ESTAMP:
        rows.append((t["id"], t["name"], f"{t['family']} {D.FAMILIES[t['family']]}", t["layout"], t["title"],
                     t["dn"], t["article"], estamp_condition(t), t["exec_role"] or "-", t["claim_role"] or "-",
                     {"immovable": "Required (immovable)", "movable": "Required (movable)", "optional": "Optional",
                      "none": "-"}[t["schedule"]] if t["layout"] != "A" else "-",
                     "Yes" if t["consideration"] else "No", len(field_rows(t)),
                     os.path.relpath(generated[t["id"]], HERE)))
    M.write_sheet(ws, ["Template ID", "Document type", "Family", "Layout", "Page header title", "Nature code (DN)",
                       "KSA Article", "Registration (per Legal_Document_Template_Fields.xlsx)", "Executant role", "Claimant role",
                       "Property schedule", "Consideration", "Total fields", "Template file"], rows,
                  [11, 38, 24, 7, 30, 10, 16, 50, 22, 22, 18, 12, 9, 70])

    ws = wb.create_sheet("Common_Fields")
    rows = []
    used = OrderedDict()
    for t in ESTAMP:
        for bid, _ in estamp_blocks(t):
            used[bid] = True
    for bid in sorted(used, key=lambda b: (PART_ORDER.get(BLOCKS[b]["part"], 3), list(BLOCKS).index(b))):
        blk = BLOCKS[bid]
        for i, f in enumerate(blk["fields"], 1):
            rows.append((bid, blk["name"], "Part III" if bid == "A1" else blk["part"], f"{bid}-{i:02d}",
                         M.placeholder(f), f["group"] or "-", f["label"], f["desc"], f["dtype"], f["fmt"], f["req"],
                         f["src"], f["sample"] if f["sample"] is not None else ""))
    M.write_sheet(ws, ["Block", "Block name", "Part", "Field ID", "Placeholder", "Repeating group", "Field label",
                       "Description / rule", "Data type", "Format / allowed values", "Req", "Source", "Sample"], rows,
                  [7, 34, 10, 9, 30, 13, 34, 40, 13, 40, 6, 7, 40], req_col=10, freeze="E2")

    ws = wb.create_sheet("Template_Fields")
    rows = []
    for t in ESTAMP:
        for i, f in enumerate(t["fields"], 1):
            rows.append((t["id"], t["name"], f"{t['id']}-F{i:02d}", M.placeholder(f), f["group"] or "-", f["label"],
                         f["desc"], f["dtype"], f["fmt"], f["req"], f["src"],
                         f["sample"] if f["sample"] is not None else ""))
    M.write_sheet(ws, ["Template ID", "Document type", "Field ID", "Placeholder", "Repeating group", "Field label",
                       "Description / rule", "Data type", "Format / allowed values", "Req", "Source", "Sample"], rows,
                  [11, 34, 13, 34, 12, 40, 30, 12, 50, 6, 7, 34], req_col=9, freeze="D2")

    ws = wb.create_sheet("Full_Field_List")
    rows = []
    for t in ESTAMP:
        for r in field_rows(t):
            rows.append((t["id"], t["name"], r["part"], r["block"], r["block_name"], r["appl"], r["fid"],
                         M.placeholder(r), r["group"] or "-", r["label"], r["dtype"], r["fmt"], r["req"], r["src"]))
    M.write_sheet(ws, ["Template ID", "Document type", "Part", "Block", "Block name", "Block applicability",
                       "Field ID", "Placeholder", "Repeating group", "Field label", "Data type",
                       "Format / allowed values", "Req", "Source"], rows,
                  [11, 32, 10, 7, 34, 11, 13, 32, 12, 38, 12, 44, 6, 7], req_col=12, freeze="H2")

    ws = wb.create_sheet("Clauses")
    rows = []
    for t in ESTAMP:
        op = M.opener(t) if t["layout"] != "A" else None
        if op:
            rows.append((t["id"], t["name"], "Opening", 0, op))
        for i, r in enumerate(t["recitals"], 1):
            rows.append((t["id"], t["name"], "Recital", i, r))
        for i, c in enumerate(t["clauses"], 1):
            rows.append((t["id"], t["name"], "Clause", i, c))
    M.write_sheet(ws, ["Template ID", "Document type", "Type", "No.", "Text (with placeholders)"], rows,
                  [11, 34, 10, 5, 120], freeze="C2")
    wb.save(WORKBOOK)


# ---------------------------------------------------------------------------
def main():
    for d in [d for d in os.listdir(HERE) if re.match(r"^\d\d ", d)]:
        shutil.rmtree(os.path.join(HERE, d))
    generated = {}
    tmp = tempfile.mkdtemp(prefix="kaveri_est_")
    for t in ESTAMP:
        path = template_path(t)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        build(t).save(path)
        generated[t["id"]] = path
        text = M.render(path, build_context(t), os.path.join(tmp, f"{t['id']}.docx"))
        if "Part V" in text or "Part IV" not in text:
            raise RuntimeError(f"{t['id']}: expected Part I to Part IV only")
    shutil.rmtree(tmp, ignore_errors=True)

    os.makedirs(SAMPLE_DIR, exist_ok=True)
    sample_id = "T-LSE-02"
    sample_out = os.path.join(SAMPLE_DIR, f"{sample_id} - rendered with sample data.docx")
    M.render(generated[sample_id], build_context(D.TEMPLATES[[t["id"] for t in D.TEMPLATES].index(sample_id)]),
             sample_out)
    build_workbook(generated)

    by_layout = defaultdict(int)
    for t in ESTAMP:
        by_layout[t["layout"]] += 1
    print(f"Digital E-Stamp templates: {len(generated)} (Part I-IV, all rendered OK with StrictUndefined) "
          f"by layout {dict(by_layout)}")
    print(f"Workbook: {WORKBOOK}")
    print(f"Sample  : {sample_out}")


if __name__ == "__main__":
    main()
