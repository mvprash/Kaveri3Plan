"""Builds the Kaveri legal document templates (.docx) and Legal_Document_Template_Fields.xlsx
from _template_data.py, maps every source format in ../LEGAL FORMATS to a template, and
test-renders each template with docxtpl (StrictUndefined) to prove every placeholder is defined.

Run:  python _make_templates.py
"""

import copy
import os
import re
import shutil
import sys
import tempfile
from collections import Counter, OrderedDict, defaultdict

import jinja2
from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docxtpl import DocxTemplate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import _template_data as D  # noqa: E402

TEMPLATE_ROOT = os.path.dirname(HERE)
SOURCE_ROOT = os.path.join(os.path.dirname(TEMPLATE_ROOT), "LEGAL FORMATS")
WORKBOOK = os.path.join(TEMPLATE_ROOT, "Legal_Document_Template_Fields.xlsx")
SAMPLE_DIR = os.path.join(HERE, "_Sample_Rendered")

TPL = {t["id"]: t for t in D.TEMPLATES}
ALIAS = D.GROUP_ALIAS

VERIFY_NOTE = ("NOTE : Please verify authenticity & genuineness of this digital stamp paper online by visiting "
               "https://kaveri.karnataka.gov.in or scanning the QR Code. Digital Stamp paper not issued by Stamps & "
               "Registration Department, Govt. of Karnataka is INVALID and PERSONS USING ARE LIABLE FOR CRIMINAL CASE.")


# ---------------------------------------------------------------------------
# Field catalogue per template
# ---------------------------------------------------------------------------
def template_blocks(t):
    """[(block_id, applicability)] for a template. applicability: 'Required' / 'Optional'."""
    out = []
    for b in D.LAYOUT_BLOCKS[t["layout"]]:
        if b == "K2":
            if t["schedule"] == "immovable":
                out.append(("K2", "Required"))
            elif t["schedule"] == "optional":
                out.append(("K2", "Optional"))
            elif t["schedule"] == "movable":
                out.append(("K2M", "Required"))
            continue
        if b == "K3":
            if t["consideration"]:
                out.append(("K3", "Required"))
            continue
        out.append((b, "Required"))
    return out


def placeholder(f):
    if f["group"] is None:
        return "{{ %s }}" % f["key"]
    a = ALIAS[f["group"]]
    if f["key"] is None:
        return "{%% for %s in %s %%}{{ %s }}{%% endfor %%}" % (a, f["group"], a)
    return "{{ %s.%s }}" % (a, f["key"])


def field_rows(t):
    """All fields of a template: list of dicts with block info + field id."""
    rows = []
    for bid, appl in template_blocks(t):
        blk = D.BLOCKS[bid]
        for i, f in enumerate(blk["fields"], 1):
            r = dict(f)
            r.update(fid=f"{bid}-{i:02d}", block=bid, block_name=blk["name"], part=blk["part"], appl=appl)
            if appl == "Optional" and r["req"] == "M":
                r["req"] = "C"
            rows.append(r)
    for i, f in enumerate(t["fields"], 1):
        r = dict(f)
        r.update(fid=f"{t['id']}-F{i:02d}", block="SPEC", block_name="Template specific", part=spec_part(t),
                 appl="Required")
        rows.append(r)
    return rows


def spec_part(t):
    return "Part III" if t["layout"] in ("K",) else "Body"


# ---------------------------------------------------------------------------
# docx helpers
# ---------------------------------------------------------------------------
def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    run._r.append(fld)


def para(doc_or_cell, text="", bold=False, italic=False, size=None, align=None, space_after=4, color=None):
    p = doc_or_cell.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    return p


def tag(doc, text):
    """Paragraph-level Jinja tag ({%p ... %}) - docxtpl removes the paragraph."""
    return para(doc, text, space_after=0)


def cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(0)


def cell_lines(cell, lines, size=9):
    cell.text = ""
    first = True
    for ln in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        r = p.add_run(ln)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(0)


def table(doc, ncols, widths=None):
    tb = doc.add_table(rows=0, cols=ncols)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    tb._widths = widths
    return tb


def add_row(tb, values, bold=False, header=False, size=9):
    cells = tb.add_row().cells
    for c, v in zip(cells, values):
        if isinstance(v, (list, tuple)):
            cell_lines(c, v, size)
        else:
            cell_text(c, v, bold=bold or header, size=size)
        if header:
            set_cell_bg(c, "D9E2F3")
    if tb._widths:
        for c, w in zip(cells, tb._widths):
            c.width = Cm(w)
    return cells


def kv_table(doc, pairs, widths=(6.5, 10.5)):
    tb = table(doc, 2, widths)
    for k, v in pairs:
        cells = add_row(tb, [k, v])
        set_cell_bg(cells[0], "F2F2F2")
        cells[0].paragraphs[0].runs[0].bold = True
    para(doc, space_after=2)
    return tb


def loop_table(doc, headers, group, cells, widths=None, total=None):
    """Header row + docxtpl row loop. cells: list of str or list-of-lines per column."""
    tb = table(doc, len(headers), widths)
    add_row(tb, headers, header=True)
    add_row(tb, ["{%%tr for %s in %s %%}" % (ALIAS[group], group)] + [""] * (len(headers) - 1))
    add_row(tb, cells)
    add_row(tb, ["{%tr endfor %}"] + [""] * (len(headers) - 1))
    if total:
        add_row(tb, total, bold=True)
    para(doc, space_after=2)
    return tb


def new_document(t):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run("{{ document_title }}")
    r.bold = True
    r.font.size = Pt(13)

    fp = sec.footer.paragraphs[0]
    if t["layout"] in ("K", "S"):
        r = fp.add_run("[QR]  " + VERIFY_NOTE)
        r.font.size = Pt(7)
        f2 = sec.footer.add_paragraph()
        r = f2.add_run("{{ footer_party_names }}")
        r.font.size = Pt(8)
        f3 = sec.footer.add_paragraph()
    else:
        f3 = fp
    f3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = f3.add_run("Page ")
    r.font.size = Pt(8)
    add_page_field(f3)
    return doc


def part_heading(doc, part, subtitle):
    para(doc, part, bold=True, size=12, align="c", space_after=0)
    para(doc, subtitle, italic=True, align="c", space_after=6)


def end_of_part(doc, part):
    para(doc, f"--------------- End of {part} ---------------", align="c", size=9, space_after=8)


def heading(doc, text):
    para(doc, text, bold=True, size=11, space_after=4)


# ---------------------------------------------------------------------------
# Template-specific particulars
# ---------------------------------------------------------------------------
def specific_section(doc, t, title="Particulars of the Transaction"):
    scalars = [f for f in t["fields"] if f["group"] is None]
    groups = OrderedDict()
    for f in t["fields"]:
        if f["group"]:
            groups.setdefault(f["group"], []).append(f)
    if not scalars and not groups:
        return
    heading(doc, title)
    if scalars:
        kv_table(doc, [(f["label"], placeholder(f)) for f in scalars])
    for g, fs in groups.items():
        label = g.replace("_", " ").title()
        para(doc, label, bold=True, space_after=2)
        a = ALIAS[g]
        if fs[0]["key"] is None:
            tag(doc, "{%%p for %s in %s %%}" % (a, g))
            para(doc, "{{ loop.index }}. {{ %s }}" % a, align="j")
            tag(doc, "{%p endfor %}")
        else:
            loop_table(doc, ["Sl.No"] + [f["label"] for f in fs], g,
                       ["{{ loop.index }}"] + ["{{ %s.%s }}" % (a, f["key"]) for f in fs])


# ---------------------------------------------------------------------------
# Layout K - Kaveri registrable instrument: Digital Execution (Part I - IV) + Admission and Registration (Part V)
# ---------------------------------------------------------------------------
def part_i(doc):
    part_heading(doc, "Part I", "(Document Header)")
    tb = table(doc, 4, (4.2, 4.3, 4.2, 4.3))
    rows = [
        ("Digital e-Stamp Certificate No.", "{{ estamp_cert_no }}", "Executee/Claimant Details", "{{ claimant_names }}"),
        ("Issue Date", "{{ estamp_issue_date }}", "Executor/Executant Details", "{{ executant_names }}"),
        ("Description of Document", "{{ document_description }}", "Stamp Duty Amount", "₹ {{ stamp_duty_amount }}"),
        ("As Per Stamp Act", "{{ stamp_article }} {{ stamp_article_description }}", "Challan Number",
         "{{ challan_number }}"),
    ]
    for r in rows:
        cells = add_row(tb, r)
        for i in (0, 2):
            set_cell_bg(cells[i], "F2F2F2")
            cells[i].paragraphs[0].runs[0].bold = True
    para(doc, space_after=2)
    end_of_part(doc, "Part I")


def part_ii(doc, t):
    part_heading(doc, "Part II", "(Schedules)")
    heading(doc, "Schedule Details")
    if t["schedule"] in ("immovable", "optional"):
        if t["schedule"] == "optional":
            tag(doc, "{%p if schedules %}")
        tag(doc, "{%p for s in schedules %}")
        para(doc, "Schedule {{ loop.index }} of {{ schedules|length }} ({{ s.source }})", bold=True)
        tb = table(doc, 4)
        add_row(tb, ["District:", "Taluka:", "Hobli:", "Village:"], header=True)
        add_row(tb, ["{{ s.district }}", "{{ s.taluk }}", "{{ s.hobli }}", "{{ s.village }}"])
        para(doc, space_after=2)
        tb = table(doc, 7)
        add_row(tb, ["Sy.No", "Site Area ({{ s.area_unit }})", "Builtup Area", "North", "South", "East", "West"],
                header=True)
        add_row(tb, ["{{ s.survey_no }}", "{{ s.site_area }}", "{{ s.builtup_area }}", "{{ s.north }}",
                     "{{ s.south }}", "{{ s.east }}", "{{ s.west }}"])
        para(doc, space_after=2)
        kv_table(doc, [
            ("Owner Name", "{{ s.owner_name }}"),
            ("{{ s.property_category }} Property Type", "{{ s.property_type }}"),
            ("{{ s.property_category }} Property Usage", "{{ s.property_usage }}"),
            ("First Sale", "{{ s.first_sale }}"),
            ("Local Body Limits", "{{ s.local_body_limit }}"),
            ("Market Value (Rs.)", "{{ s.market_value }}"),
        ])
        para(doc, "Property Number Details", bold=True, space_after=2)
        tb = table(doc, 2, (6.5, 10.5))
        add_row(tb, ["Type", "Current Number"], header=True)
        add_row(tb, ["{{ s.property_number_type }}", "{{ s.property_number }}"])
        para(doc, space_after=2)
        para(doc, "RDPR PID {{ s.pid }}", bold=True)
        para(doc, "Description:", bold=True, space_after=0)
        para(doc, "{{ s.description }}", align="j")
        tag(doc, "{%p endfor %}")
        if t["schedule"] == "optional":
            tag(doc, "{%p else %}")
            para(doc, "No property schedule is annexed to this instrument.", italic=True)
            tag(doc, "{%p endif %}")
    elif t["schedule"] == "movable":
        para(doc, "Schedule of Movable Property / Subject Matter", bold=True)
        loop_table(doc, ["Sl.No", "Description", "Identification", "Quantity", "Value (Rs.)"], "items",
                   ["{{ loop.index }}", "{{ it.description }}", "{{ it.identification }}", "{{ it.quantity }}",
                    "{{ it.value }}"], widths=(1.3, 6.5, 4, 2, 3.2))
    else:
        para(doc, "Not applicable - this instrument does not convey or affect a scheduled property.", italic=True)
    end_of_part(doc, "Part II")


def opener(t):
    if t["id"].startswith("T-WIL-01"):
        return None
    e = t["exec_role"] or "Executant"
    if t["claim_role"]:
        return ("This {{ document_title|title }} is made and executed at {{ execution_place }} on "
                "{{ execution_date }} by {{ executant_names }} (hereinafter called the '%s', which expression shall "
                "include heirs, legal representatives and assigns) in favour of {{ claimant_names }} (hereinafter "
                "called the '%s', which expression shall include heirs, legal representatives and assigns)." % (e, t["claim_role"]))
    return ("This {{ document_title|title }} is made and executed at {{ execution_place }} on {{ execution_date }} by "
            "{{ executant_names }} (hereinafter called the '%s')." % e)


def clauses_block(doc, t, heading_text="Terms and Conditions", with_additional=True):
    heading(doc, heading_text)
    if t["layout"] == "K":
        para(doc, "NOW THIS INSTRUMENT WITNESSETH AS FOLLOWS:", bold=True)
    n = 0
    for n, c in enumerate(t["clauses"], 1):
        para(doc, f"{n}. {c}", align="j")
    if with_additional:
        tag(doc, "{%p for t in terms %}")
        para(doc, "{{ loop.index + %d }}. {{ t }}" % n, align="j")
        tag(doc, "{%p endfor %}")


def part_iii(doc, t):
    part_heading(doc, "Part III", "(Title Flow/Consideration Payments/Terms & Conditions)")
    heading(doc, "Title Flow")
    op = opener(t)
    if op:
        para(doc, op, align="j")
    for r in t["recitals"]:
        para(doc, r, align="j")
    para(doc, "{{ title_flow }}", align="j")
    specific_section(doc, t)
    if t["consideration"]:
        heading(doc, "Consideration Payment Details")
        loop_table(doc, ["Date of Payment", "Payment Mode", "Reference Number", "Bank Name", "Amount", "Remarks"],
                   "payments",
                   ["{{ p.date }}", "{{ p.mode }}", "{{ p.reference_no }}", "{{ p.bank_name }}", "₹ {{ p.amount }}",
                    "{{ p.remarks }}"],
                   total=["Total", "", "", "", "₹ {{ payments_total }}", ""])
        para(doc, "Consideration: ₹ {{ consideration_amount }} ({{ consideration_in_words }}); Market value "
                  "(Sec 45-A): ₹ {{ market_value_total }}", size=9)
    clauses_block(doc, t)
    end_of_part(doc, "Part III")


def party_rows(doc, t):
    loop_table(doc, ["Sl.No", "Type of Executant", "Name", "Photo", "UIDAI Ref No.", "e-Sign"], "parties", [
        "{{ loop.index }}",
        ["{{ pt.party_type }}", "({{ pt.side }})"],
        ["Name as in ID Proof: {{ pt.name }}",
         "{{ pt.relation_type }} {{ pt.relation_name }}",
         "Category: {{ pt.party_category }}",
         "DOB: {{ pt.dob }}   Age: {{ pt.age }}   Gender: {{ pt.gender }}",
         "Address: {{ pt.address }}",
         "PinCode: {{ pt.pincode }}",
         "ID Proof: {{ pt.id_proof_type }} {{ pt.id_proof_no }}   PAN: {{ pt.pan }}",
         "Mobile: {{ pt.mobile }}   E-mail: {{ pt.email }}",
         "Executing capacity: {{ pt.capacity }}{% if pt.capacity != 'Self' %}, represented by {{ pt.represented_by }} "
         "({{ pt.authority_reference }}){% endif %}",
         "Aadhaar e-KYC Status: {{ pt.ekyc_status }}"],
        "{{ pt.photo }}", "{{ pt.uidai_ref }}", "{{ pt.esign_field }}",
    ], widths=(1.2, 2.6, 7.4, 1.8, 2.0, 2.0))


def witness_rows(doc):
    heading(doc, "Witness Details")
    loop_table(doc, ["Sl.No", "Name", "Photo", "UIDAI Ref No.", "e-Sign"], "witnesses", [
        "{{ loop.index }}",
        ["Name as in ID Proof: {{ w.name }}", "DOB: {{ w.dob }}", "Gender: {{ w.gender }}",
         "Address: {{ w.address }}", "PinCode: {{ w.pincode }}", "ID Proof Number: {{ w.id_proof_no }}",
         "Aadhaar e-KYC Status: {{ w.ekyc_status }}"],
        "[Photo]", "", "{{ w.esign_field }}",
    ], widths=(1.2, 9.6, 2.0, 2.2, 2.0))


def part_iv(doc, t):
    part_heading(doc, "Part IV", "(Executants and Witnesses)")
    heading(doc, "Executant Details")
    party_rows(doc, t)
    witness_rows(doc)
    end_of_part(doc, "Part IV")


def part_v(doc, standalone=False):
    if standalone:
        kv_table(doc, [
            ("Document", "{{ document_title }}"),
            ("Kaveri application no.", "{{ kaveri_application_no }}"),
            ("e-Stamp / stamp certificate no.", "{{ estamp_cert_no }}"),
            ("Executed on / at", "{{ execution_date }} / {{ execution_place }}"),
            ("Digital execution document (Part I-IV) reference", "{{ execution_document_ref }}"),
        ])
    else:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    part_heading(doc, "Part V", "Endorsement By Sub Registrar")
    heading(doc, "1. Presentation")
    para(doc, "Presented on {{ presentation_date }} at the Office of the {{ sro_office }} Sub-Registrar on "
              "{{ presentation_date }} at {{ presentation_time }} by {{ presenter_name }}.", align="j")
    heading(doc, "2. Fee Details")
    loop_table(doc, ["Sl.No", "Description", "Fees in Rs."], "fees",
               ["{{ loop.index }}", "{{ f.description }}", "{{ f.amount }}"],
               widths=(1.5, 11, 4.5), total=["", "Total", "{{ fees_total }}"])
    heading(doc, "3. Presented by {{ presenter_name }}")
    tb = table(doc, 5, (1.2, 8.8, 2.4, 2.2, 2.4))
    add_row(tb, ["Sl.No", "Name", "Captured Photo", "Thumb", "eSign"], header=True)
    add_row(tb, ["1", ["{{ presenter_name }}", "{{ presenter_address }}"], "", "", "{{ presenter_esign }}"])
    para(doc, space_after=2)
    heading(doc, "4. Document execution admitted by")
    loop_table(doc, ["Sl.No", "Type of Executant", "Admission Date", "Aadhar eKYC Details", "Captured Photo", "Thumb",
                     "eSign"], "admissions",
               ["{{ loop.index }}", "{{ a.party_type }}", "{{ a.admission_datetime }}",
                ["Name: {{ a.name }},", "DOB: {{ a.dob }},", "Address: {{ a.address }}"], "", "", "{{ a.esign }}"],
               widths=(1.1, 2.3, 2.3, 5.3, 2.0, 1.6, 2.4))
    heading(doc, "5. Executants Identified by")
    loop_table(doc, ["Sl.No", "Identifier Name", "Photo", "eSign"], "identifiers",
               ["{{ loop.index }}", "{{ idf.name }} (Identifier)", "", "{{ idf.esign }}"],
               widths=(1.2, 9.4, 2.6, 3.8))
    para(doc, "Registered as Document No. {{ registration_number }} in {{ book_no }} on {{ registration_date }}.",
         bold=True)
    para(doc, "SRO Signature  {{ sro_signature }}", bold=True, align="r")


def build_k(t):
    """Part I - Part IV: digital execution template (parties execute by e-Sign before presentation)."""
    doc = new_document(t)
    part_i(doc)
    part_ii(doc, t)
    part_iii(doc, t)
    part_iv(doc, t)
    return doc


def build_k_registration(t):
    """Part V: admission and registration template (SRO presentation, admission, identification, registration)."""
    doc = new_document(t)
    hp = doc.sections[0].header.paragraphs[0]
    r = hp.add_run(" - ADMISSION AND REGISTRATION")
    r.bold = True
    r.font.size = Pt(13)
    part_v(doc, standalone=True)
    return doc


SPLIT_KINDS = OrderedDict([
    ("exec", "Part I-IV Digital Execution"),
    ("reg", "Part V Admission and Registration"),
])


# ---------------------------------------------------------------------------
# Layout S - simple instrument
# ---------------------------------------------------------------------------
def stamp_strip(doc):
    kv_table(doc, [("e-Stamp / Stamp Certificate No.", "{{ estamp_cert_no }} dated {{ estamp_issue_date }}"),
                   ("Description / Article", "{{ document_description }} - Art {{ stamp_article }} "
                                             "{{ stamp_article_description }}"),
                   ("Stamp Duty / Challan", "₹ {{ stamp_duty_amount }} / {{ challan_number }}")])


def build_s(t):
    doc = new_document(t)
    stamp_strip(doc)
    op = opener(t)
    if op:
        para(doc, op, align="j")
    specific_section(doc, t, "Particulars")
    clauses_block(doc, t, "Terms", with_additional=False)
    para(doc, "Place: {{ execution_place }}          Date: {{ execution_date }}")
    heading(doc, "Parties")
    party_rows(doc, t)
    witness_rows(doc)
    return doc


# ---------------------------------------------------------------------------
# Layout A - affidavit
# ---------------------------------------------------------------------------
def build_a(t):
    doc = new_document(t)
    stamp_strip(doc)
    tag(doc, "{%p if affidavit_before %}")
    para(doc, "BEFORE {{ affidavit_before|upper }}", bold=True, align="c")
    tag(doc, "{%p endif %}")
    para(doc, "{{ case_reference }}", align="c")
    para(doc, "{{ document_title }}", bold=True, size=12, align="c")
    specific_section(doc, t, "Particulars")
    para(doc, "I, {{ deponent_name }}, {{ deponent_relation_type }} {{ deponent_relation_name }}, aged about "
              "{{ deponent_age }} years, {{ deponent_occupation }}, residing at {{ deponent_address }} "
              "(ID: {{ deponent_id_no }}), do hereby solemnly affirm and state on oath as follows:", align="j")
    tag(doc, "{%p for st in statements %}")
    para(doc, "{{ loop.index }}. {{ st }}", align="j")
    tag(doc, "{%p endfor %}")
    para(doc, "This affidavit is made {{ purpose }}.", align="j")
    para(doc, "DEPONENT", bold=True, align="r")
    heading(doc, "VERIFICATION")
    para(doc, "I, the deponent above named, do hereby verify that the contents of paragraphs 1 to "
              "{{ statements|length }} of this affidavit are true and correct to the best of my knowledge and belief "
              "and nothing material has been concealed therefrom. Verified at {{ verification_place }} on "
              "{{ verification_date }}.", align="j")
    para(doc, "DEPONENT", bold=True, align="r")
    para(doc, "Sworn / affirmed before me: {{ attesting_officer }} {{ attesting_officer_name }} "
              "(Reg. No. {{ attesting_reg_no }})", size=9)
    return doc


# ---------------------------------------------------------------------------
# Layout C / F - court
# ---------------------------------------------------------------------------
def cause_title(doc):
    para(doc, "IN THE {{ court_name|upper }} AT {{ court_place|upper }}", bold=True, align="c")
    para(doc, "{{ case_type }} No. {{ case_number }} of {{ case_year }}", bold=True, align="c")
    para(doc, "(Under {{ filed_under }})", italic=True, align="c")
    tag(doc, "{%p if related_case and related_case != '-' %}")
    para(doc, "In {{ related_case }}", align="c")
    tag(doc, "{%p endif %}")
    para(doc, "BETWEEN:", bold=True)
    tag(doc, "{%p for pe in petitioners %}")
    para(doc, "{{ loop.index }}. {{ pe.name }}, {{ pe.relation_type }} {{ pe.relation_name }}, aged about {{ pe.age }} "
              "years, {{ pe.occupation }}, residing at {{ pe.address }}", align="j")
    tag(doc, "{%p endfor %}")
    para(doc, "... {{ petitioner_label }}(s)", bold=True, align="r")
    para(doc, "AND", bold=True, align="c")
    tag(doc, "{%p for re in respondents %}")
    para(doc, "{{ loop.index }}. {{ re.name }}, {{ re.relation_type }} {{ re.relation_name }}, aged about {{ re.age }} "
              "years, {{ re.occupation }}, residing at {{ re.address }}", align="j")
    tag(doc, "{%p endfor %}")
    para(doc, "... {{ respondent_label }}(s)", bold=True, align="r")


def build_c(t):
    doc = new_document(t)
    cause_title(doc)
    para(doc, "{{ document_heading }}", bold=True, size=11, align="c")
    specific_section(doc, t, "Particulars")
    para(doc, "The {{ petitioner_label }} above named most respectfully submits as follows:", align="j")
    heading(doc, "FACTS")
    tag(doc, "{%p for fa in facts %}")
    para(doc, "{{ loop.index }}. {{ fa }}", align="j")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p if grounds %}")
    heading(doc, "GROUNDS")
    tag(doc, "{%p for gr in grounds %}")
    para(doc, "{{ loop.index }}. {{ gr }}", align="j")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endif %}")
    for h, k in (("CAUSE OF ACTION", "cause_of_action"), ("JURISDICTION", "jurisdiction_statement"),
                 ("LIMITATION", "limitation_statement"), ("VALUATION AND COURT FEE", "valuation_court_fee")):
        tag(doc, "{%%p if %s %%}" % k)
        heading(doc, h)
        para(doc, "{{ %s }}" % k, align="j")
        tag(doc, "{%p endif %}")
    heading(doc, "PRAYER")
    para(doc, "Wherefore, the {{ petitioner_label }} respectfully prays that this Hon'ble Court may be pleased to:",
         align="j")
    tag(doc, "{%p for pr in prayers %}")
    para(doc, "({{ loop.index }}) {{ pr }}", align="j")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p if interim_relief and interim_relief != '-' %}")
    para(doc, "Interim relief: {{ interim_relief }}", align="j")
    tag(doc, "{%p endif %}")
    para(doc, "and grant such other relief(s) as this Hon'ble Court deems fit in the interest of justice and equity.",
         align="j")
    tb = table(doc, 2, (8.5, 8.5))
    add_row(tb, [["Place: {{ filing_place }}", "Date: {{ filing_date }}", "", "Advocate for {{ petitioner_label }}",
                  "{{ advocate_name }} ({{ advocate_enrolment_no }})", "{{ advocate_address }}",
                  "{{ advocate_contact }}"],
                 ["", "", "", "Signature of {{ petitioner_label }}", "{{ verifier_name }}"]], size=10)
    para(doc, space_after=4)
    heading(doc, "VERIFICATION")
    para(doc, "I, {{ verifier_name }}, the {{ petitioner_label }} above named, do hereby verify that the contents of "
              "paragraphs 1 to {{ facts|length }} above are true to my personal knowledge and belief and the rest are "
              "based on legal advice believed to be true. Verified at {{ filing_place }} on {{ filing_date }}.",
         align="j")
    para(doc, "{{ petitioner_label }}", bold=True, align="r")
    tag(doc, "{%p if documents %}")
    heading(doc, "LIST OF DOCUMENTS")
    loop_table(doc, ["Sl.No", "Description of document", "Date"], "documents",
               ["{{ loop.index }}", "{{ d.description }}", "{{ d.date }}"], widths=(1.5, 11.5, 4))
    tag(doc, "{%p endif %}")
    return doc


def build_f(t):
    doc = new_document(t)
    cause_title(doc)
    para(doc, "{{ form_title }}", bold=True, size=11, align="c")
    specific_section(doc, t, "Particulars")
    para(doc, "To,", space_after=0)
    para(doc, "{{ addressee_name }}", bold=True, space_after=0)
    para(doc, "{{ addressee_address }}")
    tag(doc, "{%p for op in order_paragraphs %}")
    para(doc, "{{ op }}", align="j")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p if hearing_date and hearing_date != '-' %}")
    para(doc, "Date fixed for hearing / compliance: {{ hearing_date }}", bold=True)
    tag(doc, "{%p endif %}")
    para(doc, "Given under my hand and the seal of the Court, this {{ issue_date }}.", align="j")
    para(doc, "(Seal)")
    para(doc, "{{ presiding_officer }}", bold=True, align="r")
    return doc


# ---------------------------------------------------------------------------
# Layout N / G / R
# ---------------------------------------------------------------------------
def build_n(t):
    doc = new_document(t)
    para(doc, "{{ sender_name }}", bold=True, size=12, align="c", space_after=0)
    para(doc, "{{ sender_address }}", align="c", space_after=0)
    para(doc, "{{ sender_contact }}", align="c")
    tb = table(doc, 2, (8.5, 8.5))
    add_row(tb, ["Ref: {{ notice_ref_no }}", "Date: {{ notice_date }}"], size=10)
    para(doc, "BY {{ delivery_mode|upper }}", bold=True)
    para(doc, "To,", space_after=0)
    tag(doc, "{%p for ad in addressees %}")
    para(doc, "{{ loop.index }}. {{ ad.name }}, {{ ad.address }}", space_after=0)
    tag(doc, "{%p endfor %}")
    para(doc, "")
    para(doc, "{{ document_title }}", bold=True, size=12, align="c")
    para(doc, "Sub: {{ subject }}", bold=True)
    specific_section(doc, t, "Particulars")
    tag(doc, "{%p if client_name %}")
    para(doc, "Under instructions from and on behalf of my client {{ client_name }}, residing at {{ client_address }}, "
              "I hereby serve upon you the following notice:", align="j")
    tag(doc, "{%p endif %}")
    tag(doc, "{%p for pg in paragraphs %}")
    para(doc, "{{ loop.index }}. {{ pg }}", align="j")
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p if demand_text %}")
    para(doc, "I therefore call upon you to {{ demand_text }} within {{ compliance_days }} days from the date of receipt "
              "of this notice, failing which {{ consequence_text }}", align="j")
    tag(doc, "{%p endif %}")
    para(doc, "A copy of this notice is retained in my office for record and further action.", align="j")
    para(doc, "{{ sender_name }}", bold=True, align="r")
    return doc


def build_g(t):
    doc = new_document(t)
    para(doc, "From:", bold=True, space_after=0)
    para(doc, "{{ applicant_name }}, {{ applicant_relation }}", space_after=0)
    para(doc, "{{ applicant_address }}", space_after=0)
    para(doc, "{{ applicant_contact }}")
    para(doc, "To,", bold=True, space_after=0)
    para(doc, "{{ to_designation }}", space_after=0)
    para(doc, "{{ to_office }}", space_after=0)
    para(doc, "{{ to_address }}")
    para(doc, "Sub: {{ subject }}", bold=True, space_after=0)
    para(doc, "Ref: {{ reference }}")
    specific_section(doc, t, "Particulars")
    para(doc, "Respected Sir / Madam,")
    tag(doc, "{%p for pg in paragraphs %}")
    para(doc, "{{ loop.index }}. {{ pg }}", align="j")
    tag(doc, "{%p endfor %}")
    para(doc, "I therefore request you to {{ request_text }}", align="j")
    tag(doc, "{%p if enclosures %}")
    para(doc, "Enclosures:", bold=True, space_after=0)
    tag(doc, "{%p for en in enclosures %}")
    para(doc, "{{ loop.index }}. {{ en }}", space_after=0)
    tag(doc, "{%p endfor %}")
    tag(doc, "{%p endif %}")
    para(doc, "")
    tb = table(doc, 2, (8.5, 8.5))
    add_row(tb, [["Place: {{ application_place }}", "Date: {{ application_date }}"],
                 ["Yours faithfully,", "", "{{ applicant_name }}"]], size=10)
    return doc


def build_r(t):
    doc = new_document(t)
    para(doc, "{{ company_name|upper }}", bold=True, size=13, align="c", space_after=0)
    para(doc, "CIN / Regn. No.: {{ cin }}", align="c", space_after=0)
    para(doc, "Registered Office: {{ registered_office }}", align="c")
    para(doc, "CERTIFIED TRUE COPY OF THE {{ resolution_kind|upper }} PASSED AT THE {{ meeting_type|upper }} OF "
              "{{ company_name|upper }} HELD ON {{ meeting_date }} AT {{ meeting_time }} AT {{ meeting_place|upper }}",
         bold=True, align="j")
    para(doc, "Chairperson: {{ chairperson }}")
    specific_section(doc, t, "Particulars")
    para(doc, "{{ resolution_title|upper }}", bold=True, align="c")
    tag(doc, "{%p for rs in resolutions %}")
    para(doc, "{{ rs }}", align="j")
    tag(doc, "{%p endfor %}")
    para(doc, "RESOLVED FURTHER THAT the following persons be and are hereby severally authorised to sign and execute "
              "all documents and do all acts necessary to give effect to this resolution:", align="j")
    loop_table(doc, ["Sl.No", "Name", "Designation", "DIN / ID"], "signatories",
               ["{{ loop.index }}", "{{ sg.name }}", "{{ sg.designation }}", "{{ sg.din }}"], widths=(1.5, 7, 5, 3.5))
    para(doc, "Certified to be a true copy", italic=True, space_after=0)
    para(doc, "For {{ company_name }}", bold=True)
    para(doc, "")
    para(doc, "{{ certified_by_name }}", bold=True, space_after=0)
    para(doc, "{{ certified_by_designation }}   DIN: {{ certified_by_din }}", space_after=0)
    para(doc, "Place: {{ certification_place }}   Date: {{ certification_date }}")
    return doc


BUILDERS = {"K": build_k, "S": build_s, "A": build_a, "C": build_c, "F": build_f, "N": build_n, "G": build_g,
            "R": build_r}


def safe_name(s):
    return re.sub(r"[\\/:*?\"<>|]", "-", s).strip()


def template_path(t, kind=None):
    fam = f"{t['family']} {D.FAMILIES[t['family']]}"
    suffix = f" - {SPLIT_KINDS[kind]}" if kind else ""
    return os.path.join(HERE, fam, f"{t['id']} {safe_name(t['name'])}{suffix}.docx")


def template_files(t):
    """{kind: (path, builder)} - layout K is split into digital execution and admission & registration."""
    if t["layout"] == "K":
        return OrderedDict([("exec", (template_path(t, "exec"), build_k)),
                            ("reg", (template_path(t, "reg"), build_k_registration))])
    return OrderedDict([(None, (template_path(t), BUILDERS[t["layout"]]))])


# ---------------------------------------------------------------------------
# Sample contexts
# ---------------------------------------------------------------------------
def sample_value(f, idx=0):
    if f["sample"] not in (None, ""):
        return f["sample"]
    return f"<{f['label']}>"


def build_context(t):
    rows = field_rows(t)
    ctx = {}
    groups = defaultdict(list)
    for r in rows:
        if r["group"]:
            groups[r["group"]].append(r)
        else:
            ctx[r["key"]] = sample_value(r)
    for g, fs in groups.items():
        if fs[0]["key"] is None:
            ctx[g] = [sample_value(fs[0]), sample_value(fs[0]) + " (2)"]
        else:
            items = []
            for i in range(2):
                items.append({f["key"]: sample_value(f) for f in fs})
            ctx[g] = items
    ctx["document_title"] = t["title"]
    if "parties" in ctx:
        e = dict(ctx["parties"][0], side="Executant", party_type=(t["exec_role"] or "Executant").split(" /")[0])
        c = dict(ctx["parties"][1], side="Claimant", party_type=(t["claim_role"] or "Claimant").split(" /")[0],
                 name="Priya Singh", esign_field="SigClaim_370521_PartIV")
        ctx["parties"] = [e, c] if t["claim_role"] else [e]
    return ctx


def sale_deed_reference_context():
    """Values transcribed from Sale Deed Filled-Template/Endorsement_740241_20260905124528380 (1).pdf."""
    ctx = build_context(TPL["T-SAL-01"])
    ctx.update({
        "document_title": "DEED OF SALE", "kaveri_application_no": "740241", "execution_date": "05-09-2026",
        "execution_place": "Bengaluru", "executant_names": "DHANUSH C V", "claimant_names": "Priya Singh; Priya Singh",
        "footer_party_names": "1). DHANUSH C V, 2). Priya Singh, 3). Priya Singh",
        "estamp_cert_no": "UL0YR415HXISVBO", "estamp_issue_date": "05-Sep-2026", "document_description": "Sale",
        "stamp_duty_amount": "60730.00", "stamp_article": "20 (1)",
        "stamp_article_description": "Conveyance: Sale of property (market value), not charged under No.52.",
        "challan_number": "PAY-20260905123811-639683",
        "title_flow": "WHEREAS sdfdsf sdfdsfdsfdsfsd sdfdsfsdfs dfsfds f ds",
        "terms": [], "consideration_amount": "100000.00", "consideration_in_words": "Rupees One Lakh Only",
        "market_value_total": "607300.00", "payments_total": "100000.00",
        "payments": [{"date": "05-Sep-2026", "mode": "cash", "reference_no": "N/A", "bank_name": "N/A",
                      "amount": "100000.00", "remarks": "N/A"}],
        "instrument_subtype": "Absolute sale", "sale_consideration": "100000.00", "advance_paid": "0.00",
        "agreement_reference": "-", "possession_date": "05-09-2026", "encumbrance_status": "Free from encumbrances",
        "previous_deed_reference": "title as recorded in e-Swathu (Khata No. 111)", "khata_transfer": "Yes",
        "execution_document_ref": "UL0YR415HXISVBO-EXEC",
        "sro_office": "Gandhinagara", "presentation_date": "05-09-2026", "presentation_time": "12:40 PM",
        "presenter_name": "DHANUSH C V", "presenter_address": "345 Marathahalli, Bangalore, Karnataka",
        "presenter_esign": "SigPresenter_370520_Endorsement",
        "fees": [{"description": "Stamp Duty", "amount": "60,730.00"}, {"description": "Stamp Duty", "amount": "60,730.00"}],
        "fees_total": "121,460.00", "registration_number": "-", "registration_date": "-", "book_no": "Book I",
        "sro_signature": "SROSignaturefield",
    })
    ctx["schedules"] = [{
        "source": "Imported from ESwathu", "district": "Bangalore Rural", "taluk": "Nelamangala1",
        "hobli": "Nelamangala Hobli", "village": "Nelamangala", "survey_no": "N/A", "site_area": "121.46",
        "area_unit": "Sq.m", "builtup_area": "0", "north": "ಸೈಟ್ ನಂ 106", "south": "-", "east": "ಸೈಟ್ ನಂ 108",
        "west": "-", "owner_name": "DHANUSH C V", "property_category": "Non Agriculture", "property_type": "Vacant",
        "property_usage": "Residential", "first_sale": "No", "property_number_type": "Khata No",
        "property_number": "111", "pid": "150300300700820626", "local_body_limit": "LL-3",
        "market_value": "607300.00",
        "description": "All that piece and parcel of non-agricultural property, situated at Bangalore Rural District, "
                       "Nelamangala1 Taluk, and bounded on the: East:-; West:-; North:-; South:-.",
    }]
    base = dict(party_category="Individual", relation_type="", relation_name="", gender="", id_proof_type="Aadhaar",
                pan="", mobile="", email="", capacity="Self", represented_by="", authority_reference="",
                photo="[Photo]", uidai_ref="", ekyc_status="True")
    ctx["parties"] = [
        dict(base, side="Executant", party_type="Seller", name="DHANUSH C V", dob="28/08/1976", age="50",
             address="345 Marathahalli, Bangalore, Karnataka", pincode="560001", id_proof_no="1",
             esign_field="SigExec_370520_PartIV"),
        dict(base, side="Claimant", party_type="Purchaser", name="Priya Singh", dob="15/07/1985", age="41",
             address="456 Brigade Road, Pune, Maharashtra", pincode="411001", id_proof_no="2333",
             esign_field="SigClaim_370521_PartIV"),
        dict(base, side="Claimant", party_type="Purchaser", name="Priya Singh", dob="15/07/1985", age="41",
             address="456 Brigade Road, Pune, Maharashtra", pincode="411001", id_proof_no="2333",
             esign_field="SigClaim_370522_PartIV"),
    ]
    ctx["witnesses"] = [
        dict(name="Rajesh Kumar Kumar", dob="14/12/1953", gender="", address="890 Koramangala", pincode="700001",
             id_proof_no="1111", ekyc_status="True", esign_field="SigWitness_140143_PartIV"),
        dict(name="Anjali Kumar Nair", dob="04/09/2007", gender="", address="123 MG Road", pincode="400001",
             id_proof_no="432", ekyc_status="True", esign_field="SigWitness_140144_PartIV"),
    ]
    ctx["admissions"] = [
        dict(party_type=p["party_type"], admission_datetime="05-09-2026 12:40", name=p["name"],
             dob=p["dob"].replace("/", "-"), address=p["address"],
             esign=f"SigParty_{370520 + i}_Endorsement") for i, p in enumerate(ctx["parties"])]
    ctx["identifiers"] = [dict(name="Rajesh Kumar Kumar", esign="SigWitness_140143_Endorsement"),
                          dict(name="Anjali Kumar Nair", esign="SigWitness_140144_Endorsement")]
    return ctx


def render(path, ctx, out):
    doc = DocxTemplate(path)
    env = jinja2.Environment(undefined=jinja2.StrictUndefined, autoescape=True)
    doc.render(ctx, jinja_env=env)
    doc.save(out)
    text = "\n".join(p.text for p in Document(out).paragraphs)
    for tb in Document(out).tables:
        for row in tb.rows:
            for c in row.cells:
                text += "\n" + c.text
    left = re.findall(r"\{\{|\{%|%\}|\}\}", text)
    if left:
        raise RuntimeError(f"Unrendered tags remain in {out}: {left[:5]}")
    return text


# ---------------------------------------------------------------------------
# Source mapping
# ---------------------------------------------------------------------------
KW = [(re.compile(rx), tid) for rx, tid in D.KEYWORD_RULES]


def classify(folder, fname):
    stem, ext = os.path.splitext(fname)
    up = stem.upper().strip()
    if fname.startswith("~") or ext.lower() in (".tmp", ".rar"):
        return "IGNORE", "Temp / archive file"
    for f, prefix, tid in D.OVERRIDES:
        if f == folder and up.startswith(prefix.upper()):
            if tid is None:
                break
            return tid, "Override" if prefix else "Folder rule"
    for rx, tid in KW:
        if rx.search(up):
            if tid == "IGNORE":
                return tid, "Temp / archive file"
            return tid, "Keyword"
    return D.FOLDER_DEFAULTS.get(folder.split(os.sep)[0], "T-REF-01"), "Folder default"


def source_mapping():
    rows = []
    for folder in sorted(os.listdir(SOURCE_ROOT), key=str.lower):
        p = os.path.join(SOURCE_ROOT, folder)
        if not os.path.isdir(p):
            continue
        for dp, _dn, fns in os.walk(p):
            for fn in sorted(fns, key=str.lower):
                rel_dir = os.path.relpath(dp, SOURCE_ROOT)
                tid, basis = classify(folder, fn)
                rows.append((rel_dir, fn, os.path.splitext(fn)[1].lower(), tid, basis))
    return rows


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
GRP_FILL = PatternFill("solid", fgColor="DDEBF7")
M_FILL = PatternFill("solid", fgColor="E2EFDA")
C_FILL = PatternFill("solid", fgColor="FFF2CC")
O_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_sheet(ws, headers, rows, widths, req_col=None, freeze="B2"):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 32
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.alignment, cell.border = WRAP, BORDER
        if req_col is not None:
            v = row[req_col].value
            row[req_col].fill = {"M": M_FILL, "C": C_FILL, "O": O_FILL}.get(v, PatternFill())
            row[req_col].alignment = Alignment(horizontal="center", vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def build_workbook(mapping, generated):
    wb = Workbook()
    counts = Counter(r[3] for r in mapping)
    folders = defaultdict(set)
    for r in mapping:
        folders[r[3]].add(r[0])

    # README
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Kaveri 3.0 - Legal Document Templates : Field Workbook", True),
        ("", False),
        ("Purpose", True),
        ("Describes every field to be captured for each legal document template generated from the formats in "
         "'Legal Formats/LEGAL FORMATS' (%d source files in %d folders). Templates follow the registered Sale Deed "
         "(Endorsement_740241) structure: Part I Document Header, Part II Schedules, Part III Title Flow / "
         "Consideration / Terms, Part IV Executants & Witnesses, Part V Endorsement by Sub-Registrar."
         % (len(mapping), len({r[0].split(os.sep)[0] for r in mapping})), False),
        ("", False),
        ("Document Registration flow - two templates per registrable instrument (layout K)", True),
        ("1. Digital Execution template (Part I - Part IV): generated after data entry and stamp duty payment; the "
         "executants, claimants and witnesses e-Sign it before presentation. File suffix '- %s'." % SPLIT_KINDS["exec"],
         False),
        ("2. Admission and Registration template (Part V): generated at the Sub-Registrar office after presentation, "
         "fee collection, admission of execution and identification; signed by the SRO on registration. It carries "
         "the application no., e-Stamp no. and the reference / hash of the e-Signed Part I-IV document so the two "
         "are bound together in the registered record. File suffix '- %s'." % SPLIT_KINDS["reg"], False),
        ("Other layouts (S, A, C, F, N, G, R) have no Part V and remain a single template file.", False),
        ("", False),
        ("Sheets", True),
        ("Template_Index - one row per document type template, with stamp article, registration need, roles, file.", False),
        ("Layout_Blocks - which common field blocks apply to each template.", False),
        ("Common_Fields - fields of the shared blocks (Part I, II, III, IV, V, court, notice, affidavit, resolution).", False),
        ("Template_Fields - fields specific to each template (Particulars section).", False),
        ("Full_Field_List - every field of every template (common + specific). Filter on Template ID.", False),
        ("Clauses - standard recitals and clauses printed in each template, with placeholders.", False),
        ("Source_Mapping - every source format file mapped to the template that replaces it.", False),
        ("Lookups - dropdown values referenced by the fields.", False),
        ("", False),
        ("Placeholders", True),
        ("Templates are docxtpl / Jinja2 templates. Scalar fields print as {{ field }}. Repeating groups (parties, "
         "witnesses, schedules, payments, fees ...) are table-row loops; inside a loop a field prints as "
         "{{ alias.field }} (e.g. {{ pt.name }} inside parties). Plain list groups (terms, facts, prayers ...) print "
         "one paragraph per item.", False),
        ("", False),
        ("Requirement codes", True),
        ("M = Mandatory, C = Conditional (rule in Description / Format), O = Optional. A block marked Optional for a "
         "template turns its mandatory fields into C (required only if the block is used).", False),
        ("", False),
        ("Source codes", True),
    ] + [(f"{k} = {v}", False) for k, v in D.SOURCES.items()] + [
        ("", False),
        ("Layouts", True),
    ] + [(f"{k} = {v}", False) for k, v in D.LAYOUTS.items()] + [
        ("", False),
        ("Notes", True),
        ("Stamp articles (Karnataka Stamp Act, 1957) and registration requirement are indicative; the Stamp Duty "
         "Article module (Acts_Rules/Document) makes the final determination. Nature codes DN-xx match "
         "Stamp_Duty_Article_Determination_Inputs.xlsx.", False),
        ("Regenerate templates and this workbook with: python _make_templates.py", False),
    ]
    for i, (txt, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True, size=12 if i == 1 else 11, color="1F4E78")
    ws.column_dimensions["A"].width = 130

    # Template_Index
    ws = wb.create_sheet("Template_Index")
    rows = []
    def rel(t, kind):
        p = generated.get(t["id"], {}).get(kind)
        return os.path.relpath(p, TEMPLATE_ROOT) if p else "-"

    for t in D.TEMPLATES:
        sub = next((f["fmt"] for f in t["fields"] if f["key"] == "instrument_subtype"), "")
        files = generated.get(t["id"], {})
        if not files:
            exec_file, reg_file = "(no template)", "-"
        elif "reg" in files:
            exec_file, reg_file = rel(t, "exec"), rel(t, "reg")
        else:
            exec_file, reg_file = rel(t, None), "- (no Part V for this layout)"
        rows.append((t["id"], t["name"], f"{t['family']} {D.FAMILIES[t['family']]}", t["layout"], t["title"], t["dn"],
                     t["article"], t["registration"], t["exec_role"] or "-", t["claim_role"] or "-",
                     {"immovable": "Required (immovable)", "movable": "Required (movable)", "optional": "Optional",
                      "none": "-"}[t["schedule"]] if t["layout"] == "K" else "-",
                     "Yes" if t["consideration"] else "No", sub, exec_file, reg_file,
                     counts.get(t["id"], 0), ", ".join(sorted(folders.get(t["id"], [])))))
    write_sheet(ws, ["Template ID", "Document type", "Family", "Layout", "Page header title", "Nature code (DN)",
                     "KSA Article", "Registration", "Executant role", "Claimant role", "Property schedule",
                     "Consideration", "Sub-types (dropdown)",
                     "Digital Execution template (Part I-IV) / single template file",
                     "Admission and Registration template (Part V) file", "Source formats mapped",
                     "Source folders"], rows,
                [11, 38, 24, 7, 30, 10, 16, 34, 22, 22, 18, 12, 60, 55, 55, 10, 60])

    # Layout_Blocks
    ws = wb.create_sheet("Layout_Blocks")
    bids = list(D.BLOCKS.keys())
    rows = []
    for t in D.TEMPLATES:
        appl = dict(template_blocks(t))
        rows.append([t["id"], t["name"], t["layout"]] + [appl.get(b, "-") for b in bids] + [len(t["fields"])])
    write_sheet(ws, ["Template ID", "Document type", "Layout"] + [f"{b}\n{D.BLOCKS[b]['name']}" for b in bids]
                + ["Specific fields"], rows, [11, 38, 7] + [14] * len(bids) + [10], freeze="D2")
    ws.row_dimensions[1].height = 80
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=3 + len(bids)):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="top")
            c.fill = M_FILL if c.value == "Required" else C_FILL if c.value == "Optional" else PatternFill()

    # Common_Fields
    ws = wb.create_sheet("Common_Fields")
    rows = []
    for bid, blk in D.BLOCKS.items():
        for i, f in enumerate(blk["fields"], 1):
            rows.append((bid, blk["name"], blk["part"], f"{bid}-{i:02d}", placeholder(f), f["group"] or "-",
                         f["label"], f["desc"], f["dtype"], f["fmt"], f["req"], f["src"],
                         f["sample"] if f["sample"] is not None else ""))
    write_sheet(ws, ["Block", "Block name", "Part / section", "Field ID", "Placeholder", "Repeating group", "Field label",
                     "Description / rule", "Data type", "Format / allowed values", "Req", "Source",
                     "Sample (Sale Deed 740241 where available)"], rows,
                [7, 30, 14, 9, 30, 13, 34, 40, 13, 40, 6, 7, 40], req_col=10, freeze="E2")

    # Template_Fields
    ws = wb.create_sheet("Template_Fields")
    rows = []
    for t in D.TEMPLATES:
        for i, f in enumerate(t["fields"], 1):
            rows.append((t["id"], t["name"], f"{t['id']}-F{i:02d}", placeholder(f), f["group"] or "-", f["label"],
                         f["desc"], f["dtype"], f["fmt"], f["req"], f["src"],
                         f["sample"] if f["sample"] is not None else ""))
    write_sheet(ws, ["Template ID", "Document type", "Field ID", "Placeholder", "Repeating group", "Field label",
                     "Description / rule", "Data type", "Format / allowed values", "Req", "Source", "Sample"], rows,
                [11, 34, 13, 34, 12, 40, 30, 12, 50, 6, 7, 34], req_col=9, freeze="D2")

    # Full_Field_List
    ws = wb.create_sheet("Full_Field_List")
    rows = []
    for t in D.TEMPLATES:
        for r in field_rows(t):
            if t["layout"] != "K":
                tfile = "Single template"
            elif r["part"] == "Part V":
                tfile = "Admission and Registration"
            else:
                tfile = "Digital Execution"
            rows.append((t["id"], t["name"], r["part"], tfile, r["block"], r["block_name"], r["appl"], r["fid"],
                         placeholder(r), r["group"] or "-", r["label"], r["dtype"], r["fmt"], r["req"], r["src"]))
    write_sheet(ws, ["Template ID", "Document type", "Part / section", "Template file", "Block", "Block name",
                     "Block applicability", "Field ID", "Placeholder", "Repeating group", "Field label", "Data type",
                     "Format / allowed values", "Req", "Source"], rows,
                [11, 32, 13, 22, 7, 30, 11, 13, 32, 12, 38, 12, 44, 6, 7], req_col=13, freeze="I2")

    # Clauses
    ws = wb.create_sheet("Clauses")
    rows = []
    for t in D.TEMPLATES:
        op = opener(t) if t["layout"] in ("K", "S") else None
        if op:
            rows.append((t["id"], t["name"], "Opening", 0, op))
        for i, r in enumerate(t["recitals"], 1):
            rows.append((t["id"], t["name"], "Recital", i, r))
        for i, c in enumerate(t["clauses"], 1):
            rows.append((t["id"], t["name"], "Clause", i, c))
    write_sheet(ws, ["Template ID", "Document type", "Type", "No.", "Text (with placeholders)"], rows,
                [11, 34, 10, 5, 120], freeze="C2")

    # Source_Mapping
    ws = wb.create_sheet("Source_Mapping")
    rows = [(r[0], r[1], r[2], r[3], TPL[r[3]]["name"] if r[3] in TPL else "Ignored", r[4]) for r in mapping]
    write_sheet(ws, ["Source folder", "Source file", "Ext", "Template ID", "Document type", "Mapping basis"], rows,
                [36, 80, 7, 11, 44, 14], freeze="C2")

    # Lookups
    ws = wb.create_sheet("Lookups")
    names = list(D.LOOKUPS.keys())
    ws.append(names)
    for c in range(1, len(names) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = 34
    for c, n in enumerate(names, 1):
        for r, v in enumerate(D.LOOKUPS[n], 2):
            ws.cell(row=r, column=c, value=v).alignment = WRAP
    ws.freeze_panes = "A2"

    wb.save(WORKBOOK)


# ---------------------------------------------------------------------------
def validate_data():
    errs = []
    ids = [t["id"] for t in D.TEMPLATES]
    if len(ids) != len(set(ids)):
        errs.append("Duplicate template ids")
    for t in D.TEMPLATES:
        for f in t["fields"]:
            if f["group"] and f["group"] not in ALIAS:
                errs.append(f"{t['id']}: group {f['group']} has no alias")
    for _f, _p, tid in D.OVERRIDES:
        if tid and tid not in TPL:
            errs.append(f"Override -> unknown template {tid}")
    for _rx, tid in D.KEYWORD_RULES:
        if tid != "IGNORE" and tid not in TPL:
            errs.append(f"Keyword rule -> unknown template {tid}")
    for k, tid in D.FOLDER_DEFAULTS.items():
        if tid not in TPL:
            errs.append(f"Folder default {k} -> unknown template {tid}")
    if errs:
        raise SystemExit("\n".join(errs))


def main():
    validate_data()
    for fam_dir in [d for d in os.listdir(HERE) if re.match(r"^\d\d ", d)]:
        shutil.rmtree(os.path.join(HERE, fam_dir))
    generated = {}
    tmp = tempfile.mkdtemp(prefix="kaveri_tpl_")
    n_files = 0
    for t in D.TEMPLATES:
        if t["layout"] == "X":
            continue
        generated[t["id"]] = OrderedDict()
        for kind, (path, builder) in template_files(t).items():
            os.makedirs(os.path.dirname(path), exist_ok=True)
            builder(t).save(path)
            generated[t["id"]][kind] = path
            text = render(path, build_context(t), os.path.join(tmp, f"{t['id']}-{kind}.docx"))
            has_v = "Part V" in text
            if (kind == "reg") != has_v or (kind == "exec" and "Part IV" not in text):
                raise RuntimeError(f"{t['id']} {kind}: unexpected parts in rendered output")
            n_files += 1
    os.makedirs(SAMPLE_DIR, exist_ok=True)
    for f in os.listdir(SAMPLE_DIR):
        os.remove(os.path.join(SAMPLE_DIR, f))
    ctx = sale_deed_reference_context()
    text = ""
    sample_outs = []
    for kind, label in SPLIT_KINDS.items():
        out = os.path.join(SAMPLE_DIR, f"T-SAL-01 Sale Deed - {label} - rendered with Endorsement 740241 data.docx")
        text += render(generated["T-SAL-01"][kind], ctx, out)
        sample_outs.append(out)
    for must in ("UL0YR415HXISVBO", "PAY-20260905123811-639683", "150300300700820626", "SigClaim_370522_PartIV",
                 "SigWitness_140144_Endorsement", "121,460.00", "Gandhinagara"):
        if must not in text:
            raise RuntimeError(f"Sale deed sample missing '{must}'")
    shutil.rmtree(tmp, ignore_errors=True)

    mapping = source_mapping()
    build_workbook(mapping, generated)

    c = Counter(r[3] for r in mapping)
    b = Counter(r[4] for r in mapping)
    n_split = sum(1 for v in generated.values() if "reg" in v)
    print(f"Templates generated : {len(generated)} templates, {n_files} files "
          f"({n_split} split into Part I-IV Digital Execution + Part V Admission and Registration); "
          "all rendered OK with StrictUndefined")
    print(f"Source files mapped : {len(mapping)}  basis: {dict(b)}")
    unused = [t['id'] for t in D.TEMPLATES if c.get(t['id'], 0) == 0]
    print(f"Templates with no source file: {unused}")
    print(f"Workbook            : {WORKBOOK}")
    for s in sample_outs:
        print(f"Sale deed sample    : {s}")


if __name__ == "__main__":
    main()
