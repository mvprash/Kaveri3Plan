# -*- coding: utf-8 -*-
"""Restructure complete In Scope section from CSG Documents only.

Sources under CSG Documents/:
  FRS Documents/* (UM, Marriage, Document/master FRS, Firm, Paperless, Digital e-Stamp,
                   GIS valuation, EC/CC in master FRS, Suo Moto CC, Accounts, Audit,
                   Round-robin, K1 pending, Will, Book-I filings, Re-registration, Aadhaar note)
  Design Documents/Application (KaveriArchitecture.pdf, Software Design Document)
  Design Documents/Database Design (schemas / ERDs)
  QA Files/Test_Cases (testing coverage)

Output:
  IT Project Scope/IT_Project_Scope_Document_v1.4.docx
  + Template1 + CSG Documents copy
"""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC_CANDIDATES = [
    SCOPE_DIR / "IT_Project_Scope_Document_Template1.docx",
    SCOPE_DIR / "IT_Project_Scope_Document_v1.3.docx",
    SCOPE_DIR / "IT_Project_Scope_Document_v1.2.docx",
]
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v1.4.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v1.4.docx"

# Complete restructured In Scope — every row filled from CSG Documents
SCOPE_ROWS: list[tuple[str, str, list[str]]] = [
    (
        "IS-01",
        "User Management",
        [
            "Source: User management FRS V2.6; citizen registration in FRS – Kaveri 2.0 v0.2.38.",
            "",
            "Citizen / Applicant:",
            "Account creation, login, forgot password; Email / Mobile OTP verification.",
            "Mobile OTP-based login (also required for Digital e-Stamp) alongside email login.",
            "Aadhaar-based e-KYC for citizens where mandated by flow.",
            "",
            "Department users:",
            "Admin hierarchy – Super Admin (DIGR Admin), DRO Admin, SRO Admin.",
            "Masters – Office, Post, User Type, Role; Map modules / sub-modules / functions to Roles (RBAC).",
            "User lifecycle – Create / View / Edit user (KGID for Government Appointee); Activate / Deactivate.",
            "Deputation / Leave management; Transfer; Promotion.",
            "Authentication – KGID / biometric / face as applicable for department users.",
            "",
            "Boundary: Access control foundation for all other modules. External dept users (IT, Bank, Lawyer) via admin-provisioned accounts where applicable.",
        ],
    ),
    (
        "IS-02",
        "Marriage Registration",
        [
            "Source: Marriage Registration Module FRS V4.0.",
            "",
            "Marriage types:",
            "Hindu Marriage – Online and Offline (Hindu Marriage Act, 1955).",
            "Special Marriage (Intended Marriage Notice) – Offline notice, solemnization, certificate (Special Marriage Act, 1954 / Karnataka Rules).",
            "Special Marriage of Other Forms – Online and Offline.",
            "Parsi Marriage (Parsi Marriage and Divorce Act, 1936) – as covered in marriage module scope.",
            "",
            "Application has to support both Online and Offline Process.",
            "Online – via Kaveri Application / Portal.",
            "Offline – at SR / Marriage Office; registration by SRO / MO / FDA / SDA.",
            "",
            "Core capabilities: application capture, fee payment, scrutiny, joint photo / witness capture, marriage certificate generation and delivery, MIS; bilingual (Kannada / English).",
            "Boundary: Marriage Certificate Certified Copy issuance aligns with IS-06 where applicable.",
        ],
    ),
    (
        "IS-03",
        "Document Registration",
        [
            "Source: FRS – Kaveri 2.0 v0.2.38; Paperless registration-I FRS; Re-registration FRS v1.4;",
            "K1 pending document FRS; Will after death of testator FRS; SRO Round-robin FRS;",
            "Book-I Part-I Other Memo / Part III(b) / Rule 17(3) Filing FRSs; PoA / Witness eKYC / related notes.",
            "",
            "2 types of flows – Paper and Paperless.",
            "Paper – ink-signed deed, physical presentation, scan & store.",
            "Paperless – digital execution (e-Sign / DSC), biometric e-KYC, digital endorsements & Sec. 60.",
            "",
            "Online and Offline Process (Portal / SRO Kiosk / DEO / FDA / SDA / SR).",
            "",
            "Citizen pre-registration flow:",
            "Type of Document (Sale, Mortgage, Lease, DoT, Partition, Gift, GPA, Will, etc.); Property Search (Bhoomi, e-Swathu, e-Aasthi, ULMS); Property Schedule;",
            "Market Valuation & Fee Calculation (detail under IS-04); Party Information; Document Summary; SRO Confirmation;",
            "Payment; Appointment; Anywhere registration with Round-robin SRO scrutiny allocation.",
            "",
            "Department flow: allocation to DEO; biometrics; Aadhaar auth scenarios (Aadhaar Authentication Requirement note);",
            "Register / Keep Pending / Withdraw / Refuse / Send back; Index I–II; J-slip to Bhoomi; scan & digital sign.",
            "",
            "Extended: Re-registration on DR/Court order; K1 pending release; Will after death (Book 3); Book-I filings; exceptional cases (delay condonation, private attendance, undervaluation/impound, PoA, etc.);",
            "Court Order, Liability, Cross Reference, Sec. 68(2) Correction, Filing Document supporting modules.",
            "",
            "Boundary: EC (IS-05), CC (IS-06), Firm filing at DRO (IS-07), Digital e-Stamp issue (IS-08) are separate.",
        ],
    ),
    (
        "IS-04",
        "Market Valuation",
        [
            "Source: FRS – Kaveri 2.0 (Market Valuation / Fee Calculation); GIS based valuation FRS v2.1;",
            "UPOR FRS; related valuation / fee notes in FRS Documents.",
            "",
            "Guideline / market value determination for stamp duty and registration fee.",
            "Agricultural properties – GIS-based automated valuation via Bhoomi / KSRSAC attribute fetch (GIS FRS v2.1); reduce manual road/zone discretion.",
            "Non-agricultural – existing market valuation logic retained unless separately changed.",
            "Citizen / SR ability to estimate EMV, Stamp Duty and Registration Fee; Form-1 / schedule based valuation at SRO.",
            "Article-based fee calculation integrated with Document Registration and Digital e-Stamp.",
            "",
            "Boundary: Valuation supports IS-03 and IS-08; CVC guideline value master / fixation process as per department valuation policy reflected in GIS/UPOR FRSs.",
        ],
    ),
    (
        "IS-05",
        "Encumbrance Certificate",
        [
            "Source: FRS – Kaveri 2.0 v0.2.38 (Encumbrance Certificate); KaveriArchitecture (EC volume / K1 EC).",
            "",
            "EC shows registered transactions over a property (Form 22 application; Form 15 / nil Form 16).",
            "Search – property parameters (district / taluk / hobli / village / survey / property no. / extent / boundaries); seller / purchaser; Blockchain key where applicable;",
            "Kaveri-1 EC retrieval via Kaveri-1 integration.",
            "",
            "Online application → payment → digitally signed EC (SDA / FDA / SR digital sign chain).",
            "Offline / department path at SRO; Index-II based generation.",
            "Transaction EC printed as part of Document Registration completion (Form 15).",
            "Suo Moto EC payment mode reference for related department searches.",
            "",
            "Boundary: Certified Copy is IS-06. EC search of classified Book-3/4 documents follows Suo Moto rules under IS-06 where applicable.",
        ],
    ),
    (
        "IS-06",
        "Certified Copy",
        [
            "Source: FRS – Kaveri 2.0 v0.2.38 (Certified Copy); Note on Suo Moto CC (Book-3 and Book-4).",
            "",
            "Digitally signed copy of registered scanned document (Book-1 and related book types; Marriage / Marriage Notice as applicable).",
            "Search by Document details – Document Type, District, SRO, Book Type (Book-1 parts, Rule 17(2)/(3), etc.), Document Number, Year.",
            "Online application → stamp duty payment (Sec. 10A / Appendix BA) → SDA / FDA / SR digital sign → citizen download.",
            "",
            "Suo Moto CC (Book-3 / Book-4): citizen presents original at SRO; department CC search without citizen self-service for classified books;",
            "SDA prepare & sign → FDA verify → SR approve & sign → print & deliver; payment mode same as Suo Moto EC.",
            "",
            "Boundary: EC is IS-05. Firm Form-C / register extracts also available under Firm module (IS-07).",
        ],
    ),
    (
        "IS-07",
        "Firm Registration",
        [
            "Source: Firm registration FRS V1.7.",
            "",
            "Services:",
            "New Firm Registration",
            "Amendment – firm name / principal place; other places of business; partner name / address; notice of election on attaining majority",
            "Reconstitution – incoming / outgoing partner",
            "Dissolution of Firm",
            "Protest Entry; Index Search; Certified Copy of Register of Firms entries",
            "",
            "Online and Offline at DRO; approval by District Registrar.",
            "Two-step: partnership / reconstitution / dissolution deed at SRO (IS-03) then firm filing at DRO (Indian Partnership Act, 1932).",
            "Name uniqueness / objectionable names (Emblems and Names Act); fee, Form C, Register of Firms & Index; DigiLocker e-Sign; SAKALA; Aadhaar dependency; K1→K2 firm migration.",
            "",
            "Boundary: Deed registration under IS-03. Society Registration out of this item. K1 Form C shown as-is; full Form I–VI / Form A for K1 firms not available (Firm FRS limitation).",
        ],
    ),
    (
        "IS-08",
        "Digital E-Stamp",
        [
            "Source: Digital e-stamp FRS v1 3.",
            "",
            "Digital e-Stamp for optionally registrable articles (stamp duty payable; registration under Registration Act not compulsory).",
            "Citizen applies on Kaveri 2.0; stamp duty calculated from article / property / location; pay via UPI / net banking / cards;",
            "Instant unique e-Stamp certificate; download digitally signed e-Stamp from dashboard.",
            "",
            "Two-level article mapping (Nature of document → Stamp Sub-article); only optionally registrable articles;",
            "Schedule selection (Agricultural / Non-Agricultural / Miscellaneous); property validations (Bhoomi / e-Aasthi / e-Swathu / ULMS);",
            "Party details with Aadhaar e-KYC / OTP; name matching (C-DAC threshold); Mobile OTP login + e-KYC mandatory.",
            "",
            "Boundary: Mandatorily registrable articles remain under Document Registration (IS-03). Physical stamp paper / franking outside digital e-Stamp path.",
        ],
    ),
    (
        "IS-09",
        "Accounts",
        [
            "Source: Accounts Module FRS V.01.",
            "",
            "Internal accounting for IGR / DR / SR – remittance to Head of Account, bank challan and treasury reconciliation.",
            "Types: Add Funds (IGR – HOA, Quarter, Amount, Upload Order);",
            "Allocate Funds (IGR→District; DR→SRO);",
            "Expenditure (IGR / DR / SR);",
            "Reconciliation (SR / DR with treasury / bank).",
            "",
            "SRO daily accounting → DRO consolidation → IGR; statements prepared and remitted.",
            "Boundary: Citizen payment capture remains in service modules (IS-03/05/06/07/08); Khajane gateway under IS-11.",
        ],
    ),
    (
        "IS-10",
        "Audit",
        [
            "Source: Audit Module FRS V.02.",
            "",
            "Internal Audit – DR / IGR / Regional Committee / HQA / Deputy Commissioner inspections (DR 50%, HQA 100%, IGR one office per district).",
            "External Audit – Auditor General / AG office annual inspection; Local Audit Report; Audit Para query & response.",
            "AG / auditor login creation by office head (SR / DR / division head); AG email + ID card; period-based fetch of registered documents;",
            "Request to view document details → office head approval → digitally accessible audit trail instead of manual consolidation.",
            "",
            "Boundary: Separate login for audit on KAVERI reports portal as specified. Does not replace statutory AG process outside the application.",
        ],
    ),
    (
        "IS-11",
        "External Integrations",
        [
            "Source: KaveriArchitecture.pdf (Integrations); FRS cross-module integration notes.",
            "",
            "Kaveri 3.0 shall connect to send/receive data with (as deployed / available):",
            "Bhoomi & Mojini – agri property pull; mutation push.",
            "eSwathu / RDPR – rural non-agri pull; mutation / eKhata push.",
            "eAasthi & BBMP – urban property pull; mutation / eKhata push.",
            "ULMS – non-eAasthi urban properties.",
            "FRUITS / eSaala – agri mortgage / release filing; 1-day EC push.",
            "Kaveri 1 – EC/CC and pending/completed registration data.",
            "Khajane – Payment Gateway.",
            "eKYC & eSign (Aadhaar / DigiLocker / CDAC as applicable).",
            "Sakala – service timelines; Scality – document store; SMS; Email.",
            "KGIS – map layer; NGDRS – user details push on need basis.",
            "UPOR / PAN / NESL / other FRS-noted integrations as approved.",
            "",
            "P.S.: External systems are owned by respective agencies; Kaveri integrates via APIs / scheduled jobs.",
        ],
    ),
    (
        "IS-12",
        "Reporting / Dashboard",
        [
            "Source: FRS – Kaveri 2.0 MIS sections; Department Login & dashboard; KaveriArchitecture (reporting / batch); Audit reports portal.",
            "",
            "Citizen – application status tracking, download certificates / e-Stamp / EC / CC.",
            "Department dashboards – SR / DR / IGR worklists (scrutiny, registration, EC/CC approval, firm, audit requests).",
            "MIS – registrations by article / book / office; revenue by mode/channel; pending / refused / withdrawn / undervaluation;",
            "Marriage / Firm statistics; SLA / Sakala oriented views; bilingual labels where required.",
            "Scheduled / batch jobs for reconciliation and reporting (Architecture).",
            "",
            "Boundary: Operational observability metrics are under IS-16; Audit Para workflows under IS-10.",
        ],
    ),
    (
        "IS-13",
        "Security",
        [
            "Source: KaveriArchitecture.pdf; Software Design Document – Kaveri2.0 v2.2; User management FRS; Aadhaar Authentication Requirement note.",
            "",
            "Reverse proxy – SSL offloading, rate limiting, bot blocking, load balancing.",
            "API Gateway – authN/Z on every request, IP/route rate limits, Redis session, request signature validation (WIP→Prod).",
            "RBAC via User Management (IS-01); department interface only on KSWAN.",
            "Aadhaar e-KYC / e-Sign; consent refusal / service-down / mismatch handling per Aadhaar note.",
            "Digital signatures on registered documents, EC, CC, e-Stamp, firm certificates.",
            "Reduced attack surface by not exposing internal microservices directly (API Gateway pattern – SDD).",
            "",
            "Boundary: Infrastructure WAF/SDC controls owned by SDC; application implements gateway and service-level controls.",
        ],
    ),
    (
        "IS-14",
        "Environment & Deployment",
        [
            "Source: KaveriArchitecture.pdf §5 Deployment and Environments; Software Design Document.",
            "",
            "Stack: Angular SPA (citizen & department UI); .NET Core / Python / Go services; PostgreSQL; Redis cache.",
            "Environments: Development (Azure) → Staging (SDC) → Production (SDC).",
            "Deployment units – services, web app, DB migrations; automated Dev pipelines; Staging via scripts/Jenkins/Portainer; Prod via scripts (promote Portainer/Docker Swarm).",
            "Post-deploy checks – service up, DB connectivity, downstream dependencies (gateway/cache).",
            "Rollback – retain stable artifacts; DB rollback scripts mandatory; low downtime tolerance for user-facing workflows.",
            "",
            "Boundary: SDC hosting and WAF HTTP/2 issues resolved with SDC/WAF vendor (Architecture known item).",
        ],
    ),
    (
        "IS-15",
        "Database and Migration",
        [
            "Source: Design Documents/Database Design (kaveri.sql, kaverimig, kaverimis, kavericdc, kaveridelapp, kaverimrg, kaveri_analytics, ERDs);",
            "KaveriArchitecture (shared DB, replicas, legacy schema); Firm FRS (K1→K2 firm migration).",
            "",
            "Shared PostgreSQL across services with logical domain grouping (document, marriage, e-Stamp, firm, etc.); read replicas for offload.",
            "Schemas / DBs for registration, migration, MIS, CDC, analytics, backup, marriage, firm, e-stamp, K1 as per Design ERDs.",
            "Kaveri-1 compatibility – pending documents, EC/CC, firm Form C limitations; migration and CDC paths via kaverimig / kavericdc.",
            "Schema evolution with impact analysis; data ownership ideally per service (Architecture focus / known shared-DB limitation).",
            "",
            "Boundary: Physical SDC DB infrastructure owned by ops; application owns schema migrations and data contracts.",
        ],
    ),
    (
        "IS-16",
        "Observability, Monitoring & Testing",
        [
            "Source: KaveriArchitecture (observability across services/environments; ops scale); QA Files/Test_Cases and Release_Notes;",
            "KaveriMaintenanceServer.pdf (maintenance/ops).",
            "",
            "Observability present across services and environments; monitor latency-sensitive workflows and peak concurrency (~4.5–5k concurrent).",
            "Transaction scale awareness (Doc Reg / EC / CC volumes per Architecture).",
            "Testing in scope: functional / regression coverage aligned to CSG test cases – User Registration, Document Registration,",
            "Document Info & Property Search, SRO Confirmation, Paperless, Firm (New/Amendment/Reconstitution), Marriage Online/Offline, and related module suites;",
            "UAT / Prod release notes process under QA Files/Release_Notes.",
            "",
            "Boundary: Infrastructure monitoring tools at SDC may be shared; application instrumentation and QA evidence are in scope.",
        ],
    ),
    (
        "IS-17",
        "AI",
        [
            "Source: No dedicated AI FRS / design baseline found under CSG Documents as of this scope version.",
            "",
            "AI capabilities for Kaveri 3.0 are not defined in the current CSG FRS / Design document set.",
            "Any AI features (assistive drafting, classification, anomaly detection, etc.) require a separate approved CSG requirement note / FRS before inclusion in delivery scope.",
            "",
            "Boundary: Out of baseline until CSG publishes AI requirements. Existing C-DAC name-matching and rule-based valuations are not classified as AI modules.",
        ],
    ),
]


def set_cell_lines(cell, lines: list[str]) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    if not lines:
        first.add_run("")
        return
    first.add_run(lines[0])
    for line in lines[1:]:
        cell.add_paragraph(line)


def set_cell_text(cell, text: str) -> None:
    set_cell_lines(cell, [text] if text else [""])


def ensure_rows(table, needed: int) -> None:
    """Ensure table has header + needed data rows."""
    while len(table.rows) < needed + 1:
        tbl = table._tbl
        last = table.rows[-1]._tr
        tbl.append(deepcopy(last))


def main() -> None:
    base = next((p for p in SRC_CANDIDATES if p.exists()), None)
    if not base:
        raise SystemExit("No base scope document found")

    shutil.copy2(base, OUT)
    doc = Document(str(OUT))
    table = doc.tables[0]
    ensure_rows(table, len(SCOPE_ROWS))

    # Clear any extra rows beyond SCOPE_ROWS
    while len(table.rows) > len(SCOPE_ROWS) + 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)

    for i, (sid, item, lines) in enumerate(SCOPE_ROWS, start=1):
        row = table.rows[i]
        set_cell_text(row.cells[0], sid)
        set_cell_text(row.cells[1], item)
        set_cell_lines(row.cells[2], lines)

    doc.save(str(OUT))

    # Sync Template1 if writable
    tpl = SCOPE_DIR / "IT_Project_Scope_Document_Template1.docx"
    try:
        shutil.copy2(OUT, tpl)
        tpl_msg = f"Updated: {tpl}"
    except PermissionError:
        tpl_msg = f"Skipped Template1 (locked): {tpl}"

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, CSG_OUT)

    # Verify
    v = Document(str(OUT))
    print(f"Rows (incl header): {len(v.tables[0].rows)}")
    for row in v.tables[0].rows[1:]:
        sid = row.cells[0].text.strip()
        item = row.cells[1].text.strip()
        desc_len = len(row.cells[2].text.strip())
        print(f"  {sid:6} {item:35} {desc_len:5} chars")

    print(f"\nSaved: {OUT}")
    print(f"Copied: {CSG_OUT}")
    print(tpl_msg)


if __name__ == "__main__":
    main()
