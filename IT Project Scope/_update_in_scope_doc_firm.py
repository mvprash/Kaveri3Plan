# -*- coding: utf-8 -*-
"""Update In Scope IS-03 (Document Registration) and IS-06 (Firm Registration)
from RFP/Document&FirmRegistration.doc and Acts_Rules/Document; write new version
under CSG Documents.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SRC = ROOT / "IT Project Scope" / "IT_Project_Scope_Document_Template1.docx"
OUT_DIR = ROOT / "CSG Documents" / "IT Project Scope"
OUT = OUT_DIR / "IT_Project_Scope_Document_v1.1.docx"

# Match Marriage (IS-02) style: short lines / bullets as separate paragraphs
DOC_REG_LINES = [
    "2 types of registration flows – Paper and Paperless.",
    "Paper – ink-signed deed, physical presentation, scan & store after registration.",
    "Paperless – digital execution (e-Sign / DSC), biometric e-KYC and digital endorsements.",
    "",
    "Application has to support both Online and Offline Process.",
    "Online – via Kaveri Application / Portal / BangaloreOne, Karnataka One, Nemmadi.",
    "Offline – at SRO Kiosk; data entry and registration by SRO / FDA / SDA.",
    "",
    "Core capabilities (RFP MOD 01):",
    "EMV, Stamp Duty (SD) and Registration Fee (RF) estimation; valuation / guidance slip.",
    "Application capture (e-Form / upload), appointment booking and service token.",
    "SR scrutiny checklist – classification, property schedule, testimonials, court orders, Index-II duplicate check.",
    "Payment capture (cash / DD / challan / e-Stamp) with receipt; fine / penalty / delay condonation where applicable.",
    "Photograph and thumb impression of parties / witnesses; Aadhaar authentication where applicable.",
    "Summary report, registration endorsement and Section 60 certificate of registration.",
    "Unique registration number; auto entries in A-Register, Register Books and Index I–IV.",
    "Scan, digital sign and upload of registered document; delivery acknowledgement.",
    "Outcomes – Admit / Refuse (Book II) / Keep Pending (Minute Book) / Withdraw / Impound (undervaluation Sec. 45-A).",
    "Cross-link amending / cancelling documents (Rule 123); PoA presentation; private attendance / commission; memo under Sec. 64–67; re-registration on DR / Court order.",
    "J-slip / property mutation intimation to Bhoomi / ULB / UPOR; e-Stamp validation.",
    "",
    "Legal basis (Acts_Rules/Document):",
    "The Registration Act, 1908 (as amended in Karnataka) – presentation, enquiry, endorsement, certificate, books & indexes.",
    "The Karnataka Registration Rules, 1965.",
    "The Karnataka Stamp Act, 1957 (including e-Stamping and undervaluation / Sec. 45-A).",
    "",
    "Boundary: Document registration at SRO / DRO. Encumbrance Certificate (IS-04) and Certified Copy (IS-05) are separate In-Scope items.",
]

FIRM_REG_LINES = [
    "Firm services (RFP MOD 03 / Indian Partnership Act, 1932):",
    "New Firm Registration",
    "Amendment – firm name / principal place of business; other places of business; partner name / address; notice of election on attaining majority",
    "Reconstitution – incoming / outgoing partner",
    "Dissolution of Firm",
    "Protest Entry against Register of Firms entry",
    "Index Search of registered firms",
    "Certified Copy of Register of Firms entries / related forms",
    "",
    "Application has to support both Online and Offline Process.",
    "Online – via Kaveri Application / Portal / BangaloreOne, Karnataka One, Nemmadi.",
    "Offline – at District Registrar Office (DRO); registration / approval by District Registrar.",
    "",
    "Two-step linkage with Document Registration (IS-03):",
    "1. Partnership / reconstitution / dissolution deed registered at SRO (Karnataka Stamp Act – partnership instruments).",
    "2. Firm filing at DRO under the Indian Partnership Act, 1932.",
    "",
    "Core capabilities:",
    "Firm name and address uniqueness search; block duplicate / objectionable names.",
    "Fee calculation; acknowledgement and fee receipt.",
    "Unique Firm Registration Number; update Register of Firms and its Index.",
    "Print Certificate of Registration / endorsements (e.g. Form C) as per application type.",
    "Scan and upload of application attachments (Form I, partnership deed copy, declaration, ID / address proof).",
    "Automatic A-Register accounting entry; MIS reports; filing of returns / information via portal or offline at DRO.",
    "DR authentication via digital signature / biometric as prescribed.",
    "",
    "Legal basis (Acts_Rules/Document + Partnership law):",
    "Indian Partnership Act, 1932.",
    "The Karnataka Stamp Act, 1957 – stamp duty on partnership / reconstitution / dissolution deeds.",
    "The Registration Act, 1908 and The Karnataka Registration Rules, 1965 – deed registration at SRO.",
    "",
    "Boundary: Firm module at DRO. Deed registration itself is under IS-03 Document Registration. Society Registration is out of this item.",
]


def set_cell_lines(cell, lines: list[str]) -> None:
    """Replace cell content with one paragraph per line (preserves first para formatting)."""
    # Ensure at least one paragraph exists
    if not cell.paragraphs:
        cell.add_paragraph()

    # Clear all but first paragraph
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)

    first = cell.paragraphs[0]
    # Clear runs in first paragraph
    for r in list(first.runs):
        r._element.getparent().remove(r._element)

    if lines:
        first.add_run(lines[0])
        for line in lines[1:]:
            cell.add_paragraph(line)
    else:
        first.add_run("")


def find_row(table, scope_id: str):
    for row in table.rows:
        if row.cells[0].text.strip() == scope_id:
            return row
    raise KeyError(scope_id)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)

    doc = Document(str(OUT))
    table = doc.tables[0]

    row_doc = find_row(table, "IS-03")
    set_cell_lines(row_doc.cells[2], DOC_REG_LINES)

    row_firm = find_row(table, "IS-06")
    set_cell_lines(row_firm.cells[2], FIRM_REG_LINES)

    doc.save(str(OUT))

    # Also refresh Template1 working copy so it stays aligned
    shutil.copy2(OUT, SRC)

    # Verify
    verify = Document(str(OUT))
    t = verify.tables[0]
    for sid in ("IS-03", "IS-06"):
        for row in t.rows:
            if row.cells[0].text.strip() == sid:
                text = row.cells[2].text
                print(f"=== {sid} ({len(text)} chars, {text.count(chr(10))+1} lines) ===")
                print(text[:500])
                print("...")
                break

    print(f"\nSaved: {OUT}")
    print(f"Updated: {SRC}")


if __name__ == "__main__":
    main()
