# -*- coding: utf-8 -*-
"""Create IT_Project_Scope_Document_v1.6 from v1.5 — strengthen IS-13 Security with HTTPS/SSL."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC = SCOPE_DIR / "IT_Project_Scope_Document_v1.5.docx"
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v1.6.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v1.6.docx"

IS13_LINES = [
    "HTTPS enforced for all citizen and department web / API traffic (TLS end-to-end to the edge).",
    "SSL / TLS certificates managed for public and internal endpoints; reverse proxy performs SSL offloading where deployed.",
    "Reverse proxy – SSL offloading, rate limiting, bot blocking, load balancing.",
    "",
    "API Gateway – authN/Z on every request, IP/route rate limits, Redis session,",
    "RBAC via User Management; department interface only on KSWAN.",
    "Digital signatures (e-Sign/DSC) on registered documents, EC, CC, e-Stamp, firm certificates.",
    "Reduced attack surface by not exposing internal microservices directly (API Gateway pattern).",
]


def set_cell_lines(cell, lines: list[str]) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    if not lines:
        first.add_run("")
        return
    first.add_run(lines[0])
    for line in lines[1:]:
        cell.add_paragraph(line)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    table = doc.tables[0]

    found = False
    for row in table.rows:
        if row.cells[0].text.strip() == "IS-13":
            set_cell_lines(row.cells[2], IS13_LINES)
            found = True
            break
    if not found:
        raise SystemExit("IS-13 not found")

    doc.save(str(OUT))

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(OUT, CSG_OUT)
        csg_msg = f"Copied: {CSG_OUT}"
    except PermissionError:
        csg_msg = f"Skipped CSG copy (locked): {CSG_OUT}"

    verify = Document(str(OUT))
    for row in verify.tables[0].rows:
        if row.cells[0].text.strip() == "IS-13":
            print(f"Created: {OUT}")
            print(csg_msg)
            print("--- IS-13 Security ---")
            print(row.cells[2].text)
            break


if __name__ == "__main__":
    main()
