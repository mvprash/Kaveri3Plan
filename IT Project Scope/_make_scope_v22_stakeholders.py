# -*- coding: utf-8 -*-
"""Create IT_Project_Scope_Document_v2.2 — add missing partner stakeholders."""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC = SCOPE_DIR / "IT_Project_Scope_Document_v2.1.docx"
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v2.2.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v2.2.docx"

NAVY = "1F4E79"
NAVY_RGB = RGBColor(0x1F, 0x4E, 0x79)
ALT_ROW = "F2F6FA"
WHITE = "FFFFFF"
BORDER = "8FA3B8"
BODY_RGB = RGBColor(0x24, 0x24, 0x24)

# Insert after matching stakeholder name (exact first-cell text)
# Order of tuples = insert sequence (each after its anchor on the growing table)
NEW_ROWS: list[tuple[str, tuple[str, str, str, str, str]]] = [
    # After Citizens — professional facilitators
    (
        "Citizens / Applicants",
        (
            "Advocate / Deed Writers",
            "Legal practitioners and deed writers facilitating citizen registration",
            "Accurate deed drafting, party guidance and presentation support for Document Registration and related services",
            "Workshops / circulars; portal guidance; UAT sampling where workflows touch facilitators",
            "None (service facilitators; statutory acts remain with SR / parties)",
        ),
    ),
    # After Eaasthi — ULMS property system
    (
        "Eaasthi",
        (
            "ULMS",
            "Urban land / property management system (non-eAasthi urban properties)",
            "Property pull and mutation support for urban properties not covered by eAasthi",
            "Sandbox SIT; integration war-room; mutation reconciliation",
            "Interface SLAs and change windows (jointly with DSR / ULB)",
        ),
    ),
    # After Income Tax — partner organisations (before IT Cell)
    (
        "Income Tax",
        (
            "NIC",
            "National Informatics Centre — custodian / partner for eSwathu, eAasthi and related state platforms",
            "Availability and change control of NIC-hosted property and citizen platforms used by Kaveri integrations",
            "Integration war-rooms; change windows; joint SIT with RDPR / UDD / DSR",
            "NIC platform SLAs and change windows (jointly with DSR)",
        ),
    ),
    (
        "NIC",
        (
            "CSG (Kaveri 2.0)",
            "Centre for Smart Governance — legacy Kaveri 2.0 design, FRS and operational knowledge",
            "As-is behaviour, FRS / design baseline, KT for migration and parity with Kaveri 2.0",
            "KT sessions; FRS clarification workshops; migration dual-run support",
            "None on Kaveri 3.0 product decisions (advisory / handover)",
        ),
    ),
    (
        "CSG (Kaveri 2.0)",
        (
            "CEG (Centre for e-Governance)",
            "Centre for e-Governance — Aadhaar e-KYC, e-Sign, DigiLocker, FRUITS, Sakala, ULMS and related e-Gov services",
            "Reliable statewide e-Gov platform services consumed by Kaveri 3.0 integrations",
            "Integration war-rooms; sandbox SIT; SLA / change-window coordination",
            "e-Gov platform SLAs and change windows (jointly with DSR)",
        ),
    ),
    (
        "CEG (Centre for e-Governance)",
        (
            "CMS (System Integrator for Kaveri 2.0)",
            "Current system integrator and O&M partner for Kaveri 2.0",
            "Production as-is support, data extracts, dual-run and cutover coordination during Kaveri 3.0 transition",
            "KT / war-rooms; migration freeze windows; hypercare handoff",
            "None on Kaveri 3.0 build decisions (transition support within agreed SI scope)",
        ),
    ),
    (
        "CMS (System Integrator for Kaveri 2.0)",
        (
            "Kaveri Call Center",
            "Citizen / department helpdesk for Kaveri services",
            "First-line support, ticket triage and citizen guidance across modules",
            "SOP / training before each phase Go-Live; L1–L2 handoff; hypercare staffing",
            "None (operational support within published SOPs)",
        ),
    ),
]


def set_run_font(run, *, name="Calibri", size_pt=8, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def set_cell_borders(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)


def clear_paragraph(para) -> None:
    for r in list(para.runs):
        r._element.getparent().remove(r._element)


def rewrite_cell(cell, text: str, *, bold=False, size_pt=8) -> None:
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    para = cell.paragraphs[0]
    clear_paragraph(para)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    set_run_font(run, size_pt=size_pt, bold=bold, color=NAVY_RGB if bold else BODY_RGB)


def style_body_row(row, alt: bool) -> None:
    fill = ALT_ROW if alt else WHITE
    for ci, cell in enumerate(row.cells):
        set_cell_shading(cell, fill)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        for para in cell.paragraphs:
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)


def insert_row_after(table, after_index: int) -> int:
    src_tr = table.rows[after_index]._tr
    new_tr = deepcopy(src_tr)
    src_tr.addnext(new_tr)
    return after_index + 1


def find_row_index(table, name: str) -> int:
    for i, row in enumerate(table.rows):
        if row.cells[0].text.strip() == name:
            return i
    raise SystemExit(f"Anchor row not found: {name!r}")


def find_stakeholders_table(doc: Document):
    for table in doc.tables:
        if not table.rows:
            continue
        h = " | ".join(c.text.strip() for c in table.rows[0].cells).lower()
        if "stakeholder" in h and "role" in h:
            return table
    raise SystemExit("Stakeholders table not found")


def update_version_banner(doc: Document) -> None:
    for para in doc.paragraphs:
        if para.text.strip().startswith("Document version"):
            clear_paragraph(para)
            run = para.add_run(
                "Document version 2.2  ·  In Scope (with Actors) & Key Stakeholders"
            )
            set_run_font(run, size_pt=9, bold=False, color=RGBColor(0x5A, 0x6A, 0x7A))
            para.paragraph_format.space_after = Pt(12)
            return


def reshade_all_data_rows(table) -> None:
    for ri, row in enumerate(table.rows[1:], start=1):
        style_body_row(row, alt=(ri % 2 == 0))


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    update_version_banner(doc)
    table = find_stakeholders_table(doc)

    for anchor, row_data in NEW_ROWS:
        # Skip if already present
        existing = {r.cells[0].text.strip() for r in table.rows}
        if row_data[0] in existing:
            print(f"Skip (already present): {row_data[0]}")
            continue
        idx = find_row_index(table, anchor)
        new_idx = insert_row_after(table, idx)
        for ci, val in enumerate(row_data):
            if ci < len(table.rows[new_idx].cells):
                rewrite_cell(table.rows[new_idx].cells[ci], val, bold=(ci == 0))

    reshade_all_data_rows(table)
    doc.save(str(OUT))

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(OUT, CSG_OUT)
        csg_msg = f"Copied: {CSG_OUT}"
    except PermissionError:
        csg_msg = f"Skipped CSG copy (locked): {CSG_OUT}"

    v = Document(str(OUT))
    t = find_stakeholders_table(v)
    print(f"Created: {OUT}")
    print(csg_msg)
    print(f"Stakeholder rows: {len(t.rows) - 1}")
    added = {
        "Advocate / Deed Writers",
        "NIC",
        "CSG (Kaveri 2.0)",
        "CEG (Centre for e-Governance)",
        "ULMS",
        "CMS (System Integrator for Kaveri 2.0)",
        "Kaveri Call Center",
    }
    names = [r.cells[0].text.strip() for r in t.rows[1:]]
    for i, n in enumerate(names, 1):
        mark = " *" if n in added else ""
        print(f"  {i}. {n}{mark}")


if __name__ == "__main__":
    main()
