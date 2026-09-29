# -*- coding: utf-8 -*-
"""Create IT_Project_Scope_Document_v1.5 from v1.4_stakeholders_updated.

Change: replace the single 'External integration owners' row with one row
per external integration (from In Scope IS-11 in the same document).
"""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC = SCOPE_DIR / "IT_Project_Scope_Document_v1.4_stakeholders_updated.docx"
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v1.5.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v1.5.docx"

# One row per integration — aligned to IS-11 External Integrations
# Columns: Stakeholder / Group | Role | Interest / impact | Engagement method | Decision authority
INTEGRATION_ROWS: list[tuple[str, str, str, str, str]] = [
    (
        "e-KYC / UIDAI (Aadhaar)",
        "Aadhaar e-KYC / Face Authentication partner",
        "Citizen and party identity verification for registration and related services",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and authN change windows (jointly with DSR)",
    ),
    (
        "e-Sign / DSC provider",
        "Aadhaar-based eSign and digital signature services",
        "Citizen eSign and officer DSC on certificates, deeds, EC/CC and e-Stamp artefacts",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and signing-service change windows (jointly with DSR)",
    ),
    (
        "DigiLocker",
        "Document wallet / issued-document partner",
        "Fetch / push of citizen documents and issued certificates where enabled",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "CDAC",
        "Name-matching / supporting e-KYC–eSign technical services (as applicable)",
        "Reliable matching and related technical services used by registration flows",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "Khajane",
        "Payment gateway (Government of Karnataka treasury)",
        "Fee / stamp / registration payment capture and reconciliation",
        "Sandbox SIT; integration war-room; reconciliation reviews",
        "Payment interface SLAs and change windows (jointly with DSR / Treasury)",
    ),
    (
        "Sakala",
        "Karnataka Guarantee of Services to Citizen platform",
        "Service timelines, GSC issuance and statutory status sync",
        "Sandbox SIT; integration war-room; SLA reviews",
        "Service-timeline interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "SMS gateway",
        "Transactional / OTP SMS provider",
        "OTP, status and notification delivery to citizens and officers",
        "Sandbox SIT; delivery monitoring; change windows",
        "SMS provider SLAs and template/change windows (jointly with DSR)",
    ),
    (
        "Email gateway",
        "Transactional email provider",
        "OTP (where used), receipts, certificates and status notifications",
        "Sandbox SIT; delivery monitoring; change windows",
        "Email provider SLAs and template/change windows (jointly with DSR)",
    ),
    (
        "Bhoomi & Mojini",
        "Agricultural property and mutation systems",
        "Agri property pull; mutation push after registration",
        "Sandbox SIT; integration war-room; mutation reconciliation",
        "Property/mutation interface SLAs and change windows (jointly with DSR / Revenue)",
    ),
    (
        "eSwathu / RDPR",
        "Rural non-agricultural property systems",
        "Rural non-agri property pull; mutation / eKhata push",
        "Sandbox SIT; integration war-room; mutation reconciliation",
        "Property/mutation interface SLAs and change windows (jointly with DSR / RDPR)",
    ),
    (
        "eAasthi & BBMP",
        "Urban property systems (BBMP / eAasthi)",
        "Urban property pull; mutation / eKhata push",
        "Sandbox SIT; integration war-room; mutation reconciliation",
        "Property/mutation interface SLAs and change windows (jointly with DSR / ULB)",
    ),
    (
        "ULMS",
        "Urban land / property system for non-eAasthi urban properties",
        "Property pull for non-eAasthi urban properties; mutation push where applicable",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR / ULB)",
    ),
    (
        "FRUITS / eSaala",
        "Agricultural mortgage / release filing systems",
        "Agri mortgage / release filing; 1-day EC push",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "Kaveri 1.0",
        "Legacy registration / EC / CC data source",
        "EC/CC and pending/completed registration data for migration and dual-run needs",
        "Migration workstream; dual-run / cutover windows; reconciliation",
        "Cutover and freeze windows (jointly with DSR / Steering)",
    ),
    (
        "Scality",
        "Document / object store",
        "Storage and retrieval of scanned deeds, certificates and related artefacts",
        "Ops coordination; capacity and change windows",
        "Storage capacity and change windows (jointly with DSR / SDC)",
    ),
    (
        "KGIS",
        "Karnataka GIS map-layer provider",
        "Map layer support for GIS / valuation related flows",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "NGDRS",
        "National Generic Document Registration System (as applicable)",
        "User / registration detail push on need basis",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "UPOR",
        "Urban property / ownership records integration (as approved)",
        "Property / ownership data exchange where FRS-approved",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "PAN validation / SFT",
        "PAN verification and SFT-related services (as approved)",
        "PAN validation along with statutory reporting where mandated",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "NESL",
        "National E-Governance Services Ltd. / related API (as approved)",
        "NESL API exchange for approved document / liability use cases",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
]


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    first.add_run(text or "")


def find_stakeholders_table(doc: Document):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " | ".join(c.text.strip() for c in table.rows[0].cells).lower()
        if "stakeholder" in header and "role" in header:
            return table
    return None


def insert_row_after(table, after_index: int):
    """Clone the row at after_index and insert the clone immediately after it."""
    src_tr = table.rows[after_index]._tr
    new_tr = deepcopy(src_tr)
    src_tr.addnext(new_tr)
    # python-docx table.rows is refreshed from XML; return new row index
    return after_index + 1


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    table = find_stakeholders_table(doc)
    if table is None:
        raise SystemExit("Stakeholders table not found")

    # Find External integration owners row
    ext_idx = None
    for i, row in enumerate(table.rows):
        name = row.cells[0].text.strip().lower()
        if name.startswith("external integration"):
            ext_idx = i
            break
    if ext_idx is None:
        raise SystemExit("Could not find 'External integration owners' row")

    # Replace that one row with N integration rows
    # First overwrite the existing row with INTEGRATION_ROWS[0]
    # Then insert additional rows after it for the rest
    first = INTEGRATION_ROWS[0]
    for ci, val in enumerate(first):
        set_cell_text(table.rows[ext_idx].cells[ci], val)

    insert_at = ext_idx
    for row_data in INTEGRATION_ROWS[1:]:
        insert_at = insert_row_after(table, insert_at)
        for ci, val in enumerate(row_data):
            set_cell_text(table.rows[insert_at].cells[ci], val)

    doc.save(str(OUT))

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(OUT, CSG_OUT)
        csg_msg = f"Copied: {CSG_OUT}"
    except PermissionError:
        csg_msg = f"Skipped CSG copy (locked): {CSG_OUT}"

    # Verify
    verify = Document(str(OUT))
    vt = find_stakeholders_table(verify)
    print(f"Created: {OUT}")
    print(csg_msg)
    print(f"Stakeholder rows (excl. header): {len(vt.rows) - 1}")
    print("--- Stakeholders ---")
    for i, row in enumerate(vt.rows[1:], 1):
        print(f"  {i}. {row.cells[0].text}")


if __name__ == "__main__":
    main()
