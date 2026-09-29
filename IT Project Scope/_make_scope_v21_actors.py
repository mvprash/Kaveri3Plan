# -*- coding: utf-8 -*-
"""Create IT_Project_Scope_Document_v2.1 from v2.0 — add Actors per scoped item."""
from __future__ import annotations

import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC = SCOPE_DIR / "IT_Project_Scope_Document_v2.0.docx"
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v2.1.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v2.1.docx"

NAVY = "1F4E79"
NAVY_RGB = RGBColor(0x1F, 0x4E, 0x79)
ALT_ROW = "F2F6FA"
WHITE = "FFFFFF"
BORDER = "8FA3B8"
BODY_RGB = RGBColor(0x24, 0x24, 0x24)

# Actors per scope ID — only where end-user / business / ops actors apply
ACTORS: dict[str, list[str]] = {
    "IS-03": [
        "Citizen / Applicant (parties, witnesses, PoA holder)",
        "DEO",
        "FDA / SDA",
        "Sub-Registrar (SR / SRO)",
        "District Registrar (DRO) — undervaluation, re-registration, adjudication",
        "IGRO — appeal (where applicable)",
    ],
    "IS-04": [
        "Citizen / Applicant (estimate Guidance Value, SD, RF)",
        "Sub-Registrar (SR) — Form-1 / schedule valuation at SRO",
        "CVC / DIGR CVC (where CVC / GIS valuation workflows apply)",
    ],
    "IS-05": [
        "Citizen / Applicant",
        "SDA / FDA",
        "Sub-Registrar (SR) — digital sign chain",
        "DEO (offline / department path at SRO, where used)",
    ],
    "IS-06": [
        "Citizen / Applicant",
        "SDA — prepare & sign",
        "FDA — verify",
        "Sub-Registrar (SR) — approve & sign / digital sign",
        "DEO (Suo Moto / counter path, where used)",
    ],
    "IS-07": [
        "Citizen / Applicant / Partners",
        "Sub-Registrar (SRO) — partnership / reconstitution / dissolution deed",
        "District Registrar (DRO) — firm filing approval",
        "FDA / SDA (DRO office support, where used)",
    ],
    "IS-08": [
        "Citizen / Applicant",
    ],
    "IS-10": [
        "Office head (SR / DR) — AG login creation & document-view approval",
        "DR / IGR / Regional Committee / HQA / Deputy Commissioner — internal audit",
        "Auditor General / AG auditor — external audit",
    ],
    "IS-11": [
        "Citizen / Applicant (consuming integrated services)",
        "DSR Officers (SR / DEO / FDA / SDA / DRO / IGR as per module)",
        "Other Department users (e.g. Income Tax)",
        "External system owners (per integration row in §2)",
    ],
    "IS-12": [
        "Citizen / Applicant — status & downloads",
        "SR / DR / IGR — worklists & MIS",
        "AIGR / Admin — executive / statutory dashboards",
        "BI / Analytics (Kaveri IT Cell) — Power BI packs",
    ],
    "IS-13": [
        "Application Admin / KPMU — RBAC & access policy",
        "Security Specialist (Kaveri IT Cell)",
        "DevOps & Release Manager",
        "SDC / Infrastructure operations",
    ],
    "IS-14": [
        "DevOps & Release Manager",
        "Kaveri IT Cell delivery team",
        "SDC / Infrastructure operations",
    ],
    "IS-15": [
        "Database Administrator",
        "Data Migration Specialist",
        "Domain Expert (cutover / reconciliation sign-off)",
        "Product Owner / Steering (freeze & Go-Live cutover)",
    ],
    "IS-16": [
        "QA / Test Engineers",
        "Performance & Security Test Lead",
        "DevOps (observability stack)",
        "L2 Support (hypercare monitoring)",
    ],
    "IS-17": [
        "Not applicable — AI scope not defined; actors TBD when requirements are approved",
    ],
}


def set_run_font(run, *, name="Calibri", size_pt=10, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill_hex: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def set_cell_borders(cell, color_hex: str = BORDER, sz: str = "4") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color_hex)


def set_cell_margins(cell, top=40, bottom=40, left=60, right=60) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = tcMar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcMar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


def set_row_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


def set_table_fixed_layout(table) -> None:
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_col_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for i, w_cm in enumerate(widths_cm):
            if i >= len(row.cells):
                continue
            tcPr = row.cells[i]._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(w_cm * 567)))
            tcW.set(qn("w:type"), "dxa")


def clear_paragraph(para) -> None:
    for r in list(para.runs):
        r._element.getparent().remove(r._element)


def rewrite_cell_plain(cell, text: str, *, bold=False, size_pt=9, center=False) -> None:
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    para = cell.paragraphs[0]
    clear_paragraph(para)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    set_run_font(run, size_pt=size_pt, bold=bold, color=BODY_RGB)


def rewrite_cell_lines(cell, lines: list[str], *, size_pt=9, bullet: bool = False) -> None:
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    clear_paragraph(first)

    if not lines:
        return

    for idx, line in enumerate(lines):
        para = first if idx == 0 else cell.add_paragraph()
        clear_paragraph(para)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(2 if line.strip() else 4)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if not line.strip():
            continue

        display = f"• {line}" if bullet and not line.startswith("•") else line
        m = re.match(r"^(.*?)(\s+[—–-]\s+)(.*)$", line)
        if m and len(m.group(1)) < 80 and not bullet:
            r1 = para.add_run(m.group(1))
            set_run_font(r1, size_pt=size_pt, bold=True, color=NAVY_RGB)
            r2 = para.add_run(m.group(2) + m.group(3))
            set_run_font(r2, size_pt=size_pt, bold=False, color=BODY_RGB)
        else:
            r = para.add_run(display)
            set_run_font(r, size_pt=size_pt, bold=False, color=BODY_RGB)


def style_header_cell(cell, text: str) -> None:
    set_cell_shading(cell, NAVY)
    set_cell_borders(cell, color_hex="163A5F", sz="6")
    set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    para = cell.paragraphs[0]
    clear_paragraph(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(text)
    set_run_font(run, size_pt=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))


def style_body_cell(cell, *, alt: bool = False) -> None:
    set_cell_shading(cell, ALT_ROW if alt else WHITE)
    set_cell_borders(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def add_column_to_table(table) -> None:
    """Append one column by cloning the last cell of each row."""
    for row in table.rows:
        last_tc = row.cells[-1]._tc
        new_tc = deepcopy(last_tc)
        # clear text in cloned cell
        for p in new_tc.iterchildren(qn("w:p")):
            for r in list(p.iterchildren(qn("w:r"))):
                p.remove(r)
            # ensure at least empty paragraph remains
        # if no paragraphs, add one
        if new_tc.find(qn("w:p")) is None:
            new_tc.append(OxmlElement("w:p"))
        row._tr.append(new_tc)


def update_version_banner(doc: Document) -> None:
    for para in doc.paragraphs:
        t = para.text.strip()
        if t.startswith("Document version"):
            clear_paragraph(para)
            run = para.add_run(
                "Document version 2.1  ·  In Scope (with Actors) & Key Stakeholders"
            )
            set_run_font(run, size_pt=9, bold=False, color=RGBColor(0x5A, 0x6A, 0x7A))
            para.paragraph_format.space_after = Pt(12)
            break


def beautify_scope_table(table) -> None:
    set_table_fixed_layout(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # ID | Scope item | Description | Actors
    set_col_widths(table, [1.6, 2.8, 7.6, 4.5])

    headers = ["ID", "Scope item", "Description / boundary", "Actors"]
    for cell, text in zip(table.rows[0].cells, headers):
        style_header_cell(cell, text)
    set_row_repeat_header(table.rows[0])

    for ri, row in enumerate(table.rows[1:], start=1):
        alt = ri % 2 == 0
        sid = row.cells[0].text.strip()
        item = row.cells[1].text.strip()
        desc_lines = [p.text for p in row.cells[2].paragraphs]
        actors = ACTORS.get(sid, ["TBD"])

        rewrite_cell_plain(row.cells[0], sid, bold=True, size_pt=9, center=True)
        rewrite_cell_plain(row.cells[1], item, bold=True, size_pt=9)
        rewrite_cell_lines(row.cells[2], desc_lines, size_pt=8)
        rewrite_cell_lines(row.cells[3], actors, size_pt=8, bullet=True)

        for ci, cell in enumerate(row.cells):
            style_body_cell(cell, alt=alt)
            if ci == 0:
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size_pt=9, bold=True, color=NAVY_RGB)
            if ci == 1:
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size_pt=9, bold=True, color=BODY_RGB)


def beautify_stakeholders_table(table) -> None:
    set_table_fixed_layout(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [3.6, 3.2, 3.4, 3.2, 3.1])

    headers = [c.text.strip() for c in table.rows[0].cells]
    for cell, text in zip(table.rows[0].cells, headers):
        style_header_cell(cell, text)
    set_row_repeat_header(table.rows[0])

    for ri, row in enumerate(table.rows[1:], start=1):
        alt = ri % 2 == 0
        values = [c.text.strip() for c in row.cells]
        for ci, (cell, text) in enumerate(zip(row.cells, values)):
            rewrite_cell_plain(cell, text, bold=(ci == 0), size_pt=8)
            style_body_cell(cell, alt=alt)
            if ci == 0:
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size_pt=8, bold=True, color=NAVY_RGB)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    update_version_banner(doc)

    scope = doc.tables[0]
    # Ensure 4 columns
    while len(scope.rows[0].cells) < 4:
        add_column_to_table(scope)

    beautify_scope_table(scope)
    if len(doc.tables) > 1:
        beautify_stakeholders_table(doc.tables[1])

    doc.save(str(OUT))

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(OUT, CSG_OUT)
        csg_msg = f"Copied: {CSG_OUT}"
    except PermissionError:
        csg_msg = f"Skipped CSG copy (locked): {CSG_OUT}"

    v = Document(str(OUT))
    print(f"Created: {OUT}")
    print(csg_msg)
    t = v.tables[0]
    print(f"Scope table: {len(t.rows)}x{len(t.columns)}")
    print("HDR:", [c.text for c in t.rows[0].cells])
    for row in t.rows[1:]:
        sid = row.cells[0].text.strip()
        actors = row.cells[3].text.replace("\n", " | ")
        print(f"  {sid}: {actors[:120]}")


if __name__ == "__main__":
    main()
