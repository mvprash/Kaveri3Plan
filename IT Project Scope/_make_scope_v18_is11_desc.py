# -*- coding: utf-8 -*-
"""Create IT_Project_Scope_Document_v1.8 — add descriptions to each IS-11 integration."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC = SCOPE_DIR / "IT_Project_Scope_Document_v1.7.docx"
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v1.8.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v1.8.docx"

IS11_LINES = [
    "Kaveri 3.0 shall connect to send/receive data with the following external systems (as deployed / available):",
    "",
    "EKYC – Aadhar Verification — Real-time Aadhaar e-KYC / Face Authentication for citizen and party identity verification during registration, marriage and related flows.",
    "Khazane – Payment systems — Government payment gateway for collection and reconciliation of stamp duty, registration fees and other statutory payments.",
    "DigiLocker — Fetch / push of citizen documents and issued certificates (where enabled) via DigiLocker APIs.",
    "SMS — Transactional SMS for OTP, application status, appointment and certificate / receipt notifications.",
    "EMail — Transactional email for OTP (where used), receipts, certificates and status notifications.",
    "Kutumba — Family / household database lookup and data exchange for eligible Kaveri citizen services.",
    "Civil Registration System — Birth / death and related civil-registration data exchange for marriage and linked workflows.",
    "Labour Department — Data exchange with Labour Department systems where mandated for verification or registration use cases.",
    "Sakala — Karnataka Guarantee of Services to Citizen: GSC issuance, service timelines and statutory status sync.",
    "e-Sign — Aadhaar-based electronic signature for citizens and officers on applications, deeds, certificates, EC/CC and related artefacts.",
    "Eswathu — Rural non-agricultural property records: property pull and mutation / eKhata push after registration.",
    "Eaasthi — Urban property records: property pull and mutation / eKhata push after registration.",
    "Passport — Passport number / details validation for party identity where Passport is selected as proof of identity.",
    "PAN — Real-time PAN authentication for parties (including mandatory checks when market / consideration value thresholds apply; Form 60/61 where PAN is unavailable).",
    "Bhoomi — Agricultural land records: property schedule pull and mutation intimation / push after registration.",
    "KSRSAC — Karnataka State Remote Sensing Applications Centre: GIS / map layers for valuation and spatial property support.",
    "FRUITS — Agricultural mortgage / release filing integration and related EC / notification exchange with FRUITS.",
    "Mogini — Survey / agri property data exchange alongside Bhoomi for property identification and mutation support.",
    "Income Tax — Inter-department access and data exchange for Income Tax Department use cases on Kaveri (Other Department users / approved APIs).",
    "",
    "P.S.: External systems are owned by respective agencies; Kaveri integrates via SOAP/REST APIs / scheduled jobs.",
]


def set_cell_lines(cell, lines: list[str]) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    if not lines:
        first.add_run("")
        return
    first.add_run(lines[0])
    for line in lines[1:]:
        cell.add_paragraph(line)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))

    found = False
    for row in doc.tables[0].rows:
        if row.cells[0].text.strip() == "IS-11":
            set_cell_lines(row.cells[2], IS11_LINES)
            found = True
            break
    if not found:
        raise SystemExit("IS-11 not found")

    doc.save(str(OUT))

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(OUT, CSG_OUT)
        csg_msg = f"Copied: {CSG_OUT}"
    except PermissionError:
        csg_msg = f"Skipped CSG copy (locked): {CSG_OUT}"

    verify = Document(str(OUT))
    print(f"Created: {OUT}")
    print(csg_msg)
    print("--- IS-11 ---")
    for row in verify.tables[0].rows:
        if row.cells[0].text.strip() == "IS-11":
            print(row.cells[2].text)
            break


if __name__ == "__main__":
    main()
