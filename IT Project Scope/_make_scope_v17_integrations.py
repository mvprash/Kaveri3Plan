# -*- coding: utf-8 -*-
"""Create IT_Project_Scope_Document_v1.7 from v1.6.

- Rewrite IS-11 External Integrations to the approved list.
- Align §8.1 stakeholder integration rows 1:1 with that list.
"""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
SRC = SCOPE_DIR / "IT_Project_Scope_Document_v1.6.docx"
OUT = SCOPE_DIR / "IT_Project_Scope_Document_v1.7.docx"
CSG_OUT = ROOT / "CSG Documents" / "IT Project Scope" / "IT_Project_Scope_Document_v1.7.docx"

IS11_LINES = [
    "Kaveri 3.0 shall connect to send/receive data with the following external systems (as deployed / available):",
    "",
    "EKYC – Aadhar Verification",
    "Khazane – Payment systems",
    "DigiLocker",
    "SMS",
    "EMail",
    "Kutumba",
    "Civil Registration System",
    "Labour Department",
    "Sakala",
    "e-Sign",
    "Eswathu",
    "Eaasthi",
    "Passport",
    "PAN",
    "Bhoomi",
    "KSRSAC",
    "FRUITS",
    "Mogini",
    "Income Tax",
    "",
    "P.S.: External systems are owned by respective agencies; Kaveri integrates via APIs / scheduled jobs.",
]

# One stakeholder row per IS-11 integration
INTEGRATION_STAKEHOLDERS: list[tuple[str, str, str, str, str]] = [
    (
        "EKYC – Aadhar Verification",
        "Aadhaar e-KYC / identity verification partner",
        "Citizen and party identity verification for registration and related services",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and authN change windows (jointly with DSR)",
    ),
    (
        "Khazane – Payment systems",
        "Government payment gateway / treasury payment systems",
        "Fee, stamp and registration payment capture and reconciliation",
        "Sandbox SIT; integration war-room; reconciliation reviews",
        "Payment interface SLAs and change windows (jointly with DSR / Treasury)",
    ),
    (
        "DigiLocker",
        "Document wallet / issued-document partner",
        "Fetch / push of citizen documents and issued certificates where enabled",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "SMS",
        "Transactional / OTP SMS provider",
        "OTP, status and notification delivery to citizens and officers",
        "Sandbox SIT; delivery monitoring; change windows",
        "SMS provider SLAs and template/change windows (jointly with DSR)",
    ),
    (
        "EMail",
        "Transactional email provider",
        "OTP (where used), receipts, certificates and status notifications",
        "Sandbox SIT; delivery monitoring; change windows",
        "Email provider SLAs and template/change windows (jointly with DSR)",
    ),
    (
        "Kutumba",
        "Kutumba family / household database integration",
        "Family / household data exchange for eligible Kaveri services",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "Civil Registration System",
        "Civil Registration System (birth / death / related) partner",
        "Civil registration data exchange for marriage and related workflows",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "Labour Department",
        "Labour Department systems",
        "Labour-department data exchange where mandated for registration / verification",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR / Labour)",
    ),
    (
        "Sakala",
        "Karnataka Guarantee of Services to Citizen platform",
        "Service timelines, GSC issuance and statutory status sync",
        "Sandbox SIT; integration war-room; SLA reviews",
        "Service-timeline interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "e-Sign",
        "Aadhaar-based eSign / digital signature services",
        "Citizen eSign and officer digital signature on certificates, deeds, EC/CC and related artefacts",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and signing-service change windows (jointly with DSR)",
    ),
    (
        "Eswathu",
        "Rural property records system (eSwathu)",
        "Rural property pull; mutation / eKhata push where applicable",
        "Sandbox SIT; integration war-room; mutation reconciliation",
        "Property/mutation interface SLAs and change windows (jointly with DSR / RDPR)",
    ),
    (
        "Eaasthi",
        "Urban property records system (eAasthi)",
        "Urban property pull; mutation / eKhata push where applicable",
        "Sandbox SIT; integration war-room; mutation reconciliation",
        "Property/mutation interface SLAs and change windows (jointly with DSR / ULB)",
    ),
    (
        "Passport",
        "Passport verification / related services",
        "Passport data validation for eligible Kaveri identity / party flows",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "PAN",
        "PAN verification services",
        "PAN validation for parties where mandated (including SFT-related use where approved)",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "Bhoomi",
        "Agricultural land records system (Bhoomi)",
        "Agri property pull; mutation push after registration",
        "Sandbox SIT; integration war-room; mutation reconciliation",
        "Property/mutation interface SLAs and change windows (jointly with DSR / Revenue)",
    ),
    (
        "KSRSAC",
        "Karnataka State Remote Sensing Applications Centre (GIS / map services)",
        "Map / GIS layer support for valuation and related spatial flows",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "FRUITS",
        "FRUITS agri mortgage / release filing system",
        "Agri mortgage / release filing and related EC / notification exchange",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR)",
    ),
    (
        "Mogini",
        "Mogini (survey / agri property) system",
        "Survey / agri property data exchange alongside Bhoomi flows",
        "Sandbox SIT; integration war-room; change windows",
        "Interface SLAs and change windows (jointly with DSR / Revenue)",
    ),
    (
        "Income Tax",
        "Income Tax Department systems / access",
        "Inter-department access and data exchange for Income Tax use cases on Kaveri",
        "Admin provisioning; sandbox SIT where APIs apply; change windows",
        "Interface / access SLAs and change windows (jointly with DSR / ITD)",
    ),
]

# Non-integration stakeholders kept from v1.6 (leading + trailing)
CORE_BEFORE = [
    "IGR & Commissioner of Stamps / Steering Committee",
    "AIGR (Computers) & Requirements Committee",
    "KPMU / Application Admin",
    "Product Owner (Kaveri IT Cell)",
    "Project Manager (Kaveri IT Cell)",
    "Domain Expert (Stamps & Registration)",
    "Citizens / Applicants",
    "DSR Officers (SR / DEO / FDA / SDA / DRO / DIGR / AIGR / HQA)",
    "Other Department users",
]
CORE_AFTER = [
    "Kaveri IT Cell delivery & ops",
    "SDC / Infrastructure operations",
]


def set_cell_lines(cell, lines: list[str]) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    if not lines:
        first.add_run("")
        return
    first.add_run(lines[0])
    for line in lines[1:]:
        cell.add_paragraph(line)


def set_cell_text(cell, text: str) -> None:
    set_cell_lines(cell, [text] if text else [""])


def insert_row_after(table, after_index: int) -> int:
    src_tr = table.rows[after_index]._tr
    new_tr = deepcopy(src_tr)
    src_tr.addnext(new_tr)
    return after_index + 1


def find_stakeholders_table(doc: Document):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " | ".join(c.text.strip() for c in table.rows[0].cells).lower()
        if "stakeholder" in header and "role" in header:
            return table
    return None


def rebuild_stakeholder_integrations(table) -> None:
    """Keep core before/after rows; replace middle with INTEGRATION_STAKEHOLDERS."""
    # Capture trailing core row content from current table
    after_rows_data: list[list[str]] = []
    before_end = None
    after_start = None

    for i, row in enumerate(table.rows):
        name = row.cells[0].text.strip()
        if name in CORE_BEFORE:
            before_end = i
        if name in CORE_AFTER and after_start is None:
            after_start = i
        if name in CORE_AFTER:
            after_rows_data.append([c.text.strip() for c in row.cells])

    if before_end is None or after_start is None:
        raise SystemExit(
            f"Could not locate core stakeholder anchors "
            f"(before_end={before_end}, after_start={after_start})"
        )

    # Delete rows between before_end and after_start (old integrations)
    # Delete from after_start-1 down to before_end+1
    for idx in range(after_start - 1, before_end, -1):
        tr = table.rows[idx]._tr
        tr.getparent().remove(tr)

    # after_start shifts: new insert point is before_end
    insert_at = before_end
    for row_data in INTEGRATION_STAKEHOLDERS:
        insert_at = insert_row_after(table, insert_at)
        for ci, val in enumerate(row_data):
            if ci < len(table.rows[insert_at].cells):
                set_cell_text(table.rows[insert_at].cells[ci], val)


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Source not found: {SRC}")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))

    # IS-11
    found = False
    for row in doc.tables[0].rows:
        if row.cells[0].text.strip() == "IS-11":
            set_cell_lines(row.cells[2], IS11_LINES)
            found = True
            break
    if not found:
        raise SystemExit("IS-11 not found")

    # Stakeholders — align integrations 1:1
    st = find_stakeholders_table(doc)
    if st is None:
        raise SystemExit("Stakeholders table not found")
    rebuild_stakeholder_integrations(st)

    doc.save(str(OUT))

    CSG_OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(OUT, CSG_OUT)
        csg_msg = f"Copied: {CSG_OUT}"
    except PermissionError:
        csg_msg = f"Skipped CSG copy (locked): {CSG_OUT}"

    verify = Document(str(OUT))
    print(f"Created: {OUT}")
    print(csg_msg)
    print("--- IS-11 ---")
    for row in verify.tables[0].rows:
        if row.cells[0].text.strip() == "IS-11":
            print(row.cells[2].text)
            break
    print("--- Stakeholders ---")
    vt = find_stakeholders_table(verify)
    for i, row in enumerate(vt.rows[1:], 1):
        print(f"  {i}. {row.cells[0].text}")


if __name__ == "__main__":
    main()
