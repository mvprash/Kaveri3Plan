# -*- coding: utf-8 -*-
"""Inspect formatting of IT_Project_Scope_Document_v1.8.docx"""
from __future__ import annotations

import sys
from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

p = r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\IT Project Scope\IT_Project_Scope_Document_v1.8.docx"
doc = Document(p)

print("=== PARAS ===")
for i, para in enumerate(doc.paragraphs):
    t = para.text
    style = para.style.name if para.style else None
    print(f"[{i}] style={style} align={para.paragraph_format.alignment} text={t!r}")
    for r in para.runs[:4]:
        f = r.font
        color = None
        try:
            color = f.color.rgb
        except Exception:
            pass
        print(f"    run={r.text[:50]!r} bold={f.bold} size={f.size} name={f.name} color={color}")

print("\n=== TABLES ===", len(doc.tables))
for ti, table in enumerate(doc.tables):
    print(f"--- Table {ti}: {len(table.rows)}x{len(table.columns)} style={table.style.name if table.style else None} ---")
    for ci, cell in enumerate(table.rows[0].cells):
        tcPr = cell._tc.tcPr
        w = None
        shd = None
        if tcPr is not None:
            if tcPr.tcW is not None:
                w = tcPr.tcW.get(qn("w:w"))
            shd_el = tcPr.find(qn("w:shd"))
            if shd_el is not None:
                shd = shd_el.get(qn("w:fill"))
        print(f"  H col{ci} w={w} shd={shd} text={cell.text[:60]!r}")
    # first data row sample
    if len(table.rows) > 1:
        row = table.rows[1]
        for ci, cell in enumerate(row.cells[:3]):
            print(f"  D1 col{ci} paras={len(cell.paragraphs)} text_len={len(cell.text)}")
            if cell.paragraphs:
                r0 = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else None
                if r0:
                    print(f"    run bold={r0.font.bold} size={r0.font.size} name={r0.font.name}")

print("\n=== SECTION ===")
for s in doc.sections:
    print(f"page {s.page_width.inches:.2f}x{s.page_height.inches:.2f}")
    print(f"margins L={s.left_margin.inches:.2f} R={s.right_margin.inches:.2f} T={s.top_margin.inches:.2f} B={s.bottom_margin.inches:.2f}")

# Check for weird chars
print("\n=== ODD CHARS ===")
for i, para in enumerate(doc.paragraphs):
    if any(ord(c) > 0x2000 for c in para.text):
        print(f"para[{i}]: {[hex(ord(c)) for c in para.text if ord(c)>0x2000]} {para.text!r}")
