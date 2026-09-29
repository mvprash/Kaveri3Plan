# -*- coding: utf-8 -*-
"""Fill §8.1 Key Stakeholders in IT Project Scope v1.4 from Kaveri 3 Plan.

Sources:
  Project_Plan_Kaveri_3.0_Programme_v0.4.md (§14 Governance, §15 Training, §17 Acceptance)
  Finalized BRD/User Management/BRD_User_Management_v1.6.md (§4 Stakeholders)
  Finalized BRD/Marriage (actors)
  IT_Project_Scope_Document_v1.4 In-Scope (IS-01..IS-11 actors / integrations)
  Requirement Discussions daily reports (Committee, AIGR Comp, Kaveri IT Cell, KPMU)

Updates every IT_Project_Scope_Document_*.docx under:
  IT Project Scope/
  CSG Documents/IT Project Scope/
and Template1 when writable.
"""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan")
SCOPE_DIR = ROOT / "IT Project Scope"
CSG_SCOPE = ROOT / "CSG Documents" / "IT Project Scope"

# Columns: Stakeholder / Group | Role | Interest / impact | Engagement method | Decision authority
STAKEHOLDERS: list[tuple[str, str, str, str, str]] = [
    (
        "IGR & Commissioner of Stamps / Steering Committee",
        "Programme sponsor; head of Department of Stamps and Registration (DSR)",
        "Programme outcomes, statutory compliance, phase Go-Live readiness and residual risk",
        "Fortnightly Steering; Go-Live (G-GL) checklist; escalations",
        "Scope change / swap-in; phase Go-Live approval; residual risk acceptance",
    ),
    (
        "AIGR (Computers) & Requirements Committee",
        "Domain and requirements authority for Kaveri 3.0",
        "Correctness and completeness of BR across Marriage, UM, Document, EC/CC, Firm, e-Stamp, Audit",
        "BR workshops per schedule v3; daily requirement discussions with Kaveri IT Cell",
        "BRD / domain sign-off; promote or defer scope items",
    ),
    (
        "KPMU / Application Admin",
        "Programme Management Unit; system super-admin / IGR nominee",
        "Delivery governance; designated creator roles; sanctioned posts and system-wide configuration",
        "Steering nominee; admin / configuration workshops; UAT sign-off support",
        "Designate creator roles; sanctioned-posts master; KPMU programme decisions",
    ),
    (
        "Product Owner (Kaveri IT Cell)",
        "Product ownership for Kaveri 3.0",
        "Roadmap, backlog, sprint goals and stakeholder alignment within approved scope",
        "Sprint review / retro; backlog refinement; Steering pack",
        "Prioritisation and MVP boundary inside Steering-approved scope",
    ),
    (
        "Project Manager (Kaveri IT Cell)",
        "Programme / project delivery management",
        "Plan, milestones, risks, vendor/SLA and O&M reporting across 11-month window",
        "Daily stand-ups; risk reviews; Steering reporting",
        "Schedule and mitigation within Steering bounds; release train coordination",
    ),
    (
        "Domain Expert (Stamps & Registration)",
        "Department SME (≥20 yrs) for Acts, Rules, forms and fees",
        "Legal fidelity of forms, fees, registers and process rules in software",
        "BR workshops; design / UAT reviews",
        "Domain acceptance of rules and forms implemented in Kaveri 3.0",
    ),
    (
        "Citizens / Applicants",
        "Public end users of Kaveri portal services",
        "Self-service for Marriage, Document Registration, EC, CC, Firm, Digital e-Stamp and related services",
        "Portal help / KB; sampled UAT; post go-live feedback",
        "None (service consumers)",
    ),
    (
        "DSR Officers (SR / DEO / FDA / SDA / DRO / DIGR / AIGR / HQA)",
        "Departmental operational users across SRO / DRO / IGRO",
        "Scrutiny, registration, DEO upload, firm/DRO workflows, MIS, audit and appeals",
        "UAT; training before each phase Go-Live; hypercare support",
        "Operational decisions within statutory powers (admit / refuse / pending / etc.)",
    ),
    (
        "Other Department users",
        "Officers / staff of government departments other than DSR with assigned Kaveri access",
        "Controlled access to assigned modules (e.g. Income Tax, Banking, Govt. Lawyer, Revenue)",
        "Admin-provisioned accounts; targeted UAT where modules are shared",
        "None beyond assigned module actions",
    ),
    (
        "External integration owners",
        "Partner systems: e-KYC/UIDAI, Khajane, DigiLocker, eSign/DSC, Sakala, SMS/Email, Bhoomi / e-Swathu / e-Aasthi / ULMS, Kutumba, CRS, Labour, CDAC, Passport, etc.",
        "Reliable authN, payment, notification, mutation and statutory-service data exchange",
        "Integration war-rooms; sandbox SIT; change windows",
        "Interface SLAs and change windows (jointly with DSR / SDC)",
    ),
    (
        "Kaveri IT Cell delivery & ops",
        "Architect, Tech Leads, BA, Full Stack, Integration, UI/UX, DBA, Migration, QA/Test, Perf/Security, DevOps, Security, BI, Content, Transition, L2",
        "Design, build, migrate, test, deploy, observe and hypercare Kaveri 3.0",
        "Architecture sync; SIT/UAT; release trains; L2 hypercare",
        "Technical design within AD/NFR; release-readiness recommendation to Steering",
    ),
    (
        "SDC / Infrastructure operations",
        "State Data Centre hosting, network (incl. KSWAN for department UI) and WAF",
        "Environments (Staging/Prod), capacity, connectivity and infra change windows",
        "Ops coordination; deployment and rollback windows",
        "Infra capacity, network and change-window approval",
    ),
]


def set_cell_text(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    first.add_run(text or "")


def ensure_rows(table, needed_data_rows: int) -> None:
    while len(table.rows) < needed_data_rows + 1:
        tbl = table._tbl
        last = table.rows[-1]._tr
        tbl.append(deepcopy(last))
    while len(table.rows) > needed_data_rows + 1:
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)


def find_stakeholders_table(doc: Document):
    for table in doc.tables:
        if not table.rows:
            continue
        header = " | ".join(c.text.strip() for c in table.rows[0].cells).lower()
        if "stakeholder" in header and "role" in header:
            return table
    return None


def update_doc(path: Path) -> str:
    if not path.exists():
        return f"Skip (missing): {path}"
    try:
        doc = Document(str(path))
    except Exception as e:
        return f"Skip (open failed): {path} — {e}"

    table = find_stakeholders_table(doc)
    if table is None:
        return f"Skip (no stakeholders table): {path}"

    ensure_rows(table, len(STAKEHOLDERS))
    for i, row_data in enumerate(STAKEHOLDERS, start=1):
        row = table.rows[i]
        for ci, value in enumerate(row_data):
            if ci < len(row.cells):
                set_cell_text(row.cells[ci], value)

    try:
        doc.save(str(path))
    except PermissionError:
        alt = path.with_name(path.stem + "_stakeholders_updated.docx")
        doc.save(str(alt))
        return f"Locked; wrote: {alt}"
    return f"Updated: {path}"


def main() -> None:
    targets: list[Path] = []
    for folder in (SCOPE_DIR, CSG_SCOPE):
        if not folder.exists():
            continue
        for p in sorted(folder.glob("IT_Project_Scope_Document*.docx")):
            name = p.name
            if name.startswith("~$"):
                continue
            targets.append(p)

    # Prefer updating v1.4 first, then other versions / template
    targets.sort(key=lambda p: (0 if "v1.4" in p.name else 1, str(p)))

    seen: set[Path] = set()
    results: list[str] = []
    for path in targets:
        if path in seen:
            continue
        seen.add(path)
        results.append(update_doc(path))

    # Sync Template1 from v1.4 when present
    v14 = SCOPE_DIR / "IT_Project_Scope_Document_v1.4.docx"
    tpl = SCOPE_DIR / "IT_Project_Scope_Document_Template1.docx"
    if v14.exists() and tpl.exists():
        try:
            shutil.copy2(v14, tpl)
            results.append(f"Synced Template1 from v1.4: {tpl}")
        except PermissionError:
            results.append(f"Skipped Template1 sync (locked): {tpl}")

    for line in results:
        print(line)
    print(f"Done. Stakeholders rows: {len(STAKEHOLDERS)}")


if __name__ == "__main__":
    main()
