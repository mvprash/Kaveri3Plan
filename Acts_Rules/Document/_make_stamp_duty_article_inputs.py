"""Builds Stamp_Duty_Article_Determination_Inputs.xlsx from _stamp_article_inputs_data.py."""

import os
import re
import sys
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _stamp_article_inputs_data as D  # noqa: E402

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Stamp_Duty_Article_Determination_Inputs.xlsx")

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
GRP_FILL = PatternFill("solid", fgColor="DDEBF7")
ID_FILL = PatternFill("solid", fgColor="E2EFDA")
CALC_FILL = PatternFill("solid", fgColor="FFF2CC")
BOTH_FILL = PatternFill("solid", fgColor="F8CBAD")
EX_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

NATURE = {n[0]: n for n in D.DOCUMENT_NATURES}
INTENT = dict(D.TRANSACTION_INTENTS)
INPUT = {i[0]: i for i in D.INPUTS}
RULE_KEYS = [r["key"] for r in D.RULES]


def validate():
    errors = []
    if len(set(RULE_KEYS)) != len(RULE_KEYS):
        dup = [k for k in RULE_KEYS if RULE_KEYS.count(k) > 1]
        errors.append(f"Duplicate rule keys: {sorted(set(dup))}")
    for r in D.RULES:
        for i in r["id_in"] + r["calc_in"] + r["adj_in"]:
            if i not in INPUT:
                errors.append(f"Rule {r['key']} references unknown input {i}")
        if r["nature"] not in NATURE:
            errors.append(f"Rule {r['key']} unknown nature {r['nature']}")
    referenced = set()
    for row in flow_rows():
        for tok in re.findall(r"=>\s*([0-9A-Za-z\-\(\)]+(?:\([^)]*\))*)", row[6]):
            if tok in RULE_KEYS:
                referenced.add(tok)
    unreached = [k for k in RULE_KEYS if k not in referenced]
    if errors:
        raise SystemExit("\n".join(errors))
    return unreached


def flow_rows():
    """(Nature code, Nature, Step, Input ID, Question, Answer, Result)."""
    rows = []
    in_flow = OrderedDict()
    for f in D.FLOW:
        in_flow.setdefault(f[0], []).append(f)
    by_nature = defaultdict(list)
    for r in D.RULES:
        by_nature[r["nature"]].append(r)

    for code, name, _intent, _art in D.DOCUMENT_NATURES:
        if code in in_flow:
            for f in in_flow[code]:
                rows.append((code, name, str(f[1]), f[2], f[3], f[4], f[5]))
        elif code in D.SIMPLE_THRESHOLD_FLOWS:
            inp, q, opts = D.SIMPLE_THRESHOLD_FLOWS[code]
            for ans, key in opts:
                rows.append((code, name, "1", inp, q, ans, f"=> {key}"))
        else:
            rules = by_nature.get(code, [])
            if len(rules) == 1:
                rows.append((code, name, "1", "-", "No further question - nature decides the article", "-",
                             f"=> {rules[0]['key']}"))
            else:
                for r in rules:
                    rows.append((code, name, "1", ", ".join(r["id_in"]) or "-", r["cond"], "Condition met",
                                 f"=> {r['key']}"))
    return rows


def style_header(ws, row, ncols, height=32):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = height


def write_table(ws, headers, rows, widths, start_row=1, freeze="B2", filt=True):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=c, value=h)
    style_header(ws, start_row, len(headers))
    for r_i, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c, value=v)
            cell.alignment = WRAP
            cell.border = BORDER
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    if freeze:
        ws.freeze_panes = freeze
    if filt and rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def sheet_readme(wb, unreached):
    ws = wb.active
    ws.title = "ReadMe"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120
    lines = [
        ("Title", "Stamp Duty - Article Determination & Input Capture Matrix (Kaveri 3.0 - Document Registration)"),
        ("Purpose", "Stamp duty depends on the Article of the Schedule under which the instrument falls. This workbook "
                    "lists every input the system must capture, the rules that map those inputs to an Article / "
                    "sub-clause, and the duty basis, so that the system can determine the Article automatically "
                    "instead of relying on the citizen / operator to pick it."),
        ("Legal source", "Karnataka Stamp Act, 1957 - Sections 4, 5, 6, 45-A and the Schedule (consolidated up to 2022), "
                         "files in Acts_Rules/Document: 'THE KARNATAKA STAMP ACT 1957.pdf', 'Karnataka Stamp Act 1957 "
                         "Schedule 2022.pdf', Amendment Acts 2011 / 2012 / 2014."),
        ("Rate caution", "Duty rates are as printed in the 2022 consolidated Schedule in this folder. Any later amendment "
                         "(e.g. Karnataka Stamp (Amendment) Act 2023 / 2024 revisions) is NOT in the folder and must be "
                         "applied by the Department before configuring rates. Rows marked 'confirm with Department' had "
                         "illegible / omitted text in the source copy."),
        ("", ""),
        ("How the system decides", "1. Ask GEN-01 Transaction intent -> filter GEN-02 Nature of document.\n"
                                   "2. Ask the questions in Decision_Flow for that nature, in step order.\n"
                                   "3. The first terminal answer ('=> rule') gives the Article / sub-clause (Article_Rules).\n"
                                   "4. Capture the Calculation inputs listed for that rule and compute the duty.\n"
                                   "5. Apply provisos / adjustments (earlier duty paid, caps) and exemptions.\n"
                                   "6. Apply general rules (General_Rules sheet): several matters (Sec 5) = aggregate; "
                                   "multiple descriptions (Sec 6) = highest; several instruments of one sale (Sec 4) = "
                                   "principal pays full, others Rs.100; Sec 45-A market value check."),
        ("", ""),
        ("Sheet: Lookup_Lists", "Dropdown values: Transaction intent, Nature of document (DN codes -> Article), Agreement "
                                "subject, POA purpose, Conveyance type, Local body limits, Relationship, Allotting "
                                "authority, Form of security."),
        ("Sheet: Input_Catalogue", "Every field to capture: ID, group, label, question text, data type, allowed values, "
                                   "source (User / System), purpose (I = identifies the Article, C = needed to calculate "
                                   "duty, E = exemption / adjustment check) and the Articles that use it."),
        ("Sheet: Article_Rules", "One row per Article / sub-clause: identification conditions, identification inputs, "
                                 "calculation inputs, duty basis and rate, min / max, provisos, exemptions, Sec 45-A flag."),
        ("Sheet: Article_Input_Matrix", "Grid of Article rule vs Input ID. I = needed to identify, C = needed to calculate, "
                                        "I+C = both, E = exemption / adjustment check. Use it to build the dynamic "
                                        "capture screen per Article."),
        ("Sheet: Decision_Flow", "Question sequence per Nature of document with branching answers ending in '=> Article'."),
        ("Sheet: Family_Definitions", "Family definition differs by Article (Gift, Settlement, Release, Lease, POA). The "
                                      "system must check the relationship against the correct column."),
        ("Sheet: Jurisdiction_Slab", "Fixed duty by local body limits for family Gift / Settlement / Release / Lease. "
                                     "If properties fall in more than one limit, the maximum applies."),
        ("Sheet: General_Rules", "Act-level rules affecting Article selection and calculation."),
        ("Sheet: Worked_Examples", "Test cases: inputs -> expected Article and duty (for UAT)."),
        ("", ""),
        ("Counts", f"{len(D.TRANSACTION_INTENTS)} intents, {len(D.DOCUMENT_NATURES)} document natures, "
                   f"{len(D.INPUTS)} inputs, {len(D.RULES)} Article / sub-clause rules, {len(D.EXAMPLES)} worked examples."),
        ("Rules not reached by flow", ", ".join(unreached) if unreached else "None - every rule is reachable from Decision_Flow."),
        ("Generator", "Regenerate with: python Acts_Rules/Document/_make_stamp_duty_article_inputs.py"),
    ]
    for r, (a, b) in enumerate(lines, 1):
        ca = ws.cell(row=r, column=1, value=a)
        cb = ws.cell(row=r, column=2, value=b)
        ca.font = Font(bold=True)
        ca.alignment = WRAP
        cb.alignment = WRAP
        if r == 1:
            ca.font = Font(bold=True, size=13, color="1F4E78")
            cb.font = Font(bold=True, size=13, color="1F4E78")
        if a == "Rate caution":
            cb.font = Font(bold=True, color="C00000")


def sheet_lookups(wb):
    ws = wb.create_sheet("Lookup_Lists")
    blocks = [
        ("Transaction Intent (GEN-01)", ["Code", "Intent"], D.TRANSACTION_INTENTS),
        ("Nature of Document (GEN-02)", ["Code", "Nature of Document", "Intent", "Article(s)"],
         [(c, n, f"{i} - {INTENT[i]}", a) for c, n, i, a in D.DOCUMENT_NATURES]),
        ("Agreement Subject (AGR-01)", ["Code", "Agreement relates to", "Article"], D.AGREEMENT_SUBJECTS),
        ("POA Purpose (POA-01)", ["Code", "Purpose", "Article"], D.POA_PURPOSES),
        ("Conveyance Type (Art 20)", ["Code", "Conveyance type", "Article"], D.CONVEYANCE_TYPES),
        ("Local Body Limits (PRP-03)", ["Code", "Limits"], D.LOCAL_LIMITS),
        ("Relationship (GEN-07)", ["Relationship"], [(r,) for r in D.RELATIONSHIPS]),
        ("Allotting Authority (AGR-04)", ["Authority"], [(a,) for a in D.ALLOTTING_AUTHORITIES]),
        ("Form of Security (LON-03)", ["Form"],
         [(s,) for s in ["Deposit of title deeds", "Pawn / pledge of movable property"] + D.SECURITY_FORMS
          + ["Mortgage of crop", "Further charge"]]),
    ]
    row = 1
    for title, headers, data in blocks:
        t = ws.cell(row=row, column=1, value=title)
        t.font = Font(bold=True, size=12, color="1F4E78")
        row += 1
        for c, h in enumerate(headers, 1):
            ws.cell(row=row, column=c, value=h)
        style_header(ws, row, len(headers), height=20)
        for rec in data:
            row += 1
            for c, v in enumerate(rec, 1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.alignment = WRAP
                cell.border = BORDER
        row += 2
    for c, w in enumerate([12, 70, 55, 16], 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def sheet_inputs(wb):
    ws = wb.create_sheet("Input_Catalogue")
    used = defaultdict(list)
    for r in D.RULES:
        for i in OrderedDict.fromkeys(r["id_in"] + r["calc_in"] + r["adj_in"]):
            used[i].append(r["key"])
    for i in D.ALL_DOC_INPUTS:
        used[i].insert(0, "All documents")
    used["GEN-11"].append("Sale / Mortgage / Settlement (Sec 4)")
    used["GEN-12"].append("Adjustment provisos")
    purpose_txt = {"I": "Identification", "C": "Calculation", "E": "Exemption / Adjustment"}
    headers = ["Input ID", "Group", "Field label", "Question to user", "Data type", "Allowed values / format",
               "Source", "Purpose", "No. of Article rules using it", "Article rules using it"]
    rows = []
    for i in D.INPUTS:
        u = used.get(i[0], [])
        rows.append(list(i[:7]) + [purpose_txt.get(i[7], i[7]), len([x for x in u if x[0].isdigit()]) or "",
                                   ", ".join(u)])
    write_table(ws, headers, rows, [10, 16, 32, 60, 18, 48, 20, 16, 12, 60], freeze="C2")
    prev = None
    for r_i in range(2, len(rows) + 2):
        grp = ws.cell(row=r_i, column=2).value
        if grp != prev:
            for c in range(1, 4):
                ws.cell(row=r_i, column=c).fill = GRP_FILL
        prev = grp
        ws.cell(row=r_i, column=1).font = Font(bold=True)


def sheet_rules(wb):
    ws = wb.create_sheet("Article_Rules")
    headers = ["Rule / Article ref", "Article", "Nature code", "Nature of document", "Transaction intent",
               "Instrument (Schedule description)", "Identification conditions (system logic)",
               "Identification inputs", "Calculation inputs", "Exemption / adjustment inputs", "Duty basis",
               "Proper stamp duty (Schedule)", "Minimum", "Maximum", "Provisos / adjustments", "Exemptions",
               "Sec 45-A MV check"]
    rows = []
    for r in D.RULES:
        n = NATURE[r["nature"]]
        rows.append([r["key"], r["art"], r["nature"], n[1], INTENT[n[2]], r["instrument"], r["cond"],
                     ", ".join(r["id_in"]), ", ".join(r["calc_in"]), ", ".join(r["adj_in"]), r["basis"], r["duty"],
                     r["min"], r["max"], r["proviso"], r["exempt"], r["s45a"]])
    write_table(ws, headers, rows, [14, 8, 9, 28, 26, 38, 55, 26, 30, 22, 14, 55, 12, 18, 50, 40, 10], freeze="B2")
    for r_i in range(2, len(rows) + 2):
        ws.cell(row=r_i, column=1).font = Font(bold=True)
        if ws.cell(row=r_i, column=17).value == "Yes":
            ws.cell(row=r_i, column=17).fill = CALC_FILL


def sheet_matrix(wb):
    ws = wb.create_sheet("Article_Input_Matrix")
    used_ids = [i[0] for i in D.INPUTS
                if any(i[0] in r["id_in"] + r["calc_in"] + r["adj_in"] for r in D.RULES)]
    headers = ["Rule / Article ref", "Instrument"] + used_ids
    ws.cell(row=1, column=1, value="Legend:")
    ws.cell(row=1, column=2, value="I = identifies the Article   C = needed for duty calculation   I+C = both   "
                                   "E = exemption / adjustment check   (" + ", ".join(D.ALL_DOC_INPUTS)
                                   + " are captured for every document)")
    ws.cell(row=1, column=1).font = Font(bold=True)
    for c, h in enumerate(headers, 1):
        ws.cell(row=2, column=c, value=h)
    style_header(ws, 2, len(headers), height=70)
    for c in range(3, len(headers) + 1):
        ws.cell(row=2, column=c).alignment = Alignment(text_rotation=90, horizontal="center", vertical="center")
    for r_i, r in enumerate(D.RULES, 3):
        ws.cell(row=r_i, column=1, value=r["key"]).font = Font(bold=True)
        ws.cell(row=r_i, column=2, value=r["instrument"]).alignment = WRAP
        for c, iid in enumerate(used_ids, 3):
            a, b = iid in r["id_in"], iid in r["calc_in"]
            val, fill = ("I+C", BOTH_FILL) if a and b else ("I", ID_FILL) if a else ("C", CALC_FILL) if b else ("", None)
            if not val and iid in r["adj_in"]:
                val, fill = "E", GRP_FILL
            cell = ws.cell(row=r_i, column=c, value=val or None)
            cell.alignment = CENTER
            cell.border = BORDER
            if fill:
                cell.fill = fill
        ws.cell(row=r_i, column=1).border = BORDER
        ws.cell(row=r_i, column=2).border = BORDER
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 45
    for c in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 4.2
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(D.RULES) + 2}"


def sheet_flow(wb):
    ws = wb.create_sheet("Decision_Flow")
    pre = [
        ("ALL", "All documents", "0", "GEN-01", "What does this document primarily do?", "Intent INT-01..INT-12",
         "Filter Nature list"),
        ("ALL", "All documents", "0", "GEN-02", "Select nature of document", "DN-01..DN-59", "Go to that nature's steps"),
        ("ALL", "All documents", "0", "GEN-10",
         "Several distinct matters in one document?", "Yes", "Run the flow for each matter; duty = aggregate (Sec 5)"),
        ("ALL", "All documents", "0", "-", "Document matches two or more descriptions?", "Yes",
         "Compute each; charge the highest (Sec 6)"),
    ]
    headers = ["Nature code", "Nature of document", "Step", "Input ID", "Question", "Answer", "Next step / Result"]
    rows = pre + flow_rows()
    write_table(ws, headers, rows, [10, 34, 7, 18, 60, 45, 40], freeze="C2")
    prev = None
    band = False
    for r_i in range(2, len(rows) + 2):
        code = ws.cell(row=r_i, column=1).value
        if code != prev:
            band = not band
        prev = code
        res = ws.cell(row=r_i, column=7)
        if str(res.value).startswith("=>"):
            res.font = Font(bold=True, color="1F4E78")
        if band:
            for c in range(1, 3):
                ws.cell(row=r_i, column=c).fill = GRP_FILL


def sheet_family(wb):
    ws = wb.create_sheet("Family_Definitions")
    headers = ["Relationship (claimant to executant)"] + [f"{a}\n{d}" for a, d in D.FAMILY_MATRIX_ARTICLES]
    rows = [[rel] + D.FAMILY_MATRIX[rel] for rel in D.RELATIONSHIPS]
    write_table(ws, headers, rows, [34, 20, 20, 22, 22, 26], freeze="B2", filt=False)
    ws.row_dimensions[1].height = 60
    for r_i in range(2, len(rows) + 2):
        for c in range(2, 7):
            cell = ws.cell(row=r_i, column=c)
            cell.alignment = CENTER
            if cell.value == "Y":
                cell.fill = ID_FILL
    n = len(rows) + 3
    notes = [
        "Y = relation is 'family' for that Article. Blank = not family (non-family rule applies).",
        "28(b) Gift & 48-A(ii) Settlement: father, mother, husband, wife, son, daughter, daughter-in-law, brothers, "
        "sisters and grand children.",
        "45(b) Release: husband, wife, son, daughter, father, mother, brother, wife / children of predeceased brother, "
        "sister, husband / children of predeceased sister, wife of predeceased son, children of predeceased son or "
        "daughter (source text partly garbled - confirm exact list with Department).",
        "30(1) 2nd proviso Lease: wife, husband, father, mother, son, daughter, brother, sister.",
        "41(eb) POA to sell: if attorney is father, mother, wife, husband, son, daughter, brother or sister, Art 41(eb) "
        "does NOT apply (then 41(b) / (c) / (d) / (h) as per Decision_Flow), unless given for consideration (41(e)).",
    ]
    for k, t in enumerate(notes):
        c = ws.cell(row=n + k, column=1, value=t)
        c.alignment = Alignment(wrap_text=False)
        c.font = Font(italic=True) if k else Font(bold=True)


def sheet_slab(wb):
    ws = wb.create_sheet("Jurisdiction_Slab")
    names = dict(D.LOCAL_LIMITS)
    headers = ["Limit code", "Local body limits"] + D.JURISDICTION_SLAB_COLS
    rows = [[c, names[c]] + list(v) for c, *v in D.JURISDICTION_SLAB]
    write_table(ws, headers, rows, [10, 60, 18, 22, 20, 26], freeze="C2", filt=False)
    for r_i in range(2, len(rows) + 2):
        for c in range(3, 7):
            ws.cell(row=r_i, column=c).number_format = '"Rs. "#,##0'
    ws.cell(row=len(rows) + 3, column=1,
            value="If the properties are situated in a combination of limits, the duty payable is the maximum of the "
                  "applicable amounts (proviso to 28(b), 30(1), 45(b), 48-A(ii)). PRP-03 must therefore be multi-select "
                  "and derived from the property location master.").font = Font(italic=True)


def sheet_general(wb):
    ws = wb.create_sheet("General_Rules")
    rows = [
        ("GR-01", "Sec 5", "Instrument relating to several distinct matters",
         "Chargeable with the aggregate of duties for each matter as separate instruments.",
         "Capture GEN-10; run Decision_Flow per matter; sum the duties."),
        ("GR-02", "Sec 6", "Instrument falling under two or more descriptions",
         "Chargeable only with the highest of such duties (subject to Sec 5).",
         "Evaluate each matching rule; select maximum."),
        ("GR-03", "Sec 4", "Several instruments for one sale / mortgage / settlement",
         "Principal instrument pays Schedule duty (highest of the set); every other instrument Rs. 100.",
         "Capture GEN-11; if Secondary -> Rs. 100; parties may designate principal."),
        ("GR-04", "Sec 6 proviso / Art 22", "Counterpart / duplicate",
         "Counterpart of an instrument on which proper duty is paid: Art 22.",
         "Validate original document reference (MSC-09)."),
        ("GR-05", "Sec 45-A / 45-B", "Market value check",
         "For conveyance, gift 28(a), exchange, settlement 48-A(i), partnership 40-B(a) / 40-C(a), agreement 5(e)(i) "
         "/ 5(f), lease 30(1)(vi), POA 41(e)/(ea)/(eb), release 45(a), decree conveyance, award 11(a), trust 54(iii), "
         "TDR 20(7): duty on higher of declared value and guidance MV published by CVC.",
         "GEN-09 auto-fetched from Guidance Value module; 'Sec 45-A MV check' column in Article_Rules = Yes."),
        ("GR-06", "Schedule", "'or part thereof' rounding",
         "Wherever duty is 'per Rs. X or part thereof', the base is rounded UP to the next multiple of X.",
         "Implement CEILING(base / X) x rate."),
        ("GR-07", "Schedule", "Minimum / maximum caps",
         "Apply the min / max stated for the clause after computing the ad valorem amount.",
         "Article_Rules columns Minimum / Maximum."),
        ("GR-08", "Art 5 Expl. II, 41 provisos, 20(1) provisos", "Adjustment of duty already paid",
         "Duty paid on agreement 5(e)/5(h), POA 41(e)/(eb), lease-cum-sale 5(d)/5(da) is adjustable against the later "
         "conveyance / mortgage between the same parties for the same property.",
         "Capture GEN-12 / POA-04 with document reference; system verifies the earlier e-stamp / registration."),
        ("GR-09", "Art 5(e), 5(g), 14", "Cancellation of agreement",
         "Deed cancelling a duly stamped 5(e)(i) / 5(g)(i) agreement between the same parties: max Rs. 500.",
         "AGR-15 flag."),
        ("GR-10", "Art 30, 32-A", "Average annual rent",
         "Average annual rent = total rent for the term / number of years (for term < 1 year, rent for the term).",
         "Derived from LSE-03 rent schedule (captures escalation)."),
        ("GR-11", "Art 30, 32-A, 5(f), 41(ea)", "Money advanced",
         "Includes security deposit, whether refundable or adjustable.", "LSE-05 / AGR-13 always captured."),
        ("GR-12", "Art 41 Explanation", "Counting attorneys", "Persons belonging to the same firm count as one person.",
         "POA-02 validation."),
        ("GR-13", "Sec 3", "Exemptions",
         "Instruments listed as exempt under an Article (and Govt-executed instruments exempt under Sec 3) carry no duty.",
         "Exemption inputs (Purpose = E) are checked before calculation."),
    ]
    write_table(ws, ["Rule", "Section / Article", "Topic", "Legal provision (summary)", "System implementation"],
                rows, [8, 26, 34, 80, 60], freeze="B2")


def sheet_examples(wb):
    ws = wb.create_sheet("Worked_Examples")
    rows = []
    rule = {r["key"]: r for r in D.RULES}
    for ex in D.EXAMPLES:
        r = rule.get(ex[3])
        rows.append(list(ex[:4]) + [r["instrument"] if r else "?", ex[4]])
    write_table(ws, ["Example", "Scenario", "Key inputs captured", "Expected Article", "Instrument", "Expected duty"],
                rows, [9, 55, 70, 14, 40, 45], freeze="B2")
    for r_i in range(2, len(rows) + 2):
        ws.cell(row=r_i, column=4).font = Font(bold=True, color="1F4E78")
        ws.cell(row=r_i, column=4).fill = EX_FILL


def main():
    unreached = validate()
    missing_ex = [e[3] for e in D.EXAMPLES if e[3] not in RULE_KEYS]
    if missing_ex:
        raise SystemExit(f"Examples reference unknown rules: {missing_ex}")
    wb = Workbook()
    sheet_readme(wb, unreached)
    sheet_lookups(wb)
    sheet_inputs(wb)
    sheet_rules(wb)
    sheet_matrix(wb)
    sheet_flow(wb)
    sheet_family(wb)
    sheet_slab(wb)
    sheet_general(wb)
    sheet_examples(wb)
    wb.save(OUT)
    print(f"Saved {OUT}")
    print(f"Inputs: {len(D.INPUTS)}  Rules: {len(D.RULES)}  Flow rows: {len(flow_rows())}")
    print(f"Rules not reached by flow: {unreached}")


if __name__ == "__main__":
    main()
