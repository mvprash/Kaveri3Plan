"""One-off extractor for Acts_Rules PDFs/DOCX used in guest-services legal review."""
from pathlib import Path
from pypdf import PdfReader
from docx import Document

base = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Acts_Rules")
out = Path(r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\_tmp_extract")
out.mkdir(exist_ok=True)

docs = [
    base / "Document" / "the_registration_act,_1908.pdf",
    base / "Document" / "THE KARNATAKA STAMP ACT 1957.pdf",
    base / "Document" / "Karnataka Stamp Act 1957 Schedule 2022.pdf",
    base / "Document" / "The Karnataka Registration Rules 1965.pdf",
    base / "Document" / "THE KARNATAKA STAMP- Payment of Stamp Duty by means of e-Stamping.pdf",
    base / "Document" / "KarnatakaSocietiesRegistration(Amendment)Rules,2021.pdf",
    base / "Document" / "RegistrationActNotification.pdf",
    base / "Document" / "Lo-61-2019-20 (2).pdf",
    base / "Marriage" / "The Special Marriage Act, 1954.pdf",
    base / "Marriage" / "SpecialMarriage(Karnataka)Rules1961.pdf",
    base / "Marriage" / "Hindu Marriage Act, 1955.pdf",
    base / "Marriage" / "RD48MNMU2023-Notification-marriage.pdf",
]

docx_files = [
    base
    / "Document"
    / "Karnataka Stamp (Constitution of Central Valuation Committee for Estimation, Publication and Revision of Market Value Guidelines of Properties) Rules, 2003.docx",
    base / "Document" / "Karnataka Stamp (Franking Impression Of Stamps) Rules, 2000.docx",
    base / "Marriage" / "REGISTRATIONOFHINDUMARRIAGE_KARNATAKARULES_1966.docx",
    base / "Marriage" / "SpecialMarriageFees.docx",
    base / "Document" / "INSTRUMENTS GOVERNED BY THE STAMP ACT 1899.docx",
]


def extract_pdf(p: Path):
    r = PdfReader(str(p))
    texts = []
    for i, page in enumerate(r.pages):
        try:
            t = page.extract_text() or ""
        except Exception as e:
            t = f"[extract error page {i}: {e}]"
        texts.append(f"\n---PAGE {i+1}---\n{t}")
    return "\n".join(texts), len(r.pages)


def extract_docx(p: Path):
    d = Document(str(p))
    paras = [para.text for para in d.paragraphs if para.text.strip()]
    for ti, table in enumerate(d.tables):
        paras.append(f"\n---TABLE {ti+1}---")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            paras.append(" | ".join(cells))
    return "\n".join(paras)


for p in docs:
    if not p.exists():
        print("MISSING", p)
        continue
    text, pages = extract_pdf(p)
    dest = out / (p.stem[:80] + ".txt")
    dest.write_text(text, encoding="utf-8", errors="replace")
    print(f"PDF {p.name}: {pages} pages -> {dest.name} ({len(text)} chars)")

for p in docx_files:
    if not p.exists():
        print("MISSING", p)
        continue
    text = extract_docx(p)
    dest = out / (p.stem[:80] + ".txt")
    dest.write_text(text, encoding="utf-8", errors="replace")
    print(f"DOCX {p.name}: -> {dest.name} ({len(text)} chars)")
