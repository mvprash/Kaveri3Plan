# -*- coding: utf-8 -*-
"""Create Document_Registration_requirement_07092026_v1.1.docx

Adds Acts, sections, Rules and notifications for:
  - Registration Appeal
  - Will After the death of the testator
sourced from Acts_Rules/Document/.
"""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(
    r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Requirement Discussions\Daily Reports"
)
SRC = BASE / "Document_Registration_requirement_07092026.docx"
DST = BASE / "Document_Registration_requirement_07092026_v1.1.docx"

HEADING_FONT = "Segoe UI"
HEADING_SIZE = Pt(14.5)
SUBHEAD_SIZE = Pt(13.5)


def set_cell_text(cell, text: str) -> None:
    paras = cell.paragraphs
    first = paras[0]
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        first.add_run(text)


def add_heading(doc: Document, text: str, size=HEADING_SIZE):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.name = HEADING_FONT
    run.font.size = size
    return para


def add_note(doc: Document, bold_prefix: str, rest: str):
    para = doc.add_paragraph()
    r1 = para.add_run(bold_prefix)
    r1.bold = True
    r1.font.name = HEADING_FONT
    r1.font.size = SUBHEAD_SIZE
    r2 = para.add_run(rest)
    r2.font.name = HEADING_FONT
    r2.font.size = SUBHEAD_SIZE
    return para


def make_3col_template(doc: Document) -> Table:
    style_name = doc.tables[0].style.name if doc.tables and doc.tables[0].style else None
    tbl = doc.add_table(rows=2, cols=3)
    if style_name:
        try:
            tbl.style = style_name
        except KeyError:
            pass
    for i, h in enumerate(["A", "B", "C"]):
        cell = tbl.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.makeelement(
            qn("w:shd"),
            {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "D9E2F3"},
        )
        tc_pr.append(shd)
    for i, v in enumerate(["x", "y", "z"]):
        cell = tbl.rows[1].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(v)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
    tbl._tbl.getparent().remove(tbl._tbl)
    return tbl


def clone_table_after(paragraph, doc: Document, template: Table,
                      headers: list[str], rows: list[list[str]]) -> Table:
    new_tbl = deepcopy(template._tbl)
    paragraph._p.addnext(new_tbl)
    table = Table(new_tbl, paragraph._parent)

    data_tr = table.rows[1]._tr if len(table.rows) > 1 else table.rows[0]._tr
    for row in list(table.rows)[1:]:
        new_tbl.remove(row._tr)

    for i, head in enumerate(headers):
        if i < len(table.rows[0].cells):
            set_cell_text(table.rows[0].cells[i], head)

    template_tr = deepcopy(data_tr)
    for _ in range(len(rows)):
        new_tbl.append(deepcopy(template_tr))

    for ri, values in enumerate(rows, start=1):
        for ci, value in enumerate(values):
            if ci < len(table.rows[ri].cells):
                set_cell_text(table.rows[ri].cells[ci], value)
    return table


def add_version_row(doc: Document, version_text: str) -> None:
    meta = doc.tables[0]
    meta._tbl.append(deepcopy(meta.rows[-1]._tr))
    set_cell_text(meta.rows[-1].cells[0], "Version")
    set_cell_text(meta.rows[-1].cells[1], version_text)


ACTS_ROWS = [
    [
        "The Registration Act, 1908 (Central Act 16 of 1908)",
        "Primary",
        "Registration Appeal (Part XII — Secs. 71–77); Wills presentation & deposit "
        "(Parts VIII–IX — Secs. 40–46); wills may be presented/deposited at any time "
        "(Sec. 27); optional registration of wills (Sec. 18)",
    ],
    [
        "The Registration (Karnataka Amendment) Act, 2023 (Karnataka Act 47 of 2024)",
        "Related — Appeal",
        "Sec. 22-D — appeal against District Registrar’s order cancelling registration "
        "under Sec. 22-C (forged / prohibited documents)",
    ],
    [
        "The Karnataka Registration Rules, 1965",
        "Primary (Rules)",
        "Appeals & enquiries (Ch. XXV — Rules 175–191); Wills / authorities to adopt "
        "(Ch. XIV — Rules 83–86); sealed covers containing wills (Ch. XV — Rules 87–93); "
        "withdrawal of sealed covers (Ch. XXXII — Rule 213)",
    ],
    [
        "The Indian Succession Act, 1925",
        "Related — Will after death (reference)",
        "Probate / letters of administration and succession proof often accompany "
        "post-death will registration or opening of deposited sealed covers — not in "
        "Acts_Rules/Document; confirm with Domain Expert whether Kaveri 3.0 captures "
        "probate/court order as a prerequisite",
    ],
]

SECTIONS_APPEAL = [
    [
        "Sec. 71",
        "Reasons for refusal to register to be recorded",
        "SRO must record reasons — starting point for appeal / application",
    ],
    [
        "Sec. 72",
        "Appeal to Registrar — refusal on ground other than denial of execution",
        "Primary Registration Appeal path to District Registrar",
    ],
    [
        "Sec. 73",
        "Application to Registrar — refusal on ground of denial of execution",
        "Separate application (not styled as Sec. 72 appeal) when execution is denied",
    ],
    [
        "Sec. 74",
        "Procedure of Registrar on such application",
        "Enquiry workflow at DR login",
    ],
    [
        "Sec. 75",
        "Order by Registrar to register and procedure thereon",
        "Direction to register; re-presentation / status 'Ordered to register'",
    ],
    [
        "Sec. 76",
        "Order of refusal by Registrar",
        "DR refusal — end of departmental appeal (unless Sec. 22-D path applies)",
    ],
    [
        "Sec. 77",
        "Suit in case of order of refusal by Registrar",
        "Civil Court suit within 30 days — court-ordered registration status",
    ],
    [
        "Sec. 22-D (Kar. Amd. 2023)",
        "Appeal against District Registrar cancellation under Sec. 22-C",
        "Appeal of forged/prohibited document cancellation — parallel appeal track",
    ],
    [
        "Sec. 68",
        "Power of Registrar to superintend and control Sub-Registrars",
        "DR oversight of SRO refusal / appeal handling",
    ],
]

SECTIONS_WILL = [
    [
        "Sec. 18",
        "Documents of which registration is optional — includes wills",
        "Will registration is optional (not compulsory under Sec. 17)",
    ],
    [
        "Sec. 27",
        "Wills may be presented or deposited at any time",
        "No 4-month presentation bar for wills — including after death of testator",
    ],
    [
        "Sec. 40",
        "Persons entitled to present wills and authorities to adopt",
        "Who may present after death (executor / legatee / person claiming under will)",
    ],
    [
        "Sec. 41",
        "Registration of wills and authorities to adopt",
        "Procedure for registering a will when presented (Book 3)",
    ],
    [
        "Sec. 42",
        "Deposit of wills",
        "Sealed-cover deposit with Registrar during testator’s lifetime",
    ],
    [
        "Sec. 43",
        "Procedure on deposit of wills",
        "Receipt / entry of deposited sealed cover",
    ],
    [
        "Sec. 44",
        "Withdrawal of sealed cover deposited under Sec. 42",
        "Testator may withdraw sealed cover during lifetime",
    ],
    [
        "Sec. 45",
        "Proceedings on death of depositor",
        "Core section for Will after death of testator — opening / delivery of "
        "deposited sealed cover to Court or entitled person as prescribed",
    ],
    [
        "Sec. 46",
        "Saving of certain enactments and powers of Courts",
        "Court powers over wills / probate not affected",
    ],
    [
        "Sec. 51",
        "Register-books — Book 3 (wills and authorities to adopt)",
        "Statutory book for will registration entries",
    ],
]

RULES_APPEAL = [
    [
        "Rule 171–174 (Ch. XXIV)",
        "Refusal to register — reasons; partial refusal; etc.",
        "Pre-appeal refusal record at SRO",
    ],
    [
        "Rule 175",
        "Appeal against refusal",
        "Intake of Sec. 72 appeal",
    ],
    [
        "Rule 176",
        "Appeal by whom to be preferred",
        "Who can file the appeal",
    ],
    [
        "Rule 177",
        "Persons who can appear in an enquiry connected with a will or authority to adopt",
        "Cross-link — will-related enquiry appearance",
    ],
    [
        "Rule 179–180",
        "Procedure of disposing appeal; endorsement on order",
        "DR enquiry and order endorsement",
    ],
    [
        "Rule 181",
        "Appeal against refusal to register a will",
        "Specific appeal path when the refused document is a will",
    ],
    [
        "Rule 182–183",
        "Refusal based on non-appearance; communication of orders",
        "Special refusal / notice workflow",
    ],
    [
        "Rule 184 / 188",
        "Registration ordered by Registrar or Court; order directing registration after enquiry",
        "Re-presentation and registration after successful appeal / court order",
    ],
    [
        "Rule 185–187, 190–191",
        "File of appeal orders; refusal orders; no appeal when returned at presentant’s "
        "request; limitation on appeals",
        "Appeal file maintenance, limitation and exclusions",
    ],
    [
        "Rule 108–109",
        "Endorsement on document registered under Sec. 74; presented by order of "
        "Registrar or Court",
        "Endorsement templates after appeal / court direction",
    ],
]

RULES_WILL = [
    [
        "Rule 83",
        "Registration of a Will or authority to adopt",
        "SRO/DR procedure when will is presented for registration",
    ],
    [
        "Rule 84",
        "Return of Will or authority to adopt after the death of Testator — unregistered",
        "Directly covers Will after death of the testator when still unregistered",
    ],
    [
        "Rule 85",
        "Registration of revocation or cancellation of a Will or authority to adopt",
        "Revocation / cancellation registration",
    ],
    [
        "Rule 86",
        "Unclaimed Wills",
        "Custody / disposal of unclaimed wills",
    ],
    [
        "Rules 87–89",
        "Sealed covers — manner of entries; deposit by persons; wills sent by post",
        "Deposit workflow (lifetime deposit under Sec. 42)",
    ],
    [
        "Rules 90–93",
        "Endorsements when sealed cover is sent to Court; forwardal; opening procedure",
        "Post-death opening / Court production of deposited will (Sec. 45)",
    ],
    [
        "Rule 213 (Ch. XXXII)",
        "Withdrawal of sealed covers",
        "Withdrawal under Sec. 44",
    ],
    [
        "Rule 177 / 181",
        "Enquiry appearance for will; appeal against refusal to register a will",
        "Overlap with Registration Appeal when will registration is refused",
    ],
]

NOTIF_ROWS = [
    [
        "The Karnataka Registration Rules, 1965 (under Registration Act Sec. 69)",
        "Parent subordinate legislation for appeal and will procedures",
        "Both topics",
    ],
    [
        "RGN 2/2002-03 (1 Apr 2002; w.e.f. 4 Apr 2002)",
        "Document sheets; photograph / digital photo at presentation (Rule 40) — "
        "applies when a will is presented for registration",
        "Will registration",
    ],
    [
        "The Registration (Karnataka Amendment) Act, 2023 — Gazette Extra-ordinary "
        "No. 480 (19 Oct 2024)",
        "Secs. 22-B–22-D, 81-A–81-B — refusal/cancellation of forged documents and "
        "appeal under Sec. 22-D",
        "Registration Appeal (cancellation track)",
    ],
    [
        "RD 403 ESR 85 / RD/46/MNMU/2025 (registration fee table under Sec. 78)",
        "Fee for appeal / application / will registration or deposit as per notified "
        "Table of Fees",
        "Both topics — fee masters",
    ],
    [
        "Karnataka Registration (Amendment) Rules / GSR notifications cited in Rules "
        "1965 (incl. 1971 onwards)",
        "Amendments to will / sealed-cover / appeal procedure text embedded in Rules PDF",
        "Both topics",
    ],
]


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)
    shutil.copy2(SRC, DST)
    doc = Document(str(DST))

    add_version_row(
        doc,
        "1.1 (08-09-2026) — Acts, sections, Rules and notifications added for "
        "Registration Appeal and Will After the death of the testator",
    )

    template = make_3col_template(doc)

    doc.add_paragraph()
    add_note(
        doc,
        "Scope: ",
        "Discussion topics dated 07–08 Sep 2026 — (1) Registration Appeal against "
        "refusal / related Registrar orders; (2) Will after the death of the testator "
        "(presentation / registration and proceedings on deposited sealed covers). "
        "Source folder: Acts_Rules/Document/.",
    )

    doc.add_paragraph()
    h = add_heading(doc, "1. Primary Acts")
    clone_table_after(
        h, doc, template, ["Act / instrument", "Role", "Relevance to topics"], ACTS_ROWS
    )

    doc.add_paragraph()
    add_heading(doc, "2. Relevant sections")

    doc.add_paragraph()
    h = add_heading(doc, "2.1 Registration Appeal", SUBHEAD_SIZE)
    clone_table_after(
        h, doc, template, ["Section", "Topic", "BRD / system relevance"], SECTIONS_APPEAL
    )

    doc.add_paragraph()
    h = add_heading(doc, "2.2 Will After the death of the testator", SUBHEAD_SIZE)
    clone_table_after(
        h, doc, template, ["Section", "Topic", "BRD / system relevance"], SECTIONS_WILL
    )

    doc.add_paragraph()
    add_heading(doc, "3. Relevant Rules — Karnataka Registration Rules, 1965")

    doc.add_paragraph()
    h = add_heading(doc, "3.1 Registration Appeal", SUBHEAD_SIZE)
    clone_table_after(
        h, doc, template, ["Rule", "Requirement", "System feature"], RULES_APPEAL
    )

    doc.add_paragraph()
    h = add_heading(doc, "3.2 Will After the death of the testator", SUBHEAD_SIZE)
    clone_table_after(
        h, doc, template, ["Rule", "Requirement", "System feature"], RULES_WILL
    )

    doc.add_paragraph()
    h = add_heading(doc, "4. Notifications / amendments")
    clone_table_after(
        h,
        doc,
        template,
        ["Instrument", "Effect", "Topics"],
        NOTIF_ROWS,
    )

    doc.add_paragraph()
    add_note(
        doc,
        "Note: ",
        "Schedule mapping — Will after death of testator aligns to modules list #10 "
        "(Sr.17). Registration Appeal under Secs. 72–77 is SRO→DR refusal appeal; "
        "property-valuation / IGRO appeal (modules list #45 / Sr.23) is a separate "
        "track under Stamp Act undervaluation and is not expanded here unless the "
        "committee treats it as in-scope for this discussion.",
    )

    doc.save(str(DST))
    print(f"Wrote {DST}")

    check = Document(str(DST))
    print("tables=", len(check.tables))
    for i, t in enumerate(check.tables):
        print(f"  {i}: {len(t.rows)}x{len(t.columns)} | {t.rows[0].cells[0].text.strip()[:40]}")


if __name__ == "__main__":
    main()
