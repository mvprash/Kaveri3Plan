# -*- coding: utf-8 -*-
"""Create Document_Registration_requirement_09092026_v1.1.docx

Schedule Sr.19 (09-09-2026 to 10-09-2026) —
  Integration module; Integration exemption; Court entry; Liability.
Format aligned to Document_Registration_requirement_08092026_v1.1.docx.
Source: Acts_Rules/Document/; modules list #3, #16, #17, #18.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(
    r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Requirement Discussions\Daily Reports"
)
DST = BASE / "Document_Registration_requirement_09092026_v1.1.docx"

HEADING_FONT = "Segoe UI"
HEADING_SIZE = Pt(14.5)
SUBHEAD_SIZE = Pt(13.5)
BODY_FONT = "Times New Roman"


def set_run_font(run, *, name=BODY_FONT, size=Pt(12), bold=False):
    run.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = size


def set_cell_text(cell, text: str, *, bold=False, size=Pt(10)) -> None:
    paras = cell.paragraphs
    first = paras[0]
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    first.text = ""
    run = first.add_run(text)
    set_run_font(run, size=size, bold=bold)


def shade_header_row(table: Table, fill="D9E2F3") -> None:
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.makeelement(
            qn("w:shd"),
            {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill},
        )
        tc_pr.append(shd)
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True


def add_heading(doc: Document, text: str, size=HEADING_SIZE):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, name=HEADING_FONT, size=size, bold=True)
    return para


def add_note(doc: Document, bold_prefix: str, rest: str):
    para = doc.add_paragraph()
    r1 = para.add_run(bold_prefix)
    set_run_font(r1, name=HEADING_FONT, size=SUBHEAD_SIZE, bold=True)
    r2 = para.add_run(rest)
    set_run_font(r2, name=HEADING_FONT, size=SUBHEAD_SIZE)
    return para


def add_3col_table(doc: Document, headers: list[str], rows: list[list[str]]) -> Table:
    tbl = doc.add_table(rows=1 + len(rows), cols=3)
    tbl.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(tbl.rows[0].cells[i], h, bold=True, size=Pt(12))
    shade_header_row(tbl)
    for ri, values in enumerate(rows, start=1):
        for ci, value in enumerate(values):
            set_cell_text(tbl.rows[ri].cells[ci], value, size=Pt(10))
    return tbl


def add_story(doc: Document, sid: str, actor: str, want: str, so_that: str) -> None:
    para = doc.add_paragraph()
    r0 = para.add_run(f"{sid}. ")
    set_run_font(r0, name=HEADING_FONT, size=Pt(11), bold=True)
    article = "an" if actor[:1].lower() in "aeiou" else "a"
    r1 = para.add_run(f"As {article} {actor}, I want to {want}, so that {so_that}.")
    set_run_font(r1, name=HEADING_FONT, size=Pt(11))


ACTS_ROWS = [
    [
        "The Registration Act, 1908 (Central Act 16 of 1908)",
        "Primary — Court entry; Integration context; Liability visibility",
        "Sec. 89 — Court / revenue / loan certificates and instruments to be sent "
        "and filed in Book 1 (Court entry + institutional intake); Secs. 64–67 — "
        "cross-office memoranda/copies (pattern for outbound integration payloads); "
        "Secs. 17 / 18 — decrees & orders of Court; Sec. 90–91 — exemption of certain "
        "Govt documents from registration (Integration exemption / non-push cases); "
        "Sec. 57 & EC practice — liability notes visible on search / EC",
    ],
    [
        "The Karnataka Registration Rules, 1965",
        "Primary (Rules)",
        "Rule 17 — Book 1 supplements incl. Court / Revenue sale certificates and "
        "institutional filings (Court entry); Rules 148–154 — Certificate of "
        "encumbrance (Liability notes must surface on EC); Rules 184 / 188 — "
        "registration ordered by Registrar or Court; Rule 166 — copy of memo / "
        "Court decree; Ch. XXVII files of Court orders",
    ],
    [
        "The Karnataka Stamp Act, 1957 (+ Schedule)",
        "Related — Integration exemption (stamp-exempt instruments)",
        "Sec. 3 and Schedule exemptions / remissions — instruments that may be "
        "100% stamp-duty exempt and need clear receipt / challan / integration "
        "handling; undervaluation / charge context for property liabilities",
    ],
]

SECTIONS_INTEG = [
    [
        "Secs. 64–67",
        "Memoranda / copies to other registration offices after registration",
        "Statutory outbound transmission pattern — Integration module mirrors "
        "this for Bhoomi / RDPR / other land-record systems (J-slip / XML push)",
    ],
    [
        "Sec. 60–61",
        "Certificate of registration; completion and return of document",
        "Integration trigger after registration is complete / digitally signed",
    ],
    [
        "Sec. 51 / Indexes",
        "Register books and indexes as source of truth for push payloads",
        "Survey / party / property data sent via Integration Report must match "
        "registered entry",
    ],
]

SECTIONS_EXEMPT = [
    [
        "Sec. 90",
        "Exemption of certain documents executed by or in favour of Government",
        "Instruments/maps that need not be registered — candidates for Integration "
        "exemption / non-push configuration",
    ],
    [
        "Sec. 91",
        "Inspection and copies of Sec. 90 documents",
        "Exempted records remain inspectable / copyable under notified fees",
    ],
    [
        "Sec. 17 proviso / State exemptions",
        "State Government may exempt certain leases etc. by notification",
        "Configurable exemption master for article / deed types not requiring "
        "external push",
    ],
    [
        "Stamp Act Sec. 3 + Schedule exemptions",
        "Instruments exempt / remitted from stamp duty",
        "Stamp-exempt registration flows must still complete without false "
        "challan / integration failures (ServiceDesk 31768)",
    ],
]

SECTIONS_COURT = [
    [
        "Sec. 89(2)",
        "Court granting certificate of sale of immovable property — send copy to "
        "registering officer; file in Book 1",
        "Primary Court entry statutory path for sale certificates",
    ],
    [
        "Sec. 89(1),(3),(4)",
        "Loan orders / mortgage instruments / Revenue sale certificates filed in Book 1",
        "Parallel institutional entry paths alongside Court entry",
    ],
    [
        "Sec. 17(1) / (2)",
        "Decrees and orders of Court affecting immovable property",
        "Court decrees presented for registration or filed as directed",
    ],
    [
        "Secs. 75 / 77",
        "Registrar or Court order directing registration after appeal / suit",
        "Court-ordered registration status (distinct from Sec. 89 filing)",
    ],
]

SECTIONS_LIAB = [
    [
        "Sec. 57 / Rules 148–154 (context)",
        "Inspection, search and Certificate of encumbrance",
        "Liability filings must appear (or be removable) on EC / property search "
        "for the correct survey / village / FY",
    ],
    [
        "Sec. 68 / Rule 123 (related)",
        "DR control; cross-notes on modification of earlier entries",
        "Liability order / note may need cross-link to property and later removal "
        "or rectification under DR control",
    ],
    [
        "Stamp Act — charge / undervaluation context",
        "Instruments creating rights/liabilities; Sec. 45-A references",
        "Business liability notes are departmental overlays on property — confirm "
        "with Domain Expert the exact statutory instrument for Kaveri Liability Filing",
    ],
]

RULES_ROWS = [
    [
        "Rule 17 (esp. Part I)",
        "Court / Revenue sale certificates; LA statements; institutional filings "
        "into Book 1 supplements",
        "Court entry + institutional intake storage",
    ],
    [
        "Rules 148–154",
        "Certificate of encumbrance — complete list of acts and encumbrances",
        "Liability notes must be included / excluded correctly on EC",
    ],
    [
        "Rules 184 / 188",
        "Registration ordered by Registrar or Court; endorsement after enquiry",
        "Court-directed registration workflow",
    ],
    [
        "Rule 166",
        "Copy of memo or decree of a Court",
        "Court decree copy handling when property is in another district",
    ],
    [
        "Rule 123",
        "Cross-notes when a communication modifies / cancels an earlier entry",
        "Related when Court / liability communication affects a prior filing",
    ],
    [
        "Rules 202 / Sec. 88",
        "Govt officers / public functionaries exempt from personal appearance",
        "Related exemption practice for institutional presentants (not Integration "
        "exemption per se)",
    ],
]

NOTIF_ROWS = [
    [
        "The Karnataka Registration Rules, 1965 (under Registration Act Sec. 69)",
        "Parent rules for Rule 17 Court/institutional filing and EC rules",
        "Court entry; Liability (EC)",
    ],
    [
        "RD 403 ESR 85 / RD/46/MNMU/2025 (Table of Fees under Sec. 78)",
        "Fees for filing / copies / EC; exempted instruments fee handling",
        "Integration exemption; Court entry",
    ],
    [
        "Bhoomi / RDPR / land-record integration MoUs & technical specs "
        "(departmental — outside Acts_Rules/Document PDF set)",
        "J-slip / XML push contracts, retry, partial upload — Integration module",
        "Integration module",
    ],
    [
        "Stamp Schedule / exemption notifications under Stamp Act",
        "100% stamp-exempt document classes and receipt display rules",
        "Integration exemption",
    ],
]

STORIES_INTEG = [
    (
        "US-INT-01",
        "Sub-Registrar / DEO",
        "see a registered document in the Integration Report after registration "
        "completes and push it to Bhoomi / RDPR (or other configured systems) "
        "with correct survey, party and office payloads",
        "land-record systems receive a complete J-slip / XML without manual rework",
    ),
    (
        "US-INT-02",
        "Sub-Registrar / support user",
        "re-push or diagnose a failed / missing Integration Report entry "
        "(partial upload, reference number not generated, document not listed)",
        "stuck integrations can be cleared without data loss",
    ),
    (
        "US-INT-03",
        "system / integration service",
        "build outbound payloads from the registered Book 1 entry and indexes "
        "(Secs. 51, 60–61 pattern) and acknowledge success/failure per target system",
        "audit trail shows what was sent and when",
    ),
]

STORIES_EXEMPT = [
    (
        "US-IEX-01",
        "Application Admin / Domain Expert",
        "maintain an Integration exemption master (deed type / article / Sec. 90 "
        "class / stamp-exempt class) so selected registrations skip external push",
        "exempt instruments do not clog Integration Report or fail on missing "
        "challan data",
    ),
    (
        "US-IEX-02",
        "Sub-Registrar / citizen",
        "complete registration of a 100% stamp-duty exempt document with clear "
        "receipt details and without false challan / integration errors",
        "exempt flows finish cleanly (ServiceDesk 31768 class of issues)",
    ),
]

STORIES_COURT = [
    (
        "US-CRT-01",
        "Court / departmental user",
        "enter or file a Court sale certificate / Court order under Sec. 89(2) "
        "and Rule 17 so it is stored in the correct Book 1 supplement",
        "the Court entry is searchable and linked to the right property",
    ),
    (
        "US-CRT-02",
        "Sub-Registrar / FDA",
        "capture Court order property particulars (survey, village, parties) once "
        "so the same data appears consistently in citizen and SR logins",
        "citizens see the Court entry and wrong Sy. No. mismatches are avoided",
    ),
    (
        "US-CRT-03",
        "Sub-Registrar",
        "cancel or correct a Court entry when authorised, with OTP / audit where "
        "required",
        "erroneous Court entries do not remain visible on citizen search",
    ),
    (
        "US-CRT-04",
        "citizen / presentant",
        "re-present a document for registration when a Court (or Registrar) has "
        "ordered registration under Rules 184 / 188",
        "Court-directed registration completes with the correct endorsement",
    ),
]

STORIES_LIAB = [
    (
        "US-LIA-01",
        "District Registrar",
        "file a Liability note / order against the correct district, SRO, village "
        "and survey number for a chosen financial year",
        "the liability attaches only to the intended property",
    ),
    (
        "US-LIA-02",
        "Sub-Registrar / citizen (EC user)",
        "see active Liability filings on Certificate of encumbrance / property "
        "search (Rules 148–154) and not see removed or wrong-property liabilities",
        "EC reflects true departmental liability status",
    ),
    (
        "US-LIA-03",
        "District Registrar / Sub-Registrar",
        "search Liability by document number or property number, including legacy "
        "Kaveri 1 records, and remove or correct a wrongly filed liability",
        "FY / jurisdiction / fetch failures do not block liability maintenance",
    ),
]


def build_meta(doc: Document) -> None:
    tbl = doc.add_table(rows=4, cols=2)
    tbl.style = "Table Grid"
    rows = [
        ("Date", "09-09-2026 to 10-09-2026"),
        (
            "Topics",
            "Integration module; Integration exemption; Court entry; Liability "
            "(Liabality) — Schedule Sr.19; modules #3, #16, #17, #18",
        ),
        (
            "Attendees",
            "Kaveri IT Cell, AIGR Computers team, Domain Expert and committee members",
        ),
        (
            "Version",
            "1.1 (15-09-2026) — Acts, sections, Rules, notifications, pain points "
            "and user stories for Integration, Integration exemption, Court entry "
            "and Liability",
        ),
    ]
    for i, (a, b) in enumerate(rows):
        set_cell_text(tbl.rows[i].cells[0], a, bold=True, size=Pt(12))
        set_cell_text(tbl.rows[i].cells[1], b, size=Pt(12))


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    build_meta(doc)

    doc.add_paragraph()
    add_note(
        doc,
        "Scope: ",
        "Discussion topics dated 09–10 Sep 2026 (Schedule Sr.19) — (1) Integration "
        "module — post-registration push to Bhoomi / RDPR / related systems via "
        "Integration Report; (2) Integration exemption — deed / Sec. 90 / "
        "stamp-exempt classes that skip or specially handle external push; "
        "(3) Court entry — Sec. 89 / Rule 17 Court certificates and Court-order "
        "capture; (4) Liability — DR liability filing that must appear correctly "
        "on EC / property search. Source folder: Acts_Rules/Document/. FRUITS "
        "intake (Sr.16) is adjacent but only cross-referenced where Court / "
        "institutional filing overlaps Rule 17.",
    )

    doc.add_paragraph()
    add_heading(doc, "1. Primary Acts")
    add_3col_table(
        doc, ["Act / instrument", "Role", "Relevance to topics"], ACTS_ROWS
    )

    doc.add_paragraph()
    add_heading(doc, "2. Relevant sections")

    doc.add_paragraph()
    add_heading(doc, "2.1 Integration module", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Section", "Topic", "BRD / system relevance"], SECTIONS_INTEG
    )

    doc.add_paragraph()
    add_heading(doc, "2.2 Integration exemption", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Section", "Topic", "BRD / system relevance"], SECTIONS_EXEMPT
    )

    doc.add_paragraph()
    add_heading(doc, "2.3 Court entry", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Section", "Topic", "BRD / system relevance"], SECTIONS_COURT
    )

    doc.add_paragraph()
    add_heading(doc, "2.4 Liability", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Section", "Topic", "BRD / system relevance"], SECTIONS_LIAB
    )

    doc.add_paragraph()
    add_heading(doc, "3. Relevant Rules — Karnataka Registration Rules, 1965")
    add_3col_table(
        doc, ["Rule", "Requirement", "System feature"], RULES_ROWS
    )

    doc.add_paragraph()
    add_heading(doc, "4. Notifications / amendments")
    add_3col_table(
        doc, ["Instrument", "Effect", "Topics"], NOTIF_ROWS
    )

    doc.add_paragraph()
    add_heading(doc, "5. Pain Points")
    add_note(
        doc,
        "Source: ",
        "ServiceDeskIssuesList.xlsx (OverallList + Categorized).",
    )

    add_heading(doc, "5.1 Integration module", SUBHEAD_SIZE)
    for rid, subj in [
        ("11870 / 18110", "Integration Report / Document number not reflected in Integration Report"),
        ("16847 / 19405 / 19669", "Unable to push from Integration Report to Bhoomi / J-slip not showing"),
        ("23558 / 23977 / 29214 / 30472", "Document not in Integration Report / reconveyance not showing"),
        ("28334 / 30496 / 30742", "RDPR ref not received; XML not sent; J-slip upload from Integration Report"),
        ("Categorized Bhoomi", "J-slip not generated / partial upload / data not found (large pending counts)"),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{rid}: ")
        set_run_font(r1, name=HEADING_FONT, size=Pt(10.5), bold=True)
        r2 = p.add_run(subj)
        set_run_font(r2, name=HEADING_FONT, size=Pt(10.5))

    add_heading(doc, "5.2 Integration exemption", SUBHEAD_SIZE)
    for rid, subj in [
        (
            "31768",
            "Stamp duty 100% exempted documents — challan reference not showing in receipt details",
        ),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{rid}: ")
        set_run_font(r1, name=HEADING_FONT, size=Pt(10.5), bold=True)
        r2 = p.add_run(subj)
        set_run_font(r2, name=HEADING_FONT, size=Pt(10.5))

    add_heading(doc, "5.3 Court entry", SUBHEAD_SIZE)
    for rid, subj in [
        ("5184", "Court entry still not reflecting in citizen login after dept entry"),
        ("21896", "Court entry to be cancelled"),
        ("27219 / 27231", "Court Order Issue"),
        ("29744", "Court order Sy. No. mismatch between FDA and SR login"),
        ("94903 / 93570…", "OTP / cancel court case failures (Categorized)"),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{rid}: ")
        set_run_font(r1, name=HEADING_FONT, size=Pt(10.5), bold=True)
        r2 = p.add_run(subj)
        set_run_font(r2, name=HEADING_FONT, size=Pt(10.5))

    add_heading(doc, "5.4 Liability", SUBHEAD_SIZE)
    for rid, subj in [
        ("20687", "Liability Filing issue"),
        ("29127", "Liability financial year issue / FY 2015-16 not reflecting in DR login"),
        ("29332", "Liability need remove in EC"),
        ("30317", "Liability Entry wrongly displaying for different Sy. No. / village"),
        ("95242", "Wrong district SRO names while filing Liability Note in DR login"),
        ("95253 / 94932…", "Kaveri 1 documents not fetched for Liability Filing"),
        ("94109", "Unable to fetch liability data in SR and DR login"),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{rid}: ")
        set_run_font(r1, name=HEADING_FONT, size=Pt(10.5), bold=True)
        r2 = p.add_run(subj)
        set_run_font(r2, name=HEADING_FONT, size=Pt(10.5))

    doc.add_paragraph()
    add_heading(doc, "6. User Stories")
    intro = doc.add_paragraph()
    r = intro.add_run(
        "Derived from Schedule Sr.19 Acts, sections, Rules and ServiceDesk "
        "(modules #3 Integration, #16 Integration exemption, #17 Court entry, "
        "#18 Liability)."
    )
    set_run_font(r, name=HEADING_FONT, size=Pt(11))

    doc.add_paragraph()
    add_heading(doc, "6.1 Integration module", SUBHEAD_SIZE)
    for story in STORIES_INTEG:
        add_story(doc, *story)

    doc.add_paragraph()
    add_heading(doc, "6.2 Integration exemption", SUBHEAD_SIZE)
    for story in STORIES_EXEMPT:
        add_story(doc, *story)

    doc.add_paragraph()
    add_heading(doc, "6.3 Court entry", SUBHEAD_SIZE)
    for story in STORIES_COURT:
        add_story(doc, *story)

    doc.add_paragraph()
    add_heading(doc, "6.4 Liability", SUBHEAD_SIZE)
    for story in STORIES_LIAB:
        add_story(doc, *story)

    doc.add_paragraph()
    add_note(
        doc,
        "Note: ",
        "Schedule mapping — Integration module #3; Integration exemption #16; "
        "Court entry #17; Liability (modules list spelling “Liabality”) #18 — all "
        "under Sr.19. Integration push to Bhoomi/RDPR is largely departmental/"
        "technical; statutory anchors are Secs. 64–67 (outbound memo pattern) and "
        "post–Sec. 60 completion. Confirm with Domain Expert (1) the authoritative "
        "exemption matrix for Integration exemption, and (2) the exact legal "
        "instrument governing Liability Filing beyond EC practice under Rules "
        "148–154. FRUITS bank e-filing remains primarily Sr.16 / Rule 17 Part IV–V.",
    )

    doc.save(str(DST))
    print(f"Wrote {DST}")
    check = Document(str(DST))
    print("tables=", len(check.tables))
    stories = sum(
        1
        for p in check.paragraphs
        if p.text.strip().startswith(("US-INT-", "US-IEX-", "US-CRT-", "US-LIA-"))
    )
    print(f"user_stories={stories}")


if __name__ == "__main__":
    main()
