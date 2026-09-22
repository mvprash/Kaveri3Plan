# -*- coding: utf-8 -*-
"""Create Document_Registration_requirement_08092026_v1.1.docx

Schedule Sr.18 (08-09-2026) — 68(2) correction and Cross-reference Rule 123.
Format aligned to Document_Registration_requirement_11092026_v1.1.docx.
Source: Acts_Rules/Document/; modules list #14, #15.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path(
    r"E:\MVP\Kaveri 3.0\Source Code\Kaveri 3 Plan\Requirement Discussions\Daily Reports"
)
DST = BASE / "Document_Registration_requirement_08092026_v1.1.docx"

HEADING_FONT = "Segoe UI"
HEADING_SIZE = Pt(14.5)
SUBHEAD_SIZE = Pt(13.5)
BODY_FONT = "Times New Roman"


def set_run_font(run, *, name=BODY_FONT, size=Pt(12), bold=False):
    run.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = size


def set_cell_text(cell, text: str, *, bold=False, size=Pt(10)) -> None:
    paras = cell.paragraphs
    first = paras[0]
    for extra in paras[1:]:
        extra._element.getparent().remove(extra._element)
    first.text = ""
    run = first.add_run(text)
    set_run_font(run, size=size, bold=bold)


def shade_header_row(table: Table, fill="D9E2F3") -> None:
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.makeelement(
            qn("w:shd"),
            {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill},
        )
        tc_pr.append(shd)
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True


def add_heading(doc: Document, text: str, size=HEADING_SIZE):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, name=HEADING_FONT, size=size, bold=True)
    return para


def add_note(doc: Document, bold_prefix: str, rest: str):
    para = doc.add_paragraph()
    r1 = para.add_run(bold_prefix)
    set_run_font(r1, name=HEADING_FONT, size=SUBHEAD_SIZE, bold=True)
    r2 = para.add_run(rest)
    set_run_font(r2, name=HEADING_FONT, size=SUBHEAD_SIZE)
    return para


def add_3col_table(doc: Document, headers: list[str], rows: list[list[str]]) -> Table:
    tbl = doc.add_table(rows=1 + len(rows), cols=3)
    tbl.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(tbl.rows[0].cells[i], h, bold=True, size=Pt(12))
    shade_header_row(tbl)
    for ri, values in enumerate(rows, start=1):
        for ci, value in enumerate(values):
            set_cell_text(tbl.rows[ri].cells[ci], value, size=Pt(10))
    return tbl


def add_story(doc: Document, sid: str, actor: str, want: str, so_that: str) -> None:
    para = doc.add_paragraph()
    r0 = para.add_run(f"{sid}. ")
    set_run_font(r0, name=HEADING_FONT, size=Pt(11), bold=True)
    r1 = para.add_run(f"As a {actor}, I want to {want}, so that {so_that}.")
    set_run_font(r1, name=HEADING_FONT, size=Pt(11))


ACTS_ROWS = [
    [
        "The Registration Act, 1908 (Central Act 16 of 1908)",
        "Primary — Sec. 68(2) correction",
        "Sec. 68(2) — District Registrar may issue any order consistent with the Act "
        "for rectification of any error regarding the book or the office in which a "
        "document has been registered (on complaint or otherwise); Sec. 68(1) — "
        "SRO under DR superintendence; Secs. 51, 54–55 — books and indexes that "
        "must stay consistent after correction",
    ],
    [
        "The Karnataka Registration Rules, 1965",
        "Primary (Rules) — 68(2) correction & Rule 123 cross-reference",
        "Ch. XXIII Errors in Registration (Rules 167–169) — wrong book / wrong "
        "office correction under Sec. 68 direction; Rule 123 — cross-notes on "
        "revocation / cancellation / rectification / modification of a previously "
        "registered or Rule 17-filed document, including Index II / index item notes; "
        "Rule 165 — erratum and cross-references in indexes for memorandum corrections",
    ],
]

SECTIONS_68 = [
    [
        "Sec. 68(1)",
        "Registrar superintendence and control of Sub-Registrars",
        "DR login owns oversight of SRO registration acts / omissions feeding "
        "correction cases",
    ],
    [
        "Sec. 68(2)",
        "Registrar’s order to rectify error regarding book or office of registration",
        "Core 68(2) correction module — DR order (on complaint or suo motu) to "
        "fix wrong book / wrong office / related registration error; triggers "
        "re-registration / re-copy / re-index workflow without fresh fee where Rules "
        "so provide",
    ],
    [
        "Sec. 51",
        "Register-books to be kept",
        "Target book / volume / page after correction must match statutory books",
    ],
    [
        "Secs. 54–55",
        "Indexes Nos. I–IV",
        "Index entries must be corrected / cross-referenced when book or particulars "
        "are rectified (see Rules 123(ii), 165, 167(iv))",
    ],
    [
        "Secs. 64–66",
        "Memoranda / copies to other offices",
        "Wrong-office / wrong-book corrections may require free re-forwardal of "
        "memo/copy (Rule 168–169)",
    ],
]

SECTIONS_123 = [
    [
        "Sec. 89 / Rule 17 filing (context)",
        "Documents / returns filed under Sec. 89 and Rule 17 supplements",
        "Rule 123 also applies when a later document or communication affects a "
        "Sec. 89 / Rule 17 filed paper (e.g. land acquisition return)",
    ],
    [
        "Sec. 17 / optional & compulsory registration (context)",
        "Later deed that revokes, cancels, rectifies or modifies an earlier deed",
        "New registration under ordinary presentation rules; Rule 123 then places "
        "cross-notes on both entries",
    ],
]

RULES_68 = [
    [
        "Rule 167 (Ch. XXIII)",
        "Document registered / copied in a wrong Book — DR sanction; fresh copy in "
        "proper book; red-ink footnotes both ways; indexes not cancelled but "
        "cross-referenced",
        "68(2) wrong-book correction workflow; dual book footnotes; re-index with "
        "cross-references",
    ],
    [
        "Rule 168",
        "Correction when memorandum / copy under Secs. 64–67 was for wrong Book — "
        "error notice / fresh memo free of cost",
        "Cross-office memo repair after wrong-book registration",
    ],
    [
        "Rule 169",
        "Document registered in a wrong office — parties advised to apply to "
        "Registrar under Sec. 68; fresh registration without fee; original office "
        "forwards free memo/copy to proper office",
        "Core 68(2) wrong-office correction path",
    ],
    [
        "Rule 165",
        "Corrections in memoranda — erratum filed in Supplement Book 1 Part I; "
        "note on original; indexes corrected with cross-references",
        "Memo / index correction supporting 68(2) and cross-office consistency",
    ],
]

RULES_123 = [
    [
        "Rule 123(i)",
        "On registration of a document (or receipt of Revenue / Court "
        "communication) that revokes, cancels, rectifies an error in, or modifies "
        "a previously registered / Rule 17-filed document — enter mutual footnotes "
        "on the later and earlier entries",
        "Cross-reference module — bidirectional notes linking Document No. / "
        "volume / page / supplement part",
    ],
    [
        "Rule 123(ii)",
        "If immovable property is affected — corresponding note in Index No. II; "
        "if an index item is rectified — note in Index I/II/III/IV against that item",
        "Index cross-reference and index-item rectification notes",
    ],
    [
        "Rule 124 (related)",
        "Note when Court declares registered document a forgery / false personation",
        "Related foot-note path (not the primary Rule 123 revoke/modify track)",
    ],
    [
        "Rule 17 (related filing)",
        "Supplement parts that Rule 123 may point to (filed returns / documents)",
        "Cross-link targets include Rule 17 filed papers and Sec. 89 returns",
    ],
]

NOTIF_ROWS = [
    [
        "The Karnataka Registration Rules, 1965 (under Registration Act Sec. 69)",
        "Parent rules for Errors in Registration (Ch. XXIII) and Rule 123 "
        "cancellation / rectification / modification notes",
        "Both topics",
    ],
    [
        "RD 403 ESR 85 / RD/46/MNMU/2025 (Table of Fees under Sec. 78)",
        "Fresh registration under Rule 169 after Sec. 68 direction is without levy "
        "of fee; other correction / note actions as per fee table practice",
        "68(2) correction — fee exception",
    ],
    [
        "ServiceDesk / Kaveri 2.0 68(2) correction module practice",
        "Operational 68(2) correction, re-scan, index correction and verify steps "
        "used in current system — to be re-engineered under Kaveri 3.0",
        "68(2) correction — as-is pain points",
    ],
]

STORIES_68 = [
    (
        "US-682-01",
        "citizen / executant / claimant",
        "complain to the District Registrar under Sec. 68(2) that my document was "
        "registered in the wrong book or wrong office (or has a related registration "
        "error)",
        "the Registrar can issue a rectification order consistent with the Act",
    ),
    (
        "US-682-02",
        "District Registrar",
        "issue a Sec. 68(2) order (on complaint or otherwise) directing "
        "rectification of the book or office error, and track the case to closure",
        "SRO action and status are auditable against the Registrar’s order",
    ),
    (
        "US-682-03",
        "Sub-Registrar",
        "after DR sanction under Rule 167, re-copy a wrong-book entry into the "
        "proper book with red-ink footnotes both ways and create proper index "
        "entries with cross-references (without cancelling the original serial)",
        "both books and indexes remain consistent and searchable",
    ),
    (
        "US-682-04",
        "Sub-Registrar / citizen",
        "complete wrong-office correction under Rule 169 — advise parties, receive "
        "Sec. 68 direction, re-register without fee with endorsement citing the "
        "order, and forward free memo/copy to the proper office",
        "the document is correctly recorded in the proper office without double fee",
    ),
    (
        "US-682-05",
        "Sub-Registrar / DEO",
        "send an erratum for a memorandum/copy under Rule 165 / 168, file it in "
        "Supplement Book 1 Part I, correct indexes with cross-references, and "
        "re-scan / verify corrected pages where the module requires it",
        "cross-office copies and citizen CC/EC views show the corrected record",
    ),
]

STORIES_123 = [
    (
        "US-123-01",
        "Sub-Registrar / DEO",
        "auto-prompt and save Rule 123(i) mutual footnotes on both entries "
        "(Document No., volume, page, supplement part) when registering a deed "
        "that revokes, cancels, rectifies or modifies an earlier registered or "
        "Rule 17-filed document",
        "anyone reading either entry sees the linked revoke/modify relationship",
    ),
    (
        "US-123-02",
        "Sub-Registrar / DEO",
        "on receipt of a Revenue Officer or Court communication that similarly "
        "revokes / cancels / rectifies / modifies a filed or registered paper, "
        "file the communication and apply the same Rule 123 cross-notes",
        "institutional modifications are visible from both the old and new records",
    ),
    (
        "US-123-03",
        "Sub-Registrar / DEO",
        "for immovable-property cases, write the corresponding Index No. II note, "
        "and when an index item is rectified write the note in Index I/II/III/IV "
        "under Rule 123(ii)",
        "property and name indexes stay aligned with the register footnotes",
    ),
    (
        "US-123-04",
        "citizen / search user",
        "see Rule 123 cross-references when I search, take EC, or view a certified "
        "copy of either the original or the later modifying document",
        "I am not misled by an obsolete register entry that was later modified",
    ),
]


def build_meta(doc: Document) -> None:
    tbl = doc.add_table(rows=4, cols=2)
    tbl.style = "Table Grid"
    rows = [
        ("Date", "08-09-2026"),
        (
            "Topics",
            "Sec. 68(2) correction and Cross-reference Rule 123 — Schedule Sr.18; "
            "modules #14, #15",
        ),
        (
            "Attendees",
            "Kaveri IT Cell, AIGR Computers team, Domain Expert and committee members",
        ),
        (
            "Version",
            "1.1 (15-09-2026) — Acts, sections, Rules, notifications, pain points "
            "and user stories for Sec. 68(2) correction and Rule 123 cross-reference",
        ),
    ]
    for i, (a, b) in enumerate(rows):
        set_cell_text(tbl.rows[i].cells[0], a, bold=True, size=Pt(12))
        set_cell_text(tbl.rows[i].cells[1], b, size=Pt(12))


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    build_meta(doc)

    doc.add_paragraph()
    add_note(
        doc,
        "Scope: ",
        "Discussion topic dated 08-09-2026 (Schedule Sr.18) — (1) Sec. 68(2) "
        "correction of errors regarding the book or office in which a document "
        "was registered (wrong book / wrong office / related rectification under "
        "District Registrar’s order); (2) Cross-reference under Rule 123 when a "
        "later document or communication revokes, cancels, rectifies or modifies "
        "a previously registered or Rule 17-filed entry, including index notes. "
        "Source folder: Acts_Rules/Document/.",
    )

    doc.add_paragraph()
    add_heading(doc, "1. Primary Acts")
    add_3col_table(
        doc, ["Act / instrument", "Role", "Relevance to topics"], ACTS_ROWS
    )

    doc.add_paragraph()
    add_heading(doc, "2. Relevant sections")

    doc.add_paragraph()
    add_heading(doc, "2.1 Sec. 68(2) correction", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Section", "Topic", "BRD / system relevance"], SECTIONS_68
    )

    doc.add_paragraph()
    add_heading(doc, "2.2 Cross-reference Rule 123 (context sections)", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Section", "Topic", "BRD / system relevance"], SECTIONS_123
    )

    doc.add_paragraph()
    add_heading(doc, "3. Relevant Rules — Karnataka Registration Rules, 1965")

    doc.add_paragraph()
    add_heading(doc, "3.1 Sec. 68(2) correction / Errors in Registration", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Rule", "Requirement", "System feature"], RULES_68
    )

    doc.add_paragraph()
    add_heading(doc, "3.2 Cross-reference Rule 123", SUBHEAD_SIZE)
    add_3col_table(
        doc, ["Rule", "Requirement", "System feature"], RULES_123
    )

    doc.add_paragraph()
    add_heading(doc, "4. Notifications / amendments")
    add_3col_table(
        doc, ["Instrument", "Effect", "Topics"], NOTIF_ROWS
    )

    doc.add_paragraph()
    add_heading(doc, "5. Pain Points")
    add_note(
        doc,
        "Source: ",
        "ServiceDeskIssuesList.xlsx (OverallList) — tickets naming 68(2) / "
        "68 correction.",
    )
    pain = [
        ("8788", "68(2) note Re-scanning issue"),
        ("15135", "68(2) correction module Hobli name not working"),
        ("18321", "68(2) Verification issue"),
        ("19748", "68 correction not able to verify document"),
        ("20158", "68(2) Rescan completed but not reflecting in KOS while applying for CC"),
        ("20763 / 20772", "68(2) correction module issue / MODULE error"),
        ("21319 / 27261 / 28486", "68 correction error / unable to update / section correction issue"),
        ("29418", "68(2) Index Correction"),
        ("29586 / 31576", "68(2) correction Issue (e.g. SRO Nelamangala) / 68(2) correction"),
    ]
    for rid, subj in pain:
        p = doc.add_paragraph(style="List Bullet")
        r1 = p.add_run(f"{rid}: ")
        set_run_font(r1, name=HEADING_FONT, size=Pt(10.5), bold=True)
        r2 = p.add_run(subj)
        set_run_font(r2, name=HEADING_FONT, size=Pt(10.5))

    doc.add_paragraph()
    add_heading(doc, "6. User Stories")
    intro = doc.add_paragraph()
    r = intro.add_run(
        "Derived from Schedule Sr.18 Acts, sections and Rules "
        "(68(2) correction — modules #14; Cross-reference Rule 123 — modules #15)."
    )
    set_run_font(r, name=HEADING_FONT, size=Pt(11))

    doc.add_paragraph()
    add_heading(doc, "6.1 Sec. 68(2) correction", SUBHEAD_SIZE)
    for story in STORIES_68:
        add_story(doc, *story)

    doc.add_paragraph()
    add_heading(doc, "6.2 Cross-reference Rule 123", SUBHEAD_SIZE)
    for story in STORIES_123:
        add_story(doc, *story)

    doc.add_paragraph()
    add_note(
        doc,
        "Note: ",
        "Schedule mapping — Sec. 68(2) correction = modules list #14; "
        "Cross-reference Rule 123 = modules list #15 (both under Sr.18). "
        "Rule 123 is about mutual footnotes / index notes when a later instrument "
        "or communication affects an earlier entry — distinct from Sec. 68(2) "
        "rectification of wrong book/office. Pre-registration “send back for "
        "correction” in the SR approval workflow (discussed on 28-08) is a "
        "separate intake path and is not expanded here. Sec. 22-B/22-C forged/"
        "prohibited cancellation (Kar. Amd. 2023) is a parallel track under "
        "Sr.12/appeal discussions, not the Rule 123 revoke/modify note.",
    )

    doc.save(str(DST))
    print(f"Wrote {DST}")
    check = Document(str(DST))
    print("tables=", len(check.tables))
    stories = sum(
        1
        for p in check.paragraphs
        if p.text.strip().startswith(("US-682-", "US-123-"))
    )
    print(f"user_stories={stories}")


if __name__ == "__main__":
    main()
