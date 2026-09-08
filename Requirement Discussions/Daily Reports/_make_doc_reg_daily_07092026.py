# -*- coding: utf-8 -*-
"""Create Document_Registration_requirement_07092026.docx from the working
02-09 template (keeps Word styles/theme that open correctly), replacing body
with Acts/Rules/sections/notifications for:
  - Registration Appeal
  - Will after the death of the testator
"""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.stdout.reconfigure(encoding="utf-8")

BASE = Path("/workspace/Requirement Discussions/Daily Reports")
SRC = BASE / "Document_Registration_requirement_02092026_v1.1.docx"
DST = BASE / "Document_Registration_requirement_07092026.docx"

HEADER_FILL = "1F4E79"


def clear_body(doc: Document) -> None:
    body = doc.element.body
    # keep sectPr at end
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is sectPr:
            continue
        body.remove(child)


def insert_paragraph(doc: Document, text: str = "", style: str | None = None) -> Paragraph:
    p = doc.add_paragraph(text, style=style)
    return p


def set_cell_text(cell, text: str, *, header: bool = False, size_pt: float = 9.5) -> None:
    # Clear existing paragraphs except first
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    p = cell.paragraphs[0]
    # clear runs
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.size = Pt(size_pt)
    run.bold = bool(header)
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)
        from docx.oxml import OxmlElement

        tcPr = cell._tc.get_or_add_tcPr()
        # remove old shd
        for old in tcPr.findall(qn("w:shd")):
            tcPr.remove(old)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), HEADER_FILL)
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> Table:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], h, header=True, size_pt=10)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            set_cell_text(table.rows[i + 1].cells[j], val, header=False)
    return table


def main() -> None:
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    shutil.copy2(SRC, DST)
    doc = Document(str(DST))
    clear_body(doc)

    # Meta
    add_table(
        doc,
        ["Field", "Value"],
        [
            ["Date", "07-09-2026"],
            [
                "Topics",
                "Registration Appeal; Will after the death of the testator",
            ],
            [
                "Attendees",
                "Kaveri IT Cell, AIGR Computers team, Committee members (DSR)",
            ],
            [
                "Schedule ref",
                "Sr.17 / Sub-module #10 (Will after Death of testator); "
                "Registration Appeal — Registration Act Part XII (Secs. 71–77) "
                "and Karnataka Rules Ch. XXIV–XXV; also Kar. Amendment Sec. 22-D",
            ],
            ["Version", "1.1 (08-09-2026) — rebuilt from working template"],
        ],
    )
    insert_paragraph(doc, "")

    insert_paragraph(doc, "1. Primary Acts", style="Heading 3")
    insert_paragraph(
        doc,
        "Statutory basis for Registration Appeal and for registration / opening "
        "of wills after the death of the testator.",
    )

    insert_paragraph(doc, "A. Registration Appeal", style="Heading 4")
    add_table(
        doc,
        ["Act", "Sections", "Relevance"],
        [
            [
                "The Registration Act, 1908 (Central Act 16 of 1908)",
                "Part XII — Sec. 71",
                "Sub-Registrar refusing registration (except wrong jurisdiction) "
                "must make a refusal order and record reasons in Book No. 2; "
                "endorse refusal and give a copy on demand.",
            ],
            [
                "",
                "Sec. 72",
                "Appeal to Registrar from Sub-Registrar’s refusal on any ground "
                "other than denial of execution — within 30 days of the order.",
            ],
            [
                "",
                "Sec. 73",
                "Where refusal is on denial of execution: application to Registrar "
                "(not an appeal under Sec. 72) by any person claiming under the "
                "document, within 30 days, verified like a plaint.",
            ],
            [
                "",
                "Sec. 74",
                "Registrar’s enquiry: (a) whether the document was executed; "
                "(b) whether requirements of law for registration were complied with.",
            ],
            [
                "",
                "Sec. 75",
                "If satisfied, Registrar orders registration; document to be "
                "presented within 30 days of the order; registration has effect "
                "as of the original presentation date.",
            ],
            [
                "",
                "Sec. 76",
                "Registrar’s own refusal order (or refusal to direct registration "
                "under Sec. 72 / 75) — reasons in Book No. 2; copy on demand.",
            ],
            [
                "",
                "Sec. 77",
                "Suit in Civil Court within 30 days of Registrar’s refusal order "
                "under Sec. 72 or Sec. 76; if decree directs registration, present "
                "within 30 days of decree.",
            ],
            [
                "",
                "Sec. 68–69 (supporting)",
                "Registrar’s superintendence / control over Sub-Registrars; "
                "IGR’s power to make rules (including appeal procedure).",
            ],
            [
                "Registration (Karnataka Amendment) Act, 2023 "
                "(Karnataka Act 47 of 2024)",
                "Sec. 22-B",
                "Mandatory refusal to register forged documents, prohibited "
                "transactions, attached property transfers, and notified classes.",
            ],
            [
                "",
                "Sec. 22-C",
                "District Registrar may cancel registration made in contravention "
                "of Sec. 22-B (suo motu or on complaint), after notice — "
                "cancellation entered in books/indexes.",
            ],
            [
                "",
                "Sec. 22-D",
                "Appeal to Inspector General of Registration within 30 days from "
                "cancellation under Sec. 22-C; IGR may confirm, modify or cancel "
                "the District Registrar’s order.",
            ],
            [
                "",
                "Sec. 69(1)(m) (inserted)",
                "IGR rule-making power for the process of cancellation under Sec. 22-C.",
            ],
        ],
    )
    insert_paragraph(doc, "")

    insert_paragraph(doc, "B. Will after the death of the testator", style="Heading 4")
    add_table(
        doc,
        ["Act", "Sections", "Relevance"],
        [
            [
                "The Registration Act, 1908",
                "Sec. 27",
                "A will may be presented for registration or deposited at any time "
                "(no ordinary four-month presentation bar).",
            ],
            [
                "",
                "Sec. 40",
                "Who may present: the testator; or after his death any person "
                "claiming as executor or otherwise under the will. (Authority to "
                "adopt: donor / after death the donee or adoptive son.)",
            ],
            [
                "",
                "Sec. 41(1)",
                "Will presented by the testator — registered like any other document.",
            ],
            [
                "",
                "Sec. 41(2)",
                "Will presented after death by a person under Sec. 40 — register only "
                "if the Registering Officer is satisfied that: (a) the will was "
                "executed by the testator; (b) the testator is dead; and (c) the "
                "presentant is entitled under Sec. 40.",
            ],
            [
                "",
                "Sec. 42–44",
                "Deposit of wills in sealed cover with Registrar (Book 5); "
                "withdrawal only by living testator / authorised agent.",
            ],
            [
                "",
                "Sec. 45",
                "On death of depositor: if application is made and death is "
                "satisfied, Registrar opens the sealed cover in applicant’s "
                "presence, copies contents into Book No. 3 at applicant’s expense, "
                "and re-deposits the original will.",
            ],
            [
                "",
                "Sec. 46",
                "Saving for Succession / Probate enactments and Court’s power to "
                "compel production of any will deposited / registered.",
            ],
            [
                "",
                "Sec. 51 / Books 2, 3, 5",
                "Book 2 — reasons for refusal; Book 3 — Register of Wills and "
                "authorities to adopt; Book 5 — deposits of wills (with index).",
            ],
            [
                "",
                "Sec. 71–77 (cross-link)",
                "Refusal to register a will after death can be challenged under "
                "Part XII; Rule 181 specially allows appeal/application by any "
                "executor named in the will.",
            ],
            [
                "Indian Succession Act, 1925 (supporting)",
                "Cross-ref via Reg. Act Sec. 46",
                "Registration / deposit of a will does not replace probate where "
                "required; Court production of deposited wills remains available.",
            ],
        ],
    )
    insert_paragraph(doc, "")

    insert_paragraph(
        doc,
        "2. Primary Rules — Karnataka Registration Rules, 1965",
        style="Heading 3",
    )

    insert_paragraph(doc, "A. Registration Appeal (Ch. XXIV–XXV)", style="Heading 4")
    add_table(
        doc,
        ["Rule", "Module mapping", "What it covers"],
        [
            [
                "Rule 171 (Ch. XXIV)",
                "Refusal to register",
                "Reasons for refusal recorded at once in Book 2 — language "
                "(Sec. 19), interlineations (Sec. 20), description (Secs. 21–22 / "
                "Rule 15), maps (Sec. 21(4)), date of execution (Rule 50), time "
                "(Secs. 23–26, 72, 75, 77), appearance / denial (Secs. 34–35), "
                "identity, will/authority after death where Sec. 41(2) not "
                "satisfied, fees (Sec. 80), etc.",
            ],
            [
                "Rules 172–174",
                "Refusal variants",
                "Executants appearing at different times; partial refusal; "
                "documents executed by registering officers.",
            ],
            [
                "Rule 175",
                "Appeal / Sec. 73 application intake",
                "Appeal under Sec. 72 or application under Sec. 73 in writing to "
                "the District Registrar (or officer in charge), with copy of "
                "refusal order and the original document (time may be given to "
                "produce it).",
            ],
            [
                "Rule 176",
                "Who may prefer",
                "Sec. 72 appeal by executant / claimant or duly authorised agent / "
                "vakil; Sec. 73 application by person claiming under the document "
                "or Sec. 33 agent. Not accepted by post.",
            ],
            [
                "Rule 177",
                "Enquiry appearance (also wills)",
                "In Sec. 41(2) will enquiry, Sec. 72/73 appeal/application, or "
                "Sec. 74 enquiry — only Advocates Act-qualified vakils (or "
                "authenticated PoA agents) may appear.",
            ],
            [
                "Rules 178–180",
                "Appeal disposal",
                "Unverified Sec. 73 applications may be returned for verification; "
                "hearing dates; refusal to direct registration recorded in Book 2; "
                "order on appeal/application is not endorsed on the document but "
                "kept with case records.",
            ],
            [
                "Rule 181",
                "Appeal — will after death",
                "Appeal/application against refusal to register a will presented "
                "after the testator’s death may be preferred by any executor "
                "appointed under the will.",
            ],
            [
                "Rules 182–188",
                "Orders & post-order registration",
                "Non-appearance of executant (Sec. 72 vs Sec. 73 path); "
                "communication of orders; registration when ordered by Registrar "
                "or Court (endorsements); file of appeal orders/judgments; "
                "Book 2 notes when Registrar refuses to direct registration.",
            ],
            [
                "Rules 189–191",
                "Limits on Registrar",
                "No power on appeal to call for property description; no appeal "
                "when document returned at presentant’s request; limitation on "
                "appeals to Registrar against Sub-Registrar.",
            ],
            [
                "Rule 18 / Rule 109 (supporting)",
                "Court / Registrar-ordered registration",
                "Filing of Registrar’s order or Court decree; endorsement when "
                "document is presented under such order.",
            ],
            [
                "Rule 23 (Minute Book)",
                "Office record",
                "Note suspension, refusal, summons, withdrawal and related events.",
            ],
        ],
    )
    insert_paragraph(doc, "")

    insert_paragraph(
        doc, "B. Will after the death of the testator (Ch. XIV–XV)", style="Heading 4"
    )
    add_table(
        doc,
        ["Rule", "Module mapping", "What it covers"],
        [
            [
                "Rule 83",
                "Will / authority after death — Sec. 41(2) enquiry",
                "Registering Officer must be satisfied on: (a) execution by "
                "testator/donor; (b) death; (c) presentant is claimant/executor "
                "(or adoptee). Record deposition of presentant and identifying "
                "witnesses on the will; prescribed satisfaction endorsement.",
            ],
            [
                "Rule 84",
                "Unregistered will after death",
                "Return of will / authority to adopt after death of testator when "
                "still unregistered — custody / return procedure.",
            ],
            [
                "Rule 85",
                "Revocation / cancellation of will",
                "Revocation or cancellation of a will or authority to adopt is "
                "treated as a distinct instrument and registered in Book 3.",
            ],
            [
                "Rule 86",
                "Unclaimed wills",
                "Wills registered or refused that remain unclaimed beyond the "
                "prescribed period — transfer / safe custody (see also Rule 223).",
            ],
            [
                "Rules 87–88 (Ch. XV)",
                "Sealed covers — Book 5",
                "Entries under Sec. 43 in Book 5; joint wills deposited by both "
                "testators; withdrawal endorsements.",
            ],
            [
                "Rule 89",
                "Wills by post",
                "Wills sent by post are not “presented” or “deposited” under the "
                "Act — return unopened / retain only per prescribed exceptions.",
            ],
            [
                "Rules 90–93",
                "Opening sealed cover on death / Court",
                "Endorsements when opened under Sec. 45 (death) or Sec. 46 "
                "(Court order); copy into Book 3; forwardal to Court with fee "
                "memo and acknowledgment.",
            ],
            [
                "Index No. III (Rules on indexes)",
                "Search / certified copy",
                "Index III for wills — executor / guardian columns; death "
                "particulars when known.",
            ],
            [
                "Rule 223",
                "Safe custody of belated wills",
                "Unclaimed wills / cancellations over two years in Sub-Registry "
                "to be transferred for safe custody.",
            ],
        ],
    )
    insert_paragraph(doc, "")

    insert_paragraph(doc, "3. Notifications / amendments", style="Heading 3")
    add_table(
        doc,
        ["Notification / instrument", "Effect for this topic"],
        [
            [
                "Karnataka Registration Rules, 1965 (under Registration Act Sec. 69)",
                "Parent instrument for Ch. XIV (Wills), Ch. XV (Sealed covers), "
                "Ch. XXIV (Refusal) and Ch. XXV (Appeals and Enquiries).",
            ],
            [
                "Registration (Karnataka Amendment) Act, 2023 "
                "(Karnataka Act 47 of 2024) — Gazette Extra-ordinary No. 480, "
                "Part IVA, 19-10-2024 (President’s assent 08-10-2024)",
                "Inserts Secs. 22-B, 22-C, 22-D (refusal / cancellation / IGR "
                "appeal), Sec. 69(1)(m), and Secs. 81-A / 81-B. Commencement on "
                "date appointed by State Government notification.",
            ],
            [
                "Sec. 69(1)(m) rules (to be notified)",
                "Rules regulating cancellation process under Sec. 22-C — confirm "
                "whether State notification / IGR rules have been issued before "
                "building the cancellation–appeal workflow.",
            ],
            [
                "Table of Registration Fees — Notification No. RD 403 ESR 85 "
                "(27 May 1986) as amended",
                "Fee articles for registration, appeals/applications, searches "
                "and copies (Book 2 / Book 3 / Book 5 extracts).",
            ],
            [
                "RD/46/MNMU/2025 (29-08-2025; w.e.f. 31-08-2025)",
                "Substitutes fee rates in Article I(4)(a) and III(a)(i)/(ii) of "
                "the Table of Fees — applies to fee columns on appeal / will / "
                "certified-copy screens.",
            ],
            [
                "RGN 2/2002-03 (1 Apr 2002; w.e.f. 4 Apr 2002) and related "
                "document-sheet notifications",
                "Document sheets / photography / biometric capture rules that "
                "also apply when a will is presented after death or when "
                "registration is ordered after appeal.",
            ],
        ],
    )
    insert_paragraph(doc, "")

    insert_paragraph(doc, "Pain Points (ServiceDesk — indicative)", style="Heading 3")
    insert_paragraph(
        doc,
        "Few tickets are explicitly titled “appeal”; will / Book-3 issues "
        "surface mainly as payment and certified-copy failures:",
    )
    add_table(
        doc,
        ["Pain point", "Evidence"],
        [
            ["Payment related issue for WILL", "28795"],
            [
                "Unable to download Book 3 and Book 4 documents "
                "(certified copies of wills / miscellaneous register)",
                "Categorized: 92626, 87534, 82435, 78789; also 29879, 29938, 31252",
            ],
            ["Book-3 CC Issue", "26371"],
        ],
    )
    insert_paragraph(doc, "")
    p = insert_paragraph(doc, "")
    r0 = p.add_run("Note:")
    r0.bold = True
    p.add_run(
        " Discussion focus: map Kaveri 3.0 workflows for (1) Sub-Registrar "
        "refusal → Book 2 → Sec. 72 appeal / Sec. 73 application → Registrar "
        "enquiry → order to register or refuse → optional Sec. 77 suit; "
        "(2) Sec. 22-C cancellation → Sec. 22-D IGR appeal; and (3) will "
        "presentation after death under Sec. 40 / 41(2) (and sealed-cover "
        "opening under Sec. 45), including Rule 83 enquiry and Rule 181 "
        "appeal by executor."
    )

    doc.save(str(DST))
    print(f"Wrote {DST} size={DST.stat().st_size}")

    # Verify
    check = Document(str(DST))
    texts = [p.text for p in check.paragraphs if p.text.strip()]
    print("headings/paras:", len(texts))
    for t in texts[:8]:
        print(" -", t[:90])
    print("tables:", len(check.tables))
    print("meta topics:", check.tables[0].rows[2].cells[1].text[:80])


if __name__ == "__main__":
    main()
