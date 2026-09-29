"""Generate Introduction_to_Kaveri_2.0.pptx (Kaveri 2.0 overview deck)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("Introduction_to_Kaveri_2.0.pptx")

NAVY = RGBColor(0x0B, 0x3D, 0x5C)
TEAL = RGBColor(0x1A, 0x6B, 0x7A)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
GREEN = RGBColor(0x1B, 0x7A, 0x4E)
PURPLE = RGBColor(0x5B, 0x3F, 0x8C)
LIGHT = RGBColor(0xF0, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1E, 0x29, 0x33)
MUTED = RGBColor(0x5A, 0x6A, 0x78)
SUB = RGBColor(0xB8, 0xD0, 0xDC)

W, H = Inches(13.333), Inches(7.5)
FOOTER = "Introduction to Kaveri 2.0  |  Department of Stamps & Registration, Government of Karnataka"

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
_pages = []


def rect(slide, l, t, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None):
    sh = slide.shapes.add_shape(shape, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    sh.shadow.inherit = False
    return sh


def text(slide, l, t, w, h, lines, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullet=False, space=6):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        r = p.add_run()
        r.text = ("\u2022  " + line) if bullet else line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, "Calibri"
        r.font.color.rgb = color
    return box


def shape_text(sh, lines, size=13, color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.name, r.font.color.rgb = Pt(size), "Calibri", color
        r.font.bold = bold if i == 0 else False


def new_slide(title=None, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    if title:
        rect(s, 0, 0, W, Inches(0.95), NAVY)
        rect(s, 0, Inches(0.95), W, Inches(0.07), TEAL)
        text(s, Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.5), title, 26, WHITE, True)
        if subtitle:
            text(s, Inches(0.5), Inches(0.56), Inches(12.3), Inches(0.35), subtitle, 13, SUB)
    _pages.append(s)
    return s


def card(slide, l, t, w, h, title, lines, color=TEAL, size=12):
    rect(slide, l, t, w, h, LIGHT)
    rect(slide, l, t, Inches(0.08), h, color)
    text(slide, l + Inches(0.2), t + Inches(0.1), w - Inches(0.3), Inches(0.4), title, 15, color, True)
    text(slide, l + Inches(0.2), t + Inches(0.5), w - Inches(0.3), h - Inches(0.6), lines, size,
         DARK, bullet=True, space=4)


def section(num, title, points, color):
    s = new_slide()
    rect(s, 0, 0, W, H, NAVY)
    rect(s, 0, Inches(5.2), W, Inches(0.08), color)
    text(s, Inches(0.8), Inches(1.4), Inches(3), Inches(1.4), f"0{num}", 96, color, True)
    text(s, Inches(0.8), Inches(3.0), Inches(11.5), Inches(1.2), title, 36, WHITE, True)
    text(s, Inches(0.8), Inches(5.5), Inches(11.5), Inches(1.2), points, 15, SUB, space=4)
    return s


def flow(slide, top, steps, colors, h=Inches(1.25)):
    n = len(steps)
    gap = Inches(0.12)
    left0 = Inches(0.4)
    w = int((W - 2 * left0 - gap * (n - 1)) / n)
    for i, (head, body) in enumerate(steps):
        l = left0 + i * (w + gap)
        sh = rect(slide, l, top, w, h, colors[i % len(colors)],
                  MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
        sh.adjustments[0] = 0.22
        notch = int(h * 0.22)
        inset = Inches(0.05) if i == 0 else notch
        text(slide, l + inset, top, w - inset - notch, h, head, 12 if n <= 5 else 11, WHITE, True,
             PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, space=0)
        text(slide, l, top + h + Inches(0.1), w, Inches(1.6), body, 11, DARK, align=PP_ALIGN.CENTER,
             space=2)


# 1. Title
s = new_slide()
rect(s, 0, 0, W, H, NAVY)
rect(s, 0, Inches(4.55), W, Inches(0.08), ACCENT)
text(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.0), "Introduction to Kaveri 2.0", 48, WHITE, True)
text(s, Inches(0.8), Inches(2.7), Inches(11.5), Inches(0.8),
     "Karnataka's Digital Property & Document Registration Platform", 24, SUB)
text(s, Inches(0.8), Inches(4.9), Inches(11.5), Inches(1.2),
     ["Department of Stamps & Registration, Government of Karnataka",
      "Architecture  |  Workflow & Integrations  |  Cybersecurity  |  Capacity Building"], 16, WHITE, space=8)

# 2. Agenda
s = new_slide("Agenda", "Four themes covered in this session")
items = [
    ("1", "Architecture and Functioning of the Digital Registration Platform", TEAL),
    ("2", "Workflow Automation and Backend Integration with Revenue & Land Records Systems", ACCENT),
    ("3", "Cybersecurity Measures and Data Protection Protocols", GREEN),
    ("4", "Capacity Building and Change Management for Field-level Functionaries", PURPLE),
]
for i, (n, t, c) in enumerate(items):
    top = Inches(1.5) + i * Inches(1.35)
    circ = rect(s, Inches(0.9), top, Inches(0.95), Inches(0.95), c, MSO_SHAPE.OVAL)
    shape_text(circ, n, 28, bold=True)
    rect(s, Inches(2.1), top, Inches(10.3), Inches(0.95), LIGHT)
    rect(s, Inches(2.1), top, Inches(0.08), Inches(0.95), c)
    text(s, Inches(2.4), top, Inches(9.9), Inches(0.95), t, 20, DARK, True, anchor=MSO_ANCHOR.MIDDLE)

# 3. At a glance
s = new_slide("Kaveri 2.0 at a Glance", "Online-first, end-to-end registration services for citizens and the Department")
text(s, Inches(0.5), Inches(1.25), Inches(12.3), Inches(0.9),
     "Kaveri 2.0 is the web-based application of the Department of Stamps & Registration that moves "
     "registration from counter-based data entry to citizen-driven online applications, with the Sub-Registrar "
     "Office (SRO) visit reduced to verification, biometrics and final registration.", 15, DARK)
svc = [
    ("Document Registration", "Sale, gift, release, settlement, mortgage and other deeds under the Registration Act"),
    ("Marriage Registration", "Registration of marriages with online application and SRO verification"),
    ("Encumbrance Certificate", "Form 15 EC generated from Index II data; online EC search (Form 22)"),
    ("Certified Copy", "Certified copies of registered documents for banks, IT dept. and citizens"),
    ("Firm Registration", "Registration of partnership firms"),
    ("Property Valuation", "Guidance-value based market valuation, stamp duty & fee calculation"),
]
colors = [TEAL, ACCENT, GREEN, PURPLE, NAVY, TEAL]
for i, (h_, b_) in enumerate(svc):
    col, row = i % 3, i // 3
    l = Inches(0.5) + col * Inches(4.15)
    t = Inches(2.35) + row * Inches(1.75)
    rect(s, l, t, Inches(3.95), Inches(1.6), LIGHT)
    rect(s, l, t, Inches(3.95), Inches(0.08), colors[i])
    text(s, l + Inches(0.15), t + Inches(0.15), Inches(3.7), Inches(0.4), h_, 16, colors[i], True)
    text(s, l + Inches(0.15), t + Inches(0.6), Inches(3.7), Inches(0.9), b_, 12, DARK)
text(s, Inches(0.5), Inches(5.95), Inches(12.3), Inches(0.9),
     ["Two user classes: Guest users (stamp duty, guidance value, EC search) and Registered users (OTP-verified "
      "account for filing applications).",
      "Department users: Sub-Registrar (SR), First/Second Division Assistants (FDA/SDA), Data Entry Operators (DEO), "
      "DR/DIG and head-office administrators."], 12, MUTED, bullet=True, space=3)

# ---- Section 1 ----
section(1, "Architecture and Functioning of the Digital Registration Platform",
        ["Layered, microservices-based design  \u2022  API Gateway  \u2022  Containerised deployment  \u2022  "
         "Functional modules"], TEAL)

# 5. Layered architecture
s = new_slide("Solution Architecture", "Layered microservices architecture (per Kaveri 2.0 Software Design Document v2.2)")
layers = [
    ("Users", "Citizens / Advocates / Document writers  \u2022  SRO staff (SR, FDA, SDA, DEO)  \u2022  DR / DIG / IGR", MUTED),
    ("Presentation Layer", "Angular (TypeScript) single-page application  \u2022  Citizen portal & Departmental portal", TEAL),
    ("API Gateway", "Single entry point  \u2022  Reverse proxy / routing  \u2022  AuthN/AuthZ  \u2022  SSL termination  "
                    "\u2022  Rate limiting  \u2022  Logging & correlation  \u2022  IP allow-listing", ACCENT),
    ("Service & Business Layer", "ASP.NET Core (C#) microservices per business capability  \u2022  Business workflow, "
                                 "components & entities  \u2022  Service agents for external systems", NAVY),
    ("Data Layer", "PostgreSQL (Npgsql, Dapper)  \u2022  Index II / EC data  \u2022  Object storage for scanned "
                   "deeds, photos & biometrics", GREEN),
]
for i, (name, desc, c) in enumerate(layers):
    t = Inches(1.3) + i * Inches(1.15)
    sh = rect(s, Inches(0.5), t, Inches(2.6), Inches(1.0), c)
    shape_text(sh, name, 15, bold=True)
    b = rect(s, Inches(3.2), t, Inches(7.1), Inches(1.0), LIGHT)
    shape_text(b, desc, 13, DARK, align=PP_ALIGN.LEFT)
side = rect(s, Inches(10.45), Inches(2.45), Inches(2.4), Inches(4.45), PURPLE)
shape_text(side, ["Cross-cutting", "", "Centralised logging", "Exception management",
                  "Authentication & authorisation", "Caching", "Monitoring"], 12, bold=True)
sb = rect(s, Inches(10.45), Inches(1.3), Inches(2.4), Inches(1.0), ACCENT)
shape_text(sb, ["External systems", "Bhoomi, e-Swathu, e-Aasthi, Khajane-II, eKYC"], 11, bold=True)

# 6. Technology stack
s = new_slide("Technology Stack & Deployment", "Open, container-ready stack designed for independent scaling of services")
card(s, Inches(0.5), Inches(1.3), Inches(4.0), Inches(2.6), "Front End",
     ["Angular component-based SPA", "Modules, components, services, dependency injection",
      "Separate citizen and departmental interfaces", "Responsive web access for citizens"], TEAL)
card(s, Inches(4.65), Inches(1.3), Inches(4.0), Inches(2.6), "Back End",
     ["ASP.NET Core Web APIs in C#", "Microservices per bounded context", "API Gateway pattern (BFF)",
      "Synchronous HTTPS + asynchronous messaging"], NAVY)
card(s, Inches(8.8), Inches(1.3), Inches(4.0), Inches(2.6), "Data & Storage",
     ["PostgreSQL with Npgsql provider", "Dapper for data access & transactions",
      "Object storage (Scality) for images & scans", "Separate reporting / search database"], GREEN)
card(s, Inches(0.5), Inches(4.1), Inches(6.1), Inches(2.6), "Deployment Model",
     ["Services packaged as Docker container images", "Independent deployment & rollback per service",
      "Horizontal scale-out of high-load services (search, valuation)",
      "UAT and Production environments with versioned releases"], ACCENT)
card(s, Inches(6.75), Inches(4.1), Inches(6.05), Inches(2.6), "Design Principles",
     ["Separation of concerns, single responsibility", "Don't repeat yourself; composition over inheritance",
      "Fault isolation between services", "Automated QA: unit tests, static code analysis"], PURPLE)

# 7. Functional modules
s = new_slide("Functional Modules of Kaveri 2.0", "Citizen-facing and departmental modules that together make up the registration lifecycle")
mods_c = ["User Registration & Login (OTP)", "Document Info & Property Search", "Property Schedule",
          "Market Valuation & Fee Calculation", "Party Information", "Document for Approval",
          "Payment (Khajane-II / e-Stamp)", "Schedule Appointment", "EC Search & Certified Copy"]
mods_d = ["Departmental Login (Biometric)", "Application Allocation (FIFO) to FDA/SDA",
          "SR Review & Remarks", "Allocation to DEO", "Party Verification: eKYC, Photo, Biometric",
          "Document Scanning & Upload", "SR Digital Signature", "EC (Form 15) & Acknowledgement",
          "Court Order / Litigation Entry"]
for col, (hdr, mods, c) in enumerate([("Citizen Portal", mods_c, TEAL), ("Departmental Portal (SRO)", mods_d, NAVY)]):
    l = Inches(0.5) + col * Inches(6.25)
    hb = rect(s, l, Inches(1.3), Inches(6.05), Inches(0.55), c)
    shape_text(hb, hdr, 16, bold=True)
    for i, m in enumerate(mods):
        r_, c_ = i // 3, i % 3
        bx = rect(s, l + c_ * Inches(2.03), Inches(2.0) + r_ * Inches(1.0), Inches(1.95), Inches(0.9), LIGHT, line=c)
        shape_text(bx, m, 12, DARK, bold=True)
text(s, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.5),
     ["Property types supported: Agricultural (Bhoomi), Non-agricultural rural (e-Swathu), urban (e-Aasthi), "
      "BBMP, BDA, UPOR and 'Others'.",
      "Multi-property and mixed transactions (e.g., Bhoomi + e-Aasthi) are handled within one application.",
      "Rule-driven article selection determines stamp duty, registration fee, cess, surcharge, scanning and mutation fee."],
     13, DARK, bullet=True, space=5)

# ---- Section 2 ----
section(2, "Workflow Automation and Backend Integration with Revenue & Land Records Systems",
        ["End-to-end registration workflow  \u2022  Bhoomi / e-Swathu / e-Aasthi / UPOR  \u2022  Khajane-II  "
         "\u2022  eKYC  \u2022  Mutation"], ACCENT)

# 9. End-to-end workflow
s = new_slide("End-to-End Document Registration Workflow", "From online application to registered document and Encumbrance Certificate")
steps = [
    ("Online Application", "Citizen logs in, selects article & document type"),
    ("Property Search", "Auto-fetch from Bhoomi / e-Swathu / e-Aasthi; restriction & litigation check"),
    ("Valuation & Fees", "Guidance value, stamp duty, registration fee auto-computed"),
    ("SRO Pre-Scrutiny", "SR allocates to FDA/SDA (FIFO); SR approves or returns with remarks"),
    ("Payment", "Khajane-II challan / SHCIL e-Stamp verified online"),
    ("Book Slot", "Appointment booked at SRO; token generated"),
    ("Registration at SRO", "DEO: eKYC, photo, biometrics; scan & upload"),
    ("Sign & Issue", "SR digitally signs; Index II, EC Form 15, acknowledgement"),
]
flow(s, Inches(1.5), steps, [TEAL, NAVY, ACCENT, GREEN, PURPLE, TEAL, NAVY, ACCENT])
rect(s, Inches(0.5), Inches(4.6), Inches(12.3), Inches(1.9), LIGHT)
text(s, Inches(0.75), Inches(4.8), Inches(11.9), Inches(1.6),
     ["Pre-registration scrutiny happens online, so the SRO visit is limited to identity verification, admission of "
      "execution and signing.",
      "The applicant sees only the SR's decision and remarks; FDA/SDA remarks are internal. Returned applications reopen "
      "with previously entered data pre-filled.",
      "One hour of daily slots is reserved at the SR's discretion for emergency registrations (e.g., wills)."],
     13, DARK, bullet=True, space=6)

# 10. Integrations
s = new_slide("Backend Integration with Revenue, Land Records & Payment Systems",
              "Kaveri 2.0 validates data at source instead of relying on manual entry")
integ = [
    ("Bhoomi (Revenue Dept.)", ["Agricultural land search by district/taluk/hobli/village/survey no.",
                                "Owner, extent, hissa and Pyki / 11E sketch check",
                                "Restriction flag blocks registration", "Mutation data flows back to land records"], TEAL),
    ("e-Swathu / e-Aasthi", ["e-Swathu: Gram Panchayat properties", "e-Aasthi: urban local body properties",
                             "Search by location & PID; property PDF shown", "BBMP, BDA, UPOR and 'Others' handled separately"], GREEN),
    ("Khajane-II & SHCIL", ["Challan number, amount and date verified online", "Deficit amount prompts an additional challan",
                            "e-Stamp (SHCIL) accepted for stamp duty", "Payment unlocks appointment booking"], ACCENT),
    ("Identity & Notifications", ["Aadhaar-based eKYC for presenter, executants, claimants",
                                  "Photo and fingerprint biometrics captured at SRO",
                                  "Mobile / email OTP for citizen accounts", "SMS / email alerts at each workflow stage"], PURPLE),
]
for i, (h_, lines, c) in enumerate(integ):
    col, row = i % 2, i // 2
    card(s, Inches(0.5) + col * Inches(6.2), Inches(1.3) + row * Inches(2.8), Inches(6.0), Inches(2.6), h_, lines, c, 13)

# 11. Automation benefits
s = new_slide("What the Workflow Automates", "Rules and checks that were earlier manual are now enforced by the system")
auto = [
    ("Property validation", "Ownership, extent and restrictions validated against the source land record", TEAL),
    ("Litigation / court-order check", "Applications halt if an attachment or court order is recorded against the property", ACCENT),
    ("Fee computation", "Duty = Govt. duty + surcharge + cess + scanning fee + mutation fee \u2212 adjustments, plus registration fee", GREEN),
    ("Work allocation", "FIFO-based allocation of applications to FDA/SDA and slot-based allocation to DEOs", NAVY),
    ("Payment reconciliation", "Online challan verification removes manual cross-checking of receipts", PURPLE),
    ("Record generation", "Index II update, EC (Form 15), endorsements and acknowledgement produced automatically", TEAL),
]
for i, (h_, b_, c) in enumerate(auto):
    col, row = i % 2, i // 2
    l = Inches(0.5) + col * Inches(6.2)
    t = Inches(1.3) + row * Inches(1.85)
    circ = rect(s, l, t + Inches(0.3), Inches(0.9), Inches(0.9), c, MSO_SHAPE.OVAL)
    shape_text(circ, str(i + 1), 24, bold=True)
    rect(s, l + Inches(1.05), t, Inches(4.95), Inches(1.6), LIGHT)
    text(s, l + Inches(1.2), t + Inches(0.12), Inches(4.7), Inches(0.4), h_, 16, c, True)
    text(s, l + Inches(1.2), t + Inches(0.55), Inches(4.7), Inches(1.0), b_, 13, DARK)

# ---- Section 3 ----
section(3, "Cybersecurity Measures and Data Protection Protocols",
        ["Identity & access  \u2022  Application & network security  \u2022  Data protection  \u2022  Threat model"], GREEN)

# 13. Security controls
s = new_slide("Security Controls Across the Platform", "Defence in depth: identity, application, transport and data layers")
card(s, Inches(0.5), Inches(1.3), Inches(4.0), Inches(2.7), "Identity & Access",
     ["Citizen accounts verified by mobile & email OTP", "Departmental login: credentials + biometric",
      "Role-based access (SR, FDA, SDA, DEO, DR, admin)", "Office-level data segregation per SRO"], GREEN)
card(s, Inches(4.65), Inches(1.3), Inches(4.0), Inches(2.7), "Application & Gateway",
     ["Only the API Gateway is exposed; services stay internal", "Centralised AuthN/AuthZ at the gateway",
      "Rate limiting, throttling, IP allow-listing", "Input validation and exception handling in every layer"], TEAL)
card(s, Inches(8.8), Inches(1.3), Inches(4.0), Inches(2.7), "Transport & Integrity",
     ["HTTPS / TLS for all client traffic (SSL termination)", "TLS-capable PostgreSQL connections",
      "SR digitally signs scanned deeds & annexures", "eKYC + biometrics establish party identity"], NAVY)
card(s, Inches(0.5), Inches(4.2), Inches(12.3), Inches(2.5), "Operational Security",
     ["Centralised, correlated logging across layers for audit trails and incident investigation",
      "Controlled release management: UAT sign-off and versioned production releases with release notes",
      "Security vulnerabilities, OS and hardware patches tracked as part of runtime-environment hardening",
      "Separation of duties: FDA/SDA scrutiny, DEO data capture and SR approval are distinct roles in the workflow"],
     PURPLE, 13)

# 14. Data protection & threat model
s = new_slide("Data Protection & Threat Model", "Protecting a statewide legal record of property ownership")
text(s, Inches(0.5), Inches(1.2), Inches(6.0), Inches(0.4), "Data assets inherited from Kaveri 1.0 (June 2022)", 16, NAVY, True)
data = [("SR/DR registration database", "2.5 TB"), ("Centralised EC data", "2 TB"),
        ("Reporting / search database", "1 TB"), ("Scanned images & documents (SAN)", "50 TB"),
        ("Photos, thumbprints, EC/CC copies, Form 22", "File server")]
for i, (k, v) in enumerate(data):
    t = Inches(1.7) + i * Inches(0.62)
    rect(s, Inches(0.5), t, Inches(4.4), Inches(0.55), LIGHT)
    text(s, Inches(0.65), t, Inches(4.2), Inches(0.55), k, 13, DARK, anchor=MSO_ANCHOR.MIDDLE)
    vb = rect(s, Inches(4.95), t, Inches(1.4), Inches(0.55), TEAL)
    shape_text(vb, v, 13, bold=True)
text(s, Inches(0.5), Inches(4.95), Inches(5.9), Inches(1.8),
     ["Personal data (Aadhaar-linked eKYC, photos, biometrics) handled only within the departmental workflow",
      "Registered documents are permanent legal records: integrity, backup and long-term retention are mandatory"],
     12, DARK, bullet=True, space=4)
text(s, Inches(6.8), Inches(1.2), Inches(6.0), Inches(0.4), "Threat categories addressed in the design", 16, NAVY, True)
threats = [
    ("Runtime environment", "OS & hardware vulnerabilities, insecure drivers, untested code, endpoint compromise", TEAL),
    ("Communication protocol", "DDoS, DNS & routing attacks, open ports/IPs, ransomware, SSL/TCP weaknesses", ACCENT),
    ("Cryptographic", "Weak keys, flawed key generation, private-key security, algorithm vulnerabilities", GREEN),
    ("NIST risk-based security", "Categorise, value assets, identify threats, treat risk, authorise & monitor", PURPLE),
]
for i, (h_, b_, c) in enumerate(threats):
    t = Inches(1.7) + i * Inches(1.25)
    rect(s, Inches(6.8), t, Inches(6.0), Inches(1.1), LIGHT)
    rect(s, Inches(6.8), t, Inches(0.08), Inches(1.1), c)
    text(s, Inches(7.0), t + Inches(0.05), Inches(5.7), Inches(0.4), h_, 14, c, True)
    text(s, Inches(7.0), t + Inches(0.45), Inches(5.7), Inches(0.65), b_, 12, DARK)

# ---- Section 4 ----
section(4, "Capacity Building and Change Management Strategies for Field-level Functionaries",
        ["Stakeholder roles  \u2022  Training approach  \u2022  Rollout & hand-holding  \u2022  Sustaining adoption"], PURPLE)

# 16. Roles & training
s = new_slide("Who Needs to Be Trained, and on What", "Role-based capacity building aligned to the Kaveri 2.0 workflow")
roles = [
    ("Sub-Registrar (SR)", ["Biometric login, application allocation", "Review & remarks, approval / rejection",
                            "Digital signing, court-order entry", "Emergency slot management"], NAVY),
    ("FDA / SDA", ["FIFO work queue handling", "Scrutiny of deed, parties, schedule, valuation",
                   "Verifying document summary", "Internal remarks to SR"], TEAL),
    ("Data Entry Operator", ["Party appearance & eKYC", "Photo and fingerprint capture standards",
                             "Scanning, re-scanning & upload", "Offline payment entry"], ACCENT),
    ("DR / DIG / Admin", ["Monitoring dashboards & reports", "User & office management",
                          "Escalations and exceptions", "Audit and compliance review"], GREEN),
    ("Citizens & Document Writers", ["Account creation & online application", "Property search & fee payment",
                                     "Appointment booking", "Handling SR remarks"], PURPLE),
]
cw = Inches(2.38)
for i, (h_, lines, c) in enumerate(roles):
    l = Inches(0.5) + i * (cw + Inches(0.1))
    hb = rect(s, l, Inches(1.3), cw, Inches(0.7), c)
    shape_text(hb, h_, 14, bold=True)
    rect(s, l, Inches(2.0), cw, Inches(2.6), LIGHT)
    text(s, l + Inches(0.1), Inches(2.1), cw - Inches(0.2), Inches(2.5), lines, 12, DARK, bullet=True, space=5)
text(s, Inches(0.5), Inches(4.8), Inches(12.3), Inches(0.4), "Training approach", 16, NAVY, True)
text(s, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.6),
     ["Train-the-trainer model: master trainers per district cascade training to every SRO",
      "Hands-on sessions in a UAT / sandbox environment using real-life deed scenarios",
      "Role-wise SOPs, quick-reference guides, video tutorials and Kannada-language material",
      "Refresher training delivered alongside each release using release notes"], 13, DARK, bullet=True, space=4)

# 17. Change management
s = new_slide("Change Management Strategy", "Moving field offices from counter-based data entry to an online-first workflow")
phases = [
    ("Awareness", "Communicate why: transparency, fewer visits, faster registration"),
    ("Pilot", "Go-live in select SROs; capture issues early"),
    ("Hand-holding", "On-site support teams & helpdesk during rollout"),
    ("Statewide Rollout", "Phased extension district by district"),
    ("Reinforce", "Monitor KPIs, recognise champions, refine SOPs"),
]
flow(s, Inches(1.45), phases, [TEAL, NAVY, ACCENT, GREEN, PURPLE], h=Inches(1.1))
card(s, Inches(0.5), Inches(4.2), Inches(6.05), Inches(2.5), "Addressing Resistance",
     ["Involve SRs and staff in UAT and feedback loops", "Designate SRO-level Kaveri champions",
      "Clear escalation path via helpdesk and release fixes", "Show time savings from online pre-scrutiny"], ACCENT)
card(s, Inches(6.75), Inches(4.2), Inches(6.05), Inches(2.5), "Measuring Adoption",
     ["Share of applications filed online by citizens", "Turnaround time: submission to registration",
      "Applications returned with SR remarks", "Helpdesk tickets and recurring issues per release"], GREEN)

# 18. Key takeaways
s = new_slide("Key Takeaways", "Kaveri 2.0 in summary")
kt = [
    ("Modern architecture", "Layered microservices behind an API Gateway, Angular front end, PostgreSQL and containerised services", TEAL),
    ("Connected government", "Property, payment and identity data validated at source through Bhoomi, e-Swathu, e-Aasthi, Khajane-II and eKYC", ACCENT),
    ("Secure by design", "OTP, biometrics, role-based access, digital signatures, TLS and centralised audit logging", GREEN),
    ("People-centred adoption", "Role-based training, phased rollout and continuous feedback drive field-level adoption", PURPLE),
]
for i, (h_, b_, c) in enumerate(kt):
    t = Inches(1.35) + i * Inches(1.35)
    sh = rect(s, Inches(0.5), t, Inches(3.3), Inches(1.15), c)
    shape_text(sh, h_, 17, bold=True)
    b = rect(s, Inches(3.9), t, Inches(8.9), Inches(1.15), LIGHT)
    shape_text(b, b_, 15, DARK, align=PP_ALIGN.LEFT)

# 19. Thank you
s = new_slide()
rect(s, 0, 0, W, H, NAVY)
rect(s, 0, Inches(4.3), W, Inches(0.08), ACCENT)
text(s, Inches(0.8), Inches(2.3), Inches(11.5), Inches(1.2), "Thank You", 54, WHITE, True, PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(3.4), Inches(11.5), Inches(0.7), "Questions & Discussion", 24, SUB, align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(4.7), Inches(11.5), Inches(0.6),
     "Department of Stamps & Registration, Government of Karnataka", 16, WHITE, align=PP_ALIGN.CENTER)

total = len(_pages)
for idx, sl in enumerate(_pages, 1):
    if idx in (1, total):
        continue
    rect(sl, 0, H - Inches(0.35), W, Inches(0.35), NAVY)
    text(sl, Inches(0.4), H - Inches(0.34), Inches(10), Inches(0.3), FOOTER, 10, WHITE, anchor=MSO_ANCHOR.MIDDLE)
    text(sl, W - Inches(1.3), H - Inches(0.34), Inches(0.9), Inches(0.3), f"{idx}/{total}", 10, WHITE,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print(f"Saved {OUT} ({total} slides)")
