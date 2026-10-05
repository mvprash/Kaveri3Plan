import docx
d=docx.Document(r'Finalized BRD/Marriage/RFP/BRD_Marriage_BRD_v8.docx')
for i,t in enumerate(d.tables):
    rows=t.rows
    print(i, len(rows),'x',len(t.columns), 'style=',t.style.name if t.style else None)
    for r in rows[:2]:
        print('   ', ' || '.join(c.text.replace('\n',' / ')[:70] for c in r.cells))
s=d.sections[0]
print('page',s.page_width,s.page_height,s.left_margin,s.orientation)
for p in s.header.paragraphs: print('HDR',p.text)
for p in s.footer.paragraphs: print('FTR',p.text)
